import json

from openpyxl import Workbook

from tools import build_vector_scalar_gap_coverage as coverage_builder
from tools import export_switch_matrix_xls as export
from tools import vector_scalar_gap_reporting as reporting


def test_norm_value_decodes_json_encoded_categorical_values():
    assert reporting.norm_value('"BREAKOUT"') == "breakout"
    assert reporting.norm_value("true") == "true"
    assert reporting.norm_value("1.50") == "1.5"
    assert export.norm_val('"30"') == "30"
    assert export.norm_val('"false"') == "false"


def test_normalizer_accepts_entry_and_exit_proxy_metadata_shapes():
    entry = reporting._normalise_detail_row(
        {
            "vector_evidence_class": "VEC_APPROX",
            "approximation": {
                "vector_target_fields": ["ENTRY_SCORE"],
                "approximation_confidence": "LOW",
            },
        }
    )
    exit_row = reporting._normalise_detail_row(
        {
            "vector_evidence_class": "VEC_APPROX",
            "proxy_kind": "BOOLEAN_ENABLE",
            "confidence": "LOW",
            "mismatch_class": "BOOLEAN_FAMILY_EVENT_PROXY",
            "formula": "direct boolean proxy",
        }
    )

    assert entry["proxy_field"] == "ENTRY_SCORE"
    assert exit_row["proxy_field"] == "BOOLEAN_ENABLE"
    assert exit_row["approximation_confidence"] == "LOW"
    assert (
        exit_row["approximation_mismatch_class"]
        == "BOOLEAN_FAMILY_EVENT_PROXY"
    )


def test_normalizer_isolates_reviewed_scalar_gap_adapter_for_amber_display():
    row = reporting._normalise_detail_row(
        {
            "campaign": "stocks_repaired_20260730_c5_1yr_vec_gap",
            "tier": "VEC_DIAGNOSTIC",
            "exact_completion_credit": False,
            "promotion_allowed": False,
            "vector_adapter": "BREAKOUT_SIZE_SMA200_T2_MULT",
            "value_json": "1.5",
        }
    )

    assert row["source_tier"] == "VEC_DIAGNOSTIC"
    assert row["tier"] == "VEC_APPROX"
    assert row["vector_evidence_class"] == "VEC_APPROX"
    assert row["proxy_field"] == "BREAKOUT_SIZE_SMA200_T2_MULT"
    assert row["proxy_value"] == "1.5"
    assert row["required_next_stage"] == "EXACT_V8"
    assert row["engine_ranking_allowed"] is False
    assert row["db_engine_write_allowed"] is False
    assert row["live_config_write_allowed"] is False
    assert reporting._detail_contract_ok(row) is True


def _write_fixture(root, *, mismatch=False, malformed_approx=False):
    reports = root / "data" / "reports"
    detail = reports / "vectorized_scalar_gap"
    detail.mkdir(parents=True)
    generated = "2026-07-31T02:00:00Z"
    coverage = {
        "schema_version": 1,
        "generated_at": generated,
        "tier": "VEC_DIAGNOSTIC",
        "display_contract": "SEPARATE_AMBER_ITALIC",
        "exact_completion_credit": False,
        "engine_ranking_allowed": False,
        "promotion_allowed": False,
        "approved_vector_adapters": ["APPROVED_ALIAS"],
        "screened_cells": 4 if mismatch else 3,
        "per_key": {
            "MU_LONG": {
                "screened_cells": 4 if mismatch else 3,
                "moved": 2,
                "inert": 1,
                "zero_trade": 0,
                "native_vector_cells": 1,
                "sampled_exact_parity_adapter_cells": 1,
                "vec_native_cells": 1,
                "vec_parity_cells": 1,
                "vec_approx_cells": 1,
                "vec_approx_moved": 1,
                "exact_replay_priority_cells": 2,
                "approximation_confidence": {"MEDIUM": 1},
            }
        },
    }
    (reports / "VECTOR_SCALAR_GAP_COVERAGE_20260731.json").write_text(
        json.dumps(coverage)
    )
    (reports / "VECTOR_ADAPTER_PARITY_AUDIT_20260731.json").write_text(
        json.dumps(
            {
                "approved_params": ["APPROVED_ALIAS"],
                "rejected_params": ["RETIRED_ALIAS"],
                "per_param": {"RETIRED_ALIAS": {"compared": 15}},
            }
        )
    )
    common = {
        "key": "MU_LONG",
        "tier": "VEC_DIAGNOSTIC",
        "exact_completion_credit": False,
        "promotion_allowed": False,
        "generated_at": "2026-07-31T01:00:00Z",
        "trades": 4,
    }
    rows = [
        {
            **common,
            "param": "NATIVE",
            "value_json": "true",
            "status": "MOVED",
            "delta_gain_mo_vs_bh_diagnostic": 1.0,
        },
        {
            **common,
            "param": "APPROVED_ALIAS",
            "value_json": "1.5",
            "vector_adapter": "VEC_ALIAS",
            "status": "INERT",
        },
        {
            **common,
            "param": "APPROX_SIZE",
            "value_json": "200",
            "status": "MOVED",
            "tier": "VEC_APPROX",
            "vector_evidence_class": "VEC_APPROX",
            "approximation_confidence": (
                "UNKNOWN" if malformed_approx else "MEDIUM"
            ),
            "approximation_mismatch_class": (
                "PATH_SPECIFIC_SIZING_COLLAPSED_TO_SIDE_SIZE_MULT"
            ),
            "source_group": "SIZING",
            "action_group": "SIZING",
            "source_rank": 0.5,
            "proxy_field": "LONG_SIZE_MULT",
            "proxy_value": 1.25,
            "required_next_stage": "EXACT_V8",
            "engine_ranking_allowed": False,
            "live_config_write_allowed": False,
            "db_engine_write_allowed": False,
            "result_uniqueness_status": "PASS",
            "delta_gain_mo_vs_bh_diagnostic": 0.4,
        },
        {
            **common,
            "param": "RETIRED_ALIAS",
            "value_json": "2",
            "vector_adapter": "OLD_ALIAS",
            "status": "MOVED",
        },
        {
            **common,
            "param": "PARITY_ONLY",
            "value_json": "3",
            "status": "MOVED",
            "protected_exact_replay_diagnostic": True,
        },
        {
            **common,
            "param": "EXACT_ALREADY_OWNS_CELL",
            "value_json": '"BREAKOUT"',
            "status": "MOVED",
            "protected_exact_present": True,
        },
    ]
    (detail / "MU_LONG.jsonl").write_text(
        "".join(json.dumps(row) + "\n" for row in rows)
    )
    (detail / "MU_LONG_summary.json").write_text(
        json.dumps(
            {
                "eligible_vector_cells": 331,
                "protected_exact_cells": 2,
            }
        )
    )


def test_loader_keeps_native_parity_and_contract_valid_approx(tmp_path):
    _write_fixture(tmp_path)

    payload = reporting.load(tmp_path)

    assert payload["available"] is True
    assert payload["screened_cells"] == 3
    assert {row["param"] for row in payload["rows"]} == {
        "NATIVE",
        "APPROVED_ALIAS",
        "APPROX_SIZE",
    }
    per_key = payload["per_key"]["MU_LONG"]
    assert per_key["vec_native_cells"] == 1
    assert per_key["vec_parity_cells"] == 1
    assert per_key["vec_approx_cells"] == 1
    assert per_key["approximation_confidence"] == {"MEDIUM": 1}
    assert per_key["approximation_mismatch_classes"] == {
        "PATH_SPECIFIC_SIZING_COLLAPSED_TO_SIDE_SIZE_MULT": 1
    }
    assert payload["excluded_detail_rows"]["parity_rejected_adapter"] == 1
    assert payload["excluded_detail_rows"]["protected_parity_only"] == 1
    assert payload["excluded_detail_rows"]["protected_exact_present"] == 1
    assert all(row["exact_completion_credit"] is False for row in payload["rows"])
    assert all(row["engine_ranking_allowed"] is False for row in payload["rows"])


def test_loader_fails_closed_when_receipt_count_disagrees(tmp_path):
    _write_fixture(tmp_path, mismatch=True)

    payload = reporting.load(tmp_path)

    assert payload["available"] is False
    assert "detail/coverage mismatch" in payload["reason"]


def test_loader_fails_closed_on_malformed_approx_authority(tmp_path):
    _write_fixture(tmp_path, malformed_approx=True)

    payload = reporting.load(tmp_path)

    assert payload["available"] is False
    assert "detail/coverage mismatch" in payload["reason"]


def test_mapping_explains_vector_boundary_without_exact_credit(tmp_path):
    _write_fixture(tmp_path)
    payload = reporting.load(tmp_path)

    mapping = reporting.scalar_grid_mapping(tmp_path, payload)

    row = mapping["per_key"]["MU_LONG"]
    assert row["eligible_vector_cells"] == 331
    assert row["pre_parity_vector_candidate_cells"] == 346
    assert row["parity_rejected_adapter_cells"] == 15
    assert row["protected_exact_cells_at_screen_time"] == 2
    assert mapping["exact_completion_credit"] is False
    assert mapping["engine_ranking_allowed"] is False


def test_xls_sheet_is_separate_amber_italic_and_nonpromotable(tmp_path):
    _write_fixture(tmp_path)
    payload = reporting.load(tmp_path)
    wb = Workbook()

    ws = export.write_vector_scalar_diagnostics_sheet(wb, payload)

    assert ws.title == "VEC Scalar Diagnostics"
    assert ws["A3"].value == "MU_LONG"
    assert ws["U3"].value == "VEC_DIAGNOSTIC"
    assert ws["V3"].value is False
    assert ws["W3"].value is False
    assert ws["X3"].value is False
    assert ws["Y3"].value is False
    assert ws["Z3"].value is False
    assert ws["A3"].font.italic is True
    assert ws["A3"].font.color.rgb.endswith("9A6B00")
    assert ws["A3"].fill.fgColor.rgb.endswith("FFF2CC")
    approx_row = next(
        row for row in ws.iter_rows(min_row=3)
        if row[2].value == "APPROX_SIZE"
    )
    assert approx_row[9].value == "VEC_APPROX"
    assert approx_row[13].value == "MEDIUM"
    assert approx_row[15].value == 0.5
    assert approx_row[16].value == "LONG_SIZE_MULT"
    assert approx_row[18].value is True


def test_main_xls_approx_overlay_fills_only_blank_cells(tmp_path):
    _write_fixture(tmp_path)
    payload = reporting.load(tmp_path)
    item = next(
        row for row in payload["rows"]
        if row["vector_evidence_class"] == "VEC_APPROX"
    )
    wb = Workbook()
    ws = wb.active
    blank = ws["A1"]
    exact = ws["A2"]
    exact.value = 9.5

    assert export.apply_vector_approx_overlay(blank, item) is True
    assert export.apply_vector_approx_overlay(exact, item) is False

    assert blank.value == 0.4
    assert blank.font.italic is True
    assert blank.font.color.rgb.endswith("9A6B00")
    assert blank.fill.fgColor.rgb.endswith("FFF2CC")
    assert "underlying ENGINE cell is blank" in blank.comment.text
    assert "DB ENGINE write allowed: false" in blank.comment.text
    assert exact.value == 9.5
    assert exact.comment is None


def test_email_html_states_zero_exact_credit(tmp_path):
    _write_fixture(tmp_path)

    section = reporting.html_section(tmp_path)

    assert "Exact completion credit: none" in section
    assert "ENGINE/database writes, live config writes" in section
    assert "ranking, and promotion are forbidden" in section
    assert "1</td><td>1</td>" in section


def test_merged_coverage_discovers_all_approximate_lane_directories():
    names = {str(path) for path in coverage_builder.DEFAULT_APPROX_DIRS}
    assert "data/reports/vector_approx_scalar" in names
    assert "data/reports/vectorized_exit_reentry_approx" in names


def test_builder_discovers_new_priority_keys_from_approx_ledgers(tmp_path):
    reports = tmp_path / "data/reports"
    exit_dir = reports / "vectorized_exit_reentry_approx"
    exit_dir.mkdir(parents=True)
    receipt = reports / "VECTOR_SCALAR_GAP_COVERAGE_20260731.json"
    receipt.write_text(
        json.dumps({"approved_vector_adapters": [], "per_key": {}})
    )
    row = {
        "key": "IBIT_LONG",
        "tier": "VEC_APPROX",
        "vector_evidence_class": "VEC_APPROX",
        "status": "MOVED",
        "param": "EXIT_PROXY",
        "value_json": "true",
        "source_rank": 0.5,
        "proxy_field": "BOOLEAN_ENABLE",
        "proxy_value": True,
        "approximation_confidence": "LOW",
        "approximation_mismatch_class": "BOOLEAN_FAMILY_EVENT_PROXY",
        "required_next_stage": "EXACT_V8",
        "exact_completion_credit": False,
        "engine_ranking_allowed": False,
        "db_engine_write_allowed": False,
        "live_config_write_allowed": False,
        "promotion_allowed": False,
    }
    (exit_dir / "IBIT_LONG.jsonl").write_text(json.dumps(row) + "\n")

    payload = coverage_builder.build(tmp_path, receipt)

    assert payload["per_key"]["IBIT_LONG"]["vec_approx_cells"] == 1
