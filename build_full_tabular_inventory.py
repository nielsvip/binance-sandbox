from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import build_path_explanations as base


ROOT = Path('/Users/niels/Documents/binance')
INPUT = ROOT / 'outputs/019fcd2f-b9de-78d3-a50a-398223e8822c/SWITCH_MATRIX_TRB_S1_PATH_EXPLANATIONS.xlsx'
OUTPUT = ROOT / 'outputs/019fcd2f-b9de-78d3-a50a-398223e8822c/SWITCH_MATRIX_TRB_S1_COMPLETE_PATH_INVENTORY.xlsx'


def files():
    out = []
    for p in ROOT.rglob('*'):
        if not p.is_file() or 'outputs' in p.parts or '.git' in p.parts:
            continue
        # The 504 MB bloated/corrupt/vector-filled artifacts are preserved as source
        # files but are not re-read for analysis; they contain duplicated worksheet XML.
        if any(tag in p.name.lower() for tag in ('bloated', 'corrupt', 'vector_filled')):
            continue
        if p.suffix.lower() in {'.xlsx', '.xlsm', '.xls', '.csv', '.tsv'}:
            out.append(p)
    return sorted(out)


def text_row(values):
    return ' | '.join(str(v).strip() for v in values if v is not None and str(v).strip())


def key_candidates(text):
    # Config-style keys are the strongest path identifiers; retain explicit reason/path labels too.
    keys = re.findall(r'\b[A-Z][A-Z0-9_]{4,}\b', text)
    for k in re.findall(r'(?i)(?:entry|exit|reentry|open|close)[A-Za-z0-9_./:-]{2,}', text):
        if k.upper() not in keys:
            keys.append(k)
    return keys[:12]


def category(text):
    low = text.lower()
    if any(x in low for x in ('exit', 'close', 'stop', 'take_profit', 'reduce')):
        return 'Exit'
    if any(x in low for x in ('entry', 'open', 'reentry', 'augment', 'force_open')):
        return 'Entry'
    return 'Other'


def relevant(text, keys):
    if keys:
        return True
    low = text.lower()
    return any(x in low for x in ('entry', 'exit', 'reentry', 'open', 'close', 'path', 'switch', 'reason', 'function'))


def scan():
    source_rows = []
    coverage = []
    for path in files():
        rel = str(path.relative_to(ROOT))
        if path.suffix.lower() in {'.csv', '.tsv'}:
            total = 0
            relevant_count = 0
            try:
                with path.open('r', encoding='utf-8-sig', errors='replace', newline='') as f:
                    reader = csv.reader(f, delimiter='\t' if path.suffix.lower() == '.tsv' else ',')
                    for row_num, row in enumerate(reader, 1):
                        total += 1
                        text = text_row(row)
                        keys = key_candidates(text)
                        if relevant(text, keys):
                            relevant_count += 1
                            source_rows.append((rel, 'CSV', row_num, category(text), keys, text[:1600]))
            except Exception as e:
                coverage.append((rel, 'CSV', 0, 0, 0, f'read error: {e}'))
                continue
            coverage.append((rel, 'CSV', total, 1, relevant_count, ''))
        else:
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    total = 0
                    relevant_count = 0
                    for row_num, row in enumerate(ws.iter_rows(values_only=True), 1):
                        vals = list(row)
                        if not any(v is not None and str(v).strip() for v in vals):
                            continue
                        total += 1
                        text = text_row(vals)
                        keys = key_candidates(text)
                        if relevant(text, keys):
                            relevant_count += 1
                            source_rows.append((rel, ws.title, row_num, category(text), keys, text[:1600]))
                    coverage.append((rel, ws.title, total, ws.max_column or 0, relevant_count, ''))
            except Exception as e:
                coverage.append((rel, '<workbook>', 0, 0, 0, f'read error: {e}'))
    return coverage, source_rows


def style_sheet(ws, widths):
    fill = PatternFill('solid', fgColor='1F4E78')
    white = Font(color='FFFFFF', bold=True)
    thin = Side(style='thin', color='D9E2F3')
    for cell in ws[1]:
        cell.fill = fill
        cell.font = white
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=thin)
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False


def main():
    wb = load_workbook(INPUT)
    for name in ('Source Coverage', 'All Path Evidence'):
        if name in wb.sheetnames:
            del wb[name]
    coverage, source_rows = scan()
    code_index = base.build_code_index()

    cov = wb.create_sheet('Source Coverage')
    cov.append(['Source file', 'Sheet / type', 'Non-empty rows', 'Columns', 'Relevant path rows', 'Read note'])
    for row in coverage:
        cov.append(row)
    style_sheet(cov, [58, 32, 16, 12, 20, 50])

    ev = wb.create_sheet('All Path Evidence')
    ev.append(['Source file', 'Sheet / type', 'Source row', 'Area', 'Extracted path keys', 'Engine function / file:line', 'Source row contents'])
    for rel, sheet, row_num, area, keys, text in source_rows:
        traces = []
        for key in keys:
            trace, _ = base.trace_for(key, code_index)
            if trace not in traces:
                traces.append(trace)
        ev.append([rel, sheet, row_num, area, ', '.join(keys), ' || '.join(traces[:6]), text])
    style_sheet(ev, [58, 32, 12, 10, 48, 72, 150])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f'Wrote {OUTPUT}; coverage rows={len(coverage)}, evidence rows={len(source_rows)}')


if __name__ == '__main__':
    main()
