import json

from tools.export_switch_matrix_xls import (
    apply_preserved_v8_overlay,
    apply_vector_approx_overlay,
    audit_post_precedence_vector_overlays,
    audit_vector_overlay_uniqueness,
    load_active_vector_cell_audit_overlay,
)
from tools.audit_stock_matrix_workbooks import audit_vector_display_uniqueness


def row(metric, fingerprint, *, status="MOVED", trades=3):
    return {
        "status": status,
        "delta_gain_mo_vs_bh_diagnostic": metric,
        "action_fingerprint": fingerprint,
        "trades": trades,
    }


def test_duplicate_numbers_are_never_matrix_fills():
    source = {
        ("ENTRY_A", "true", "MU_LONG"): row(4.25, "fp-a"),
        ("EXIT_B", "false", "MU_LONG"): row(4.25, "fp-b"),
        ("ENTRY_A", "true", "NVDA_LONG"): row(4.25, "fp-c"),
    }
    audited, report = audit_vector_overlay_uniqueness(source)
    assert audited[("ENTRY_A", "true", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"
    assert audited[("EXIT_B", "false", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"
    # Uniqueness is per symbol/side; another key has a different B&H contract.
    assert audited[("ENTRY_A", "true", "NVDA_LONG")]["result_uniqueness_status"] == "PASS"
    assert report["passed_numeric_overlays"] == 1
    assert report["reason_counts"]["DUPLICATE_RESULT_NUMBER"] == 2


def test_duplicate_action_fingerprint_is_quarantined_even_when_metrics_differ():
    source = {
        ("ENTRY_A", "true", "MU_LONG"): row(3.1, "same-actions"),
        ("ENTRY_A", "false", "MU_LONG"): row(3.2, "same-actions"),
    }
    audited, report = audit_vector_overlay_uniqueness(source)
    assert all(item["result_uniqueness_status"] == "QUARANTINED" for item in audited.values())
    assert report["reason_counts"]["DUPLICATE_ACTION_FINGERPRINT"] == 2


def test_zero_inert_and_zero_close_rows_fail_closed():
    source = {
        ("A", "1", "MU_LONG"): row(0.0, "fp-a"),
        ("B", "1", "MU_LONG"): row(1.0, "fp-b", status="INERT"),
        ("C", "1", "MU_LONG"): row(2.0, "fp-c", trades=0),
        ("D", "1", "MU_LONG"): row(3.0, "fp-d", trades=2),
    }
    audited, report = audit_vector_overlay_uniqueness(source)
    assert audited[("A", "1", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"
    assert audited[("B", "1", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"
    assert audited[("C", "1", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"
    assert audited[("D", "1", "MU_LONG")]["result_uniqueness_status"] == "PASS"
    assert report["passed_numeric_overlays"] == 1


def test_preserved_v8_precedes_vector_and_collision_stays_visible_red():
    from openpyxl import Workbook

    cell = Workbook().active["A1"]
    preserved = {
        "delta_gain_mo_vs_bh": 12.5,
        "gain_per_mo": 18.0,
        "campaign": "stocks_repaired_20260725_c2",
        "ts": "2026-07-25T00:00:00Z",
        "source_file": "param_matrix_daemon/result.json",
        "preservation_display_status": "QUARANTINED",
        "preservation_quarantine_reasons": ["DUPLICATE_RESULT_NUMBER"],
        "result_collision_count": 2,
    }
    vector = {
        "vector_evidence_class": "VEC_APPROX",
        "delta_gain_mo_vs_bh_diagnostic": 99.0,
        "status": "MOVED",
        "result_uniqueness_status": "PASS",
        "exact_completion_credit": False,
        "engine_ranking_allowed": False,
        "promotion_allowed": False,
        "live_config_write_allowed": False,
        "db_engine_write_allowed": False,
    }
    # Collided historical evidence stays in Evidence Provenance and cannot
    # occupy a physical numeric result cell.  A lower behavior-unique vector
    # diagnostic may compete for the still-blank surface.
    assert apply_preserved_v8_overlay(cell, preserved) is False
    assert cell.value is None
    assert apply_vector_approx_overlay(cell, vector) is True
    assert cell.value == 99.0


def test_current_engine_cell_cannot_be_overwritten_by_preserved_v8():
    from openpyxl import Workbook

    cell = Workbook().active["A1"]
    cell.value = 7.0
    assert apply_preserved_v8_overlay(cell, {"delta_gain_mo_vs_bh": 12.5}) is False
    assert cell.value == 7.0


def test_overlay_index_builder_stops_campaign_on_broadcast_results(tmp_path, monkeypatch):
    import json
    import sys

    from tools import build_full_matrix_vector_overlay_index as builder

    rows = [
        {
            "param": "A", "value_json": "true", "key": "MU_LONG",
            "status": "MOVED", "delta_gain_mo_vs_bh_diagnostic": 4.2,
            "action_fingerprint": "actions-a", "trades": 3,
        },
        {
            "param": "B", "value_json": "false", "key": "MU_LONG",
            "status": "MOVED", "delta_gain_mo_vs_bh_diagnostic": 4.2,
            "action_fingerprint": "actions-b", "trades": 3,
        },
    ]
    (tmp_path / "all_cells.jsonl").write_text(
        "".join(json.dumps(row) + "\n" for row in rows)
    )
    monkeypatch.setattr(sys, "argv", ["builder", "--out-dir", str(tmp_path)])
    assert builder.main() == 2
    receipt = json.loads((tmp_path / "overlay_index_receipt.json").read_text())
    assert receipt["campaign_must_stop"] is True
    assert receipt["uniqueness_contract"]["quarantined_numeric_overlays"] == 2


def test_canonical_export_source_excludes_old_full_blank_broadcast():
    import inspect
    from tools import export_switch_matrix_xls as exporter

    source = inspect.getsource(exporter.main)
    assert "vector_approx_index = {}" in source
    assert "vector_approx_index = load_full_matrix_vector_overlay" not in source


def test_historical_export_is_rejected_before_surface_lock():
    import inspect
    from tools import export_switch_matrix_xls as exporter

    source = inspect.getsource(exporter.main)
    reject = source.index("FROZEN_HISTORICAL_EXPORT_BLOCKED")
    lock = source.index("acquire_surface_lock()")
    assert reject < lock
    assert "ALLOW_FROZEN_HISTORICAL_MATRIX_EXPORT" in source


def test_strict_overlay_loader_selects_largest_verified_union_not_newest_mtime(tmp_path):
    import hashlib
    import os

    reports = tmp_path / "data" / "reports"
    reports.mkdir(parents=True)

    def detail(key, param, value, metric, fingerprint):
        return {
            "key": key,
            "param": param,
            "value_json": value,
            "status": "MOVED",
            "delta_gain_mo_vs_bh_diagnostic": metric,
            "action_fingerprint": fingerprint,
            "trades": 3,
            "campaign": "stocks_repaired_20260730_c5_1yr_vec_gap",
            "tier": "VEC_DIAGNOSTIC",
            "vector_adapter": "test_adapter",
            "exact_completion_credit": False,
            "promotion_allowed": False,
        }

    def write_receipt(name, rows):
        ledgers = []
        hashes = {}
        for index, item in enumerate(rows):
            rel = f"data/reports/{name}_{index}.jsonl"
            path = tmp_path / rel
            path.write_text(json.dumps(item) + "\n")
            ledgers.append(rel)
            hashes[rel] = hashlib.sha256(path.read_bytes()).hexdigest()
        receipt = {
            "contract": "STRICT_PER_KEY_UNIQUE_METRIC_AND_ACTION_NO_DUPLICATES",
            "exact_completion_credit": False,
            "promotion_allowed": False,
            "source_ledgers": ledgers,
            "source_ledger_sha256": hashes,
            "raw_rows_read": len(rows),
            "passed": len(rows),
            "quarantined": 0,
        }
        path = reports / f"TRB_ACTIVE_VECTOR_CELL_AUDIT_{name}.json"
        path.write_text(json.dumps(receipt))
        return path

    larger = write_receipt(
        "LARGER",
        [
            detail("MU_LONG", "P1", True, 1.0, "fp1"),
            detail("NVDA_LONG", "P2", False, 2.0, "fp2"),
        ],
    )
    smaller = write_receipt(
        "SMALLER", [detail("MU_LONG", "P1", True, 1.0, "fp1")]
    )
    newer = larger.stat().st_mtime + 60
    os.utime(smaller, (newer, newer))

    overlay, audit = load_active_vector_cell_audit_overlay(tmp_path)
    assert audit["selection_rule"] == "LARGEST_HASH_VERIFIED_SOURCE_UNION"
    assert audit["report"].endswith("_LARGER.json")
    assert audit["source_ledgers"] == 2
    assert audit["passed"] == 2
    assert len(overlay) == 2


def test_one_isolated_two_value_plateau_is_the_only_accepted_duplicate():
    source = {
        ("THRESHOLD", "0", "MU_LONG"): row(1.25, "same-actions"),
        ("THRESHOLD", "20", "MU_LONG"): row(1.25, "same-actions"),
        ("THRESHOLD", "40", "MU_LONG"): row(2.0, "actions-40"),
        ("THRESHOLD", "60", "MU_LONG"): row(3.0, "actions-60"),
        ("THRESHOLD", "80", "MU_LONG"): row(4.0, "actions-80"),
    }
    audited, report = audit_vector_overlay_uniqueness(source)
    assert audited[("THRESHOLD", "0", "MU_LONG")]["result_uniqueness_status"] == "PASS"
    assert audited[("THRESHOLD", "20", "MU_LONG")]["result_uniqueness_status"] == "PASS"
    assert all(item["result_uniqueness_status"] == "PASS" for item in audited.values())
    assert report["allowed_isolated_metric_plateau_groups"] == 1
    assert report["allowed_isolated_fingerprint_plateau_groups"] == 1


def test_two_value_plateau_without_three_other_unique_values_is_rejected():
    source = {
        ("THRESHOLD", "0", "MU_LONG"): row(1.25, "same-actions"),
        ("THRESHOLD", "20", "MU_LONG"): row(1.25, "same-actions"),
        ("THRESHOLD", "40", "MU_LONG"): row(2.0, "actions-40"),
        ("THRESHOLD", "60", "MU_LONG"): row(3.0, "actions-60"),
    }
    audited, _report = audit_vector_overlay_uniqueness(source)
    assert audited[("THRESHOLD", "0", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"
    assert audited[("THRESHOLD", "20", "MU_LONG")]["result_uniqueness_status"] == "QUARANTINED"


def test_post_precedence_reaudit_quarantines_a_partially_hidden_plateau():
    source = {
        ("THRESHOLD", str(value), "MU_LONG"): row(
            1.25 if value in (0, 20) else float(value),
            "same-actions" if value in (0, 20) else f"actions-{value}",
        )
        for value in (0, 20, 40, 60, 80)
    }
    audited, _ = audit_vector_overlay_uniqueness(source)
    matrix = [["THRESHOLD", "", str(value)] for value in (0, 20, 40, 60, 80)]
    states = [[None], [None], ["green"], ["green"], ["green"]]
    overlays = [[audited[("THRESHOLD", str(value), "MU_LONG")]] for value in (0, 20, 40, 60, 80)]
    preserved = [[None] for _ in matrix]

    report = audit_post_precedence_vector_overlays(
        matrix, states, overlays, preserved, ["MU_LONG"]
    )

    assert overlays[0][0]["result_uniqueness_status"] == "QUARANTINED"
    assert overlays[1][0]["result_uniqueness_status"] == "QUARANTINED"
    assert report["quarantined_numeric_overlays"] == 2


def test_serialized_workbook_rejects_broadcast_vector_numbers(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Entry"
    ws.append(["main_switch", "MU_LONG"])
    for name in ("A", "B"):
        ws.append([name, 4.25])
        ws.cell(ws.max_row, 2).comment = Comment(
            "VEC_APPROX diagnostic overlay — the underlying ENGINE cell is blank.\n"
            "status: MOVED\nresult uniqueness: PASS\naction fingerprint: fp-1",
            "test",
        )
    wb.save(path)
    assert any(
        "VECTOR_DISPLAY_DUPLICATE_RESULT:MU_LONG:4.25" in error
        for error in audit_vector_display_uniqueness(path)
    )


def test_serialized_workbook_accepts_unique_moved_vector_numbers(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Entry"
    ws.append(["main_switch", "MU_LONG"])
    for index, value in enumerate((4.25, 5.5), 1):
        ws.append([f"P{index}", value])
        ws.cell(ws.max_row, 2).comment = Comment(
            "VEC_APPROX diagnostic overlay — the underlying ENGINE cell is blank.\n"
            f"status: MOVED\nresult uniqueness: PASS\naction fingerprint: fp-{index}",
            "test",
        )
    wb.save(path)
    assert audit_vector_display_uniqueness(path) == []


def test_serialized_workbook_accepts_only_one_bounded_pair_plateau(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Entry"
    ws.append(["main_switch", "sub_setting", "value", "description", "status", "MU_LONG"])
    for index, (axis_value, result, fingerprint) in enumerate(
        ((0, 1.25, "same"), (20, 1.25, "same"), (40, 2.0, "fp40"),
         (60, 3.0, "fp60"), (80, 4.0, "fp80")),
        2,
    ):
        ws.append(["THRESHOLD", "", axis_value, "test", "OK", result])
        ws.cell(index, 6).comment = Comment(
            "VEC_APPROX diagnostic overlay — the underlying ENGINE cell is blank.\n"
            "status: MOVED\nresult uniqueness: PASS\n"
            f"action fingerprint: {fingerprint}",
            "test",
        )
    wb.save(path)
    assert audit_vector_display_uniqueness(path) == []


def test_serialized_workbook_forward_fills_grouped_parameter_axis(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sizing"
    ws.append(["main_switch", "sub_setting", "value", "description", "status", "MU_LONG"])
    for index, (axis_value, result, fingerprint) in enumerate(
        ((0.75, 1.25, "same"), (1.875, 2.0, "fp1875"),
         (2.25, 3.0, "fp225"), (1.5, 4.0, "fp15"),
         (1.125, 1.25, "same")),
        2,
    ):
        ws.append([
            "BREAKOUT_SIZE_EMA200_T2_PCT" if index == 2 else None,
            None,
            axis_value,
            "test",
            "OK",
            result,
        ])
        ws.cell(index, 6).comment = Comment(
            "VEC_APPROX diagnostic overlay — the underlying ENGINE cell is blank.\n"
            "status: MOVED\nresult uniqueness: PASS\n"
            f"action fingerprint: {fingerprint}",
            "test",
        )
    wb.save(path)
    assert audit_vector_display_uniqueness(path) == []


def test_serialized_workbook_uses_strict_receipt_when_exact_cells_hide_axis(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sizing"
    ws.append(["main_switch", "sub_setting", "value", "description", "status", "MU_LONG"])
    for index, axis_value in enumerate((0.75, 1.875, 2.25, 1.5, 1.125), 2):
        # The middle three cells represent exact-owned values and deliberately
        # carry no VEC comment; only the bounded vector pair is displayed.
        result = 1.25 if axis_value in (0.75, 1.125) else index * 10.0
        ws.append([
            "BREAKOUT_SIZE_EMA200_T2_PCT" if index == 2 else None,
            None,
            axis_value,
            "test",
            "OK",
            result,
        ])
        if axis_value in (0.75, 1.125):
            ws.cell(index, 6).comment = Comment(
                "VEC_APPROX diagnostic overlay — the underlying ENGINE cell is blank.\n"
                "status: MOVED\nresult uniqueness: PASS\n"
                "action fingerprint: same",
                "test",
            )
    wb.save(path)
    receipt_path = tmp_path / "VECTOR_OVERLAY_UNIQUENESS_AUDIT.json"
    receipt = {
        "old_full_blank_broadcast_display_excluded": True,
        "strict_active_vector_audit": {"available": True, "passed": 5},
        "matrix_numeric_fill_requires_pass": True,
        "engine_cells_modified": False,
        "allowed_plateau_rule": (
            "ONE_TWO_VALUE_PAIR_PER_5PLUS_VALUE_PARAMETER_AXIS; "
            "ALL_OTHER_VALUES_UNIQUE; NO_CROSS_PARAMETER_COLLISION"
        ),
        "allowed_plateaus": [{
            "key": "MU_LONG",
            "param": "BREAKOUT_SIZE_EMA200_T2_PCT",
            "values": ["0.75", "1.125"],
            "metric": 1.25,
            "behavior_fingerprint": "same",
            "measured_axis_values": 5,
        }],
    }
    receipt_path.write_text(json.dumps(receipt))
    path.touch()
    assert audit_vector_display_uniqueness(path) == []

    # The receipt permits only the exact source-audited tuple. Any broader or
    # stale authority must fail closed.
    for field, bad_value in (
        ("values", ["0.75", "1.5"]),
        ("metric", 1.2501),
        ("behavior_fingerprint", "tampered"),
    ):
        tampered = json.loads(json.dumps(receipt))
        tampered["allowed_plateaus"][0][field] = bad_value
        receipt_path.write_text(json.dumps(tampered))
        path.touch()
        assert audit_vector_display_uniqueness(path)

    receipt_path.write_text("{}")
    path.touch()
    assert audit_vector_display_uniqueness(path)

    import os
    import time
    receipt_path.write_text(json.dumps(receipt))
    old = time.time() - 1200
    os.utime(receipt_path, (old, old))
    path.touch()
    assert audit_vector_display_uniqueness(path)
