import json
import sqlite3

from openpyxl import Workbook, load_workbook

from tools import param_results_store as prs


def _memory_store():
    con = sqlite3.connect(":memory:")
    con.executescript(prs.SCHEMA)
    for migration in prs.MIGRATIONS:
        try:
            con.execute(migration)
        except sqlite3.OperationalError:
            pass
    return con


def _insert(con, table, values):
    cols = list(values)
    con.execute(
        f"INSERT INTO {table} ({','.join(cols)}) VALUES ({','.join('?' for _ in cols)})",
        [values[col] for col in cols],
    )


def test_restored_param_baseline_sheets_are_contract_tier_isolated(tmp_path, monkeypatch):
    (tmp_path / "data").mkdir()
    (tmp_path / "reports").mkdir()
    (tmp_path / "symbols_trb_long.json").write_text('["MU"]')
    (tmp_path / "symbols_trb_short.json").write_text('["HAO"]')
    registry = {
        "tradier": {
            "WT_DC_ENTRY_ENABLED": {
                "family": "WT_DC_ENTRY", "group": "ENTRY", "role": "MAIN_SWITCH",
                "per_sym": True,
            },
            "WT_DC_ENTRY_THRESHOLD": {
                "family": "WT_DC_ENTRY", "group": "ENTRY", "role": "THRESHOLD",
                "per_sym": True,
            },
            "WT_DC_EXIT_ENABLED": {
                "family": "WT_DC_EXIT", "group": "EXIT", "role": "MAIN_SWITCH",
                "per_sym": True,
            },
        }
    }
    manifest = {
        "params": {
            "WT_DC_ENTRY_ENABLED": {
                "default": True, "test_values": [False, True], "sweepable": True,
            },
            "WT_DC_ENTRY_THRESHOLD": {
                "default": 45, "test_values": [35, 45, 55], "sweepable": True,
            },
            "WT_DC_EXIT_ENABLED": {
                "default": True, "test_values": [False, True], "sweepable": True,
            },
        }
    }
    (tmp_path / "data" / "knob_registry.json").write_text(json.dumps(registry))
    (tmp_path / "data" / "param_sweep_manifest_tradier.json").write_text(json.dumps(manifest))

    con = _memory_store()
    common = {
        "mode": "tradier", "symbol": "MU", "side": "LONG",
        "campaign": prs.REPAIRED_CAMPAIGN, "ts": "2026-07-26T07:00:00Z",
        "tier": "ENGINE", "validation_status": "PASS",
        "contract_fingerprint": "contract-mu", "real_closes": 2,
        "reentry_violations": 0, "requested_fill_ratio": 1.0,
        "size_clamp_count": 0,
    }
    _insert(con, "key_baseline", {
        **common, "gain_per_mo": 6.0, "bh_per_mo": 4.0,
        "delta_gain_mo_vs_bh": 2.0, "trades": 8, "pool_sharpe": 0.8,
        "time_in_mkt_pct": 60.0, "capture_vs_bh": 1.5,
    })
    _insert(con, "param_cells", {
        **common, "param": "WT_DC_ENTRY_THRESHOLD", "value_json": "45",
        "gain_per_mo": 9.0, "delta_gain_mo_vs_bh": 5.0, "trades": 9,
        "pool_sharpe": 1.0, "time_in_mkt_pct": 55.0, "inert": 0,
    })
    # Tempting but invalid evidence must never displace the repaired cell.
    _insert(con, "param_cells", {
        **{**common, "campaign": "stocks_baseline_v2_s4h", "tier": "VEC",
           "validation_status": None, "contract_fingerprint": None},
        "param": "WT_DC_ENTRY_THRESHOLD", "value_json": "999",
        "gain_per_mo": 999.0, "delta_gain_mo_vs_bh": 995.0, "trades": 999,
        "pool_sharpe": 9.0, "time_in_mkt_pct": 99.0, "inert": 0,
    })
    con.commit()

    wb = Workbook()
    wb.active.title = "Baselines"
    coverage = prs._write_restored_workbook_sheets(wb, con, "tradier", tmp_path)
    output = tmp_path / "PARAM_BASELINE_STOCKS.xlsx"
    wb.save(output)
    check = load_workbook(output, data_only=True)

    assert {"Workbook Guide", "PerSym Results", "Entry Paths", "Exit Paths"} <= set(check.sheetnames)
    assert coverage == {
        "keys": 2, "repaired_baselines": 1, "repaired_cells": 1,
        "entry_paths": 1, "exit_paths": 1,
    }
    per_sym = check["PerSym Results"]
    assert per_sym["A2"].value == "MU_LONG"
    assert per_sym["V2"].value == "ENTRY:WT_DC_ENTRY"
    assert per_sym["W2"].value == "WT_DC_ENTRY_THRESHOLD=45"
    assert per_sym["Y2"].value == 5.0
    assert "PROMOTABLE" in per_sym["Z2"].value
    assert "999" not in str(per_sym["W2"].value)

    entry = check["Entry Paths"]
    assert entry["F1"].value == "WT_DC_ENTRY"
    assert "ENTRY MAIN SWITCH" in entry["F2"].value
    assert "SETTINGS:" in entry["F2"].value
    assert "default=45" in entry["F2"].value
    assert "grid=35,45,55" in entry["F2"].value
    assert entry["F3"].value == "tested settings"
    assert "WT_DC_ENTRY_THRESHOLD=45" in entry["F4"].value
    assert entry["G4"].value == "WT_DC_ENTRY_THRESHOLD=45"
    assert entry["H4"].value == 5.0
    assert entry["J4"].value == "PROMOTABLE > B&H"
    assert check["Exit Paths"]["F4"].value is None
    assert check["Exit Paths"]["J4"].value == "PENDING"

    # The real recurring exporter starts from a fresh Workbook on every refresh.
    # Prove the restored sheets are regenerated (and therefore cannot disappear)
    # on two consecutive full exports, not merely added by the helper once.
    monkeypatch.setattr(prs, "BASE", tmp_path)
    recurring = tmp_path / "PARAM_BASELINE_RECURRING.xlsx"
    seed = Workbook()
    seed.active.title = "Matrix_Top"
    seed.active.append(["key", "proof"])
    seed.active.append(["MU_LONG", "preserve-me"])
    seed.save(recurring)
    prs.export_xlsx(con, recurring, "tradier", None)
    prs.export_xlsx(con, recurring, "tradier", None)
    refreshed = load_workbook(recurring, read_only=True, data_only=True)
    assert refreshed.sheetnames[:4] == [
        "Workbook Guide", "PerSym Results", "Entry Paths", "Exit Paths",
    ]
    assert refreshed["Matrix_Top"]["B2"].value == "preserve-me"
