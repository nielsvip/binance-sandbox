import openpyxl


def _rows(tpl, tab):
    wb = openpyxl.load_workbook("SPREADSHEETS/%s" % tpl, read_only=True, data_only=False)
    ws = wb[tab]
    out = []
    for r in range(3, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and a.strip():
            out.append((r, a.strip(), ws.cell(row=r, column=2).value, ws.cell(row=r, column=12).value))
    return out


def test_eod_lookback_default_30_everywhere():
    from v12_quick_engine import QuickConfig
    from config_tradier import TradierConfig
    assert int(QuickConfig().GAP_PER_SYMBOL_LOOKBACK_DAYS) == 30
    assert int(TradierConfig.GAP_PER_SYMBOL_LOOKBACK_DAYS) == 30


def test_eod_loop_knob_defaults():
    from v12_quick_engine import QuickConfig
    c = QuickConfig()
    assert int(c.GAP_MOC_WINDOW_MINUTES) == 90
    assert bool(c.GAP_MOC_REQUIRE_TOP) is True
    assert bool(c.GAP_MORNING_REENTRY_ENABLED) is True
    assert int(c.GAP_MORNING_REENTRY_MINUTES_AFTER_OPEN) == 120
    assert float(c.GAP_MOC_REENTRY_SIZE_MULT) == 1.25


def _eod_counts(**kw):
    import numpy as np
    from v12_quick_engine import QuickConfig, simulate_one
    d = dict(np.load("backtest_v8/indicators/AAPL.npz", allow_pickle=True))
    cfg = QuickConfig()
    cfg.MODE = "tradier"
    for k, v in kw.items():
        setattr(cfg, k, v)
    led = simulate_one(d, "AAPL", True, cfg).get("ledger", []) or []
    n_exit = sum(1 for e in led if e.get("type") == "CLOSE" and "GAP_MOC" in str(e.get("reason", "")))
    n_rebuy = sum(1 for e in led if e.get("type") == "OPEN" and "GAP_MORNING" in str(e.get("reason", "")))
    return n_exit, n_rebuy


def test_eod_gap_exit_and_rebuy_fire():
    n_exit, n_rebuy = _eod_counts()
    assert n_exit > 0, "gap exit never fires at defaults"
    assert n_rebuy > 0, "morning rebuy never fires at defaults"


def test_eod_masters_disable():
    assert _eod_counts(GAP_MOC_EXIT_ENABLED=False) == (0, 0)
    n_exit, n_rebuy = _eod_counts(GAP_MORNING_REENTRY_ENABLED=False)
    assert n_exit > 0 and n_rebuy == 0


def test_eod_template_yes_rows():
    for tpl in ("TEMPLATE_STOCKS_LONG.xlsx", "TEMPLATE_STOCKS_SHORT.xlsx"):
        lb = [(b, L) for (_, a, b, L) in _rows(tpl, "GLOBAL_RISK_GATES") if a == "GAP_PER_SYMBOL_LOOKBACK_DAYS"]
        assert ("30" in [str(b) for b, _ in lb]) or (30 in [b for b, _ in lb])
        yes = [b for b, L in lb if str(L).upper() == "YES"]
        assert len(yes) == 1 and int(float(str(yes[0]))) == 30, (tpl, lb)
        rw = [(a, b, L) for (_, a, b, L) in _rows(tpl, "REENTRY_WINDOWED")]
        ex = [(a, b, L) for (_, a, b, L) in _rows(tpl, "EXIT_STRUCTURAL")]
        sz = [(a, b, L) for (_, a, b, L) in _rows(tpl, "STDEV_SLOPE_SIZING")]
        assert ("GAP_MORNING_REENTRY_ENABLED", "True", "YES") in rw, tpl
        assert ("GAP_MORNING_REENTRY_MINUTES_AFTER_OPEN", 120, "YES") in rw, tpl
        assert ("GAP_MOC_WINDOW_MINUTES", 90, "YES") in ex, tpl
        assert ("GAP_MOC_REQUIRE_TOP", "True", "YES") in ex, tpl
        assert ("GAP_MOC_REENTRY_SIZE_MULT", "1.25", "YES") in sz, tpl
