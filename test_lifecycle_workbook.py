import hashlib
import json

from openpyxl import Workbook

from tools import current_matrix_reporting as cmr
from tools.lifecycle_workbook import (
    COLUMNS, assert_lifecycle_rows, install_lifecycle_sheet, write_lifecycle_sheet,
)


def _hash_json(value):
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _receipt(root, *, local_files=True):
    artifact = root / "data/reports/vec_research/xyz_lifecycle_test"
    artifact.mkdir(parents=True)
    npz = artifact / "XYZ.npz"
    raw = artifact / "raw_results.jsonl"
    if local_files:
        npz.write_bytes(b"causal")
        raw.write_text('{"row":1}\n')
    npz_sha = hashlib.sha256(b"causal").hexdigest()
    raw_sha = hashlib.sha256(b'{"row":1}\n').hexdigest()
    manifest = {
        "schema": "TRB_LIFECYCLE_COMBO_BEAM_V1", "symbol": "XYZ", "side": "SHORT",
        "source_npz": {"path": str(npz), "sha256": npz_sha},
        "source_code": {"path": "/remote/run.py", "sha256": "a" * 64},
        "window": {"start": "2026-01-01", "end_exclusive": "2026-02-01"},
        "commission_bps_one_way": 0, "slippage_bps_one_way": 5,
    }
    manifest["manifest_sha256"] = _hash_json(manifest)
    best = {
        "symbol": "XYZ", "side": "SHORT", "tier": "VECTOR_LIFECYCLE",
        "manifest_sha256": manifest["manifest_sha256"], "gain_pct_per_month": 8,
        "bh_gain_pct_per_month": 2, "delta_gain_mo_vs_bh": 6,
        "no_lookahead_future_htf_count": 0, "recipe_id": "recipe-1",
        "recipe": {"exit": 1}, "paths": {"exit": {"label": "WT"}},
    }
    best["result_sha256"] = _hash_json(best)
    result = {
        "status": "COMPLETE", "symbol": "XYZ", "side": "SHORT",
        "manifest_sha256": manifest["manifest_sha256"], "best": best,
        "raw_results_sha256": raw_sha, "rows_evaluated": 10,
    }
    result["receipt_sha256"] = _hash_json(result)
    (artifact / "campaign_manifest.json").write_text(json.dumps(manifest))
    (artifact / "result.json").write_text(json.dumps(result))
    return artifact


def test_local_verified_lifecycle_is_written_to_separate_non_exact_sheet(tmp_path):
    _receipt(tmp_path, local_files=True)
    rows = cmr.lifecycle_workbook_rows(tmp_path)
    assert rows[0]["evidence_tier"] == "VECTOR_LIFECYCLE_HASH_VERIFIED"
    wb = Workbook()
    wb.active.title = "Matrix"
    wb["Matrix"]["A1"] = 123.456
    assert write_lifecycle_sheet(wb, tmp_path, rows) == 1
    assert wb["Matrix"]["A1"].value == 123.456
    ws = wb["Vector Lifecycle"]
    assert ws.cell(3, COLUMNS.index("result_sha256") + 1).value == rows[0]["result_sha256"]
    assert ws.cell(3, COLUMNS.index("exact_completion_credit") + 1).value is False
    assert ws.cell(3, COLUMNS.index("scalar_matrix_written") + 1).value is False


def test_remote_receipt_is_retained_unverified_but_tampering_is_rejected(tmp_path):
    artifact = _receipt(tmp_path, local_files=False)
    rows = cmr.lifecycle_workbook_rows(tmp_path)
    assert rows[0]["evidence_tier"] == "VECTOR_LIFECYCLE_REMOTE_RECEIPT_UNVERIFIED"
    result = json.loads((artifact / "result.json").read_text())
    result["best"]["gain_pct_per_month"] = 999
    (artifact / "result.json").write_text(json.dumps(result))
    assert cmr.lifecycle_workbook_rows(tmp_path) == []


def test_duplicate_result_receipt_is_rejected():
    row = {
        "key": "XYZ_SHORT", "evidence_tier": "VECTOR_LIFECYCLE_HASH_VERIFIED",
        "manifest_sha256": "a" * 64, "result_sha256": "b" * 64,
        "receipt_sha256": "c" * 64, "raw_results_sha256": "d" * 64,
        "source_npz": {"sha256": "e" * 64}, "recipe": {"exit": 1},
        "paths": {"exit": {}}, "costs": {"slippage_bps_one_way": 5},
    }
    try:
        assert_lifecycle_rows([row, dict(row)])
    except RuntimeError as exc:
        assert "duplicate lifecycle" in str(exc)
    else:
        raise AssertionError("duplicate result must fail closed")


def test_atomic_install_preserves_every_non_lifecycle_cell(tmp_path):
    _receipt(tmp_path, local_files=True)
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "Matrix"
    wb["Matrix"]["A1"] = 12.5
    wb["Matrix"]["B2"] = "do-not-change"
    wb.save(path)
    receipt = install_lifecycle_sheet(path, tmp_path)
    assert receipt["rows"] == 1
    from openpyxl import load_workbook
    out = load_workbook(path, read_only=True)
    assert out["Matrix"]["A1"].value == 12.5
    assert out["Matrix"]["B2"].value == "do-not-change"
    assert out["Vector Lifecycle"]["A3"].value == "XYZ_SHORT"
    out.close()
