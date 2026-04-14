#!/usr/bin/env python3
"""Auto-update XLSX from sweep results. Run via cron every 5 min."""
import csv, json, time
from pathlib import Path
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, numbers
except ImportError:
    print("openpyxl not installed"); exit(1)

BASE = Path("/Users/niels/Documents/binance")
CSV_FILE = BASE / "sweep_mega_results.csv"
XLSX_FILE = BASE / "sweep_mega_results.xlsx"

if not CSV_FILE.exists():
    print("No results yet"); exit(0)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "TRADIER Sweep"

# Headers
headers = ["Config", "NOLOSS", "Ablation", "Trades", "Realized PnL", "Win Rate %", "IBS Exits", "WT Exits", "Other Exits", "Avg Win %", "Avg Loss %", "Time (s)"]
green = PatternFill(start_color="C6EFCE", fill_color="C6EFCE", fill_type="solid")
red = PatternFill(start_color="FFC7CE", fill_color="FFC7CE", fill_type="solid")
bold = Font(bold=True)

for i, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = bold

with open(CSV_FILE) as f:
    reader = csv.reader(f)
    next(reader)  # skip header
    for row_idx, row in enumerate(reader, 2):
        for col_idx, val in enumerate(row, 1):
            try:
                val = float(val)
            except:
                pass
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            # Color PnL column
            if col_idx == 5 and isinstance(val, (int, float)):
                c.fill = green if val > 0 else red

# Auto-width
for col in ws.columns:
    max_len = max(len(str(c.value or "")) for c in col)
    ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

wb.save(XLSX_FILE)
rows = ws.max_row - 1
print(f"XLSX updated: {XLSX_FILE} ({rows} configs)")
