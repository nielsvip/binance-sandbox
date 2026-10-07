#!/usr/bin/env python3
"""v15_yellow_paint_from_proposal — paint the filter cells of ANY TEMPLATE_*.xlsx from the discovery proposal (Agent Y, 2026-10-01).

Reads switch name / option from col A / B (row 3+), filter header 'FILTER=opt' from row 2 (columns after POS_SYM) — NO tab or row numbers, so a repaired / grown
template works unchanged. Evidence = data/yellow_discovery/<date>/yellow_proposal_<cat>.json (tools/v15_yellow_discovery_proposal.py).
  BRIGHT FFFF00  any evaluated non-zero effect for (SWITCH=cand, FILTER=opt) [Stage B cell, or Stage-A effect inherited by a non-binding evaluated row, or 365D audit]
  none           evaluated at least once and always exactly zero / inert (an existing yellow/light fill on such a cell is REMOVED)
  LIGHT  FFF2CC  UNKNOWN: never evaluated (filter not wired in the engine used / row or pair never evaluated) — NEVER 0, never a result colour
Only the fill of filter cells that are currently yellow / light yellow / empty is changed; other fills, all values, fonts, rows, order untouched (checked with
tools/template_row_guard.py: every row fingerprint over all non-filter columns must be identical before/after).
DRY-RUN by default. --out-dir DIR writes painted COPIES (no freeze check). --apply paints the live template in place (backup to backups/ first) and is refused
while SPREADSHEETS/TEMPLATES_FROZEN exists unless TEMPLATES_UNFREEZE=1.
  python tools/v15_yellow_paint_from_proposal.py --date 20261001 [--templates SPREADSHEETS] [--out-dir DIR | --apply]
"""
import argparse
import datetime
import json
import os
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import template_row_guard as G  # noqa: E402

CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
BRIGHT, LIGHT = "FFFFFF00", "FFFFF2CC"
HDR = 2


def hx(v):
    return int(v, 16) if isinstance(v, str) else int(v or 0)


def load_prop(day_dir: Path, cs):
    j = json.loads((day_dir / f"yellow_proposal_{cs}.json").read_text())
    filt = {h: {"eff": hx(e["A_effect"]), "valid": hx(e["A_valid"])} for h, e in j["filters"].items()}
    rows = {r: {"eval": hx(m["eval"]), "bind": hx(m["bind"])} for r, m in j["rows"].items()}
    return filt, rows, set(j["B_bright"]), set(j["B_evaluated"])


def decide(row_key, sw, hdr, filt, rows, b_bright, b_eval):
    f = filt.get(hdr)
    k = f"{row_key}\t{hdr}"
    if k in b_bright:
        return "B"
    r = rows.get(row_key)
    nb = (r["eval"] & ~r["bind"]) if r else 0
    if f and (f["eff"] & nb):
        return "B"
    if hdr.split("=", 1)[0].strip() == sw:
        return "N"  # the switch row itself sets this key: not a combination
    if k in b_eval or (f and (f["valid"] & nb)):
        return "N"
    return "L"


def paint(path: Path, cs, day_dir: Path, out_path: Path, write: bool):
    filt, rows, b_bright, b_eval = load_prop(day_dir, cs)
    before = openpyxl.load_workbook(str(path))
    wb = openpyxl.load_workbook(str(path))
    stats = {"bright": 0, "light": 0, "none": 0, "changed": 0, "other_fill_kept": 0}
    fcols = {}
    for tab in SWITCH_SHEETS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        pos = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(row=HDR, column=c).value or "").strip().upper() == "POS_SYM"), 14)
        hdrs = {c: str(ws.cell(row=HDR, column=c).value).strip() for c in range(pos + 1, ws.max_column + 1) if isinstance(ws.cell(row=HDR, column=c).value, str) and "=" in ws.cell(row=HDR, column=c).value}
        fcols[tab] = set(hdrs)
        for r in range(HDR + 1, ws.max_row + 1):
            a, b = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
            if a in (None, "") or b in (None, ""):
                continue
            sw = str(a).strip()
            rk = f"{sw}={str(b).strip()}"
            for c, h in hdrs.items():
                d = decide(rk, sw, h, filt, rows, b_bright, b_eval)
                want = BRIGHT if d == "B" else LIGHT if d == "L" else None
                cell = ws.cell(row=r, column=c)
                cur = str(cell.fill.fgColor.rgb or "").upper() if cell.fill is not None and cell.fill.fill_type == "solid" else None
                stats["bright" if want == BRIGHT else "light" if want == LIGHT else "none"] += 1
                if cur not in (None, BRIGHT, LIGHT):
                    stats["other_fill_kept"] += 1
                    continue
                if cur == want:
                    continue
                cell.fill = PatternFill(start_color=want, end_color=want, fill_type="solid") if want else PatternFill(fill_type=None)
                stats["changed"] += 1
    for tab, cols in fcols.items():  # row integrity: all NON-filter columns identical (value + fill)
        skip = set(cols)  # cols already holds the filter COLUMN INDICES
        G.assert_rows_intact(before[tab], wb[tab], skip)
        # values of the filter columns must be unchanged too
        for r in range(HDR + 1, wb[tab].max_row + 1):
            for c in skip:
                if before[tab].cell(row=r, column=c).value != wb[tab].cell(row=r, column=c).value:
                    raise AssertionError(f"{tab} r{r} c{c}: value changed")
    if write:
        tmp = out_path.with_suffix(".tmp.xlsx")
        wb.save(str(tmp))
        openpyxl.load_workbook(str(tmp))
        tmp.replace(out_path)
    return stats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default="20261001")
    ap.add_argument("--templates", default=str(ROOT / "SPREADSHEETS"), help="dir holding TEMPLATE_{CAT}.xlsx")
    ap.add_argument("--out-dir")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--cat-side", action="append", choices=CAT_SIDES)
    a = ap.parse_args()
    if a.apply and (ROOT / "SPREADSHEETS" / "TEMPLATES_FROZEN").exists() and not os.environ.get("TEMPLATES_UNFREEZE"):
        sys.exit("REFUSED: SPREADSHEETS/TEMPLATES_FROZEN exists (set TEMPLATES_UNFREEZE=1 only on the user's explicit approval)")
    day_dir = ROOT / "data" / "yellow_discovery" / a.date
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for cs in (a.cat_side or CAT_SIDES):
        src = Path(a.templates) / f"TEMPLATE_{cs}.xlsx"
        if a.apply:
            shutil.copy2(src, ROOT / "backups" / f"before_yellow_paint_{ts}_TEMPLATE_{cs}.xlsx")
            out, write = src, True
        elif a.out_dir:
            Path(a.out_dir).mkdir(parents=True, exist_ok=True)
            out, write = Path(a.out_dir) / src.name, True
        else:
            out, write = src, False
        st = paint(src, cs, day_dir, out, write)
        print(f"[{cs}] {st} {'-> ' + str(out) if write else '(dry-run)'}")


if __name__ == "__main__":
    main()
