#!/usr/bin/env python3
"""Manifest-driven health audit for repaired exact-matrix workers on S1."""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import sqlite3
import sys


CAMPAIGN = "stocks_repaired_20260725_c2"
EXPECTED_SHEETS = [
    "Engine Coverage",
    "Entry",
    "Exit",
    "Sizing",
    "Other",
    "Ladder",
    "BandLadder",
    "DC4 Diagnostics",
    "Coverage",
    "Exact Engine Evidence",
    "Baselines",
    "Inventory",
]


def _processes() -> dict[int, dict]:
    out = {}
    for proc in Path("/proc").iterdir():
        if not proc.name.isdigit():
            continue
        try:
            argv = [
                item.decode(errors="replace")
                for item in (proc / "cmdline").read_bytes().split(b"\0")
                if item
            ]
            stat = (proc / "stat").read_text().split()
        except OSError:
            continue
        out[int(proc.name)] = {
            "pid": int(proc.name),
            "ppid": int(stat[3]),
            "sid": int(stat[5]),
            "argv": argv,
        }
    return out


def _script_index(argv: list[str], script: str) -> int | None:
    for index, arg in enumerate(argv):
        if Path(arg).name == script:
            return index
    return None


def _arg(argv: list[str], name: str) -> str | None:
    try:
        return argv[argv.index(name) + 1]
    except (ValueError, IndexError):
        return None


def _manifest(sbx: Path) -> tuple[dict[str, tuple[str, str]], dict]:
    path = sbx / "data/matrix_worker_manifest.json"
    if not path.exists():
        raise FileNotFoundError(path)
    payload = json.loads(path.read_text())
    expected = {
        str(row["tag"]): (
            str(row["symbol"]).upper(),
            str(row["side"]).upper(),
        )
        for row in payload.get("workers", [])
    }
    if not expected:
        raise ValueError("matrix worker manifest has no workers")
    return expected, payload


def classify_workers(processes: dict[int, dict], expected: dict):
    """Discover every safe worker; extras remain owners but are flagged."""
    workers = {}
    worker_pids = set()
    extras = []
    failures = []
    for proc in processes.values():
        index = _script_index(proc["argv"], "param_matrix_daemon.py")
        if index is None:
            continue
        argv = proc["argv"][index + 1 :]
        if "--safe-contract" not in argv:
            continue
        tag = _arg(argv, "--tag")
        observed = (_arg(argv, "--only"), _arg(argv, "--side"))
        row = {
            "pid": proc["pid"],
            "ppid": proc["ppid"],
            "sid": proc["sid"],
            "symbol": observed[0],
            "side": observed[1],
            "manifested": tag in expected,
        }
        workers.setdefault(tag or "<missing-tag>", []).append(row)
        worker_pids.add(proc["pid"])
        if tag not in expected:
            extras.append({"tag": tag, **row})
        elif observed != expected[tag]:
            failures.append(
                f"{tag}: expected {expected[tag][0]}_{expected[tag][1]}, "
                f"observed {observed}"
            )
        if proc["ppid"] != 1 or proc["sid"] != proc["pid"]:
            failures.append(
                f"{tag}: not independently detached "
                f"(pid={proc['pid']} ppid={proc['ppid']} sid={proc['sid']})"
            )
    for tag in expected:
        count = len(workers.get(tag, []))
        if count != 1:
            failures.append(f"{tag}: expected one worker, observed {count}")
    return workers, worker_pids, extras, failures


def audit(sbx: Path) -> dict:
    processes = _processes()
    failures = []
    warnings = []
    try:
        expected, manifest = _manifest(sbx)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        expected, manifest = {}, {}
        failures.append(f"invalid worker manifest: {exc}")
    workers, worker_pids, extras, worker_failures = classify_workers(
        processes, expected
    )
    failures.extend(worker_failures)
    if extras:
        warnings.append(
            f"{len(extras)} detached safe-contract worker(s) are not in the "
            "current manifest; they still own their engine descendants"
        )
    campaign = str(manifest.get("campaign") or CAMPAIGN)
    npz_fragment = str(
        manifest.get("npz_dir") or f"data/matrix_npz/{campaign}"
    ).rstrip("/")

    repaired_engines = []
    for proc in processes.values():
        index = _script_index(proc["argv"], "backtest_v8_engine.py")
        if index is None or not any(
            npz_fragment in arg for arg in proc["argv"]
        ):
            continue
        ancestor = proc["ppid"]
        seen = set()
        while ancestor in processes and ancestor not in seen and ancestor not in worker_pids:
            seen.add(ancestor)
            ancestor = processes[ancestor]["ppid"]
        owned = ancestor in worker_pids
        repaired_engines.append(
            {"pid": proc["pid"], "ppid": proc["ppid"], "worker_owned": owned}
        )
        if not owned:
            failures.append(
                f"repaired exact engine {proc['pid']} has no live repaired-worker ancestor"
            )

    claim_path = sbx / "data" / "param_matrix_claims.db"
    claims = []
    dead_claims = []
    if not claim_path.exists():
        failures.append(f"missing claims DB: {claim_path}")
    else:
        con = sqlite3.connect(f"file:{claim_path}?mode=ro", uri=True)
        try:
            for unit, worker, ts in con.execute(
                "SELECT unit,worker,ts FROM claims ORDER BY unit"
            ):
                owner = str(worker).rsplit(":", 1)[-1]
                live = owner.isdigit() and int(owner) in processes
                row = {"unit": unit, "worker": worker, "ts": ts, "live": live}
                claims.append(row)
                if not live:
                    dead_claims.append(row)
        finally:
            con.close()
        if dead_claims:
            failures.append(f"{len(dead_claims)} claims belong to dead workers")

    report = sbx / "data" / "reports" / "SWITCH_MATRIX_TRB.xlsx"
    sheets = []
    if not report.exists():
        failures.append(f"missing canonical workbook: {report}")
    else:
        try:
            from openpyxl import load_workbook

            book = load_workbook(report, read_only=True, data_only=False)
            sheets = book.sheetnames
            book.close()
        except Exception as exc:
            failures.append(f"cannot read canonical workbook: {exc}")
        if sheets != EXPECTED_SHEETS:
            failures.append(
                f"canonical workbook sheet contract changed: "
                f"expected {EXPECTED_SHEETS}, observed {sheets}"
            )

    return {
        "status": (
            "FAIL"
            if failures
            else ("PASS_WITH_UNMANIFESTED_WORKERS" if extras else "PASS")
        ),
        "campaign": campaign,
        "manifest_expected": {
            tag: {"symbol": value[0], "side": value[1]}
            for tag, value in expected.items()
        },
        "workers": workers,
        "unmanifested_workers": extras,
        "repaired_engines": repaired_engines,
        "claims_total": len(claims),
        "dead_claims": dead_claims,
        "canonical_workbook": {
            "path": str(report),
            "size": report.stat().st_size if report.exists() else None,
            "mtime": report.stat().st_mtime if report.exists() else None,
            "sheet_count": len(sheets),
            "sheets": sheets,
        },
        "warnings": warnings,
        "failures": failures,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--sbx",
        type=Path,
        default=Path(os.environ.get("V8_SBX", "/home/niels/binance-sandbox")),
    )
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = audit(args.sbx)
    payload = json.dumps(result, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload, encoding="utf-8")
    sys.stdout.write(payload)
    return 0 if not result["failures"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
