#!/usr/bin/env python3
"""Defaults diff (Agent DEF 2026-10-01): current templates vs TRAINWRECK vs pre-trainwreck backup vs configs. Read-only; writes data/wiring/defaults/defaults_diff.csv"""
import csv, json, re, sys
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
CS = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
SRC = {"current": "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_{c}.xlsx", "trainwreck": "SPREADSHEETS/TRAINWRECK_20261001/TEMPLATE_{c}.xlsx", "clean2110": "backups/before_daily_template_update_202609302110_TEMPLATE_{c}.xlsx"}
SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
def norm(v):
    s = str(v).strip()
    if s.lower() in ("true", "false"): return s.capitalize()
    try:
        f = float(s); return str(int(f)) if f == int(f) else repr(round(f, 9))
    except Exception: return s
def tdefaults(path):
    out = {}
    wb = openpyxl.load_workbook(path, read_only=False)
    for tab in SHEETS:
        if tab not in wb.sheetnames: continue
        ws = wb[tab]
        for r in range(3, ws.max_row + 1):
            a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
            if a in (None, "") or b in (None, ""): continue
            if str(ws.cell(r, 12).value).strip().upper() == "YES" or (ws.cell(r, 2).font.b and ws.cell(r, 12).value is None):
                out.setdefault(str(a).strip(), {})[tab] = norm(b)
    return out
def cfg_vals(fn, cls_re):
    t = Path(ROOT / fn).read_text(errors="ignore"); d = {}
    for m in re.finditer(r"^    ([A-Z][A-Z0-9_]+)\s*(?::[^=\n]+)?=\s*([^#\n]+)", t, re.M):
        d.setdefault(m.group(1), norm(m.group(2).strip().strip("'\"")))
    return d
def main():
    cat = json.loads((ROOT / "data/per_sym_settings.json").read_text())
    cfgc, cfgt, qc = cfg_vals("config.py", 0), cfg_vals("config_tradier.py", 0), cfg_vals("v12_quick_engine.py", 0)
    rows = []
    for c in CS:
        T = {k: tdefaults(ROOT / v.format(c=c)) for k, v in SRC.items()}
        keys = set().union(*[set(x) for x in T.values()])
        cfg = cfgc if c.startswith("CRYPTO") else cfgt
        for k in sorted(keys):
            vals = {s: sorted(set(T[s].get(k, {}).values())) for s in T}
            cur, tw, cl = vals["current"], vals["trainwreck"], vals["clean2110"]
            cd = norm(cat[c][k]) if k in cat[c] else ""
            cv, qv = cfg.get(k, ""), qc.get(k, "")
            if cur and (tw and cur != tw or cl and cur != cl or cd and [cd] != cur or cv and [cv] != cur):
                src = "trainwreck" if tw and cur != tw else "clean2110" if cl and cur != cl else "cat_side" if cd and [cd] != cur else "config"
                old = (tw or cl or ([cd] if cd else [cv]))
                rows.append([k, c, "|".join(old), "|".join(cur), src, "RESTORE_OLD" if old else "REVIEW", "|".join(tw), "|".join(cl), cd, cv, qv])
            elif not cur:
                rows.append([k, c, "|".join(tw or cl), "", "removed_in_current", "REVIEW", "|".join(tw), "|".join(cl), cd, cv, qv])
    out = ROOT / "data/wiring/defaults/defaults_diff.csv"
    with out.open("w", newline="") as f:
        w = csv.writer(f); w.writerow(["key", "cat_side", "old", "new", "source", "action", "trainwreck", "clean2110", "cat_side_json", "config", "quickconfig"]); w.writerows(rows)
    print(len(rows), out)
main()
