#!/usr/bin/env python3
"""AVG2: export AVG_DELTA / POS_SYM of the templates to data/avg_delta_pos_sym.json for the sampler (reads the template columns, NO-LIES: blank stays None)."""
import json, sys, datetime
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
TEMPL = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
HDR = 2
def main():
    tdir = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "SPREADSHEETS"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data" / "avg_delta_pos_sym.json"
    res = {"_meta": {"at": datetime.datetime.utcnow().isoformat() + "Z", "source": str(tdir)}}
    for cs, fn in TEMPL.items():
        wb = openpyxl.load_workbook(str(tdir / fn), read_only=False)
        d = {}
        for tab in SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            cols = {str(ws.cell(row=HDR, column=c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
            ca, cp = cols.get("AVG_DELTA"), cols.get("POS_SYM")
            if not (ca and cp):
                raise SystemExit(f"{cs}/{tab}: AVG_DELTA/POS_SYM header missing")
            for r in range(HDR + 1, ws.max_row + 1):
                a, b = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
                if a in (None, "") or b in (None, ""):
                    continue
                av, pv = ws.cell(row=r, column=ca).value, ws.cell(row=r, column=cp).value
                d[f"{tab}!{str(a).strip()}={str(b).strip()}"] = {"avg_delta": av if isinstance(av, (int, float)) else None, "pos_sym": int(pv) if isinstance(pv, (int, float)) else None}
        res[cs] = d
    out.write_text(json.dumps(res, indent=1))
    for cs in TEMPL:
        v = [x["pos_sym"] for x in res[cs].values()]
        print(cs, "rows", len(v), "pos_sym None", sum(1 for x in v if x is None), "0", sum(1 for x in v if x == 0), "1", sum(1 for x in v if x == 1), "2", sum(1 for x in v if x == 2), "3", sum(1 for x in v if x == 3), "4+", sum(1 for x in v if x is not None and x >= 4))
if __name__ == "__main__":
    main()
