"""Parse TEMPLATE_*.xlsx into candidate universes. Run once per template change. Read-only."""
import json, sys
sys.path.insert(0, "/Users/niels/Documents/binance")
from openpyxl import load_workbook

SKIP = {"LEGEND_FILTERS", "INSTRUCTIONS", "FILTERS_EXPLAINED", "INSTRUCTIONS_V2", "Results_Deltas", "FILTER_DICTIONARY_V2", "Results_30d_Deltas", "WIRING_INVENTORY", "FINAL_FILTER_RECHECK", "COMPLIANCE_REPAIR"}
ENTRY_TABS = {"ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES"}
EXIT_TABS = {"EXIT_STRUCTURAL", "EXIT_VELOCITY", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"}


def coerce(v):
    if v is None or v is True or v is False:
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "True":
        return True
    if s == "False":
        return False
    if s in ("", "None", "OFF"):
        return s if s == "OFF" else None
    try:
        return int(s) if "." not in s and "p" not in s else float(s)
    except ValueError:
        return s


def parse(path, out):
    wb = load_workbook(path, read_only=True, data_only=False)
    uni = {"entry_switches": {}, "entry_yellows": set(), "exit_switches": {}, "defaults": {}}
    for sn in wb.sheetnames:
        if "BASELINE" in sn or sn in SKIP:
            continue
        ws = wb[sn]
        hdr = {}
        for c in ws[2]:
            if c.value:
                hdr[c.column] = c.value
        if sn in ENTRY_TABS:
            for c in sorted(hdr):
                if c >= 15 and hdr[c] and "=" in str(hdr[c]):
                    uni["entry_yellows"].add(str(hdr[c]).strip())
        for row in ws.iter_rows(min_row=3, values_only=True):
            a = row[0]
            if not a or not isinstance(a, str) or not a.strip():
                continue
            sw, b = a.strip(), coerce(row[1])
            if b is None:
                continue
            if sn in ENTRY_TABS:
                uni["entry_switches"].setdefault(sw, set()).add(json.dumps(b))
            elif sn in EXIT_TABS:
                uni["exit_switches"].setdefault(sw, set()).add(json.dumps(b))
            lv = row[11] if len(row) > 11 else None
            if str(lv).upper() == "YES":
                uni["defaults"][sw] = b
    uni["entry_yellows"] = sorted(uni["entry_yellows"])
    for k in ("entry_switches", "exit_switches"):
        uni[k] = {sw: sorted(v) for sw, v in uni[k].items()}
    json.dump(uni, open(out, "w"), indent=1)
    print(f"{path}: entry_sw={len(uni['entry_switches'])} entry_y={len(uni['entry_yellows'])} exit_sw={len(uni['exit_switches'])} defaults={len(uni['defaults'])} -> {out}")


if __name__ == "__main__":
    for cat_side in ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"):
        parse(f"SPREADSHEETS/TEMPLATE_{cat_side}.xlsx", f"data/reports/gain_pusher/universe_{cat_side}.json")
