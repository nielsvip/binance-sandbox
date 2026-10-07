"""Paint template filter cells from finalized-only pos_sym evidence (USER 2026-10-07).

Reads data/cat_avg/priority_evidence_{cat}.json (built by tools/v15_cat_avg_matrix.py
from finalized *_bh*_gain*_t*_30d_matrix.xlsx, 1 per sym_side) and, for each of the
8 template files (main + FINAL_NORM x 4 cat_sides):

  - filter data cell (row 3+, FILTER=opt header col): pos>=4 -> bright FFFF00,
    1-3 -> light FFF2CC, 0/NONE -> clear yellow fill (other fills untouched).
  - row POS_SYM column (by header): pos count int, or blank when the row was
    never numeric (mirrors the AVG workbook N column by (switch, setting) key).

Row-2 headers are never touched. M AVG_DELTA untouched. DRY-RUN by default;
--apply writes (per-file backup + atomic tmp/zip-verify/replace).

Usage: python3 tools/v15_yellow_paint_pos.py [--apply] [--sets main,norm] [--cats ...]
"""
import argparse
import copy
import datetime
import json
import os
import shutil
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from tools.v15_cat_avg_matrix import TABS, norm_cand  # noqa: E402

BRIGHT = "FFFFFF00"
LIGHT = "FFFFF2CC"
YELLOWS = ("FFFF00", "FFF2CC")


def is_yellow_fill(fill) -> bool:
    try:
        if fill is None or getattr(fill, "fill_type", None) != "solid":
            return False
        rgb = str(getattr(getattr(fill, "fgColor", None), "rgb", "") or "")
        return rgb.upper().endswith(YELLOWS)
    except Exception:
        return False


def chain_active() -> bool:
    """True when the S1 daily chain holds its lock (non-blocking probe, never held).
    Painting under a running chain is a read-modify-write race the chain wins."""
    try:
        import fcntl
        _lk = open("/tmp/v15_daily_chain.lock", "w")
        try:
            fcntl.flock(_lk, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except (BlockingIOError, OSError):
            return True
        fcntl.flock(_lk, fcntl.LOCK_UN)
        return False
    except Exception:
        return False


def paint_file(path, rows_ev, cells_ev, cat, apply, stamp):
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.load_workbook(path)
    stats = {"bright": 0, "light": 0, "cleared": 0, "already": 0, "n_written": 0, "n_blanked": 0, "n_same": 0, "rows_missing_ev": 0, "refused": False}
    if apply and chain_active():
        wb.close()
        stats["refused"] = True
        return stats
    try:
        for tab in TABS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            hdr2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
            try:
                c_pos = next(i + 1 for i, h in enumerate(hdr2[:14]) if isinstance(h, str) and h.strip() == "POS_SYM")
            except StopIteration:
                print(f"  {Path(path).name} {tab}: no POS_SYM header, skipped", flush=True)
                continue
            fcols = []
            for c in range(15, ws.max_column + 1):
                h = hdr2[c - 1]
                if h is None or (isinstance(h, str) and not h.strip()):
                    continue
                hs = str(h).strip()
                if hs.startswith("WHAT SWITCH"):
                    break
                if "=" not in hs:
                    continue
                fcols.append((c, hs))
            for r in range(3, ws.max_row + 1):
                a = ws.cell(row=r, column=1).value
                if a is None or (isinstance(a, str) and not a.strip()):
                    continue
                sw = str(a).strip()
                cand = norm_cand(ws.cell(row=r, column=2).value)
                rk = f"{tab}!{sw}={cand}"
                rev = rows_ev.get(rk)
                want_n = rev[1] if rev is not None else None
                cur_n = ws.cell(row=r, column=c_pos).value
                cur_n = int(cur_n) if isinstance(cur_n, bool) is False and isinstance(cur_n, (int, float)) else (None if cur_n in (None, "") else cur_n)
                if want_n != cur_n:
                    if apply:
                        ws.cell(row=r, column=c_pos).value = want_n
                    stats["n_written" if want_n is not None else "n_blanked"] += 1
                else:
                    stats["n_same"] += 1
                if rev is None:
                    stats["rows_missing_ev"] += 1
                for c, h in fcols:
                    ev = cells_ev.get(f"{rk}@{h}")
                    pos = ev[1] if ev is not None else None
                    cell = ws.cell(row=r, column=c)
                    if pos is not None and pos >= 4:
                        want = BRIGHT
                    elif pos is not None and pos >= 1:
                        want = LIGHT
                    else:
                        want = None
                    if want is not None:
                        cur = str(getattr(getattr(cell.fill, "fgColor", None), "rgb", "") or "").upper()
                        if cur == want and getattr(cell.fill, "fill_type", None) == "solid":
                            stats["already"] += 1
                        else:
                            if apply:
                                cell.fill = PatternFill("solid", fgColor=want)
                            stats["bright" if want == BRIGHT else "light"] += 1
                    elif is_yellow_fill(cell.fill):
                        if apply:
                            cell.fill = PatternFill(fill_type=None)
                        stats["cleared"] += 1
        if apply:
            bkp = ROOT / "backups" / f"before_yellow_paint_pos_{stamp}_{Path(path).parent.name}_{Path(path).name}"
            shutil.copy2(path, bkp)
            tmp = str(path) + ".painttmp"
            wb.save(tmp)
            with zipfile.ZipFile(tmp) as z:
                assert len(z.namelist()) >= 10, f"truncated save? {len(z.namelist())} entries"
            os.replace(tmp, str(path))
    finally:
        wb.close()
    return stats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write files (default: dry-run report only)")
    ap.add_argument("--sets", default="main,norm")
    ap.add_argument("--cats", default=",".join(["CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"]))
    ap.add_argument("--evidence-dir", default=str(ROOT / "data" / "cat_avg"))
    a = ap.parse_args()
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%d%H%M")
    dirs = {"main": ROOT / "SPREADSHEETS", "norm": ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"}
    if a.apply and chain_active():
        print("[paint] REFUSED: daily chain is running (lock held) — repaint after it finishes", flush=True)
        return
    for cat in [c.strip() for c in a.cats.split(",") if c.strip()]:
        ev = json.loads((Path(a.evidence_dir) / f"priority_evidence_{cat}.json").read_text())
        rows_ev, cells_ev = ev.get("rows") or {}, ev.get("cells") or {}
        print(f"[paint] {cat}: rows_ev={len(rows_ev)} cells_ev={len(cells_ev)}", flush=True)
        for s in [s.strip() for s in a.sets.split(",") if s.strip()]:
            path = dirs[s] / f"TEMPLATE_{cat}.xlsx"
            if not path.exists():
                print(f"  {s}: missing {path}, skipped", flush=True)
                continue
            st = paint_file(str(path), rows_ev, cells_ev, cat, a.apply, stamp)
            if st.get("refused"):
                print(f"  {s}/{path.name}: REFUSED (chain took the lock mid-run) — rerun after the chain", flush=True)
                return
            print(f"  {s}/{path.name}: bright={st['bright']} light={st['light']} cleared={st['cleared']} already={st['already']} n_written={st['n_written']} n_blanked={st['n_blanked']} n_same={st['n_same']} rows_missing_ev={st['rows_missing_ev']}", flush=True)
    print("[paint] DRY-RUN (no writes)" if not a.apply else "[paint] APPLIED", flush=True)


if __name__ == "__main__":
    main()
