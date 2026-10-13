#!/usr/bin/env python3
"""V17 AUGMENT PROBE — measure why augment switches move nothing (Task A, 2026-10-13).

Real engine only: tools.opt.evaluate_v12.prepare + _evaluate_prepared_nocache (include_ledger, no eval cache) on worker NPZ.
Per sym_side: base (prepared cat_side defaults) + every non-default AUGMENT_TREND / AUGMENT_RISK_SIZING template candidate
+ forced probe sets. Records trades, gain_pct, AUGMENT event count + reasons, and an augment GATE FUNNEL obtained by
wrapping (not reimplementing) the engine's own gate functions:
  gl_calls/gl_fire          vec_decisions.gain_ladder_augment.gain_ladder_fire   (ladder source, every open bar)
  gl_profit_bars            bars where raw gain > 0 (ladder profit gate passes)
  gl_peak_ge_min            bars where peak gain >= effective min (pullback leg reachable)
  uag_calls/uag_pass        vec_decisions.uagain_gate.uagain_gate_pass          (bounce/pyramid/fast-riser sources)
  agg_calls/agg_block       vec_decisions.aug_gain_gate.passes                  (crypto live AUGMENT gain gate twin)
  allowed_calls/allowed_no  v12_quick_engine._augment_allowed                    (final profit/loss gate)
Usage (worker): python3 tools/v17_augment_probe.py --symsides ADAUSDC_LONG,SOLUSDC_SHORT --out data/v17_augment_probe.json
"""
import argparse
import collections
import concurrent.futures as cf
import dataclasses
import json
import os
import sys
import time

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ.setdefault("V15_EVAL_CACHE", "0")
TABS = ("AUGMENT_TREND", "AUGMENT_RISK_SIZING")
FORCED = {
    "F_LADDER_MIN_2p5": {"AUGMENT_MIN_GAIN_PCT": 2.5, "MIN_GAIN_TO_BUY_AGGRESSIVELY": 2.5},
    "F_NO_UAG": {"UNIVERSAL_AUGMENT_GAIN_GATE_ENABLED": False},
    "F_NO_AGG": {"AUG_GAIN_GATE_LIVE_TWIN_ENABLED": False},
    "F_PYR_BOUNCE_ON": {"PYRAMID_ENABLED": True, "BOUNCE_AUGMENT_ENABLED": True},
    "F_PYR_BOUNCE_ON_NO_UAG_NO_AGG": {"PYRAMID_ENABLED": True, "BOUNCE_AUGMENT_ENABLED": True, "UNIVERSAL_AUGMENT_GAIN_GATE_ENABLED": False, "AUG_GAIN_GATE_LIVE_TWIN_ENABLED": False},
    "F_PROFIT_FLAG_OFF_ONLY_TRADIER": {"AUGMENT_ONLY_WHEN_PROFITABLE_TRADIER": False},
    "F_ABL_FILTER_AUGMENT": {"ABLATION_DISABLE_FILTER_AUGMENT": True},
    "F_ABL_FILTER_ALL": {"ABLATION_DISABLE_FILTER_ENTRY": True, "ABLATION_DISABLE_FILTER_MTF_HTF": True, "ABLATION_DISABLE_FILTER_REENTRY": True, "ABLATION_DISABLE_FILTER_AUGMENT": True},
    "F_MAXAUG_BIG": {"MAX_AUGMENTS_PER_POSITION": 999999},
    "F_STOCK_GUARD_OFF": {"STOCK_LIVE_AUGMENT_ONLY": False, "MAX_AUGMENTS_PER_POSITION": 999999},
}


def template_candidates(cat_side):
    import openpyxl
    p = os.path.join(ROOT, "SPREADSHEETS", "TEMPLATE_FINAL_NORM", f"TEMPLATE_{cat_side}.xlsx")
    wb = openpyxl.load_workbook(p, read_only=True)
    out = []
    for tab in TABS:
        ws = wb[tab]
        hdr = [c for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
        idef = hdr.index("is_default") if "is_default" in hdr else None
        for r in ws.iter_rows(min_row=3, values_only=True):
            if not r or not r[0]:
                continue
            if idef is not None and str(r[idef]).strip().upper() == "YES":
                continue
            out.append((tab, str(r[0]).strip(), r[1]))
    return out


class Funnel:
    def __init__(self):
        self.c = collections.Counter()
        self.min_gain = None


def _install(F):
    import v12_quick_engine as V
    import vec_decisions.gain_ladder_augment as GL
    import vec_decisions.uagain_gate as UG
    import vec_decisions.aug_gain_gate as AG
    o_gl, o_ug, o_ag, o_al = GL.gain_ladder_fire, UG.uagain_gate_pass, AG.passes, V._augment_allowed

    def gl(config, is_long, px, last_aug_px, raw, mx):
        r = o_gl(config, is_long, px, last_aug_px, raw, mx)
        F.c["gl_calls"] += 1
        mg = GL.effective_min_gain(config)
        F.min_gain = mg
        if raw > 0:
            F.c["gl_profit_bars"] += 1
        if raw >= mg:
            F.c["gl_raw_ge_min"] += 1
        if mx >= mg:
            F.c["gl_peak_ge_min"] += 1
        if r[0]:
            F.c["gl_fire"] += 1
        return r

    def ug(*a, **k):
        r = o_ug(*a, **k)
        F.c["uag_calls"] += 1
        F.c["uag_pass"] += int(bool(r))
        return r

    def ag(*a, **k):
        r = o_ag(*a, **k)
        F.c["agg_calls"] += 1
        F.c["agg_block"] += int(not r)
        return r

    def al(cfg, pnl):
        r = o_al(cfg, pnl)
        F.c["allowed_calls"] += 1
        F.c["allowed_no"] += int(not r)
        return r

    GL.gain_ladder_fire, UG.uagain_gate_pass, AG.passes, V._augment_allowed = gl, ug, ag, al


def _run(prep, ov, F):
    from tools.opt import evaluate_v12 as E
    from tools.opt import v12_pilot as VP
    F.c.clear()
    t0 = time.time()
    r = E._evaluate_prepared_nocache(prep, VP._venue_augment_guard(prep, dict(ov)), include_ledger=True)
    led = r.get("execution_ledger") or r.get("ledger") or []
    augs = [e for e in led if isinstance(e, dict) and e.get("type") == "AUGMENT"]
    reasons = collections.Counter(str(e.get("reason", ""))[:24] for e in augs)
    return {
        "gain_pct": round(float(r.get("gain_pct") or 0.0), 4),
        "trades": int(r.get("trades") or 0),
        "valid": bool(r.get("valid")),
        "invalid_reason": r.get("invalid_reason"),
        "augments": len(augs),
        "aug_reasons": dict(reasons.most_common(6)),
        "funnel": dict(F.c),
        "eff_min_gain": F.min_gain,
        "secs": round(time.time() - t0, 2),
    }


def probe_symside(ss):
    import v12_quick_engine as V
    import v15_pilot as P
    from tools.opt import evaluate_v12 as E
    import cat_side_defaults as CSD
    F = Funnel()
    _install(F)
    prep = E.prepare(ss, 30)
    if prep is None:
        return ss, {"error": "prepare failed"}
    qd = {f.name: f.default for f in dataclasses.fields(V.QuickConfig)}
    sym, _, side = ss.rpartition("_")
    cat = CSD.cat_side_of(sym, side, prep.get("mode"))
    base = _run(prep, {}, F)
    bcfg = prep.get("base_cfg_dict") or {}
    base["base_values"] = {k: bcfg.get(k) for k in ("AUGMENT_ONLY_WHEN_PROFITABLE", "AUGMENT_ONLY_WHEN_PROFITABLE_TRADIER", "AUGMENT_MIN_GAIN_PCT", "MIN_GAIN_TO_BUY_AGGRESSIVELY", "MAX_AUGMENTS_PER_POSITION", "AUGMENTATION_COOLDOWN_SECONDS", "UNIVERSAL_AUGMENT_GAIN_GATE_ENABLED", "PYRAMID_ENABLED", "BOUNCE_AUGMENT_ENABLED", "STOCK_LIVE_AUGMENT_ONLY", "BOUNCE_AUGMENT_MIN_LOSS_PCT", "PULLBACK_AUGMENT_ENABLED")}
    out = {"cat_side": cat, "mode": prep.get("mode"), "base": base, "flips": {}, "forced": {}}
    for tab, sw, val in template_candidates(cat):
        key = f"{tab}|{sw}={val}"
        try:
            pv = P._parse_opt_value(val, qd.get(sw))
            ov = P._switch_overrides(sw, pv)
        except Exception as e:
            out["flips"][key] = {"err": str(e)[:80]}
            continue
        if sw.startswith("ABLATION_") and pv is True:
            out["flips"][key] = {"skip": "ABLATION_NEVER_ON (pilot rule)"}
            continue
        if all(str(bcfg.get(k)) == str(v) for k, v in ov.items()):
            out["flips"][key] = {"identity": True}
            continue
        rec = _run(prep, ov, F)
        rec["d_gain"] = round(rec["gain_pct"] - base["gain_pct"], 4)
        rec["d_trades"] = rec["trades"] - base["trades"]
        rec["d_augments"] = rec["augments"] - base["augments"]
        out["flips"][key] = rec
    for name, ov in FORCED.items():
        rec = _run(prep, ov, F)
        rec["d_gain"] = round(rec["gain_pct"] - base["gain_pct"], 4)
        rec["d_augments"] = rec["augments"] - base["augments"]
        out["forced"][name] = rec
    return ss, out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--symsides", required=True)
    ap.add_argument("--out", default=os.path.join(ROOT, "data", "v17_augment_probe.json"))
    ap.add_argument("--workers", type=int, default=6)
    a = ap.parse_args()
    sss = [s.strip() for s in a.symsides.split(",") if s.strip()]
    res = {}
    import multiprocessing as mp
    with cf.ProcessPoolExecutor(max_workers=min(a.workers, len(sss)), mp_context=mp.get_context("fork")) as ex:
        for ss, r in ex.map(probe_symside, sss):
            res[ss] = r
            print(ss, "done", flush=True)
    res["_meta"] = {"at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "host": os.uname().nodename, "engine": "tools.opt.evaluate_v12._evaluate_prepared_nocache (real v12_quick_engine.simulate_one), eval cache off"}
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    json.dump(res, open(a.out, "w"), indent=1, default=str)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
