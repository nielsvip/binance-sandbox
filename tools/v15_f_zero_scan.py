#!/usr/bin/env python3
"""v15_f_zero_scan — report (and with --fix blank) column F cells that are an exact 0 / 0.0 in sweep xlsx (NO-LIES: F is blank unless a real hustle delta).
Real exact-zero deltas are only legitimate when the JSON progress record says delta_vs_initial == 0.0; anything else is blanked with --fix.
  python tools/v15_f_zero_scan.py DIR [--progress DIR] [--fix]"""
import argparse, glob, json, os, sys
import openpyxl
SW = {'STDEV_SLOPE_SIZING','ENTRY_REVERSAL_BOUNCE','ENTRY_BREAKOUT_CHANNEL','ENTRY_CONFIRMATION_GATES','EXIT_STRUCTURAL','EXIT_VELOCITY','REENTRY_WINDOWED','REENTRY_ADAPTIVE','AUGMENT_TREND','AUGMENT_RISK_SIZING','REDUCE_PROFIT_LOCK','REDUCE_SIGNAL_RATER','GLOBAL_RISK_GATES'}
ap = argparse.ArgumentParser(); ap.add_argument("dir"); ap.add_argument("--progress"); ap.add_argument("--fix", action="store_true"); ap.add_argument("--limit", type=int, default=0); a = ap.parse_args()
tot = zeros = fixed = 0
for p in (sorted(glob.glob(os.path.join(a.dir, "*_matrix.xlsx")))[:a.limit] if a.limit else sorted(glob.glob(os.path.join(a.dir, "*_matrix.xlsx")))):
    try: wb = openpyxl.load_workbook(p)
    except Exception: continue
    ss = os.path.basename(p).split("_30d")[0]
    real = set()
    if a.progress:
        try:
            j = json.load(open(os.path.join(a.progress, f"{ss}_v14_progress.json")))
            for k, v in (j.get("done") or {}).items():
                if isinstance(v, dict) and v.get("delta_vs_initial") == 0.0 and not v.get("is_running"): real.add(k.split(":", 1)[0])  # key = "TAB!row:SWITCH=cand" -> "TAB!row"
        except Exception: pass
    dirty = False
    for ws in wb.worksheets:
        if ws.title not in SW: continue
        for r in range(3, ws.max_row + 1):
            v = ws.cell(row=r, column=6).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                tot += 1
                if v == 0:
                    if f"{ws.title}!{r}" in real: continue
                    zeros += 1
                    if a.fix: ws.cell(row=r, column=6).value = None; dirty = True; fixed += 1
    if dirty:
        tmp = p + ".fzero.tmp"; wb.save(tmp); openpyxl.load_workbook(tmp); os.replace(tmp, p)
print(f"F numeric={tot} F_exact_zero_unproven={zeros} blanked={fixed}")
