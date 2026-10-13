#!/usr/bin/env python3
"""v17_combo_engine — LEARNED + EXPLORED switch/filter/TF combination search (BIBLE §97, USER 2026-10-13: "v17 ... LEARNS FROM
PAST DATA AND STARTS IDENTIFYING AND EMPLOYING TENS OF THOUSANDS OF CUSTOM SWITCH/FILTER/TF COMBINATIONS").

API (called by the V17 per-board SURGEON in v15_pilot.py inside its pinpoint budget, or standalone):

    from tools import v17_combo_engine as CE
    ranked = CE.search(prepared, sym_side, base_overrides, symptoms=None, regime=None, budget_s=300, workers=4,
                       stats_out=None, evaluator=None, library=None, seed=None)

    prepared        tools.opt.v12_pilot.prepare_batch(sym_side, 30) (None -> prepared here)
    base_overrides  the board's current best set (full set, as the pilot keeps cumulative_overrides)
    symptoms        pilot surgeon shape {name: {"n", "size"}} (missed_moves, losers, premature_exits, giveback, missed_augment,
                    churn, tim_low, tim_high, below_bh, old_set_recent); None -> diagnosed here (1 ledger eval + v15_trade_autopsy)
    regime          pilot _v17_regime() dict (optional; only labels/features are read)
    -> ranked list of VERIFIED combos, best first:
       {"kind": "joint"|"combo", "overrides_delta", "gain", "base_gain", "delta", "trades", "tim", "dd", "wr",
        "recent_segment_gain", "recent_delta", "n_evals", "origin", "why", "flags", "tf", ...}
       kind "joint" (first, only when it lifts) = the greedy-combined set, freshly re-evaluated with the eval cache OFF on 30D
       and on the recent 7D segment; every number is a real v12_quick_engine eval (evaluate_prepared_sanitized).

Every value comes from the four TEMPLATE_* candidate sets (rows + yellow filter headers), numeric neighbourhoods inside the
template span of that switch, FILTER_EXEMPTIONS from data/v17_filter_exemptions.json / v17_priors blocks, or the config field
type (AUGMENT_INTO_REENTRY family) — each validated by tools/switch_types.coerce (fail-closed: refused -> combo dropped).
Promotion gates = the pilot's (imported from v15_pilot, never re-implemented): valid (sanitized floor, TIM<=80, DD<=30),
trades >= ADAPT_FLOOR_TRADES, tim_guard_veto / trades_guard_veto vs the chain, promotion_block_reason, MAXAUG never 0,
ABLATION never on. Library: data/v17_combo_library.sqlite (build-library / merge / info CLI below).

CLI:
  python3 tools/v17_combo_engine.py build-library [--out PATH] [--progress-glob G ...] [--delta-glob G ...]
  python3 tools/v17_combo_engine.py merge --into PATH --from PATH [...]
  python3 tools/v17_combo_engine.py info [--lib PATH]
  python3 tools/v17_combo_engine.py space --cat-side CRYPTO_LONG
  python3 tools/v17_combo_engine.py proof --sym-side RLCUSDT_LONG [--budget 300] [--workers 4] [--out JSON]
"""
import argparse
import datetime
import glob
import hashlib
import json
import math
import os
import random
import re
import sqlite3
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
LIB_PATH = Path(os.environ.get("V17_COMBO_LIB") or (ROOT / "data" / "v17_combo_library.sqlite"))
TF_SEM_PATH = Path(os.environ.get("V17_TF_SEMANTICS") or (ROOT / "data" / "v17_tf_semantics.json"))
EXEMPT_PATH = Path(os.environ.get("V15_V17_EXEMPTIONS_JSON") or (ROOT / "data" / "v17_filter_exemptions.json"))
SWITCH_SHEETS = ("STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES")
CATEGORY_TABS = {"entry": ("ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES"), "reentry": ("REENTRY_WINDOWED", "REENTRY_ADAPTIVE"), "augment": ("AUGMENT_TREND", "AUGMENT_RISK_SIZING"), "exit": ("EXIT_STRUCTURAL", "EXIT_VELOCITY", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER"), "sizing": ("STDEV_SLOPE_SIZING",), "global": ("GLOBAL_RISK_GATES",)}
TAB_CATEGORY = {t: c for c, ts in CATEGORY_TABS.items() for t in ts}
SYMPTOM_CATS = {"missed_moves": ("entry", "reentry"), "tim_low": ("entry", "reentry"), "below_bh": ("entry", "reentry", "augment"), "premature_exits": ("exit",), "giveback": ("exit",), "churn": ("exit",), "losers": ("entry", "exit"), "tim_high": ("entry",), "missed_augment": ("augment",), "old_set_recent": ("entry", "reentry", "augment", "exit")}
LOOSEN = frozenset({"missed_moves", "tim_low", "below_bh", "missed_augment"})
TIGHTEN = frozenset({"losers", "tim_high", "churn"})
TF_SPAN = {"15m": 1, "30m": 2, "1h": 4, "2h": 8, "4h": 16, "6h": 24, "8h": 32, "12h": 48, "D": 96, "1d": 96, "3d": 288, "W": 672, "1w": 672}
AIR_FAMILY = {"REENTRY_AUGMENT_SIZE_MULT": (0.5, 1.0, 1.5, 2.0), "REENTRY_AUGMENT_REQUIRE_HIGHER_LOW": (True, False), "REENTRY_AUGMENT_HL_REF": ("entry", "low"), "REENTRY_AUGMENT_WINDOW_SEC": (0.0, 43200.0, 86400.0, 172800.0), "REENTRY_AUGMENT_REQUIRE_NO_PRIOR_AUGMENT": (True, False)}
NUM_GRID = (0.5, 0.75, 1.25, 1.5, 2.0)
ORIGIN_QUOTA = {"tf_coherent": 0.05, "lib_regime": 0.06, "lib_same_sym": 0.12, "lib_neighbour": 0.16, "lib_symptom": 0.08, "prior_synergy": 0.04, "prior_block_exempt": 0.04, "symptom_path": 0.12, "filter_path": 0.08, "exempt_path": 0.06, "numeric": 0.05, "air": 0.02, "pair": 0.08, "explore": 0.15}
RECENT_DAYS = 7  # the only short window evaluate_v12._exact_30d_slice honours (crypto 7 calendar days / stocks 7 sessions); 8 = full 30D


def _env_f(name, default):
    try:
        return float(os.environ.get(name, default))
    except (TypeError, ValueError):
        return float(default)


def _utc():
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def cat_side_of(sym_side):
    s = str(sym_side).upper()
    base = s[:-5] if s.endswith("_LONG") else s[:-6] if s.endswith("_SHORT") else s
    crypto = base.endswith(("USDT", "USDC", "USD1", "BUSD", "FDUSD", "TUSD", "DAI"))
    return f"{'CRYPTO' if crypto else 'STOCKS'}_{'LONG' if s.endswith('_LONG') else 'SHORT'}"


def venue_of(cat_side):
    return "CRYPTO" if str(cat_side).upper().startswith("CRYPTO") else "STOCKS"


def canon(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dict, list)):
        return json.dumps(v, sort_keys=True, default=str)
    s = str(v).strip() if v is not None else ""
    if s.lower() in ("true", "false", "off", "on", "none"):
        return s.upper()
    try:
        f = float(s)
        return str(int(f)) if f.is_integer() else repr(f)
    except (TypeError, ValueError):
        return s


def combo_key(delta):
    return "|".join(f"{k}={canon(delta[k])}" for k in sorted(delta))


def same_val(a, b):
    return canon(a) == canon(b)


def merge_exemptions(base, add):
    out = {}
    for src in (base, add):
        if not isinstance(src, dict):
            continue
        for path, fl in src.items():
            cur = out.setdefault(str(path), [])
            for f in (fl if isinstance(fl, (list, tuple)) else [fl]):
                if f not in cur:
                    cur.append(f)
    return out


def apply_delta(base, delta):
    out = dict(base)
    for k, v in delta.items():
        out[k] = merge_exemptions(base.get(k), v) if k == "FILTER_EXEMPTIONS" and isinstance(v, dict) else v
    return out


def effective_delta(base, delta):
    """keys of delta that really change the base (FILTER_EXEMPTIONS: only when the union adds something)."""
    out = {}
    for k, v in delta.items():
        if k == "FILTER_EXEMPTIONS":
            if canon(merge_exemptions(base.get(k), v)) != canon(merge_exemptions(base.get(k), {})):
                out[k] = v
        elif k not in base or not same_val(base.get(k), v):
            out[k] = v
    return out


# ----------------------------------------------------------------------------------------------------------- pilot bridge
class _Pilot:
    """v15_pilot's own sanitize / defaults / guards (imported, never re-implemented). Import failure = refuse (fail-closed)."""
    _mod = None

    @classmethod
    def m(cls):
        if cls._mod is None:
            import v15_pilot as P  # noqa: N813
            cls._mod = P
        return cls._mod


_DEF_CACHE = {}


def defaults_for(sym_side):
    cs = cat_side_of(sym_side)
    if cs not in _DEF_CACHE:
        _DEF_CACHE[cs] = _Pilot.m().get_defaults_for_symside(sym_side)
    return _DEF_CACHE[cs]


def sanitize(ov, sym_side):
    return _Pilot.m().sanitize_overrides(dict(ov), defaults_for(sym_side))[0]


def is_filter_key(k):
    return bool(re.search(r"FILTER|GATE|BLOCK|VETO|CONFIRM|REQUIRE|GUARD|MIN_TFS", str(k or "").upper()))


def blocked_reason(delta, tabs=()):
    P = _Pilot.m()
    for k, v in (delta or {}).items():
        k = str(k).strip()
        if k == "MAX_AUGMENTS_PER_POSITION" and P._maxaug_is_zero(str(v)):
            return "MAXAUG_NEVER_ZERO"
        if k.startswith("ABLATION_") and str(v).strip().lower() in ("true", "1", "yes", "on"):
            return "ABLATION_NEVER_ON"
        if k in P.GLOBAL_ONLY_KEYS:
            return "GLOBAL_ONLY"
        b = P.promotion_block_reason(k, None)
        if b:
            return b
    for t in tabs or ():
        b = P.promotion_block_reason("", t) if t in getattr(P, "DEAD_VEC_TABS", ()) else ""
        if b:
            return b
    return ""


def gate_reason(res, chain_res, delta_ov=None):
    """'' = promotable vs the chain (the pilot's _v17_ok without the delta>0 part)."""
    P = _Pilot.m()
    if not res:
        return "NO_RESULT"
    if not res.get("valid"):
        return "INVALID " + str(res.get("reason") or "")[:60]
    if int(res.get("trades") or 0) < P.ADAPT_FLOOR_TRADES:
        return "TRADES<FLOOR"
    b = blocked_reason(delta_ov or {})
    if b:
        return b
    t = P.tim_guard_veto(res.get("tim"), (chain_res or {}).get("tim"))
    if t:
        return t
    t = P.trades_guard_veto(res.get("trades"), (chain_res or {}).get("trades"))
    return t or ""


# ------------------------------------------------------------------------------------------------------------------ types
def coerce_value(field, value, venue):
    """(ok, native) via tools/switch_types.coerce; unknown fields refused (strict) — a combo key must exist on a surface."""
    from tools import switch_types as ST
    try:
        ok, val, _why = ST.coerce(field, value, venue, strict_unknown=True)
    except Exception:
        return False, value
    return bool(ok), val


# ------------------------------------------------------------------------------------------------------------ TF semantics
_SEM = {"memo": None, "doc": None}


def tf_semantics():
    """data/v17_tf_semantics.json (TF study lane): {"keys": {KEY: {role, tf, span_bars_15m, window_key}}, "rules": [...]}; {} absent."""
    try:
        st = TF_SEM_PATH.stat()
        memo = (st.st_mtime_ns, st.st_size)
        if _SEM["memo"] != memo:
            _SEM.update(memo=memo, doc=json.loads(TF_SEM_PATH.read_text()))
        return _SEM["doc"] or {}
    except (OSError, ValueError):
        return {}


def _tf_token(v):
    s = str(v).strip()
    if s.upper() == "OFF":
        return "OFF"
    for t in TF_SPAN:
        if s.lower() == t.lower():
            return "D" if t in ("1d", "D") else t
    return None


def _key_tf(key):
    k = str(key).upper()
    for suf, tf in (("_15M", "15m"), ("_1H", "1h"), ("_4H", "4h"), ("_D", "D"), ("_1D", "D")):
        if k.endswith(suf) or (suf + "_") in k:
            return tf
    return None


_ROLE_MEMO = {}


def key_role(key, sem=None):
    sem = sem if sem is not None else tf_semantics()
    mk = (id(sem), key)
    if mk in _ROLE_MEMO:
        return _ROLE_MEMO[mk]
    r = _key_role(key, sem)
    if len(_ROLE_MEMO) > 200000:
        _ROLE_MEMO.clear()
    _ROLE_MEMO[mk] = r
    return r


def _key_role(key, sem):
    e = ((sem.get("keys") or {}).get(key)) if isinstance(sem, dict) else None
    if isinstance(e, dict) and e.get("role"):
        return str(e["role"])
    k = str(key).upper()
    if re.search(r"MAX_BARS_AGO|_WINDOW|LOOKBACK_BARS", k):
        return "window"
    if "COOLDOWN" in k:
        return "cooldown"
    if re.search(r"MIN_HOLD|MIN_TFS|ALIGNMENT_MIN", k):
        return "confirm"
    if is_filter_key(k):
        return "permission"
    if k.startswith(("EXIT", "REDUCE")) or "_EXIT" in k or "STOP" in k:
        return "exit"
    if "SIZ" in k or "MULT" in k:
        return "sizing"
    return "trigger"


def tf_features(delta, full=None):
    """TF structure of a combo (stored in the library so learning generalises by structure, not key names):
    htf = largest TF among permission/trigger values, ltf_confirm = smallest TF among confirm/permission values,
    window_bars = largest validity-window value (bars), rearm = cooldown/consume keys touched."""
    sem = tf_semantics()
    tfs, htf, ltf, win, rearm = [], None, None, None, []
    for k, v in (delta or {}).items():
        if k == "FILTER_EXEMPTIONS":
            continue
        role = key_role(k, sem)
        e = ((sem.get("keys") or {}).get(k)) if isinstance(sem, dict) else None
        tf = _tf_token(v) or (str(e.get("tf")) if isinstance(e, dict) and e.get("tf") in TF_SPAN else None) or _key_tf(k)
        if tf and tf != "OFF":
            tfs.append(tf)
            span = TF_SPAN.get(tf, 1)
            if role in ("permission", "trigger", "structure") and (htf is None or span > TF_SPAN[htf]):
                htf = tf
            if role in ("confirm", "permission") and (ltf is None or span < TF_SPAN[ltf]):
                ltf = tf
        if role == "window" and isinstance(v, (int, float)) and not isinstance(v, bool):
            win = max(win or 0.0, float(v))
        if role in ("cooldown", "consume"):
            rearm.append(k)
    return {"tfs": sorted(set(tfs), key=lambda t: TF_SPAN.get(t, 0)), "htf": htf, "ltf_confirm": ltf, "window_bars": win, "rearm": sorted(rearm)}


C2_LTF_CONFIRMS = ("DC_BREAK_FILTER_TF", "DC_MOMENTUM_BOTA_SCORER_FILTER_TF", "CIRCUIT_SHARPE_GATES_FILTER_TF", "BT_WT_CROSS_LADDER_FILTER_TF", "BAR_PATTERNS_FILTER_TF", "BREAKOUT_RETEST_FILTER_TF", "BB_RECOVERY_FILTER_TF", "EMA_BLANKET_FILTER_FILTER_TF", "DELTA_ENGINE_FILTER_TF")


def _span(tf, venue=None):
    """15m bars per HTF bar (data/v17_tf_semantics.json C5/C6): crypto D 96 / 4h 16 / 1h 4; stocks D 26 RTH bars, 4h ~13."""
    if str(venue or "").upper().startswith("STOCK"):
        return {"D": 26, "1d": 26, "4h": 13}.get(tf, TF_SPAN.get(tf, 1))
    return TF_SPAN.get(tf, 1)


def htf_state_triggers(full):
    """HTF (D/4h) STATE triggers switched on in a full set (C1: re-fire on every flat 15m bar inside the HTF bar)."""
    out = []
    if full.get("DC_BREAKOUT_ENTRY_ENABLED") is True and _tf_token(full.get("DC_BREAKOUT_TF_EXPANDED")) in ("D", "4h"):
        out.append(("DC_BREAKOUT_TF_EXPANDED", _tf_token(full.get("DC_BREAKOUT_TF_EXPANDED"))))
    if _tf_token(full.get("BB_BREAKOUT_ENTRY_TF")) in ("D", "4h"):
        out.append(("BB_BREAKOUT_ENTRY_TF", _tf_token(full.get("BB_BREAKOUT_ENTRY_TF"))))
    return out


def coherence_reason(delta, full, sem=None, venue=None):
    """'' = coherent, else the rule that rejects the combo BEFORE evaluation (rules + evidence: data/v17_tf_semantics.json,
    docs/v17_tf_semantics.md; only ERROR-class rules reject, WARN-class rules only shape generation — coherence_warnings):
    C4  BB_BOUNCE_ENTRY_TF != OFF is also a global AND entry filter (generic_filter_tf.py:34/90-98) — measured opens 133->0;
        an HTF event trigger that lives one 15m bar (C3) — rejected as a path change.
    C9  WT_15M_BOUNCE_MAX_BARS_AGO has no decision reader (delta 0 by construction) — never spend an eval on it.
    C5  a validity window in BARS set by the combo must not exceed the span of the TF that arms it (stocks D = 26, C6)."""
    sem = sem if sem is not None else tf_semantics()
    if "BB_BOUNCE_ENTRY_TF" in delta and _tf_token(delta["BB_BOUNCE_ENTRY_TF"]) not in (None, "OFF"):
        return f"C4 BB_BOUNCE_ENTRY_TF={delta['BB_BOUNCE_ENTRY_TF']} is a one-bar HTF event AND a global entry filter"
    if "WT_15M_BOUNCE_MAX_BARS_AGO" in delta:
        return "C9 WT_15M_BOUNCE_MAX_BARS_AGO has no decision reader"
    keys = (sem.get("keys") or {}) if isinstance(sem, dict) else {}
    wins = [(k, v) for k, v in delta.items() if "BARS" in str(k).upper() and not isinstance(v, bool) and isinstance(v, (int, float)) and key_role(k, sem) == "window"]
    if not wins:
        return ""
    arm = [_tf_token(v) for k, v in (full or delta).items() if isinstance(v, str) and key_role(k, sem) == "trigger" and _tf_token(v) not in (None, "OFF")]
    arm_tf = max(arm, key=lambda t: TF_SPAN.get(t, 0)) if arm else None
    for k, v in wins:
        e = keys.get(k) or {}
        tf = _tf_token(e.get("tf")) or _key_tf(k) or arm_tf
        if tf and float(v) > _span(tf, venue):
            return f"C5 {k}={v:g} bars > {tf} span {_span(tf, venue)}"
    return ""


def coherence_warnings(delta, full, venue=None):
    """WARN-class rules (shape generation, never reject): C1 HTF state trigger with COOLDOWN 0 = unbounded re-open inside one
    HTF instance; C2 HTF state trigger without any 15m timing confirm."""
    out = []
    trig = htf_state_triggers(full or {})
    if not trig:
        return out
    cdk = "COOLDOWN_BARS_TRADIER" if str(venue or "").upper().startswith("STOCK") else "COOLDOWN_BARS"
    try:
        cd = float((full or {}).get(cdk) or 0)
    except (TypeError, ValueError):
        cd = 0.0
    if cd <= 0 and not (full or {}).get("COOLDOWN_FROM_LIVE_SECONDS_ENABLED"):
        out.append("C1_UNBOUNDED_REOPEN")
    if not any(_tf_token((full or {}).get(k)) == "15m" for k in C2_LTF_CONFIRMS) and (full or {}).get("WT_15M_BOUNCE_OPEN_ENABLED") is not True:
        out.append("C2_NO_LTF_CONFIRM")
    return out

# --------------------------------------------------------------------------------------------------------------- the space
def template_path(cat_side):
    for d in (ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM", ROOT / "SPREADSHEETS"):
        p = d / f"TEMPLATE_{cat_side}.xlsx"
        if p.exists():
            return p
    raise FileNotFoundError(f"TEMPLATE_{cat_side}.xlsx")


def _file_md5(p):
    h = hashlib.md5()
    with open(p, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def read_template(cat_side):
    """{rows: [[tab, switch, raw]], headers: {tab: [FILTER=opt]}} from TEMPLATE_{cat_side}.xlsx, cached per template md5."""
    tp = template_path(cat_side)
    md5 = _file_md5(tp)
    cache = Path(os.environ.get("V17_COMBO_CACHE_DIR") or (ROOT / "data" / "v17_combo_cache")) / f"template_{cat_side}_{md5[:12]}.json"
    try:
        return json.loads(cache.read_text())
    except (OSError, ValueError):
        pass
    import openpyxl
    wb = openpyxl.load_workbook(str(tp), read_only=True, data_only=True)
    rows, headers = [], {}
    for tab in SWITCH_SHEETS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            if i == 0:
                headers[tab] = [str(h).strip() for h in r[14:] if h is not None and "=" in str(h)]
                continue
            if r and r[0] not in (None, "") and len(r) > 1:
                raw = r[1]
                rows.append([tab, str(r[0]).strip(), raw if isinstance(raw, (bool, int, float)) else ("" if raw is None else str(raw).strip())])
    out = {"template": str(tp), "md5": md5, "rows": rows, "headers": headers}
    try:
        cache.parent.mkdir(parents=True, exist_ok=True)
        tmp = cache.with_suffix(".tmp")
        tmp.write_text(json.dumps(out, default=str))
        os.replace(tmp, cache)
    except OSError:
        pass
    return out


class ComboSpace:
    """Typed candidate values of one cat_side: template rows (path switches + filter rows), yellow filter headers per tab,
    numeric spans, exemption map. Every value already passed switch_types.coerce (refused values are counted, never used)."""

    def __init__(self, cat_side, sym_side=None):
        self.cat_side = cat_side
        self.venue = venue_of(cat_side)
        tpl = read_template(cat_side)
        self.template = tpl["template"]
        self.refused = 0
        self.values = {}  # key -> [native values] (ordered, unique)
        self.key_tabs = {}  # key -> set(tabs)
        self.filters_by_tab = {}  # tab -> [(filter, native)]
        P = _Pilot.m()
        comp = getattr(P, "COMPOSITE_SWITCHES", {})
        for tab, sw, raw in tpl["rows"]:
            if raw in ("", None):
                continue
            ok, val = (True, raw) if sw in comp else coerce_value(sw, raw, self.venue)
            if not ok:
                self.refused += 1
                continue
            self._add(sw, val, tab)
        for tab, hdrs in tpl["headers"].items():
            seen = set()
            for h in hdrs:
                f, _, o = h.partition("=")
                f, o = f.strip(), o.strip()
                ok, val = coerce_value(f, o, self.venue)
                if not ok:
                    self.refused += 1
                    continue
                ck = (f, canon(val))
                if ck in seen or blocked_reason({f: val}):
                    continue
                seen.add(ck)
                self.filters_by_tab.setdefault(tab, []).append((f, val))
                self._add(f, val, None)
        for k, vals in AIR_FAMILY.items():
            for v in vals:
                ok, val = coerce_value(k, v, self.venue)
                if ok:
                    self._add(k, val, "AUGMENT_TREND")
        ok, val = coerce_value("AUGMENT_INTO_REENTRY_ENABLED", True, self.venue)
        self.air_ok = bool(ok)
        if ok:
            self._add("AUGMENT_INTO_REENTRY_ENABLED", val, "AUGMENT_TREND")
        try:
            self.exemptions = json.loads(EXEMPT_PATH.read_text()) if EXEMPT_PATH.exists() else {}
        except (OSError, ValueError):
            self.exemptions = {}

    def _add(self, key, val, tab):
        if blocked_reason({key: val}):
            self.blocked = getattr(self, "blocked", 0) + 1
            return
        lst = self.values.setdefault(key, [])
        if all(canon(x) != canon(val) for x in lst):
            lst.append(val)
        if tab:
            self.key_tabs.setdefault(key, set()).add(tab)

    def category(self, key):
        cats = {TAB_CATEGORY.get(t) for t in self.key_tabs.get(key, ())}
        cats.discard(None)
        return sorted(cats)[0] if cats else ("augment" if key.startswith(("AUGMENT_INTO_REENTRY", "REENTRY_AUGMENT")) else None)

    def path_keys(self, cats=None):
        out = []
        for k, tabs in self.key_tabs.items():
            if is_filter_key(k):
                continue
            c = self.category(k)
            if cats is None or c in cats:
                out.append(k)
        return sorted(out)

    def filter_keys(self):
        return sorted({f for lst in self.filters_by_tab.values() for f, _ in lst} | {k for k in self.key_tabs if is_filter_key(k)})

    def to_delta(self, key, val):
        P = _Pilot.m()
        return P._switch_overrides(key, val) if key in getattr(P, "COMPOSITE_SWITCHES", {}) else {key: val}

    def numeric_neighbours(self, key, base_val):
        """base x {0.5,0.75,1.25,1.5,2} clipped to the template candidate span of this key; typed via coerce."""
        nums = [float(v) for v in self.values.get(key, []) if isinstance(v, (int, float)) and not isinstance(v, bool)]
        if len(nums) < 2 or not isinstance(base_val, (int, float)) or isinstance(base_val, bool) or base_val == 0:
            return []
        lo, hi = min(nums), max(nums)
        out = []
        for m in NUM_GRID:
            x = float(base_val) * m
            if not (lo <= x <= hi):
                continue
            if isinstance(base_val, int) or all(float(n).is_integer() for n in nums):
                x = float(round(x))
            ok, val = coerce_value(key, x, self.venue)
            if ok and not same_val(val, base_val) and all(not same_val(val, o) for o in out):
                out.append(val)
        return out

    def size(self):
        n_rows = sum(len(v) for v in self.values.values())
        n_f = sum(len(v) for v in self.filters_by_tab.values())
        return {"keys": len(self.values), "values": n_rows, "filter_cells": n_f, "refused": self.refused, "blocked": getattr(self, "blocked", 0), "path_keys": len(self.path_keys()), "exemption_keys": len(self.exemptions)}


# ------------------------------------------------------------------------------------------------------------- the library
SCHEMA = """
CREATE TABLE IF NOT EXISTS boards (board_id INTEGER PRIMARY KEY, bkey TEXT UNIQUE, sym_side TEXT, cat_side TEXT, src TEXT, run TEXT,
  ts_last REAL, gain REAL, bh REAL, tim REAL, dd REAL, trades REAL, side_net_pct REAL, vol_bar_pct REAL, chop REAL, range_pct REAL,
  regime TEXT, symptoms TEXT, engine_md5 TEXT, added_utc TEXT);
CREATE INDEX IF NOT EXISTS ix_boards_cs ON boards(cat_side);
CREATE INDEX IF NOT EXISTS ix_boards_ss ON boards(sym_side);
CREATE TABLE IF NOT EXISTS combos (combo_id INTEGER PRIMARY KEY, ckey TEXT UNIQUE, combo TEXT, n_changes INTEGER, cats TEXT,
  htf TEXT, ltf_confirm TEXT, window_bars REAL, rearm TEXT, tfs TEXT);
CREATE TABLE IF NOT EXISTS evidence (board_id INTEGER, combo_id INTEGER, source TEXT, origin TEXT, symptom TEXT, window_days INTEGER,
  n INTEGER, delta REAL, best_delta REAL, gain REAL, base_gain REAL, trades REAL, tim REAL, dd REAL, wr REAL, valid INTEGER,
  recent_delta REAL, ts TEXT, engine_md5 TEXT, base_fp TEXT, PRIMARY KEY (board_id, combo_id, source, window_days));
CREATE INDEX IF NOT EXISTS ix_ev_combo ON evidence(combo_id);
CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT);
CREATE TABLE IF NOT EXISTS regime_evidence (board_id INTEGER, combo_id INTEGER, regime TEXT, n_seg INTEGER, d_pnl REAL, d_dd REAL,
  base_pnl REAL, cand_pnl REAL, window_days INTEGER, ts TEXT, PRIMARY KEY (board_id, combo_id, regime, window_days));
CREATE INDEX IF NOT EXISTS ix_rev_regime ON regime_evidence(regime);
"""
FEATS = ("side_net_pct", "vol_bar_pct", "chop", "range_pct", "bh", "tim", "trades")


class ComboLibrary:
    def __init__(self, path=None):
        self.path = Path(path or LIB_PATH)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(str(self.path), timeout=60)
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.execute("PRAGMA busy_timeout=60000")
        self.db.executescript(SCHEMA)
        self._cid = {}

    def close(self):
        try:
            self.db.commit()
            self.db.execute("PRAGMA wal_checkpoint(TRUNCATE)")
            self.db.close()
        except sqlite3.Error:
            pass

    def board(self, bkey, **f):
        row = self.db.execute("SELECT board_id FROM boards WHERE bkey=?", (bkey,)).fetchone()
        if row:
            return row[0]
        cols = ["bkey", "added_utc"] + [k for k in f]
        vals = [bkey, _utc()] + [json.dumps(v, default=str) if isinstance(v, (dict, list)) else v for v in f.values()]
        cur = self.db.execute(f"INSERT INTO boards ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals)
        return cur.lastrowid

    def combo_id(self, delta):
        ck = combo_key(delta)
        if ck in self._cid:
            return self._cid[ck]
        row = self.db.execute("SELECT combo_id FROM combos WHERE ckey=?", (ck,)).fetchone()
        if row:
            self._cid[ck] = row[0]
            return row[0]
        tf = tf_features(delta)
        cats = sorted({c for c in (TAB_CATEGORY_GUESS(k) for k in delta) if c})
        cur = self.db.execute("INSERT INTO combos (ckey, combo, n_changes, cats, htf, ltf_confirm, window_bars, rearm, tfs) VALUES (?,?,?,?,?,?,?,?,?)", (ck, json.dumps(delta, sort_keys=True, default=str), len(delta), ",".join(cats), tf["htf"], tf["ltf_confirm"], tf["window_bars"], ",".join(tf["rearm"]), ",".join(tf["tfs"])))
        self._cid[ck] = cur.lastrowid
        return cur.lastrowid

    def add_evidence(self, board_id, delta, source, delta_val, window_days=30, origin="", symptom="", gain=None, base_gain=None, trades=None, tim=None, dd=None, wr=None, valid=None, recent_delta=None, ts=None, engine_md5=None, base_fp=None):
        cid = self.combo_id(delta)
        row = self.db.execute("SELECT n, delta, best_delta FROM evidence WHERE board_id=? AND combo_id=? AND source=? AND window_days=?", (board_id, cid, source, window_days)).fetchone()
        if row:
            n, d0, b0 = row
            nd = (d0 * n + float(delta_val)) / (n + 1)
            self.db.execute("UPDATE evidence SET n=?, delta=?, best_delta=?, gain=COALESCE(?, gain), trades=COALESCE(?, trades), tim=COALESCE(?, tim), dd=COALESCE(?, dd), recent_delta=COALESCE(?, recent_delta), ts=? WHERE board_id=? AND combo_id=? AND source=? AND window_days=?", (n + 1, nd, max(b0, float(delta_val)), gain, trades, tim, dd, recent_delta, ts or _utc(), board_id, cid, source, window_days))
            return
        self.db.execute("INSERT INTO evidence VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (board_id, cid, source, origin, symptom, window_days, 1, float(delta_val), float(delta_val), gain, base_gain, trades, tim, dd, wr, None if valid is None else int(bool(valid)), recent_delta, ts or _utc(), engine_md5, base_fp))

    def commit(self):
        self.db.commit()

    def add_regime_attr(self, board_id, delta, attr, window_days=365):
        """per-regime P&L/DD of a combo vs its base on the long window (regime fixes reusable across sym_sides)."""
        cid = self.combo_id(delta)
        for reg, a in (attr or {}).items():
            self.db.execute("INSERT OR REPLACE INTO regime_evidence VALUES (?,?,?,?,?,?,?,?,?,?)", (board_id, cid, reg, a.get("n_seg"), a.get("d_pnl"), a.get("d_dd"), a.get("base_pnl"), a.get("cand_pnl"), window_days, _utc()))

    def regime_fixes(self, cat_side, regime, limit=400):
        """combos that improved P&L (or cut DD) inside segments of this regime label on other boards of the cat_side."""
        cur = self.db.execute("SELECT c.combo, AVG(r.d_pnl), AVG(r.d_dd), COUNT(*) FROM regime_evidence r JOIN combos c ON c.combo_id=r.combo_id JOIN boards b ON b.board_id=r.board_id WHERE b.cat_side=? AND r.regime=? GROUP BY r.combo_id HAVING AVG(r.d_pnl) > 1e-9 OR AVG(r.d_dd) < -1e-9 ORDER BY COUNT(*) DESC, AVG(r.d_pnl) DESC LIMIT ?", (cat_side, regime, limit))
        return [{"combo": json.loads(r[0]), "d_pnl": r[1], "d_dd": r[2], "n": r[3]} for r in cur.fetchall()]

    def boards_of(self, cat_side):
        cur = self.db.execute(f"SELECT board_id, sym_side, ts_last, {','.join(FEATS)}, symptoms FROM boards WHERE cat_side=?", (cat_side,))
        return [dict(zip(("board_id", "sym_side", "ts_last") + FEATS + ("symptoms",), r)) for r in cur.fetchall()]

    def neighbours(self, cat_side, feats, k=25, exclude_sym=None):
        """nearest boards of the cat_side on z-scored regime/result features (nan-aware), same-sym excluded on request."""
        bs = self.boards_of(cat_side)
        if not bs:
            return []
        stats = {}
        for f in FEATS:
            xs = [_ftr(b[f], f) for b in bs if b[f] is not None]
            if len(xs) >= 3:
                mu = sum(xs) / len(xs)
                sd = (sum((x - mu) ** 2 for x in xs) / len(xs)) ** 0.5 or 1.0
                stats[f] = (mu, sd)
        out = []
        for b in bs:
            if exclude_sym and b["sym_side"] == exclude_sym:
                continue
            d2, n = 0.0, 0
            for f, (mu, sd) in stats.items():
                a, c = feats.get(f), b[f]
                if a is None or c is None:
                    continue
                d2 += ((_ftr(a, f) - _ftr(c, f)) / sd) ** 2
                n += 1
            if n >= 3:
                out.append((math.sqrt(d2 / n), b))
        out.sort(key=lambda t: t[0])
        return out[:k]

    def positives(self, board_ids, min_delta=1e-9, limit=4000):
        if not board_ids:
            return []
        q = ",".join("?" * len(board_ids))
        cur = self.db.execute(f"SELECT e.board_id, c.combo, e.best_delta, e.delta, e.symptom, e.source, e.recent_delta FROM evidence e JOIN combos c ON c.combo_id=e.combo_id WHERE e.board_id IN ({q}) AND e.best_delta > ? AND (e.valid IS NULL OR e.valid=1) ORDER BY e.best_delta DESC LIMIT ?", list(board_ids) + [min_delta, limit])
        return [{"board_id": r[0], "combo": json.loads(r[1]), "best_delta": r[2], "delta": r[3], "symptom": r[4], "source": r[5], "recent_delta": r[6]} for r in cur.fetchall()]

    def symptom_fixes(self, cat_side, symptom, limit=600):
        cur = self.db.execute("SELECT c.combo, AVG(e.best_delta), COUNT(*) FROM evidence e JOIN combos c ON c.combo_id=e.combo_id JOIN boards b ON b.board_id=e.board_id WHERE b.cat_side=? AND e.best_delta > 1e-9 AND (e.symptom=? OR b.symptoms LIKE ?) GROUP BY e.combo_id ORDER BY COUNT(*) DESC, AVG(e.best_delta) DESC LIMIT ?", (cat_side, symptom, f'%"{symptom}"%', limit))
        return [{"combo": json.loads(r[0]), "mean_delta": r[1], "n": r[2]} for r in cur.fetchall()]

    def structure_stats(self, cat_side, limit=200):
        """positive-rate by TF structure (htf, ltf_confirm) of the cat_side — generalises across key names."""
        cur = self.db.execute("SELECT c.htf, c.ltf_confirm, COUNT(*), AVG(e.best_delta) FROM evidence e JOIN combos c ON c.combo_id=e.combo_id JOIN boards b ON b.board_id=e.board_id WHERE b.cat_side=? AND e.best_delta > 1e-9 GROUP BY c.htf, c.ltf_confirm ORDER BY COUNT(*) DESC LIMIT ?", (cat_side, limit))
        return {(r[0], r[1]): (r[2], r[3]) for r in cur.fetchall()}

    def info(self):
        q = lambda s: self.db.execute(s).fetchone()[0]  # noqa: E731
        out = {"path": str(self.path), "bytes": self.path.stat().st_size if self.path.exists() else 0, "boards": q("SELECT COUNT(*) FROM boards"), "combos": q("SELECT COUNT(*) FROM combos"), "evidence": q("SELECT COUNT(*) FROM evidence"), "engine_evals": q("SELECT COALESCE(SUM(n),0) FROM evidence WHERE source='v17_combo_engine'")}
        out["by_source"] = dict(self.db.execute("SELECT source, COUNT(*) FROM evidence GROUP BY source").fetchall())
        out["by_cat_side"] = dict(self.db.execute("SELECT cat_side, COUNT(*) FROM boards GROUP BY cat_side").fetchall())
        out["boards_with_features"] = q("SELECT COUNT(*) FROM boards WHERE vol_bar_pct IS NOT NULL")
        return out


def TAB_CATEGORY_GUESS(key):  # noqa: N802
    k = str(key).upper()
    if k.startswith(("AUGMENT", "REENTRY_AUGMENT")):
        return "augment"
    if k.startswith("REENTRY"):
        return "reentry"
    if k.startswith(("EXIT", "REDUCE")) or "_EXIT" in k:
        return "exit"
    if is_filter_key(k):
        return "filter"
    return "entry"


def _ftr(x, f):
    x = float(x)
    if f == "chop":
        return math.log1p(min(abs(x), 999.0))
    if f == "trades":
        return math.log1p(max(0.0, x))
    return x


# -------------------------------------------------------------------------------------------------------- regime features
def seg_features(close, is_long):
    """trend / volatility / chop of a close segment (same definitions as the pilot's _v17_seg_feats)."""
    cl = [float(x) for x in close if x is not None and math.isfinite(float(x)) and float(x) > 0]
    if len(cl) < 8:
        return {}
    rets = [(cl[i] / cl[i - 1] - 1.0) * 100.0 for i in range(1, len(cl))]
    net = (cl[-1] / cl[0] - 1.0) * 100.0
    mu = sum(rets) / len(rets)
    vol = (sum((r - mu) ** 2 for r in rets) / len(rets)) ** 0.5
    chop = (sum(abs(r) for r in rets) / abs(net)) if abs(net) > 1e-9 else 999.0
    rng = (max(cl) - min(cl)) / (sum(cl) / len(cl)) * 100.0
    return {"bars": len(cl), "net_pct": round(net, 4), "side_net_pct": round(net if is_long else -net, 4), "vol_bar_pct": round(vol, 5), "chop": round(chop, 3), "range_pct": round(rng, 4)}


def classify_regime(f, ref=None):
    if not f:
        return "UNKNOWN"
    hv = bool(ref) and f.get("vol_bar_pct", 0) > 1.5 * max(1e-9, ref.get("vol_bar_pct", 0))
    sq = bool(ref) and f.get("vol_bar_pct", 0) < 0.6 * max(1e-9, ref.get("vol_bar_pct", 0))
    lab = ("TREND_UP" if f["net_pct"] > 0 else "TREND_DOWN") if f.get("chop", 999) < 12 and abs(f.get("net_pct", 0)) > 3 * f.get("vol_bar_pct", 0) else "RANGE"
    return lab + ("+HIGH_VOL" if hv else ("+SQUEEZE" if sq else ""))


def npz_window_close(sym, ts_last, crypto, days=30):
    """30D window close (15m-compacted, last-closed-bar, no look-ahead) ending at ts_last from the host NPZ; None if absent."""
    import numpy as np
    from tools.opt import evaluate_v12 as E
    npz_sym, _tok = E.resolve_tokenised(sym)
    for pfx in ("backtest_v8", "backtest_v7"):
        p = ROOT / pfx / "indicators" / f"{npz_sym}.npz"
        if p.exists():
            break
    else:
        return None
    try:
        z = np.load(str(p), allow_pickle=True)
        ts = np.asarray(z["timestamps"], dtype="float64")
        sec = ts / (1000.0 if ts[-1] > 1e11 else 1.0)
        right = int(np.searchsorted(sec, float(ts_last) + 1.0, side="left")) if ts_last else len(sec)
        if right < 50:
            return None
        d = {"timestamps": ts[:right], "close": np.asarray(z["close"], dtype="float64")[:right]}
        if "timestamp_15m" in z.files:
            d["timestamp_15m"] = np.asarray(z["timestamp_15m"], dtype="float64")[:right]
        sl, _pol = E._exact_30d_slice(d, crypto, days)
        if "timestamp_15m" in sl:
            try:
                sl = E._compact_to_15m(sl)
            except ValueError:
                pass
        return [float(x) for x in sl["close"]]
    except Exception:
        return None


# --------------------------------------------------------------------------------------------------- building the library
def _num(x):
    try:
        f = float(x)
        return f if math.isfinite(f) else None
    except (TypeError, ValueError):
        return None


def _board_metrics(d):
    gain = _num(d.get("final_gain_fresh_vec")) if d.get("final_gain_fresh_vec") is not None else _num(d.get("cumulative_gain"))
    tim = trades = dd = None
    for src in ((d.get("diagnose_repair") or {}).get("after"), ((d.get("redo_history") or [{}])[-1] or {}).get("result") if d.get("redo_history") else None, d.get("result")):
        if isinstance(src, dict):
            tim = tim if tim is not None else _num(src.get("tim") if src.get("tim") is not None else src.get("tim_pct"))
            trades = trades if trades is not None else _num(src.get("trades"))
            dd = dd if dd is not None else _num(src.get("dd") if src.get("dd") is not None else src.get("max_dd_pct"))
    if trades is None and isinstance(d.get("tier"), dict):
        trades = _num(d["tier"].get("trades"))
    return gain, tim, trades, dd


def _parse_done_key(k):
    """'TAB!row:SWITCH=cand' -> (tab, switch, cand) or None."""
    try:
        tab, rest = str(k).split("!", 1)
        rest = rest.split(":", 1)[1] if ":" in rest else rest
        sw, _, cand = rest.partition("=")
        if not sw or not _:
            return None
        return tab, sw.strip(), cand.strip()
    except ValueError:
        return None


def _typed_delta(parts, venue, P=None):
    out = {}
    comp = getattr(P, "COMPOSITE_SWITCHES", {}) if P else {}
    for k, v in parts:
        k = str(k).strip()
        if k in comp:
            try:
                out.update(comp[k](v))
                continue
            except Exception:
                return None
        ok, val = coerce_value(k, v, venue)
        if not ok:
            return None
        out[k] = val
    return out or None


_DL_POS = re.compile(rb'"delta": (?!-|0\.0,|null|0,)')
_LABEL_SKIP = re.compile(r"EXEMPT|PROXY|STRIP|ABL_|\+exempt|\+proxy|MODE=|ABLATION", re.I)


def build_library(lib, progress_globs, delta_globs, log=print, max_files=None):
    """Ingest past fleet data: (a) progress done rows (naked + every yellow) with delta>0 per board, (b) delta-log positive rows
    of non-row phases (V17/FOCUS/SURGEON/ENDGAME/FINAL_RECHECK...), (c) v17_priors synergy/blocks + (d) fleet endgame knowledge
    stored as cat_side pseudo-boards. Boards get regime features from the host NPZ at their ts_last (no look-ahead)."""
    P = _Pilot.m()
    t0 = time.time()
    files = sorted({f for g in progress_globs for f in glob.glob(os.path.expanduser(g), recursive=True)})
    if max_files:
        files = files[: int(max_files)]
    n_ev = n_b = 0
    run_board = {}
    close_cache = {}
    for i, f in enumerate(files):
        try:
            d = json.loads(Path(f).read_text())
        except (OSError, ValueError):
            continue
        ss = str(d.get("symside") or Path(f).name.replace("_v14_progress.json", ""))
        if not ss.endswith(("_LONG", "_SHORT")):
            continue
        cs = cat_side_of(ss)
        venue = venue_of(cs)
        nid = d.get("npz_id") if isinstance(d.get("npz_id"), dict) else {}
        ts_last = _num(nid.get("ts_last"))
        ts_last = ts_last / 1000.0 if ts_last and ts_last > 1e12 else ts_last
        gain, tim, trades, dd = _board_metrics(d)
        bkey = f"{ss}|{nid.get('md5') or ts_last}|{None if gain is None else round(gain, 4)}"
        run = Path(f).parent.parent.name
        feats = {}
        sym = ss.rsplit("_", 1)[0]
        ck = (sym, ts_last)
        if ck not in close_cache:
            close_cache[ck] = npz_window_close(sym, ts_last, cs.startswith("CRYPTO"))
            if len(close_cache) > 4000:
                close_cache.clear()
        cl = close_cache.get(ck)
        if cl:
            feats = seg_features(cl, ss.endswith("_LONG"))
        syms = []
        if tim is not None and tim < 20:
            syms.append("tim_low")
        if tim is not None and tim > 80:
            syms.append("tim_high")
        if gain is not None and _num(d.get("bh")) is not None and gain < float(d["bh"]):
            syms.append("below_bh")
        bid = lib.board(bkey, sym_side=ss, cat_side=cs, src=f, run=run, ts_last=ts_last, gain=gain, bh=_num(d.get("bh")), tim=tim, dd=dd, trades=trades, side_net_pct=feats.get("side_net_pct"), vol_bar_pct=feats.get("vol_bar_pct"), chop=feats.get("chop"), range_pct=feats.get("range_pct"), regime=classify_regime(feats), symptoms=syms, engine_md5=d.get("result_engine_md5"))
        run_board[(run, ss)] = bid
        n_b += 1
        for k, e in (d.get("done") or {}).items():
            if not isinstance(e, dict):
                continue
            pk = _parse_done_key(k)
            if not pk:
                continue
            tab, sw, cand = pk
            base_parts = [(sw, cand)]
            nd = _num(e.get("naked_delta"))
            if nd is not None and nd > 1e-9:
                dl = _typed_delta(base_parts, venue, P)
                if dl:
                    lib.add_evidence(bid, dl, "progress_done", nd, origin=tab, trades=_num(e.get("trades")), ts=d.get("result_done_utc"), engine_md5=d.get("result_engine_md5"))
                    n_ev += 1
            for lab, yd in (e.get("yellows") or {}).items():
                yd = _num(yd)
                if yd is None or yd <= 1e-9 or "=" not in str(lab) or _LABEL_SKIP.search(str(lab)):
                    continue
                parts = base_parts + [tuple(p.split("=", 1)) for p in str(lab).split(" + ") if "=" in p]
                dl = _typed_delta(parts, venue, P)
                if dl:
                    lib.add_evidence(bid, dl, "progress_done", yd, origin=tab, ts=d.get("result_done_utc"), engine_md5=d.get("result_engine_md5"))
                    n_ev += 1
        if i % 25 == 0:
            lib.commit()
            log(f"[build] progress {i + 1}/{len(files)} boards={n_b} evidence_rows={n_ev} {time.time() - t0:.0f}s")
    lib.commit()
    dfiles = sorted({f for g in delta_globs for f in glob.glob(os.path.expanduser(g), recursive=True)})
    n_dl = 0
    for j, f in enumerate(dfiles):
        run = Path(f).parent.parent.parent.name
        ss = re.sub(r"_(jump|sequential|hustle|seq)$", "", Path(f).stem)
        if not ss.endswith(("_LONG", "_SHORT")):
            continue
        cs = cat_side_of(ss)
        venue = venue_of(cs)
        bid = run_board.get((run, ss))
        try:
            fh = open(f, "rb")
        except OSError:
            continue
        with fh:
            for line in fh:
                if not _DL_POS.search(line):
                    continue
                try:
                    r = json.loads(line)
                except ValueError:
                    continue
                dv = _num(r.get("delta"))
                if dv is None or dv <= 1e-9 or not r.get("valid") or r.get("err"):
                    continue
                if str(r.get("sheet") or "") in SWITCH_SHEETS and bid is not None:
                    continue  # covered by the progress done rows of the same run
                sw, cand, lab = str(r.get("switch") or "").strip(), r.get("cand"), str(r.get("label") or "")
                if not sw or sw.upper().startswith(("VERIFY", "BASE")) or _LABEL_SKIP.search(lab):
                    continue
                symptom = ""
                m = re.match(r"SURGEON:([a-z_]+):", lab)
                if m:
                    symptom = m.group(1)
                    lab = lab.split(":", 2)[2]
                parts = [(sw, cand)]
                for p in re.split(r" \+ ", lab):
                    if "=" in p and not p.startswith(("NAKED", "naked")):
                        a, _, b = p.partition("=")
                        if a.strip() and a.strip() != sw and re.match(r"^[A-Z0-9_]+$", a.strip()):
                            parts.append((a.strip(), b.strip()))
                dl = _typed_delta(parts, venue, P)
                if not dl:
                    continue
                if bid is None:
                    tsv = r.get("ts")
                    try:
                        tse = datetime.datetime.strptime(str(tsv)[:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=datetime.timezone.utc).timestamp()
                    except ValueError:
                        tse = None
                    bid = lib.board(f"{ss}|{run}|deltalog", sym_side=ss, cat_side=cs, src=f, run=run, ts_last=tse)
                    run_board[(run, ss)] = bid
                lib.add_evidence(bid, dl, "delta_log", dv, origin=str(r.get("sheet") or ""), symptom=symptom, gain=_num(r.get("gain_pct")), base_gain=_num(r.get("cum_before")), trades=_num(r.get("trades")), tim=_num(r.get("tim")), valid=True, ts=r.get("ts"))
                n_dl += 1
        if j % 10 == 0:
            lib.commit()
            log(f"[build] delta logs {j + 1}/{len(dfiles)} rows={n_dl} {time.time() - t0:.0f}s")
    lib.commit()
    n_pr = ingest_priors(lib)
    lib.db.execute("INSERT OR REPLACE INTO meta VALUES ('built_utc', ?)", (_utc(),))
    lib.commit()
    log(f"[build] done boards={n_b} progress_evidence={n_ev} delta_log={n_dl} priors={n_pr} {time.time() - t0:.0f}s")
    return {"boards": n_b, "progress_evidence": n_ev, "delta_log": n_dl, "priors": n_pr, "secs": round(time.time() - t0, 1)}


def ingest_priors(lib):
    """v17_priors synergy pairs + block->exemption pairs and fleet endgame knowledge as cat_side pseudo-boards (sym_side '*')."""
    n = 0
    try:
        from tools import v17_priors as VP
    except Exception:
        VP = None
    try:
        from tools.v15_endgame_knowledge_merge import load_ev
    except Exception:
        load_ev = None
    for cs in ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"):
        venue = venue_of(cs)
        bid = lib.board(f"*{cs}|priors", sym_side="*", cat_side=cs, src="v17_priors+endgame")
        p = VP.load(cs) if VP else {}
        for s in (p.get("synergy") or []):
            if float(s.get("mean_joint") or 0) <= 0:
                continue
            parts = [tuple(str(x).split("=", 1)) for x in (s.get("a"), s.get("b")) if "=" in str(x)]
            dl = _typed_delta(parts, venue, _Pilot.m())
            if dl:
                lib.add_evidence(bid, dl, "prior_synergy", float(s["mean_joint"]), origin="synergy")
                n += 1
        for b in (p.get("blocks") or []):
            ex = b.get("exemption_override")
            if not isinstance(ex, dict) or "=" not in str(b.get("switch")):
                continue
            dl = _typed_delta([tuple(str(b["switch"]).split("=", 1))], venue, _Pilot.m())
            if not dl:
                continue
            dl = dict(dl)
            for k, v in ex.items():
                dl[k] = v
            lib.add_evidence(bid, dl, "prior_block_exempt", float(b.get("mean_unlock") or 0.0), origin="blocks")
            n += 1
        if load_ev:
            try:
                ev = load_ev(cs) or {}
            except Exception:
                ev = {}
            for key, e in ev.items():
                v = e[1] * e[2] if isinstance(e, (list, tuple)) and len(e) >= 3 else (e if isinstance(e, (int, float)) else None)
                if v is None or v <= 0 or "=" not in str(key):
                    continue
                dl = _typed_delta([tuple(str(key).split("=", 1))], venue, _Pilot.m())
                if dl:
                    lib.add_evidence(bid, dl, "endgame_knowledge", float(v), origin="endgame")
                    n += 1
    lib.commit()
    return n


def merge_library(into, srcs, log=print):
    """Union of other host libraries: boards by bkey, combos by ckey, evidence max-merged."""
    lib = ComboLibrary(into)
    for s in srcs:
        o = sqlite3.connect(str(s))
        bmap = {}
        cols = [r[1] for r in o.execute("PRAGMA table_info(boards)").fetchall()]
        for row in o.execute("SELECT * FROM boards").fetchall():
            rec = dict(zip(cols, row))
            old = rec.pop("board_id")
            bkey = rec.pop("bkey")
            rec.pop("added_utc", None)
            bmap[old] = lib.board(bkey, **{k: (json.loads(v) if k == "symptoms" and isinstance(v, str) and v.startswith("[") else v) for k, v in rec.items()})
        cmap = {}
        for cid, combo in o.execute("SELECT combo_id, combo FROM combos").fetchall():
            cmap[cid] = lib.combo_id(json.loads(combo))
        ecols = [r[1] for r in o.execute("PRAGMA table_info(evidence)").fetchall()]
        n = 0
        for row in o.execute("SELECT * FROM evidence").fetchall():
            e = dict(zip(ecols, row))
            e["board_id"], e["combo_id"] = bmap.get(e["board_id"]), cmap.get(e["combo_id"])
            if e["board_id"] is None or e["combo_id"] is None:
                continue
            ex = lib.db.execute("SELECT n, best_delta FROM evidence WHERE board_id=? AND combo_id=? AND source=? AND window_days=?", (e["board_id"], e["combo_id"], e["source"], e["window_days"])).fetchone()
            if ex:
                if e["best_delta"] is not None and e["best_delta"] > (ex[1] or -1e18):
                    lib.db.execute("UPDATE evidence SET best_delta=? WHERE board_id=? AND combo_id=? AND source=? AND window_days=?", (e["best_delta"], e["board_id"], e["combo_id"], e["source"], e["window_days"]))
                continue
            lib.db.execute(f"INSERT INTO evidence ({','.join(ecols)}) VALUES ({','.join('?' * len(ecols))})", [e[c] for c in ecols])
            n += 1
        lib.commit()
        o.close()
        log(f"[merge] {s}: +{n} evidence rows")
    info = lib.info()
    lib.close()
    return info


# ----------------------------------------------------------------------------------------------------------- generation
def _cand(delta, origin, why, score, cats=()):
    return {"delta": delta, "origin": origin, "why": why, "score": float(score), "cats": tuple(cats)}


def generate(space, sym_side, base, symptoms, feats, lib=None, rng=None, max_n=40000, log=None):
    """Ranked, typed, coherent, deduplicated candidate combos (origin-tagged) for one board."""
    rng = rng or random.Random(0)
    cs = space.cat_side
    venue = space.venue
    out, seen = [], set()
    stats = {"generated": 0, "refused_type": 0, "incoherent": 0, "blocked": 0, "identity": 0, "by_origin": {}}
    sym_size = {k: float((v or {}).get("size") or 0.0) for k, v in (symptoms or {}).items()}
    hot_cats = {}
    for s, sz in sym_size.items():
        for c in SYMPTOM_CATS.get(s, ()):
            hot_cats[c] = hot_cats.get(c, 0.0) + max(0.0, sz)
    try:
        from tools import v17_priors as VP
        pri = VP.load(cs)
    except Exception:
        VP, pri = None, {}
    single = pri.get("single") or {}
    dep = {d.get("key") for d in (pri.get("deprioritise") or [])}

    def prior_ev(k, v):
        e = single.get(f"{k}={v}") or (VP.single(cs, f"{k}={v}") if VP else None)
        if not isinstance(e, dict) or not e.get("n"):
            return 0.0
        return float(e.get("mean") or 0.0) * float(e.get("pos") or 0) / max(1.0, float(e.get("n")))

    def add(delta, origin, why, score):
        stats["generated"] += 1
        if not delta:
            return
        typed = {}
        for k, v in delta.items():
            if k == "FILTER_EXEMPTIONS":
                ok, val = coerce_value(k, v, venue)
                if not ok or not isinstance(val, dict):
                    stats["refused_type"] += 1
                    return
                typed[k] = val
                continue
            ok, val = coerce_value(k, v, venue)
            if not ok:
                stats["refused_type"] += 1
                return
            typed[k] = val
        eff = effective_delta(base, typed)
        if not eff:
            stats["identity"] += 1
            return
        if len([k for k in eff if k != "FILTER_EXEMPTIONS"]) > 4:
            return
        ck = combo_key(eff)
        if ck in seen:
            return
        if blocked_reason(eff):
            stats["blocked"] += 1
            return
        full = apply_delta(base, eff)
        inc = coherence_reason(eff, full, venue=venue)
        if inc:
            stats["incoherent"] += 1
            return
        if any(f"{k}={v}" in dep for k, v in eff.items() if not isinstance(v, dict)) and origin not in ("explore", "lib_same_sym"):
            score -= 1.0
        seen.add(ck)
        cats = sorted({c for c in (space.category(k) for k in eff) if c})
        sc = score + 0.15 * sum(hot_cats.get(c, 0.0) for c in cats) / max(1.0, sum(hot_cats.values()) or 1.0)
        warn = coherence_warnings(eff, full, venue) if any(k in eff for k in ("DC_BREAKOUT_TF_EXPANDED", "BB_BREAKOUT_ENTRY_TF", "DC_BREAKOUT_ENTRY_ENABLED")) else []
        out.append(_cand(eff, origin, why, sc - 0.3 * len(warn), cats))
        stats["by_origin"][origin] = stats["by_origin"].get(origin, 0) + 1
        if warn and origin != "tf_coherent":
            tf_fix.append((eff, warn, sc))

    tf_fix = []
    # (1) library: same sym_side history, regime neighbours, symptom fixes
    if lib is not None:
        same = [b["board_id"] for b in lib.boards_of(cs) if b["sym_side"] == sym_side]
        for r in lib.positives(same, limit=3000):
            add(r["combo"], "lib_same_sym", f"positive on {sym_side} before (best {r['best_delta']:+.2f}, {r['source']})", 3.0 + min(5.0, r["best_delta"]))
        nn = lib.neighbours(cs, feats, k=int(_env_f("V17_COMBO_NN_K", 30)), exclude_sym=sym_side)
        wmap = {b["board_id"]: (dist, b["sym_side"]) for dist, b in nn}
        for r in lib.positives(list(wmap), limit=6000):
            dist, nsym = wmap[r["board_id"]]
            add(r["combo"], "lib_neighbour", f"positive on regime-neighbour {nsym} (dist {dist:.2f}, {r['best_delta']:+.2f})", 2.0 + min(5.0, r["best_delta"]) / (1.0 + dist))
        for s, sz in sorted(sym_size.items(), key=lambda t: -t[1]):
            if sz <= 0:
                continue
            for r in lib.symptom_fixes(cs, s, limit=400):
                add(r["combo"], "lib_symptom", f"fixed {s} on {r['n']} boards (mean {r['mean_delta']:+.2f})", 1.5 + 0.2 * r["n"] + min(3.0, r["mean_delta"]))
        rl = feats.get("regime_now")
        if rl and rl != "UNKNOWN":
            for r in lib.regime_fixes(cs, rl, limit=400):
                add(r["combo"], "lib_regime", f"improved {rl} segments on {r['n']} boards (d_pnl {r['d_pnl']:+.2f}, d_dd {r['d_dd']:+.2f})", 1.5 + 0.2 * r["n"] + min(3.0, max(0.0, r["d_pnl"] or 0.0)))
        star = [b["board_id"] for b in lib.boards_of(cs) if b["sym_side"] == "*"]
        for r in lib.positives(star, limit=3000):
            o = {"prior_synergy": "prior_synergy", "prior_block_exempt": "prior_block_exempt"}.get(r["source"], "lib_neighbour")
            add(r["combo"], o, f"{r['source']} (cat_side {cs}, {r['best_delta']:+.2f})", 1.0 + min(3.0, r["best_delta"]))
    # (2) symptom paths: path switches of the symptom's categories (+ loosened / tightened filters)
    singles = []
    for k in space.path_keys():
        cat = space.category(k)
        for v in space.values.get(k, []):
            if same_val(base.get(k), v):
                continue
            ev = prior_ev(k, v)
            hot = hot_cats.get(cat, 0.0)
            singles.append((ev + 0.05 * hot, k, v, cat))
    singles.sort(key=lambda t: -t[0])
    for sc, k, v, cat in singles:
        o = "symptom_path" if hot_cats.get(cat, 0) > 0 else "symptom_path"
        add(space.to_delta(k, v), o, f"path {cat}:{k}={v} (prior ev {sc:+.2f})", 1.0 + sc)
    for f in space.filter_keys():
        cur = base.get(f)
        for v in space.values.get(f, []):
            off = (v is False) or (str(v).upper() == "OFF")
            want = "loosen" if off else "tighten"
            sc = prior_ev(f, v)
            if any(s in LOOSEN and sym_size.get(s, 0) > 0 for s in sym_size) and off:
                sc += 0.5
            if any(s in TIGHTEN and sym_size.get(s, 0) > 0 for s in sym_size) and not off:
                sc += 0.5
            if not same_val(cur, v):
                add({f: v}, "filter_path", f"filter {want} {f}={v}", 0.8 + sc)
    # (3) path + specific filter applied/removed (row x yellow cells of the path's own tab) and path + exemption
    top_paths = [(k, v, cat) for _sc, k, v, cat in singles[: int(_env_f("V17_COMBO_TOP_PATHS", 400))]]
    for k, v, cat in top_paths:
        for tab in sorted(space.key_tabs.get(k, ())):
            for f, fv in space.filters_by_tab.get(tab, [])[:: max(1, len(space.filters_by_tab.get(tab, [])) // 40)]:
                if f == k:
                    continue
                d = dict(space.to_delta(k, v))
                d[f] = fv
                add(d, "filter_path", f"{k}={v} with {f}={fv}", 0.6 + prior_ev(k, v) + 0.5 * prior_ev(f, fv))
        exm = space.exemptions.get(k) or {}
        if cat in ("entry", "reentry", "augment") and isinstance(exm, dict):
            for fk, ov in list(exm.items())[:8]:
                if not isinstance(ov, dict):
                    continue
                d = dict(space.to_delta(k, v))
                for ek, evv in ov.items():
                    d[ek] = merge_exemptions(d.get(ek), evv) if ek == "FILTER_EXEMPTIONS" else evv
                add(d, "exempt_path", f"{k}={v} exempt from {fk}", 0.9 + prior_ev(k, v))
    # (4) numeric neighbourhoods of the numeric keys of the current set
    for k in sorted(space.values):
        b = base.get(k)
        for nv in space.numeric_neighbours(k, b):
            add({k: nv}, "numeric", f"{k} {b} -> {nv} (x{(float(nv) / float(b)):.2f})", 0.5 + prior_ev(k, nv))
    # (5) AUGMENT_INTO_REENTRY family (path = ENABLED + up to 2 params)
    if space.air_ok:
        add({"AUGMENT_INTO_REENTRY_ENABLED": True}, "air", "AUGMENT_INTO_REENTRY on (defaults)", 1.0 + 0.1 * hot_cats.get("augment", 0.0))
        keys = list(AIR_FAMILY)
        for i, a in enumerate(keys):
            for va in AIR_FAMILY[a]:
                add({"AUGMENT_INTO_REENTRY_ENABLED": True, a: va}, "air", f"AIR on + {a}={va}", 0.9)
                for b in keys[i + 1:]:
                    for vb in AIR_FAMILY[b]:
                        add({"AUGMENT_INTO_REENTRY_ENABLED": True, a: va, b: vb}, "air", f"AIR on + {a}={va} + {b}={vb}", 0.7)
    # (6) cross-category pairs / triples of the top singles
    pool = [c for c in out if c["origin"] in ("lib_same_sym", "lib_neighbour", "symptom_path", "lib_symptom") and len(c["delta"]) == 1][:300]
    pool.sort(key=lambda c: -c["score"])
    pool = pool[:160]
    for i, a in enumerate(pool):
        for b in pool[i + 1:]:
            if set(a["delta"]) & set(b["delta"]) or (a["cats"] and a["cats"] == b["cats"]):
                continue
            d = dict(a["delta"])
            d.update(b["delta"])
            add(d, "pair", f"pair {combo_key(a['delta'])} + {combo_key(b['delta'])}", 0.5 * (a["score"] + b["score"]))
        if len(out) >= max_n:
            break
    # (6b) TF-coherent variants (USER: HTF = permission/validity, LTF = timing, one open per HTF instance): every HTF state
    # trigger combo flagged C1/C2 is re-offered WITH a consumption rule (cooldown <= the HTF span, template values) and/or an
    # explicit 15m confirm (C2 list, template values).
    cdk = "COOLDOWN_BARS_TRADIER" if venue == "STOCKS" else "COOLDOWN_BARS"
    for eff, warn, sc in tf_fix[: int(_env_f("V17_COMBO_TF_FIX_N", 300))]:
        trig = htf_state_triggers(apply_delta(base, eff))
        span = max((_span(t, venue) for _k, t in trig), default=16)
        cds = [v for v in space.values.get(cdk, []) if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 < float(v) <= span]
        confs = [(f, v) for f in C2_LTF_CONFIRMS for v in space.values.get(f, []) if _tf_token(v) == "15m"][:3]
        if "C1_UNBOUNDED_REOPEN" in warn:
            for v in cds[:3]:
                add(dict(eff, **{cdk: v}), "tf_coherent", f"{combo_key(eff)[:60]} + consumption {cdk}={v} (<= {span}-bar HTF span)", sc + 0.4)
        if "C2_NO_LTF_CONFIRM" in warn:
            for f, v in confs:
                d = dict(eff, **{f: v})
                add(d, "tf_coherent", f"{combo_key(eff)[:60]} + LTF confirm {f}=15m", sc + 0.3)
                if cds and "C1_UNBOUNDED_REOPEN" in warn:
                    add(dict(d, **{cdk: cds[0]}), "tf_coherent", f"{combo_key(eff)[:60]} + 15m confirm + {cdk}={cds[0]}", sc + 0.5)
    # (7) exploration: random valid 1-3 changes from different categories (+ optional filter, + optional exemption)
    keys = space.path_keys()
    fkeys = space.filter_keys()
    n_explore = int(max(200, _env_f("V17_COMBO_EXPLORE", 0.15) * max_n))
    tries = 0
    while stats["by_origin"].get("explore", 0) < n_explore and tries < n_explore * 6 and keys:
        tries += 1
        d = {}
        for _ in range(rng.choice((1, 1, 2, 2, 3))):
            k = rng.choice(keys)
            vals = space.values.get(k) or []
            if vals:
                d.update(space.to_delta(k, rng.choice(vals)))
        if fkeys and rng.random() < 0.5:
            f = rng.choice(fkeys)
            if space.values.get(f):
                d[f] = rng.choice(space.values[f])
        add(d, "explore", "exploration (random valid combo)", rng.random() * 0.5)
    stats["n"] = len(out)
    if log:
        log(f"[combo-gen] {sym_side}: {len(out)} combos {stats['by_origin']} refused_type={stats['refused_type']} incoherent={stats['incoherent']} blocked={stats['blocked']} identity={stats['identity']}")
    return out, stats


def schedule(cands, explore_frac=None, quota=None):
    """Interleave origins by quota (exploration = a fixed share), each origin in score order."""
    quota = dict(quota or ORIGIN_QUOTA)
    if explore_frac is not None:
        quota["explore"] = float(explore_frac)
    by = {}
    for c in cands:
        by.setdefault(c["origin"], []).append(c)
    for v in by.values():
        v.sort(key=lambda c: -c["score"])
    credit = {o: 0.0 for o in by}
    idx = {o: 0 for o in by}
    out = []
    total = len(cands)
    while len(out) < total:
        live = [o for o in by if idx[o] < len(by[o])]
        if not live:
            break
        for o in live:
            credit[o] += quota.get(o, 0.03)
        o = max(live, key=lambda x: credit[x])
        credit[o] -= 1.0
        out.append(by[o][idx[o]])
        idx[o] += 1
    return out


# ----------------------------------------------------------------------------------------------------------- evaluation
_PREP = {}


def _compact(res, ledger=False):
    if not isinstance(res, dict):
        return None
    out = {"gain": res.get("gain_pct"), "trades": res.get("trades"), "tim": res.get("tim_pct"), "dd": res.get("max_dd_pct"), "wr": res.get("wr_pct"), "valid": bool(res.get("valid")), "reason": str(res.get("invalid_reason") or "")[:120], "fp": (res.get("behavior_fingerprint") or "")[:16], "bh": res.get("bh_pct"), "pool_sharpe": res.get("pool_sharpe")}
    if ledger:
        out["trades_list"] = [(int(t.get("bar_entry", -1)), int(t.get("bar_exit", -1)), round(float(t.get("pnl_pct") or 0.0), 4), round(float(t.get("pnl_dollars") or 0.0), 4)) for t in (res.get("ledger") or []) if isinstance(t, dict) and "pnl_dollars" in t]
    return out


def _worker(chunk):
    from tools.opt.v12_pilot import evaluate_prepared_sanitized as eps
    out = []
    for key, window, ov, ledger in chunk:
        c0, t0 = time.process_time(), time.time()
        try:
            r = _compact(eps(_PREP[window], ov, window, bool(ledger)), ledger)
            err = ""
        except Exception as e:  # one bad eval never poisons siblings
            r, err = None, f"ERR {e}"[:120]
        out.append((key, r, err, time.process_time() - c0, time.time() - t0))
    return out


def _worker_sig_default():
    """Pool initializer: SIGTERM/SIGINT back to SIG_DFL (forked workers must never run a host's shutdown handler on teardown)."""
    import signal
    try:
        signal.signal(signal.SIGTERM, signal.SIG_DFL)
        signal.signal(signal.SIGINT, signal.SIG_DFL)
    except Exception:
        pass


class PoolEvaluator:
    """Fork pool over the prepared 30D + recent-7D slices (inherited copy-on-write, pre-forked single-threaded)."""

    def __init__(self, preps, workers=4, chunk=None):
        import concurrent.futures as cf
        import multiprocessing as mp
        _PREP.clear()
        _PREP.update(preps)
        self.workers = max(1, int(workers))
        self.chunk = int(chunk or os.environ.get("V17_COMBO_CHUNK", "4"))
        self.cpu_s = 0.0
        self.n = 0
        self.ex = None
        if self.workers > 1:
            self.ex = cf.ProcessPoolExecutor(max_workers=self.workers, mp_context=mp.get_context(os.environ.get("V17_COMBO_MP", "fork")), initializer=_worker_sig_default)
            list(self.ex.map(abs, range(self.workers * 4)))

    def __call__(self, items, deadline=None):
        """items: [(key, window, overrides, ledger)] -> {key: (result|None, err)} (keys not reached by the deadline omitted)."""
        import concurrent.futures as cf
        out = {}
        chunks = [items[i:i + self.chunk] for i in range(0, len(items), self.chunk)]
        if self.ex is None:
            for ch in chunks:
                if deadline and time.time() > deadline:
                    break
                for key, r, err, cpu, _w in _worker(ch):
                    out[key] = (r, err)
                    self.cpu_s += cpu
                    self.n += 1
            return out
        pending = set()
        it = iter(chunks)
        for _ in range(self.workers * 2):
            ch = next(it, None)
            if ch is None:
                break
            pending.add(self.ex.submit(_worker, ch))
        while pending:
            tmo = None if not deadline else max(0.1, deadline - time.time() + 30.0)
            done, pending = cf.wait(pending, timeout=tmo, return_when=cf.FIRST_COMPLETED)
            if not done:
                break
            for fu in done:
                try:
                    for key, r, err, cpu, _w in fu.result():
                        out[key] = (r, err)
                        self.cpu_s += cpu
                        self.n += 1
                except Exception:
                    pass
                if not deadline or time.time() < deadline:
                    ch = next(it, None)
                    if ch is not None:
                        pending.add(self.ex.submit(_worker, ch))
        for fu in pending:
            fu.cancel()
        return out

    def close(self):
        if self.ex is not None:
            procs = list((getattr(self.ex, "_processes", None) or {}).values())
            try:
                self.ex.shutdown(wait=False, cancel_futures=True)
            except Exception:
                pass
            for p in procs:
                try:
                    p.terminate()
                except Exception:
                    pass


# ---------------------------------------------------------------------------------------------------------------- search
def auto_symptoms(evaluate, base_ov, prepared, sym_side, base30):
    """pilot-shaped symptoms from one real ledger eval of the base (tools/v15_trade_autopsy, same classes as the surgeon)."""
    from tools import v15_trade_autopsy as TA
    r = evaluate([("ledger", 30, base_ov, True)]).get("ledger", (None, ""))[0] or {}
    rows = [{"be": t[0], "bx": t[1], "pnl": t[2], "pnl_pct": t[2], "entry_reason": "", "exit_reason": "", "entry_price": 0.0, "exit_price": 0.0, "n_aug": 0} for t in (r.get("trades_list") or [])]
    npz = (prepared or {}).get("npz_prepared") or {}
    close = [float(x) for x in (npz.get("close") if npz.get("close") is not None else [])]
    is_long = sym_side.upper().endswith("_LONG")
    sym = {}
    if close and rows:
        base = TA.classify(rows, close, is_long)
        moves = TA.missed_moves(base, close, is_long)
        bad = lambda c: [t for t in base if c in (t.get("cls") or [])]  # noqa: E731
        churn = [t for t in base if int(t["bx"]) - int(t["be"]) <= 1]
        sym = {"missed_moves": {"n": len(moves), "size": round(sum(m["move_pct"] for m in moves), 3)}, "losers": {"n": len(bad("LOSER")), "size": round(-sum(t["pnl_pct"] for t in bad("LOSER")), 3)}, "premature_exits": {"n": len(bad("PREMATURE_EXIT")), "size": round(sum(t.get("post_pct", 0.0) for t in bad("PREMATURE_EXIT")), 3)}, "giveback": {"n": len(bad("GIVEBACK")), "size": round(sum(t.get("mfe_pct", 0.0) - t.get("pnl_pct", 0.0) for t in bad("GIVEBACK")), 3)}, "missed_augment": {"n": len(bad("MISSED_AUGMENT")), "size": round(sum(t.get("mfe_pct", 0.0) for t in bad("MISSED_AUGMENT")), 3)}, "churn": {"n": len(churn), "size": round(-sum(min(0.0, t["pnl_pct"]) for t in churn), 3)}}
    tim, g, bh = base30.get("tim"), base30.get("gain"), base30.get("bh")
    if tim is not None and tim < 20:
        sym["tim_low"] = {"n": 1, "size": round(20 - tim, 3)}
    if tim is not None and tim > 80:
        sym["tim_high"] = {"n": 1, "size": round(tim - 80, 3)}
    if g is not None and bh is not None and g < bh:
        sym["below_bh"] = {"n": 1, "size": round(bh - g, 3)}
    return sym


def _changed_trades(a, b):
    sa, sb = set(map(tuple, a or [])), set(map(tuple, b or []))
    diff = sa ^ sb
    pn = [abs(t[2]) for t in diff]
    tot = sum(pn)
    return len(diff), (max(pn) / tot if tot > 0 else 0.0)


def regime_segments(prepared, is_long, seg_days=None):
    """long-window regime segmentation: consecutive seg_days (default 14) calendar blocks of the prepared (15m-compacted,
    last-closed-bar) NPZ, each labelled with classify_regime(seg features, ref = whole-window features). [(b0, b1, label)]."""
    seg_days = float(seg_days or _env_f("V17_COMBO_SEG_DAYS", 14))
    npz = (prepared or {}).get("npz_prepared") or {} if isinstance(prepared, dict) else {}
    ts, cl = npz.get("timestamps"), npz.get("close")
    if ts is None or cl is None or len(ts) < 50:
        return []
    ts = [float(t) / (1000.0 if float(t) > 1e11 else 1.0) for t in ts]
    cl = [float(x) for x in cl]
    ref = seg_features(cl, is_long)
    out, b0 = [], 0
    for i in range(1, len(ts) + 1):
        if i == len(ts) or ts[i] - ts[b0] >= seg_days * 86400.0:
            f = seg_features(cl[b0:i], is_long)
            out.append((b0, i - 1, classify_regime(f, ref) if f else "UNKNOWN"))
            b0 = i
    return out


def _seg_pnl_dd(trades, gain, segs):
    """per-segment P&L (gain-pp units: pnl$ scaled so the sum = gain_pct, as v15_trade_autopsy.scaled_rows) and trade-level DD."""
    tl = [t for t in (trades or []) if len(t) >= 3]
    dol = [float(t[3]) if len(t) >= 4 else float(t[2]) for t in tl]
    tot = sum(dol)
    sc = (float(gain) / tot) if (gain is not None and abs(tot) > 1e-12) else 0.0
    out = []
    for b0, b1, _lab in segs:
        pn = [d * sc for t, d in zip(tl, dol) if b0 <= int(t[1]) <= b1]
        eq = peak = dd = 0.0
        for x in pn:
            eq += x
            peak = max(peak, eq)
            dd = max(dd, peak - eq)
        out.append((sum(pn), dd))
    return out


def regime_attribution(base_res, cand_res, segs):
    """{regime: {n_seg, base_pnl, cand_pnl, d_pnl, d_dd}} of a candidate vs its base on the long window (real ledgers only)."""
    if not segs or not base_res or not cand_res:
        return {}
    a = _seg_pnl_dd(base_res.get("trades_list"), base_res.get("gain"), segs)
    b = _seg_pnl_dd(cand_res.get("trades_list"), cand_res.get("gain"), segs)
    out = {}
    for (b0, b1, lab), (pa, da), (pb, dbb) in zip(segs, a, b):
        e = out.setdefault(lab, {"n_seg": 0, "base_pnl": 0.0, "cand_pnl": 0.0, "d_dd": 0.0})
        e["n_seg"] += 1
        e["base_pnl"] += pa
        e["cand_pnl"] += pb
        e["d_dd"] += dbb - da
    for e in out.values():
        e["d_pnl"] = round(e["cand_pnl"] - e["base_pnl"], 4)
        e["base_pnl"], e["cand_pnl"], e["d_dd"] = round(e["base_pnl"], 4), round(e["cand_pnl"], 4), round(e["d_dd"], 4)
    return out


def default_accept(cand, base):
    """dual-window default (caller may pass its own): 30D already passed the pilot gates with delta > 0; the long window must
    produce a result and either raise gain or lower DD vs the same base on that window (no regression on both)."""
    c, b = cand.get("w365") or {}, base.get("w365") or {}
    if c.get("gain") is None:
        return False
    if b.get("gain") is None:
        return True
    return float(c["gain"]) >= float(b["gain"]) - 1e-9 or (c.get("dd") is not None and b.get("dd") is not None and float(c["dd"]) < float(b["dd"]) - 1e-9)


def search(prepared, sym_side, base_overrides, symptoms=None, regime=None, budget_s=300, workers=4, stats_out=None, evaluator=None, library=None, seed=None, recent_prepared=None, log=print, max_gen=None, record=True, windows=(30,), prepared_365=None, accept=None):
    t0 = time.time()
    deadline = t0 + float(budget_s)
    verify_reserve = max(min(8.0, 0.2 * float(budget_s)), min(40.0, 0.08 * float(budget_s)))
    cs = cat_side_of(sym_side)
    st = {"sym_side": sym_side, "cat_side": cs, "budget_s": float(budget_s), "workers": int(workers), "started_utc": _utc()}
    own_eval = evaluator is None
    if own_eval and os.environ.get("V17_COMBO_IN_PILOT") == "1":
        raise RuntimeError("v17_combo_engine.search inside v15_pilot needs evaluator= (the pilot's pool); never forks its own")
    if prepared is None and own_eval:
        from tools.opt.v12_pilot import prepare_batch
        prepared = prepare_batch(sym_side, 30)
    if own_eval:
        if recent_prepared is None:
            from tools.opt.v12_pilot import prepare_batch
            recent_prepared = prepare_batch(sym_side, RECENT_DAYS)
        preps = {30: prepared, RECENT_DAYS: recent_prepared}
        if 365 in tuple(windows or ()):
            if prepared_365 is None:
                from tools.opt.v12_pilot import prepare_batch
                prepared_365 = prepare_batch(sym_side, 365)
            preps[365] = prepared_365
        evaluator = PoolEvaluator(preps, workers=workers)
    lib = library if library is not None else (ComboLibrary() if os.environ.get("V17_COMBO_NO_LIB") != "1" else None)
    n_evals = [0]

    def ev(items, dl=None):
        r = evaluator(items, dl)
        n_evals[0] += len(r)
        return r

    try:
        base = sanitize(dict(base_overrides), sym_side) if own_eval else dict(base_overrides)
        dual = 365 in tuple(windows or ())
        accept = accept or default_accept
        segs = regime_segments(prepared_365, sym_side.upper().endswith("_LONG")) if dual else []
        t_b = time.time()
        b = ev([("b30", 30, base, False), ("b7", RECENT_DAYS, base, False)] + ([("b365", 365, base, True)] if dual else []))
        t365 = (time.time() - t_b) if dual else 0.0
        if dual:  # room for the fresh 365D verify + one chained 365D check (crypto 365D ledger evals are the expensive part)
            verify_reserve = max(verify_reserve, 2.5 * t365)
            st["t365_s"] = round(t365, 1)
        base365 = (b.get("b365") or (None,))[0] if dual else None
        cur365 = base365
        if dual:
            st["base365"] = None if not base365 else {k: base365.get(k) for k in ("gain", "trades", "tim", "dd", "valid", "reason")}
            st["segments"] = {lab: sum(1 for s_ in segs if s_[2] == lab) for lab in sorted({s_[2] for s_ in segs})}

        def long_check(full_set, ref365, delta):
            """365D eval (ledger, for regime attribution) of full_set vs ref365 + the caller's accept rule."""
            g = ev([("l", 365, sanitize(full_set, sym_side) if own_eval else full_set, True)], None).get("l", (None,))[0]
            okk = bool(accept({"w365": g, "delta": delta}, {"w365": ref365}))
            return g, okk, regime_attribution(ref365, g, segs)
        base30, base7 = (b.get("b30") or (None,))[0], (b.get("b7") or (None,))[0]
        if not base30 or base30.get("gain") is None:
            st["error"] = f"base eval failed {(b.get('b30') or (None, ''))[1]}"
            return []
        bg, bg7 = float(base30["gain"]), (None if not base7 or base7.get("gain") is None else float(base7["gain"]))
        st["base"] = {"gain": bg, "trades": base30.get("trades"), "tim": base30.get("tim"), "dd": base30.get("dd"), "valid": base30.get("valid"), "recent_gain": bg7, "bh": base30.get("bh")}
        npz = (prepared or {}).get("npz_prepared") or {} if isinstance(prepared, dict) else {}
        close = [float(x) for x in (npz.get("close") if npz.get("close") is not None else [])]
        feats = seg_features(close, sym_side.upper().endswith("_LONG")) if close else {}
        feats.update({"bh": base30.get("bh"), "tim": base30.get("tim"), "trades": base30.get("trades")})
        feats["regime_now"] = (regime or {}).get("regime_now") if isinstance(regime, dict) and regime.get("regime_now") not in (None, "UNKNOWN") else None
        if not feats["regime_now"] and close:
            tsx = npz.get("timestamps")
            try:
                tl = [float(t) / (1000.0 if float(t) > 1e11 else 1.0) for t in tsx]
                cut = next(k for k, t in enumerate(tl) if t >= tl[-1] - RECENT_DAYS * 86400.0)
                feats["regime_now"] = classify_regime(seg_features(close[cut:], sym_side.upper().endswith("_LONG")), seg_features(close[:cut], sym_side.upper().endswith("_LONG")))
            except (StopIteration, TypeError, ValueError):
                pass
        st["features"] = {k: feats.get(k) for k in FEATS}
        if symptoms is None:
            try:
                symptoms = auto_symptoms(lambda it: ev(it), base, prepared, sym_side, base30)
            except Exception as e:
                symptoms = {}
                st["symptom_error"] = str(e)[:120]
        st["symptoms"] = symptoms
        t_gen = time.time()
        space = ComboSpace(cs, sym_side)
        st["space"] = space.size()
        rng = random.Random(seed if seed is not None else int(hashlib.md5(f"{sym_side}|{_utc()[:10]}".encode()).hexdigest()[:8], 16))
        cands, gst = generate(space, sym_side, base, symptoms, feats, lib, rng, int(max_gen or _env_f("V17_COMBO_MAX_GEN", 40000)), log)
        st["generated"] = gst
        st["gen_s"] = round(time.time() - t_gen, 1)
        order = schedule(cands, _env_f("V17_COMBO_EXPLORE", 0.15))
        # ---- successive-halving ROUNDS until the budget is used: each round screens the next slice of the ranked order on the
        # round's base (stage 1, full 30D = the gate window), takes the top ~5% to the recent 7D segment (stage 2) and greedy-
        # combines non-conflicting winners on the CURRENT best set (stage 3, re-evaluated every adoption); the next round's base
        # is the improved set. One fresh NO-LIES verify of the final joint set vs the ORIGINAL base at the end.
        tol = _env_f("V15_V17_RECENT_TOL", 0.25)
        adopt_old = os.environ.get("V17_COMBO_ADOPT_OLD_REGIME", "0") == "1"  # pilot ANTICIPATE parity: 30D-only winners are flagged, not adopted
        max_adopt = int(_env_f("V17_COMBO_MAX_ADOPT", 12))
        min_round = min(_env_f("V17_COMBO_MIN_ROUND_S", 12.0), 0.2 * float(budget_s))
        wave = max(8, evaluator.workers * 6 if hasattr(evaluator, "workers") else 64)
        cur, cur_res, cur_g, cur_g7 = dict(base), dict(base30), bg, bg7
        res1, top_all, adopted, rounds_log = [], [], [], []
        i = 0
        rnd = 0
        while i < len(order) and len(adopted) < max_adopt and time.time() < deadline - verify_reserve - min_round:
            rnd += 1
            rb, rb_res, rb_g, rb_g7 = dict(cur), dict(cur_res), cur_g, cur_g7
            now = time.time()
            s1_end = now + (deadline - verify_reserve - now) * _env_f("V17_COMBO_S1_SHARE", 0.55)
            t_s1 = time.time()
            r1 = []
            while i < len(order) and time.time() < s1_end - 1.0:
                part = []
                while i < len(order) and len(part) < wave:
                    c = order[i]
                    i += 1
                    eff = effective_delta(rb, c["delta"])
                    if not eff or (adopted and (set(eff) - {"FILTER_EXEMPTIONS"}) & {k for a in adopted for k in a["delta"] if k != "FILTER_EXEMPTIONS"}):
                        continue
                    part.append(dict(c, delta=eff))
                if not part:
                    break
                items = [(j, 30, sanitize(apply_delta(rb, c["delta"]), sym_side) if own_eval else apply_delta(rb, c["delta"]), False) for j, c in enumerate(part)]
                got = ev(items, s1_end)
                for j, c in enumerate(part):
                    r, err = got.get(j, (None, "deadline"))
                    if r is None and err == "deadline":
                        continue
                    d = None if not r or r.get("gain") is None else float(r["gain"]) - rb_g
                    r1.append(dict(c, res=r, err=err, d=d, round=rnd, base_gain=rb_g, gate=gate_reason(r, rb_res, c["delta"]) if r else "NO_RESULT"))
            res1.extend(r1)
            pos = sorted([x for x in r1 if x["d"] is not None and x["d"] > 1e-9 and not x["gate"]], key=lambda x: -x["d"])
            k2 = int(min(_env_f("V17_COMBO_S2_MAX", 300), max(_env_f("V17_COMBO_S2_MIN", 20), _env_f("V17_COMBO_TOP_FRAC", 0.05) * len(r1))))
            top = pos[:k2]
            t_s2 = time.time()
            if top and rb_g7 is not None:
                got = ev([(j, RECENT_DAYS, sanitize(apply_delta(rb, x["delta"]), sym_side) if own_eval else apply_delta(rb, x["delta"]), False) for j, x in enumerate(top)], deadline - verify_reserve)
                for j, x in enumerate(top):
                    r7 = (got.get(j) or (None,))[0]
                    x["recent_gain"] = None if not r7 or r7.get("gain") is None else float(r7["gain"])
                    x["recent_delta"] = None if x["recent_gain"] is None else x["recent_gain"] - rb_g7
            for x in top:
                x["flags"] = ["OLD_REGIME_ONLY"] if x.get("recent_delta") is not None and x["recent_delta"] < -tol else []
            if dual and top:
                n365 = int(_env_f("V17_COMBO_365_MAX", 24))
                adopt_old_ = os.environ.get("V17_COMBO_ADOPT_OLD_REGIME", "0") == "1"
                chk = [x for x in top if adopt_old_ or "OLD_REGIME_ONLY" not in x["flags"]]
                n_ok = n_chk = 0
                rej = []  # canonical item sets of long-window rejects: a superset combo is skipped (pruned, flagged unchecked)
                for x in chk:
                    if time.time() > deadline - verify_reserve - 1.5 * t365 or n_ok >= n365 or n_chk >= 3 * n365:
                        break
                    xi = {f"{k}={canon(v)}" for k, v in x["delta"].items()}
                    if any(r_ <= xi for r_ in rej):
                        continue
                    n_chk += 1
                    g, okk, attr = long_check(apply_delta(rb, x["delta"]), cur365, x["delta"])
                    x["w365"], x["regime_attr"] = (None if not g else {k: g.get(k) for k in ("gain", "trades", "tim", "dd", "valid")}), attr
                    if not okk:
                        x["flags"].append("LONG_WINDOW_REJECT")
                        rej.append(xi)
                    else:
                        n_ok += 1
                for x in top:
                    if "w365" not in x and "LONG_WINDOW_REJECT" not in x["flags"]:
                        x["flags"].append("LONG_WINDOW_UNCHECKED")
            top_all.extend(top)
            t_s3 = time.time()
            pool = [x for x in top if "OLD_REGIME_ONLY" not in x["flags"]] + ([x for x in top if "OLD_REGIME_ONLY" in x["flags"]] if adopt_old else [])
            if dual:
                pool = [x for x in pool if "LONG_WINDOW_REJECT" not in x["flags"] and "LONG_WINDOW_UNCHECKED" not in x["flags"]]
            n_ad0 = len(adopted)
            first = True
            while pool and len(adopted) < max_adopt and time.time() < deadline - verify_reserve:
                if not first:
                    fixed = {k for a in adopted for k in a["delta"] if k != "FILTER_EXEMPTIONS"}
                    pool = [x for x in pool if not (set(x["delta"]) - {"FILTER_EXEMPTIONS"}) & fixed]
                    got = ev([(j, 30, sanitize(apply_delta(cur, x["delta"]), sym_side) if own_eval else apply_delta(cur, x["delta"]), False) for j, x in enumerate(pool)], deadline - verify_reserve)
                    nxt = []
                    for j, x in enumerate(pool):
                        r = (got.get(j) or (None,))[0]
                        if r is None or r.get("gain") is None:
                            continue
                        x = dict(x, res_cur=r, d_cur=float(r["gain"]) - cur_g, gate_cur=gate_reason(r, cur_res, x["delta"]))
                        if x["d_cur"] > 1e-9 and not x["gate_cur"]:
                            nxt.append(x)
                    pool = sorted(nxt, key=lambda x: (("OLD_REGIME_ONLY" in x["flags"]), -x["d_cur"]))
                else:
                    for x in pool:
                        x.update(res_cur=x["res"], d_cur=x["d"], gate_cur=x["gate"])
                    first = False
                if not pool:
                    break
                ch = pool.pop(0)
                if dual and len(adopted) > n_ad0:
                    g, okk, attr = long_check(apply_delta(cur, ch["delta"]), cur365, ch["delta"])
                    if not okk:
                        continue
                    ch = dict(ch, w365_chain={k: (g or {}).get(k) for k in ("gain", "trades", "tim", "dd")}, regime_attr_chain=attr, _g365=g)
                elif dual:
                    ch = dict(ch, _g365=None)
                cur = apply_delta(cur, ch["delta"])
                cur_res, cur_g = ch["res_cur"], float(ch["res_cur"]["gain"])
                if ch.get("recent_gain") is not None and len(adopted) == n_ad0:
                    cur_g7 = ch["recent_gain"]
                adopted.append(dict(ch, before=cur_g - ch["d_cur"], after=cur_g))
                if dual:
                    g_ = ch.get("_g365")
                    if g_ is None:
                        g_ = ev([("l", 365, sanitize(cur, sym_side) if own_eval else cur, True)], None).get("l", (None,))[0]
                    cur365 = g_ or cur365
                log(f"[combo-adopt] {sym_side} r{rnd} +{ch['d_cur']:.4f} -> {cur_g:.4f} {combo_key(ch['delta'])[:160]} ({ch['origin']}: {ch['why'][:80]})")
            if len(adopted) > n_ad0 + 1 and cur_g7 is not None:
                g7 = ev([("r7", RECENT_DAYS, sanitize(cur, sym_side) if own_eval else cur, False)]).get("r7", (None,))[0]
                cur_g7 = None if not g7 or g7.get("gain") is None else float(g7["gain"])
            rounds_log.append({"round": rnd, "screened": len(r1), "positive": len(pos), "survivors": len(top), "recent_hold": sum(1 for x in top if "OLD_REGIME_ONLY" not in x["flags"]), "adopted": len(adopted) - n_ad0, "base_gain": round(rb_g, 4), "end_gain": round(cur_g, 4), "s1_s": round(t_s2 - t_s1, 1), "s2_s": round(t_s3 - t_s2, 1), "s3_s": round(time.time() - t_s3, 1)})
            log(f"[combo-round] {sym_side} r{rnd} screened={len(r1)} pos={len(pos)} top={len(top)} adopted={len(adopted) - n_ad0} {rb_g:.4f}->{cur_g:.4f}")
        top = top_all
        st["rounds"] = rounds_log
        st["stage1"] = {"screened": len(res1), "positive": sum(1 for x in res1 if x["d"] is not None and x["d"] > 1e-9 and not x["gate"]), "changed": sum(1 for x in res1 if x["res"] and x["res"].get("fp") != base30.get("fp")), "order_len": len(order), "order_pos": i}
        st["stage2"] = {"survivors": len(top_all), "recent_hold": sum(1 for x in top_all if "OLD_REGIME_ONLY" not in x["flags"])}
        st["stage3"] = {"adopted": len(adopted)}
        # ---- NO-LIES verify: fresh evals with the eval cache OFF (30D + 7D) + ledgers for the few-trades guard
        out = []
        verified = None
        if adopted:
            prev_cache = os.environ.get("V15_EVAL_CACHE")
            os.environ["V15_EVAL_CACHE"] = "0"
            try:
                full_cur = sanitize(cur, sym_side) if own_eval else cur
                items = [("v30", 30, full_cur, True), ("v7", RECENT_DAYS, full_cur, False), ("vb", 30, base, True)]
                steps = [base]
                for a in adopted[:-1]:
                    steps.append(apply_delta(steps[-1], a["delta"]))
                for si, s in enumerate(steps[1:], start=1):
                    items.append((f"s{si}", 30, sanitize(s, sym_side) if own_eval else s, True))
                vv = evaluator(items, None)
                n_evals[0] += len(vv)
            finally:
                if prev_cache is None:
                    os.environ.pop("V15_EVAL_CACHE", None)
                else:
                    os.environ["V15_EVAL_CACHE"] = prev_cache
            v30, v7, vb = (vv.get("v30") or (None,))[0], (vv.get("v7") or (None,))[0], (vv.get("vb") or (None,))[0]
            ledgers = [(vb or {}).get("trades_list")] + [((vv.get(f"s{si}") or (None,))[0] or {}).get("trades_list") for si in range(1, len(adopted))] + [(v30 or {}).get("trades_list")]
            for ai, a in enumerate(adopted):
                n_ch, conc = _changed_trades(ledgers[ai], ledgers[ai + 1])
                a["changed_trades"], a["delta_concentration"] = n_ch, round(conc, 3)
                if n_ch <= int(_env_f("V17_COMBO_FEW_TRADES", 3)) or conc > 0.7:
                    a["flags"] = list(a.get("flags") or []) + ["FEW_TRADES"]
            ok = bool(v30 and v30.get("gain") is not None and abs(float(v30["gain"]) - cur_g) < 1e-6 and float(v30["gain"]) > bg + 1e-9 and not gate_reason(v30, base30, effective_delta(base, cur)))
            v365 = attr365 = None
            if dual:
                os.environ["V15_EVAL_CACHE"] = "0"
                try:
                    v365 = evaluator([("v365", 365, sanitize(cur, sym_side) if own_eval else cur, True)], None).get("v365", (None,))[0]
                    n_evals[0] += 1
                finally:
                    if prev_cache is None:
                        os.environ.pop("V15_EVAL_CACHE", None)
                    else:
                        os.environ["V15_EVAL_CACHE"] = prev_cache
                attr365 = regime_attribution(base365, v365, segs)
                ok = ok and bool(accept({"w365": v365, "delta": effective_delta(base, cur)}, {"w365": base365}))
            verified = {"w365": None if not v365 else {k: v365.get(k) for k in ("gain", "trades", "tim", "dd", "valid")}, "base365": st.get("base365"), "ok": ok, "gain": None if not v30 else v30.get("gain"), "greedy_gain": cur_g, "match": bool(v30 and v30.get("gain") is not None and abs(float(v30["gain"]) - cur_g) < 1e-6), "gate": gate_reason(v30, base30, effective_delta(base, cur)) if v30 else "NO_RESULT", "base_fresh_gain": None if not vb else vb.get("gain")}
            st["verify"] = verified
            if ok:
                eff = effective_delta(base, cur)
                rg = None if not v7 or v7.get("gain") is None else float(v7["gain"])
                ps = None
                try:
                    import metrics_guard as MG
                    ps = MG.validate_and_format_sharpe(float(v30.get("pool_sharpe") or 0.0), label="pool_sharpe", n_syms=1, years=30 / 365.25, trades=int(v30.get("trades") or 0), mode="crypto" if cs.startswith("CRYPTO") else "stocks")
                except Exception as e:
                    ps = f"pool_sharpe [UNVERIFIED: {str(e)[:60]}]"
                out.append({"kind": "joint", "overrides_delta": eff, "gain": float(v30["gain"]), "base_gain": bg, "delta": float(v30["gain"]) - bg, "trades": v30.get("trades"), "tim": v30.get("tim"), "dd": v30.get("dd"), "wr": v30.get("wr"), "pool_sharpe": ps, "recent_segment_gain": rg, "recent_delta": None if rg is None or bg7 is None else rg - bg7, "n_evals": n_evals[0], "origin": "+".join(sorted({a["origin"] for a in adopted})), "why": f"{len(adopted)} adopted combos, fresh no-cache verify", "flags": sorted({f for a in adopted for f in a.get("flags") or []}), "tf": tf_features(eff), "w365": verified.get("w365"), "base365": st.get("base365"), "regime_attr_365": attr365, "steps": [{"overrides_delta": a["delta"], "d_on_chain": round(a["d_cur"], 4), "d_vs_round_base": round(a["d"], 4), "round": a.get("round"), "recent_delta": a.get("recent_delta"), "origin": a["origin"], "why": a["why"], "changed_trades": a.get("changed_trades"), "delta_concentration": a.get("delta_concentration"), "flags": a.get("flags") or []} for a in adopted]})
        for x in top:
            r = x["res"]
            out.append({"kind": "combo", "overrides_delta": x["delta"], "gain": r.get("gain"), "base_gain": x.get("base_gain", bg), "round": x.get("round"), "delta": x["d"], "trades": r.get("trades"), "tim": r.get("tim"), "dd": r.get("dd"), "wr": r.get("wr"), "recent_segment_gain": x.get("recent_gain"), "recent_delta": x.get("recent_delta"), "n_evals": 2 if x.get("recent_gain") is not None else 1, "origin": x["origin"], "why": x["why"], "flags": x.get("flags") or [], "tf": tf_features(x["delta"]), "w365": x.get("w365"), "regime_attr_365": x.get("regime_attr"), "adopted": any(a is x or a.get("delta") == x["delta"] for a in adopted)})
        wall = time.time() - t0
        cpu = getattr(evaluator, "cpu_s", None)
        st.update({"wall_s": round(wall, 1), "n_evals": n_evals[0], "evals_per_s": round(n_evals[0] / max(1e-6, wall), 2), "evals_per_s_per_core": round(n_evals[0] / max(1e-6, wall) / max(1, getattr(evaluator, "workers", 1)), 3), "cpu_s_per_eval": None if not cpu else round(cpu / max(1, getattr(evaluator, "n", 1)), 4), "combos_per_min": round(len(res1) / max(1e-6, wall) * 60.0, 1), "lift": None if not out or out[0]["kind"] != "joint" else round(out[0]["delta"], 4)})
        log(f"[combo-search] {sym_side} base={bg:.4f} screened={len(res1)} pos={st['stage1']['positive']} adopted={len(adopted)} lift={st['lift']} evals={n_evals[0]} wall={wall:.0f}s ({st['evals_per_s']}/s)")
        if record and lib is not None:
            try:
                _record(lib, sym_side, cs, feats, symptoms, base30, bg, bg7, res1, top, out, base365)
            except sqlite3.Error as e:
                st["library_error"] = str(e)[:120]
        return out
    finally:
        if stats_out is not None:
            stats_out.update(st)
        if own_eval:
            evaluator.close()
        if library is None and lib is not None:
            lib.close()


def _engine_md5():
    try:
        return _file_md5(ROOT / "v12_quick_engine.py")
    except OSError:
        return None


def _record(lib, sym_side, cs, feats, symptoms, base30, bg, bg7, res1, top, out, base365=None):
    em = _engine_md5()
    syms = [k for k, v in (symptoms or {}).items() if float((v or {}).get("size") or 0) > 0]
    bid = lib.board(f"{sym_side}|combo_engine|{_utc()}|{round(bg, 4)}", sym_side=sym_side, cat_side=cs, src="v17_combo_engine", run="v17_combo_engine", ts_last=time.time(), gain=bg, bh=base30.get("bh"), tim=base30.get("tim"), dd=base30.get("dd"), trades=base30.get("trades"), side_net_pct=feats.get("side_net_pct"), vol_bar_pct=feats.get("vol_bar_pct"), chop=feats.get("chop"), range_pct=feats.get("range_pct"), regime=classify_regime(feats), symptoms=syms, engine_md5=em)
    top_d = {combo_key(x["delta"]): x for x in top}
    sym0 = max(syms, key=lambda s: float(symptoms[s].get("size") or 0)) if syms else ""
    for x in res1:
        r = x.get("res")
        if not r or x.get("d") is None:
            continue
        t = top_d.get(combo_key(x["delta"]))
        lib.add_evidence(bid, x["delta"], "v17_combo_engine", x["d"], origin=x["origin"], symptom=sym0 if x["d"] > 0 else "", gain=r.get("gain"), base_gain=x.get("base_gain", bg), trades=r.get("trades"), tim=r.get("tim"), dd=r.get("dd"), wr=r.get("wr"), valid=r.get("valid") and not x.get("gate"), recent_delta=(t or {}).get("recent_delta"), engine_md5=em)
    b365 = (base365 or {}).get("gain")
    for x in top:
        w = x.get("w365") or {}
        if w.get("gain") is not None and b365 is not None:
            lib.add_evidence(bid, x["delta"], "v17_combo_engine", float(w["gain"]) - float(b365), window_days=365, origin=x["origin"], gain=w.get("gain"), base_gain=b365, trades=w.get("trades"), tim=w.get("tim"), dd=w.get("dd"), valid=w.get("valid"), engine_md5=em)
        if x.get("regime_attr"):
            lib.add_regime_attr(bid, x["delta"], x["regime_attr"])
    if out and out[0]["kind"] == "joint":
        j = out[0]
        if j.get("regime_attr_365"):
            lib.add_regime_attr(bid, j["overrides_delta"], j["regime_attr_365"])
        lib.add_evidence(bid, j["overrides_delta"], "v17_combo_engine_joint", j["delta"], origin=j["origin"], symptom=sym0, gain=j["gain"], base_gain=bg, trades=j["trades"], tim=j["tim"], dd=j["dd"], wr=j["wr"], valid=True, recent_delta=j.get("recent_delta"), engine_md5=em)
    lib.commit()


# ------------------------------------------------------------------------------------------------------------------- CLI
def latest_published_set(sym_side):
    """newest progress JSON (fleet run dirs, imports, lifecycle_pilot) with a non-empty cumulative_overrides -> (path, doc)."""
    pats = ["~/v15_run*/progress/{s}_v14_progress.json", "~/v15_import*/progress/{s}_v14_progress.json", "~/v15_run_import*/progress/{s}_v14_progress.json", str(ROOT / "data" / "reports" / "lifecycle_pilot" / "{s}_v14_progress.json")]
    files = sorted({f for p in pats for f in glob.glob(os.path.expanduser(p.format(s=sym_side)))}, key=lambda f: -os.path.getmtime(f))
    best = None
    for f in files:
        try:
            d = json.loads(Path(f).read_text())
        except (OSError, ValueError):
            continue
        if not d.get("cumulative_overrides"):
            continue
        done = str(d.get("result_done_utc") or "")
        if best is None or done > best[2]:
            best = (f, d, done)
    return (best[0], best[1]) if best else (None, None)


def _default_globs():
    pg = ["~/v15_run*/progress/*_v14_progress.json", "~/v15_import*/progress/*_v14_progress.json", "~/v15_run_import*/progress/*_v14_progress.json", str(ROOT / "data" / "reports" / "lifecycle_pilot" / "*_v14_progress.json")]
    dg = ["~/v15_run*/progress/v15_delta_log/*.jsonl", "~/v15_import*/progress/v15_delta_log/*.jsonl", "~/v15_run_import*/progress/v15_delta_log/*.jsonl"]
    return pg, dg


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build-library")
    b.add_argument("--out", default=str(LIB_PATH))
    b.add_argument("--progress-glob", action="append")
    b.add_argument("--delta-glob", action="append")
    b.add_argument("--max-files", type=int)
    m = sub.add_parser("merge")
    m.add_argument("--into", required=True)
    m.add_argument("--from", dest="srcs", action="append", required=True)
    i = sub.add_parser("info")
    i.add_argument("--lib", default=str(LIB_PATH))
    s = sub.add_parser("space")
    s.add_argument("--cat-side", required=True)
    s.add_argument("--sym-side")
    p = sub.add_parser("proof")
    p.add_argument("--sym-side", required=True)
    p.add_argument("--budget", type=float, default=300.0)
    p.add_argument("--workers", type=int, default=4)
    p.add_argument("--out")
    p.add_argument("--seed", type=int)
    p.add_argument("--dual", action="store_true", help="dual-window mode: 30D + 365D (default accept: 365D gain up or DD down)")
    a = ap.parse_args(argv)
    if a.cmd == "build-library":
        pg, dg = _default_globs()
        lib = ComboLibrary(a.out)
        r = build_library(lib, a.progress_glob or pg, a.delta_glob or dg, max_files=a.max_files)
        r["info"] = lib.info()
        lib.close()
        print(json.dumps(r, default=str))
    elif a.cmd == "merge":
        print(json.dumps(merge_library(a.into, a.srcs), default=str))
    elif a.cmd == "info":
        lib = ComboLibrary(a.lib)
        print(json.dumps(lib.info(), indent=1, default=str))
        lib.close()
    elif a.cmd == "space":
        sp = ComboSpace(a.cat_side)
        print(json.dumps(sp.size()))
        if a.sym_side:
            src, d = latest_published_set(a.sym_side)
            base = sanitize(d["cumulative_overrides"], a.sym_side)
            c, gst = generate(sp, a.sym_side, base, {}, {}, ComboLibrary() if LIB_PATH.exists() else None, random.Random(1), 40000, print)
            print(json.dumps({"src": src, "generated": gst}, default=str))
    elif a.cmd == "proof":
        src, d = latest_published_set(a.sym_side)
        if not d:
            print(json.dumps({"sym_side": a.sym_side, "error": "no published set"}))
            return 2
        st = {}
        ranked = search(None, a.sym_side, d["cumulative_overrides"], budget_s=a.budget, workers=a.workers, stats_out=st, seed=a.seed, windows=(30, 365) if a.dual else (30,))
        rep = {"sym_side": a.sym_side, "start_src": src, "start_published_gain": d.get("final_gain_fresh_vec", d.get("cumulative_gain")), "start_npz": d.get("npz_id"), "stats": st, "ranked": ranked[:40]}
        try:
            from tools.opt.evaluate_v12 import _npz_path_for
            npp = _npz_path_for(a.sym_side.rsplit("_", 1)[0])
            rep["npz_md5_now"] = _file_md5(npp) if npp else None
        except Exception:
            pass
        txt = json.dumps(rep, default=str, indent=1)
        if a.out:
            Path(a.out).write_text(txt)
        print(json.dumps({k: rep[k] for k in ("sym_side", "start_src", "start_published_gain")} | {"lift": st.get("lift"), "screened": (st.get("stage1") or {}).get("screened"), "evals_per_s": st.get("evals_per_s"), "verify": st.get("verify")}, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
