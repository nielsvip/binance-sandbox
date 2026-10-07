#!/usr/bin/env python3
"""Read-only: for every template switch (col A) and filter header (FILTER=opt), where is it really connected?
config.py / config_tradier.py / QuickConfig field, and REAL reads in ez_manage.py, tradier_manage.py and the sweep engine
(v12_quick_engine.py + vec_decisions/). Registry-list lines ('NAME',) and dead-farm `_=getattr(...)` reads are not reads."""
import csv
import dataclasses
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import openpyxl  # noqa: E402

TEMPLATES = ["TEMPLATE_CRYPTO_LONG", "TEMPLATE_CRYPTO_SHORT", "TEMPLATE_STOCKS_LONG", "TEMPLATE_STOCKS_SHORT"]
SWITCH_SHEETS = ["ENTRY_REVERSAL_BOUNCE", "STDEV_SLOPE_SIZING", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
DEAD = re.compile(r"^\s*(['\"][A-Z0-9_]+['\"]\s*,?\s*(#.*)?$|_\s*=\s*getattr\(|#)")


def real_reads(lines, name):
    pat = re.compile(r"(?<![A-Z0-9_])" + re.escape(name) + r"(?![A-Z0-9_])")
    n = 0
    for ln in lines:
        if pat.search(ln) and not DEAD.match(ln) and not re.match(r"^\s*" + re.escape(name) + r"\s*:", ln):
            n += 1
    return n


def main():
    import config
    import config_tradier
    import v12_quick_engine as V
    q = {f.name for f in dataclasses.fields(V.QuickConfig)}
    ez = (ROOT / "ez_manage.py").read_text(errors="ignore").splitlines()
    tr = (ROOT / "tradier_manage.py").read_text(errors="ignore").splitlines()
    vec = (ROOT / "v12_quick_engine.py").read_text(errors="ignore").splitlines()
    for p in sorted((ROOT / "vec_decisions").glob("*.py")):
        if not p.name.startswith("test_"):
            vec += p.read_text(errors="ignore").splitlines()
    names = {}
    for t in TEMPLATES:
        wb = openpyxl.load_workbook(ROOT / "SPREADSHEETS" / f"{t}.xlsx", read_only=True)
        for s in SWITCH_SHEETS:
            if s not in wb.sheetnames:
                continue
            rows = wb[s].iter_rows(min_row=2, values_only=True)
            hdr = next(rows)
            for h in hdr[14:]:
                if isinstance(h, str) and "=" in h:
                    names.setdefault(h.split("=", 1)[0].strip(), set()).add(t[9:] + ":filter")
            for row in rows:
                if row and isinstance(row[0], str) and row[0].strip() and row[0].strip().upper() == row[0].strip():
                    names.setdefault(row[0].strip(), set()).add(t[9:])
    out = ROOT / "data" / "reports" / "switch_connection_audit_20260930.csv"
    counts = {}
    with out.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["name", "used_in", "config_py", "config_tradier", "quickconfig", "ez_reads", "tradier_reads", "engine_reads", "status"])
        for name in sorted(names):
            c, t, qq = hasattr(config.Config, name), hasattr(config_tradier.TradierConfig, name), name in q
            e, r, v = real_reads(ez, name), real_reads(tr, name), real_reads(vec, name)
            missing = [lbl for lbl, ok in (("config", c or t), ("QuickConfig", qq), ("ez_", e > 0), ("tradier_", r > 0), ("engine", v > 0)) if not ok]
            status = "CONNECTED" if not missing else "MISSING:" + "+".join(missing)
            counts[status.split(":")[0] if status == "CONNECTED" else "MISSING"] = counts.get(status.split(":")[0] if status == "CONNECTED" else "MISSING", 0) + 1
            w.writerow([name, ";".join(sorted(names[name])), c, t, qq, e, r, v, status])
    print(f"{len(names)} names -> {out}")
    print(counts)


if __name__ == "__main__":
    main()
