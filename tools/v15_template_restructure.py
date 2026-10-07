#!/usr/bin/env python3
"""v15_template_restructure — PILOT/PREVIEW: rearrange a template's switch rows worst_first by avg_delta,
keeping every row wholly intact (values + bold + fill + number_format + all yellow L:BI cells move together).
Orange (GENERAL, FFE699-filled) switch-groups go below the white ones, same worst_first order.

SAFE: reads SPREADSHEETS/TEMPLATE_{cat_side}.xlsx + v15_avg_delta_latest.xlsx; writes a REARRANGED COPY to
SPREADSHEETS/V15_TEMPLATE_PREVIEW/ (originals untouched) + a human-readable report. No defaults changed here
(that's the separate promote step); this pilot is purely to verify the rearrange doesn't corrupt content.
Switch-tab data rows are VALUES not formulas (CLAUDE.md §TEMPLATE r>=3), so moving rows is formula-safe.
"""
import argparse, copy, json, pathlib
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
SPREAD = ROOT / "SPREADSHEETS"
PREVIEW = SPREAD / "V15_TEMPLATE_PREVIEW"
AVGDELTA = SPREAD / "v15_avg_delta_latest.xlsx"
TEMPLATES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx",
             "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
# ONLY these 13 SWITCH_SHEETS may be reordered — never the dictionary/reference/results tabs
# (reordering FILTER_DICTIONARY_V2 etc. would scramble yellow-map lookups).
SWITCH_SHEETS = {"STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL",
                 "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED",
                 "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK",
                 "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"}
ORANGE_FILL = "FFE699"
HDR_ROWS = 2  # rows 1-2 are the tab title + column headers, kept fixed


def load_avg_delta(cat_side):
    if not AVGDELTA.exists():
        return {}
    wb = openpyxl.load_workbook(str(AVGDELTA), data_only=True)
    if cat_side not in wb.sheetnames:
        return {}
    ws = wb[cat_side]
    hdr = [c.value for c in ws[1]]
    try:
        i_name, i_avg = hdr.index("name"), hdr.index("avg_delta")
    except ValueError:
        return {}
    acc = {}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(row=r, column=i_name + 1).value
        av = ws.cell(row=r, column=i_avg + 1).value
        if nm is not None and av is not None:
            # stats are per "SWITCH=value" (v15_avg_delta_rebuild): a switch group sorts by the mean of its values
            acc.setdefault(str(nm).split("=", 1)[0].strip(), []).append(float(av))
    wb.close()
    return {k: sum(v) / len(v) for k, v in acc.items()}


def snapshot_row(ws, r, ncol):
    """capture (value,font,fill,number_format,alignment) per column so we can rewrite it elsewhere intact."""
    cells = []
    for c in range(1, ncol + 1):
        cc = ws.cell(row=r, column=c)
        cells.append((cc.value, copy.copy(cc.font), copy.copy(cc.fill), cc.number_format, copy.copy(cc.alignment)))
    return cells


def write_row(ws, r, cells):
    for c, (val, font, fill, nfmt, align) in enumerate(cells, start=1):
        cc = ws.cell(row=r, column=c)
        cc.value = val
        cc.font = font
        cc.fill = fill
        cc.number_format = nfmt
        cc.alignment = align


def is_orange(ws, r):
    f = ws.cell(row=r, column=1).fill
    return bool(f and f.fgColor and ORANGE_FILL in str(f.fgColor.rgb or ""))


def restructure_tab(ws, avg):
    ncol = ws.max_column
    # gather data rows (r>HDR) grouped by (switch, colour class) in document order. The class is decided PER ROW:
    # a switch whose rows are mixed white/orange splits into a white group and an orange group, so a white row can
    # never be carried below the orange block (TEMPLATE_STOCKS_SHORT EXIT_VELOCITY MTF_EXIT_USE_COMPOUND, 2026-09-29).
    groups = []          # list of dict{switch, rows:[snapshots], orange, avgd}
    by_key = {}
    last = None
    for r in range(HDR_ROWS + 1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            if last is not None:
                last["rows"].append(snapshot_row(ws, r, ncol))  # trailing blank stays with group
            continue
        a = str(a).strip()
        k = (a, is_orange(ws, r))
        if k not in by_key:
            by_key[k] = {"switch": a, "rows": [], "orange": k[1], "avgd": avg.get(a, None)}
            groups.append(by_key[k])
        last = by_key[k]
        last["rows"].append(snapshot_row(ws, r, ncol))
    if not groups:
        return []
    # worst_first: most-negative avg_delta first; None (no data) sinks to the end of its class, stable
    BIG = float("inf")
    white = [g for g in groups if not g["orange"]]
    orange = [g for g in groups if g["orange"]]
    white.sort(key=lambda g: (g["avgd"] is None, g["avgd"] if g["avgd"] is not None else BIG))
    orange.sort(key=lambda g: (g["avgd"] is None, g["avgd"] if g["avgd"] is not None else BIG))
    ordered = white + orange
    # rewrite rows HDR+1.. in the new order
    r = HDR_ROWS + 1
    order_report = []
    for g in ordered:
        order_report.append((g["switch"], "orange" if g["orange"] else "white", g["avgd"]))
        for cells in g["rows"]:
            write_row(ws, r, cells)
            r += 1
    bad = white_below_orange(ws)
    if bad:
        raise RuntimeError(f"{ws.title}: white rows {bad[:10]} below the first orange row after restructure")
    return order_report


def white_below_orange(ws):
    """rows (r>HDR) holding a white switch row below the first orange filter row — must always be empty (BIBLE §56 R5)"""
    first = next((r for r in range(HDR_ROWS + 1, ws.max_row + 1) if ws.cell(row=r, column=1).value not in (None, "") and is_orange(ws, r)), None)
    if first is None:
        return []
    return [r for r in range(first, ws.max_row + 1) if ws.cell(row=r, column=1).value not in (None, "") and not is_orange(ws, r)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cat-side", default="STOCKS_LONG", help="pilot one cat_side; 'ALL' for all four")
    ap.add_argument("--tab", default=None, help="restrict to one tab (pilot)")
    args = ap.parse_args()
    PREVIEW.mkdir(parents=True, exist_ok=True)
    cats = list(TEMPLATES) if args.cat_side == "ALL" else [args.cat_side]
    report = {}
    for cs in cats:
        src = SPREAD / TEMPLATES[cs]
        if not src.exists():
            print(f"[{cs}] MISSING {src}"); continue
        avg = load_avg_delta(cs)
        wb = openpyxl.load_workbook(str(src), data_only=False)
        tabreport = {}
        for tab in wb.sheetnames:
            if tab not in SWITCH_SHEETS:
                continue
            if args.tab and tab != args.tab:
                continue
            order = restructure_tab(wb[tab], avg)
            tabreport[tab] = order
            top = [f"{sw}({'' if av is None else round(av,3)})" for sw, kind, av in order[:5]]
            print(f"[{cs}/{tab}] {len(order)} groups; worst-first top5: {top}")
        outp = PREVIEW / TEMPLATES[cs]
        wb.save(str(outp))
        wb.close()
        report[cs] = {t: [(s, k, a) for s, k, a in o] for t, o in tabreport.items()}
        print(f"[{cs}] preview -> {outp}")
    (PREVIEW / "restructure_report.json").write_text(json.dumps(report, indent=1, default=str))
    print(f"[report] {PREVIEW/'restructure_report.json'}")


if __name__ == "__main__":
    main()
