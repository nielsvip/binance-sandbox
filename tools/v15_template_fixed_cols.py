#!/usr/bin/env python3
"""v15_template_fixed_cols — restore the fixed columns between PER_ROW_FILTERS and the yellow block (USER 2026-09-30).

Layout every SWITCH_SHEETS tab must have (as STDEV_SLOPE_SIZING kept it; BIBLE §56 R10 yellows start at O):
    K PER_ROW_FILTERS | L is_default | M AVG_DELTA | N POS_SYM | O.. yellow "FILTER=opt" columns
Found broken: yellow headers started at L (AVG_DELTA/POS_SYM/is_default headers lost), the avg/pos VALUES sat in M/N inside
yellow cells, is_default was written far right. Per tab: the yellow block (headers, fills, notes, cells) is shifted so it
starts at O, the avg/pos values go to M/N (and are cleared from the yellow cells they polluted), is_default moves to L and
its far-right copy is removed. Report only unless --apply (backup first).
"""
import argparse
import datetime
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
SPREAD = ROOT / "SPREADSHEETS"
TEMPLATES = ["TEMPLATE_CRYPTO_LONG.xlsx", "TEMPLATE_CRYPTO_SHORT.xlsx", "TEMPLATE_STOCKS_LONG.xlsx", "TEMPLATE_STOCKS_SHORT.xlsx"]
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
K, L, M, N, O = 11, 12, 13, 14, 15
HFONT = Font(name="Arial", size=10, bold=True)


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def fix_tab(ws, apply: bool) -> str:
    hdr = {c: ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)}
    assert str(hdr.get(K) or "").strip() == "PER_ROW_FILTERS", f"{ws.title}: K is {hdr.get(K)!r}"
    first_y = next((c for c in range(L, ws.max_column + 1) if isinstance(hdr[c], str) and "=" in hdr[c]), None)
    isd = [c for c, v in hdr.items() if isinstance(v, str) and v.strip().lower().startswith("is_default")]
    avg_h = next((c for c, v in hdr.items() if isinstance(v, str) and v.strip().upper().replace(" ", "_") == "AVG_DELTA"), None)
    pos_h = next((c for c, v in hdr.items() if isinstance(v, str) and v.strip().upper() == "POS_SYM"), None)
    rows = range(3, ws.max_row + 1)
    if first_y is None:
        # no yellow block (STDEV): only normalise the three fixed headers in place
        msg = f"no yellows; L={hdr.get(L)!r} M={hdr.get(M)!r} N={hdr.get(N)!r}"
        if apply and isd == [L] and avg_h == M and pos_h == N:
            for c, name in ((L, "is_default"), (M, "AVG_DELTA"), (N, "POS_SYM")):
                ws.cell(row=2, column=c).value = name
                ws.cell(row=2, column=c).font = HFONT
        elif isd != [L] or avg_h != M or pos_h != N:
            msg += "  !! UNEXPECTED — not touched"
        return msg
    # avg/pos values: the adjacent numeric column pair in L..first_y+3 (the stray avg/pos write), unless real headers exist
    if avg_h and pos_h:
        src_avg, src_pos = avg_h, pos_h
    else:
        counts = {c: sum(1 for r in rows if num(ws.cell(row=r, column=c).value)) for c in range(L, first_y + 4)}
        pair = [c for c in range(L, first_y + 3) if counts.get(c) and counts.get(c + 1) == counts.get(c)]
        src_avg, src_pos = (pair[0], pair[0] + 1) if pair else (None, None)
    vals = {r: (ws.cell(row=r, column=src_avg).value if src_avg else None, ws.cell(row=r, column=src_pos).value if src_pos else None) for r in rows}
    isd_src = isd[-1] if isd else None
    isd_vals = {r: ws.cell(row=r, column=isd_src).value for r in rows} if isd_src else {}
    shift = O - first_y
    msg = f"yellows start col {first_y} -> O (shift {shift:+d}); avg/pos values from cols {src_avg},{src_pos} ({sum(1 for v in vals.values() if num(v[0]))} rows); is_default from col {isd_src}"
    if not apply:
        return msg
    if src_avg:
        for r in rows:  # stray avg/pos cleared from the cells they polluted (fills stay)
            ws.cell(row=r, column=src_avg).value = None
            ws.cell(row=r, column=src_pos).value = None
    if isd_src and isd_src >= first_y:
        isd_src_after = isd_src + shift
    else:
        isd_src_after = isd_src
    if shift > 0:
        ws.insert_cols(first_y, shift)
    elif shift < 0:
        # the columns first_y+shift .. first_y-1 hold no yellow headers (checked below) -> removed
        for c in range(O, first_y):
            assert not (isinstance(hdr.get(c), str) and "=" in hdr[c]), f"{ws.title}: col {c} is a yellow header"
        ws.delete_cols(O, -shift)
    if isd_src_after and isd_src_after not in (L, M, N):
        ws.delete_cols(isd_src_after, 1)
    for c, name in ((L, "is_default"), (M, "AVG_DELTA"), (N, "POS_SYM")):
        h = ws.cell(row=2, column=c)
        h.value, h.font, h.fill = name, HFONT, openpyxl.styles.PatternFill(fill_type=None)
        for r in rows:
            ws.cell(row=r, column=c).fill = openpyxl.styles.PatternFill(fill_type=None)
    for r in rows:
        ws.cell(row=r, column=L).value = isd_vals.get(r)
        ws.cell(row=r, column=L).font = Font(name="Arial", size=10, bold=isd_vals.get(r) == "YES")
        ws.cell(row=r, column=M).value = vals[r][0]
        ws.cell(row=r, column=N).value = vals[r][1]
    return msg


def check(ws) -> list:
    bad = []
    want = {K: "PER_ROW_FILTERS", L: "is_default", M: "AVG_DELTA", N: "POS_SYM"}
    for c, name in want.items():
        if str(ws.cell(row=2, column=c).value or "").strip() != name:
            bad.append(f"col {c} = {ws.cell(row=2, column=c).value!r} (want {name})")
    ys = [c for c in range(1, ws.max_column + 1) if isinstance(ws.cell(row=2, column=c).value, str) and "=" in ws.cell(row=2, column=c).value]
    if ys and ys[0] != O:
        bad.append(f"first yellow header at col {ys[0]} (want O=15)")
    if sum(1 for c in range(1, ws.max_column + 1) if str(ws.cell(row=2, column=c).value or "").strip().lower().startswith("is_default")) != 1:
        bad.append("is_default header count != 1")
    for c in ys:
        for r in range(3, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v not in (None, ""):
                bad.append(f"yellow cell {ws.cell(row=r, column=c).coordinate} holds {v!r} in a TEMPLATE")
                break
    return bad


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for name in TEMPLATES:
        path = SPREAD / name
        wb = openpyxl.load_workbook(str(path))
        for tab in SWITCH_SHEETS:
            if tab in wb.sheetnames:
                print(f"[{name[9:-5]}] {tab}: {fix_tab(wb[tab], args.apply)}")
        if args.apply:
            shutil.copy2(path, ROOT / "backups" / f"before_fixed_cols_{ts}_{name}")
            tmp = path.with_suffix(".tmp.xlsx")
            wb.save(str(tmp))
            wb2 = openpyxl.load_workbook(str(tmp))
            bad = {t: check(wb2[t]) for t in SWITCH_SHEETS if t in wb2.sheetnames}
            bad = {t: b for t, b in bad.items() if b}
            if bad:
                tmp.unlink()
                print(f"[{name}] CHECK FAILED — not saved: {list(bad.items())[:4]}")
                continue
            tmp.replace(path)
            print(f"[{name}] saved, layout check OK")


if __name__ == "__main__":
    main()
