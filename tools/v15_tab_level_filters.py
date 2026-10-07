#!/usr/bin/env python3
"""v15_tab_level_filters — USER 2026-10-01: a filter whose yellow cells cover more than 20% of a tab's switch rows is applied to the ENTIRE tab:
it becomes orange filter rows at the bottom of that tab (tested once in the pilot's row chain) and the pilot stops testing it per switch row (the cell
keeps its yellow status). Unit = FILTER (all its option columns): flagged when ANY option column is yellow (bright FFFF00 by default) on > THRESH of the
tab's white switch rows. Default option = the bold header (== config default), so adding the rows never shifts a baseline.
  python tools/v15_tab_level_filters.py [--templates SPREADSHEETS] [--thresh 0.2] [--basis bright|any]
Writes data/wiring/tab_filters/tab_level_filters.json (read by v15_pilot.py) + tab_level_filters.csv (per column)."""
import argparse
import csv
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
BRIGHT, LIGHT = "FFFFFF00", "FFFFF2CC"
ORANGE = "FFE699"
OUT_JSON = ROOT / "data" / "wiring" / "tab_filters" / "tab_level_filters.json"
OUT_CSV = ROOT / "data" / "wiring" / "tab_filters" / "tab_level_filters.csv"


def rgb(cell):
    f = cell.fill
    return str(f.fgColor.rgb or "").upper() if f is not None and f.fill_type == "solid" else ""


def tab_columns(ws):
    pos = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(row=2, column=c).value or "").strip().upper() == "POS_SYM"), 14)
    return {c: ws.cell(row=2, column=c).value.strip() for c in range(pos + 1, ws.max_column + 1) if isinstance(ws.cell(row=2, column=c).value, str) and "=" in ws.cell(row=2, column=c).value}


def white_rows(ws):
    return [r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value not in (None, "") and ORANGE not in rgb(ws.cell(row=r, column=1))]


def compute(templates: Path, thresh: float, basis: str):
    spec, rows = {}, []
    want = {BRIGHT} if basis == "bright" else {BRIGHT, LIGHT}
    for cs in CAT_SIDES:
        wb = openpyxl.load_workbook(str(templates / f"TEMPLATE_{cs}.xlsx"))
        spec[cs] = {}
        for tab in SWITCH_SHEETS:
            ws = wb[tab] if tab in wb.sheetnames else None
            if ws is None:
                continue
            hd, wr = tab_columns(ws), white_rows(ws)
            if not hd or not wr:
                continue
            by_f = {}
            for c, h in hd.items():
                f, o = h.split("=", 1)
                n = sum(1 for r in wr if rgb(ws.cell(row=r, column=c)) in want)
                bold = bool(ws.cell(row=2, column=c).font is not None and ws.cell(row=2, column=c).font.b)
                by_f.setdefault(f.strip(), []).append({"col": c, "opt": o.strip(), "n": n, "bold": bold})
                rows.append([cs, tab, h, len(wr), n, round(n / len(wr), 3)])
            for f, cols in by_f.items():
                share = max(x["n"] for x in cols) / len(wr)
                if share > thresh:
                    bolds = [x["opt"] for x in cols if x["bold"]]
                    spec[cs].setdefault(tab, {})[f] = {"options": [x["opt"] for x in cols], "default": bolds[0] if len(bolds) == 1 else None, "share": round(share, 3), "n_white_rows": len(wr)}
    return spec, rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--templates", default=str(ROOT / "SPREADSHEETS"))
    ap.add_argument("--thresh", type=float, default=0.2)
    ap.add_argument("--basis", choices=("bright", "any"), default="bright")
    a = ap.parse_args()
    spec, rows = compute(Path(a.templates), a.thresh, a.basis)
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps({"thresh": a.thresh, "basis": a.basis, "unit": "FILTER (all option columns)", "cats": spec}, indent=1))
    flagged = {(cs, t, f) for cs, d in spec.items() for t, fs in d.items() for f in fs}
    with open(OUT_CSV, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["cat_side", "tab", "filter_col", "n_white_rows", "n_yellow_rows", "share", "tab_level"])
        for cs, t, h, n, k, s in rows:
            w.writerow([cs, t, h, n, k, s, "YES" if (cs, t, h.split("=", 1)[0].strip()) in flagged else ""])
    for cs, d in spec.items():
        print(cs, {t: len(fs) for t, fs in d.items()}, "no_default:", sum(1 for fs in d.values() for v in fs.values() if not v["default"]))
    print(OUT_JSON, "filters flagged:", len(flagged))


if __name__ == "__main__":
    main()
