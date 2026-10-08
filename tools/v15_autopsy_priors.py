#!/usr/bin/env python3
"""v15_autopsy_priors — cat_side PRIORS from the autopsy-first bases (USER 2026-10-08 approved #4: "cat_side priors from the
autopsies, free gain, no compute").

Every `{SS}_autopsy_base.json` carries `autopsy.combo_verified`-derived kept switches: engine-verified, cumulative
`gain_after - gain_before` on that sym_side's TEMPLATE-bold base (tools/v15_trade_autopsy_run). This tool folds that
evidence into the daily avg_delta aggregate in the exact partial shape `v15_vector_delta_rebuild.merge_partials` consumes:
  {cat_side: {"TAB\\tswitch\\tSWITCH=value": [delta, ...]}, "_symsides": {cat_side: [sym_side, ...]}, "_source": "autopsy"}
NO-LIES rules: one value per sym_side per (tab, switch=value) (max over its kept entries); only sym_sides WITHOUT a
30D progress file in the current aggregate selection (`--exclude-symsides` / selection json) so a symbol is never
counted twice for the breadth gate; the tab is the template tab that holds the switch row (first of SWITCH_SHEETS);
kept = engine said gain went up and stayed valid. Nothing is annualised, averaged across cat_sides or invented.
  python tools/v15_autopsy_priors.py --bases ~/v15_autopsy_first --out data/autopsy_priors_partial.json [--exclude-selection data/avg_delta_selection.json]
"""
import argparse
import glob
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY",
                 "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
CRYPTO_SUFFIX = ("USDT", "USDC", "USD1", "BUSD", "FDUSD", "TUSD", "DAI")
# never evidence for a cat_side default: venue/mode selectors, the price>0 pseudo-gate, ablation kills (engine-level artefacts, not strategy)
PRIOR_DENY = {"MODE", "SIMPLE_PRICE_GT0_ENABLED", "VENUE", "IS_TRADIER"}
PRIOR_DENY_PREFIX = ("ABLATION_",)


def cat_side_of(ss):
    sym, side = ss.rsplit("_", 1)
    return f"{'CRYPTO' if sym.upper().endswith(CRYPTO_SUFFIX) else 'STOCKS'}_{side.upper()}"


def switch_tabs(cat_side, tdir):
    """switch -> first SWITCH_SHEETS tab holding a row for it, from this cat_side's template (read once)."""
    import openpyxl
    cat, side = cat_side.split("_")
    p = ROOT / tdir / f"TEMPLATE_{cat}_{side}.xlsx"
    out = {}
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    for tab in SWITCH_SHEETS:
        if tab not in wb.sheetnames:
            continue
        for row in wb[tab].iter_rows(min_row=3, max_col=1, values_only=True):
            a = row[0]
            if a and str(a).strip() and str(a).strip() not in out:
                out[str(a).strip()] = tab
    wb.close()
    return out


def fold(bases_dir, exclude, tdir):
    agg = {cs: {} for cs in CAT_SIDES}
    seen = {cs: set() for cs in CAT_SIDES}
    tabs = {}
    stats = {"bases": 0, "excluded_has_board": 0, "no_kept": 0, "used": 0, "no_tab": 0}
    for p in sorted(glob.glob(os.path.join(bases_dir, "*_autopsy_base.json"))):
        ss = os.path.basename(p)[: -len("_autopsy_base.json")]
        stats["bases"] += 1
        if ss in exclude:
            stats["excluded_has_board"] += 1
            continue
        cs = cat_side_of(ss)
        full = Path(p).with_name(f"{ss}_autopsy.json")
        try:
            rep = json.load(open(full))
        except Exception:
            continue
        kept = [v for v in (rep.get("combo_verified") or []) if v.get("kept") and v.get("valid")]
        if not kept:
            stats["no_kept"] += 1
            continue
        if cs not in tabs:
            tabs[cs] = switch_tabs(cs, tdir)
        best = {}
        for v in kept:
            try:
                d = float(v["gain_after"]) - float(v["gain_before"])
            except Exception:
                continue
            if d <= 1e-9:
                continue  # kept means the engine saw a rise; a non-rise is never evidence
            sw = str(v["switch"]).split("=", 1)[0].strip()
            if sw in PRIOR_DENY or sw.startswith(PRIOR_DENY_PREFIX):
                stats["denied"] = stats.get("denied", 0) + 1
                continue
            tab = tabs[cs].get(sw)
            if not tab:
                stats["no_tab"] += 1
                continue
            k = f"{tab}\tswitch\t{str(v['switch']).strip()}"
            best[k] = max(best.get(k, 0.0), d)
        if not best:
            continue
        for k, d in best.items():
            agg[cs].setdefault(k, []).append(round(d, 6))
        seen[cs].add(ss)
        stats["used"] += 1
    out = {cs: agg[cs] for cs in CAT_SIDES}
    out["_symsides"] = {cs: sorted(seen[cs]) for cs in CAT_SIDES}
    out["_source"] = "autopsy combo_verified (engine gain_after-gain_before on the template-bold base); sym_sides without a board in the selection only"
    out["_stats"] = stats
    return out


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--bases", default=os.path.expanduser("~/v15_autopsy_first"))
    ap.add_argument("--out", default=str(ROOT / "data" / "autopsy_priors_partial.json"))
    ap.add_argument("--exclude-selection", default=None, help="avg_delta selection json ({sym_side: {...}}) — those sym_sides already have a board")
    ap.add_argument("--exclude-symsides", default="", help="comma list")
    ap.add_argument("--template-dir", default="SPREADSHEETS/TEMPLATE_FINAL_NORM")
    a = ap.parse_args(argv)
    exclude = {s for s in a.exclude_symsides.split(",") if s}
    if a.exclude_selection and os.path.exists(a.exclude_selection):
        try:
            exclude |= set(json.load(open(a.exclude_selection)).keys())
        except Exception:
            pass
    out = fold(a.bases, exclude, a.template_dir)
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(out, indent=1))
    print(f"[autopsy-priors] {a.out}: " + ", ".join(f"{cs} keys={len(out[cs])} sym_sides={len(out['_symsides'][cs])}" for cs in CAT_SIDES) + f" | {out['_stats']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
