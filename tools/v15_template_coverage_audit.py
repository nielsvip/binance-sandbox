#!/usr/bin/env python3
"""v15_template_coverage_audit — every switch has rows in EXACTLY the cat_side templates it applies to.

USER 2026-10-15: "make sure every switch and all its options have a row in only the cat_side templates it applies to."

Venue a switch applies to is decided by the surface that DEFINES it, never by guesswork:
    config.Config only          -> CRYPTO       (ez_manage / live crypto)
    config_tradier only         -> STOCKS       (tradier_manage / live stocks)
    both                        -> BOTH
    QuickConfig only            -> VECTOR_ONLY  (no live surface reads it: either wire it live or delete it)
    nothing                     -> ORPHAN       (a template row for a switch no code defines)
A `_TRADIER`-suffixed name is STOCKS even when the crypto config mirrors it, because only the tradier surface reads it.

The template IS the side (USER 2026-10-15: "the template dictates the side, so in the _LONG templates only long
trades can take place and all functions apply to long only"). So the question per template is simply: can a trade
placed under this cat_side be influenced by this switch? If yes it needs a row here. The only switches that are
one-direction-by-name are those whose NAME hard-codes it (BB_ENTRY_LONG_THRESHOLD vs BB_ENTRY_SHORT_THRESHOLD are
two separate config fields, each read only when trading that direction); everything else needs a row in BOTH the
LONG and the SHORT template of every venue it applies to.

Findings per switch:
    NO_ROW_ANYWHERE  defined in config/config_tradier/QuickConfig but in NO template  -> never swept; live and the
                     board each fall back to their own class default, which is exactly where they drift apart
    MISSING_ROW      applies to a cat_side that has no row for it     -> rows to add (tools/v15_switch_add.py)
                     gap_kind: VENUE (missing a whole venue) | SIDE (one direction within a venue) | MIXED | ALL
    NO_BOLD_DEFAULT  has a row but no cat_side default was extracted  -> the shared fallback layer is blind to it,
                     so both engines use their own class default even though the row exists
    EXTRA_ROW        has rows in a cat_side it does not apply to      -> rows to remove (tools/v15_switch_remove.py)
    ORPHAN           no config/QuickConfig field at all               -> dead row
    VECTOR_ONLY      QuickConfig-only: the board can sweep what live cannot read
    ONE_OPTION       a single row (no alternative to sweep against)   -> record-only
Read-only. Writes the report to --out (default data/reports/template_coverage_audit.json) and prints the summary.

Usage: python3 tools/v15_template_coverage_audit.py [--out PATH] [--templates-dir DIR] [--quiet]
"""
import argparse
import dataclasses
import json
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
CRYPTO_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT")
STOCKS_SIDES = ("STOCKS_LONG", "STOCKS_SHORT")


def config_surfaces():
    import v12_quick_engine as V
    from config import Config
    from config_tradier import TradierConfig
    cc, ct = Config(), TradierConfig()
    qf = {f.name for f in dataclasses.fields(V.QuickConfig)}
    cck = {k for k in dir(cc) if not k.startswith("_") and not callable(getattr(cc, k, None))}
    ctk = {k for k in dir(ct) if not k.startswith("_") and not callable(getattr(ct, k, None))}
    return qf, cck, ctk


def side_of(key):
    """LONG / SHORT / BOTH from the NAME's own tokens — a switch called *_SHORT_* can only apply to shorts.
    Token match only: 'LONG_STRUCT_EXIT_TF' is long-side, 'LH_HL_FILTER' is neither."""
    toks = set(key.upper().split("_"))
    long_t = bool(toks & {"LONG", "LONGS"})
    short_t = bool(toks & {"SHORT", "SHORTS"})
    if long_t and not short_t:
        return "LONG"
    if short_t and not long_t:
        return "SHORT"
    return "BOTH"


def venue_of(key, qf, cck, ctk):
    """(venue, applies_to_cat_sides) — venue from the defining surface, side from the name."""
    if key.endswith("_TRADIER"):
        venue, sides = ("STOCKS", STOCKS_SIDES) if key in ctk or key in qf else ("ORPHAN", ())
    else:
        in_c, in_t = key in cck, key in ctk
        if in_c and in_t:
            venue, sides = "BOTH", CAT_SIDES
        elif in_c:
            venue, sides = "CRYPTO", CRYPTO_SIDES
        elif in_t:
            venue, sides = "STOCKS", STOCKS_SIDES
        elif key in qf:
            venue, sides = "VECTOR_ONLY", CAT_SIDES
        else:
            venue, sides = "ORPHAN", ()
    s = side_of(key)
    if s != "BOTH":
        sides = tuple(cs for cs in sides if cs.endswith(s))
        venue = f"{venue}/{s}"
    return venue, sides


def template_rows(templates_dir, switch_sheets):
    """{cat_side: {switch: {"rows": n, "tabs": [...], "options": [...]}}}."""
    import openpyxl
    from v15_switch_add import FILES
    out = {}
    for cs, fname in FILES.items():
        path = Path(templates_dir) / fname
        per = defaultdict(lambda: {"rows": 0, "tabs": [], "options": []})
        wb = openpyxl.load_workbook(str(path), read_only=True)
        try:
            for tab in switch_sheets:
                if tab not in wb.sheetnames:
                    continue
                ws = wb[tab]
                for r in ws.iter_rows(min_row=3, values_only=True):
                    if not r or not r[0] or not isinstance(r[0], str):
                        continue
                    name = r[0].strip()
                    if not name or name == "Switch":
                        continue
                    e = per[name]
                    e["rows"] += 1
                    if tab not in e["tabs"]:
                        e["tabs"].append(tab)
                    if len(r) > 1:
                        e["options"].append(r[1])
        finally:
            wb.close()
        out[cs] = {k: dict(v) for k, v in per.items()}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--templates-dir", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--out", default=str(ROOT / "data" / "reports" / "template_coverage_audit.json"))
    ap.add_argument("--quiet", action="store_true")
    a = ap.parse_args()

    from v15_switch_add import SWITCH_SHEETS
    qf, cck, ctk = config_surfaces()
    rows = template_rows(a.templates_dir, SWITCH_SHEETS)
    # enumerate from BOTH ends: every switch that has a row somewhere, AND every switch the code defines. A switch the
    # code reads but no template carries is invisible from the template side — and it is the most dangerous case,
    # because live and the board then each fall back to their own class default.
    import per_sym_settings as ps
    cat_defaults = {cs: set(ps.defaults(cs)) for cs in CAT_SIDES}
    every = sorted({k for cs in CAT_SIDES for k in rows[cs]} | qf | cck | ctk)

    findings = {}
    counts = defaultdict(int)
    for key in every:
        venue, applies = venue_of(key, qf, cck, ctk)
        present = [cs for cs in CAT_SIDES if key in rows[cs]]
        missing = [cs for cs in applies if cs not in present]
        extra = [cs for cs in present if cs not in applies]
        n_rows = {cs: rows[cs][key]["rows"] for cs in present}
        no_bold = [cs for cs in present if key not in cat_defaults[cs]]
        flags = []
        if venue.startswith("ORPHAN"):
            flags.append("ORPHAN")
        if applies and not present:
            flags.append("NO_ROW_ANYWHERE")
        elif missing:
            flags.append("MISSING_ROW")
        if no_bold:
            flags.append("NO_BOLD_DEFAULT")
        if extra:
            flags.append("EXTRA_ROW")
        if venue.startswith("VECTOR_ONLY"):
            flags.append("VECTOR_ONLY")
        if present and max(n_rows.values()) < 2:
            flags.append("ONE_OPTION")
        for f in flags:
            counts[f] += 1
        if flags:
            cry_m, stk_m = set(missing) & set(CRYPTO_SIDES), set(missing) & set(STOCKS_SIDES)
            both_venues = bool(cry_m) and bool(stk_m)
            gap_kind = ("ALL" if applies and not present else
                        "VENUE" if not both_venues and (cry_m == set(CRYPTO_SIDES) or stk_m == set(STOCKS_SIDES)) else
                        "SIDE" if len(missing) == 1 else "MIXED" if missing else "")
            findings[key] = {"venue": venue, "applies_to": list(applies), "present_in": present,
                             "missing_row": missing, "gap_kind": gap_kind, "extra_row": extra,
                             "no_bold_default_in": no_bold, "n_rows": n_rows,
                             "tabs": {cs: rows[cs][key]["tabs"] for cs in present}, "flags": flags}
    report = {"_meta": {"templates_dir": a.templates_dir, "switch_sheets": list(SWITCH_SHEETS),
                        "n_switches": len(every), "counts": dict(counts),
                        "per_cat_side_switch_counts": {cs: len(rows[cs]) for cs in CAT_SIDES}},
              "findings": findings}
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(report, indent=1, sort_keys=True))
    if not a.quiet:
        print(json.dumps(report["_meta"], indent=1))
        for flag in ("ORPHAN", "NO_ROW_ANYWHERE", "MISSING_ROW", "NO_BOLD_DEFAULT", "EXTRA_ROW", "VECTOR_ONLY"):
            hits = [k for k, v in findings.items() if flag in v["flags"]]
            print(f"\n{flag}: {len(hits)}")
            for k in hits[:20]:
                v = findings[k]
                print(f"   {k:46s} venue={v['venue']:14s} kind={v.get('gap_kind',''):6s} has={sorted(v['present_in'])} needs={sorted(v['missing_row'])}")
        print(f"\nreport -> {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
