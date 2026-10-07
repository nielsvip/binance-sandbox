#!/usr/bin/env python3
"""
Monitor for ZECUSDC_LONG/SHORT sheets and verify E column is properly filled.
Runs continuously, checks for new .xlsx files every 30 seconds.
"""

import time
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime

def check_e_column_fill(xlsx_file: Path) -> dict:
    """Check if E column is properly initialized in a sheet."""
    try:
        wb = load_workbook(xlsx_file, data_only=False)

        # Check BASELINE_METRICS
        baseline_sheet = [s for s in wb.sheetnames if "BASELINE_METRICS" in s]
        if not baseline_sheet:
            return {"file": xlsx_file.name, "status": "❌ NO BASELINE_METRICS", "e3": None}

        bm = wb[baseline_sheet[0]]
        baseline_gain = bm.cell(row=2, column=2).value

        if baseline_gain is None:
            return {"file": xlsx_file.name, "status": "❌ BASELINE_METRICS empty", "e3": None, "baseline": None}

        # Check first data sheet (usually ENTRY_REVERSAL_BOUNCE)
        data_sheet = None
        for sheet_name in wb.sheetnames:
            if "BASELINE" not in sheet_name and "FILTER" not in sheet_name and "LEGEND" not in sheet_name:
                data_sheet = sheet_name
                break

        if not data_sheet:
            return {"file": xlsx_file.name, "status": "❌ No data sheet", "e3": None}

        ws = wb[data_sheet]
        e2 = ws.cell(row=2, column=5).value
        e3 = ws.cell(row=3, column=5).value

        wb.close()

        # Check results
        if e2 != "BASELINE":
            return {"file": xlsx_file.name, "status": "❌ E2 not 'BASELINE'", "e2": e2, "e3": e3}

        if e3 is None:
            return {"file": xlsx_file.name, "status": "❌ E3 EMPTY (FIX FAILED)", "baseline": baseline_gain, "e3": None}

        if abs(float(e3) - float(baseline_gain)) < 1e-9:
            return {"file": xlsx_file.name, "status": "✅ E3 CORRECT", "baseline": baseline_gain, "e3": e3}
        else:
            return {"file": xlsx_file.name, "status": "⚠️ E3 mismatch", "baseline": baseline_gain, "e3": e3}

    except Exception as e:
        return {"file": xlsx_file.name, "status": f"❌ ERROR: {str(e)[:50]}", "e3": None}


def monitor_zecusdc():
    """Monitor ZECUSDC sheets for E column fix."""
    xlsx_dir = Path("SPREADSHEETS/V15_V16_CELL_BY_CELL")
    checked = set()

    print(f"[monitor] Starting E column fill monitor for ZECUSDC_LONG/SHORT")
    print(f"[monitor] Will check new files every 30s")
    print("="*80)

    while True:
        try:
            # Find all ZECUSDC files
            zec_files = list(xlsx_dir.glob("ZECUSDC_LONG_bh*_30d_matrix.xlsx")) + \
                       list(xlsx_dir.glob("ZECUSDC_SHORT_bh*_30d_matrix.xlsx"))

            for xlsx_file in sorted(zec_files, key=lambda x: x.stat().st_mtime, reverse=True)[:5]:
                if xlsx_file not in checked:
                    checked.add(xlsx_file)
                    result = check_e_column_fill(xlsx_file)

                    timestamp = datetime.now().strftime("%H:%M:%S")
                    status = result.get("status", "?")
                    baseline = result.get("baseline", "?")
                    e3 = result.get("e3", "?")

                    print(f"[{timestamp}] {result['file']:<50} | {status:<30} | baseline={baseline} e3={e3}")

                    if "CORRECT" in status:
                        print(f"{'':40} 🎉 FIX WORKING!")
                    elif "EMPTY" in status or "FAILED" in status:
                        print(f"{'':40} 🔴 FIX NOT WORKING - ROLLBACK NEEDED")

            time.sleep(30)

        except KeyboardInterrupt:
            print("\n[monitor] Stopped by user")
            break
        except Exception as e:
            print(f"[monitor] Error: {e}")
            time.sleep(30)


if __name__ == "__main__":
    monitor_zecusdc()
