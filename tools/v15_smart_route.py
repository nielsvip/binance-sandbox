#!/usr/bin/env python3
"""v15 SMART ROUTE (USER 2026-10-13): ablation-first, abbreviated worst_first, enlarged PINPOINT, self-improving.

Evidence (real finished boards, mined with `mine` below — never estimated):
  KMNOUSDT_SHORT    3164 rows: 2717 rows G == 0.0, 0 of them carried a positive yellow; 106,656 yellow evals -> 30 positive rows.
  MELANIAUSDT_SHORT 3164 rows: 2576 rows G == 0.0, 0 of them carried a positive yellow;  81,807 yellow evals -> 39 positive rows.
  KMNO started from a 0-trade baseline: entry filters masked every path, so almost every exit/augment/reduce row was
  non-binding and its ~40-200 yellows were pure waste. The gains came from what moved the ledger.

Route (v15_pilot._spec_fill_workbook, kill switch V15_SMART_ROUTE=0):
  ROUND 1a ABLATION   leave-one-out strip of every active entry-blocking filter/gate of the running set (real evals);
                      blocking = stripping it adds trades; the UNMASKED base U = running set minus every blocking filter.
                      U (or a single strip) is adopted only on a real positive delta (greedy law §14.2).
  ROUND 1a PATH PROBE entry/exit/reentry rows evaluated naked on U: rows that move on U are MASKED PATHS (they never had
                      a chance on the masked chain) -> full yellows in round 1 + joint unmask in round 2.
  ROUND 1b ABBREVIATED WORST_FIRST  every row keeps its NAKED eval (§77); yellows run only on HOT rows (knowledge says
                      they ever won / masked path / sacred / essential / new / rotation slice). Every other row is
                      ROUTE-deferred (settled, not hollow) and REVIVED inline when its naked flip binds.
  ROUND 2 PINPOINT    masked-path joint unmask -> final-chain filter re-ablation -> enlarged FOCUS (full yellow matrix on
                      movers, numeric grid, pairwise synergy).
  SELF-IMPROVEMENT    every board appends data/smart_routes/route_ledger.jsonl and merges route_knowledge.json (per
                      cat_side: row wins, blocking filters, masked-path wins, stage seconds). The next board loads it at
                      entry (no look-ahead). `mine` bootstraps/refreshes it from finished xlsx + progress JSONs.

Everything here is pure (no engine import) so it is unit-tested in tests/test_v15_smart_route.py.
"""
from __future__ import annotations

import argparse
import datetime
import glob
import hashlib
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ROUTE_DIR = Path(os.environ.get("V15_ROUTE_DIR") or (ROOT / "data" / "smart_routes"))
KNOWLEDGE_NAME = "route_knowledge.json"
LEDGER_NAME = "route_ledger.jsonl"
SWITCH_TABS = ("STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES")
PATH_TABS = {"ENTRY_REVERSAL_BOUNCE": "entry", "ENTRY_BREAKOUT_CHANNEL": "entry", "ENTRY_CONFIRMATION_GATES": "entry", "EXIT_STRUCTURAL": "exit", "EXIT_VELOCITY": "exit", "REENTRY_WINDOWED": "reentry", "REENTRY_ADAPTIVE": "reentry"}
BLOCKER_TOKENS = ("FILTER", "GATE", "BLOCK", "REQUIRE", "REQUIRED", "CONFIRM", "VETO", "GUARD", "NOLOSS")
NEVER_STRIP_PREFIXES = ("ABLATION_", "FIXED_STOP", "GOLDEN_RULE_", "ULTIMATE_DC_4H", "HARD_STOP", "MAX_AUGMENTS_PER_POSITION", "MAX_ORDER_VALUE", "KILL_SWITCH", "HEDGE")
OFF_WORDS = ("off", "none", "false", "0", "", "disabled", "no")
TF_WORDS = ("1m", "3m", "5m", "15m", "30m", "1h", "2h", "4h", "6h", "12h", "d", "1d", "w", "1w")


def utcnow() -> str:
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def canon_val(v) -> str:
    """Twin of v15_pilot._canon_val (identity = FULL NAME + VALUE, never row/col)."""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    s = str(v).strip() if v is not None else ""
    if s.lower() in ("true", "false", "off", "on", "none"):
        return s.upper()
    try:
        f = float(s)
        return str(int(f)) if f.is_integer() else repr(f)
    except Exception:
        return s


def canon_key(tab: str, switch: str, cand) -> str:
    return f"{str(tab).strip()}!{str(switch).strip().upper()}={canon_val(cand)}"


def canon_filter(key: str, val) -> str:
    return f"{str(key).strip().upper()}={canon_val(val)}"


def cat_side_of(symside: str) -> str:
    """Twin of v15_pilot.map_key_for_symside."""
    s = str(symside).upper()
    base = s[:-5] if s.endswith("_LONG") else s[:-6] if s.endswith("_SHORT") else s
    is_crypto = base.endswith(("USDT", "USDC", "USD1", "BUSD", "FDUSD", "TUSD", "DAI"))
    return f"{'CRYPTO' if is_crypto else 'STOCKS'}_{'LONG' if s.endswith('_LONG') else 'SHORT'}"


def is_off(v) -> bool:
    if isinstance(v, bool):
        return not v
    if v is None:
        return True
    if isinstance(v, (int, float)):
        return float(v) == 0.0
    return str(v).strip().lower() in OFF_WORDS


def is_tf_value(v) -> bool:
    if isinstance(v, bool) or v is None:
        return False
    parts = [p.strip().lower() for p in str(v).split(",") if p.strip()]
    return bool(parts) and all(p in TF_WORDS for p in parts)


def is_blocker_name(key: str) -> bool:
    toks = {t for t in str(key).upper().split("_") if t}
    return bool(toks & set(BLOCKER_TOKENS))


def off_value_for(key: str, cur):
    """The OFF value that removes this filter/gate, or None when it has no unambiguous off (numeric thresholds)."""
    if isinstance(cur, bool):
        return False if cur else None
    if isinstance(cur, str):
        low = cur.strip().lower()
        if low in ("true", "on", "yes"):
            return False
        if is_tf_value(cur):
            return "OFF"
    return None


def ablation_candidates(running: dict, defaults: dict, filt_names=None, knowledge_filters: dict | None = None, cap: int = 240) -> list:
    """Active entry-blocking filters/gates of the effective running set -> [(key, cur, off_value)], most promising first.
    A key qualifies when it is a template filter (orange row / yellow header base) or carries a blocker token, it is ON
    (bool True / a timeframe), and it is not a safety catch (NEVER_STRIP_PREFIXES — stops, ablation artefacts, caps) or a
    portfolio-level hedge rule. Order: proven unmask wins, proven blockers, template filters, blocker-token names."""
    filt_names = set(filt_names or ())
    eff = dict(defaults or {})
    eff.update(running or {})
    out = []
    for k in sorted(eff):
        ks = str(k).strip()
        if ks.startswith(NEVER_STRIP_PREFIXES) or "HEDGE" in ks.upper():
            continue
        if ks not in filt_names and not is_blocker_name(ks):
            continue
        cur = eff[k]
        if is_off(cur):
            continue
        off = off_value_for(ks, cur)
        if off is None:
            continue
        out.append((ks, cur, off))
    kf = knowledge_filters or {}

    def _rank(t):
        e = kf.get(canon_filter(t[0], t[1])) or {}
        return (-int(e.get("unmask_win") or 0), -int(e.get("blocking") or 0), 0 if t[0] in filt_names else 1, 0 if is_blocker_name(t[0]) else 1, t[0])

    out.sort(key=_rank)
    return out[: max(0, int(cap))]


def _num(v, default=None):
    try:
        if v is None or isinstance(v, bool):
            return default
        return float(v)
    except Exception:
        return default


def classify_strip(base_res: dict, strip_res: dict, trade_floor: int = 10) -> dict:
    """One leave-one-out strip vs the running base -> {d_gain, d_trades, d_tim, valid, blocking, helps}.
    blocking: stripping the filter ADDS trades (>= max(2, 10% of base), or any trade when the base is under the floor) —
    it was masking entry/exit paths. helps: a valid strip with a real positive gain delta."""
    bg, bt = _num((base_res or {}).get("gain_pct"), 0.0), int(_num((base_res or {}).get("trades"), 0) or 0)
    sg, st = _num((strip_res or {}).get("gain_pct")), int(_num((strip_res or {}).get("trades"), 0) or 0)
    valid = bool((strip_res or {}).get("valid")) and sg is not None
    d_gain = (sg - bg) if sg is not None else None
    d_trades = st - bt
    btim, stim = _num((base_res or {}).get("tim_pct")), _num((strip_res or {}).get("tim_pct"))
    d_tim = (stim - btim) if (stim is not None and btim is not None) else None
    need = max(2, int(round(0.10 * bt)))
    blocking = d_trades >= need or (bt < trade_floor and d_trades > 0)
    helps = valid and d_gain is not None and d_gain > 1e-9
    return {"d_gain": d_gain, "d_trades": d_trades, "d_tim": d_tim, "valid": valid, "blocking": bool(blocking), "helps": bool(helps)}


def unmask_order(strips: dict) -> list:
    """Keys to strip for the UNMASKED base: every blocking or helping filter, most trades unlocked first."""
    keys = [k for k, v in (strips or {}).items() if v.get("blocking") or v.get("helps")]
    keys.sort(key=lambda k: (-(strips[k].get("d_trades") or 0), -(strips[k].get("d_gain") or 0.0), k))
    return keys


def is_enabling(cand, running_val) -> bool:
    """A row that turns a path ON (or onto a timeframe) while the running value is OFF."""
    return (not is_off(cand)) and is_off(running_val)


def probe_rows(rows: list, running: dict, defaults: dict, knowledge_rows: dict | None = None, cap: int = 600) -> list:
    """Path rows for the probe on the unmasked base. rows = [(tab, rr, switch, cand)] (white rows only).
    Order: knowledge masked-path wins, row wins, then ENABLING rows (paths the masked chain never opened), then the rest."""
    kr = knowledge_rows or {}
    eff = dict(defaults or {})
    eff.update(running or {})
    out = []
    for tab, rr, sw, cand in rows:
        if tab not in PATH_TABS:
            continue
        cur = eff.get(str(sw).strip())
        if canon_val(cur) == canon_val(cand):
            continue
        e = kr.get(canon_key(tab, sw, cand)) or {}
        out.append((-int(e.get("masked_win") or 0), -int(e.get("pos") or 0), 0 if is_enabling(cand, cur) else 1, tab, rr, sw, cand))
    out.sort(key=lambda t: t[:3] + (SWITCH_TABS.index(t[3]) if t[3] in SWITCH_TABS else 99, t[4]))
    return [(t[3], t[4], t[5], t[6]) for t in out[: max(0, int(cap))]]


def row_is_hot(entry: dict | None) -> bool:
    """Knowledge says this exact SWITCH=cand ever earned its yellows: it promoted, carried a positive yellow, or won as
    a masked path on some board of this cat_side."""
    if not isinstance(entry, dict):
        return False
    return int(entry.get("promoted") or 0) > 0 or int(entry.get("yellow_pos") or 0) > 0 or int(entry.get("masked_win") or 0) > 0


def rotation_hit(rr: int, tab: str, seq: int, every: int) -> bool:
    """Deterministic 1/every slice (same input -> same answer within a round; rotates with seq)."""
    if every <= 1:
        return True
    h = int(hashlib.md5(f"{tab}|{int(rr)}|{int(seq)}".encode()).hexdigest()[:8], 16)
    return (h % int(every)) == 0


def route_row_policy(knowledge_rows: dict | None, tab: str, switch, cand, hdrs, rr: int, seq: int, every: int, exempt: bool, masked: set | None = None) -> tuple:
    """Pure: -> (mode, kept_hdrs, deferred_hdrs, meta). mode 'full' | 'abbrev'.
    full: exempt (sacred/essential/new/coarse-expand), no yellows, masked path, hot by knowledge, rotation slice.
    abbrev: naked only now; every yellow ROUTE-deferred (settled), revived inline when the naked flip binds."""
    hdrs = list(hdrs or [])
    if exempt or not hdrs:
        return "full", hdrs, [], {"why": "exempt" if exempt else "no_yellows"}
    ck = canon_key(tab, switch, cand)
    if masked and ck in masked:
        return "full", hdrs, [], {"why": "ROUTE_MASKED_PATH"}
    e = (knowledge_rows or {}).get(ck)
    if row_is_hot(e):
        return "full", hdrs, [], {"why": f"ROUTE_HOT promoted={e.get('promoted')} yellow_pos={e.get('yellow_pos')} masked_win={e.get('masked_win')} n={e.get('n')}"}
    if rotation_hit(rr, tab, seq, every):
        return "full", hdrs, [], {"why": f"ROUTE_ROTATE 1/{every}"}
    return "abbrev", [], hdrs, {"why": "ROUTE_ABBREV: naked first, yellows only if the naked flip binds" + (f" (n={e.get('n')} pos={e.get('pos')})" if e else " (no knowledge)")}


def should_revive(naked_delta, binding, revive_min: float = -1.0, sacred: bool = False) -> bool:
    """ROUTE-REVIVED: the naked flip moved the ledger on THIS sym_side (binding) and is not deeply negative.
    A non-binding flip (same ledger fingerprint, delta 0) makes switch+filter == filter alone, which the filter's own
    orange row already measures — its yellows stay deferred."""
    if naked_delta is None:
        return False
    try:
        nd = float(naked_delta)
    except Exception:
        return False
    moved = abs(nd) > 1e-9 or bool(binding)
    return moved and (nd > float(revive_min) or sacred)


def empty_knowledge() -> dict:
    return {"version": 1, "updated": None, "boards": {}, "cat_sides": {}}


def load_knowledge(path: Path | None = None) -> dict:
    p = Path(path or (ROUTE_DIR / KNOWLEDGE_NAME))
    try:
        if p.exists():
            j = json.loads(p.read_text())
            if isinstance(j, dict) and isinstance(j.get("cat_sides"), dict):
                j.setdefault("boards", {})
                return j
    except Exception:
        pass
    return empty_knowledge()


def knowledge_for(knowledge: dict, cat_side: str) -> dict:
    c = ((knowledge or {}).get("cat_sides") or {}).get(cat_side) or {}
    return {"rows": c.get("rows") or {}, "filters": c.get("filters") or {}, "boards": c.get("boards") or {}}


def _atomic_write(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    with open(tmp, "w") as fh:
        fh.write(text)
        fh.flush()
        os.fsync(fh.fileno())
    os.replace(tmp, path)


def save_knowledge(knowledge: dict, path: Path | None = None):
    knowledge["updated"] = utcnow()
    _atomic_write(Path(path or (ROUTE_DIR / KNOWLEDGE_NAME)), json.dumps(knowledge, indent=1, sort_keys=True, default=str))


def board_id(record: dict) -> str:
    return f"{record.get('symside')}|{record.get('result_key') or record.get('source') or ''}"


def merge_record(knowledge: dict, record: dict) -> bool:
    """Fold one board record into the knowledge. Idempotent per board id (same board twice = no double count).
    Returns True when merged."""
    bid = board_id(record)
    if bid in (knowledge.get("boards") or {}):
        return False
    cat = record.get("cat_side") or cat_side_of(record.get("symside") or "")
    c = knowledge.setdefault("cat_sides", {}).setdefault(cat, {"rows": {}, "filters": {}, "boards": {"n": 0}})
    rows, filters, boards = c.setdefault("rows", {}), c.setdefault("filters", {}), c.setdefault("boards", {"n": 0})
    for ck, r in (record.get("rows") or {}).items():
        e = rows.setdefault(ck, {"n": 0, "pos": 0, "neg": 0, "zero": 0, "promoted": 0, "yellow_pos": 0, "masked": 0, "masked_win": 0, "sum_pos": 0.0, "sum_neg": 0.0})
        e["n"] += 1
        d = _num(r.get("delta"))
        if d is not None:
            if d > 1e-9:
                e["pos"] += 1
                e["sum_pos"] = round(e["sum_pos"] + d, 6)
            elif d < -1e-9:
                e["neg"] += 1
                e["sum_neg"] = round(e["sum_neg"] + d, 6)
            else:
                e["zero"] += 1
        e["promoted"] += 1 if r.get("promoted") else 0
        e["yellow_pos"] += 1 if (r.get("pos_yellows") or []) else 0
        e["masked"] += 1 if r.get("masked") else 0
        e["masked_win"] += 1 if r.get("masked_win") else 0
        for h in r.get("pos_yellows") or []:
            fe = filters.setdefault(str(h).upper(), {"n": 0, "blocking": 0, "strip_pos": 0, "unmask_win": 0, "yellow_pos": 0, "reapply_pos": 0})
            fe["yellow_pos"] += 1
    for fk, f in (record.get("filters") or {}).items():
        fe = filters.setdefault(fk, {"n": 0, "blocking": 0, "strip_pos": 0, "unmask_win": 0, "yellow_pos": 0, "reapply_pos": 0})
        fe["n"] += 1
        fe["blocking"] += 1 if f.get("blocking") else 0
        fe["strip_pos"] += 1 if f.get("helps") else 0
        fe["unmask_win"] += 1 if f.get("unmask_win") else 0
        fe["reapply_pos"] += 1 if f.get("reapply_pos") else 0
    boards["n"] = int(boards.get("n") or 0) + 1
    for k in ("baseline_gain", "final_gain", "bh"):
        v = _num(record.get(k))
        if v is not None:
            boards[f"sum_{k}"] = round(float(boards.get(f"sum_{k}") or 0.0) + v, 6)
    if _num(record.get("final_gain")) is not None and _num(record.get("bh")) is not None:
        boards["beat_bh"] = int(boards.get("beat_bh") or 0) + (1 if float(record["final_gain"]) > float(record["bh"]) else 0)
    for st, secs in (record.get("stage_secs") or {}).items():
        s = boards.setdefault("stage_secs", {}).setdefault(st, {"n": 0, "sum": 0.0})
        s["n"] += 1
        s["sum"] = round(s["sum"] + float(secs or 0.0), 3)
    knowledge.setdefault("boards", {})[bid] = {"at": utcnow(), "cat_side": cat, "final_gain": record.get("final_gain"), "source": record.get("source")}
    return True


def append_ledger(record: dict, path: Path | None = None):
    p = Path(path or (ROUTE_DIR / LEDGER_NAME))
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "a") as fh:
        fh.write(json.dumps(record, default=str) + "\n")


def record_from_progress(symside: str, progress: dict, source: str = "") -> dict:
    """Board record from a pilot progress JSON (the source of truth): done rows, smart_route, focus, endgame, final."""
    prog = progress or {}
    rec = {"symside": symside, "cat_side": cat_side_of(symside), "source": source or "progress", "at": utcnow(), "rows": {}, "filters": {}, "baseline_gain": prog.get("initial_baseline_gain", prog.get("baseline_gain")), "final_gain": prog.get("final_gain", prog.get("cumulative_gain")), "bh": prog.get("bh", prog.get("bh_pct"))}
    sr = prog.get("smart_route") or {}
    masked = set(sr.get("masked_paths") or [])
    masked_won = set(sr.get("masked_wins") or [])
    for k, v in (prog.get("done") or {}).items():
        if not isinstance(v, dict) or "!" not in str(k) or ":" not in str(k):
            continue
        tab, rest = str(k).split("!", 1)
        if tab not in SWITCH_TABS:
            continue
        swc = rest.split(":", 1)[1]
        if swc.endswith(":REPASS") or "=" not in swc:
            continue
        sw, cand = swc.split("=", 1)
        ck = canon_key(tab, sw, cand)
        ys = v.get("yellows") or {}
        pos_y = sorted(str(h) for h, d in ys.items() if _num(d) is not None and float(d) > 1e-9)
        rec["rows"][ck] = {"delta": v.get("delta"), "naked": v.get("naked_delta"), "promoted": bool(v.get("promoted")), "pos_yellows": pos_y, "masked": ck in masked, "masked_win": ck in masked_won}
    for fk, f in (sr.get("strips") or {}).items():
        rec["filters"][fk] = {"blocking": bool(f.get("blocking")), "helps": bool(f.get("helps")), "unmask_win": fk in set(sr.get("unmask_adopted") or [])}
    eg = prog.get("endgame") or {}
    prev = _num(eg.get("strip_gain"))
    final_ov = prog.get("cumulative_overrides") or {}
    for step in eg.get("reapplied") or []:
        try:
            f, g = step[0], _num(step[1])
        except Exception:
            continue
        if g is None:
            continue
        fk = canon_filter(f, final_ov.get(f))
        rec["filters"].setdefault(fk, {})["reapply_pos"] = prev is not None and g > prev + 1e-9
        prev = g
    rec["stage_secs"] = dict(sr.get("stage_secs") or {})
    rec["result_key"] = hashlib.md5(json.dumps(sorted((str(k), str(v)) for k, v in (prog.get("cumulative_overrides") or {}).items())).encode()).hexdigest()[:16]
    return rec


def record_from_workbook(xlsx_path: str) -> dict | None:
    """Board record from a finished *_30d_matrix.xlsx (read-only, never saved): every SWITCH tab row's G and yellows,
    ENDGAME reapply steps, DIAGNOSE row recommendations and FOCUS/ROUTE promotions. Rows with numeric G only."""
    try:
        import openpyxl
    except Exception:
        return None
    name = Path(xlsx_path).name
    symside = "_".join(name.split("_")[:2])
    try:
        wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    except Exception:
        return None
    rec = {"symside": symside, "cat_side": cat_side_of(symside), "source": f"xlsx:{name}", "at": utcnow(), "rows": {}, "filters": {}, "stats": {"rows": 0, "g_numeric": 0, "g_pos": 0, "g_zero": 0, "yellow_cells": 0, "yellow_pos": 0, "zero_rows_with_pos_yellow": 0}}
    try:
        for tab in SWITCH_TABS:
            if tab not in wb.sheetnames:
                continue
            it = wb[tab].iter_rows(min_row=2, values_only=True)
            try:
                hdr = next(it)
            except StopIteration:
                continue
            ycols = [(i, str(h)) for i, h in enumerate(hdr) if isinstance(h, str) and "=" in h]
            for r in it:
                if not r or r[0] in (None, "") or len(r) < 7:
                    continue
                rec["stats"]["rows"] += 1
                g = _num(r[6])
                ys = {h: r[i] for i, h in ycols if i < len(r) and _num(r[i]) is not None}
                pos_y = sorted(h for h, d in ys.items() if float(d) > 1e-9)
                rec["stats"]["yellow_cells"] += len(ys)
                rec["stats"]["yellow_pos"] += len(pos_y)
                if g is None:
                    continue
                rec["stats"]["g_numeric"] += 1
                if g > 1e-9:
                    rec["stats"]["g_pos"] += 1
                elif abs(g) <= 1e-9:
                    rec["stats"]["g_zero"] += 1
                    rec["stats"]["zero_rows_with_pos_yellow"] += 1 if pos_y else 0
                c = r[2]
                promoted = g > 1e-9 and c not in (None, "")
                rec["rows"][canon_key(tab, r[0], r[1])] = {"delta": g, "promoted": promoted, "pos_yellows": pos_y}
        if "ENDGAME_FILTER_CYCLE" in wb.sheetnames:
            prev = None
            for r in wb["ENDGAME_FILTER_CYCLE"].iter_rows(values_only=True):
                if r and r[0] == "REAPPLY" and isinstance(r[1], int):
                    ga = _num(r[4])
                    if ga is not None:
                        rec["filters"].setdefault(canon_filter(r[2], r[3]), {})["reapply_pos"] = prev is not None and ga > prev + 1e-9
                        prev = ga
                elif r and r[0] == "STRIP" and isinstance(r[2], str) and "->" in r[2]:
                    prev = _num(r[2].split("->")[-1])
        if "DIAGNOSE_REPAIR" in wb.sheetnames:
            on = False
            for r in wb["DIAGNOSE_REPAIR"].iter_rows(values_only=True):
                if r and r[0] == "ROW RECOMMENDATIONS":
                    on = True
                    continue
                if on and (not r or r[0] in (None, "")):
                    break
                if on and len(r) > 12 and r[1] in SWITCH_TABS and _num(r[11]) is not None and float(r[11]) > 1e-9 and r[12] is True:
                    e = rec["rows"].setdefault(canon_key(r[1], r[0], r[3]), {"delta": None, "promoted": False, "pos_yellows": []})
                    e["diag_win"] = True
                    e["promoted"] = True
        for tab in ("FOCUS_ROUND", "ROUTE_PLAN"):
            if tab in wb.sheetnames:
                for r in wb[tab].iter_rows(values_only=True):
                    if r and isinstance(r[1], str) and "=" in r[1] and _num(r[2]) is not None and float(r[2]) > 1e-9 and "!" in str(r[1]):
                        tb, swc = str(r[1]).split("!", 1)
                        sw, cd = swc.split("=", 1)
                        e = rec["rows"].setdefault(canon_key(tb, sw, cd.split("@")[0]), {"delta": None, "promoted": False, "pos_yellows": []})
                        e["promoted"] = True
                        e["masked_win"] = str(r[0]).startswith("masked")
        bm = [s for s in wb.sheetnames if s.endswith("_BASELINE_METRICS")]
        if bm:
            for r in wb[bm[0]].iter_rows(values_only=True):
                if r and r[0] == "gain_pct":
                    rec["baseline_gain"] = _num(r[1])
                elif r and r[0] == "bh_pct":
                    rec["bh"] = _num(r[1])
        if "DIAGNOSE_REPAIR" in wb.sheetnames:
            for r in wb["DIAGNOSE_REPAIR"].iter_rows(max_row=4, values_only=True):
                if r and r[0] == "gain":
                    rec["final_gain"] = _num(r[2])
    finally:
        wb.close()
    rec["result_key"] = hashlib.md5(name.encode()).hexdigest()[:16]
    return rec


def mine(xlsx_globs: list, progress_globs: list, knowledge_path: Path, ledger_path: Path | None = None, dry: bool = False, rebuild: bool = False) -> dict:
    """Fold boards into the knowledge. rebuild=True starts from empty (daily S1 chain: one clean union of every finished
    board, no drift); otherwise incremental and idempotent per board id."""
    know = empty_knowledge() if rebuild else load_knowledge(knowledge_path)
    report = {"merged": 0, "skipped_dup": 0, "failed": 0, "boards": []}
    seen = set()
    paths = []
    for g in progress_globs or []:
        paths += [("p", p) for p in sorted(glob.glob(os.path.expanduser(g)))]
    for g in xlsx_globs or []:
        paths += [("x", p) for p in sorted(glob.glob(os.path.expanduser(g))) if not Path(p).name.startswith(("TEMPLATE", "~$"))]
    from_progress = set()
    for kind, p in paths:
        if p in seen:
            continue
        seen.add(p)
        try:
            if kind == "x":
                if "_".join(Path(p).name.split("_")[:2]) in from_progress:
                    continue  # the progress JSON (source of truth) of this sym_side was already folded
                rec = record_from_workbook(p)
            else:
                prog = json.loads(Path(p).read_text())
                if not isinstance(prog, dict) or prog.get("final_gain") is None or prog.get("needs_redo"):
                    continue
                rec = record_from_progress(Path(p).name.replace("_v14_progress.json", ""), prog, source=f"progress:{p}")
                from_progress.add(rec["symside"])
        except Exception:
            rec = None
        if not rec:
            report["failed"] += 1
            continue
        if merge_record(know, rec):
            report["merged"] += 1
            if ledger_path is not None and not dry:
                append_ledger(rec, ledger_path)
        else:
            report["skipped_dup"] += 1
        report["boards"].append({"source": rec.get("source"), "cat_side": rec.get("cat_side"), "rows": len(rec.get("rows") or {}), "stats": rec.get("stats")})
    if not dry:
        save_knowledge(know, knowledge_path)
    report["knowledge"] = {c: {"rows": len(v.get("rows") or {}), "hot_rows": sum(1 for e in (v.get("rows") or {}).values() if row_is_hot(e)), "filters": len(v.get("filters") or {}), "boards": (v.get("boards") or {}).get("n")} for c, v in (know.get("cat_sides") or {}).items()}
    return report


def render_report(knowledge: dict) -> str:
    lines = ["# SMART ROUTE KNOWLEDGE", "", f"updated {knowledge.get('updated')} — boards {len(knowledge.get('boards') or {})}", ""]
    for cat, c in sorted((knowledge.get("cat_sides") or {}).items()):
        rows = c.get("rows") or {}
        hot = sorted(((k, e) for k, e in rows.items() if row_is_hot(e)), key=lambda t: (-(t[1].get("sum_pos") or 0.0), t[0]))
        filt = c.get("filters") or {}
        blk = sorted(((k, e) for k, e in filt.items() if e.get("blocking")), key=lambda t: (-t[1].get("blocking", 0), t[0]))
        b = c.get("boards") or {}
        lines += [f"## {cat}", "", f"boards {b.get('n')} · rows known {len(rows)} · HOT rows {len(hot)} ({(100.0 * len(hot) / max(1, len(rows))):.1f}%) · blocking filters {len(blk)} · beat B&H {b.get('beat_bh', 0)}", ""]
        lines += ["| hot row | n | pos | promoted | yellow_pos | masked_win | sum_pos |", "|---|---|---|---|---|---|---|"]
        for k, e in hot[:40]:
            lines.append(f"| `{k}` | {e.get('n')} | {e.get('pos')} | {e.get('promoted')} | {e.get('yellow_pos')} | {e.get('masked_win')} | {e.get('sum_pos')} |")
        if blk:
            lines += ["", "| blocking filter | n | blocking | strip_pos | unmask_win |", "|---|---|---|---|---|"]
            for k, e in blk[:30]:
                lines.append(f"| `{k}` | {e.get('n')} | {e.get('blocking')} | {e.get('strip_pos')} | {e.get('unmask_win')} |")
        st = b.get("stage_secs") or {}
        if st:
            lines += ["", "| stage | boards | mean s |", "|---|---|---|"]
            for k, s in sorted(st.items()):
                lines.append(f"| {k} | {s.get('n')} | {(s.get('sum', 0.0) / max(1, s.get('n', 1))):.1f} |")
        lines.append("")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="v15 SMART ROUTE knowledge: mine finished boards, report")
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("mine", help="fold finished xlsx boards + final progress JSONs into route_knowledge.json")
    m.add_argument("--xlsx", action="append", default=[], help="glob of finished *_30d_matrix.xlsx (repeatable)")
    m.add_argument("--progress", action="append", default=[], help="glob of *_v14_progress.json with a top-level final_gain (repeatable)")
    m.add_argument("--knowledge", default=str(ROUTE_DIR / KNOWLEDGE_NAME))
    m.add_argument("--ledger", default=str(ROUTE_DIR / LEDGER_NAME))
    m.add_argument("--dry-run", action="store_true")
    m.add_argument("--rebuild", action="store_true", help="start from empty knowledge (daily chain), else incremental")
    r = sub.add_parser("report", help="markdown report of the knowledge")
    r.add_argument("--knowledge", default=str(ROUTE_DIR / KNOWLEDGE_NAME))
    r.add_argument("--out", default=None)
    a = ap.parse_args(argv)
    if a.cmd == "mine":
        xl = a.xlsx or [str(ROOT / "SPREADSHEETS" / "V15_V16_CELL_BY_CELL" / "*_30d_matrix.xlsx"), str(ROOT / "SPREADSHEETS" / "*_30d_matrix.xlsx")]
        pg = a.progress or [str(ROOT / "data" / "reports" / "lifecycle_pilot" / "*_v14_progress.json"), "~/v15_run*/progress/*_v14_progress.json", "~/v15_run*/*_v14_progress.json"]
        rep = mine(xl, pg, Path(a.knowledge), None if a.rebuild else Path(a.ledger), dry=a.dry_run, rebuild=a.rebuild)
        print(json.dumps({k: v for k, v in rep.items() if k != "boards"}, indent=1))
        for b in rep["boards"]:
            print(json.dumps(b))
        return 0
    know = load_knowledge(Path(a.knowledge))
    txt = render_report(know)
    if a.out:
        _atomic_write(Path(a.out), txt)
    print(txt)
    return 0


if __name__ == "__main__":
    sys.exit(main())
