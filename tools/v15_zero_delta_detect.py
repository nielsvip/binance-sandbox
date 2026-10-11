"""v15_zero_delta_detect — sidecar: scan progress JSON + manifest, paint published XLSX.

USER 2026-10-11: multiple same/0 deltas -> red flag + red cell + red tab, fixed
ASAP, never slowing calculations. This runs OUTSIDE the hot loop (cron/manual);
the live pilot also evaluates the same rules inside _paint_tab_status.

Identity is canonical TAB!SWITCH=value (row numbers shift daily — never identity).
Within ONE file, row numbers order the SAME_RUN scan (ordering, not identity).

Usage:
  python3 tools/v15_zero_delta_detect.py --sym-side ENAUSDC_LONG [--no-paint]
  python3 tools/v15_zero_delta_detect.py --all [--no-paint]
Exit 1 when any RED alert (cron-alertable).
"""
from __future__ import annotations
import argparse
import glob
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
PROGRESS_DIR = Path(os.environ["V15_PROGRESS_DIR"]) if os.environ.get("V15_PROGRESS_DIR") else ROOT / "data" / "reports" / "lifecycle_pilot"
CELL_DIR = ROOT / "SPREADSHEETS" / "V15_V16_CELL_BY_CELL"


def _load(symside: str):
    p = PROGRESS_DIR / f"{symside}_v14_progress.json"
    if not p.exists():
        return None, f"no progress {p.name}"
    try:
        d = json.loads(p.read_text())
    except Exception as e:
        return None, f"bad json: {e}"
    return d, ""


def scan_symside(symside: str):
    from v15_pilot import SWITCH_SHEETS, _canon_key_of_done, _zero_bh_eval, _zero_rules_eval
    d, err = _load(symside)
    if d is None:
        return [], err
    from v15_pilot import _rec_measured
    done = d.get("done") or {}
    order = {s: i for i, s in enumerate(SWITCH_SHEETS)}
    per_tab: dict = {}
    skipped_tab: dict = {}
    for k, rec in done.items():
        ck = _canon_key_of_done(k)
        if ck is None:
            continue
        tab = ck.split("!", 1)[0]
        sw = ck.split("!", 1)[1].split("=", 1)[0]
        if not _rec_measured(rec):
            skipped_tab.setdefault(tab, {}).setdefault(sw, []).append(ck)
            continue
        try:
            row = int(str(k).split("!", 1)[1].split(":", 1)[0])
        except Exception:
            row = 10**9
        g = rec.get("delta")
        g = float(g) if isinstance(g, (int, float)) else None
        per_tab.setdefault(tab, []).append((order.get(tab, 999), row, ck, sw, g))
    alerts = []
    for tab, rows in per_tab.items():
        rows.sort()
        ordered = [(ck, g) for _, _, ck, _, g in rows]
        by_sw: dict = {}
        for _, _, ck, sw, g in rows:
            by_sw.setdefault(sw, []).append((ck, g))
        alerts.extend(_zero_rules_eval(tab, ordered, by_sw, skipped_tab.get(tab)))
    for tab, sk in skipped_tab.items():
        if tab not in per_tab:
            alerts.extend(_zero_rules_eval(tab, [], {}, sk))
    gain = d.get("cumulative_gain")
    bh = d.get("bh")
    man = None
    for mp in sorted(CELL_DIR.glob(f"{symside}_*_manifest.json"), key=lambda p: p.stat().st_mtime, reverse=True)[:1]:
        try:
            man = json.loads(mp.read_text()).get("metrics") or {}
            gain = man.get("gain_pct", gain)
            bh = man.get("bh", bh)
        except Exception:
            pass
    zb = _zero_bh_eval(gain, bh)
    if zb is not None:
        alerts.append(zb)
    return alerts, ""


def paint_symside(symside: str, alerts: list) -> str:
    from v15_pilot import _atomic_save, _write_zero_alerts_sheet
    import openpyxl
    cands = [p for p in CELL_DIR.glob(f"{symside}_*matrix.xlsx")]
    if not cands:
        return "no published matrix"
    x = max(cands, key=lambda p: p.stat().st_mtime)
    wb = openpyxl.load_workbook(str(x))
    try:
        tabs = {a.get("tab") for a in alerts if a.get("tab") in wb.sheetnames}
        for t in tabs:
            wb[t].sheet_properties.tabColor = "FF0000"
        if any(a.get("tab") == "BASELINE_METRICS" for a in alerts):
            for s in wb.sheetnames:
                if "BASELINE" in s:
                    wb[s].sheet_properties.tabColor = "FF0000"
        _write_zero_alerts_sheet(wb, alerts)
        _atomic_save(wb, x)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return f"painted {x.name}"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sym-side", default="")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--no-paint", action="store_true")
    a = ap.parse_args()
    if a.all:
        syms = sorted(p.name[:-len("_v14_progress.json")] for p in PROGRESS_DIR.glob("*_v14_progress.json"))
    elif a.sym_side:
        syms = [a.sym_side.strip()]
    else:
        ap.error("--sym-side or --all required")
    red = 0
    for ss in syms:
        alerts, err = scan_symside(ss)
        if err:
            print(f"{ss}: SKIP {err}", flush=True)
            continue
        if not alerts:
            print(f"{ss}: OK no alerts", flush=True)
            continue
        red += 1
        rules = {}
        for al in alerts:
            rules.setdefault(al["rule"], []).append(al["tab"])
        print(f"{ss}: RED " + " ".join(f"{r}x{len(set(t))}" for r, t in sorted(rules.items())), flush=True)
        for al in alerts[:12]:
            print(f"  [{al['rule']}] {al.get('tab')}: {al.get('detail')}", flush=True)
        if not a.no_paint:
            try:
                print(f"  {paint_symside(ss, alerts)}", flush=True)
            except Exception as e:
                print(f"  paint FAILED: {e}", flush=True)
    return 1 if red else 0


if __name__ == "__main__":
    sys.exit(main())
