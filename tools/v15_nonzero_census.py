#!/usr/bin/env python3
"""v15_nonzero_census — USER acceptance metric (2026-09-29): count NON-ZERO,
NON-REPEAT delta cells across finished sheets. A cell counts when |value|>1e-9
AND it is not an echo (same value repeated across a row's yellow cells counts
ONCE). G/F and yellow columns are counted; E/C are not deltas."""
import sys, glob
import openpyxl

SW = ["STDEV_SLOPE_SIZING","ENTRY_REVERSAL_BOUNCE","ENTRY_BREAKOUT_CHANNEL","ENTRY_CONFIRMATION_GATES","EXIT_STRUCTURAL","EXIT_VELOCITY","REENTRY_WINDOWED","REENTRY_ADAPTIVE","AUGMENT_TREND","AUGMENT_RISK_SIZING","REDUCE_PROFIT_LOCK","REDUCE_SIGNAL_RATER","GLOBAL_RISK_GATES"]

def census(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    nz = 0
    for sn in wb.sheetnames:
        if sn not in SW: continue
        ws = wb[sn]
        for row in ws.iter_rows(min_row=3):
            v = [c.value for c in row] + [None]*70
            if v[0] in (None, ""): continue
            row_counted = set()
            for x in (v[5], v[6]):  # F, G
                if isinstance(x, (int, float)) and abs(x) > 1e-9:
                    key = round(float(x), 10)
                    if key not in row_counted:
                        row_counted.add(key); nz += 1
            for x in v[11:70]:      # yellows: same-value repeats in a row count once
                if isinstance(x, (int, float)) and abs(x) > 1e-9:
                    key = round(float(x), 10)
                    if key not in row_counted:
                        row_counted.add(key); nz += 1
    wb.close()
    return nz

total = 0
for p in sorted(set(sum((glob.glob(g) for g in sys.argv[1:]), []))):
    n = census(p)
    total += n
    print(f"{n:6d}  {p.split('/')[-1]}")
print(f"TOTAL NON-ZERO NON-REPEAT CELLS: {total}")
