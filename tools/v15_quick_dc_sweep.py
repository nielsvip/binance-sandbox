#!/usr/bin/env python3
"""v15_quick_dc_sweep — quick DC-channel stop/target sweep from BEST baselines.

User 2026-09-26: Backtests 55h stale (Sept 19 NPZ predicted +5.86% XLM but bear). 
Fix: rerun QUICK tests from settings in SPREADSHEETS/BEST/ (cumulative overrides as baseline)
with DC-based stop/target instead of fixed %:
  LONG stop = dc_low_TF * (1 - 0.25%), SHORT stop = dc_high_TF * (1+0.25%)
  LONG target = dc_high_TF * (1 - 0.10%), SHORT target = dc_low_TF * (1+0.10%)
TF options: 3m (covers 3/5m), 15m, 1h — each vs cumulative_before, greedy not needed, just delta.

This is NOT the huge 12-sheet TEMPLATE sweep (3043 rows). This is 6 DC variants + baseline,
0.07s each, <5s per sym_side. Starting point = BEST global cumulative overrides (GLOBAL_RISK_GATES final C).

After quick DC proof, the full previously-tested scope can be rerun via v15_pilot with fresh NPZ.

Usage (on S1 with fresh NPZ):
  python3 tools/v15_quick_dc_sweep.py --sym XLMUSDT_LONG --window 30
  python3 tools/v15_quick_dc_sweep.py --all --venue crypto --limit 20
  python3 tools/v15_quick_dc_sweep.py --all --fresh-npz   # regen stale NPZ first
"""
from __future__ import annotations
import argparse, json, re, sys, time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import numpy as np
import openpyxl

BEST_DIRS = [ROOT / "SPREADSHEETS" / "BEST" / d for d in ("CRYPTO_LONG","CRYPTO_SHORT","STOCKS_LONG","STOCKS_SHORT")]
V15_DIR = ROOT / "SPREADSHEETS" / "V15_V16_CELL_BY_CELL"

# DC sweep candidates per user spec
DC_STOP_TFS = ["3m", "15m", "1h"]   # 3m covers 3/5m
DC_TARGET_TFS = ["3m", "15m", "1h"]
STOP_BUF = 0.25   # %
TARGET_BUF = 0.10 # %


def parse_overrides_string(s: str) -> dict:
    """'K=V + K2=V2 + ...' -> dict. Handles bool/float."""
    out = {}
    if not s or not isinstance(s, str):
        return out
    # split by ' + ' but also handle ' +'
    parts = [p.strip() for p in s.split(" + ")]
    if len(parts) == 1 and " + " not in s and "+" in s:
        parts = [p.strip() for p in s.split("+")]
    for p in parts:
        if "=" not in p:
            continue
        k, v = p.split("=", 1)
        k = k.strip(); v = v.strip()
        if v == "True": v = True
        elif v == "False": v = False
        elif v == "OFF": v = "OFF"
        else:
            try:
                # keep string TFs as string
                if re.match(r"^-?\d+\.?\d*$", v):
                    v = float(v) if "." in v else int(v)
            except: pass
        out[k] = v
    return out


def extract_best_overrides(sym_side: str) -> dict:
    """Find BEST xlsx for sym_side and extract cumulative overrides from GLOBAL_RISK_GATES final C."""
    # locate file
    for d in BEST_DIRS:
        if not d.exists():
            continue
        for p in d.glob(f"{sym_side}_*.xlsx"):
            try:
                wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
                if "GLOBAL_RISK_GATES" not in wb.sheetnames:
                    wb.close(); continue
                ws = wb["GLOBAL_RISK_GATES"]
                # find last row with C override
                best = ""
                for r in range(ws.max_row, 2, -1):
                    v = ws.cell(r, 3).value
                    if v and isinstance(v, str) and "=" in v:
                        best = v
                        break
                wb.close()
                if best:
                    return parse_overrides_string(best)
            except Exception:
                continue
        # also try V15_V16 as fallback (freshest)
        for p in (V15_DIR.glob(f"{sym_side}_*.xlsx")):
            try:
                wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
                if "GLOBAL_RISK_GATES" not in wb.sheetnames:
                    wb.close(); continue
                ws = wb["GLOBAL_RISK_GATES"]
                best = ""
                for r in range(ws.max_row, 2, -1):
                    v = ws.cell(r, 3).value
                    if v and isinstance(v, str) and "=" in v:
                        best = v; break
                wb.close()
                if best:
                    return parse_overrides_string(best)
            except Exception:
                continue
    return {}


def eval_gain(npz, sym, is_long, overrides, window_days=30):
    """Evaluate gain via v12_quick_engine on sliced 30d window with overrides as QuickConfig."""
    import v12_quick_engine as V
    from tools.opt.evaluate_v12 import _exact_30d_slice
    # slice to 30d window anchored to NPZ last bar
    crypto = sym.upper().endswith(("USDT","USDC","USD1"))
    try:
        sliced, _ = _exact_30d_slice(npz, crypto, window_days)
    except Exception as e:
        return None, f"slice fail {e}"
    cfg = V.QuickConfig()
    if crypto:
        pass
    else:
        cfg.apply_tradier_defaults()
    # apply overrides
    for k, v in overrides.items():
        if hasattr(cfg, k):
            cur = getattr(cfg, k)
            if isinstance(cur, bool):
                setattr(cfg, k, bool(v))
            elif isinstance(cur, (int,float)) and not isinstance(cur, bool):
                try:
                    setattr(cfg, k, type(cur)(v))
                except: setattr(cfg, k, v)
            else:
                setattr(cfg, k, v)
        else:
            # allow new DC fields even if not yet in class (set attr anyway)
            setattr(cfg, k, v)
    # ensure venue mode
    cfg.MODE = "crypto" if crypto else "tradier"
    r = V.simulate_one(sliced, sym, is_long, cfg)
    if r is None:
        return None, "simulate_one None"
    return r, None


def run_sym(sym_side: str, window_days: int = 30, verbose: bool = True):
    sym, side = (sym_side[:-5], "LONG") if sym_side.endswith("_LONG") else (sym_side[:-6], "SHORT")
    is_long = side == "LONG"
    overrides = extract_best_overrides(sym_side)
    if verbose:
        print(f"[{sym_side}] best overrides: {len(overrides)} keys")
        if len(overrides) < 5:
            print(f"  raw sample: {list(overrides.items())[:5]}")
    # load NPZ
    import v12_quick_engine as V
    mode = "crypto" if sym.upper().endswith(("USDT","USDC","USD1")) else "tradier"
    try:
        stores = V.load_npz(mode, [sym], "2024-01-01")
        npz = stores.get(sym) if isinstance(stores, dict) else stores
        if npz is None:
            # try tokenised strip
            base = sym[:-4] if sym.endswith("USDT") else sym
            stores = V.load_npz(mode, [base], "2024-01-01")
            npz = stores.get(base) if isinstance(stores, dict) else stores
        if npz is None or len(npz.get("timestamps", [])) < 10:
            return {"sym_side": sym_side, "error": "npz missing"}
    except Exception as e:
        return {"sym_side": sym_side, "error": f"load_npz {e}"}
    # baseline
    base_r, err = eval_gain(npz, sym, is_long, overrides, window_days)
    if base_r is None:
        return {"sym_side": sym_side, "error": err, "overrides": len(overrides)}
    base_gain = base_r.get("gain_pct_2000norm", 0.0)
    base_bh = None
    try:
        from tools.opt.evaluate_v12 import _bh
        # bh from sliced npz window
        from tools.opt.evaluate_v12 import _exact_30d_slice as _slice
        crypto = sym.upper().endswith(("USDT","USDC","USD1"))
        sliced, _ = _slice(npz, crypto, window_days)
        c = np.asarray(sliced.get("close", []), dtype=float)
        c = c[np.isfinite(c) & (c > 0)]
        if len(c) >= 2:
            raw = (c[-1]-c[0])/c[0]*100
            base_bh = raw if is_long else max(-100, -raw)
    except Exception:
        pass
    variants = []
    # sweep DC stop/target TFs — each as single-switch delta vs cumulative_before
    for tf in DC_STOP_TFS:
        ov = dict(overrides)
        # daytrade stop TF + technical stop TF both set to same TF (user spec: all three use same logic)
        ov["DAYTRADE_DC_STOP_TF"] = tf
        ov["DAYTRADE_DC_STOP_BUFFER_PCT"] = STOP_BUF
        ov["TECHNICAL_DC_STOP_TF"] = tf
        ov["TECHNICAL_DC_STOP_BUFFER_PCT"] = STOP_BUF
        r, _ = eval_gain(npz, sym, is_long, ov, window_days)
        if r is not None:
            g = r.get("gain_pct_2000norm", 0.0)
            variants.append({"variant": f"STOP_{tf}_buf{STOP_BUF}", "gain": g, "delta": g - base_gain, "trades": r.get("trades",0), "tim": r.get("tim_pct",0)})
    for tf in DC_TARGET_TFS:
        ov = dict(overrides)
        ov["DAYTRADE_DC_TARGET_TF"] = tf
        ov["DAYTRADE_DC_TARGET_BUFFER_PCT"] = TARGET_BUF
        ov["TECHNICAL_DC_TARGET_TF"] = tf
        ov["TECHNICAL_DC_TARGET_BUFFER_PCT"] = TARGET_BUF
        r, _ = eval_gain(npz, sym, is_long, ov, window_days)
        if r is not None:
            g = r.get("gain_pct_2000norm", 0.0)
            variants.append({"variant": f"TARGET_{tf}_buf{TARGET_BUF}", "gain": g, "delta": g - base_gain, "trades": r.get("trades",0), "tim": r.get("tim_pct",0)})
    # combined: best stop + best target (both)
    if variants:
        # pick best stop and best target individually
        best_stop = max([v for v in variants if v["variant"].startswith("STOP")], key=lambda x: x["delta"], default=None)
        best_tgt = max([v for v in variants if v["variant"].startswith("TARGET")], key=lambda x: x["delta"], default=None)
        if best_stop and best_tgt and best_stop["delta"] > 0 and best_tgt["delta"] > 0:
            tf_s = best_stop["variant"].split("_")[1]
            tf_t = best_tgt["variant"].split("_")[1]
            ov = dict(overrides)
            ov["DAYTRADE_DC_STOP_TF"] = tf_s; ov["TECHNICAL_DC_STOP_TF"] = tf_s
            ov["DAYTRADE_DC_TARGET_TF"] = tf_t; ov["TECHNICAL_DC_TARGET_TF"] = tf_t
            ov["DAYTRADE_DC_STOP_BUFFER_PCT"] = STOP_BUF; ov["TECHNICAL_DC_STOP_BUFFER_PCT"] = STOP_BUF
            ov["DAYTRADE_DC_TARGET_BUFFER_PCT"] = TARGET_BUF; ov["TECHNICAL_DC_TARGET_BUFFER_PCT"] = TARGET_BUF
            r, _ = eval_gain(npz, sym, is_long, ov, window_days)
            if r is not None:
                g = r.get("gain_pct_2000norm", 0.0)
                variants.append({"variant": f"COMBO_STOP_{tf_s}+TARGET_{tf_t}", "gain": g, "delta": g - base_gain, "trades": r.get("trades",0), "tim": r.get("tim_pct",0)})
    # sort by delta
    variants.sort(key=lambda x: x["delta"], reverse=True)
    best = variants[0] if variants else None
    result = {
        "sym_side": sym_side,
        "window_days": window_days,
        "base_gain": round(base_gain, 4),
        "base_bh": round(base_bh,4) if base_bh is not None else None,
        "base_trades": base_r.get("trades",0),
        "base_tim": round(base_r.get("tim_pct",0),2),
        "overrides_keys": len(overrides),
        "variants": variants,
        "best": best,
        "npz_bars": len(npz.get("timestamps", [])),
        "npz_last_ts": float(np.asarray(npz.get("timestamps", [0]))[-1]) if len(np.asarray(npz.get("timestamps", [0]))) else 0,
    }
    if verbose:
        print(f"  baseline gain {base_gain:.2f}% bh {(base_bh if base_bh is not None else 0):.2f} trades {base_r.get('trades',0)}")
        for v in variants[:5]:
            print(f"    {v['variant']:25s} gain {v['gain']:+7.2f} delta {v['delta']:+6.2f} trades {v['trades']:3d} tim {v['tim']:.1f}%")
        if best and best["delta"] > 1e-9:
            print(f"  => BEST {best['variant']} delta {best['delta']:+.2f}")
        else:
            print(f"  => no positive DC variant")
    return result


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--sym", default="", help="sym_side e.g. XLMUSDT_LONG")
    p.add_argument("--all", action="store_true", help="run all BEST symbols")
    p.add_argument("--venue", choices=["crypto","stocks","both"], default="both")
    p.add_argument("--window", type=int, default=30, help="30d window")
    p.add_argument("--limit", type=int, default=0, help="limit N symbols when --all")
    p.add_argument("--out", default="data/reports/v15_quick_dc_sweep.json")
    p.add_argument("--fresh-npz", action="store_true", help="regen NPZ on S1 before sweep (calls backtest_v8_precompute)")
    args = p.parse_args()
    if args.fresh_npz:
        import subprocess as sp
        print("[fresh-npz] regenerating indicators on S1 (this may take minutes)...")
        # local regen if on S1 (has sandbox), else warn
        if Path("/home/niels/binance-sandbox/backtest_v8/indicators").exists():
            sp.run([sys.executable, "backtest_v8_precompute.py", "--mode", "crypto"], check=False)
            sp.run([sys.executable, "backtest_v8_precompute.py", "--mode", "tradier"], check=False)
        else:
            print("[fresh-npz] not on S1, skip regen — ensure S1 NPZ are fresh via ssh s1-int")
    results = []
    if args.sym:
        results.append(run_sym(args.sym, args.window))
    elif args.all:
        syms = []
        for d in BEST_DIRS:
            if not d.exists(): continue
            if args.venue != "both":
                if args.venue == "crypto" and "CRYPTO" not in str(d): continue
                if args.venue == "stocks" and "STOCKS" not in str(d): continue
            for pth in sorted(d.glob("*.xlsx")):
                name = pth.stem
                # stem is SYM_SIDE_bh.._gain.._30d_matrix -> extract SYM_SIDE
                # SYM may contain underscores? No, SYM_SIDE is SYM_LONG/SHORT
                if "_LONG_" in name:
                    ss = name.split("_LONG_")[0] + "_LONG"
                elif "_SHORT_" in name:
                    ss = name.split("_SHORT_")[0] + "_SHORT"
                else:
                    continue
                if ss not in syms:
                    syms.append(ss)
        if args.limit:
            syms = syms[:args.limit]
        print(f"[all] {len(syms)} symbols from BEST")
        for ss in syms:
            r = run_sym(ss, args.window)
            results.append(r)
            time.sleep(0.05)
    else:
        p.print_help(); return
    # write out
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(results, indent=2))
    print(f"[done] wrote {out} ({len(results)} syms)")
    # summary
    pos = sum(1 for r in results if r.get("best") and r["best"]["delta"] > 0)
    print(f"  positive DC variants: {pos}/{len(results)}")
    if results:
        avg_delta = sum(r.get("best",{}).get("delta",0) for r in results if r.get("best"))/max(1,len([r for r in results if r.get("best")]))
        print(f"  avg best delta: {avg_delta:+.2f}pp")
    # stress note per backtest-expert: this is quick 30d validation, not 1y robustness
    print("[note] backtest-expert: quick 30d is indication only — full scope needs 1y + walk-forward before promotion")

if __name__ == "__main__":
    main()
