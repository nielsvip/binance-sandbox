#!/usr/bin/env python3
"""The TEMPLATE dictates. The cat_side layer adopts it verbatim — value AND type. Nothing more, nothing less.

USER 2026-10-21, after repeating it many times: "WT_DIV_EXIT_ENABLED template False / config True needs to be
IMPOSSIBLE as is 30 vs 30.0 or ANYTHING ELSE that does not BLINDLY use the values from the
TEMPLATE_CRYPTO/STOCK_LONG_SHORT (ONLY ONE EVER APPLIES!!) into its config. NO overrides NO exceptions NO
time int float whatever confusions JUST FKN USE WHAT IT DICTATES".

WHY THIS LAYER AND NOT config.py. One `config.py` serves BOTH crypto lanes and one `config_tradier.py` serves both
stock lanes, so a single global value CANNOT equal two templates whose bold defaults differ — which is why
comparing templates against the venue configs produced ~180 "mismatches" that are not fixable there. Exactly one
template applies to any sym_side, and the layer with one slot per cat_side is `data/cat_side_defaults_4.json`
(CRYPTO_LONG / CRYPTO_SHORT / STOCKS_LONG / STOCKS_SHORT), which the resolver consults as
per-sym > cat_side > global. So that file is where "only one ever applies" is expressible, and this makes it an
exact projection of the four workbooks.

MEASURED BEFORE THE FIRST APPLY (2026-10-21), per lane, bold `is_default` rows vs the stored lane:
  CRYPTO_LONG   1092 defaults: 68 MISSING, 58 VALUE_DIFF, 174 TYPE_ONLY
  CRYPTO_SHORT  1098 defaults: 60 MISSING, 42 VALUE_DIFF, 186 TYPE_ONLY
  STOCKS_LONG   1191 defaults: 70 MISSING, 56 VALUE_DIFF, 195 TYPE_ONLY
  STOCKS_SHORT  1171 defaults: 64 MISSING, 48 VALUE_DIFF, 199 TYPE_ONLY
  = 262 missing + 204 wrong value + 754 wrong type = 1220 deviations.
A MISSING key is the worst of the three: it falls through to the venue global, which is precisely the
`WT_DIV_EXIT_ENABLED template False / config True` case. TYPE_ONLY (`30` vs `30.0`) is included because the
operator ruled it out explicitly, and it is not cosmetic everywhere — a TF field holding `80.0` renders as the
string "80.0" and silently matches no timeframe.

TYPE RULE: the cell's own type is the truth. A bool cell becomes a real bool (including the strings "True"/"False"
that some rows store), a numeric cell keeps int as int and float as float, everything else stays the string it is,
stripped. No coercion to a "nicer" type, because that is the confusion being banned.

SCOPE: only keys the templates actually dictate are touched. Keys present in a lane but in no template are left
alone and reported — they are not template-dictated, and deleting infrastructure entries is not what was asked.

Exit 0 = every lane is an exact projection, 1 = deviations remain, 2 = could not audit. Read-only without --apply.
"""
from __future__ import annotations

import argparse
import collections
import glob
import json
import os
import pathlib
import shutil
import sys
import time

ROOT = pathlib.Path(os.environ.get("BASE_PATH") or pathlib.Path(__file__).resolve().parents[1])
LAYER = ROOT / "data" / "cat_side_defaults_4.json"
LANES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
SKIP_TABS = {"PROTOCOL", "LEGEND_FILTERS", "INSTRUCTIONS", "INSTRUCTIONS_V2", "FILTERS_EXPLAINED",
             "FILTER_DICTIONARY_V2", "Results_Deltas", "Results_30d_Deltas"}


def native(v):
    """The cell's own type, verbatim. Bools stay bools, ints stay ints, floats stay floats, text stays text."""
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    low = s.lower()
    if low == "true":
        return True
    if low == "false":
        return False
    return s


def template_defaults() -> dict:
    """{lane: {SWITCH: default}} from the bold is_default row (column L == 'YES') of every switch tab."""
    import openpyxl
    out = {}
    for p in sorted(glob.glob(str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_*.xlsx"))):
        lane = pathlib.Path(p).stem.replace("TEMPLATE_", "")
        if lane not in LANES:
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        d = {}
        for ws in wb.worksheets:
            if ws.title in SKIP_TABS or "BASELINE" in ws.title:
                continue
            for r in ws.iter_rows(min_row=3, max_row=min(ws.max_row or 1, 1200), max_col=12, values_only=True):
                if r and r[0] and r[11] and str(r[11]).strip().upper() == "YES":
                    d[str(r[0]).strip().upper()] = native(r[1])
        wb.close()
        out[lane] = d
    return out


def deviations(tmpl: dict, layer: dict) -> dict:
    """Per lane: missing / value_diff / type_diff against the template."""
    rep = {}
    for lane in LANES:
        T, L = tmpl.get(lane, {}), layer.get(lane, {}) or {}
        missing, valdiff, typediff = [], [], []
        for k, want in T.items():
            if k not in L:
                missing.append(k)
                continue
            have = L[k]
            if have == want and type(have) is type(want):
                continue
            same_number = False
            if not isinstance(want, bool) and not isinstance(have, bool):
                try:
                    same_number = abs(float(str(want)) - float(str(have))) < 1e-9
                except Exception:
                    same_number = False
            if str(have).strip().lower() == str(want).strip().lower() or same_number:
                typediff.append((k, want, have))
            else:
                valdiff.append((k, want, have))
        extras = [k for k in L if k not in T]
        rep[lane] = {"n_template": len(T), "missing": missing, "value_diff": valdiff,
                     "type_diff": typediff, "n_extras_not_in_template": len(extras)}
    return rep


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="rewrite the cat_side layer to the template verbatim")
    ap.add_argument("--json", default=str(ROOT / "data" / "reports" / "cat_side_from_template.json"))
    a = ap.parse_args()

    if not LAYER.exists():
        print(f"missing {LAYER}")
        return 2
    tmpl = template_defaults()
    if len(tmpl) != 4:
        print(f"expected 4 templates, found {sorted(tmpl)}")
        return 2
    layer = json.loads(LAYER.read_text())
    rep = deviations(tmpl, layer)

    tot = collections.Counter()
    for lane in LANES:
        r = rep[lane]
        tot["missing"] += len(r["missing"])
        tot["value_diff"] += len(r["value_diff"])
        tot["type_diff"] += len(r["type_diff"])
        print(f"  {lane:14s} template_defaults={r['n_template']:4d}  MISSING={len(r['missing']):3d}  "
              f"VALUE_DIFF={len(r['value_diff']):3d}  TYPE_DIFF={len(r['type_diff']):3d}  "
              f"(lane also holds {r['n_extras_not_in_template']} keys no template dictates)")
    total = sum(tot.values())
    print(f"\n  TOTAL DEVIATIONS FROM THE TEMPLATE: {total}  {dict(tot)}")
    for lane in LANES[:1]:
        for k, want, have in rep[lane]["value_diff"][:6]:
            print(f"     {lane} {k:32s} template={want!r:12} lane={have!r}")

    if a.apply and total:
        n_before = {lane: len(layer.get(lane, {}) or {}) for lane in LANES}
        bk = ROOT / "backups" / f"cat_side_defaults_4_before_template_projection_{time.strftime('%Y%m%d%H%M', time.gmtime())}.json"
        bk.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(LAYER, bk)
        written = 0
        for lane in LANES:
            L = layer.setdefault(lane, {})
            for k, want in tmpl[lane].items():
                if k not in L or L[k] != want or type(L[k]) is not type(want):
                    L[k] = want
                    written += 1
        # WRITE GUARD. On 2026-10-20 a fixer of mine serialised a loop variable over a 96 MB live file and left
        # 3 bytes behind. Never write unless the object is still the whole mapping with no lane shrunk.
        if not isinstance(layer, dict) or any(lane not in layer for lane in LANES):
            print("REFUSING TO WRITE: layer lost a lane")
            return 2
        for lane in LANES:
            if len(layer[lane]) < n_before[lane]:
                print(f"REFUSING TO WRITE: {lane} shrank {n_before[lane]} -> {len(layer[lane])}")
                return 2
        payload = json.dumps(layer, indent=1, default=str, sort_keys=True)
        if len(payload) < 200_000:
            print(f"REFUSING TO WRITE: payload only {len(payload)} bytes")
            return 2
        layer.setdefault("_meta", {})["template_projection"] = {
            "at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "by": "tools/cat_side_from_template.py",
            "keys_written": written,
            "rule": "TEMPLATE_FINAL_NORM bold is_default row is the sole authority; value and type verbatim",
        }
        payload = json.dumps(layer, indent=1, default=str, sort_keys=True)
        tmp = LAYER.with_suffix(".json.tmp")
        tmp.write_text(payload)
        os.replace(tmp, LAYER)
        print(f"\n  APPLIED: {written} keys set to their template value+type. backup {bk}")
        layer = json.loads(LAYER.read_text())
        rep = deviations(tmpl, layer)
        total = sum(len(rep[l]["missing"]) + len(rep[l]["value_diff"]) + len(rep[l]["type_diff"]) for l in LANES)
        print(f"  RE-CHECK AFTER APPLY: {total} deviations remain")

    pathlib.Path(a.json).parent.mkdir(parents=True, exist_ok=True)
    pathlib.Path(a.json).write_text(json.dumps(
        {"ts": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "total_deviations": total,
         "per_lane": {l: {"n_template": rep[l]["n_template"], "missing": rep[l]["missing"][:200],
                          "value_diff": rep[l]["value_diff"][:200], "type_diff": rep[l]["type_diff"][:200],
                          "n_extras_not_in_template": rep[l]["n_extras_not_in_template"]} for l in LANES}},
        indent=1, default=str))
    print(f"  wrote {a.json}")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
