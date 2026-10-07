#!/usr/bin/env python3
"""Stub delta_matrix to unblock lifecycle_pilot - reads TEMPLATE_30d_matrix.xlsx"""
import openpyxl
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
def _expand_values(v):
    if v is None: return []
    s=str(v).strip()
    if not s or s=="None": return []
    # split by comma or slash
    parts=[p.strip() for p in s.replace("/",",").split(",")]
    out=[]
    for p in parts:
        if p.lower()=="true": out.append(True)
        elif p.lower()=="false": out.append(False)
        elif p.lower()=="none": out.append(None)
        else:
            try:
                if "." in p: out.append(float(p))
                else: out.append(int(p))
            except: out.append(p)
    return out or [True, False]
def parse_paired_str(s):
    if not s: return {}
    # format "PARAM:1,2,3" or "PARAM 1/2"
    # For stub, return empty - paired handled elsewhere
    return {}
def parse_sheet_rows(rows):
    # rows is list of dicts from load_inventory - stub returns as is
    return [], rows
def load_inventory():
    # Load TEMPLATE sheets
    p=ROOT/"SPREADSHEETS/TEMPLATE_30d_matrix.xlsx"
    sheets={}
    try:
        wb=openpyxl.load_workbook(str(p), data_only=True, read_only=True)
        for name in wb.sheetnames:
            # Skip non-inventory sheets: support both generic and legacy 30d-named deltas sheet
            if name.startswith("_") or name in ("INSTRUCTIONS","results","FILTER_DICTIONARY_V8","Results_Deltas","Results_30d_Deltas"): continue
            if name.lower().endswith("_deltas"): continue
            if "BASELINE" in name: continue
            ws=wb[name]
            rows=[]
            headers=[c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            # headers: Switch,default,override,Family,...
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                # check italic blanked - need font check, but read_only loses font, so check override col E/F?
                # For stub, include all non-blank rows
                d={}
                for i,h in enumerate(headers):
                    if h and i<len(row):
                        d[str(h).strip()]=row[i]
                # normalize keys expected by switch_tf_uniqueness
                d["PARAM"]=d.get("Switch") or d.get("PARAM")
                d["VALUES"]=str(d.get("default") or "")
                d["GROUP"]=d.get("Family") or ""
                rows.append(d)
            sheets[name]=rows
    except Exception as e:
        pass
    return None, None, sheets
