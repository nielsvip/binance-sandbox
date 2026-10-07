#!/usr/bin/env python3
"""build_switch_bible — deterministic scanner that writes the SWITCH BIBLE (data/SWITCH_BIBLE.json + SWITCH_BIBLE.md, repo root).

For EVERY switch/filter name that appears in a TEMPLATE_*.xlsx (white switch rows, orange filter rows, yellow FILTER=opt headers) or that is
READ by the vectorized code (v12_quick_engine.py QuickConfig fields, vec_decisions/*, tools/opt/evaluate_v12.py) it records, from real AST scans only:
  defaults        config.py Config / config_tradier.py TradierConfig / QuickConfig (+ apply_tradier_defaults) / cat_side_defaults_4 (4 cat_sides) / template bold per cat_side
  live_reads      [file:line(function)] in ez_manage.py / tradier_manage.py and the root modules they import (import closure), split ez(crypto) / tradier(stocks)
  vec_reads       [file:line(function)] in v12_quick_engine.py / vec_decisions/* / evaluate_v12, flagged reachable / dead (early-return function) / stub (`_ = getattr` no-op)
  template        {template: [{tab, rows, default, grey, orange}]} + yellow-header presence
  status          WIRED_BOTH_PARITY_PROVEN | WIRED_BOTH_UNPROVEN | VEC_ONLY | LIVE_ONLY | DEAD | NOT_IN_CONFIG   (per venue: crypto / stocks, plus overall)
  provenance      md5 of every scanned file
A "read" = attribute load (cfg.NAME), bare uppercase Name, or an uppercase string literal that is a direct argument of a call (getattr/_cfg/_psym_get/.get...).
Strings elsewhere (dict keys, lists) are weak "refs" and never make a switch WIRED. Config-class body assignments are definitions, not reads.
PARITY_PROVEN only if the name is listed in data/wiring/parity_proven.json (scalar-vs-vec parity); data/wiring/wired_proven.json (ledger flip) is recorded as `ledger_flip_proven`.
Run:  python tools/build_switch_bible.py [--no-md] [--out data/SWITCH_BIBLE.json]
"""
import argparse
import ast
import collections
import datetime
import hashlib
import json
import re
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
NAME_RE = re.compile(r"^[A-Z][A-Z0-9_]{3,}$")
TEMPLATES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
CONFIG_CLASSES = {"config.py": "Config", "config_tradier.py": "TradierConfig", "v12_quick_engine.py": "QuickConfig"}
LIVE_ROOTS = {"crypto": "ez_manage.py", "stocks": "tradier_manage.py"}
VEC_FILES_FIXED = ["v12_quick_engine.py", "tools/opt/evaluate_v12.py"]
CALL_READ_FUNCS = {"getattr", "hasattr", "_cfg", "_psym_get", "_ezm_default", "get", "_get", "getenv", "_cfg_get", "cfg_get", "_tcfg", "_tradier_cfg", "_live_cfg", "_b", "_f", "_i", "_s", "_cfg_auto", "_cget", "_c", "_g", "_gv", "_val", "_flag", "_opt", "_knob", "_cfgv", "_cfg_bool", "_cfg_float", "_cfg_str", "_cfg_int", "_et_cfg", "_gx_c", "_mu_enabled_for", "_cf", "_la_get", "_sw_get", "_thr_cfg", "_tf", "_ftf_tf"}
# 2026-10-04: `_et_cfg` (wt_dc_delta.py, unique name) is a genuine cfg.get/getattr getter; `_gx_c` (tradier grey-rewire, unique name) is a _cfg lambda alias.
# 2026-10-04: `_research_only` REMOVED from the farm regex — the single such function (tradier _apply_research_only_live_gates, called from 3 live
# sites) holds REAL gates; its `_ = _cfg(...)` admitted-coverage stubs stay excluded via _is_stub and its `and False` guards via _const_false.
FARM_FUNC_RE = re.compile(r"^_batch\d*_template|^_ensure_ez_all$|^_ensure_tradier|_wiring$")  # audit-defeating registries/stubs (BIBLE §19): reads there are NOT consumers
EXCLUDED_LIVE_IMPORTS = ("v12", "v8_", "backtest", "evaluate", "per_sym_", "seq_embed", "vec_", "config", "cat_side_defaults")


def md5(p):
    return hashlib.md5(Path(p).read_bytes()).hexdigest()


def rel(p):
    return str(Path(p).resolve().relative_to(ROOT))


# --------------------------------------------------------------------------------------------- AST scanning
def _early_return(fn):
    body = [s for s in fn.body if not (isinstance(s, ast.Expr) and isinstance(getattr(s, "value", None), ast.Constant) and isinstance(s.value.value, str))]
    return any(isinstance(s, ast.Return) and i < len(body) - 1 for i, s in enumerate(body))


def _is_stub(st):
    if isinstance(st, ast.Assign) and all(isinstance(t, ast.Name) and t.id == "_" for t in st.targets):
        return True
    return isinstance(st, ast.If) and st.body and all(isinstance(b, ast.Assign) and all(isinstance(t, ast.Name) and t.id == "_" for t in b.targets) for b in st.body) and not st.orelse


def _const_false(test):
    """`X and False`, `False and X`, `if False` — a guard that can never be true (the live farms use these)."""
    if isinstance(test, ast.Constant) and test.value is False:
        return True
    return isinstance(test, ast.BoolOp) and isinstance(test.op, ast.And) and any(isinstance(v, ast.Constant) and v.value is False for v in test.values)


def _call_name(c):
    f = c.func
    if isinstance(f, ast.Name):
        return f.id
    if isinstance(f, ast.Attribute):
        return f.attr
    return ""


class FileScan:
    """events: name -> list of dict(kind=read|write|ref|def, line, func, flag=live|dead|stub)"""

    def __init__(self, path, config_class=None):
        self.path = Path(path)
        self.rel = rel(path)
        self.events = collections.defaultdict(list)
        self.defs = {}          # config class body: name -> literal/str default
        self.stores = collections.defaultdict(dict)  # method -> {name: literal} (self.X = lit inside apply_*_defaults)
        self.funcs = {}         # qualname -> (node, calls:set)
        self.dyn = []           # (line, func, regex) read calls whose key is an f-string, e.g. get(f"{prefix}_{action}_ENABLED")
        try:
            self.tree = ast.parse(self.path.read_text(errors="replace"))
        except SyntaxError:
            self.tree = None
            return
        self.config_class = config_class
        self._walk(self.tree, [], False, False)

    def _qual(self, stack):
        return ".".join(stack) if stack else "<module>"

    def _lit(self, node):
        try:
            return ast.literal_eval(node)
        except Exception:
            try:
                return "<expr> " + ast.unparse(node)[:80]
            except Exception:
                return "<expr>"

    def _add(self, name, kind, node, stack, dead, stub):
        flag = "dead" if dead else ("stub" if stub else ("farm" if any(FARM_FUNC_RE.search(x) for x in stack) else "live"))
        self.events[name].append({"kind": kind, "line": getattr(node, "lineno", 0), "func": self._qual(stack), "flag": flag})

    def _walk(self, node, stack, dead, stub, in_config_body=False, parent=None, method_store=None):
        # Sibling-threaded deadness (2026-10-05): only statements AFTER an unconditional
        # func-top-level return/raise are dead (batch1 pattern). Guard/conditional returns
        # no longer nuke the whole function (score_exit false-positive fix).
        _d = dead
        _is_func_top = isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
        for ch in ast.iter_child_nodes(node):
            if isinstance(ch, ast.ClassDef):
                is_cfg = ch.name == self.config_class
                self._walk(ch, stack + [ch.name], _d, stub, is_cfg, ch)
                continue
            if isinstance(ch, (ast.FunctionDef, ast.AsyncFunctionDef)):
                d2 = _d or ("batch" in ch.name.lower() and "wiring" in ch.name.lower())
                q = ".".join(stack + [ch.name])
                calls = {_call_name(c) for c in ast.walk(ch) if isinstance(c, ast.Call)}
                self.funcs[q] = (ch, calls)
                self._walk(ch, stack + [ch.name], d2, False, False, ch, ch.name if in_config_body else method_store)
                continue
            if in_config_body and isinstance(ch, (ast.Assign, ast.AnnAssign)):
                targets = ch.targets if isinstance(ch, ast.Assign) else [ch.target]
                for t in targets:
                    if isinstance(t, ast.Name) and NAME_RE.match(t.id) and ch.value is not None:
                        self.defs[t.id] = self._lit(ch.value)
                        self.events[t.id].append({"kind": "def", "line": ch.lineno, "func": self._qual(stack), "flag": "live"})
                # value of a default may read other names; fall through to scan value
                if ch.value is not None:
                    self._walk_expr(ch.value, stack, _d, stub)
                continue
            if isinstance(ch, ast.stmt) and _is_stub(ch):
                self._walk(ch, stack, _d, True)
                continue
            if isinstance(ch, ast.If) and _const_false(ch.test):
                self._walk_expr(ch.test, stack, _d, True)
                self._walk(ch, stack, _d, True)
                continue
            self._visit(ch, stack, _d, stub)
            self._walk(ch, stack, _d, stub, in_config_body, ch, method_store)
            if _is_func_top and isinstance(ch, (ast.Return, ast.Raise)):
                _d = True

    def _walk_expr(self, node, stack, dead, stub):
        for n in ast.walk(node):
            self._visit(n, stack, dead, stub)

    def _visit(self, n, stack, dead, stub):
        if isinstance(n, ast.Attribute) and NAME_RE.match(n.attr):
            if isinstance(n.ctx, ast.Store):
                self._add(n.attr, "write", n, stack, dead, stub)
                # self.X = <literal> inside a config-class method (apply_tradier_defaults)
                self.stores[self._qual(stack)][n.attr] = n.lineno
            else:
                self._add(n.attr, "read", n, stack, dead, stub)
        elif isinstance(n, ast.Name) and NAME_RE.match(n.id) and isinstance(n.ctx, ast.Load):
            self._add(n.id, "read", n, stack, dead, stub)
        elif isinstance(n, ast.Call):
            cn = _call_name(n)
            if cn in CALL_READ_FUNCS and n.args and isinstance(n.args[0], ast.JoinedStr):
                parts = []
                for v in n.args[0].values:
                    parts.append(re.escape(v.value) if isinstance(v, ast.Constant) else "[A-Z0-9_]+")
                if any(isinstance(v, ast.Constant) and len(v.value) >= 4 for v in n.args[0].values):
                    self.dyn.append((n.lineno, self._qual(stack), "^" + "".join(parts) + "$"))
            for a in n.args[:3]:
                if isinstance(a, ast.Constant) and isinstance(a.value, str) and NAME_RE.match(a.value):
                    self._add(a.value, "read" if cn in CALL_READ_FUNCS else "ref", a, stack, dead, stub)
            if cn in ("setattr",) and len(n.args) >= 2 and isinstance(n.args[1], ast.Constant) and isinstance(n.args[1].value, str) and NAME_RE.match(n.args[1].value):
                self._add(n.args[1].value, "write", n.args[1], stack, dead, stub)
        elif isinstance(n, ast.Constant) and isinstance(n.value, str) and NAME_RE.match(n.value):
            self._add(n.value, "ref", n, stack, dead, stub)
        elif isinstance(n, (ast.Dict,)):
            pass


def import_closure(root_file, universe_dir=ROOT, block=()):
    """root-level repo modules statically imported (transitively) by root_file."""
    seen, todo = set(), [Path(root_file).name]
    mods = {p.stem for p in universe_dir.glob("*.py")}
    while todo:
        f = todo.pop()
        if f in seen:
            continue
        seen.add(f)
        p = universe_dir / f
        if not p.exists():
            continue
        try:
            tree = ast.parse(p.read_text(errors="replace"))
        except SyntaxError:
            continue
        for n in ast.walk(tree):
            names = []
            if isinstance(n, ast.Import):
                names = [a.name.split(".")[0] for a in n.names]
            elif isinstance(n, ast.ImportFrom) and n.module and n.level == 0:
                names = [n.module.split(".")[0]]
            for m in names:
                if m in mods and m not in block and not any(m.startswith(x) for x in EXCLUDED_LIVE_IMPORTS) and m + ".py" not in seen:
                    todo.append(m + ".py")
    return sorted(seen)


def shared_vec_modules(files, universe_dir=ROOT):
    """vec_decisions/<m>.py that the given LIVE files import (static, incl. in-function imports) -> a read there is BOTH a live and a vec site
    (N2 finding 2026-10-01: dc_channel_exits.py is the single implementation used by tradier_manage/ez_manage AND the vector engine)."""
    out = set()
    for f in files:
        p = universe_dir / f
        if not p.exists():
            continue
        try:
            tree = ast.parse(p.read_text(errors="replace"))
        except SyntaxError:
            continue
        for n in ast.walk(tree):
            if isinstance(n, ast.ImportFrom) and n.level == 0 and n.module:
                parts = n.module.split(".")
                if parts[0] == "vec_decisions":
                    if len(parts) > 1:
                        out.add(f"vec_decisions/{parts[1]}.py")
                    else:
                        out |= {f"vec_decisions/{a.name}.py" for a in n.names if (universe_dir / "vec_decisions" / f"{a.name}.py").exists()}
            elif isinstance(n, ast.Import):
                for a in n.names:
                    parts = a.name.split(".")
                    if parts[0] == "vec_decisions" and len(parts) > 1:
                        out.add(f"vec_decisions/{parts[1]}.py")
    return sorted(x for x in out if (universe_dir / x).exists())


# --------------------------------------------------------------------------------------------- vec reachability
def vec_reachability(vec_scans):
    """static BFS from simulate_one / compute_exit_signals over function names defined in the vec files. Returns (reachable_funcs{(file,qual)}, modules_called{file: bool})."""
    by_name = collections.defaultdict(list)
    for fs in vec_scans.values():
        for q, (node, calls) in fs.funcs.items():
            by_name[q.split(".")[-1]].append((fs.rel, q))
    entry = [(f, q) for n in ("simulate_one", "compute_exit_signals") for (f, q) in by_name.get(n, []) if f == "v12_quick_engine.py"]
    reach, todo = set(), list(entry)
    while todo:
        key = todo.pop()
        if key in reach:
            continue
        reach.add(key)
        fs = vec_scans.get(key[0])
        if not fs or key[1] not in fs.funcs:
            continue
        for c in fs.funcs[key[1]][1]:
            for k in by_name.get(c, []):
                if k not in reach:
                    todo.append(k)
    called = {}
    for fs in vec_scans.values():
        if fs.rel.startswith("vec_decisions/"):
            called[fs.rel] = any(k[0] == fs.rel for k in reach)
    return reach, called


# --------------------------------------------------------------------------------------------- templates
def read_templates():
    """{cat_side: {'rows': {name: [{tab,row,option,bold,is_default,grey,orange}]}, 'filters': {FILTER: {tab: [{opt,bold}]}}, 'tabs': {tab: {white, orange, first_default}}}}"""
    import openpyxl
    out = {}
    for cs, fn in TEMPLATES.items():
        import os as _os
        p = Path(_os.environ.get("SWITCH_BIBLE_TEMPLATE_DIR") or (ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM")) / fn   # USER 2026-10-07: the 4 independent templates; env override verifies a STAGED dir
        wb = openpyxl.load_workbook(str(p), read_only=False)
        rows = collections.defaultdict(list)
        filters = collections.defaultdict(lambda: collections.defaultdict(list))
        tabs = {}
        for tab in SWITCH_SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            hdr = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=2, column=c).value
                if isinstance(v, str):
                    hdr[v.strip().upper()] = c
                    if "=" in v:
                        f, o = v.split("=", 1)
                        filters[f.strip()][tab].append({"opt": o.strip(), "bold": bool(ws.cell(row=2, column=c).font and ws.cell(row=2, column=c).font.b)})
            c_isd = hdr.get("IS_DEFAULT")
            white = orange = 0
            first = None
            for r in range(3, ws.max_row + 1):
                a = ws.cell(row=r, column=1).value
                if a in (None, ""):
                    continue
                name = str(a).strip()
                fa = ws.cell(row=r, column=1).fill
                is_or = bool(fa and fa.fill_type == "solid" and "FFE699" in str(fa.fgColor.rgb))
                fc = ws.cell(row=r, column=1).font
                grey = bool(fc and fc.color is not None and str(getattr(fc.color, "rgb", "") or "").upper() in ("FFBFBFBF", "00BFBFBF"))
                isd = str(ws.cell(row=r, column=c_isd).value or "").strip().upper() if c_isd else ""
                bold = bool(ws.cell(row=r, column=2).font and ws.cell(row=r, column=2).font.b)
                if first is None:
                    first = isd == "YES"
                white, orange = (white, orange + 1) if is_or else (white + 1, orange)
                rows[name].append({"tab": tab, "row": r, "option": ws.cell(row=r, column=2).value, "bold": bold, "is_default": isd, "grey": grey, "orange": is_or})
            tabs[tab] = {"white": white, "orange": orange, "first_row_default": bool(first)}
        wb.close()
        out[cs] = {"rows": rows, "filters": filters, "tabs": tabs}
    return out


# --------------------------------------------------------------------------------------------- build
def kind_of(name, tabs):
    n = name.upper()
    tab = collections.Counter(tabs).most_common(1)[0][0] if tabs else ""
    if tab.startswith("EXIT"):
        return "exit"
    if tab.startswith("ENTRY"):
        return "entry"
    if tab.startswith("REENTRY"):
        return "reentry"
    if tab.startswith("AUGMENT"):
        return "augment"
    if tab.startswith("REDUCE"):
        return "reduce"
    if tab.startswith("STDEV"):
        return "sizing"
    if n.endswith("_FILTER_TF"):
        return "filter"
    for k, v in (("REENTRY", "reentry"), ("EXIT", "exit"), ("AUGMENT", "augment"), ("REDUCE", "reduce"), ("ENTRY", "entry"), ("SIZE", "sizing")):
        if k in n:
            return v
    return "global" if tab else "unclassified"


def engine_deploy_info():
    """data/engine_deploy/CURRENT.json if present, else the newest <ts>.json deploy note (Agent M protocol); + staged queue items (not yet deployed)."""
    d = DATA / "engine_deploy"
    cur = d / "CURRENT.json"
    pick = cur if cur.exists() else None
    if pick is None:
        notes = sorted(d.glob("*.json"), key=lambda p: p.stat().st_mtime) if d.exists() else []
        pick = notes[-1] if notes else None
    info = {"source": str(pick.relative_to(ROOT)) if pick else None, "md5": md5(pick) if pick else None}
    try:
        j = json.loads(pick.read_text()) if pick else {}
        info["files"] = j.get("files") or j.get("engine") or {}
        info["ts"] = j.get("ts_utc") or j.get("ts")
    except Exception:
        pass
    q = DATA / "wiring" / "queue"
    staged = set()
    if q.exists():
        for m in q.glob("*/*/MANIFEST.json"):
            try:
                if (m.parent / "DEPLOYED").exists():
                    continue
                staged |= {n for n in re.findall(r"[A-Z][A-Z0-9_]{5,}", m.read_text()) if NAME_RE.match(n)}
            except Exception:
                pass
    info["staged_names"] = sorted(staged)
    return info


def load_json(p, default=None):
    try:
        return json.loads(Path(p).read_text())
    except Exception:
        return default


def build():
    scans, md5s = {}, {}

    def scan(path, cls=None):
        p = ROOT / path
        if not p.exists():
            return None
        fs = FileScan(p, cls)
        scans[path] = fs
        md5s[path] = md5(p)
        return fs

    for f, c in CONFIG_CLASSES.items():
        scan(f, c)
    # live closures
    closures = {v: import_closure(ROOT / r, block=tuple(Path(o).stem for k, o in LIVE_ROOTS.items() if k != v)) for v, r in LIVE_ROOTS.items()}
    shared = {v: shared_vec_modules(closures[v]) for v in closures}
    for v in closures:
        closures[v] = sorted(set(closures[v]) | set(shared[v]))
    live_files = sorted(set(closures["crypto"]) | set(closures["stocks"]))
    for f in live_files:
        if f not in scans:
            scan(f)
    # vec files
    vec_files = list(VEC_FILES_FIXED) + sorted(rel(p) for p in (ROOT / "vec_decisions").glob("*.py"))
    for f in vec_files:
        if f not in scans:
            scan(f, CONFIG_CLASSES.get(f))
    vec_scans = {f: scans[f] for f in vec_files if f in scans}
    reach, mod_called = vec_reachability(vec_scans)

    cfg_defs = {k: scans[f].defs for f, k in CONFIG_CLASSES.items() if f in scans}  # {file: defs}
    config_defs, tradier_defs, quick_defs = (scans["config.py"].defs, scans["config_tradier.py"].defs, scans["v12_quick_engine.py"].defs)
    quick_tr_overrides = {}
    qs = scans["v12_quick_engine.py"]
    for q, d in qs.stores.items():
        if q.startswith("QuickConfig.") and "tradier" in q.lower():
            quick_tr_overrides.update(d)
    # literal values of apply_tradier_defaults assignments
    quick_tr_values = {}
    try:
        for node in ast.walk(qs.tree):
            if isinstance(node, ast.FunctionDef) and node.name == "apply_tradier_defaults":
                for st in ast.walk(node):
                    if isinstance(st, ast.Assign) and len(st.targets) == 1 and isinstance(st.targets[0], ast.Attribute) and isinstance(st.targets[0].value, ast.Name) and st.targets[0].value.id == "self":
                        quick_tr_values[st.targets[0].attr] = qs._lit(st.value)
    except Exception:
        pass

    # RUNTIME defaults (the AST sees duplicate class attributes / sub-classes; Python's real class value is the truth) — C2 bug fix 2026-10-01
    sys.path.insert(0, str(ROOT))
    try:
        import importlib
        _c = importlib.import_module("config").Config()
        _t = importlib.import_module("config_tradier").TradierConfig()
        _qm = importlib.import_module("v12_quick_engine")
        _q, _qt = _qm.QuickConfig(), _qm.QuickConfig()
        _qt.apply_tradier_defaults()
        for _obj, _dst in ((_c, config_defs), (_t, tradier_defs), (_q, quick_defs), (_qt, quick_tr_values)):
            for _k in dir(_obj):
                if NAME_RE.match(_k) and not _k.startswith("_"):
                    try:
                        _v = getattr(_obj, _k)
                    except Exception:
                        continue
                    if isinstance(_v, (bool, int, float, str, type(None))) or _dst is not quick_tr_values:
                        if isinstance(_v, (bool, int, float, str, type(None), list, tuple, dict, set)):
                            _dst[_k] = _v if isinstance(_v, (bool, int, float, str, type(None))) else repr(_v)[:80]
        defaults_source = "runtime"
    except Exception as _e:  # fall back to the AST values
        defaults_source = f"ast ({type(_e).__name__}: {str(_e)[:80]})"
    tpl = read_templates()
    csd = load_json(DATA / "cat_side_defaults_4.json", {})
    conflicts = (csd.get("_meta", {}) or {}).get("conflicts_quick_vs_live", {})
    wired_proven = set((load_json(DATA / "wiring" / "wired_proven.json", {}) or {}).get("switches", []))
    parity_proven = set(load_json(DATA / "wiring" / "parity_proven.json", []) or [])
    registry = (load_json(DATA / "wiring" / "vec_function_registry.json", {}) or {}).get("registry", {})
    unwired = load_json(DATA / "vec_unwired.json", {}) or {}
    vec_unwired = set(unwired.get("switches", []) + unwired.get("filters", []) + unwired.get("switches_manual", []) + unwired.get("filters_manual", []))
    live_gaps = load_json(DATA / "wiring" / "live_gaps.json", {}) or {}
    vec_only_json = load_json(DATA / "live_parity" / "switches_vec_only.json", {}) or {}
    staged_names = set(engine_deploy_info().get('staged_names', []))
    never_called_modules = sorted(m for m, c in mod_called.items() if not c)

    # universe
    names = set()
    for cs, t in tpl.items():
        names |= set(t["rows"]) | set(t["filters"])
    vec_read_names = set()
    for f, fs in vec_scans.items():
        for n, evs in fs.events.items():
            if any(e["kind"] == "read" for e in evs):
                vec_read_names.add(n)
    cfg_names = set(config_defs) | set(tradier_defs) | set(quick_defs)
    names |= (vec_read_names & cfg_names)
    for _v in ("crypto", "stocks"):  # config fields that live reads/refs (incl. dynamic-key proof literals) must be in the bible even if no template/vec read exists
        for _f in closures[_v]:
            _fs = scans.get(_f)
            if _fs:
                names |= {n for n, evs in _fs.events.items() if n in cfg_names and any(e["kind"] in ("read", "ref") and e["flag"] == "live" for e in evs)}
    names = {n for n in names if NAME_RE.match(n) or n in names}
    import os as _os2
    if _os2.environ.get("SWITCH_BIBLE_EXTRA_NAMES"):  # Agent G: trace names that exist only in old template versions
        names |= set(json.loads(Path(_os2.environ["SWITCH_BIBLE_EXTRA_NAMES"]).read_text()))

    def sites(name, files, kinds=("read",), only_flag=None):
        out = []
        for f in files:
            fs = scans.get(f)
            if not fs:
                continue
            for e in fs.events.get(name, []):
                if e["kind"] in kinds and (only_flag is None or e["flag"] == only_flag):
                    out.append({"file": f, "line": e["line"], "func": e["func"], "flag": e["flag"], "kind": e["kind"]})
        out.sort(key=lambda s: (s["file"], s["line"]))
        return out

    def reachable_vec(site):
        if site["flag"] != "live":
            return False
        fs = vec_scans.get(site["file"])
        if site["file"].startswith("vec_decisions/"):
            return mod_called.get(site["file"], False)
        # v12/evaluate: function-level reachability from simulate_one (class-body/def sites excluded)
        q = site["func"]
        if q == "<module>":
            return False
        return (site["file"], q) in reach or any((site["file"], q2) in reach and q2.startswith(q + ".") for q2 in [k[1] for k in reach if k[0] == site["file"]])

    sw = {}
    for name in sorted(names):
        in_cfg = {"config.py": name in config_defs, "config_tradier.py": name in tradier_defs, "QuickConfig": name in quick_defs}
        t_presence, tabs_all = {}, []
        for cs, t in tpl.items():
            lst = t["rows"].get(name, [])
            fl = t["filters"].get(name, {})
            if lst or fl:
                t_presence[cs] = {"rows": [{"tab": r["tab"], "row": r["row"], "option": r["option"], "default": r["is_default"] == "YES", "grey": r["grey"], "orange": r["orange"]} for r in lst],
                                  "yellow_header_tabs": {tb: [o["opt"] for o in v] for tb, v in fl.items()}}
                tabs_all += [r["tab"] for r in lst if not r["orange"]] or [r["tab"] for r in lst] + list(fl)
        defaults = {"config.py": config_defs.get(name, "<absent>"), "config_tradier.py": tradier_defs.get(name, "<absent>"), "QuickConfig": quick_defs.get(name, "<absent>"),
                    "QuickConfig_apply_tradier_defaults": quick_tr_values.get(name, "<n/a>"),
                    "cat_side_defaults_4": {cs: csd.get(cs, {}).get(name, "<absent>") for cs in TEMPLATES},
                    "template_bold": {cs: [r["option"] for r in tpl[cs]["rows"].get(name, []) if r["is_default"] == "YES"] for cs in TEMPLATES if name in tpl[cs]["rows"]}}
        live_all = {v: [s for f in closures[v] for s in sites(name, [f])] for v in ("crypto", "stocks")}
        live_reads = {v: [s for s in lst if s["flag"] == "live"] for v, lst in live_all.items()}
        # dynamic-key reads (N2 2026-10-01): a live f-string key pattern whose file ALSO names the full key as a literal (proof tuple) is a live read of that name
        for v in ("crypto", "stocks"):
            if live_reads[v]:
                continue
            for f in closures[v]:
                fs = scans.get(f)
                if not fs or not fs.dyn or not any(e["kind"] in ("ref", "read") for e in fs.events.get(name, [])):
                    continue
                for (ln, fn, rx) in fs.dyn:
                    if re.match(rx, name):
                        live_reads[v].append({"file": f, "line": ln, "func": fn + " [dynamic key]", "flag": "live", "kind": "read"})
                        break
        vec_all = sites(name, vec_files)
        vec_reads = [dict(s, reachable=reachable_vec(s)) for s in vec_all]
        vec_ok = [s for s in vec_reads if s["reachable"]]
        if name.startswith("TRADIER_") and not vec_ok:  # dynamic twin: get("TRADIER_" + name) in the vec module (dc_channel_exits.py _b())
            _base = name[len("TRADIER_"):]
            _bs = [dict(x, reachable=reachable_vec(x), via="TRADIER_ prefix concat") for x in sites(_base, vec_files) if x["flag"] == "live" and '"TRADIER_" +' in (ROOT / x["file"]).read_text(errors="replace")]
            vec_reads = vec_reads + [x for x in _bs if x["reachable"]]
            vec_ok = [x for x in vec_reads if x["reachable"]]
        refs_live = {v: [s for f in closures[v] for s in sites(name, [f], ("ref",))][:3] for v in ("crypto", "stocks")}
        status_v = {}
        for v in ("crypto", "stocks"):
            lv, vv = bool(live_reads[v]), bool(vec_ok)
            if not any(in_cfg.values()):
                st = "NOT_IN_CONFIG"
            elif lv and vv:
                st = "WIRED_BOTH_PARITY_PROVEN" if name in parity_proven else "WIRED_BOTH_UNPROVEN"
            elif vv:
                st = "VEC_ONLY"
            elif lv:
                st = "LIVE_ONLY"
            else:
                st = "DEAD"
            if st in ("LIVE_ONLY", "DEAD", "NOT_IN_CONFIG") and name in staged_names:
                st += "+STAGED_VEC"
            status_v[v] = st
        kind = (registry.get(name) or {}).get("kind") or kind_of(name, tabs_all)
        sw[name] = {
            "name": name, "kind": kind, "type": type(next((d for d in (config_defs.get(name), tradier_defs.get(name), quick_defs.get(name)) if d is not None), None)).__name__,
            "in_config": in_cfg, "defaults": defaults,
            "live_reads": {v: [f"{s['file']}:{s['line']}({s['func']})" for s in lst][:12] for v, lst in live_reads.items()},
            "live_read_count": {v: len(lst) for v, lst in live_reads.items()},
            "live_dead_or_stub": {v: len([s for s in live_all[v] if s["flag"] != "live"]) for v in live_all},
            "live_weak_refs": {v: [f"{s['file']}:{s['line']}" for s in lst] for v, lst in refs_live.items() if lst},
            "vec_reads": [f"{s['file']}:{s['line']}({s['func']})" + ("" if s["reachable"] else f" [{'UNREACHABLE' if s['flag']=='live' else s['flag'].upper()}]") for s in vec_reads][:14],
            "vec_read_count": {"reachable": len(vec_ok), "total": len(vec_reads)},
            "filter_tf": name.endswith("_FILTER_TF"),
            "template": t_presence,
            "sides": sorted({cs.split("_")[1] for cs in t_presence}) or [],
            "venues_in_templates": sorted({cs.split("_")[0] for cs in t_presence}),
            "status": status_v, "ledger_flip_proven": name in wired_proven, "parity_proven": name in parity_proven,
            "vec_unwired_listed": name in vec_unwired,
            "agent_c_registry": {k: (registry.get(name) or {}).get(k) for k in ("live_counterpart", "suggested_home_tab", "valid_options", "options_need_review")} if name in registry else None,
            "conflict_quick_vs_live": {cs: name in (conflicts.get(cs) or []) for cs in TEMPLATES if name in (conflicts.get(cs) or [])},
            "recently_added": (registry.get(name) or {}).get("recently_added"),
        }

    # filter-TF machinery (generic suffix handling) — real scan of v12 + vec for '_FILTER_TF' string handling
    machinery = []
    for f in vec_files:
        fs = vec_scans.get(f)
        if not fs:
            continue
        try:
            for node in ast.walk(fs.tree):
                if isinstance(node, ast.Constant) and isinstance(node.value, str) and "_FILTER_TF" in node.value and not NAME_RE.match(node.value):
                    machinery.append(f"{f}:{node.lineno}")
        except Exception:
            pass
    meta = {
        "generated": datetime.datetime.utcnow().isoformat() + "Z", "builder": "tools/build_switch_bible.py", "n_switches": len(sw), "defaults_source": defaults_source, "engine_deploy": engine_deploy_info(),
        "files_md5": md5s, "live_closure": closures, "vec_files": vec_files,
        "vec_modules_never_called_from_simulate_one": never_called_modules, "vec_modules_total": len(mod_called), "vec_modules_all": sorted(mod_called), "shared_vec_live_modules": shared,
        "filter_tf_generic_handling_sites": machinery[:40],
        "template_tabs": {cs: t["tabs"] for cs, t in tpl.items()},
        "inputs": {"wired_proven": len(wired_proven), "parity_proven": len(parity_proven), "vec_unwired": len(vec_unwired), "live_gaps": len(live_gaps.get("gaps", [])),
                   "switches_vec_only_no_live_mention": len(vec_only_json.get("no_live_mention", []))},
        "cat_side_defaults_conflicts": {cs: len(v) for cs, v in conflicts.items()},
    }
    return {"meta": meta, "switches": sw}


# --------------------------------------------------------------------------------------------- reporting
RULES = """# SWITCH BIBLE (generated by tools/build_switch_bible.py — do not hand-edit; regenerate)

## BIBLE_SWITCH_RULES — read before adding / moving / renaming ANY switch or filter
1. **Touch list (all must change in the same batch, then `python tools/build_switch_bible.py && python tools/verify_switch_bible.py` must pass):**
   `config.py` (Config) · `config_tradier.py` (TradierConfig) · `v12_quick_engine.py` QuickConfig (+ `apply_tradier_defaults` if the stock default differs) ·
   `data/cat_side_defaults_4.json` (rebuild: `tools/build_cat_side_defaults_4.py --stage-only`) · live read site in `ez_manage.py` (crypto) and/or `tradier_manage.py` (stocks) ·
   vector read site in `v12_quick_engine.simulate_one` / `compute_exit_signals` or a `vec_decisions/` module **called from simulate_one** ·
   TEMPLATE row(s) in the correct tab AND correct side/venue templates (white switch row ONCE per template; orange filter rows in every tab; yellow header columns) ·
   delta-log key `TAB!row:SWITCH=value` (nothing renames a key silently) · `data/wiring/*` status files.
2. **A switch is WIRED only when both a live read and a REACHABLE vector read exist** (this bible's scan). Stub reads (`_ = getattr(...)`), reads inside early-return/dead functions and
   vec_decisions modules never called from simulate_one do NOT count. A vector number for a switch that live never reads is a lie (BACKTEST_BIBLE §19).
3. **Renames:** keep the old name as a deprecated alias until the batch is deployed; update every surface above + template row + `data/wiring/renames.json`.
4. **Moving a template row:** whole row only (every column incl. yellow cells), via `tools/template_row_guard.py`-verified writers; a tab must start with a default row;
   every tab keeps its orange filter rows; one bold default per switch/filter.
5. **Never fabricate:** every file:line here comes from an AST scan. If the scan and your belief disagree, the scan wins — fix the code or the scan, never the bible by hand.
6. **PARITY_PROVEN** is reserved for scalar-vs-vec parity evidence (list in `data/wiring/parity_proven.json`). `ledger_flip_proven` = a real NPZ ledger flip (data/wiring/wired_proven.json) only.

"""


def short(lst, n=2):
    return "<br>".join(lst[:n]) + (f"<br>(+{len(lst) - n})" if len(lst) > n else "") if lst else "—"


def write_md(bible, path):
    sw, meta = bible["switches"], bible["meta"]
    by_tab = collections.defaultdict(list)
    for n, s in sw.items():
        tabs = collections.Counter(r["tab"] for t in s["template"].values() for r in t["rows"] if not r["orange"])
        home = tabs.most_common(1)[0][0] if tabs else ("(orange/filter only)" if s["template"] else "(not in any template)")
        by_tab[home].append(n)
    L = [RULES]
    cnt = collections.Counter()
    for s in sw.values():
        for v, st in s["status"].items():
            cnt[(v, st)] += 1
    L.append(f"_Generated {meta['generated']} · {meta['n_switches']} names · vec modules never reached from simulate_one: {len(meta['vec_modules_never_called_from_simulate_one'])}/{meta['vec_modules_total']}_\n")
    L.append("## Status counts (all names)\n\n| venue | " + " | ".join(sorted({k[1] for k in cnt})) + " |\n|---|" + "---|" * len({k[1] for k in cnt}))
    for v in ("crypto", "stocks"):
        L.append(f"| {v} | " + " | ".join(str(cnt.get((v, st), 0)) for st in sorted({k[1] for k in cnt})) + " |")
    L.append("\n## Scanned files (md5)\n\n" + "\n".join(f"- `{f}` `{m[:10]}`" for f, m in sorted(meta["files_md5"].items())) + "\n")
    for tab in sorted(by_tab):
        L.append(f"\n## {tab} ({len(by_tab[tab])})\n\n| switch | kind | defaults cfg / tradier / quick | templates | crypto | stocks | live reads (ez · tradier) | vec reads |\n|---|---|---|---|---|---|---|---|")
        for n in sorted(by_tab[tab]):
            s = sw[n]
            d = s["defaults"]
            dd = f"{d['config.py']} / {d['config_tradier.py']} / {d['QuickConfig']}"
            tp = ",".join(cs.replace("CRYPTO", "C").replace("STOCKS", "S").replace("_LONG", "L").replace("_SHORT", "S") for cs in s["template"]) or "—"
            L.append(f"| `{n}` | {s['kind']} | {str(dd)[:70]} | {tp} | {s['status']['crypto']} | {s['status']['stocks']} | {short(s['live_reads']['crypto'], 1)} · {short(s['live_reads']['stocks'], 1)} | {short(s['vec_reads'], 2)} |")
    Path(path).write_text("\n".join(L) + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(DATA / "SWITCH_BIBLE.json"))
    ap.add_argument("--md", default=str(ROOT / "SWITCH_BIBLE.md"))
    ap.add_argument("--no-md", action="store_true")
    a = ap.parse_args()
    out = Path(a.out)
    _t0 = time.time()
    if out.exists():  # keep the previous bible for --since-md5 diffs
        (out.parent / "SWITCH_BIBLE.prev.json").write_bytes(out.read_bytes())
    bible = build()
    if time.time() - _t0 > 900:  # zombie guard 2026-10-05: a wedged scan must never overwrite a fresher bible (02:40Z -12 incident)
        print(f"[STALE-SKIP] scan took {time.time() - _t0:.0f}s > 900s; not writing {out}", file=sys.stderr)
        sys.exit(2)
    tmp = out.with_suffix(".tmp")
    tmp.write_text(json.dumps(bible, indent=1, default=str, sort_keys=True))
    tmp.replace(out)
    if not a.no_md:
        write_md(bible, a.md)
    m = bible["meta"]
    print(f"[bible] {m['n_switches']} names; files scanned {len(m['files_md5'])}; live closure ez={len(m['live_closure']['crypto'])} tr={len(m['live_closure']['stocks'])}; vec modules never called {len(m['vec_modules_never_called_from_simulate_one'])}/{m['vec_modules_total']}")
    c = collections.Counter((v, s["status"][v]) for s in bible["switches"].values() for v in ("crypto", "stocks"))
    for k in sorted(c):
        print(f"  {k[0]:7s} {k[1]:26s} {c[k]}")
    print(f"[written] {out}" + ("" if a.no_md else f"  {a.md}"))


if __name__ == "__main__":
    main()
