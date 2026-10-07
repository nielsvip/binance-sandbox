#!/usr/bin/env python3
"""v15_yellow_365_audit — audit ONLY the yellow cells of the cat_side TEMPLATES on the 365D window for a handful of sym_sides per
cat_side, starting from each sym_side's latest final set of the existing runs (USER 2026-09-30).

Why: which yellow (switch=cand x FILTER=opt) cells are really worth testing? Delta per cell on 365D for 4-5 syms per cat_side gives
pos_sym per cell in hours; cells that are zero/negative on every audited sym are candidates for un-yellow, pos_sym>0 cells stay.

Evaluation = the SAME path v15_pilot uses: tools.opt.v12_pilot.evaluate_prepared_sanitized(prepared_365D, overrides, 365) on one
prepared NPZ slice per sym (forked workers share it copy-on-write). Override construction copies v15_pilot._row_plan exactly:
  start set (cumulative_overrides of the sym_side's best progress JSON) + _switch_overrides(switch, cand) [+ filter=opt], sanitized.
NO-LIES: only real engine results; invalid results are stored with their reason and never fabricated; a cell equal to the naked
switch shows delta_vs_naked 0.0 (non-binding). Deltas: delta_vs_base = cell - baseline(start set); delta_vs_naked = cell - naked
(start set + switch=cand). Dedupe: identical override sets (same key across tabs / equal-to-default options) are evaluated once.

Modes
  --dry-run                     eval counts + time estimate per cat_side (no NPZ needed; works on the Mac)
  --cat-side CRYPTO_LONG ...    run on THIS host (crypto on s1/s5, stocks on s2/s5); reads local progress JSONs + NPZ
  --sym-slice 0:2               run only syms 0..1 of the (cached, deterministic) selection (staged plan, pass 1)
  --only-pos-cells              restrict to cells that were positive (delta_vs_naked>0) on the syms already computed (pass 2)
  --summarize                   merge all <cat_side> results found in --dir into summary.json (keep-yellow / un-yellow lists)
Resumable: every eval is appended to results_<cat>__<sym>.jsonl (atomic line writes); a rerun skips what is done.
Launch later, niced:  nice -n 19 python tools/v15_yellow_365_audit.py --cat-side STOCKS_LONG --syms 5 --workers 6 --nice 19
"""
import argparse
import datetime
import hashlib
import json
import multiprocessing as mp
import os
import socket
import statistics
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
BRIGHT, LIGHT = "FFFF00", "FFF2CC"
HDR = 2
EPS = 1e-9
MIN_SPAN_DAYS = 330
MIN_BASE_TRADES_365 = 30
PROBE = {"CRYPTO_LONG": "BTCUSDC_LONG", "CRYPTO_SHORT": "BTCUSDC_SHORT", "STOCKS_LONG": "AAPL_LONG", "STOCKS_SHORT": "AAPL_SHORT"}
DEFAULT_COST = {"CRYPTO": 2.0, "STOCKS": 1.5}  # sec/eval until a real measurement exists (data/yellow_365_audit/measured_cost.json)

_PREP = None  # set in the parent before the fork pool is created (copy-on-write share)
INCLUDE_TOKEN = False  # --include-token-cells: also every filter cell whose filter name shares >=1 "_" token with the switch name (colour "token")


def out_dir(a):
    d = Path(a.dir) if a.dir else ROOT / "data" / "yellow_365_audit" / datetime.date.today().strftime("%Y%m%d")
    d.mkdir(parents=True, exist_ok=True)
    return d


def atomic_write(path: Path, obj):
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, indent=1, default=str))
    tmp.replace(path)


def _same_val(a, b) -> bool:
    # copy of the local helper in v15_pilot._spec_fill_workbook
    if isinstance(a, bool) or isinstance(b, bool) or str(a).lower() in ("true", "false") or str(b).lower() in ("true", "false"):
        return str(a).strip().lower() == str(b).strip().lower()
    try:  # USER 2026-10-07: exact-equality fast path — abs(inf-inf)=nan made inf NEVER equal inf
        fa, fb = float(a), float(b)
        return fa == fb or abs(fa - fb) < 1e-12
    except Exception:
        return str(a).strip() == str(b).strip()


def ck_of(ov: dict) -> str:
    return hashlib.md5(json.dumps(sorted((k, repr(v)) for k, v in ov.items())).encode()).hexdigest()


# ───────────────────────── template → yellow cells ─────────────────────────
def collect_cells(cs: str, P, cache_dir: Path):
    """[(tab, switch, cand, hdr, colour)] of every yellow cell (bright FFFF00 / light FFF2CC), parsed like the pilot does."""
    cf = cache_dir / (f"cells_{cs}_tok.json" if INCLUDE_TOKEN else f"cells_{cs}.json")
    tpl = ROOT / "SPREADSHEETS" / f"TEMPLATE_{cs}.xlsx"
    if cf.exists() and cf.stat().st_mtime > tpl.stat().st_mtime:
        return [tuple(x) for x in json.loads(cf.read_text())]
    import openpyxl
    wb = openpyxl.load_workbook(str(tpl))
    known = P.known_config_fields()
    cells = []
    for tab in SWITCH_SHEETS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        pos_col = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(row=HDR, column=c).value or "").strip().upper() == "POS_SYM"), 14)
        hcols = {}
        for c in range(pos_col + 1, ws.max_column + 1):
            h = ws.cell(row=HDR, column=c).value
            if isinstance(h, str) and "=" in h and h.split("=", 1)[0].strip() in known:
                hcols[c] = h.strip()
        for r in range(HDR + 1, ws.max_row + 1):
            sw, cand = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
            if sw in (None, "") or cand in (None, "") or str(cand).strip().endswith("_ALT"):
                continue
            for c, h in hcols.items():
                f = ws.cell(row=r, column=c).fill
                rgb = str(f.fgColor.rgb or "").upper() if (f is not None and f.fill_type == "solid" and f.fgColor is not None) else ""
                col = "bright" if rgb.endswith(BRIGHT) else "light" if rgb.endswith(LIGHT) else None
                if col is None and INCLUDE_TOKEN:
                    _st = {x for x in str(sw).upper().split("_") if len(x) > 1}
                    _ft = {x for x in h.split("=", 1)[0].upper().split("_") if len(x) > 1}
                    if _st & _ft:
                        col = "token"
                if col:
                    cells.append((tab, str(sw).strip(), str(cand).strip(), h, col))
    atomic_write(cf, cells)
    return cells


def unique_cells(cells):
    """(switch, cand, hdr) -> {'tabs': [...], 'colour': 'bright'|'light'} (the same cell in several tabs is evaluated once)."""
    u = {}
    for tab, sw, cand, hdr, col in cells:
        e = u.setdefault((sw, cand, hdr), {"tabs": [], "colour": col})
        e["tabs"].append(tab)
        _rank = {"bright": 0, "light": 1, "token": 2}
        if _rank.get(col, 9) < _rank.get(e["colour"], 9):
            e["colour"] = col
    return u


# ───────────────────────── eval plan (mirrors v15_pilot._row_plan) ─────────────────────────
def build_plan(start: dict, defaults: dict, ucells: dict, P):
    """returns (baseline_ov, naked{(sw,cand)->ov|None(is_running)}, cell{(sw,cand,hdr)->ov})"""
    base = P.sanitize_overrides(dict(start), defaults)[0]
    naked, cell = {}, {}
    for (sw, cand, hdr) in ucells:
        nk = (sw, cand)
        if nk not in naked:
            sw_ov = P._switch_overrides(sw, P._parse_opt_value(cand, defaults.get(sw)))
            is_running = all(_same_val(start.get(k2, defaults.get(k2)), v2) for k2, v2 in sw_ov.items())
            sv = dict(start)
            sv.update(sw_ov)
            sv = P.sanitize_overrides(sv, defaults)[0]
            naked[nk] = (None if is_running else sv, sv)
        sv = naked[nk][1]
        filt, opt = hdr.split("=", 1)
        v = dict(sv)
        v[filt.strip()] = P._parse_opt_value(opt.strip(), defaults.get(filt.strip()))
        cell[(sw, cand, hdr)] = P.sanitize_overrides(v, defaults)[0]
    return base, naked, cell


def count_evals(base, naked, cell):
    keys = {ck_of(base)}
    for nk, (ov, sv) in naked.items():
        keys.add(ck_of(sv))
    for ov in cell.values():
        keys.add(ck_of(ov))
    return len(keys)


# ───────────────────────── evaluation ─────────────────────────
def _init_worker(nice):
    if nice:
        try:
            os.nice(nice)
        except OSError:
            pass


def _work(item):
    ck, ov = item
    from tools.opt.v12_pilot import evaluate_prepared_sanitized as eps
    t0 = time.time()
    try:
        r = eps(_PREP, ov, 365) or {}
        rec = {k: r.get(k) for k in ("gain_pct", "trades", "valid", "invalid_reason", "tim_pct", "max_dd_pct", "bh_pct")}
    except Exception as e:  # recorded, never fabricated
        rec = {"gain_pct": None, "trades": None, "valid": False, "invalid_reason": f"ERR {e}"[:160]}
    rec["secs"] = round(time.time() - t0, 3)
    return ck, rec


def span_days(prepared):
    try:
        t = prepared["npz_prepared"]["timestamps"]
        t = t / 1000 if t[-1] > 1e11 else t
        return float((t[-1] - t[0]) / 86400)
    except Exception:
        return 0.0


def local_candidates(cs, P, host_hint):
    """selection entries of this cat_side whose progress JSON exists on THIS host, with its 30D final gain."""
    sel = json.loads((ROOT / "data" / "avg_delta_selection.json").read_text())
    out = []
    for ss, e in sel.items():
        if P.map_key_for_symside(ss) != cs:
            continue
        p = e.get("path")
        if not p or not os.path.exists(p):
            continue
        try:
            d = json.load(open(p))
        except Exception:
            continue
        ov = d.get("cumulative_overrides") or {}
        g = d.get("final_gain", d.get("cumulative_gain"))
        if len(ov) < 50 or g is None:  # a complete final set carries every bold default (~235 keys); tiny sets = partial/iso runs
            continue
        out.append({"symside": ss, "path": p, "gain30": float(g), "n_ov": len(ov)})
    return sorted(out, key=lambda x: x["gain30"])


def pick_pool(cands, n):
    """diverse, deterministic order: alternate best / worst / median / 2nd-best / 2nd-worst ... (mix of pos and neg 30D)."""
    if not cands:
        return []
    order, lo, hi = [], 0, len(cands) - 1
    mid = len(cands) // 2
    seq = [hi, lo, mid, hi - 1, lo + 1, mid + 1, hi - 2, lo + 2, mid - 1]
    seen = set()
    for i in seq + list(range(len(cands))):
        if 0 <= i < len(cands) and i not in seen:
            seen.add(i)
            order.append(cands[i])
    return order[: max(n * 4, 12)]


def select_syms(cs, n, P, d: Path, workers):
    sf = d / f"selection_{cs}.json"
    if sf.exists():
        return json.loads(sf.read_text())
    from tools.opt import evaluate_v12 as E
    from tools.opt.v12_pilot import evaluate_prepared_sanitized as eps
    cands = local_candidates(cs, P, socket.gethostname())
    chosen, rejected = [], []
    for c in pick_pool(cands, n):
        if len(chosen) >= n:
            break
        try:
            prep = E.prepare(c["symside"], 365)
        except Exception as e:
            rejected.append({**c, "reason": f"prepare error {e}"[:100]})
            continue
        if prep is None:
            rejected.append({**c, "reason": "no NPZ / prepare None"})
            continue
        sp = span_days(prep)
        if sp < MIN_SPAN_DAYS:
            rejected.append({**c, "reason": f"365D span {sp:.0f}d < {MIN_SPAN_DAYS}"})
            continue
        start = json.load(open(c["path"])).get("cumulative_overrides") or {}
        defaults = P.get_defaults_for_symside(c["symside"])
        r = eps(prep, P.sanitize_overrides(dict(start), defaults)[0], 365) or {}
        if int(r.get("trades") or 0) < MIN_BASE_TRADES_365:
            rejected.append({**c, "reason": f"baseline 365D trades {r.get('trades')} < {MIN_BASE_TRADES_365}"})
            continue
        chosen.append({**c, "span_days": round(sp, 1), "base365_gain": r.get("gain_pct"), "base365_trades": r.get("trades"), "base365_valid": r.get("valid")})
        del prep
    sel = {"cat_side": cs, "host": socket.gethostname(), "chosen": chosen, "rejected": rejected, "at": datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None).isoformat() + "Z"}
    atomic_write(sf, sel)
    return sel


def load_results(path: Path):
    res = {}
    if path.exists():
        for line in path.read_text().splitlines():
            try:
                j = json.loads(line)
                res[j["ck"]] = j["res"]
            except Exception:
                pass
    return res


def positive_cells(cs, d: Path, ucells, P):
    """cells positive (delta_vs_naked > 0, cell valid) on any sym already computed in this cat_side."""
    pos = set()
    for f in d.glob(f"results_{cs}__*.jsonl"):
        ss = f.stem.split("__", 1)[1]
        res = load_results(f)
        meta = json.loads((d / f"plan_{cs}__{ss}.json").read_text()) if (d / f"plan_{cs}__{ss}.json").exists() else None
        if not meta:
            continue
        for key, info in meta["cells"].items():
            c, nk = res.get(info["ck"]), res.get(meta["naked"].get(info["nk"]) or "")
            if c and nk and c.get("valid") and c.get("gain_pct") is not None and nk.get("gain_pct") is not None and c["gain_pct"] - nk["gain_pct"] > EPS:
                pos.add(tuple(key.split("\t")))
    return pos


def run_cat(cs, a, P, d: Path):
    global _PREP
    from tools.opt import evaluate_v12 as E
    cells = collect_cells(cs, P, d)
    ucells = unique_cells(cells)
    if getattr(a, "tier", "all") == "yellow":
        ucells = {k: v for k, v in ucells.items() if v["colour"] in ("bright", "light")}
    elif getattr(a, "tier", "all") == "token":
        ucells = {k: v for k, v in ucells.items() if v["colour"] == "token"}
    print(f"[{cs}] tier={getattr(a, 'tier', 'all')}: {len(ucells)} unique cells", flush=True)
    sel = select_syms(cs, a.syms, P, d, a.workers)
    lo, hi = (int(x) for x in a.sym_slice.split(":")) if a.sym_slice else (0, len(sel["chosen"]))
    chosen = sel["chosen"][lo:hi]
    restrict = positive_cells(cs, d, ucells, P) if a.only_pos_cells else None
    if restrict is not None:
        ucells = {k: v for k, v in ucells.items() if k in restrict}
        print(f"[{cs}] --only-pos-cells: {len(ucells)} cells positive on earlier syms", flush=True)
    done_evals, t_start, total_secs = 0, time.time(), []
    for c in chosen:
        ss = c["symside"]
        start = json.load(open(c["path"])).get("cumulative_overrides") or {}
        defaults = P.get_defaults_for_symside(ss)
        base, naked, cell = build_plan(start, defaults, ucells, P)
        rf = d / f"results_{cs}__{ss}.jsonl"
        res = load_results(rf)
        want = {ck_of(base): base}
        for nk, (ov, sv) in naked.items():
            want[ck_of(sv)] = sv
        for ov in cell.values():
            want[ck_of(ov)] = ov
        plan_meta = {"symside": ss, "base": ck_of(base), "naked": {f"{k[0]}={k[1]}": ck_of(v[1]) for k, v in naked.items()},
                     "cells": {"\t".join(k): {"ck": ck_of(ov), "nk": f"{k[0]}={k[1]}", "tabs": ucells[k]["tabs"], "colour": ucells[k]["colour"]} for k, ov in cell.items()},
                     "naked_is_running": [f"{k[0]}={k[1]}" for k, v in naked.items() if v[0] is None]}
        old = d / f"plan_{cs}__{ss}.json"
        if old.exists() and (a.only_pos_cells or getattr(a, "tier", "all") != "all"):  # keep the full plan from pass 1 / the other tier, merge the new cells
            prev = json.loads(old.read_text())
            prev["naked"].update(plan_meta["naked"])
            prev["cells"].update(plan_meta["cells"])
            plan_meta = prev
        atomic_write(old, plan_meta)
        pending = [(ck, ov) for ck, ov in want.items() if ck not in res]
        print(f"[{cs}] {ss}: unique evals {len(want)} done {len(want) - len(pending)} pending {len(pending)} (cells {len(cell)} naked {len(naked)}) start-set {len(start)} keys", flush=True)
        if a.max_evals and done_evals + len(pending) > a.max_evals:
            # baseline + naked first (they anchor every delta), then cells
            pending = pending[: max(0, a.max_evals - done_evals)]
        if not pending:
            continue
        prep = E.prepare(ss, 365)
        if prep is None or span_days(prep) < MIN_SPAN_DAYS:
            print(f"[{cs}] {ss}: prepare failed / span < {MIN_SPAN_DAYS}d — skipped", flush=True)
            continue
        _PREP = prep
        # baseline + naked first
        first = {ck_of(base)} | {ck_of(v[1]) for v in naked.values()}
        pending.sort(key=lambda it: (0 if it[0] in first else 1, it[0]))  # anchors first, then cells spread across switches (deterministic)
        ctx = mp.get_context("fork")
        with ctx.Pool(a.workers, initializer=_init_worker, initargs=(a.nice,)) as pool, open(rf, "a") as fh:
            for ck, rec in pool.imap_unordered(_work, pending, chunksize=1):
                fh.write(json.dumps({"ck": ck, "res": rec}, default=str) + "\n")
                fh.flush()
                res[ck] = rec
                done_evals += 1
                total_secs.append(rec["secs"])
                if done_evals % 50 == 0:
                    el = time.time() - t_start
                    print(f"[{cs}] {ss}: {done_evals} evals, mean {statistics.fmean(total_secs):.2f}s/eval (cpu), wall {el:.0f}s", flush=True)
                if a.max_secs and time.time() - t_start > a.max_secs:
                    print(f"[{cs}] --max-secs reached — stopping (resumable)", flush=True)
                    pool.terminate()
                    break
        _PREP = None
        del prep
        if a.max_secs and time.time() - t_start > a.max_secs:
            break
        if a.max_evals and done_evals >= a.max_evals:
            break
    if total_secs:
        mc = ROOT / "data" / "yellow_365_audit" / "measured_cost.json"
        cur = json.loads(mc.read_text()) if mc.exists() else {}
        cur[cs] = {"mean_cpu_sec_per_eval": round(statistics.fmean(total_secs), 3), "median": round(statistics.median(total_secs), 3), "n": len(total_secs), "workers": a.workers, "host": socket.gethostname(), "wall_s": round(time.time() - t_start, 1), "at": datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None).isoformat() + "Z"}
        mc.parent.mkdir(parents=True, exist_ok=True)
        atomic_write(mc, cur)
    print(f"[{cs}] done: {done_evals} new evals in {time.time() - t_start:.0f}s", flush=True)


# ───────────────────────── summary ─────────────────────────
def summarize(d: Path, min_n: int):
    summary = {"at": datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None).isoformat() + "Z", "min_n_for_unyellow": min_n,
               "caveat": "4-5 syms per cat_side only: pos_sym>0 cells are evidence to KEEP; zero/neg on all audited syms with n>=min_n is only a CANDIDATE for un-yellow (small sample, start sets are each sym's own winners); n<min_n = insufficient.", "cat_sides": {}}
    for cs in CAT_SIDES:
        plans = sorted(d.glob(f"plan_{cs}__*.json"))
        if not plans:
            continue
        cellrec = {}
        syms = []
        for pf in plans:
            meta = json.loads(pf.read_text())
            ss = meta["symside"]
            res = load_results(d / f"results_{cs}__{ss}.jsonl")
            base = res.get(meta["base"])
            if not base or base.get("gain_pct") is None:
                continue
            syms.append({"symside": ss, "base365_gain": base["gain_pct"], "base365_trades": base.get("trades"), "base365_valid": base.get("valid")})
            for key, info in meta["cells"].items():
                c = res.get(info["ck"])
                if c is None:
                    continue
                nk = res.get(meta["naked"].get(info["nk"]) or "")
                rec = cellrec.setdefault(key, {"tabs": info["tabs"], "colour": info["colour"], "per_sym": {}})
                zero = c.get("gain_pct") is None or int(c.get("trades") or 0) == 0
                e = {"valid": bool(c.get("valid")), "reason": c.get("invalid_reason"), "trades": c.get("trades"), "gain": c.get("gain_pct")}
                if not zero:
                    e["delta_vs_base"] = round(c["gain_pct"] - base["gain_pct"], 6)
                    if nk and nk.get("gain_pct") is not None and int(nk.get("trades") or 0) > 0:
                        e["delta_vs_naked"] = round(c["gain_pct"] - nk["gain_pct"], 6)
                        e["naked_valid"] = bool(nk.get("valid"))
                rec["per_sym"][ss] = e
        keep, unyellow, insufficient = [], [], []
        for key, rec in cellrec.items():
            vals = [(s, e) for s, e in rec["per_sym"].items() if e["valid"] and "delta_vs_naked" in e]
            rec["n"] = len(vals)
            rec["pos_sym"] = sum(1 for _, e in vals if e["delta_vs_naked"] > EPS)
            rec["avg_delta_vs_naked"] = round(statistics.fmean([e["delta_vs_naked"] for _, e in vals]), 6) if vals else None
            rec["pos_sym_vs_base"] = sum(1 for e in rec["per_sym"].values() if e["valid"] and e.get("delta_vs_base", 0) > EPS)
            if rec["pos_sym"] > 0:
                keep.append(key)
            elif rec["n"] >= min_n:
                unyellow.append(key)
            else:
                insufficient.append(key)
        summary["cat_sides"][cs] = {"syms": syms, "cells_evaluated": len(cellrec), "keep_yellow": len(keep), "unyellow_candidates": len(unyellow), "insufficient": len(insufficient)}
        atomic_write(d / f"{cs}.json", {"cat_side": cs, "syms": syms, "cells": cellrec, "keep_yellow": sorted(keep), "unyellow_candidates": sorted(unyellow), "insufficient_data": sorted(insufficient)})
        print(f"[{cs}] syms={len(syms)} cells={len(cellrec)} keep(pos_sym>0)={len(keep)} unyellow-candidates={len(unyellow)} insufficient={len(insufficient)}")
    atomic_write(d / "summary.json", summary)
    print(f"[summary] {d / 'summary.json'}")


def dry_run(a, P, d: Path):
    mc = ROOT / "data" / "yellow_365_audit" / "measured_cost.json"
    measured = json.loads(mc.read_text()) if mc.exists() else {}
    tot_h = 0.0
    for cs in a.cats:
        cells = collect_cells(cs, P, d)
        u = unique_cells(cells)
        defaults = P.get_defaults_for_symside(PROBE[cs])
        # representative eval count: start set = template defaults (no sym start set on the Mac) -> upper bound before ck-dedupe
        start = {}
        base, naked, cell = build_plan(start, defaults, u, P)
        n_evals = count_evals(base, naked, cell)
        bright = sum(1 for v in u.values() if v["colour"] == "bright")
        n_tok = sum(1 for v in u.values() if v["colour"] == "token")
        cost = (measured.get(cs) or {}).get("mean_cpu_sec_per_eval") or a.cost_sec or DEFAULT_COST[cs.split("_")[0]]
        per_sym_h = n_evals * cost / a.workers / 3600
        print(f"[{cs}] yellow cells in template {len(cells)} (unique switch=cand x filter=opt {len(u)}: bright {bright}, light {len(u) - bright - n_tok}, token-only {n_tok}) | unique evals/sym ~{n_evals} | cost {cost:.2f}s/eval cpu ({'measured' if (measured.get(cs) or {}).get('mean_cpu_sec_per_eval') else 'assumed'}) -> {per_sym_h:.2f} h/sym at {a.workers} workers -> {per_sym_h * a.syms:.1f} h for {a.syms} syms")
        tot_h += per_sym_h * a.syms
    print(f"[dry-run] total {tot_h:.1f} h for {len(a.cats)} cat_sides x {a.syms} syms on ONE box with {a.workers} workers")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cat-side", action="append", choices=CAT_SIDES, help="repeatable; default all (dry-run/summarize)")
    ap.add_argument("--syms", type=int, default=5)
    ap.add_argument("--sym-slice", help="lo:hi of the selected syms to run, e.g. 0:2 (pass 1) or 2:5 (pass 2)")
    ap.add_argument("--only-pos-cells", action="store_true", help="pass 2: only cells positive on the syms already computed")
    ap.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 4) - 2))
    ap.add_argument("--nice", type=int, default=0)
    ap.add_argument("--max-evals", type=int, default=0)
    ap.add_argument("--max-secs", type=int, default=0)
    ap.add_argument("--cost-sec", type=float, default=0.0)
    ap.add_argument("--min-n", type=int, default=3)
    ap.add_argument("--dir", default=None)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--summarize", action="store_true")
    ap.add_argument("--tier", choices=["yellow", "token", "all"], default="all", help="yellow = template bright+light cells only; token = name-token-only cells; all = both (needs --include-token-cells for token)")
    ap.add_argument("--include-token-cells", action="store_true", help="candidate set = template yellow (bright+light) PLUS name-token-related cells (>=1 shared token)")
    a = ap.parse_args()
    global INCLUDE_TOKEN
    INCLUDE_TOKEN = bool(a.include_token_cells)
    a.cats = a.cat_side or list(CAT_SIDES)
    d = out_dir(a)
    if a.summarize:
        return summarize(d, a.min_n)
    import v15_pilot as P
    if a.dry_run:
        return dry_run(a, P, d)
    if a.nice:
        try:
            os.nice(a.nice)
        except OSError:
            pass
    for cs in a.cats:
        run_cat(cs, a, P, d)
    summarize(d, a.min_n)


if __name__ == "__main__":
    main()
