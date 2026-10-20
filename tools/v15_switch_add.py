"""v15_switch_add — THE switch-adding tool (USER 2026-10-06 exception to the template one-script rule).

Adds ONE wired switch (2+ option rows) to the exact right tab of the right
TEMPLATE_FINAL_NORM file(s): whole-row insert after the last WHITE row
(before the orange block), style mirrored, default row first + bold B +
is_default YES. Refuses: unknown tab, <2 candidates, default not in
candidates, config default mismatch, unwired switches, duplicates.

Order of operations (full procedure: SWITCH_ADD_GUIDE.md):
  1. wire code first (config + QuickConfig + live fn + vec predicate + call site)
  2. registry entry + bible rebuild + verify_switch_bible green for the switch
  3. THIS TOOL --dry-run (default), review, then --apply (+ --fleet to sync)

Usage: python3 tools/v15_switch_add.py --switch NAME --tab TAB --candidates a,b
       --default a --venues both --sides both [--apply] [--fleet]
       python3 tools/v15_switch_add.py --enqueue --switch NAME --tab TAB --candidates a,b --default a [--push s1-int]
       python3 tools/v15_switch_add.py --queue data/template_row_queue.json [--templates-dirs D1,D2] [--apply]

PERSISTENT ROW INTAKE (2026-10-13, BIBLE §71 amendment / §94): rows written only on the Mac were overwritten by the
Mac-apply template pull (INV-0001 ENTRY_LH_TRIGGER_* rows lost 2026-10-09 13:49Z). The source host (S1) now owns
data/template_row_queue.json; S1 chain step 2a runs `--queue ... --apply`, which installs every queued switch that is
missing from a target file (same validation + insert + guards as the single mode, one load/save per file, flock shared
with the bookkeeper), so a queued row can never be lost again: it is re-installed on the next chain run if anything drops
it. `--enqueue` writes an inbox entry (data/template_row_queue_inbox/) instead of touching templates (use it on the Mac);
tools/v15_template_row_intake.py merges the inbox into the queue and rescues rows a pull would drop.
"""
import argparse
import ast
import copy
import csv
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import v15_template_restructure_v3 as R  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from template_row_guard import fingerprints, _cell_sig  # noqa: E402
import switch_types as ST  # noqa: E402  BIBLE §96 — the one typed-value validator

SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
FILES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
# 2026-10-13: s5/s6 were destroyed (tools/sync_indicators.sh pushes to 10.0.0.4 only). Leaving s5 here made
# every --fleet run end in "connect failed: Connection timed out / rsync error ... (code 20)" after the two
# real hosts had already succeeded, which reads like the push failed when it did not.
FLEET = ["s1-via-gateway", "s2"]



def _opt_key(v):
    """2026-10-17: compare option values numerically — a typed 50.0 candidate must match a '50' cell (extend refused it)."""
    t = str(v).strip()
    if isinstance(v, bool) or t.lower() in ("true", "false"):
        return t.lower()
    try:
        return repr(float(t))
    except ValueError:
        return t

def die(msg):
    sys.exit(f"REFUSED: {msg}")


def ast_assigns(path, cls=None):
    tree = ast.parse(Path(path).read_text())
    out = {}
    nodes = tree.body
    if cls:
        for n in tree.body:
            if isinstance(n, ast.ClassDef) and n.name == cls:
                nodes = n.body
                break
        else:
            return {}
    for n in nodes:
        if isinstance(n, ast.AnnAssign) and isinstance(n.target, ast.Name):
            try:
                out[n.target.id] = ast.literal_eval(n.value) if n.value else None
            except Exception:
                out[n.target.id] = None
        elif isinstance(n, ast.Assign):
            for t in n.targets:
                if isinstance(t, ast.Name):
                    try:
                        out[t.id] = ast.literal_eval(n.value)
                    except Exception:
                        out[t.id] = None
    return out


def parse_opt(raw, want):
    s = str(raw).strip()
    if want == bool:
        if s in ("True", "1", "1.0"):
            return True
        if s in ("False", "0", "0.0"):
            return False
        die(f"candidate {raw!r} is not a bool")
    if want == int:
        try:
            f = float(s)
        except Exception:
            die(f"candidate {raw!r} is not an int")
        if not f.is_integer():
            die(f"candidate {raw!r} is not integral (int field; rounding would invent a value)")
        return int(f)
    if want == float:
        try:
            return float(s)
        except Exception:
            die(f"candidate {raw!r} is not a float")
    return s


class Refused(Exception):
    pass


def _refuse(msg):
    raise Refused(msg)


def _type_of(dv):
    return bool if isinstance(dv, bool) else (int if isinstance(dv, int) and not isinstance(dv, bool) else (float if isinstance(dv, float) else str))


def validate_switch(sw, tab, cands, default, venues, sides, root=ROOT, bible=None):
    """Every refusal of the single mode, as a function (raises Refused). Returns (typed_ordered_opts, typed_default, sides, bible_entry)."""
    if tab not in SWITCH_SHEETS:
        _refuse(f"unknown tab {tab} (must be one of the 13 SWITCH_SHEETS)")
    cands = [str(c).strip() for c in cands if str(c).strip() != ""]
    if len(set(cands)) < 2:
        _refuse("need at least 2 distinct candidates")
    if str(default).strip() not in cands:
        _refuse("default must be one of the candidates")
    sides = list(sides)
    if "_LONG" in sw:
        sides = ["LONG"]
    if "_SHORT" in sw:
        sides = ["SHORT"]
    want = None
    for v in venues:
        cfg = ast_assigns(root / ("config.py" if v == "CRYPTO" else "config_tradier.py"), "Config" if v == "CRYPTO" else "TradierConfig")
        if sw not in cfg:
            _refuse(f"{sw} is not a field of {'config.py' if v == 'CRYPTO' else 'config_tradier.py'} — add the config field first")
        dv = cfg[sw]
        if dv is None or isinstance(dv, (tuple, list, dict)):
            _refuse(f"{sw} {v} config default {dv!r} is not a scalar — rows need a typed scalar encoding first")
        t = _type_of(dv)
        want = t if want is None else want
        try:
            ok = _parse_opt_r(default, t) == (dv if not isinstance(dv, float) else float(dv))
        except Refused:
            raise
        if not ok:
            _refuse(f"default {default!r} != {v} config default {dv!r} — align config first")
    qc = ast_assigns(root / "v12_quick_engine.py", "QuickConfig")
    if sw not in qc:
        _refuse(f"{sw} is not a QuickConfig field — wire v12_quick_engine first")
    if bible is None:
        try:
            bible = json.loads((root / "data" / "SWITCH_BIBLE.json").read_text())["switches"]
        except Exception as e:
            _refuse(f"cannot read data/SWITCH_BIBLE.json ({e}) — rebuild the bible first")
    e = bible.get(sw)
    if not e:
        _refuse(f"{sw} has no SWITCH_BIBLE entry — rebuild the bible first")
    if str(e.get("status") or "").upper().startswith("DEAD"):
        _refuse(f"{sw} bible status is dead ({e.get('status')})")
    if not (e.get("vec_reads") or e.get("vec_read_count")):
        _refuse(f"{sw} has no vec reads — wire v12_quick_engine/vec_decisions first")
    if not any((e.get("live_read_count") or {}).values()):
        _refuse(f"{sw} has no live reads — wire ez_manage/tradier_manage first")
    reg = (e.get("agent_c_registry") or {})
    if reg.get("suggested_home_tab") and reg["suggested_home_tab"] != tab:
        _refuse(f"registry suggested_home_tab is {reg['suggested_home_tab']}, not {tab}")
    typed = [_parse_opt_r(c, want) for c in cands]
    dflt = _parse_opt_r(default, want)
    for v in venues:  # BIBLE §96: the ONE validator has the last word on every value written to a template (fail-closed)
        for t in [dflt] + typed:
            ok_t, cv_t, why_t = ST.coerce(sw, t, v, root=root)
            if not ok_t or not ST.canonical(sw, t, v, root=root):
                _refuse(f"switch_types refused {sw}={t!r} for {v}: {why_t or 'not canonically typed'}")
    ordered = [dflt] + [t for t in typed if t != dflt]
    return ordered, dflt, sides, e


def _parse_opt_r(raw, want):
    try:
        return parse_opt(raw, want)
    except SystemExit as se:
        raise Refused(str(se).replace("REFUSED: ", ""))


def last_white_row(ws):
    last_w = 2
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value in (None, ""):
            continue
        if not (ws.cell(r, 1).fill.fill_type == "solid" and R.ORANGE in str(ws.cell(r, 1).fill.fgColor.rgb or "")):
            last_w = r
    return last_w


def demote_rows(ws, sw):
    """extend with a new bold default: the old bold row(s) of `sw` become plain candidates (bold off, is_default NO); never deleted."""
    hdrs = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    isd = hdrs.index("is_default") + 1
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value == sw:
            f0 = copy.copy(ws.cell(r, 2).font)
            ws.cell(r, 2).font = Font(name=f0.name, size=f0.size, bold=False, italic=f0.italic, color=f0.color)
            ws.cell(r, isd).value = "NO"


def insert_switch_rows(ws, sw, opts, dflt, ts):
    """Whole-row insert after the last WHITE row, style mirrored from that row, default first + bold + is_default YES.
    Raises Refused on a defaults-gate failure. Returns (first_row, n_rows)."""
    last_w = last_white_row(ws)
    ws.insert_rows(last_w + 1, len(opts))
    hdrs = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    isd = hdrs.index("is_default") + 1
    ref = last_w
    for i, o in enumerate(opts):
        rr = last_w + 1 + i
        for c in range(1, ws.max_column + 1):
            tgt, src = ws.cell(rr, c), ws.cell(ref, c)
            tgt.font, tgt.fill, tgt.alignment, tgt.border, tgt.number_format = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.alignment), copy.copy(src.border), src.number_format
        ws.cell(rr, 1).value = sw
        ws.cell(rr, 2).value = o
        ws.cell(rr, 2).font = Font(bold=(o == dflt))
        ws.cell(rr, isd).value = "YES" if o == dflt else "NO"
        ws.cell(rr, 1).comment = Comment(f"switch-add {ts}: wired switch staged by tools/v15_switch_add.py", "switch_add")
    grp = [(ws.cell(r, 2).font.bold, ws.cell(r, isd).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == sw]
    if sum(1 for b, y in grp if b and y == "YES") != 1 or sum(1 for b, y in grp if (b and y != "YES") or (y == "YES" and not b)):
        _refuse(f"{ws.title}: DEFAULTS-GATE failed for {sw} (need exactly one bold B == YES)")
    return last_w + 1, len(opts)


def atomic_save(wb, p):
    import zipfile
    tmp = str(p) + ".addtmp"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as z:
        if len(z.namelist()) < 10:
            os.remove(tmp)
            _refuse(f"refusing truncated save {Path(p).name}")
    os.replace(tmp, str(p))


def register_allowlist(sw, e, tab, root=ROOT):
    _al = root / "data" / "reports" / "ALL_PATHS_ALLOWLIST.csv"
    _rows = list(csv.DictReader(_al.open(newline="", encoding="utf-8")))
    if sw in {r.get("field", "").strip() for r in _rows}:
        return False
    _role = "ENTRY" if "ENTRY" in tab else ("EXIT" if "EXIT" in tab else ("REENTRY" if "REENTRY" in tab else ("REDUCE" if "REDUCE" in tab else ("AUGMENT" if "AUGMENT" in tab else "FILTER"))))
    _vec = str((e.get("vec_reads") or [""])[0])[:120] if isinstance(e.get("vec_reads"), list) else str(e.get("vec_reads", ""))[:120]
    _live = ""
    try:
        _lr = e.get("live_reads") or {}
        _live = str(next(iter(_lr.values()))[0])[:120] if isinstance(_lr, dict) and _lr else str(_lr)[:120]
    except Exception:
        _live = ""
    with _al.open("a", newline="", encoding="utf-8") as _fh:
        csv.writer(_fh).writerow([sw, _role, "", "HIGH_DIRECT_READ", f"bible:{e.get('status')};vec={_vec};live={_live};tab={tab}"[:300], "15m", "15m", _rows[0]["source_numbers_sha256"] if _rows else ""])
    return True


def switches_in_file(path):
    """{switch: tab} for every SWITCH_SHEETS row of one template (read-only)."""
    out = {}
    wb = openpyxl.load_workbook(str(path), read_only=True)
    try:
        for t in SWITCH_SHEETS:
            if t not in wb.sheetnames:
                continue
            for row in wb[t].iter_rows(min_row=3, max_col=1, values_only=True):
                if row and row[0] not in (None, ""):
                    out.setdefault(str(row[0]).strip(), t)
    finally:
        wb.close()
    return out


def _entry_venues(entry):
    v = entry.get("venues") or ["CRYPTO", "STOCKS"]
    return [x.upper() for x in v]


def _entry_cands(entry, venue):
    c = entry.get("candidates")
    return [str(x) for x in (c.get(venue) or []) ] if isinstance(c, dict) else [str(x) for x in (c or [])]


def _entry_default(entry, venue):
    d = entry.get("default")
    return str(d.get(venue)) if isinstance(d, dict) else str(d)


def _write_json_atomic(path, obj):
    tmp = str(path) + ".tmp"
    Path(tmp).write_text(json.dumps(obj, indent=1, default=str))
    os.replace(tmp, str(path))


def apply_queue(queue_path, templates_dirs, apply=False, root=ROOT, lock_wait_s=1800, ts=None):
    """Install every queued switch missing from a target template (idempotent). Returns the summary dict."""
    ts = ts or datetime.datetime.now().strftime("%Y%m%d%H%M")
    q = json.loads(Path(queue_path).read_text())
    try:
        bible = json.loads((root / "data" / "SWITCH_BIBLE.json").read_text())["switches"]
    except Exception:
        bible = None
    present_cache = {}

    def present(p):
        if p not in present_cache:
            present_cache[p] = switches_in_file(p)
        return present_cache[p]
    plan, summary = {}, {"ts": ts, "apply": bool(apply), "entries": {}, "files": {}}
    for entry in q.get("entries", []):
        sw = str(entry.get("switch") or "").strip().upper()
        if not sw or str(entry.get("status", "")).upper() in ("RETIRED", "WITHDRAWN"):
            continue
        res = {"refused": {}, "present": [], "planned": [], "installed": []}
        for venue in _entry_venues(entry):
            try:
                opts, dflt, sides, e = validate_switch(sw, entry["tab"], _entry_cands(entry, venue), _entry_default(entry, venue), [venue], [s.upper() for s in (entry.get("sides") or ["LONG", "SHORT"])], root=root, bible=bible)
            except Refused as r:
                res["refused"][venue] = str(r)
                continue
            for side in sides:
                for d in templates_dirs:
                    p = Path(d) / FILES[f"{venue}_{side}"]
                    if not p.exists():
                        res["refused"][f"{venue}_{side}@{d}"] = f"missing {p}"
                        continue
                    if sw in present(p):
                        res["present"].append(f"{Path(d).name}/{p.name}!{present(p)[sw]}")
                        continue
                    plan.setdefault(str(p), []).append((sw, entry["tab"], opts, dflt, e))
                    res["planned"].append(f"{Path(d).name}/{p.name}!{entry['tab']}")
        summary["entries"][sw] = res
    for p, items in plan.items():
        print(f"[queue] {p}: insert {len(items)} switches ({sum(len(i[2]) for i in items)} rows): {', '.join(i[0] for i in items)}", flush=True)
    if not apply:
        print(f"[queue] dry-run: {sum(len(v) for v in plan.values())} switch installs planned over {len(plan)} files — re-run with --apply", flush=True)
        return summary
    lk = None
    try:
        import fcntl
        import time as _t
        lk = open("/tmp/v15_bookkeeper.lock", "a")
        t0 = _t.time()
        while True:
            try:
                fcntl.flock(lk, fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError:
                if _t.time() - t0 > lock_wait_s:
                    _refuse("bookkeeper lock /tmp/v15_bookkeeper.lock held too long — another template writer is active")
                _t.sleep(5)
    except ImportError:
        lk = None
    try:
        for p, items in plan.items():
            p = Path(p)
            shutil.copy2(p, root / "backups" / f"before_switch_add_queue_{ts}_{p.parent.name}_{p.name}")
            wb = openpyxl.load_workbook(str(p))
            done = []
            try:
                before = {}
                for sw, tab, opts, dflt, e in items:
                    ws = wb[tab]
                    if tab not in before:
                        before[tab] = fingerprints(ws)
                    if any(ws.cell(r, 1).value == sw for r in range(3, ws.max_row + 1)):
                        continue
                    insert_switch_rows(ws, sw, opts, dflt, ts)
                    done.append((sw, tab, len(opts)))
                for tab, fp in before.items():
                    lost = sum((fp - fingerprints(wb[tab])).values())
                    if lost:
                        _refuse(f"{p.name}!{tab}: {lost} original rows changed — aborting (file untouched, backup backups/before_switch_add_queue_{ts}_{p.parent.name}_{p.name})")
                atomic_save(wb, p)
            finally:
                wb.close()
            summary["files"][str(p)] = [f"{sw}@{tab}+{n}" for sw, tab, n in done]
            for sw, tab, n in done:
                summary["entries"][sw]["installed"].append(f"{p.parent.name}/{p.name}!{tab}+{n}")
                print(f"[queue] installed {n} rows {sw} -> {p.parent.name}/{p.name}!{tab}", flush=True)
        for sw, tab, opts, dflt, e in {(i[0]): i for v in plan.values() for i in v}.values():
            try:
                if register_allowlist(sw, e, tab, root=root):
                    print(f"[allowlist] registered {sw} in ALL_PATHS_ALLOWLIST.csv (backtest override permission)", flush=True)
            except Exception as _ale:
                print(f"[allowlist] WARN {sw}: {_ale}", flush=True)
    finally:
        if lk is not None:
            lk.close()
    present_cache.clear()
    for entry in q.get("entries", []):
        sw = str(entry.get("switch") or "").strip().upper()
        res = summary["entries"].get(sw)
        if not res:
            continue
        missing = [x for x in res["planned"] if not any(x.split("!")[0] in i for i in res["installed"])]
        entry["status"] = "INSTALLED" if (not missing and (res["present"] or res["installed"])) else ("REFUSED" if res["refused"] and not (res["present"] or res["installed"]) else ("PARTIAL" if missing or res["refused"] else entry.get("status", "QUEUED")))
        entry["last_result"] = {"ts": ts, "installed": res["installed"], "present_n": len(res["present"]), "refused": res["refused"], "missing": missing}
    q["last_apply"] = {"ts": ts, "files": {k: len(v) for k, v in summary["files"].items()}}
    _write_json_atomic(queue_path, q)
    return summary


_MISS = object()


def _cat_side_defaults(cs, root=ROOT):
    try:
        import importlib.util as _ilu
        _sp = _ilu.spec_from_file_location(f"_csd_{abs(hash(str(root)))}", str(Path(root) / "cat_side_defaults.py"))
        _csd = _ilu.module_from_spec(_sp)
        sys.path.insert(0, str(root))
        _sp.loader.exec_module(_csd)
        return dict(_csd.defaults(cs) or {})
    except Exception:
        return {}


def _set_bold(cell, bold):
    f = cell.font
    cell.font = Font(name=(f.name if f else None) or "Arial", size=(f.size if f else None) or 10, bold=bold, italic=f.italic if f else None, color=f.color if f else None)


def plan_fix_types_ws(ws, venue, cs_defaults, root=ROOT):
    """BIBLE §96 type repair plan for ONE switch tab. Returns {'writes': {row: value}, 'drops': {row: reason}, 'bold_moves': [(old, new)],
    'log': [...], 'unresolved': [...]} — read-only. Rules: an unambiguous coercion is written in place (typed); an ambiguous
    candidate row is dropped (never guessed); a refused BOLD row is realigned to the cat_side default (= what live and the store
    run, BIBLE §90 builder placeholders) — bold moves onto the row already carrying that value, else the cell is rewritten to it;
    a rewritten row that now duplicates a kept row of the same switch (same typed value AND every other cell identical in value +
    fill, so no yellow/family information is lost) is dropped (bold / untouched rows are kept); a duplicate whose other cells
    differ is kept, typed."""
    hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    if not hdr or hdr[0] != "Switch" or "is_default" not in hdr:
        return {"writes": {}, "drops": {}, "bold_moves": [], "log": [], "unresolved": []}
    isd = hdr.index("is_default") + 1
    groups = {}
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and re.fullmatch(r"[A-Z0-9_]{3,}", a.strip()) and ws.cell(r, 2).value is not None:
            groups.setdefault(a.strip(), []).append(r)
    writes, drops, moves, log, unresolved = {}, {}, [], [], []
    for field, rows in groups.items():
        bold = [r for r in rows if str(ws.cell(r, isd).value or "").strip().upper() == "YES"]
        typed = {}
        for r in rows:
            raw = ws.cell(r, 2).value
            ok, cv, why = ST.to_cell(field, raw, venue, root=root)
            if ok:
                typed[r] = cv
                if not ST.canonical(field, raw, venue, root=root):
                    writes[r] = cv
                    log.append({"tab": ws.title, "row": r, "field": field, "action": "coerced", "from": repr(raw), "to": repr(cv), "bold": r in bold, "why": why})
            elif r not in bold:
                drops[r] = why
                log.append({"tab": ws.title, "row": r, "field": field, "action": "dropped", "from": repr(raw), "to": None, "bold": False, "why": why})
        for r in bold:
            if r in typed:
                continue
            raw = ws.cell(r, 2).value
            tgt = cs_defaults.get(field, _MISS)
            if tgt == "" and ST._tf_named(field) and ST.field_spec(field, venue, root)["kind"] == "str":
                ok_t, tv = True, "OFF"
            elif tgt is _MISS or tgt is None or tgt == "":
                ok_t, tv = False, None
            else:
                ok_t, tv, _ = ST.to_cell(field, tgt, venue, root=root)
            if not ok_t:
                unresolved.append({"tab": ws.title, "row": r, "field": field, "value": repr(raw), "why": f"bold value refused and no typed cat_side default ({tgt!r}) to realign to"})
                continue
            match = [x for x in rows if x != r and x in typed and x not in drops and ST.same_value(typed[x], tv)]
            if match:
                moves.append((r, match[0]))
                drops[r] = f"refused bold {raw!r} — bold moved to the cat_side default row {match[0]} ({tv!r})"
                bold = [match[0] if x == r else x for x in bold]
                log.append({"tab": ws.title, "row": r, "field": field, "action": "bold_moved_dropped", "from": repr(raw), "to": f"row {match[0]} {tv!r}", "bold": True, "why": drops[r]})
            else:
                typed[r] = tv
                writes[r] = tv
                log.append({"tab": ws.title, "row": r, "field": field, "action": "bold_realigned", "from": repr(raw), "to": repr(tv), "bold": True, "why": "refused bold value -> cat_side default (live/store value)"})
        kept, sigs = [], {}

        def _sig(x):
            if x not in sigs:
                sigs[x] = tuple(_cell_sig(ws.cell(x, c)) for c in range(1, ws.max_column + 1) if c not in (2, isd))
            return sigs[x]
        for r in rows:
            if r in drops or r not in typed:
                continue
            k = next((x for x in kept if ST.same_value(typed[x], typed[r]) and _sig(x) == _sig(r)), None)
            if k is None:
                kept.append(r)
                continue
            if r not in writes and k not in writes:
                kept.append(r)  # pre-existing canonical duplicate: not this writer's business
                continue
            loser = r
            if r in bold or (k in writes and r not in writes and k not in bold):
                loser = k
            if loser in bold:
                kept.append(r)
                continue
            drops[loser] = f"duplicate of row {r if loser == k else k} after typing ({typed[loser]!r})"
            writes.pop(loser, None)
            log.append({"tab": ws.title, "row": loser, "field": field, "action": "dropped_duplicate", "from": repr(ws.cell(loser, 2).value), "to": None, "bold": False, "why": drops[loser]})
            if loser == k:
                kept[kept.index(k)] = r
    return {"writes": writes, "drops": drops, "bold_moves": moves, "log": log, "unresolved": unresolved, "isd": isd}


def _default_violations(ws):
    hdr = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    if "is_default" not in hdr:
        return set()
    isd = hdr.index("is_default") + 1
    g = {}
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a not in (None, ""):
            g.setdefault(str(a).strip(), []).append(r)
    out = set()
    for a, rows in g.items():
        yes = [r for r in rows if str(ws.cell(r, isd).value or "").strip().upper() == "YES"]
        bold = [r for r in rows if ws.cell(r, 2).font is not None and ws.cell(r, 2).font.b]
        if len(yes) != 1 or yes != bold:
            out.add(a)
    return out


def fix_types(templates_dirs, apply=False, root=ROOT, lock_wait_s=1800, ts=None):
    """--fix-types: repair every mistyped switch value in the target templates through the ONE validator (tools/switch_types.py).
    Same guards as the queue mode: bookkeeper flock, backup per file, template_row_guard fingerprints (lost == dropped rows,
    gained == 0, columns B/is_default skipped), one-default rule not broken, post-save validator rescan, zip-validated atomic save."""
    ts = ts or datetime.datetime.now().strftime("%Y%m%d%H%M")
    paths = [Path(d) / fn for d in templates_dirs for fn in FILES.values() if (Path(d) / fn).exists()]
    summary = {"ts": ts, "apply": bool(apply), "files": {}}
    lk = None
    if apply:
        try:
            import fcntl
            import time as _t
            lk = open("/tmp/v15_bookkeeper.lock", "a")
            t0 = _t.time()
            while True:
                try:
                    fcntl.flock(lk, fcntl.LOCK_EX | fcntl.LOCK_NB)
                    break
                except OSError:
                    if _t.time() - t0 > lock_wait_s:
                        _refuse("bookkeeper lock /tmp/v15_bookkeeper.lock held too long — another template writer is active")
                    _t.sleep(5)
        except ImportError:
            lk = None
    try:
        for p in paths:
            cs = next(k for k, v in FILES.items() if v == p.name)
            venue = cs.split("_")[0]
            csd = _cat_side_defaults(cs, root)
            wb = openpyxl.load_workbook(str(p))
            try:
                plans, rep = {}, {"coerced": 0, "dropped": 0, "bold_realigned": 0, "bold_moved": 0, "unresolved": [], "log": []}
                for tab in SWITCH_SHEETS:
                    if tab not in wb.sheetnames:
                        continue
                    pl = plan_fix_types_ws(wb[tab], venue, csd, root)
                    plans[tab] = pl
                    rep["log"].extend(pl["log"])
                    rep["unresolved"].extend(pl["unresolved"])
                    rep["coerced"] += sum(1 for x in pl["log"] if x["action"] == "coerced" and x["row"] not in pl["drops"])
                    rep["dropped"] += len(pl["drops"])
                    rep["bold_realigned"] += sum(1 for x in pl["log"] if x["action"] == "bold_realigned")
                    rep["bold_moved"] += len(pl["bold_moves"])
                summary["files"][str(p)] = rep
                n_change = sum(len(pl["writes"]) + len(pl["drops"]) for pl in plans.values())
                print(f"[fix-types] {p.parent.name}/{p.name}: coerced={rep['coerced']} dropped={rep['dropped']} bold_realigned={rep['bold_realigned']} bold_moved={rep['bold_moved']} unresolved={len(rep['unresolved'])}", flush=True)
                if not apply or n_change == 0:
                    continue
                shutil.copy2(p, root / "backups" / f"before_fix_types_{ts}_{p.parent.name}_{p.name}")
                for tab, pl in plans.items():
                    if not (pl["writes"] or pl["drops"]):
                        continue
                    ws = wb[tab]
                    isd = pl["isd"]
                    skip = frozenset({2, isd})
                    before_fp = fingerprints(ws, skip)
                    before_viol = _default_violations(ws)
                    for r, v in pl["writes"].items():
                        ws.cell(r, 2).value = v
                    for old, new in pl["bold_moves"]:
                        ws.cell(new, isd).value = "YES"
                        _set_bold(ws.cell(new, 2), True)
                        _set_bold(ws.cell(new, isd), True)
                        ws.cell(old, isd).value = "NO"
                        _set_bold(ws.cell(old, 2), False)
                    for r in sorted(pl["drops"], reverse=True):
                        ws.delete_rows(r, 1)
                    after_fp = fingerprints(ws, skip)
                    lost, gained = sum((before_fp - after_fp).values()), sum((after_fp - before_fp).values())
                    if lost != len(pl["drops"]) or gained:
                        _refuse(f"{p.name}!{tab}: integrity broken (lost {lost} != {len(pl['drops'])} dropped, gained {gained}) — file untouched")
                    new_viol = _default_violations(ws) - before_viol
                    if new_viol:
                        _refuse(f"{p.name}!{tab}: one-default rule broken for {sorted(new_viol)[:5]} — file untouched")
                tmp = str(p) + ".fixtypes.xlsx"
                wb.save(tmp)
                import zipfile
                with zipfile.ZipFile(tmp) as z:
                    if len(z.namelist()) < 10:
                        os.remove(tmp)
                        _refuse(f"refusing truncated save {p.name}")
                left = [x for x in ST.scan_workbook(tmp, root=root) if not any(u["field"] == x["field"] and u["tab"] == x["tab"] for u in rep["unresolved"])]
                if left:
                    os.remove(tmp)
                    _refuse(f"{p.name}: {len(left)} mistyped value(s) remain after the repair (first {left[:3]}) — file untouched")
                os.replace(tmp, str(p))
                print(f"[fix-types] saved {p.parent.name}/{p.name} (backup backups/before_fix_types_{ts}_{p.parent.name}_{p.name})", flush=True)
            finally:
                wb.close()
    finally:
        if lk is not None:
            lk.close()
    return summary


def enqueue(entry, inbox_dir, push=None):
    Path(inbox_dir).mkdir(parents=True, exist_ok=True)
    ts = datetime.datetime.utcnow().strftime("%Y%m%d%H%M%S")
    f = Path(inbox_dir) / f"enqueue_{entry['switch']}_{ts}.json"
    _write_json_atomic(f, {"source": "v15_switch_add --enqueue", "ts": ts, "entries": [entry]})
    print(f"[enqueue] {entry['switch']} -> {f}", flush=True)
    if push:
        r = subprocess.run(["rsync", "-az", "-e", "ssh -o BatchMode=yes -o ConnectTimeout=10", str(f), f"{push}:binance-sandbox/data/template_row_queue_inbox/"], capture_output=True, text=True)
        print(f"[enqueue] push {push} rc={r.returncode} {r.stderr.strip()[:200]}", flush=True)
        return r.returncode
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--switch")
    ap.add_argument("--tab")
    ap.add_argument("--candidates")
    ap.add_argument("--default")
    ap.add_argument("--venues", default="both", choices=["crypto", "stocks", "both"])
    ap.add_argument("--sides", default="both", choices=["long", "short", "both"])
    ap.add_argument("--templates-dir", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--extend", action="store_true", help="the switch may already have rows: install ONLY the missing candidates and keep the existing bold default (USER 2026-10-16: every option must be testable, the best delta stays)")
    ap.add_argument("--fleet", action="store_true")
    ap.add_argument("--queue", help="batch mode: install every queued switch missing from the target templates")
    ap.add_argument("--templates-dirs", default=f"{ROOT / 'SPREADSHEETS' / 'TEMPLATE_FINAL_NORM'},{ROOT / 'SPREADSHEETS'}")
    ap.add_argument("--summary-out", help="batch mode: write the summary json here")
    ap.add_argument("--enqueue", action="store_true", help="write an inbox entry for the source host instead of touching templates")
    ap.add_argument("--inbox", default=str(ROOT / "data" / "template_row_queue_inbox"))
    ap.add_argument("--push", help="with --enqueue: rsync the inbox entry to HOST:binance-sandbox/data/template_row_queue_inbox/")
    ap.add_argument("--fix-types", action="store_true", help="BIBLE §96: repair mistyped switch values in --templates-dirs via tools/switch_types.py (dry-run unless --apply)")
    a = ap.parse_args()
    if a.fix_types:
        try:
            s = fix_types([d for d in a.templates_dirs.split(",") if d.strip()], apply=a.apply)
        except Refused as r:
            die(str(r))
        if a.summary_out:
            _write_json_atomic(a.summary_out, s)
        unres = [u for f in s["files"].values() for u in f["unresolved"]]
        for u in unres:
            print(f"[fix-types] UNRESOLVED {u}", flush=True)
        return 0
    if a.queue:
        try:
            s = apply_queue(a.queue, [d for d in a.templates_dirs.split(",") if d.strip()], apply=a.apply)
        except Refused as r:
            die(str(r))
        if a.summary_out:
            _write_json_atomic(a.summary_out, s)
        ref = {k: v["refused"] for k, v in s["entries"].items() if v["refused"]}
        print(f"[queue] summary: entries={len(s['entries'])} installed={sum(len(v['installed']) for v in s['entries'].values())} present={sum(len(v['present']) for v in s['entries'].values())} refused_venues={sum(len(v) for v in ref.values())}", flush=True)
        for k, v in sorted(ref.items()):
            print(f"[queue] REFUSED {k}: {v}", flush=True)
        return 0
    if not (a.switch and a.tab and a.candidates and a.default):
        die("--switch, --tab, --candidates and --default are required (or --queue)")
    sw = a.switch.strip().upper()
    if a.tab not in SWITCH_SHEETS:
        die(f"unknown tab {a.tab} (must be one of the 13 SWITCH_SHEETS)")
    cands = [c.strip() for c in a.candidates.split(",") if c.strip() != ""]
    venues = ["CRYPTO", "STOCKS"] if a.venues == "both" else [a.venues.upper()]
    sides = ["LONG", "SHORT"] if a.sides == "both" else [a.sides.upper()]
    try:
        ordered, dflt, sides, e = validate_switch(sw, a.tab, cands, a.default.strip(), venues, sides)
    except Refused as r:
        die(str(r))
    if a.enqueue:
        return enqueue({"switch": sw, "tab": a.tab, "venues": venues, "sides": sides, "candidates": cands, "default": a.default.strip(), "status": "QUEUED", "source": "v15_switch_add --enqueue", "queued_at": datetime.datetime.utcnow().isoformat() + "Z"}, a.inbox, a.push)
    files = [(f"{v}_{s}", FILES[f"{v}_{s}"]) for v in venues for s in sides]
    plan = []
    demote_cs = set()
    for cs, fn in files:
        p = Path(a.templates_dir) / fn
        if not p.exists():
            die(f"missing {p}")
        wb = openpyxl.load_workbook(str(p))
        try:
            ws = wb[a.tab]
            have = {_opt_key(ws.cell(r, 2).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == sw and ws.cell(r, 2).value is not None}
            if have and not a.extend:
                die(f"{sw} already has rows in {fn}!{a.tab} — pass --extend to install only the MISSING candidates")
            if have:
                # EXTEND (USER 2026-10-16: "ALL 4 need to be tested in TEMPLATE_*.xlsx the best delta stays"):
                # a switch installed with one option can only ever read as its own default, so the board never
                # measures the other legs. Only the missing candidates are inserted and the EXISTING bold default is
                # left untouched — the defaults gate still requires exactly one bold is_default=YES in the group.
                miss = [o for o in ordered if _opt_key(o) not in have]
                if _opt_key(dflt) not in have:
                    # 2026-10-20 USER: the live venue default is not among the rows (stale bold) -> INSERT it as the new bold row and demote the old bold to a plain candidate (never deleted).
                    demote_cs.add(cs)
                    plan.append((cs, fn, p, last_white_row(ws), miss, dflt))
                elif not miss:
                    print(f"[{cs}] {fn}!{a.tab}: all {len(ordered)} candidates already present — nothing to do", flush=True)
                else:
                    plan.append((cs, fn, p, last_white_row(ws), miss, None))
            else:
                plan.append((cs, fn, p, last_white_row(ws), ordered, dflt))
        finally:
            wb.close()
    for cs, fn, p, last_w, opts, pdflt in plan:
        print(f"[{cs}] {fn}!{a.tab}: insert {len(opts)} rows {opts} at {last_w + 1} (after last white row) default={pdflt!r}" + ("" if pdflt else " (extend: existing bold default kept)") + (" [extend: new bold default, old bold demoted]" if cs in demote_cs else ""), flush=True)
    if not plan:
        print("nothing to install — every target already carries all candidates", flush=True)
        return 0
    if not a.apply:
        print("dry-run OK — re-run with --apply to write (backs up each file first)", flush=True)
        return 0
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for cs, fn, p, last_w, opts, pdflt in plan:
        shutil.copy2(p, ROOT / "backups" / f"before_switch_add_{sw}_{ts}_{fn}")
        wb = openpyxl.load_workbook(str(p))
        try:
            ws = wb[a.tab]
            if cs in demote_cs:
                demote_rows(ws, sw)
            before = fingerprints(ws)
            try:
                insert_switch_rows(ws, sw, opts, pdflt, ts)
            except Refused as r:
                die(str(r))
            after = fingerprints(ws)
            lost = sum((before - after).values())
            if lost:
                die(f"{fn}!{a.tab}: {lost} original rows changed — aborting, restore from backups/before_switch_add_{sw}_{ts}_{fn}")
            try:
                atomic_save(wb, p)
            except Refused as r:
                die(str(r))
            print(f"[{cs}] installed {len(opts)} rows -> {fn}!{a.tab}", flush=True)
        finally:
            wb.close()
    try:
        if register_allowlist(sw, e, a.tab):
            print(f"[allowlist] registered {sw} in ALL_PATHS_ALLOWLIST.csv (backtest override permission)", flush=True)
        else:
            print(f"[allowlist] {sw} already listed", flush=True)
    except Exception as _ale:
        die(f"allowlist registration failed for {sw}: {_ale}")
    if a.fleet:
        for h in FLEET:
            for _, fn in files:
                subprocess.run(["rsync", "-az", "-e", "ssh -S none -o StrictHostKeyChecking=accept-new", str(Path(a.templates_dir) / fn), f"{h}:~/binance-sandbox/SPREADSHEETS/TEMPLATE_FINAL_NORM/{fn}"], check=False)
            r = subprocess.run(["ssh", "-S", "none", "-o", "StrictHostKeyChecking=accept-new", h, "cd ~/binance-sandbox/SPREADSHEETS/TEMPLATE_FINAL_NORM && md5sum " + " ".join(fn for _, fn in files)], capture_output=True, text=True)
            print(f"[{h}] {r.stdout.strip()}", flush=True)
        print("NOTE: a Mac-side --apply is overwritten by the next Mac-apply template pull unless the switch is also queued on S1 (--enqueue --push s1-int); the Mac-apply rescue step queues dropped rows automatically", flush=True)
    else:
        print("fleet NOT synced — re-run with --fleet (or rsync + md5 by hand); on the Mac also --enqueue --push s1-int so the S1 source host keeps the rows", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
