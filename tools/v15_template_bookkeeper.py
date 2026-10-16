"""THE ONE template-bookkeeping script (USER 2026-10-06).

NO other agent or script may touch TEMPLATE_* files. This script runs once
before market open (S1 cron 12:30 UTC), pulls zero-formula ledgers from the
fleet, and writes the ZERO_FORMULA_BOOK sheet into each of the 4 shared
templates (atomic save) + a standalone BOOK copy per cat_side.

Usage: python3 tools/v15_template_bookkeeper.py [--dry-run] [--templates DIR]
       [--ledger-root DIR] [--hosts h1,h2] [--timeout-s N]
"""
import argparse
import glob
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
SHEET = "ZERO_FORMULA_BOOK"
PROTOCOL_SHEET = "PROTOCOL"
PROTOCOL_FILE = ROOT / "data" / "TEMPLATE_PROTOCOL.md"


def write_protocol_sheet(wb):
    """USER 2026-10-14: the template protocol lives INSIDE each of the 4 templates. Replaced whole from data/TEMPLATE_PROTOCOL.md
    (one line per row, first tab) on every bookkeeper run, so it can never drift or be lost by a pull/push."""
    try:
        lines = PROTOCOL_FILE.read_text().splitlines()
    except Exception as e:
        print(f"[bookkeeper] {PROTOCOL_FILE} unreadable ({e}) — PROTOCOL tab untouched", flush=True)
        return 0
    if PROTOCOL_SHEET in wb.sheetnames:
        wb.remove(wb[PROTOCOL_SHEET])
    ws = wb.create_sheet(PROTOCOL_SHEET, 0)
    ws.column_dimensions["A"].width = 180
    for i, ln in enumerate(lines, start=1):
        ws.cell(row=i, column=1).value = ln
    return len(lines)


def merge_ledgers(ledgers):
    rows, hdrs, syms = {}, {}, set()
    for led in ledgers:
        if not isinstance(led, dict):
            continue
        syms.add(led.get("symside") or "?")
        for k, e in (led.get("rows") or {}).items():
            if not isinstance(e, dict):
                continue
            r = rows.setdefault(k, dict(e))
            r.setdefault("syms_seen", [])
            if led.get("symside") and led.get("symside") not in r["syms_seen"]:
                r["syms_seen"].append(led.get("symside"))
        for h, rks in (led.get("cells_by_header") or {}).items():
            hdrs.setdefault(h, set()).update(rks or [])
    return {"rows": rows, "headers": {h: len(s) for h, s in hdrs.items()}, "syms": sorted(s for s in syms if s != "?"), "cell_rows": sum(len(v.get("cell_rows") or {}) for v in ledgers if isinstance(v, dict))}


def write_book_sheet(wb, cat, merged, symside_ct):
    if SHEET in wb.sheetnames:
        wb.remove(wb[SHEET])
    ws = wb.create_sheet(SHEET)
    ws.append(["ZERO-FORMULA BOOK", cat, f"{symside_ct} sym_sides booked", "SKIPPED_BAD_FORMULA = skipped, booked, NOT OBSOLETE — bad formula, fix then evidence regen lifts the skip"])
    ws.append(["KIND", "TAB", "SWITCH", "CAND", "HEADER_OR_COUNT", "POS", "N", "AVG", "STATUS", "NOTE"])
    for rk in sorted((merged.get("rows") or {})):
        e = merged["rows"][rk]
        ws.cell(row=ws.max_row + 1, column=1).value = "row"
        for c, v in enumerate([e.get("tab"), e.get("switch"), e.get("cand"), "", e.get("pos_sym"), e.get("n_sym"), e.get("avg_delta"), "SKIPPED_BAD_FORMULA", "NOT OBSOLETE — bad formula"], start=2):
            ws.cell(row=ws.max_row, column=c).value = v
    top = sorted((merged.get("headers") or {}).items(), key=lambda kv: -kv[1])[:500]
    for h, n in top:
        ws.cell(row=ws.max_row + 1, column=1).value = "header_top"
        for c, v in enumerate(["", "", "", h, 0, n, "", "SKIPPED_BAD_FORMULA", f"skipped in {n} rows"], start=2):
            ws.cell(row=ws.max_row, column=c).value = v
    return ws.max_row


def atomic_save_xlsx(wb, path, min_entries=10):
    import zipfile
    tmp = str(path) + ".booktmp"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as z:
        if len(z.namelist()) < min_entries:
            raise RuntimeError(f"refusing truncated save {path}: {len(z.namelist())} zip entries")
    os.replace(tmp, str(path))


def _norm_val(v):
    s = str(v).strip()
    return {"true": "True", "false": "False", "1.0": "1", "0.0": "0"}.get(s.lower(), s)


def align_default(wb, switch, value):
    """2026-10-16: one switch whose bold default DISAGREES across tabs (the pilot's DEFAULTS-GATE refuses the whole template —
    every STOCKS_SHORT board was refused from 2026-10-15 23:22: WT_CROSSUNDER_FINAL_ENABLED bold True in
    ENTRY_CONFIRMATION_GATES, bold False in EXIT_VELOCITY). Re-bolds the EXISTING row whose value == `value` in every tab that
    carries the switch; never changes a cell value, never adds/removes/reorders rows. A tab without a row of that value is left
    untouched and reported. is_default (row-2 header) set YES/NO alongside. Returns a report; raises if a touched tab would not
    end with exactly one bold row for the switch."""
    from openpyxl.styles import Font
    import copy as _copy
    want = _norm_val(value)
    rep = {"switch": switch, "value": want, "tabs": {}, "missing": []}
    for ws in wb.worksheets:
        rows = [r for r in range(3, ws.max_row + 1) if str(ws.cell(r, 1).value or "").strip() == switch]
        if not rows:
            continue
        hit = [r for r in rows if _norm_val(ws.cell(r, 2).value) == want]
        if not hit:
            rep["missing"].append(ws.title)
            continue
        isd = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(2, c).value or "").strip().lower().startswith("is_default")), None)
        before = [(r, bool(ws.cell(r, 2).font and ws.cell(r, 2).font.b)) for r in rows]
        for r in rows:
            cell = ws.cell(r, 2)
            f = _copy.copy(cell.font) if cell.font is not None else Font(name="Arial", size=10)
            on = r == hit[0]
            cell.font = Font(name=f.name or "Arial", size=f.size or 10, bold=on, italic=f.italic, color=f.color)
            if isd:
                ws.cell(r, isd).value = "YES" if on else "NO"
        nb = sum(1 for r in rows if ws.cell(r, 2).font and ws.cell(r, 2).font.b)
        if nb != 1:
            raise RuntimeError(f"{ws.title}: {switch} ends with {nb} bold rows")
        rep["tabs"][ws.title] = {"rows": rows, "bold_before": [r for r, b in before if b], "bold_after": hit[0]}
    return rep


SWITCH_TABS = ("STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES")


def _is_yellow(cell):
    try:
        return cell.fill is not None and cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb or "").upper().endswith("FFFF00")
    except Exception:
        return False


def dedupe_rows(wb):
    """USER 2026-10-16 "duplicate contradictory fields in TEMPLATE_ ... should be impossible". Removes rows that repeat the same
    (switch, value) inside one switch tab — same row kind only (column-A fill), so a white row and its orange filter-block
    twin are never merged. Keeps the bold row, else the row with the most yellow filter cells; yellow cells that only the
    removed twin had are copied onto the kept row first, so no filter test is lost. Pilot progress is keyed by canonical
    (tab, switch, value), so no measured cell is lost either. Raises if any switch ends with a different bold count."""
    import copy as _copy
    import dataclasses as _dc
    rep = {}
    try:
        sys.path.insert(0, str(ROOT))
        import v12_quick_engine as _E
        _bool_fields = {f.name for f in _dc.fields(_E.QuickConfig) if isinstance(getattr(_E.QuickConfig(), f.name, None), bool)}
    except Exception:
        _bool_fields = set()
    for ws in wb.worksheets:
        if ws.title not in SWITCH_TABS:
            continue
        # A bool switch whose whole group is written 0/1 (generic float era) -> False/True, so the bold reads as the
        # engine's bool type (chain_harmony_check TPL_TYPE). Only when EVERY row of the group is 0 or 1.
        _bg = {}
        for r in range(3, ws.max_row + 1):
            sw = str(ws.cell(r, 1).value or "").strip()
            if sw in _bool_fields:
                _bg.setdefault(sw, []).append(r)
        for sw, rows in _bg.items():
            vals = [ws.cell(r, 2).value for r in rows]
            if all(not isinstance(v, bool) and isinstance(v, (int, float)) and v in (0, 1) for v in vals):
                for r in rows:
                    ws.cell(r, 2).value = bool(ws.cell(r, 2).value)
                rep.setdefault(ws.title, []).append({"switch": sw, "bool_normalized": rows})
        groups = {}
        for r in range(3, ws.max_row + 1):
            sw = str(ws.cell(r, 1).value or "").strip()
            if not sw or sw != sw.upper() or not sw.replace("_", "").isalnum():
                continue
            a = ws.cell(r, 1)
            kind = str(a.fill.fgColor.rgb) if (a.fill is not None and a.fill.fill_type == "solid") else "none"
            groups.setdefault((sw, _norm_val(ws.cell(r, 2).value), kind), []).append(r)
        bold_before = {}
        for (sw, _v, _k), rows in groups.items():
            bold_before[sw] = bold_before.get(sw, 0) + sum(1 for r in rows if ws.cell(r, 2).font and ws.cell(r, 2).font.b)
        drop = []
        for (sw, v, _k), rows in groups.items():
            if len(rows) < 2:
                continue
            bold = [r for r in rows if ws.cell(r, 2).font and ws.cell(r, 2).font.b]
            yel = {r: sum(1 for c in range(4, ws.max_column + 1) if _is_yellow(ws.cell(r, c))) for r in rows}
            keep = bold[0] if bold else max(rows, key=lambda r: (yel[r], -r))
            # A bool switch whose rows all carry the SAME value is a missing option, not a stray copy (e.g. two
            # LH_HL_FILTER_REQUIRE_BOTH=False rows, no True row): the first duplicate becomes the opposite value so
            # both options are testable (USER 2026-10-16 "every option must be testable"). Only when the tab has no
            # row of the opposite value at all.
            flip = None
            if v in ("True", "False"):
                opp = "False" if v == "True" else "True"
                if not any(_norm_val(ws.cell(rr, 2).value) == opp for rr in range(3, ws.max_row + 1) if str(ws.cell(rr, 1).value or "").strip() == sw):
                    flip = next((r for r in rows if r != keep), None)
            if flip is not None:
                from openpyxl.styles import Font
                c2 = ws.cell(flip, 2)
                f = _copy.copy(c2.font) if c2.font is not None else Font(name="Arial", size=10)
                c2.value = (opp == "True")
                c2.font = Font(name=f.name or "Arial", size=f.size or 10, bold=False, italic=f.italic, color=f.color)
                isd = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(2, c).value or "").strip().lower().startswith("is_default")), None)
                if isd:
                    ws.cell(flip, isd).value = "NO"
                rep.setdefault(ws.title, []).append({"switch": sw, "value": v, "flipped_to": opp, "row": flip})
            for r in rows:
                if r == keep or r == flip:
                    continue
                for c in range(4, ws.max_column + 1):
                    if _is_yellow(ws.cell(r, c)) and not _is_yellow(ws.cell(keep, c)):
                        ws.cell(keep, c).fill = _copy.copy(ws.cell(r, c).fill)
                drop.append(r)
            if len(bold) > 1:
                bold_before[sw] -= len(bold) - 1
            rep.setdefault(ws.title, []).append({"switch": sw, "value": v, "kept": keep, "removed": [r for r in rows if r not in (keep, flip)]})
        for r in sorted(drop, reverse=True):
            ws.delete_rows(r, 1)
        bold_after = {}
        for r in range(3, ws.max_row + 1):
            sw = str(ws.cell(r, 1).value or "").strip()
            if sw in bold_before and ws.cell(r, 2).font and ws.cell(r, 2).font.b:
                bold_after[sw] = bold_after.get(sw, 0) + 1
        bad = {sw: (n, bold_after.get(sw, 0)) for sw, n in bold_before.items() if n != bold_after.get(sw, 0)}
        if bad:
            raise RuntimeError(f"{ws.title}: bold count changed {list(bad.items())[:5]}")
    return rep


def single_home(wb, switch, home):
    """USER 2026-10-16 "duplicate contradictory fields ... should be impossible": a switch with regular (white) rows in two
    tabs has two places to be promoted and can end with two bold defaults (WT_CROSSUNDER_FINAL_ENABLED refused every
    STOCKS_SHORT board 2026-10-15 23:22). Removes the switch's WHITE rows from every tab except `home`; filter-block
    (filled) rows are untouched. Refuses — writing nothing — if another tab carries a value the home lacks, a yellow
    filter cell the home row of that value lacks, or a bold value that differs from the home bold."""
    import copy as _copy
    def _white(ws, r):
        a = ws.cell(r, 1)
        return not (a.fill is not None and a.fill.fill_type == "solid")
    if home not in wb.sheetnames:
        raise RuntimeError(f"{home} is not a tab")
    hw = wb[home]
    hrows = {_norm_val(hw.cell(r, 2).value): r for r in range(3, hw.max_row + 1) if str(hw.cell(r, 1).value or "").strip() == switch and _white(hw, r)}
    if not hrows:
        raise RuntimeError(f"{home} has no white {switch} row")
    hbold = [v for v, r in hrows.items() if hw.cell(r, 2).font and hw.cell(r, 2).font.b]
    rep = {"switch": switch, "home": home, "removed": {}}
    plan = []
    for ws in wb.worksheets:
        if ws.title == home or ws.title not in SWITCH_TABS:
            continue
        rows = [r for r in range(3, ws.max_row + 1) if str(ws.cell(r, 1).value or "").strip() == switch and _white(ws, r)]
        for r in rows:
            v = _norm_val(ws.cell(r, 2).value)
            if v not in hrows:
                raise RuntimeError(f"{ws.title}!{r} {switch}={v} has no row in {home} — add it there first (tools/v15_switch_add.py --extend)")
            if ws.cell(r, 2).font and ws.cell(r, 2).font.b and hbold and v not in hbold:
                raise RuntimeError(f"{ws.title}!{r} bold {switch}={v} disagrees with {home} bold {hbold}")
            for c in range(4, ws.max_column + 1):
                if _is_yellow(ws.cell(r, c)) and not _is_yellow(hw.cell(hrows[v], c)):
                    hdr = str(ws.cell(2, c).value or ws.cell(1, c).value or c)
                    hc = next((cc for cc in range(4, hw.max_column + 1) if str(hw.cell(2, cc).value or hw.cell(1, cc).value or cc) == hdr), None)
                    if hc is None:
                        raise RuntimeError(f"{ws.title}!{r} yellow filter '{hdr}' has no column in {home}")
                    plan.append(("fill", hrows[v], hc, _copy.copy(ws.cell(r, c).fill)))
            plan.append(("drop", ws.title, r))
    for p in plan:
        if p[0] == "fill":
            hw.cell(p[1], p[2]).fill = p[3]
    for title in {p[1] for p in plan if p[0] == "drop"}:
        ws = wb[title]
        drops = sorted([p[2] for p in plan if p[0] == "drop" and p[1] == title], reverse=True)
        for r in drops:
            ws.delete_rows(r, 1)
        rep["removed"][title] = drops
    return rep


def pull_fleet(ledger_root, hosts, timeout_s):
    pulled = 0
    for h in hosts:
        h = h.strip()
        if not h:
            continue
        dest = Path(ledger_root) / "_in" / h.replace(".", "_")
        dest.mkdir(parents=True, exist_ok=True)
        try:
            r = subprocess.run(["rsync", "-az", "-e", "ssh -o ConnectTimeout=10 -o StrictHostKeyChecking=accept-new", f"niels@{h}:binance-sandbox/data/zero_formula_book/", str(dest) + "/"], capture_output=True, timeout=timeout_s)
            if r.returncode == 0:
                pulled += 1
            else:
                print(f"[bookkeeper] pull {h} failed rc={r.returncode} — booking local ledgers only", flush=True)
        except Exception as e:
            print(f"[bookkeeper] pull {h} error {e} — booking local ledgers only", flush=True)
    return pulled


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--templates", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--ledger-root", default=str(ROOT / "data" / "zero_formula_book"))
    ap.add_argument("--hosts", default=os.environ.get("V15_BOOK_HOSTS", ""))
    ap.add_argument("--timeout-s", type=int, default=300)
    ap.add_argument("--protocol-only", action="store_true", help="only (re)write the PROTOCOL tab (backs up each template first)")
    ap.add_argument("--align-default", metavar="SWITCH=VALUE", help="re-bold the existing SWITCH=VALUE row in every tab of --cat-side (cross-tab default disagreement); never edits values")
    ap.add_argument("--cat-side", help="with --align-default: CRYPTO_LONG|CRYPTO_SHORT|STOCKS_LONG|STOCKS_SHORT")
    ap.add_argument("--single-home", metavar="SWITCH=TAB", help="keep SWITCH's white rows only in TAB of every template that has them there (filter-block rows untouched)")
    ap.add_argument("--dedupe-rows", action="store_true", help="remove repeated (switch, value) rows inside each switch tab of every template (yellow cells merged, bold kept)")
    a = ap.parse_args()
    try:
        import fcntl
        _lk = open("/tmp/v15_bookkeeper.lock", "w")
        fcntl.flock(_lk, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except Exception:
        print("[bookkeeper] lock held — another run active, exiting", flush=True)
        return 0
    if a.align_default:
        import datetime as _dt, shutil as _sh, openpyxl as _ox
        sw, _, val = a.align_default.partition("=")
        if a.cat_side not in TEMPLATES or not sw or not val:
            print("[bookkeeper] --align-default needs SWITCH=VALUE and --cat-side", flush=True)
            return 2
        tpath = os.path.join(a.templates, TEMPLATES[a.cat_side])
        wb = _ox.load_workbook(tpath)
        try:
            rep = align_default(wb, sw.strip(), val.strip())
            print(json.dumps(rep), flush=True)
            if not rep["tabs"]:
                print("[bookkeeper] align-default: no tab carries that value — nothing written", flush=True)
                return 2
            if not a.dry_run:
                _sh.copy2(tpath, ROOT / "backups" / f"before_bookkeeper_align_{sw.strip()}_{_dt.datetime.utcnow().strftime('%Y%m%d%H%M')}_{TEMPLATES[a.cat_side]}")
                atomic_save_xlsx(wb, tpath)
                print(f"[bookkeeper] align-default written -> {tpath}", flush=True)
        finally:
            wb.close()
        return 0
    if a.single_home:
        import datetime as _dt, shutil as _sh, openpyxl as _ox
        sw, _, home = a.single_home.partition("=")
        for cat, fname in TEMPLATES.items():
            if a.cat_side and cat != a.cat_side:
                continue
            tpath = os.path.join(a.templates, fname)
            wb = _ox.load_workbook(tpath)
            try:
                if home.strip() not in wb.sheetnames or not any(str(wb[home.strip()].cell(r, 1).value or "").strip() == sw.strip() for r in range(3, wb[home.strip()].max_row + 1)):
                    continue
                rep = single_home(wb, sw.strip(), home.strip())
                n = sum(len(v) for v in rep["removed"].values())
                print(f"[bookkeeper] single-home {cat} {sw}->{home}: removed {n} rows {rep['removed']}", flush=True)
                if n and not a.dry_run:
                    _sh.copy2(tpath, ROOT / "backups" / f"before_bookkeeper_singlehome_{sw.strip()}_{_dt.datetime.utcnow().strftime('%Y%m%d%H%M')}_{fname}")
                    atomic_save_xlsx(wb, tpath)
            except RuntimeError as e:
                print(f"[bookkeeper] single-home {cat} {sw}: REFUSED {e}", flush=True)
            finally:
                wb.close()
        return 0
    if a.dedupe_rows:
        import datetime as _dt, shutil as _sh, openpyxl as _ox
        for cat, fname in TEMPLATES.items():
            tpath = os.path.join(a.templates, fname)
            wb = _ox.load_workbook(tpath)
            try:
                rep = dedupe_rows(wb)
                n = sum(len(g.get("removed") or []) for tab in rep.values() for g in tab)
                nf = sum(1 for tab in rep.values() for g in tab if g.get("flipped_to") or g.get("bool_normalized"))
                print(f"[bookkeeper] dedupe-rows {cat}: {n} duplicate rows removed, {nf} bool rows turned into the missing opposite option, {len(rep)} tabs", flush=True)
                if (n or nf) and not a.dry_run:
                    _sh.copy2(tpath, ROOT / "backups" / f"before_bookkeeper_dedupe_{_dt.datetime.utcnow().strftime('%Y%m%d%H%M')}_{fname}")
                    atomic_save_xlsx(wb, tpath)
                    print(f"[bookkeeper] dedupe-rows written -> {tpath}", flush=True)
            finally:
                wb.close()
        return 0
    hosts = [h for h in (a.hosts or "").split(",") if h.strip()]
    if hosts and not a.dry_run:
        pull_fleet(a.ledger_root, hosts, a.timeout_s)
    import openpyxl
    for cat, fname in TEMPLATES.items():
        pats = [os.path.join(a.ledger_root, cat, "*.json")]
        if not a.dry_run:
            pats.append(os.path.join(a.ledger_root, "_in", "*", cat, "*.json"))
        leds = []
        for p in pats:
            for f in glob.glob(p):
                try:
                    leds.append(json.load(open(f)))
                except Exception:
                    continue
        merged = merge_ledgers(leds)
        nrows = len(merged["rows"])
        print(f"[bookkeeper] {cat}: {len(leds)} ledgers, {nrows} row-skips, {len(merged['headers'])} condemned headers", flush=True)
        if a.dry_run:
            continue
        tpath = os.path.join(a.templates, fname)
        if a.protocol_only:
            if os.path.exists(tpath):
                import datetime as _dt, shutil as _sh
                _sh.copy2(tpath, ROOT / "backups" / f"before_bookkeeper_protocol_{_dt.datetime.utcnow().strftime('%Y%m%d%H%M')}_{fname}")
                wb = openpyxl.load_workbook(tpath)
                try:
                    np_ = write_protocol_sheet(wb)
                    atomic_save_xlsx(wb, tpath)
                finally:
                    wb.close()
                print(f"[bookkeeper] {cat}: protocol-only — {np_} lines -> {fname}!{PROTOCOL_SHEET}", flush=True)
            continue
        if not os.path.exists(tpath):
            print(f"[bookkeeper] missing {tpath} — skipped", flush=True)
            continue
        wb = openpyxl.load_workbook(tpath)
        try:
            nlines = write_book_sheet(wb, cat, merged, len(leds))
            write_protocol_sheet(wb)
            atomic_save_xlsx(wb, tpath)
        finally:
            wb.close()
        bpath = os.path.join(a.ledger_root, f"BOOK_{cat}.xlsx")
        wb2 = openpyxl.Workbook()
        try:
            write_book_sheet(wb2, cat, merged, len(leds))
            if "Sheet" in wb2.sheetnames:
                wb2.remove(wb2["Sheet"])
            atomic_save_xlsx(wb2, bpath, min_entries=4)
        finally:
            wb2.close()
        print(f"[bookkeeper] {cat}: wrote {nlines} book lines -> {fname} + BOOK_{cat}.xlsx", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
