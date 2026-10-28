#!/usr/bin/env python3
"""Guard the one invariant that cost 18 hours of zero published boards on 2026-10-27.

INVARIANT: every switch in v15_pilot.py's ESSENTIAL_GATE_SWITCHES must be TESTABLE in every
TEMPLATE_FINAL_NORM workbook -- no grey column-A rows, and at least 2 candidate rows.

Why both halves matter:
  * GREY: v15_pilot.py:5082 never sweeps a grey-font row (GREY_SKIP_RGB; grey = template-disabled,
    and the operator confirmed 2026-10-27 that grey rows are correctly never evaluated). A grey row
    therefore produces no done-record at all.
  * <2 CANDIDATES: with a single candidate the row IS the running configuration, so its delta is 0
    by construction (RUNNING_IDENTITY) and can never be a "non-zero number".
Either way the ESSENTIAL-ROWS gate (v15_pilot.py:1596) refuses the sheet: "N protection filters
without a non-zero number". Both halves were broken on gate switches without a LOCKED_FILES.md entry:
  10-25 12:55 787085243  WT_15M_BOUNCE_FILTER_MODE + WT_15M_BOUNCE_VOLUME_MODE: 6 rows -> 2
                         (one candidate each: AND, relvol)
  10-27 02:03 c1313f5a1  those survivors GREYED
  10-27 02:21 c7d1d8e98  the grey-skip added to the uchg-locked pilot
From then on EVERY finished board was refused (s5 20/25, s6 34/39, s2 2/2), which is the operator's
"no result published since 02:14". Nothing watched this surface, so it went unnoticed for 18 hours
while four servers billed by the minute.

A switch the operator has deliberately greyed must not be a gate member at all -- that is the
corollary. OPERATOR 2026-10-27: "DAYTRADE_DC_TARGET_TF should be greyed out we do not daytrade with
fixed % targets. Same for TECHNICAL_DC_TARGET_TF. The others do need to be tested". So those two are
expected to stay 100% grey and belong OUT of ESSENTIAL_GATE_SWITCHES; until that pilot edit lands
this guard reports them, which is the correct signal. List them in --expect-ungated to acknowledge
them without muting the rest of the surface.

Read-only. Exits 1 on breach so a cron line alerts.
"""
import argparse
import json
import pathlib
import re
import sys
import time

GREY_SKIP_RGB = {"FFBFBFBF", "00BFBFBF"}
ROOT = pathlib.Path(__file__).resolve().parent.parent
SENTINELS = ("switch", "general", "blanket", "filter", "option value")


def gate_switches(pilot: pathlib.Path) -> set:
    src = pilot.read_text()
    start = src.index("ESSENTIAL_GATE_SWITCHES = frozenset({")
    depth, end = 0, start
    for i in range(start, len(src)):
        if src[i] == "{":
            depth += 1
        elif src[i] == "}":
            depth -= 1
            if depth == 0:
                end = i
                break
    names = set(re.findall(r'"([A-Z][A-Z0-9_]{3,})"', src[start:end + 1]))
    pre = src[max(0, start - 4000):start]
    if "WT CORE" in pre:
        names |= set(re.findall(r'"([A-Z][A-Z0-9_]{3,})"', pre[pre.rfind("WT CORE"):]))
    return names


def scan(template: pathlib.Path, names: set) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(template))
    rows, grey = {}, {}
    try:
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for rr in range(3, ws.max_row + 1):
                cell = ws.cell(row=rr, column=1)
                val = cell.value
                if val is None:
                    continue
                sw = str(val).strip()
                if sw.lower() in SENTINELS or sw not in names:
                    continue
                rows[sw] = rows.get(sw, 0) + 1
                font = cell.font
                rgb = str(getattr(font.color, "rgb", "") or "").upper() if (font is not None and font.color is not None) else ""
                if rgb in GREY_SKIP_RGB:
                    grey.setdefault(sw, []).append("%s!%d" % (sheet, rr))
    finally:
        wb.close()
    return {"rows": rows, "grey": grey}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--templates", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--pilot", default=str(ROOT / "v15_pilot.py"))
    ap.add_argument("--min-candidates", type=int, default=2)
    ap.add_argument("--expect-ungated", default="DAYTRADE_DC_TARGET_TF,TECHNICAL_DC_TARGET_TF",
                    help="gate members the operator has deliberately greyed; reported separately as PENDING_REMOVAL")
    ap.add_argument("--json-out", default=str(ROOT / "data" / "reports" / "gate_surface_guard.json"))
    ap.add_argument("--quiet", action="store_true")
    a = ap.parse_args()

    names = gate_switches(pathlib.Path(a.pilot))
    expect = {s.strip() for s in a.expect_ungated.split(",") if s.strip()}
    breaches, pending, absent = [], [], []
    per_template = {}
    # 2026-10-28: the glob also matched another session's in-flight temp workbook
    # (TEMPLATE_CRYPTO_LONG.vec_status.tmp.xlsx, 51 KB) and openpyxl raised BadZipFile, which killed the whole guard --
    # a guard that dies reports nothing, which is worse than a breach it could have printed. The four canonical
    # workbooks are exactly TEMPLATE_{CRYPTO,STOCKS}_{LONG,SHORT}.xlsx; anything with extra dots in the stem is a temp
    # or backup file and is skipped. An unreadable canonical workbook is reported as a breach, never swallowed.
    CANON = {"TEMPLATE_%s_%s.xlsx" % (v, s) for v in ("CRYPTO", "STOCKS") for s in ("LONG", "SHORT")}
    for tf in sorted(pathlib.Path(a.templates).glob("TEMPLATE_*.xlsx")):
        if tf.name not in CANON:
            continue
        try:
            info = scan(tf, names)
        except Exception as exc:
            breaches.append({"template": tf.name, "switch": "(whole workbook)", "rows": 0, "grey_rows": 0,
                             "grey_at": [], "why": "canonical workbook unreadable: %s: %s" % (type(exc).__name__, exc)})
            continue
        per_template[tf.name] = {"n_gate_switches_present": len(info["rows"])}
        for sw in sorted(names):
            n = info["rows"].get(sw, 0)
            g = info["grey"].get(sw, [])
            if n == 0:
                absent.append({"template": tf.name, "switch": sw})
                continue
            fully_grey = len(g) == n
            rec = {"template": tf.name, "switch": sw, "rows": n, "grey_rows": len(g), "grey_at": g[:4]}
            if fully_grey and sw in expect:
                pending.append(rec)
                continue
            if g:
                breaches.append({**rec, "why": "gate switch has grey (never-evaluated) rows"})
            if n < a.min_candidates:
                breaches.append({**rec, "why": "gate switch has <%d candidates: delta is 0 by construction" % a.min_candidates})

    out = {"ts": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "n_gate_switches": len(names),
           "templates": per_template, "n_breaches": len(breaches), "breaches": breaches,
           "n_pending_removal": len(pending), "pending_removal": pending,
           "n_absent": len(absent), "absent": absent[:20]}
    try:
        p = pathlib.Path(a.json_out)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(out, indent=1))
    except Exception:
        pass

    if breaches:
        print("[gate-surface] BREACH: %d" % len(breaches))
        for b in breaches[:24]:
            print("   %-26s %-36s rows=%-3d grey=%-3d  %s" % (b["template"], b["switch"], b["rows"], b["grey_rows"], b["why"]))
    if pending and not a.quiet:
        print("[gate-surface] PENDING_REMOVAL (operator greyed these deliberately; remove from ESSENTIAL_GATE_SWITCHES): %d" % len(pending))
        for b in pending[:8]:
            print("   %-26s %-36s rows=%d all grey" % (b["template"], b["switch"], b["rows"]))
    if absent and not a.quiet:
        print("[gate-surface] gate switches with NO row in a template: %d (first: %s)" % (len(absent), absent[0]))
    if not breaches and not a.quiet:
        print("[gate-surface] OK: every gate switch testable (no grey rows, >=%d candidates) in all templates" % a.min_candidates)
    return 1 if breaches else 0


if __name__ == "__main__":
    sys.exit(main())
