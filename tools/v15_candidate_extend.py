"""v15_candidate_extend — add candidates to a switch that already has rows (USER 2026-10-10: ENTRY_DC_TF TF variants).

Sibling of tools/v15_switch_add.py for the case the adder refuses ("already has
rows"): the switch is wired + bible-green, but the template carries only a subset
of its domain (e.g. ENTRY_DC_TF OFF without 15m/1h/4h/D). Same guards as the adder:
config/QuickConfig/bible/vec/live checks, whole-row insert, style mirror, backup per
file, fingerprints lost==0, exactly-one-bold-default gate, zip-validated atomic save.

Usage: python3 tools/v15_candidate_extend.py --switch NAME --tab TAB --candidates a,b
       --venues both --sides both [--apply] [--fleet]
"""
import argparse
import copy
import datetime
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import openpyxl  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from template_row_guard import fingerprints  # noqa: E402
from v15_switch_add import SWITCH_SHEETS, FILES, FLEET, ast_assigns, parse_opt, die  # noqa: E402


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--switch", required=True)
    ap.add_argument("--tab", required=True)
    ap.add_argument("--candidates", required=True)
    ap.add_argument("--venues", default="both", choices=["crypto", "stocks", "both"])
    ap.add_argument("--sides", default="both", choices=["long", "short", "both"])
    ap.add_argument("--templates-dir", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--fleet", action="store_true")
    a = ap.parse_args()
    sw = a.switch.strip().upper()
    if a.tab not in SWITCH_SHEETS:
        die(f"unknown tab {a.tab} (must be one of the 13 SWITCH_SHEETS)")
    cands = [c.strip() for c in a.candidates.split(",") if c.strip() != ""]
    if len(set(cands)) < 1:
        die("need at least 1 new candidate")
    venues = ["CRYPTO", "STOCKS"] if a.venues == "both" else [a.venues.upper()]
    sides = ["LONG", "SHORT"] if a.sides == "both" else [a.sides.upper()]
    if "_LONG" in sw:
        sides = ["LONG"]
    if "_SHORT" in sw:
        sides = ["SHORT"]
    want = None
    for v in venues:
        cfg = ast_assigns(ROOT / ("config.py" if v == "CRYPTO" else "config_tradier.py"), "Config" if v == "CRYPTO" else "TradierConfig")
        if sw not in cfg:
            die(f"{sw} is not a field of {'config.py' if v == 'CRYPTO' else 'config_tradier.py'} — add the config field first")
        dv = cfg[sw]
        t = bool if isinstance(dv, bool) else (int if isinstance(dv, int) and not isinstance(dv, bool) else (float if isinstance(dv, float) else str))
        want = t if want is None else want
    qc = ast_assigns(ROOT / "v12_quick_engine.py", "QuickConfig")
    if sw not in qc:
        die(f"{sw} is not a QuickConfig field — wire v12_quick_engine first")
    try:
        bible = json.loads((ROOT / "data" / "SWITCH_BIBLE.json").read_text())["switches"]
    except Exception as e:
        die(f"cannot read data/SWITCH_BIBLE.json ({e}) — rebuild the bible first")
    e = bible.get(sw)
    if not e:
        die(f"{sw} has no SWITCH_BIBLE entry — rebuild the bible first")
    if str(e.get("status") or "").upper().startswith("DEAD"):
        die(f"{sw} bible status is dead ({e.get('status')})")
    if not (e.get("vec_reads") or e.get("vec_read_count")):
        die(f"{sw} has no vec reads — wire v12_quick_engine/vec_decisions first")
    if not any((e.get("live_read_count") or {}).values()):
        die(f"{sw} has no live reads — wire ez_manage/tradier_manage first")
    reg = (e.get("agent_c_registry") or {})
    if reg.get("suggested_home_tab") and reg["suggested_home_tab"] != a.tab:
        die(f"registry suggested_home_tab is {reg['suggested_home_tab']}, not {a.tab}")
    files = [(f"{v}_{s}", FILES[f"{v}_{s}"]) for v in venues for s in sides]
    typed = [parse_opt(c, want) for c in cands]
    plan = []
    for cs, fn in files:
        p = Path(a.templates_dir) / fn
        if not p.exists():
            die(f"missing {p}")
        wb = openpyxl.load_workbook(str(p))
        try:
            ws = wb[a.tab]
            have = [(r, ws.cell(r, 2).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == sw]
            if not have:
                die(f"{sw} has NO rows in {fn}!{a.tab} — use v15_switch_add.py for new switches")
            have_vals = set()
            for _, bv in have:
                try:
                    have_vals.add(parse_opt(str(bv), want))
                except Exception:
                    have_vals.add(bv)
            dup = [t for t in typed if t in have_vals]
            if dup:
                die(f"{fn}!{a.tab}: candidates {dup} already exist for {sw} (have {[bv for _, bv in have]})")
            last_sw = max(r for r, _ in have)
            plan.append((cs, fn, p, last_sw, typed))
        finally:
            wb.close()
    for cs, fn, p, last_sw, opts in plan:
        print(f"[{cs}] {fn}!{a.tab}: insert {len(opts)} rows at {last_sw + 1} (after last {sw} row) cands={opts}", flush=True)
    if not a.apply:
        print("dry-run OK — re-run with --apply to write (backs up each file first)", flush=True)
        return 0
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for cs, fn, p, last_sw, opts in plan:
        shutil.copy2(p, ROOT / "backups" / f"before_candidate_extend_{sw}_{ts}_{fn}")
        wb = openpyxl.load_workbook(str(p))
        try:
            ws = wb[a.tab]
            before = fingerprints(ws)
            ws.insert_rows(last_sw + 1, len(opts))
            hdrs = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
            isd = hdrs.index("is_default") + 1
            ref = last_sw
            for i, o in enumerate(opts):
                rr = last_sw + 1 + i
                for c in range(1, ws.max_column + 1):
                    tgt, src = ws.cell(rr, c), ws.cell(ref, c)
                    tgt.font, tgt.fill, tgt.alignment, tgt.border, tgt.number_format = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.alignment), copy.copy(src.border), src.number_format
                ws.cell(rr, 1).value = sw
                ws.cell(rr, 2).value = o
                ws.cell(rr, 2).font = Font(bold=False)
                ws.cell(rr, isd).value = "NO"
                ws.cell(rr, 1).comment = Comment(f"candidate-extend {ts}: domain completion by tools/v15_candidate_extend.py", "candidate_extend")
            after = fingerprints(ws)
            lost = sum((before - after).values())
            if lost:
                die(f"{fn}!{a.tab}: {lost} original rows changed — aborting, restore from backups/before_candidate_extend_{sw}_{ts}_{fn}")
            grp = [(ws.cell(r, 2).font.bold, ws.cell(r, isd).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == sw]
            if sum(1 for b, y in grp if b and y == "YES") != 1 or sum(1 for b, y in grp if (b and y != "YES") or (y == "YES" and not b)):
                die(f"{fn}!{a.tab}: DEFAULTS-GATE failed for {sw} (need exactly one bold B == YES)")
            tmp = str(p) + ".exttmp"
            wb.save(tmp)
            import zipfile
            with zipfile.ZipFile(tmp) as z:
                if len(z.namelist()) < 10:
                    die(f"refusing truncated save {fn}")
            os.replace(tmp, str(p))
            print(f"[{cs}] installed {len(opts)} rows -> {fn}!{a.tab}", flush=True)
        finally:
            wb.close()
    if a.fleet:
        for h in FLEET:
            for _, fn in files:
                subprocess.run(["rsync", "-az", "-e", "ssh -S none -o StrictHostKeyChecking=accept-new", str(Path(a.templates_dir) / fn), f"{h}:~/binance-sandbox/SPREADSHEETS/TEMPLATE_FINAL_NORM/{fn}"], check=False)
            r = subprocess.run(["ssh", "-S", "none", "-o", "StrictHostKeyChecking=accept-new", h, "cd ~/binance-sandbox/SPREADSHEETS/TEMPLATE_FINAL_NORM && md5sum " + " ".join(fn for _, fn in files)], capture_output=True, text=True)
            print(f"[{h}] {r.stdout.strip()}", flush=True)
    else:
        print("fleet NOT synced — re-run with --fleet (or rsync + md5 by hand)", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
