#!/usr/bin/env python3
"""v15_template_add_missing — add the REAL switches that the SWITCH_BIBLE says are wired (live read + reachable vec read) but sit in NO template
(tools/verify_switch_bible.py COVERAGE) to the right tab of the right template(s), whole new rows only (existing rows are never touched).
Source = live templates (read), output = a STAGED dir. Placement: registry suggested_home_tab (else template_home_rules), venue = every venue whose
config has the field and where the bible status is not dead; side = both sides unless the name carries a LONG/SHORT token. Default = venue config value
(typed, per tools/v15_template_options.py), options = the registry valid_options that fit the config type (+ BOOL pair / TF domain), blank stats, no yellow.
Also puts the default row first in STDEV_SLOPE_SIZING. Decisions + evidence -> <report-dir>/added_<cat>.csv.
  python tools/v15_template_add_missing.py --src-dir SPREADSHEETS --out-dir SPREADSHEETS/TEMPLATE_STAGED/<ts> --report-dir data/template_audit/<ts>
"""
import argparse
import copy
import csv
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import v15_template_home_rules as H  # noqa: E402
import v15_template_options as O  # noqa: E402
import v15_template_restructure_v3 as R  # noqa: E402

OK_STATUS = ("WIRED", "LIVE", "VEC")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src-dir", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--report-dir", required=True)
    a = ap.parse_args()
    out, rep = Path(a.out_dir), Path(a.report_dir)
    out.mkdir(parents=True, exist_ok=True)
    rep.mkdir(parents=True, exist_ok=True)
    sys.argv = ["x"]
    r = subprocess.run([sys.executable, str(ROOT / "tools" / "verify_switch_bible.py"), "--json", "/tmp/_svb_cov.json"], capture_output=True, text=True, cwd=str(ROOT))
    cov = [x.split(":")[0] for x in json.loads(Path("/tmp/_svb_cov.json").read_text())["checks"]["COVERAGE"]]
    bible = json.loads((ROOT / "data" / "SWITCH_BIBLE.json").read_text())["switches"]
    usage_lines = {}
    try:
        usage_lines = json.loads(sorted((ROOT / "data" / "template_audit").glob("*/code_usage_lines.json"))[-1].read_text())
    except Exception:
        pass
    R.LINES.update(usage_lines)
    yfill_none = PatternFill(fill_type=None)
    summary = {}
    for cs, name in R.TEMPLATES.items():
        venue, side = cs.split("_")
        vkey = "crypto" if venue == "CRYPTO" else "stocks"
        cfgkey = "config.py" if venue == "CRYPTO" else "config_tradier.py"
        F = R.Fields(cs)
        F2 = R.Fields(("STOCKS_" if venue == "CRYPTO" else "CRYPTO_") + side)
        wb = openpyxl.load_workbook(str(Path(a.src_dir) / name))
        present = set()
        for t in H.TABS:
            for row in wb[t].iter_rows(min_row=3, max_col=1, values_only=True):
                if row[0]:
                    present.add(str(row[0]).strip())
        log, added_rows, per_tab = [], 0, {}
        for n in cov:
            e = bible.get(n) or {}
            reg = e.get("agent_c_registry") or {}
            why = None
            if n in present:
                continue
            s_side = R.H.side_of(n)
            if s_side and s_side != side:
                why = f"wrong side ({s_side}) for {cs}"
            elif not (e.get("in_config") or {}).get(cfgkey):
                why = f"not a field of {cfgkey} (venue config)"
            elif str((e.get("status") or {}).get(vkey, "")).upper().startswith(("DEAD", "NONE", "NO_")):
                why = f"status {e['status'].get(vkey)} on {vkey}"
            ref = F.get(n)
            if why is None and ref is R.MISSING:
                why = "no value in venue config object"
            if why is None and n.upper().endswith("_ALT"):
                why = "invented *_ALT"
            if why is None and (n in ("TRADIER_API_KEY",) or n.upper().endswith(("_API_KEY", "_API_SECRET", "_PRIVATE_KEY"))):
                why = "secret: never a sweep row (2026-10-04 cut5: live key was staged into workbooks+CSVs before this guard)"
            if why:
                log.append([cs, n, "", "NOT_ADDED", why, ""])
                continue
            oref = F2.get(n)
            cls = O.classify(n, ref, None if oref is R.MISSING else oref)
            if cls in ("STRUCT", "UNKNOWN"):
                log.append([cs, n, "", "NOT_ADDED", f"class {cls}: not expressible as scalar options", ""])
                continue
            okc, dv, _w = O.normalize(cls, ref, ref)
            opts = {O.key_of(dv): dv}
            if cls == "BOOL":
                opts[O.key_of(True)] = True
                opts[O.key_of(False)] = False
            elif cls == "BOOLNUM":
                for bv in (0, 1):
                    v_ = int(bv) if isinstance(ref, int) else float(bv)
                    opts[O.key_of(v_)] = v_
            else:
                for rv in reg.get("valid_options") or []:
                    ok, v_, _w2 = O.normalize(cls, ref, rv)
                    if ok and (cls not in ("INT", "FLOAT") or not O.foreign_numeric(O._num(dv), O._num(v_))):
                        opts.setdefault(O.key_of(v_), v_)
                if cls == "TF":
                    for tv in ("OFF", "15m", "1h", "4h", "D"):
                        opts.setdefault(O.key_of(tv), tv)
            note = "options from registry/type domain"
            if len(opts) < 2:
                note = "ONLY the default option exists (registry gave none; real alternatives need a decision)"
            home = reg.get("suggested_home_tab") or H.MANUAL_HOME.get(n, (None,))[0] or H.default_tab_for(H.strong_lifecycle(n) or "ENTRY", n)
            if home not in H.TABS:
                home = H.default_tab_for(H.strong_lifecycle(n) or "ENTRY", n)
            ws = wb[home]
            hdr = {str(ws.cell(2, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
            last_w = 2
            for rr in range(3, ws.max_row + 1):
                if ws.cell(rr, 1).value in (None, ""):
                    continue
                if not (ws.cell(rr, 1).fill.fill_type == "solid" and R.ORANGE in str(ws.cell(rr, 1).fill.fgColor.rgb or "")):
                    last_w = rr
            ref_row = last_w
            ws.insert_rows(last_w + 1, len(opts))
            ev = R.evidence(n)
            for i, (k, v_) in enumerate(opts.items()):
                rr = last_w + 1 + i
                for c in range(1, ws.max_column + 1):
                    ws.cell(rr, c)._style = copy.copy(ws.cell(ref_row, c)._style)
                    if c > 14 and c < hdr.get("Live Location", 10 ** 6):
                        ws.cell(rr, c).fill = copy.copy(yfill_none)
                isd = k == O.key_of(dv)
                ws.cell(rr, 1).value = n
                ws.cell(rr, 2).value = v_
                for h in ("override", "BASELINE", "HUSTLE_DELTA", "VECTOR_DELTA", "LIVE_DELTA", "LIVE_SHARPE", "REAL_COMPLETE", "PER_ROW_FILTERS", "AVG_DELTA", "POS_SYM", "Family", "Live Location", "Vec Hook", "Type", "Correct Options"):
                    if h in hdr:
                        ws.cell(rr, hdr[h]).value = None
                ws.cell(rr, hdr["is_default"]).value = "YES" if isd else "NO"
                ws.cell(rr, hdr["is_default"]).font = Font(name="Arial", size=10, bold=isd)
                bf = ws.cell(rr, 2).font
                ws.cell(rr, 2).font = Font(name=bf.name or "Arial", size=bf.size or 10, bold=isd, italic=bf.italic, color=None)
                af = ws.cell(rr, 1).font
                ws.cell(rr, 1).font = Font(name=af.name or "Arial", size=af.size or 10, bold=af.bold, italic=af.italic, color=None)
                ws.cell(rr, 1).fill = copy.copy(yfill_none)
                if i == 0:
                    ws.cell(rr, 1).comment = Comment(f"v3b: switch was wired live+vec but in no template (SWITCH_BIBLE COVERAGE). {note}", "template_add_missing")
            added_rows += len(opts)
            per_tab[home] = per_tab.get(home, 0) + 1
            log.append([cs, n, "|".join(str(v) for v in opts.values()), "ADDED", f"{home}; class {cls}; default {dv!r} (venue config); {note}; status {e.get('status')}", ev])
        # STDEV_SLOPE_SIZING: default row first
        ws = wb["STDEV_SLOPE_SIZING"]
        hdr = {str(ws.cell(2, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
        first = next((r_ for r_ in range(3, ws.max_row + 1) if ws.cell(r_, 1).value), None)
        if first and str(ws.cell(first, hdr["is_default"]).value).upper() != "YES":
            nxt = next((r_ for r_ in range(first + 1, ws.max_row + 1) if str(ws.cell(r_, hdr["is_default"]).value).upper() == "YES" and ws.cell(r_, 1).value == ws.cell(first, 1).value), None)
            if nxt:
                for c in range(1, ws.max_column + 1):
                    A, B = ws.cell(first, c), ws.cell(nxt, c)
                    A.value, B.value = B.value, A.value
                    A._style, B._style = copy.copy(B._style), copy.copy(A._style)
                log.append([cs, "STDEV_SLOPE_SIZING_ENABLED", "", "ROW_SWAP", "default row moved first in STDEV_SLOPE_SIZING (whole-row swap)", ""])
        wb.save(str(out / name))
        with open(rep / f"added_{cs}.csv", "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["template", "switch", "options", "decision", "reason", "evidence_file_line"])
            w.writerows(log)
        summary[cs] = {"switches_added": sum(1 for x in log if x[3] == "ADDED"), "rows_added": added_rows, "not_added": sum(1 for x in log if x[3] == "NOT_ADDED"), "per_tab": per_tab}
        print(cs, summary[cs], flush=True)
    (rep / "added_summary.json").write_text(json.dumps(summary, indent=1))


if __name__ == "__main__":
    main()
