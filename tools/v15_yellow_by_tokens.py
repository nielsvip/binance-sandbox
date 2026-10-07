#!/usr/bin/env python3
"""v15_yellow_by_tokens — assign yellow cells by name tokens (USER 2026-09-30, interim until the yellow agent's full map).

In every SWITCH_SHEETS tab of the 4 TEMPLATE_*.xlsx: the cell (row r, yellow column c) is YELLOW iff the row's switch name
(col A) and the column's filter name (row-2 header "FILTER=opt", part before "=") share at least 2 tokens (split on "_").
Every other cell of a yellow column gets the row's own background back (orange FFE699 on orange filter rows, none on white
rows). Only fills change: no values, no rows, no columns, no defaults. Report only unless --apply (backup first).
"""
import argparse
import datetime
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
SPREAD = ROOT / "SPREADSHEETS"
TEMPLATES = ["TEMPLATE_CRYPTO_LONG.xlsx", "TEMPLATE_CRYPTO_SHORT.xlsx", "TEMPLATE_STOCKS_LONG.xlsx", "TEMPLATE_STOCKS_SHORT.xlsx"]
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
MIN_SHARED = 2
YELLOW = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
ORANGE = PatternFill(start_color="FFFFE699", end_color="FFFFE699", fill_type="solid")
NONE = PatternFill(fill_type=None)


def tokens(name) -> set:
    return {t for t in str(name).strip().upper().split("_") if t}


def is_yellow(cell) -> bool:
    return cell.fill is not None and cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb or "").upper().endswith("FFFF00")


def is_orange(cell) -> bool:
    return cell.fill is not None and cell.fill.fill_type == "solid" and str(cell.fill.fgColor.rgb or "").upper().endswith("FFE699")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for name in TEMPLATES:
        path = SPREAD / name
        wb = openpyxl.load_workbook(str(path))
        before = after = changed = 0
        for tab in SWITCH_SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            ycols = {c: tokens(ws.cell(row=2, column=c).value.split("=", 1)[0]) for c in range(1, ws.max_column + 1) if isinstance(ws.cell(row=2, column=c).value, str) and "=" in ws.cell(row=2, column=c).value}
            for r in range(3, ws.max_row + 1):
                a = ws.cell(row=r, column=1).value
                if a in (None, ""):
                    continue
                st = tokens(a)
                row_fill = ORANGE if is_orange(ws.cell(row=r, column=1)) else NONE
                for c, ft in ycols.items():
                    cell = ws.cell(row=r, column=c)
                    was = is_yellow(cell)
                    want = len(st & ft) >= MIN_SHARED
                    before += was
                    after += want
                    if was != want:
                        changed += 1
                        if args.apply:
                            cell.fill = YELLOW if want else row_fill
        print(f"[{name}] yellow cells before={before} after={after} changed={changed}")
        if args.apply:
            shutil.copy2(path, ROOT / "backups" / f"before_yellow_by_tokens_{ts}_{name}")
            tmp = path.with_suffix(".tmp.xlsx")
            wb.save(str(tmp))
            openpyxl.load_workbook(str(tmp))
            tmp.replace(path)
            print(f"[{name}] saved")


if __name__ == "__main__":
    main()
