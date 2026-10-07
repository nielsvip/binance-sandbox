#!/usr/bin/env python3
"""v15_defaults_verify (DEF2 2026-10-01): per template tab group, exactly ONE bold default row, is_default YES only there, and its value == config default for the cat_side
(data/cat_side_defaults_4.json, fallback config values). Read-only. Writes <dir>/DEFAULTS_VERIFY.csv.   python tools/v15_defaults_verify.py [DIR]"""
import csv, json, sys, collections
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "tools"))
from v15_template_normalize_defaults import SW, norm
import build_cat_side_defaults_4 as B
d = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"
cat4 = json.loads((ROOT / "data" / "cat_side_defaults_4.json").read_text())
out = []
tot = collections.Counter()
for cs in ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"):
    typed, _l, _q = B.venue_values(cs.startswith("STOCKS")); typed = dict(typed)
    for k, v in cat4.get(cs, {}).items():
        if not isinstance(v, (dict, list, tuple, set)) and v is not None:
            typed[k] = v
    wb = openpyxl.load_workbook(str(d / f"TEMPLATE_{cs}.xlsx"), read_only=False)
    for tab in SW:
        if tab not in wb.sheetnames: continue
        ws = wb[tab]
        ci = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(2, c).value or "").strip().lower() == "is_default"), None)
        g = collections.OrderedDict()
        for r in range(3, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a not in (None, ""): g.setdefault(str(a).strip(), []).append(r)
        for sw, rows in g.items():
            grey = str(getattr(ws.cell(rows[0], 1).font.color, "rgb", "") or "").upper().endswith("BFBFBF")
            bold = [r for r in rows if ws.cell(r, 2).font and ws.cell(r, 2).font.b]
            yes = [r for r in rows if str(ws.cell(r, ci).value or "").strip().upper() == "YES"]
            cfg = typed.get(sw)
            if len(bold) != 1 or yes != bold: st = "NOT_ONE_DEFAULT"
            elif cfg is None or isinstance(cfg, (dict, list, tuple, set)): st = "NO_CONFIG_VALUE"
            elif norm(ws.cell(bold[0], 2).value) == norm(cfg): st = "OK"
            else: st = "MISMATCH_CONFIG"
            tot[(cs, st, grey)] += 1
            if st != "OK": out.append([cs, tab, sw, st, grey, ws.cell(bold[0], 2).value if len(bold) == 1 else "", cfg, len(bold), len(yes)])
with (d / "DEFAULTS_VERIFY.csv").open("w", newline="") as fh:
    w = csv.writer(fh); w.writerow(["cat_side", "tab", "switch", "status", "grey", "bold_value", "config_default", "n_bold", "n_yes"]); w.writerows(out)
for k, v in sorted(tot.items()): print(k, v)
