#!/usr/bin/env python3
"""v18_learn_common — shared, pure helpers of the V18 learning layer (doc 17 §3.2 / §3.5 / §5 ABL-0 / principles 10-13).

Used by tools/v18_ledger_replay.py, tools/v18_evidence_mine.py, tools/v18_exit_loss_rank.py, tools/v18_winners_pack.py,
tools/v18_filter_veto_map.py and tools/v18_digest.py. Stdlib only at import time (openpyxl is imported lazily by the
template reader, so the S1/S2 system python3 can import this module for the miner).

LAWS KEPT HERE (one place, so every learner applies them identically):
  * Four cat_sides, strictly apart (principle 13): `cat_side_of` refuses anything that is not *_LONG / *_SHORT.
  * Era cut at the hermetic dates (crypto 2026-10-14T13:00Z, stocks 2026-10-17T13:00Z): `era_of`.
  * One exit-family vocabulary for board ledgers and /history: `exit_family` = loss_prone.fn_family + the
    MAKER_PROFIT_EXIT_ maker-leg prefix folded onto its family (the maker leg is how the vector exit is executed).
  * Stop/execution families are reported, never repair targets of the board: `is_stop_execution`.
  * Settings never promoted: ABLATION_* (the strip lever), MODE/VENUE/IS_TRADIER/SIMPLE_PRICE_GT0_ENABLED, the pilot's
    GLOBAL_ONLY_KEYS (read from v15_pilot.py, never imported), MAXAUG sentinels, tokenized stocks, type-refused values.
  * P(pos) is a Beta posterior shrunk to the cat_side's own base rate (strength V15_BRAIN_M, default 8, as v15_brain);
    there is no GLOBAL scope because that would average across cat_sides.
"""
from __future__ import annotations

import gzip
import hashlib
import json
import os
import re
import shutil
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(os.environ.get("V18_ROOT") or Path(__file__).resolve().parents[1])
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
LEARN_DIR = Path(os.environ.get("V18_LEARN_DIR") or (ROOT / "data" / "learn"))
CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
CRYPTO_QUOTES = ("USDT", "USDC", "FDUSD", "BUSD", "TUSD")
EPS = 1e-9
M_SHRINK = float(os.environ.get("V15_BRAIN_M", "8"))
FLOOR_PER_1000 = float(os.environ.get("V18_FLOOR_PER_1000", "5"))
XPOLL_K = int(os.environ.get("V18_XPOLL_K", "5"))
WINDOW_DAYS = float(os.environ.get("V18_WINDOW_DAYS", "14"))
DENY_KEYS = frozenset({"MODE", "VENUE", "IS_TRADIER", "SIMPLE_PRICE_GT0_ENABLED"})
SWITCH_SHEETS = ("STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL",
                 "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK",
                 "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES")
TEMPLATE_DIRS = (ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM", ROOT / "SPREADSHEETS")
STOP_EXEC_PREFIXES = ("FIXED_STOP", "VEC_STALE_FLATTEN", "VEC_EXIT_OWED", "VEC_CONVERGE")
MAKER_PREFIX = "MAKER_PROFIT_EXIT_"
FIRST_YELLOW_COL = 14  # column O: first FILTER=opt header on every switch tab (A..N are fixed columns)
MAXAUG_SENTINELS = frozenset({"999999", "999999.0", "1000000", "1000000.0"})

try:  # USE the store's definitions (tools/v18_experience.py owns them); identical fallbacks if it cannot be imported
    from tools.v18_experience import HERMETIC_CUT as _HC, is_broker_echo as _echo  # noqa: E402
    HERMETIC_CUT = dict(_HC)
    is_broker_echo = _echo
except Exception:  # pragma: no cover - fallback only
    HERMETIC_CUT = {"CRYPTO": datetime(2026, 10, 14, 13, 0, tzinfo=timezone.utc).timestamp(),
                    "STOCKS": datetime(2026, 10, 17, 13, 0, tzinfo=timezone.utc).timestamp()}

    def is_broker_echo(reason):
        s = str(reason or "")
        return "BROKER_RECONCILED" in s or "broker sync" in s.lower()

ERA_CUT_ISO = {"CRYPTO": "2026-10-14T13:00:00Z", "STOCKS": "2026-10-17T13:00:00Z"}


def utc(epoch=None):
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(time.time() if epoch is None else epoch))


def stamp():
    return time.strftime("%Y%m%d%H%M", time.gmtime())


def iso_epoch(s):
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    txt = str(s).strip().replace("Z", "+00:00")
    m = re.match(r"^(.*T\d\d:\d\d:\d\d)(\.\d+)?(.*)$", txt)
    if m and m.group(2):
        txt = m.group(1) + (m.group(2) + "000000")[:7] + m.group(3)
    try:
        dt = datetime.fromisoformat(txt)
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.timestamp()


def cat_side_of(sym_side):
    """'BTCUSDC_LONG' -> 'CRYPTO_LONG'; anything that is not <SYM>_LONG/_SHORT -> None (refused, never guessed)."""
    s = str(sym_side or "").strip()
    if "_" not in s:
        return None
    symbol, side = s.rsplit("_", 1)
    side = side.upper()
    if side not in ("LONG", "SHORT") or not symbol:
        return None
    venue = "CRYPTO" if symbol.upper().endswith(CRYPTO_QUOTES) else "STOCKS"
    return f"{venue}_{side}"


def venue_of(cat_side):
    return str(cat_side).split("_", 1)[0]


def era_of(cat_side, epoch):
    if epoch is None or not cat_side:
        return "undated"
    return "hermetic" if float(epoch) >= HERMETIC_CUT[venue_of(cat_side)] else "pre_hermetic"


def is_tokenized(sym_side):
    try:
        from tools.excluded_tokenized_stocks import is_denied
        return bool(is_denied(sym_side))
    except Exception:
        return False


def _fn_family(reason):
    try:
        import loss_prone
        return loss_prone.fn_family(reason)
    except Exception:  # pragma: no cover - loss_prone is always present in the repo/sandbox
        r = re.split(r"[ |:(]", str(reason or "").strip())[0]
        return r or "UNKNOWN"


_TF_TOKENS = frozenset({"1m", "3m", "5m", "15m", "30m", "1h", "2h", "4h", "6h", "8h", "12h", "1d", "D", "W", "M"})
_PARAM_CHARS = re.compile(r"[=\[\]%<>]")
_ALIASES = (("EXIT_OWED", "VEC_EXIT_OWED"),)


def canon_family(fam):
    """The vector-label extension of fn_family (doc 17 §3.2): cut the parameter tail fn_family leaves on underscore-joined
    reasons (`DC_HOPELESS_entry=43.36_dc=[..]`, `PEAK_GIVEBACK_GAIN_EROSION_STOP_peak0.66%`, `WT_DC_EXIT_100_g=-0.96%`,
    `DAYTRADE_TARGET_dc_15m_high_-0.10%_TARGET_BUF`, `WHOLE_SHARE_REDUCE_TO_FLAT[SATOSHIT_EXIT_PARTIAL`): stop at the first
    token carrying = [ ] % < > or starting lower-case (TF tokens 15m/1h/4h... are kept), then drop trailing pure numbers."""
    f = str(fam or "")
    f = _PARAM_CHARS.split(f, 1)[0] if "[" in f else f
    out = []
    for tok in f.split("_"):
        if not tok:
            continue
        if (_PARAM_CHARS.search(tok) or (tok[0].islower() and tok not in _TF_TOKENS) or tok[0] in "-+."
                or re.fullmatch(r"\d*\.\d+\w*", tok)):
            break
        out.append(tok)
    while out and re.fullmatch(r"[-+]?\d+(\.\d+)?", out[-1]):
        out.pop()
    f = "_".join(out) or str(fam or "UNKNOWN")
    for pre, alias in _ALIASES:
        if f.startswith(pre):
            return alias
    return f


def exit_family(reason):
    """loss_prone.fn_family, then the maker-leg prefix folded onto the family it executes (MAKER_PROFIT_EXIT_EXIT_VELOCITY_WT ->
    EXIT_VELOCITY_WT, MAKER_PROFIT_EXIT_FIXED_STOP -> FIXED_STOP), then canon_family, so board ledgers and /history share one
    vocabulary."""
    fam = _fn_family(reason)
    if fam.startswith(MAKER_PREFIX) and len(fam) > len(MAKER_PREFIX):
        fam = fam[len(MAKER_PREFIX):]
    return canon_family(fam) or "UNKNOWN"


def map_to_vocab(fam, vocab):
    """Longest vocabulary family that is a token-prefix of `fam` (UNCOND_DC_LOW_15M_UNCOND -> UNCOND_DC_LOW_15M); else fam."""
    if fam in vocab:
        return fam
    best = None
    for v in vocab:
        if fam.startswith(v + "_") and (best is None or len(v) > len(best)):
            best = v
    return best or fam


def entry_family(reason):
    return _fn_family(reason)


def is_stop_execution(family):
    f = str(family or "")
    return any(f.startswith(p) for p in STOP_EXEC_PREFIXES)


_GLOBAL_ONLY = None


def global_only_keys():
    """The pilot's GLOBAL_ONLY_KEYS, read from v15_pilot.py source (never imported: importing the pilot has side effects)."""
    global _GLOBAL_ONLY
    if _GLOBAL_ONLY is None:
        keys = set()
        try:
            src = (ROOT / "v15_pilot.py").read_text(errors="replace")
            m = re.search(r"^GLOBAL_ONLY_KEYS\s*=\s*frozenset\(\{(.*?)\}\)", src, re.M | re.S)
            if m:
                keys = set(re.findall(r"\"([A-Z0-9_]+)\"", m.group(1)))
        except OSError:
            pass
        _GLOBAL_ONLY = frozenset(keys)
    return _GLOBAL_ONLY


def excluded_key_reason(switch, cand=None):
    """Why a SWITCH=cand may never be a winners/menu candidate (None = allowed). Ablation keys are the strip lever only."""
    s = str(switch or "").strip()
    if not s or not re.match(r"^[A-Z][A-Z0-9_]*$", s):
        return "not_a_switch_name"
    if s.startswith("ABLATION_"):
        return "ABLATION_lever"
    if s in DENY_KEYS:
        return "deny_listed"
    if s in global_only_keys():
        return "GLOBAL_ONLY"
    if cand is not None and str(cand).strip() in MAXAUG_SENTINELS:
        return "MAXAUG_sentinel"
    return None


def type_refused(switch, cand, cat_side):
    """tools/switch_types.coerce verdict: None = accepted, else the refusal reason (the types law, doc 17 principle 8)."""
    try:
        from tools.switch_types import coerce
    except Exception:
        return None
    venue = "tradier" if venue_of(cat_side) == "STOCKS" else "crypto"
    try:
        ok, _v, why = coerce(switch, cand, venue)
    except Exception as e:  # a crash of the type law is a refusal, never a pass
        return f"coerce_error:{e}"
    return None if ok else (why or "refused")


def p_pos_shrunk(pos, n_valid, p0, m=M_SHRINK):
    """Beta posterior mean shrunk to the cat_side prior p0 with strength m (v15_brain.Brain.p_pos, cat_side scope only)."""
    return (float(pos) + m * float(p0)) / (max(0.0, float(n_valid)) + m)


def forward_factor(n_forward, n_forward_ok):
    """Share of forward/holdout/live-attributed outcomes with delta >= 0; unknown (n=0) = 0.5, never a bonus."""
    return round(float(n_forward_ok) / float(n_forward), 4) if n_forward else 0.5


def template_path(cat_side):
    for d in TEMPLATE_DIRS:
        p = d / f"TEMPLATE_{cat_side}.xlsx"
        if p.exists():
            return p
    return None


def _md5(p, chunk=1 << 20):
    h = hashlib.md5()
    with open(p, "rb") as fh:
        for b in iter(lambda: fh.read(chunk), b""):
            h.update(b)
    return h.hexdigest()


def md5_file(p):
    return _md5(p)


def cell_str(v):
    """How a template value reads as a delta-log cand / yellow option: the pilot writes str(value) of the openpyxl cell."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "True" if v else "False"
    return str(v).strip()


def template_index(cat_side, cache=True):
    """{tabs: {tab: {"rows": [[row, switch, cand, is_default]], "yellows": [hdr...]}}, switches: {switch: [cand...]},
    defaults: {switch: cand}, md5, path} for the cat_side's TEMPLATE. Cached in data/learn/_cache by template md5."""
    p = template_path(cat_side)
    if p is None:
        raise FileNotFoundError(f"no TEMPLATE_{cat_side}.xlsx in {TEMPLATE_DIRS}")
    md5 = _md5(p)
    cp = LEARN_DIR / "_cache" / f"template_index_{cat_side}_{md5[:12]}.json"
    if cache and cp.exists():
        try:
            return json.loads(cp.read_text())
        except ValueError:
            pass
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    tabs, switches, defaults = {}, {}, {}
    try:
        for tab in SWITCH_SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            rows, yellows = [], []
            for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if i == 2:
                    yellows = [str(h).strip() for h in row[FIRST_YELLOW_COL:] if h is not None and "=" in str(h)]
                    continue
                sw = row[0] if row else None
                if sw is None or not str(sw).strip():
                    continue
                sw = str(sw).strip()
                cand = cell_str(row[1] if len(row) > 1 else None)
                isdef = str(row[11]).strip().upper() == "YES" if len(row) > 11 and row[11] is not None else False
                rows.append([i, sw, cand, isdef])
                lst = switches.setdefault(sw, [])
                if cand not in lst:
                    lst.append(cand)
                if isdef and sw not in defaults:
                    defaults[sw] = cand
            tabs[tab] = {"rows": rows, "yellows": yellows}
    finally:
        wb.close()
    out = {"cat_side": cat_side, "path": str(p), "md5": md5, "tabs": tabs, "switches": switches, "defaults": defaults}
    if cache:
        cp.parent.mkdir(parents=True, exist_ok=True)
        write_json(cp, out, indent=None)
    return out


def template_pairs(tidx):
    """Every (switch=cand, filter=opt) cell of the template: (tab, row, switch, cand, yellow_header)."""
    for tab, blk in tidx["tabs"].items():
        for _r, sw, cand, _d in blk["rows"]:
            for y in blk["yellows"]:
                yield tab, sw, cand, y


def write_json(path, obj, indent=1):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp.{os.getpid()}")
    with open(tmp, "w") as fh:
        json.dump(obj, fh, indent=indent, sort_keys=False, default=str, separators=(",", ":") if indent is None else None)
    os.replace(tmp, path)
    return path


def read_json(path, default=None):
    p = Path(path)
    if not p.exists():
        return default
    if p.suffix == ".gz":
        with gzip.open(p, "rt") as fh:
            return json.load(fh)
    return json.loads(p.read_text())


def backup_file(path, desc):
    """CLAUDE.md: cp <file> backups/before_<desc>_<YYYYMMDDHHMM>.<ext> before editing any existing file."""
    p = Path(path)
    if not p.exists():
        return None
    bdir = ROOT / "backups"
    bdir.mkdir(exist_ok=True)
    dst = bdir / f"before_{desc}_{stamp()}{''.join(p.suffixes) or '.bak'}"
    shutil.copy2(p, dst)
    return dst


def md_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    for r in rows:
        out.append("| " + " | ".join("" if v is None else (f"{v:.2f}" if isinstance(v, float) else str(v)) for v in r) + " |")
    return "\n".join(out)


def _num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def template_cand(sw, cand, tidx):
    """The template's own spelling of `cand` for switch `sw` ('1.0' -> '1' when the row reads 1), else None. Bools/strings
    must match exactly; numbers match by value (older templates wrote 1.0 where the current one writes 1)."""
    opts = (tidx.get("switches") or {}).get(sw)
    if not opts:
        return None
    c = str(cand)
    if c in opts:
        return c
    if c in ("True", "False"):
        return None
    v = _num(c)
    if v is None:
        return None
    for o in opts:
        if o in ("True", "False"):
            continue
        ov = _num(o)
        if ov is not None and abs(ov - v) <= 1e-12 * max(1.0, abs(v)):
            return o
    return None


def canonicalize_evidence(ev, tidx):
    """Re-key the miner's naked / pairs / abl / quartet entries to the TEMPLATE spelling of each switch=cand. Two log
    spellings of one template candidate ('X=1' and 'X=1.0') keep the entry with more valid sym_sides (never summed: the
    same sym_side may sit under both spellings). Keys whose cand is not a template candidate are left as they are."""
    if not ev or ev.get("_canon_md5") == tidx.get("md5"):
        return ev
    out = dict(ev)

    def ck(swc):
        sw, _, cand = swc.partition("=")
        t = template_cand(sw, cand, tidx)
        return f"{sw}={t}" if t is not None else swc

    for name, size_of in (("naked", lambda v: v[1]), ("pairs", lambda v: v[1]), ("abl", lambda v: v.get("n_syms", 0)),
                          ("quartets", lambda v: v.get("n_syms", 0)), ("pair_pos_syms", len)):
        src = ev.get(name) or {}
        dst = {}
        for k, v in src.items():
            swc, sep, rest = k.partition("|")
            nk = ck(swc) + (sep + rest if sep else "")
            if nk not in dst or size_of(v) > size_of(dst[nk]):
                dst[nk] = v
        out[name] = dst
    out["_canon_md5"] = tidx.get("md5")
    return out
