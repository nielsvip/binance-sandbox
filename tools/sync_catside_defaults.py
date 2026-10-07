#!/usr/bin/env python3
"""sync_catside_defaults — overhaul foundation (DAILY_OPTIMIZATION_PLAN.md Stage 5).

Source of truth = the TEMPLATE bold value (col B) per switch, per cat_side. This tool EXTRACTS the
per-cat_side default map from the 4 TEMPLATE_{cat_side}.xlsx (bold col-B value per switch across the
13 SWITCH_SHEETS) into data/cat_side_defaults.json, and reports one-bold-per-switch violations.

MODES:
  --extract (default): read templates → write data/cat_side_defaults.json + report. Touches NO config.
  --write-configs: (guarded, later) inject CAT_SIDE_DEFAULTS + a default_for() resolver into config.py,
     config_tradier.py, v12_quick_engine.py. NOT enabled yet — extract + operator review first.

No fabrication; a switch with no bold row is reported as MISSING (not guessed). Never deletes anything.
"""
import argparse, json, pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]
SPREAD = ROOT / "SPREADSHEETS"
CAT_SIDES = {
    "CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx",
    "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx",
    "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx",
    "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx",
}
# the 13 SWITCH_SHEETS (skip legend/instruction/filters tabs)
SKIP_TABS = {"LEGEND_FILTERS", "INSTRUCTIONS", "FILTERS_EXPLAINED", "INSTRUCTIONS_V2", "WIRING_INVENTORY"}


def coerce(v):
    if isinstance(v, str):
        s = v.strip()
        if s.lower() in ("true", "false"):
            return s.lower() == "true"
        try:
            return int(s) if s.lstrip("-").isdigit() else float(s)
        except Exception:
            return s
    return v


def extract_template(path):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=False)
    defaults = {}          # switch -> value (bold col-B)
    violations = []        # switches with 0 or >1 bold rows
    for tab in wb.sheetnames:
        if tab in SKIP_TABS:
            continue
        ws = wb[tab]
        # group rows by switch (col A), find bold col-B row(s)
        by_switch = {}      # switch -> list of (row, value, bold)
        for r in range(3, ws.max_row + 1):
            a = ws.cell(row=r, column=1).value
            if a in (None, ""):
                continue
            bcell = ws.cell(row=r, column=2)
            bold = bool((ws.cell(row=r, column=1).font and ws.cell(row=r, column=1).font.bold)
                        or (bcell.font and bcell.font.bold))
            by_switch.setdefault(str(a).strip(), []).append((r, bcell.value, bold))
        for sw, rows in by_switch.items():
            bolds = [(rw, val) for (rw, val, b) in rows if b]
            if len(bolds) == 1:
                defaults[sw] = coerce(bolds[0][1])
            elif len(bolds) == 0:
                violations.append((tab, sw, "NO_BOLD", [str(x[1]) for x in rows][:6]))
            else:
                violations.append((tab, sw, f"MULTI_BOLD({len(bolds)})", [str(v) for _, v in bolds][:6]))
                defaults[sw] = coerce(bolds[0][1])  # provisional: first bold; flagged for fix
    wb.close()
    return defaults, violations


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--extract", action="store_true", default=True)
    ap.add_argument("--out", default=str(ROOT / "data" / "cat_side_defaults.json"))
    args = ap.parse_args()
    result = {}
    all_viol = {}
    for cs, fn in CAT_SIDES.items():
        p = SPREAD / fn
        if not p.exists():
            print(f"[{cs}] MISSING TEMPLATE {p}")
            continue
        defs, viol = extract_template(p)
        result[cs] = defs
        all_viol[cs] = viol
        print(f"[{cs}] defaults={len(defs)} bold-violations={len(viol)}")
        for tab, sw, kind, vals in viol[:8]:
            print(f"    VIOLATION {kind} {tab}:{sw} -> {vals}")
    outp = pathlib.Path(args.out)
    outp.parent.mkdir(parents=True, exist_ok=True)
    outp.write_text(json.dumps({"defaults": result, "violations": all_viol}, indent=2, default=str))
    print(f"[written] {outp}")
    # cross-cat_side divergence preview: switches whose default differs across cat_sides
    keys = set().union(*[set(result.get(cs, {})) for cs in CAT_SIDES]) if result else set()
    diverge = 0
    for k in keys:
        vals = {cs: result.get(cs, {}).get(k) for cs in CAT_SIDES if k in result.get(cs, {})}
        if len(set(map(str, vals.values()))) > 1:
            diverge += 1
    print(f"[divergence] {diverge}/{len(keys)} switches already differ across cat_sides "
          f"(will grow as promotion runs per cat_side)")


if __name__ == "__main__":
    main()
