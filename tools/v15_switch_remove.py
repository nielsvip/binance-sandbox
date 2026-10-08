"""v15_switch_remove — THE switch-removing tool (USER 2026-10-08, sibling of v15_switch_add.py).

Removes EVERY row of ONE switch (column A == NAME) from every SWITCH_SHEETS tab of the
TEMPLATE_*.xlsx files in --templates-dir. Whole-row delete only; every other row must
survive byte-for-byte (template_row_guard fingerprints: lost == removed rows, gained == 0).
Refuses while the switch is still a config / QuickConfig field (remove the code first —
a template row for a live key is never deleted) and refuses truncated saves.

Order of operations:
  1. remove the key from config + QuickConfig + live/vec key lists (code first, tests green)
  2. THIS TOOL dry-run (default), review, then --apply (backs up each file first)
  3. run it on the template SOURCE host (S1) for both template sets, then --fleet / md5

Usage: python3 tools/v15_switch_remove.py --switch NAME [--templates-dir DIR] [--apply] [--fleet]
"""
import argparse
import datetime
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import openpyxl  # noqa: E402
from template_row_guard import fingerprints  # noqa: E402
from v15_switch_add import FILES, FLEET, SWITCH_SHEETS, ast_assigns, die  # noqa: E402


def plan_file(path, sw):
    wb = openpyxl.load_workbook(str(path))
    try:
        hits = {}
        for tab in SWITCH_SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            rows = [r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == sw]
            if rows:
                hits[tab] = rows
        return hits
    finally:
        wb.close()


def remove_file(path, sw, hits, ts):
    shutil.copy2(path, ROOT / "backups" / f"before_switch_remove_{sw}_{ts}_{path.name}")
    wb = openpyxl.load_workbook(str(path))
    try:
        removed = 0
        for tab, rows in hits.items():
            ws = wb[tab]
            before = fingerprints(ws)
            for r in sorted(rows, reverse=True):
                ws.delete_rows(r, 1)
            after = fingerprints(ws)
            lost, gained = sum((before - after).values()), sum((after - before).values())
            if lost != len(rows) or gained:
                die(f"{path.name}!{tab}: integrity broken (lost {lost} != {len(rows)} removed, gained {gained}) — restore backups/before_switch_remove_{sw}_{ts}_{path.name}")
            if any(ws.cell(r, 1).value == sw for r in range(3, ws.max_row + 1)):
                die(f"{path.name}!{tab}: {sw} rows survived the delete")
            removed += len(rows)
        tmp = str(path) + ".rmtmp"
        wb.save(tmp)
        with zipfile.ZipFile(tmp) as z:
            if len(z.namelist()) < 10:
                die(f"refusing truncated save {path.name}")
        os.replace(tmp, str(path))
        return removed
    finally:
        wb.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--switch", required=True)
    ap.add_argument("--templates-dir", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--fleet", action="store_true")
    a = ap.parse_args()
    sw = a.switch.strip().upper()
    for fn, cls in (("config.py", "Config"), ("config_tradier.py", "TradierConfig"), ("v12_quick_engine.py", "QuickConfig")):
        if sw in ast_assigns(ROOT / fn, cls):
            die(f"{sw} is still a {cls} field in {fn} — remove the code first, the template row of a live key is never deleted")
    plan = []
    for cs, fn in FILES.items():
        p = Path(a.templates_dir) / fn
        if not p.exists():
            die(f"missing {p}")
        hits = plan_file(p, sw)
        plan.append((cs, p, hits))
        print(f"[{cs}] {fn}: " + (", ".join(f"{t} rows {r}" for t, r in hits.items()) if hits else "no rows"), flush=True)
    if not a.apply:
        print("dry-run OK — re-run with --apply to delete (backs up each file first)", flush=True)
        return 0
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for cs, p, hits in plan:
        if hits:
            print(f"[{cs}] removed {remove_file(p, sw, hits, ts)} rows -> {p.name}", flush=True)
    if a.fleet:
        for h in FLEET:
            for cs, p, hits in plan:
                subprocess.run(["rsync", "-az", "-e", "ssh -S none -o StrictHostKeyChecking=accept-new", str(p), f"{h}:~/binance-sandbox/SPREADSHEETS/TEMPLATE_FINAL_NORM/{p.name}"], check=False)
            r = subprocess.run(["ssh", "-S", "none", "-o", "StrictHostKeyChecking=accept-new", h, "cd ~/binance-sandbox/SPREADSHEETS/TEMPLATE_FINAL_NORM && md5sum " + " ".join(p.name for _, p, _ in plan)], capture_output=True, text=True)
            print(f"[{h}] {r.stdout.strip()}", flush=True)
    else:
        print("fleet NOT synced — run on the template source host or re-run with --fleet", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
