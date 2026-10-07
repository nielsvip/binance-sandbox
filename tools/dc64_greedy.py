#!/usr/bin/env python
"""dc64_greedy.py — OLD 1-SHEET system, GREEDY + MONOTONIC.

Start from last week's BEST overrides (baseline in E/gain), then walk the switch groups
in order (exits first so TIM drops and trade count rises into 30-300, then entries/filters).
For each group pick the best option vs the CURRENT cumulative; ADD its delta ONLY if
positive (gain NEVER goes down). Write a single-sheet ledger + a REAL_ZOOMABLE chart of the
winning config's trades (bh+gain in title/filename). Honest v12 engine, no fabrication.

Usage:
  python tools/dc64_greedy.py --sym 1000000MOGUSDT_SHORT
  python tools/dc64_greedy.py --all --venue crypto --workers 12
"""
import os, sys, json, argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill
import v12_quick_engine as V
from tools import dc_simple_8_sweep as DC
from tools.opt.evaluate_v12 import _exact_30d_slice

OUT_DIR = os.path.join(ROOT, "SPREADSHEETS", "DC64_GREEDY")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
STOP_BUF, TARGET_BUF = DC.STOP_BUF, DC.TARGET_BUF
TFS = ["OFF", "15m", "1h", "4h"]

# GREEDY ORDER: exits first (raise trade count), then entries, then filters.
def _group_options():
    groups = []
    groups.append(("WT_LOWER_CROSS_EXIT", [(tf, {"WT_LOWER_CROSS_EXIT_TF": tf}) for tf in TFS]))
    dc_exit = [("OFF", {"TECHNICAL_DC_STOP_TF": "OFF", "TECHNICAL_DC_TARGET_TF": "OFF"})]
    for tf in ["15m", "1h", "4h"]:
        dc_exit.append((tf, {"TECHNICAL_DC_STOP_TF": tf, "TECHNICAL_DC_TARGET_TF": tf,
                             "TECHNICAL_DC_STOP_BUFFER_PCT": STOP_BUF, "TECHNICAL_DC_TARGET_BUFFER_PCT": TARGET_BUF}))
    groups.append(("DC_STOP_TARGET", dc_exit))
    groups.append(("DC_ENTRY", [("OFF", {"ENTRY_DC_TF": "OFF"})] +
                   [(tf, {"ENTRY_DC_TF": tf, "ENTRY_DC_BUFFER_PCT": TARGET_BUF}) for tf in ["15m", "1h", "4h"]]))
    groups.append(("WT_DC_DETAILED", [("OFF", {"WT_DC_ENABLED": False})] +
                   [(tf, {"WT_DC_ENABLED": True, "WT_DC_DETAILED_SCORER_ENABLED": True,
                          "WT_DC_TF_ENTRY": tf, "WT_DC_DC_TF": tf}) for tf in ["15m", "1h", "4h"]]))
    groups.append(("BB_SQUEEZE", [("OFF", {"BB_SQUEEZE_ENTRY_ENABLED": False})] +
                   [(tf, {"BB_SQUEEZE_ENTRY_ENABLED": True, "BB_SQUEEZE_ENTRY_TF": tf, "BB_SQUEEZE_EXIT_ENABLED": True}) for tf in ["15m", "1h", "4h"]]))
    groups.append(("EMA_9_21", [("OFF", {"EMA_9_21_FILTER_ENABLED": False})] +
                   [(tf, {"EMA_9_21_FILTER_ENABLED": True, "EMA_9_21_FILTER_FILTER_TF": tf}) for tf in ["1h", "4h", "1h,4h"]]))
    groups.append(("AF_DC", [
        ("OFF", {}),
        ("STOP_15m", {"TECHNICAL_DC_STOP_TF": "15m", "TECHNICAL_DC_TARGET_TF": "OFF", "TECHNICAL_DC_STOP_BUFFER_PCT": STOP_BUF}),
        ("STOP_1h", {"TECHNICAL_DC_STOP_TF": "1h", "TECHNICAL_DC_TARGET_TF": "OFF", "TECHNICAL_DC_STOP_BUFFER_PCT": STOP_BUF}),
        ("TARGET_15m", {"TECHNICAL_DC_STOP_TF": "OFF", "TECHNICAL_DC_TARGET_TF": "15m", "TECHNICAL_DC_TARGET_BUFFER_PCT": TARGET_BUF}),
    ]))
    return groups


def _fmt(x):
    try:
        return f"{float(x):.2f}".replace("-", "m").replace(".", "p")
    except Exception:
        return "NA"


def _base_overrides(sym_side, per_sym_map):
    sym = sym_side[:-5] if sym_side.endswith("_LONG") else sym_side[:-6]
    is_long = sym_side.endswith("_LONG")
    crypto = sym.upper().endswith(("USDT", "USDC", "USD1", "BUSD", "FDUSD", "TUSD"))
    ov = dict(DC._get_template_baseline(crypto, is_long))
    ov["KINDERGARTEN_EMA_GATE_ENABLED"] = True; ov["KINDERGARTEN_FILTER_TF"] = "4h"
    ov["EMA_9_21_FILTER_ENABLED"] = True; ov["EMA_9_21_FILTER_FILTER_TF"] = "4h"
    ov.update({k: v for k, v in per_sym_map.get(sym_side, {}).items() if k != "DAYTRDAY_DC"})
    for k in list(ov.keys()):
        if isinstance(ov[k], str) and ov[k] == "3m":
            ov[k] = "OFF"
    return ov, sym, is_long, crypto


def build_one(sym_side, per_sym_map, window_days=30):
    ov, sym, is_long, crypto = _base_overrides(sym_side, per_sym_map)
    stores = V.load_npz("crypto" if crypto else "tradier", [sym], "2024-01-01")
    npz = stores.get(sym) if isinstance(stores, dict) else stores
    if npz is None or len(npz.get("timestamps", [])) < 100:
        return sym_side, "ERROR npz missing"
    base_r, err = DC.eval_gain(npz, sym, is_long, ov, window_days)
    if base_r is None:
        return sym_side, f"ERROR {err}"
    base_gain = base_r.get("gain_pct_2000norm", 0.0)
    base_trades = base_r.get("trades", 0)
    try:
        sliced, _ = _exact_30d_slice(npz, crypto, window_days); bh = DC.bh_from_sliced(sliced, is_long)
    except Exception:
        sliced, bh = None, None
    cum_ov = dict(ov); cum_gain = base_gain; cum_trades = base_trades
    rows = [("BASE", "-", round(base_gain, 4), 0.0, base_trades, round(base_r.get("tim_pct", 0), 1))]
    for gname, opts in _group_options():
        best = None
        for label, delta_ov in opts:
            test = dict(cum_ov); test.update(delta_ov)
            r, _ = DC.eval_gain(npz, sym, is_long, test, window_days)
            if r is None:
                continue
            g = r.get("gain_pct_2000norm", 0.0)
            if best is None or g > best[1]:
                best = (label, g, r.get("trades", 0), round(r.get("tim_pct", 0), 1), delta_ov)
        if best is None:
            continue
        label, g, tr, tim, delta_ov = best
        delta = g - cum_gain
        promoted = delta > 1e-9
        if promoted:                       # GAIN NEVER GOES DOWN — only add positive
            cum_ov.update(delta_ov); cum_gain = g; cum_trades = tr
        rows.append((f"{gname}={label}", label, round(cum_gain, 4), round(delta, 4), tr, tim))
    applied = {k: cum_ov[k] for k in cum_ov if ov.get(k) != cum_ov[k]}   # exact winning switch settings (diff vs baseline)
    _write(sym_side, rows, bh, base_gain, cum_gain, cum_ov, applied, sym, is_long, crypto, sliced, npz, window_days)
    return sym_side, f"OK bh={round(bh or 0,2)} base_gain={round(base_gain,2)} FINAL_gain={round(cum_gain,2)} (monotonic +{round(cum_gain-base_gain,2)}) trades {base_trades}->{cum_trades}"


def _write(sym_side, rows, bh, base_gain, cum_gain, cum_ov, applied, sym, is_long, crypto, sliced, npz, window_days):
    os.makedirs(OUT_DIR, exist_ok=True)
    stem = f"{sym_side}_bh{_fmt(bh)}_gain{_fmt(cum_gain)}_delta{_fmt(cum_gain-base_gain)}_30d"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "DC64_GREEDY"
    for c, h in enumerate(["switch", "option", "CUM_GAIN", "delta_pp", "trades", "tim_pct"], 1):
        ws.cell(1, c, h).font = Font(bold=True)
    ws.cell(2, 8, f"bh={bh} base_gain={round(base_gain,4)} FINAL_gain={round(cum_gain,4)} monotonic")
    for i, row in enumerate(rows, start=2):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
        d = row[3]
        try:
            ws.cell(i, 4).fill = GREEN if float(d) > 1e-9 else (RED if float(d) < -1e-9 else PatternFill())
        except Exception:
            pass
    xlsx = os.path.join(OUT_DIR, stem + "_matrix.xlsx")
    wb.save(xlsx)
    # chart of winning cumulative config
    try:
        if sliced is None:
            sliced, _ = _exact_30d_slice(npz, crypto, window_days)
        cfg = V.QuickConfig(); cfg.MODE = "crypto" if crypto else "tradier"
        for k, v in cum_ov.items():
            try: setattr(cfg, k, v)
            except Exception: pass
        sim = V.simulate_one(sliced, sym, is_long, cfg) or {}
        ts = sliced.get("timestamps", sliced.get("timestamp_3m", np.array([])))
        metrics = {"trades": sim.get("trades", 0), "sharpe": sim.get("sharpe_per_trade", 0.0),
                   "dd": sim.get("max_dd_pct", 0.0), "tim": sim.get("tim_pct", 0.0), "wr": sim.get("wr", 0.0),
                   "bars": len(sliced.get("close", []))}
        _chart(sym_side, sliced.get("close", np.array([])), ts, sim.get("ledger", []), bh, base_gain, cum_gain, applied, metrics,
               os.path.join(OUT_DIR, f"{sym_side}_bh{_fmt(bh)}_gain{_fmt(cum_gain)}_30D_REAL_ZOOMABLE.html"))
        # dump winning config for trivial rich-chart regen
        with open(os.path.join(OUT_DIR, f"{sym_side}_winning.json"), "w") as jf:
            json.dump({"sym_side": sym_side, "bh": bh, "base_gain": base_gain, "final_gain": cum_gain,
                       "applied_switches": {k: (v if isinstance(v, (int, float, str, bool)) else str(v)) for k, v in applied.items()}}, jf, indent=1)
    except Exception as e:
        print(f"[chart-fail] {sym_side} {e}", flush=True)


def _chart(sym_side, close, ts, ledger, bh, base_gain, gain, applied, metrics, path):
    """ZECUSDC/hires-style: metrics bar + full-width chart, X = bar number with UTC date
    tick labels + tooltip 'bar N — date', trade table with entry/exit bar AND date columns,
    winning switch settings shown in the metrics bar."""
    n = len(close); step = max(1, n // 4000); xs = list(range(0, n, step))
    ys = [round(float(close[i]), 8) for i in xs]
    ts = [int(float(x)) for x in (ts if ts is not None else [])]
    TS = ts[::step] if len(ts) >= n else ts   # downsampled ts aligned to xs for tick labels
    def _d(bar):
        if not ts:
            return ""
        i = min(len(ts) - 1, max(0, int(bar)))
        import datetime as _dt
        return _dt.datetime.utcfromtimestamp(ts[i]).strftime("%m-%d %H:%M")
    pts = [{"be": int(t.get("bar_entry", 0)), "ep": float(t.get("entry_price", 0)), "bx": int(t.get("bar_exit", 0)),
            "xp": float(t.get("exit_price", 0)), "win": float(t.get("pnl_pct", 0)) >= 0,
            "ty": str(t.get("type", "CLOSE")), "pnl": round(float(t.get("pnl_pct", 0)), 3), "qty": round(float(t.get("qty", 0)), 4),
            "ed": _d(t.get("bar_entry", 0)), "xd": _d(t.get("bar_exit", 0)),
            "er": str(t.get("entry_reason", ""))[:44], "xr": str(t.get("exit_reason", t.get("reason", "")))[:56]} for t in ledger]
    settings = " &middot; ".join(f"{k}=<b>{v}</b>" for k, v in sorted(applied.items())) or "(none promoted — baseline unchanged)"
    imp = (gain or 0) - (base_gain or 0)
    tmpl = """<!doctype html><html><head><meta charset=utf-8><title>@@SYM@@ hires</title>
<script src='https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js'></script>
<script src='https://cdn.jsdelivr.net/npm/hammerjs@2.0.8/hammer.min.js'></script>
<script src='https://cdn.jsdelivr.net/npm/chartjs-plugin-zoom@2.0.1/dist/chartjs-plugin-zoom.min.js'></script>
<style>body{font-family:system-ui,monospace;margin:12px;background:#0a0a0a;color:#e5e5e5}
.metrics{background:#1a1a1a;border:1px solid #333;padding:8px;font-size:12px}
#wrap{height:60vh;border:1px solid #333}
table{border-collapse:collapse;width:100%;font-size:10px;margin-top:8px}
th{background:#1f2937;color:#fff;padding:4px;position:sticky;top:0}td{border:1px solid #333;padding:3px}</style></head>
<body><h2>@@SYM@@ — 30D greedy (exits→entries→filters), gain never decreases</h2>
<div class="metrics"><b>@@SYM@@ — FINAL</b> &nbsp;|&nbsp; prev(base) @@BASE@@% &rarr; <b>NEW @@GAIN@@%</b> (Δ +@@IMP@@) &nbsp; BH @@BH@@% &nbsp;|&nbsp; trades @@NTR@@ &nbsp; sharpe @@SH@@ &nbsp; DD @@DD@@% &nbsp; TIM @@TIM@@% &nbsp; WR @@WR@@% &nbsp; window @@BARS@@ bars<br>
<span style="font-size:11px;color:#4a9">WINNING SWITCHES:</span> @@SETTINGS@@<br>
<span style="font-size:10px;opacity:.8">X-axis = bar number, tick label = UTC date &middot; wheel zoom &middot; drag pan &middot; dblclick reset &middot; hover any marker for bar/date/reason</span></div>
<div id='wrap'><canvas id='c'></canvas></div>
<table id=tt><tr><th>#</th><th>type</th><th>bar_in</th><th>date_in</th><th>bar_out</th><th>date_out</th><th>entry$</th><th>exit$</th><th>pnl%</th><th>qty</th><th>entry_reason</th><th>exit_reason</th></tr></table>
<script>
const XS=@@XS@@, YS=@@YS@@, TS=@@TS@@, ST=@@STEP@@, PT=@@PT@@;
function dstr(u){if(!u)return '';const d=new Date(u*1000);return ('0'+(d.getUTCMonth()+1)).slice(-2)+'-'+('0'+d.getUTCDate()).slice(-2)+' '+('0'+d.getUTCHours()).slice(-2)+':'+('0'+d.getUTCMinutes()).slice(-2);}
function barDate(b){const i=Math.min(TS.length-1,Math.max(0,Math.round(b/ST)));return dstr(TS[i]);}
const price=XS.map((x,i)=>({x:x,y:YS[i]}));
new Chart(document.getElementById('c'),{type:'line',data:{datasets:[
{label:'price',data:price,borderColor:'rgba(155,155,155,0.95)',borderWidth:0.9,pointRadius:0,order:0,tension:0},
{label:'entry',data:PT.map(p=>({x:p.be,y:p.ep})),type:'scatter',backgroundColor:'rgba(22,163,74,1)',pointStyle:'circle',radius:5,order:5},
{label:'exit gain',data:PT.filter(p=>p.win).map(p=>({x:p.bx,y:p.xp})),type:'scatter',backgroundColor:'rgba(22,163,74,1)',pointStyle:'rect',radius:5,order:5},
{label:'exit loss',data:PT.filter(p=>!p.win).map(p=>({x:p.bx,y:p.xp})),type:'scatter',backgroundColor:'rgba(220,38,38,1)',pointStyle:'rectRot',radius:5,order:5},
{label:'augment',data:PT.filter(p=>p.ty=='AUGMENT').map(p=>({x:p.be,y:p.ep})),type:'scatter',backgroundColor:'rgba(37,99,235,1)',pointStyle:'triangle',radius:5,order:6},
{label:'reduce',data:PT.filter(p=>p.ty=='REDUCE').map(p=>({x:p.bx,y:p.xp})),type:'scatter',backgroundColor:'rgba(245,158,11,1)',pointStyle:'triangle',rotation:180,radius:5,order:6}]},
options:{parsing:false,animation:false,maintainAspectRatio:false,
scales:{x:{type:'linear',title:{display:true,text:'bar number (tick = UTC date)',color:'#e5e5e5'},ticks:{color:'#e5e5e5',maxRotation:45,minRotation:25,maxTicksLimit:14,callback:(v)=>Math.round(v)+'  '+barDate(v)}},
y:{title:{display:true,text:'price',color:'#e5e5e5'},ticks:{color:'#e5e5e5'}}},
plugins:{legend:{labels:{color:'#e5e5e5'}},zoom:{zoom:{wheel:{enabled:true},pinch:{enabled:true},mode:'x'},pan:{enabled:true,mode:'x'}},
tooltip:{callbacks:{title:(it)=>{const x=it[0].parsed.x;return 'bar '+Math.round(x)+' — '+barDate(x);},
afterBody:(it)=>{const d=it[0].dataset.label;let a=null,i=it[0].dataIndex;if(d=='entry')a=PT[i];else if(d=='exit gain')a=PT.filter(p=>p.win)[i];else if(d=='exit loss')a=PT.filter(p=>!p.win)[i];else if(d=='augment')a=PT.filter(p=>p.ty=='AUGMENT')[i];else if(d=='reduce')a=PT.filter(p=>p.ty=='REDUCE')[i];return a?['type '+a.ty,'pnl '+a.pnl+'%  qty '+a.qty,'entry: '+a.er,'exit: '+a.xr]:'';}}}}}});
document.getElementById('c').ondblclick=()=>{Chart.getChart('c').resetZoom();};
const tt=document.getElementById('tt');PT.forEach((p,i)=>{const r=tt.insertRow();r.innerHTML='<td>'+(i+1)+'</td><td>'+p.ty+'</td><td>'+p.be+'</td><td>'+p.ed+'</td><td>'+p.bx+'</td><td>'+p.xd+'</td><td>'+p.ep+'</td><td>'+p.xp+'</td><td style="color:'+(p.pnl>=0?'#16a34a':'#dc2626')+'">'+p.pnl+'</td><td>'+p.qty+'</td><td>'+p.er+'</td><td>'+p.xr+'</td>';});
</script></body></html>"""
    html = (tmpl.replace("@@SYM@@", str(sym_side)).replace("@@BH@@", f"{(bh or 0):.2f}")
            .replace("@@BASE@@", f"{(base_gain or 0):.2f}").replace("@@GAIN@@", f"{(gain or 0):.2f}")
            .replace("@@IMP@@", f"{imp:.2f}").replace("@@NTR@@", str(len(ledger)))
            .replace("@@SH@@", f"{metrics.get('sharpe',0):.3f}").replace("@@DD@@", f"{metrics.get('dd',0):.1f}")
            .replace("@@TIM@@", f"{metrics.get('tim',0):.1f}").replace("@@WR@@", f"{metrics.get('wr',0):.1f}")
            .replace("@@BARS@@", str(metrics.get('bars', n))).replace("@@SETTINGS@@", settings).replace("@@STEP@@", str(step))
            .replace("@@XS@@", json.dumps(xs)).replace("@@YS@@", json.dumps(ys)).replace("@@TS@@", json.dumps(TS)).replace("@@PT@@", json.dumps(pts)))
    with open(path, "w") as f:
        f.write(html)


def regen_rich(xlsx_path, per_sym_map, window_days=30):
    """Rebuild the RICH chart (settings panel + entry/exit reasons) from a finished ledger
    xlsx, reconstructing the winning config from its green (delta>0) rows. No re-sweep."""
    import openpyxl
    base = os.path.basename(xlsx_path)
    sym_side = base.split("_bh")[0]
    ov, sym, is_long, crypto = _base_overrides(sym_side, per_sym_map)
    lut = {gname: dict(opts) for gname, opts in _group_options()}
    wb = openpyxl.load_workbook(xlsx_path); ws = wb.active
    cum_ov = dict(ov); base_gain = None; cum_gain = None; bh = None
    for r in range(2, ws.max_row + 1):
        sw = ws.cell(r, 1).value; cg = ws.cell(r, 3).value; d = ws.cell(r, 4).value
        if sw == "BASE":
            base_gain = cg; continue
        if cg is not None:
            cum_gain = cg
        try:
            if sw and "=" in str(sw) and d is not None and float(d) > 1e-9:
                gname, label = str(sw).split("=", 1)
                if gname in lut and label in lut[gname]:
                    cum_ov.update(lut[gname][label])
        except Exception:
            pass
    b8 = ws.cell(2, 8).value or ""
    try:
        bh = float(str(b8).split("bh=")[1].split()[0])
    except Exception:
        bh = None
    stores = V.load_npz("crypto" if crypto else "tradier", [sym], "2024-01-01")
    npz = stores.get(sym)
    if npz is None:
        return sym_side, "npz missing"
    sliced, _ = _exact_30d_slice(npz, crypto, window_days)
    cfg = V.QuickConfig(); cfg.MODE = "crypto" if crypto else "tradier"
    for k, v in cum_ov.items():
        try: setattr(cfg, k, v)
        except Exception: pass
    sim = V.simulate_one(sliced, sym, is_long, cfg) or {}
    applied = {k: cum_ov[k] for k in cum_ov if ov.get(k) != cum_ov[k]}
    ts = sliced.get("timestamps", sliced.get("timestamp_3m", np.array([])))
    metrics = {"trades": sim.get("trades", 0), "sharpe": sim.get("sharpe_per_trade", 0.0), "dd": sim.get("max_dd_pct", 0.0),
               "tim": sim.get("tim_pct", 0.0), "wr": sim.get("wr", 0.0), "bars": len(sliced.get("close", []))}
    html = os.path.join(OUT_DIR, f"{sym_side}_bh{_fmt(bh)}_gain{_fmt(cum_gain)}_30D_REAL_ZOOMABLE.html")
    _chart(sym_side, sliced.get("close", np.array([])), ts, sim.get("ledger", []), bh,
           base_gain if base_gain is not None else 0.0, cum_gain if cum_gain is not None else 0.0, applied, metrics, html)
    return sym_side, f"regen OK settings={len(applied)} trades={len(sim.get('ledger', []))}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sym", default=""); ap.add_argument("--all", action="store_true")
    ap.add_argument("--regen", action="store_true", help="rebuild rich charts from finished xlsx in OUT_DIR")
    ap.add_argument("--venue", choices=["crypto", "stocks", "both"], default="both")
    ap.add_argument("--window", type=int, default=30); ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()
    pm, _, _ = DC.load_per_sym_maps()
    if a.regen:
        import glob as _g
        xs = sorted(_g.glob(os.path.join(OUT_DIR, "*_matrix.xlsx")))
        print(f"[regen] {len(xs)} sheets -> rich charts", flush=True)
        done = 0
        with ThreadPoolExecutor(max_workers=a.workers) as ex:
            futs = {ex.submit(regen_rich, x, pm, a.window): x for x in xs}
            for f in as_completed(futs):
                try:
                    ss, msg = f.result()
                except Exception as e:
                    ss, msg = os.path.basename(futs[f]), f"EXC {e}"
                done += 1
                if done % 20 == 0 or done == len(xs):
                    print(f"[regen {done}/{len(xs)}] {ss}: {msg}", flush=True)
        print(f"[regen] DONE {done}/{len(xs)}", flush=True)
        return
    if a.sym:
        targets = [a.sym]
    else:
        keys = [k for k in pm if not k.startswith("_")]
        def is_c(k):
            b = k[:-5] if k.endswith("_LONG") else k[:-6]
            return b.upper().endswith(("USDT", "USDC", "USD1", "BUSD", "FDUSD", "TUSD"))
        targets = sorted([k for k in keys if (a.venue == "both") or (is_c(k) == (a.venue == "crypto"))])
        if a.limit:
            targets = targets[:a.limit]
    print(f"[dc64_greedy] {len(targets)} sym_sides venue={a.venue} workers={a.workers} out={OUT_DIR}", flush=True)
    done = 0
    with ThreadPoolExecutor(max_workers=a.workers) as ex:
        futs = {ex.submit(build_one, ss, pm, a.window): ss for ss in targets}
        for f in as_completed(futs):
            ss = futs[f]
            try:
                _, msg = f.result()
            except Exception as e:
                msg = f"EXC {e}"
            done += 1
            print(f"[{done}/{len(targets)}] {ss}: {msg}", flush=True)
    print(f"[dc64_greedy] DONE {done}/{len(targets)} -> {OUT_DIR}", flush=True)


if __name__ == "__main__":
    main()
