#!/usr/bin/env python3
"""Un-grey (column-A font colour only) every row of the NAMED switches in TEMPLATE_*.xlsx workbooks.

USER 2026-10-27 ~20:50Z: "NONE of the 8 rows you mention here should have EVER Been greyed out make them black
again". Grey col-A font (FFBFBFBF) = v15_pilot skip-row signal (correct behaviour, kept). This tool only removes
that colour from rows whose switch is listed with --switch; values, defaults, bold markers, yellow headers and all
other rows are untouched. Edits the given (current) workbooks in place — never copies an older workbook over a newer
one. Backs up each file, zip-validates the rewrite before replacing. Dry-run unless --apply.
"""
import argparse
import datetime
import pathlib
import shutil
import sys
import zipfile

GREY_SKIP_RGB = {"FFBFBFBF", "00BFBFBF"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--switch", action="append", required=True)
    ap.add_argument("--dir", required=True)
    ap.add_argument("--backup-dir", required=True)
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    import openpyxl
    from openpyxl.styles import Font
    names = {s.strip() for s in a.switch}
    stamp = datetime.datetime.utcnow().strftime("%Y%m%d%H%M")
    for tf in sorted(pathlib.Path(a.dir).glob("TEMPLATE_*_*.xlsx")):
        wb = openpyxl.load_workbook(str(tf))
        changed = []
        for ws in wb.worksheets:
            for rr in range(3, ws.max_row + 1):
                cell = ws.cell(row=rr, column=1)
                if cell.value is None or str(cell.value).strip() not in names:
                    continue
                font = cell.font
                rgb = str(getattr(font.color, "rgb", "") or "").upper() if (font is not None and font.color is not None) else ""
                if rgb not in GREY_SKIP_RGB:
                    continue
                changed.append(f"{ws.title}!{rr} {cell.value}={ws.cell(row=rr, column=2).value}")
                if a.apply:
                    cell.font = Font(name=font.name, size=font.size, bold=font.bold, italic=font.italic, underline=font.underline, strike=font.strike, vertAlign=font.vertAlign, color=None)
        print(f"{tf.name}: {len(changed)} grey rows on named switches")
        for line in changed:
            print("   ", line)
        if not a.apply or not changed:
            wb.close()
            continue
        bdir = pathlib.Path(a.backup_dir)
        bdir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(tf, bdir / f"before_ungrey_8_switches_{stamp}_{tf.name}")
        tmp = tf.with_name(f".{tf.name}.ungrey_tmp")
        wb.save(str(tmp))
        wb.close()
        with zipfile.ZipFile(tmp) as zf:
            n, bad = len(zf.namelist()), zf.testzip()
        if n < 10 or bad is not None:
            tmp.unlink()
            print(f"    REFUSED {tf.name}: rewritten zip invalid (entries={n} bad={bad}); original untouched")
            continue
        tmp.replace(tf)
        print(f"    WROTE {tf.name} entries={n}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
