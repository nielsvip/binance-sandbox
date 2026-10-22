#!/usr/bin/env python3
"""Grey out template switch rows so the pilot never tests them again.

USER 2026-10-22: "grey them out as they have no business in the TEMPLATE_* files the positions sizes and order
values are decided in market_mode (NOTHING TO DO WITH vectorized testing) and in execute_trade_action etcetera
NOT in templates EVER, hedge_mode is OFF and not to be tested, MAX_LOSS is a LIE grey out dont ever test,
min_size is also bullshit dont test, CAT_SIDE_DEFAULTS_ENABLED is a JOKE they are the only thing that can trade if
no per_sym settings exits so DELETE".

THE MARKER IS THE PILOT'S OWN. `v15_pilot.py:110` defines `GREY_SKIP_RGB = {"FFBFBFBF", "00BFBFBF"}` and checks it
at `:1877` against the FONT COLOUR of the switch-name cell in column A — a grey switch name is a skip row and is
never evaluated. So greying is not cosmetic here, it is the supported "do not test this" instruction, and it needs
no engine change. `tools/v15_template_fix.py:183` holds the same constant but has no CLI for it, which is why this
exists.

WHY THESE SWITCHES. A board can only be reproduced live by a knob live actually resolves from the board. Position
sizes and order values are not that: `per_sym_live_set._apply_caps` clamps START_POSITION_SIZE,
MAX_POSITION_SIZE, MAX_ORDER_VALUE, MAX_ORDER_VALUE_FIN and MAX_ORDER_VALUE_MEN to the active LIGHT/NORMAL/EXTREME
mode table on EVERY read, so the mode table outranks whatever a board swept. Measured 2026-10-20: feeding the
capped values back as a board set made `live_twins.vec_exact.latest_board_ok` report
`LIVE_SET_NE_BOARD_1keys:START_POSITION_SIZE` and refuse EVERY entry on 25 of 60 crypto sym_sides. Sweeping a knob
the mode table overrides spends pilot cells to measure a delta live can never honour.

WHAT IS *NOT* GREYED, deliberately: the five FIXED_STOP_* reference switches and DELTA_ENGINE_ENABLED. They were
in the same audit bucket but they decide WHERE A STOP FREEZES / whether an engine runs — exit-bar decisions, not
quantity — and LOCKED_FILES.md 2026-10-20 made the FIXED_STOP reference a deliberately tested switch family.

ROW DELETION IS NOT DONE, and that is a judgement, not an oversight. Deleting rows renumbers every row below,
and the workbooks carry row-relative formulas (`=IF(G4="",E3,IF(G4>0,E3+G4,E3))` greedy cumulative in E, VLOOKUP
in G) plus the L:BI yellow column map the pilot indexes by position. A shift corrupts all four workbooks for the
whole fleet. Greying + clearing the candidate achieves the operative goal — the switch can never be tested or
promoted — and is reversible. CAT_SIDE_DEFAULTS_ENABLED is therefore greyed AND candidate-cleared rather than
row-deleted; say the word and physical removal can go through the S1 queue tool with a formula re-anchor.

Saves atomically and refuses to replace a workbook whose rewritten zip has fewer than 10 entries (CLAUDE.md:
truncated 225 KB BadZip saves).
"""
from __future__ import annotations

import argparse
import collections
import glob
import json
import os
import pathlib
import shutil
import sys
import time
import zipfile

ROOT = pathlib.Path(os.environ.get("BASE_PATH") or pathlib.Path(__file__).resolve().parents[1])
GREY = "FFBFBFBF"
SKIP_TABS = {"PROTOCOL", "LEGEND_FILTERS", "INSTRUCTIONS", "INSTRUCTIONS_V2", "FILTERS_EXPLAINED",
             "FILTER_DICTIONARY_V2", "Results_Deltas", "Results_30d_Deltas"}

# operator's list, 2026-10-22
GREY_SWITCHES = [
    "START_POSITION_SIZE", "MIN_POSITION_SIZE", "MAX_POSITION_SIZE",
    "MAX_ORDER_VALUE", "MAX_ORDER_VALUE_FIN", "MAX_ORDER_VALUE_MEN",
    "HIGH_GAIN_AUGMENTATION_MIN_SIZE",
    "HEDGE_MODE",
    "NEW_POSITION_MAX_LOSS_THRESHOLD",
    "CAT_SIDE_DEFAULTS_ENABLED",
]
CLEAR_CANDIDATE = {"CAT_SIDE_DEFAULTS_ENABLED"}   # operator said DELETE: also blank the candidate so no row can win


def atomic_save(wb, path: pathlib.Path) -> bool:
    tmp = path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    try:
        with zipfile.ZipFile(tmp) as z:
            n = len(z.namelist())
    except Exception as exc:
        print(f"    REFUSING: rewritten {path.name} is not a readable zip ({exc})")
        tmp.unlink(missing_ok=True)
        return False
    if n < 10:
        print(f"    REFUSING: rewritten {path.name} has only {n} zip entries")
        tmp.unlink(missing_ok=True)
        return False
    os.replace(tmp, path)
    return True


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--templates", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_*.xlsx"))
    ap.add_argument("--json", default=str(ROOT / "data" / "reports" / "template_grey_switch.json"))
    a = ap.parse_args()

    import openpyxl
    from openpyxl.styles import Font

    paths = sorted(glob.glob(a.templates))
    if not paths:
        print(f"no templates matched {a.templates}")
        return 2
    want = {s.upper() for s in GREY_SWITCHES}
    report, tot = {}, collections.Counter()

    for p in paths:
        path = pathlib.Path(p)
        wb = openpyxl.load_workbook(path)
        found, greyed, cleared, already = [], 0, 0, 0
        for ws in wb.worksheets:
            if ws.title in SKIP_TABS or "BASELINE" in ws.title:
                continue
            for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=3):
                c = row[0]
                if c.value is None:
                    continue
                key = str(c.value).strip().upper()
                if key not in want:
                    continue
                found.append(f"{ws.title}:{c.row}:{key}")
                cur = str(getattr(getattr(c.font, "color", None), "rgb", "") or "").upper()
                if cur == GREY:
                    already += 1
                else:
                    f = c.font
                    c.font = Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic, color=GREY)
                    greyed += 1
                if key in CLEAR_CANDIDATE and len(row) > 1 and row[1].value is not None:
                    row[1].value = None
                    cleared += 1
        report[path.name] = {"rows_found": len(found), "greyed": greyed, "already_grey": already,
                             "candidates_cleared": cleared, "sample": found[:8]}
        tot["rows"] += len(found)
        tot["greyed"] += greyed
        tot["already"] += already
        tot["cleared"] += cleared
        print(f"  {path.name:30s} rows={len(found):3d}  to_grey={greyed:3d}  already_grey={already:3d}  candidates_cleared={cleared}")
        if a.apply and (greyed or cleared):
            bk = ROOT / "backups" / f"{path.stem}_before_grey_{time.strftime('%Y%m%d%H%M', time.gmtime())}.xlsx"
            bk.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, bk)
            if atomic_save(wb, path):
                print(f"    saved; backup {bk.name}")
            else:
                return 2
        wb.close()

    pathlib.Path(a.json).parent.mkdir(parents=True, exist_ok=True)
    pathlib.Path(a.json).write_text(json.dumps(
        {"ts": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "grey_switches": sorted(want),
         "marker": GREY, "applied": bool(a.apply), "totals": dict(tot), "per_template": report}, indent=1))
    print(f"\n  TOTAL rows matched={tot['rows']} greyed={tot['greyed']} already={tot['already']} cleared={tot['cleared']}")
    print(f"  wrote {a.json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
