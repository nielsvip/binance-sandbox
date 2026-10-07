#!/usr/bin/env python3
"""Write the pilot's pos_sym sampler input (data/avg_delta_pos_sym.json) straight from the avg-delta aggregate
(SPREADSHEETS/v15_avg_delta_latest.xlsx: tab,name,kind,pos_sym,avg_delta,median_delta,n) — key 'TAB!NAME' with pos_sym/n_sym/avg_delta,
the exact lookup v15_pilot._ps_row_ev / _ps_filter_ev use. Director 2026-10-06 (USER: apply pos_sym sampling now)."""
import datetime, json, sys
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
agg = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "SPREADSHEETS" / "v15_avg_delta_latest.xlsx"
out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data" / "avg_delta_pos_sym.json"
wb = openpyxl.load_workbook(str(agg), read_only=True)
res = {"_meta": {"at": datetime.datetime.utcnow().isoformat() + "Z", "source": str(agg)}, "cat_sides": {}}
for cs in wb.sheetnames:
    d = {}
    for tab, name, kind, pos, avg, med, n in wb[cs].iter_rows(min_row=2, values_only=True):
        if not tab or not name:
            continue
        d[f"{tab}!{name}"] = {"pos_sym": int(pos) if isinstance(pos, (int, float)) else None, "n_sym": int(n) if isinstance(n, (int, float)) else None,
                              "avg_delta": avg if isinstance(avg, (int, float)) else None, "kind": kind}
    res["cat_sides"][cs] = d
    v = [x["pos_sym"] for x in d.values() if x["n_sym"] is not None and x["n_sym"] >= 20]
    print(cs, "keys", len(d), "with n>=20:", len(v), {k: sum(1 for x in v if (x if x < 4 else 4) == k) for k in range(5)}, "(4 = 4+)")
tmp = out.with_suffix(".tmp")
tmp.write_text(json.dumps(res))
tmp.replace(out)
print("->", out)
