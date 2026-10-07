#!/usr/bin/env python3
"""Build the switch + filter encyclopedia from the four TEMPLATE_FINAL_NORM workbooks.

Reads only: SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_*.xlsx, config.py, config_tradier.py.
Writes: SWITCH_ENCYCLOPEDIA.md (human) and data/reports/switch_encyclopedia.json (machine).
Nothing here computes a metric; it only inventories what the templates and code declare.
"""
import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"
WORKBOOKS = {
    "CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx",
    "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx",
    "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx",
    "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx",
}
SWITCH_TABS = [
    "STDEV_SLOPE_SIZING",
    "ENTRY_REVERSAL_BOUNCE",
    "ENTRY_BREAKOUT_CHANNEL",
    "ENTRY_CONFIRMATION_GATES",
    "EXIT_STRUCTURAL",
    "EXIT_VELOCITY",
    "REENTRY_WINDOWED",
    "REENTRY_ADAPTIVE",
    "AUGMENT_TREND",
    "AUGMENT_RISK_SIZING",
    "REDUCE_PROFIT_LOCK",
    "REDUCE_SIGNAL_RATER",
    "GLOBAL_RISK_GATES",
]
CONFIG_FILES = ["config.py", "config_tradier.py"]
OUT_MD = ROOT / "SWITCH_ENCYCLOPEDIA.md"
OUT_JSON = ROOT / "data" / "reports" / "switch_encyclopedia.json"
DICT_TAB = "FILTER_DICTIONARY_V2"
ASSIGN_RE = re.compile(r"^\s*([A-Z][A-Z0-9_]+)\s*=\s*(.*?)\s*#\s*(\S.*)$")
DICT_RE = re.compile(r"^\s*['\"]([A-Z][A-Z0-9_]+)['\"]\s*:\s*(.*?),?\s*#\s*(\S.*)$")
YELLOW_HEADER_ROW = 2
YELLOW_FIRST_COL = 12


def text(value):
    return "" if value is None else str(value).strip()


def load_filter_dictionary():
    filters = {}
    for venue, fname in WORKBOOKS.items():
        wb = openpyxl.load_workbook(TEMPLATE_DIR / fname, read_only=True, data_only=True)
        for row in wb[DICT_TAB].iter_rows(min_row=2, values_only=True):
            name = text(row[1])
            if not name:
                continue
            entry = filters.setdefault(
                name,
                {
                    "options": [],
                    "what_it_does": text(row[3]),
                    "live_location": text(row[4]),
                    "sheets": text(row[5]),
                    "gates_switches": set(),
                    "recommendation": text(row[7]),
                    "status": text(row[8]),
                    "venues": set(),
                },
            )
            option = text(row[2])
            if option and option not in entry["options"]:
                entry["options"].append(option)
            for token in re.split(r"[,\s]+", text(row[6])):
                if token:
                    entry["gates_switches"].add(token)
            entry["venues"].add(venue)
        wb.close()
    return filters


def load_switch_rows():
    """Return ({switch: {venue: {tab: {...}}}}, {venue: {tab: [yellow filter names]}})."""
    rows = defaultdict(lambda: defaultdict(lambda: defaultdict(dict)))
    yellow = defaultdict(dict)
    for venue, fname in WORKBOOKS.items():
        wb = openpyxl.load_workbook(TEMPLATE_DIR / fname, read_only=True, data_only=True)
        for tab in SWITCH_TABS:
            if tab not in wb.sheetnames:
                continue
            header = next(wb[tab].iter_rows(min_row=YELLOW_HEADER_ROW, max_row=YELLOW_HEADER_ROW, values_only=True))
            names = []
            for cell in header[YELLOW_FIRST_COL - 1:]:
                token = text(cell)
                if "=" in token:
                    filter_name = token.split("=", 1)[0].strip()
                    if filter_name and filter_name not in names:
                        names.append(filter_name)
            yellow[venue][tab] = names
            for row in wb[tab].iter_rows(min_row=3, max_col=4, values_only=True):
                name = text(row[0])
                if not name or name.lower() == "switch":
                    continue
                slot = rows[name][venue][tab]
                slot.setdefault("variants", [])
                value = text(row[1])
                if value and value not in slot["variants"]:
                    slot["variants"].append(value)
                if "default" not in slot:
                    slot["default"] = value
                slot["families"] = text(row[3])
        wb.close()
    return rows, yellow


def load_config_notes():
    notes = {}
    for fname in CONFIG_FILES:
        path = ROOT / fname
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            match = ASSIGN_RE.match(line) or DICT_RE.match(line)
            if match and match.group(1) not in notes:
                notes[match.group(1)] = match.group(3).strip()
    return notes


def build():
    filters = load_filter_dictionary()
    switch_rows, yellow = load_switch_rows()
    notes = load_config_notes()

    switches = {}
    for name, venues in sorted(switch_rows.items()):
        if name in filters:
            continue
        gating = sorted(
            {
                filter_name
                for venue, tabs in venues.items()
                for tab in tabs
                for filter_name in yellow.get(venue, {}).get(tab, [])
            }
        )
        switches[name] = {
            "kind": "switch",
            "tabs": sorted({tab for v in venues.values() for tab in v}),
            "venues": {
                venue: {tab: dict(slot) for tab, slot in tabs.items()}
                for venue, tabs in venues.items()
            },
            "config_note": notes.get(name, ""),
            "yellow_filters_on_its_tabs": gating,
        }

    filter_out = {}
    for name, meta in sorted(filters.items()):
        filter_out[name] = {
            "kind": "filter",
            "options": meta["options"],
            "what_it_does": meta["what_it_does"],
            "live_location": meta["live_location"],
            "sheets": meta["sheets"],
            "gates_switches": sorted(meta["gates_switches"]),
            "recommendation": meta["recommendation"],
            "status": meta["status"],
            "venues": sorted(meta["venues"]),
        }

    return {"switches": switches, "filters": filter_out}


def coverage(inventory):
    sw = inventory["switches"]
    fl = inventory["filters"]
    return {
        "switch_names": len(sw),
        "switch_names_with_config_note": sum(1 for s in sw.values() if s["config_note"]),
        "switch_names_with_no_yellow_filter": sum(1 for s in sw.values() if not s["yellow_filters_on_its_tabs"]),
        "filter_names": len(fl),
        "filter_names_with_semantics": sum(1 for f in fl.values() if f["what_it_does"]),
        "filter_names_with_status": sum(1 for f in fl.values() if f["status"]),
    }


def md_escape(value):
    return text(value).replace("|", "\\|").replace("\n", " ")


def write_markdown(inventory, stats):
    lines = [
        "# SWITCH ENCYCLOPEDIA (generated by tools/build_switch_encyclopedia.py — do not hand-edit)",
        "",
        "Inventory built from the four `SPREADSHEETS/TEMPLATE_FINAL_NORM` workbooks and `config.py` / `config_tradier.py`.",
        "Wiring truth (live read + vector read) lives in `SWITCH_BIBLE.md`; this file lists what the templates declare.",
        "Where a column is blank, the source does not say — it is not inferred here.",
        "",
        "## Coverage",
        "",
    ]
    lines += [f"- {key}: {value}" for key, value in stats.items()]
    lines += ["", "## Filters (FILTER_DICTIONARY_V2)", "",
              "| Filter | Options | What it does | Live location | Sheets | Gates switches | Recommendation | Status |",
              "|---|---|---|---|---|---|---|---|"]
    for name, f in inventory["filters"].items():
        lines.append(
            f"| {name} | {md_escape(', '.join(f['options']))} | {md_escape(f['what_it_does'])} | "
            f"{md_escape(f['live_location'])} | {md_escape(f['sheets'])} | "
            f"{md_escape(', '.join(f['gates_switches']))} | {md_escape(f['recommendation'])} | {md_escape(f['status'])} |"
        )
    lines += ["", "## Switches by tab", ""]
    by_tab = defaultdict(list)
    for name, s in inventory["switches"].items():
        for tab in s["tabs"]:
            by_tab[tab].append((name, s))
    for tab in SWITCH_TABS:
        entries = by_tab.get(tab, [])
        lines += [f"### {tab} ({len(entries)})", "",
                  "| Switch | Default by venue | Variants (CRYPTO_LONG) | Config note | Yellow filters on its tab(s) |",
                  "|---|---|---|---|---|"]
        for name, s in sorted(entries):
            defaults = []
            for venue, tabs in s["venues"].items():
                if tab in tabs:
                    defaults.append(f"{venue}={tabs[tab].get('default', '')}")
            long_slot = s["venues"].get("CRYPTO_LONG", {}).get(tab, {})
            variants = ", ".join(long_slot.get("variants", []))
            lines.append(
                f"| {name} | {md_escape('; '.join(defaults))} | {md_escape(variants)} | "
                f"{md_escape(s['config_note'])} | {md_escape(', '.join(s['yellow_filters_on_its_tabs']))} |"
            )
        lines.append("")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main():
    inventory = build()
    stats = coverage(inventory)
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps({"coverage": stats, **inventory}, indent=1, sort_keys=True), encoding="utf-8")
    write_markdown(inventory, stats)
    print(json.dumps(stats, indent=1))


if __name__ == "__main__":
    main()
