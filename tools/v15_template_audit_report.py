#!/usr/bin/env python3
"""v15_template_audit_report — write data/template_audit/<ts>/report.md from summary.json + rows_*.csv + switches_*.csv (+ live vs staged tab counts).
python tools/v15_template_audit_report.py --ts <ts>"""
import argparse
import collections
import csv
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import v15_template_home_rules as H  # noqa: E402

CS = ["CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"]
NAMED = "WT_MOMENTUM_EXIT_THRESHOLD WT_DIV_EXIT_ENABLED WT_ACCEL_EXIT_ENABLED BTC_ACCEL_RAMP_REQUIRE_POSITIVE WT_DC_DC_POS_THRESHOLD_SHORT WT_DC_DC_POS_THRESHOLD_LONG WT_DC_STOCH_THRESHOLD_SHORT WT_DC_STOCH_THRESHOLD_LONG HTF_TREND_VETO_BYPASS_REASONS HTF_GATE_D_MANDATORY HTF_TREND_VETO_BYPASS_ENABLED DELTA_GATE_BB_SQUEEZE".split()


def tabcounts(path):
    wb = openpyxl.load_workbook(str(path), read_only=True)
    out = {}
    for t in H.TABS:
        n = 0
        for r in wb[t].iter_rows(min_row=3, max_col=1, values_only=True):
            if r[0] not in (None, ""):
                n += 1
        out[t] = n
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ts", required=True)
    a = ap.parse_args()
    rep = ROOT / "data" / "template_audit" / a.ts
    stg = ROOT / "SPREADSHEETS" / "TEMPLATE_STAGED" / a.ts
    S = json.loads((rep / "summary.json").read_text())
    L = []
    w = L.append
    w(f"# TEMPLATE RESTRUCTURE AUDIT — staged {a.ts} (AGENT B)\n")
    w("Staged, NOT applied. Live templates untouched. Originals: `data/template_audit/%s/orig_TEMPLATE_*.xlsx`. Staged: `SPREADSHEETS/TEMPLATE_STAGED/%s/`. Per-row decisions: `rows_<cat>.csv`; per-switch table: `switches_<cat>.csv`; code usage evidence: `code_usage_full.json`.\n" % (a.ts, a.ts))
    w("## 1. Counts per template\n")
    keys = sorted({k for cs in CS for k in S[cs]["counts"]})
    w("| metric | " + " | ".join(CS) + " |\n|---|" + "---|" * len(CS))
    for k in keys:
        w(f"| {k} | " + " | ".join(str(S[cs]["counts"].get(k, 0)) for cs in CS) + " |")
    w("\n## 2. Rows per tab: live -> staged\n")
    w("| tab | " + " | ".join(f"{cs} live→staged" for cs in CS) + " |\n|---|" + "---|" * len(CS))
    tc = {}
    for cs in CS:
        nm = f"TEMPLATE_{cs}.xlsx"
        tc[cs] = (tabcounts(rep / ("orig_" + nm)), tabcounts(stg / nm))
    for t in H.TABS:
        w(f"| {t} | " + " | ".join(f"{tc[cs][0][t]}→{tc[cs][1][t]}" for cs in CS) + " |")
    w("\n## 3. The items the user named — where they live now (tab | options, * = bold default, G = grey)\n")
    for cs in CS:
        wb = openpyxl.load_workbook(str(stg / f"TEMPLATE_{cs}.xlsx"))
        loc = collections.defaultdict(list)
        for t in H.TABS:
            ws = wb[t]
            for r in range(3, ws.max_row + 1):
                n = ws.cell(r, 1).value
                if n in NAMED:
                    g = ws.cell(r, 1).font.color
                    loc[n].append((t, ws.cell(r, 2).value, ws.cell(r, 2).font.b, bool(g and getattr(g, "rgb", None) == "FFBFBFBF")))
        w(f"\n**{cs}**\n")
        for n in NAMED:
            l = loc.get(n)
            w(f"- `{n}`: " + ("**ABSENT** (wrong side for this template or no config field)" if not l else f"{l[0][0]} | " + " ".join(f"{o}{'*' if b else ''}{'G' if g else ''}" for _, o, b, g in l)))
    w("\n## 4. Decisions needed from the user\n")
    dec = []
    for cs in CS:
        for x in S[cs]["switches"]:
            if x["grey"].startswith("STRUCTURED"):
                dec.append((x["name"], "config value is a list/dict (e.g. HTF_TREND_VETO_BYPASS_REASONS): not sweepable as scalar options; greyed. Needs a designed option set (e.g. default list vs [] vs subsets).", cs))
            if x["imp"] == "UNWIRED":
                dec.append((x["name"], "IMPORTANT but no code path reads it (config/wiring stubs only): its delta is an honest 0 until wired — dead switch => implement it (vec + ez + tradier), not skip.", cs))
    seen = set()
    for n, t, cs in dec:
        if (n, t) in seen:
            continue
        seen.add((n, t))
        w(f"- `{n}` — {t}  ({', '.join(sorted({c for nn, tt, c in dec if nn == n and tt == t}))})")
    w("- **Grey = not calculated?** Today's pilot still calculates grey rows (only NOT_IN_CONFIG etc. are skipped, see v15_pilot ~L1809). To make NO_CONSUMER/STRUCTURED greys really 'not calculated', apply `tools/patches/v15_pilot_grey_skip.patch` (staged, not applied).")
    w("- **Numeric/float switches that had an 'OFF' option** were parked (OPTION_TYPE_MISMATCH): if OFF is a real mode it must be wired as a proper value.")
    w("- **TradierConfig stores bool switches as float** (e.g. HTF_GATE_D_MANDATORY = 0.0 in stocks, False in crypto): options follow the field type per venue; consider normalising config_tradier to bool (user-gated live config).")
    w("\n## 5. UNCLEAR — rows where it is not obvious they are important / filed correctly\n")
    for cs in CS:
        w(f"\n**{cs}**\n")
        for u in S[cs]["unclear"]:
            w(f"- `{u[1]}` ({u[2]}): {u[3]}")
        for x in S[cs]["switches"]:
            if x["imp"] == "UNCLEAR" and not x["grey"].startswith("STRUCT"):
                w(f"- `{x['name']}` → {x['home']}: {x['why']}")
    w("\n## 6. GREY lists (per template)\n")
    for cs in CS:
        ni = [x["name"] for x in S[cs]["switches"] if x["grey"] == "NOT_IN_CONFIG"]
        nc = sorted(x["name"] for x in S[cs]["switches"] if x["grey"].startswith("NO_CONSUMER"))
        st = [x["name"] for x in S[cs]["switches"] if x["grey"].startswith("STRUCT")]
        w(f"\n**{cs}** — NOT_IN_CONFIG {ni}; STRUCTURED {st}; NO_CONSUMER (contamination candidates, erase proposal) {len(nc)}: {nc}")
    w("\n## 7. Moved / merged / parked / added (summary; full list in rows_*.csv)\n")
    for cs in CS:
        rows = list(csv.reader(open(rep / f"rows_{cs}.csv")))[1:]
        c = collections.Counter(r[5] for r in rows)
        mv = sorted({(r[3], r[1], r[7]) for r in rows if r[5] == "MOVED"})
        w(f"\n**{cs}** actions: {dict(c)}")
        w("- MOVED (switch: from → to): " + "; ".join(f"{n}: {f} → {t}" for n, f, t in mv))
        ad = sorted({r[3] for r in rows if r[5] == "ADDED"})
        w(f"- ADDED own-side switches: {ad}")
        ws = sorted({r[3] for r in rows if r[5] == 'PARKED' and r[6].startswith('WRONG')})
        w(f"- WRONG_SIDE parked switches ({len(ws)}): {ws}")
    w("\n## 8. Yellow data\n")
    for cs in CS:
        c = S[cs]["counts"]
        w(f"- {cs}: ever-yellow keys {c['ever_yellow_old']} → {c['ever_yellow_new']} re-keyed ({c['ever_yellow_dropped']} dropped: duplicate broadcast copies collapse to one home row, parked options, header absent in home tab). Staged file: `yellow_ever_nonzero.staged.json`.")
    w("\n## 8b. Calculability vs wiring (not grey, but what the vector sweep can actually move)\n")
    for cs in CS:
        rows = list(csv.DictReader(open(rep / f"switches_{cs}.csv")))
        live_only = sorted(x["switch"] for x in rows if not x["grey"] and int(x["vec_consumers"]) == 0 and int(x["live_consumers"]) > 0)
        vec_only = sorted(x["switch"] for x in rows if not x["grey"] and int(x["live_consumers"]) == 0 and int(x["vec_consumers"]) > 0)
        none = sorted(x["switch"] for x in rows if not x["grey"] and int(x["live_consumers"]) == 0 and int(x["vec_consumers"]) == 0)
        w(f"\n**{cs}** LIVE_ONLY (vec cannot see -> exact 0 until vec wired; Agent A / wiring work) {len(live_only)}: {live_only}")
        w(f"- VEC_ONLY (no live counterpart; promotion must stay blocked) {len(vec_only)}: {vec_only}")
        w(f"- NO code consumer but kept calculable (important/exempt/_FILTER_TF dynamic) {len(none)}: {none}")
    w("\n## 9. Apply (needs user approval)\n```\npython tools/v15_template_staged_apply.py --ts %s            # dry-run: verify + diff summary\npython tools/v15_template_staged_apply.py --ts %s --apply    # backup live, install, sync s1/s2/s5 (+~/binance), md5 verify\n```\n" % (a.ts, a.ts))
    (rep / "report.md").write_text("\n".join(L))
    print("wrote", rep / "report.md")


if __name__ == "__main__":
    main()
