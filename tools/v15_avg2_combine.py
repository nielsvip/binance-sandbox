#!/usr/bin/env python3
"""AVG2 combiner: month history (all v15 progress dirs, 30D rows) + 365D evidence (yellow_discovery stage A / audit365) -> data/avg_delta_pos_sym.json
per cat_side per 'TAB!NAME=value': {pos_sym, n_sym, pos_30d, pos_365, pos_365_only, first_seen, last_seen, new, untested, avg_delta}.
pos_sym = number of distinct sym_sides with ANY valid positive delta (30D any run in the window, or 365D evidence). n_sym = sym_sides with any informative evaluation.
NO-LIES: exact-0 / invalid / running-default evaluations are not informative (excluded upstream); nothing is invented; untested rows have pos_sym None."""
import json, glob, os, sys, datetime, calendar, collections
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import v15_vector_delta_rebuild as V
TEMPL = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
NEW_CUT = calendar.timegm((2026, 9, 30, 20, 0, 0))
HDR = 2


def iso(t):
    return datetime.datetime.utcfromtimestamp(t).strftime("%Y-%m-%dT%H:%MZ") if t else None


def main():
    tdir = Path(sys.argv[1]); out = Path(sys.argv[2])
    hist_files = sys.argv[3:]
    # 30D history: per symside per key. Entry: [any_pos, n_files, first, last, latest, pos_clean, n_clean, pos_cont, n_cont, latest_clean]
    # (old 5-element hist files: crypto = all contaminated (pre-AUDIT/001), stocks = all clean)
    def _ext(ss, e):
        if len(e) >= 10:
            return list(e[:10])
        if V.cat_side_of(ss) and V.cat_side_of(ss).startswith("CRYPTO"):
            return [e[0], e[1], e[2], e[3], e[4], 0, 0, e[0], e[1], None]
        return [e[0], e[1], e[2], e[3], e[4], e[0], e[1], 0, 0, e[4]]
    hist = {}
    for f in hist_files:
        for ss, d in json.load(open(f)).items():
            h = hist.setdefault(ss, {})
            for k, e0 in d.items():
                e = _ext(ss, e0)
                x = h.get(k)
                if x is None:
                    h[k] = e
                else:
                    x[0] = 1 if (x[0] or e[0]) else 0; x[1] += e[1]; x[2] = min(x[2], e[2])
                    if e[3] >= x[3]:
                        x[4] = e[4]
                    x[3] = max(x[3], e[3])
                    x[5] = 1 if (x[5] or e[5]) else 0; x[6] += e[6]; x[7] = 1 if (x[7] or e[7]) else 0; x[8] += e[8]
                    if e[9] is not None:
                        x[9] = e[9]
    # 365D evidence: filter header -> per cat_side {pos:set, ev:set}
    e365 = {cs: {"pos": collections.defaultdict(set), "ev": collections.defaultdict(set)} for cs in TEMPL}
    srcs = []
    for p in glob.glob(str(ROOT / "data/yellow_discovery/20261001/fd_*/evidence_*.json")):
        cs = os.path.basename(p)[len("evidence_"):-5]
        if cs not in e365:
            continue
        j = json.load(open(p)); srcs.append(os.path.basename(os.path.dirname(p)) + "/" + os.path.basename(p))
        for hdr, d in (j.get("filters_stage_A") or {}).items():
            for sym, e in (d.get("per_sym") or {}).items():
                if isinstance(e, dict) and e.get("valid") and isinstance(e.get("delta_vs_base"), (int, float)) and e.get("effect"):
                    e365[cs]["ev"][hdr].add(sym)
                    if e["delta_vs_base"] > V.EPS:
                        e365[cs]["pos"][hdr].add(sym)
        del j
    for p in glob.glob(str(ROOT / "data/yellow_discovery/20261001/audit365_new_*/*.json")):
        cs = os.path.basename(p)[:-5]
        if cs not in e365:
            continue
        j = json.load(open(p)); srcs.append(os.path.basename(os.path.dirname(p)) + "/" + os.path.basename(p))
        for cell, d in (j.get("cells") or {}).items():
            parts = cell.split("\t")
            if len(parts) < 3:
                continue
            hdr = parts[-1]
            for sym, e in (d.get("per_sym") or {}).items():
                if isinstance(e, dict) and e.get("valid") and isinstance(e.get("delta_vs_naked"), (int, float)):
                    e365[cs]["ev"][hdr].add(sym)
                    if e["delta_vs_naked"] > V.EPS:
                        e365[cs]["pos"][hdr].add(sym)
        del j
    res = {"_meta": {"at": datetime.datetime.utcnow().isoformat() + "Z", "templates": str(tdir), "hist_files": hist_files, "evidence_365": sorted(set(srcs)),
                      "definition": "pos_sym / n_sym = distinct sym_sides with any valid positive delta / any informative evaluation, over clean AND contaminated (look-ahead crypto) results (30D, any run in window, post sign-fix; or 365D evidence for FILTER rows); *_clean / *_contaminated split them (crypto clean = written after the AUDIT/001 cutoff, stocks always clean); avg_delta_pref = mean of the latest CLEAN value per sym_side, fallback to contaminated flagged in avg_delta_pref_source. 365D per-row deltas exist only for filter rows (yellow stage A / audit365); switch rows have 30D only. avg_delta stays the 30D latest-per-sym_side mean (365D deltas are not averaged in: different scale).",
                      "new_cut": "2026-09-30T20:00Z"}}
    for cs, fn in TEMPL.items():
        wb = openpyxl.load_workbook(str(tdir / fn), read_only=False)
        by_ss = {ss: h for ss, h in hist.items() if V.cat_side_of(ss) == cs}
        d = {}
        for tab in SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            cols = {str(ws.cell(row=HDR, column=c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
            ca = cols.get("AVG_DELTA")
            for r in range(HDR + 1, ws.max_row + 1):
                a, b = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
                if a in (None, "") or b in (None, ""):
                    continue
                name = f"{str(a).strip()}={b}"
                pos30, ev30, first, last = set(), set(), None, None
                pos_c, ev_c, pos_k, ev_k, lat_c, lat_all = set(), set(), set(), set(), [], []
                for ss, h in by_ss.items():
                    for kind in ("switch", "filter"):
                        e = h.get(f"{tab}\t{kind}\t{name}")
                        if e:
                            ev30.add(ss)
                            if e[0]:
                                pos30.add(ss)
                            if e[6] > 0:
                                ev_c.add(ss)
                                if e[5]:
                                    pos_c.add(ss)
                                if e[9] is not None:
                                    lat_c.append(e[9])
                            if e[8] > 0:
                                ev_k.add(ss)
                                if e[7]:
                                    pos_k.add(ss)
                            lat_all.append(e[4])
                            first = e[2] if first is None else min(first, e[2]); last = e[3] if last is None else max(last, e[3])
                pos365 = e365[cs]["pos"].get(name, set()); ev365 = e365[cs]["ev"].get(name, set())
                pos = pos30 | pos365; ev = ev30 | ev365
                crypto = cs.startswith("CRYPTO")
                # 365D evidence: crypto yellow discovery ran before the AUDIT/001 fix = contaminated; stocks clean
                pos_c2 = pos_c | (set() if crypto else pos365); ev_c2 = ev_c | (set() if crypto else ev365)
                pos_k2 = pos_k | (pos365 if crypto else set()); ev_k2 = ev_k | (ev365 if crypto else set())
                av = ws.cell(row=r, column=ca).value if ca else None
                if lat_c:
                    pref, src = sum(lat_c) / len(lat_c), "clean"
                elif lat_all:
                    pref, src = sum(lat_all) / len(lat_all), "contaminated" if crypto else "clean"
                else:
                    pref, src = None, None
                d[f"{tab}!{name}"] = {"pos_sym": len(pos) if ev else None, "n_sym": len(ev), "pos_30d": len(pos30), "pos_365": len(pos365), "pos_365_only": len(pos365 - pos30), "first_seen": iso(first), "last_seen": iso(last),
                                      "new": bool(first is not None and first > NEW_CUT), "untested": not ev, "avg_delta": av if isinstance(av, (int, float)) else None,
                                      "n_sym_clean": len(ev_c2), "n_sym_contaminated": len(ev_k2), "pos_sym_clean": len(pos_c2), "pos_sym_contaminated": len(pos_k2),
                                      "avg_delta_pref": pref, "avg_delta_pref_source": src}
        res[cs] = d
    out.write_text(json.dumps(res, indent=1))
    for cs in TEMPL:
        v = res[cs].values()
        pc = collections.Counter((x["pos_sym"] if x["pos_sym"] is None or x["pos_sym"] < 4 else "4+") for x in v)
        print(cs, "rows", len(res[cs]), dict(pc), "new", sum(1 for x in v if x["new"]), "untested", sum(1 for x in v if x["untested"]), "pos_365_only_sum", sum(x["pos_365_only"] for x in v), "rows_with_365_only_gain", sum(1 for x in v if x["pos_365_only"] > 0))


if __name__ == "__main__":
    main()
