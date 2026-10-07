#!/usr/bin/env python3
"""Prescreen which TEMPLATE switches actually move the v12_quick_engine result.

For each sym_side, every distinct (switch, candidate) row of the 13 SWITCH_SHEETS is evaluated once, naked, against the
baseline (defaults) with the real evaluate_prepared_sanitized. A switch is LIVE-IN-VEC if any candidate changes gain_pct
or trades on any sym_side; otherwise every row of it would produce an exact 0 delta (unwired in the vector engine).
Writes the live list as the --fast-switches input for v15_pilot.py plus a CSV of every eval (no fabricated numbers).

Usage (s1/s2 only):
  python tools/v15_switch_prescreen.py --sym-sides AAVEUSDC_SHORT,BTCUSDC_LONG --template SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_SHORT.xlsx --workers 12
"""
import argparse
import concurrent.futures as cf
import csv
import json
import multiprocessing as mp
import os
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import openpyxl  # noqa: E402
import v15_pilot as P  # noqa: E402
from tools.opt.v12_pilot import evaluate_prepared_sanitized, prepare_batch  # noqa: E402

_PREPARED = None


def _eval(overrides, window_days):
    return evaluate_prepared_sanitized(_PREPARED, overrides, window_days)


def template_rows(template: Path):
    wb = openpyxl.load_workbook(str(template), read_only=True)
    rows = []
    for sheet in P.SWITCH_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        hdr = None
        for r, vals in enumerate(ws.iter_rows(min_row=1, max_col=4, values_only=True), 1):
            if r == 2:
                hdr = {str(v).strip(): i for i, v in enumerate(vals) if v is not None}
                continue
            if r < 3 or not hdr:
                continue
            switch = vals[hdr.get("Switch", 0)]
            cand = vals[hdr.get("default", 1)]
            if switch in (None, "") or cand in (None, ""):
                continue
            switch = str(switch).strip()
            if switch.lower() in ("switch", "general", "blanket", "filter", "option value") or switch.startswith("—"):
                continue
            rows.append((sheet, switch, cand))
    return rows


def main():
    global _PREPARED
    ap = argparse.ArgumentParser()
    ap.add_argument("--sym-sides", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--window-days", type=int, default=30)
    ap.add_argument("--workers", type=int, default=12)
    ap.add_argument("--out-dir", default=str(ROOT / "data" / "reports" / "v15_prescreen"))
    args = ap.parse_args()
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = template_rows(Path(args.template))
    distinct = sorted({(sw, str(c)) for _, sw, c in rows})
    print(f"[prescreen] {len(rows)} rows, {len(distinct)} distinct switch=cand", flush=True)
    live = {}
    recs = []
    for symside in args.sym_sides.split(","):
        t0 = time.time()
        _PREPARED = prepare_batch(symside, args.window_days)
        defaults = P.get_defaults_for_symside(symside)
        base = evaluate_prepared_sanitized(_PREPARED, {}, args.window_days)
        variants = []
        for sw, c in distinct:
            v = {sw: _parse(c, defaults.get(sw))}
            variants.append((sw, c, P.sanitize_overrides(v, defaults)[0]))
        with cf.ProcessPoolExecutor(max_workers=max(1, min(args.workers, (os.cpu_count() or 2) - 1)), mp_context=mp.get_context("fork")) as ex:
            futs = {ex.submit(_eval, v, args.window_days): (sw, c) for sw, c, v in variants}
            for fut in cf.as_completed(futs):
                sw, c = futs[fut]
                try:
                    res = fut.result(timeout=60)
                except Exception as e:
                    res = {"gain_pct": None, "trades": None, "invalid_reason": f"ERR {e}"[:80]}
                moved = res.get("gain_pct") is not None and (res.get("gain_pct") != base.get("gain_pct") or res.get("trades") != base.get("trades"))
                if moved:
                    live.setdefault(sw, []).append(symside)
                recs.append({"sym_side": symside, "switch": sw, "cand": c, "gain_pct": res.get("gain_pct"), "trades": res.get("trades"), "valid": res.get("valid"), "base_gain": base.get("gain_pct"), "base_trades": base.get("trades"), "delta": (res["gain_pct"] - base["gain_pct"]) if res.get("gain_pct") is not None else None, "moved": moved, "err": res.get("invalid_reason")})
        print(f"[prescreen] {symside} base gain={base.get('gain_pct'):.4f} trades={base.get('trades')} {len(variants)} evals {time.time()-t0:.1f}s live_so_far={len(live)}", flush=True)
    tag = Path(args.template).stem
    (out_dir / f"{tag}_fast_switches.json").write_text(json.dumps(sorted(live), indent=1))
    with open(out_dir / f"{tag}_prescreen.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(recs[0].keys()))
        w.writeheader()
        w.writerows(recs)
    all_sw = {sw for sw, _ in distinct}
    print(f"[prescreen] LIVE-IN-VEC {len(live)}/{len(all_sw)} switches -> {out_dir / (tag + '_fast_switches.json')}", flush=True)
    for sw in sorted(live):
        print(f"  {sw}: {','.join(live[sw])}", flush=True)


def _parse(val, default):
    if isinstance(default, bool):
        return str(val).strip().lower() == "true"
    if isinstance(default, int) and not isinstance(default, bool):
        try:
            return int(float(str(val)))
        except ValueError:
            return val
    if isinstance(default, float):
        try:
            return float(str(val))
        except ValueError:
            return val
    if str(val).strip().lower() in ("true", "false"):
        return str(val).strip().lower() == "true"
    return val


if __name__ == "__main__":
    main()
