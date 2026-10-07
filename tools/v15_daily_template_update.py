#!/usr/bin/env python3
"""v15_daily_template_update — the ONE daily writer of avg-delta results into TEMPLATE_{cat_side}.xlsx (USER 2026-09-30).

Replaces tools/v15_avg_delta_apply.py + tools/v15_vector_delta_apply.py (the latter wrote the cross-sym average INTO the
per-sym VECTOR_DELTA column; both are now refusers). Input = the merged aggregate workbook produced by
tools/v15_vector_delta_rebuild.py (SPREADSHEETS/v15_vector_delta_latest.xlsx: per cat_side tab rows tab/name/kind/pos_sym/
vector_delta(=the cross-sym_side AVERAGE)/median_delta/n; one value per sym_side per (tab,kind,name)).

Per template / per SWITCH_SHEETS tab:
  0. template columns that belong to the per-sym pilot run (override, BASELINE, HUSTLE_DELTA, VECTOR_DELTA, LIVE_DELTA,
     LIVE_SHARPE, REAL_COMPLETE, PER_ROW_FILTERS) are BLANKED — a template is a blank form; aggregate data never lives there.
  1. AVG_DELTA is REPLACED and POS_SYM is ADDED (round-idempotent via data/avg_delta_round_ledger.json) by header, per (tab, A=B).
  2. PROMOTION (the only place defaults change): per switch NAME the value with the highest POSITIVE avg (pos_sym >= --min-pos-sym)
     becomes the single bold default (+ is_default YES) in every tab that has that value; previous default regular + NO. Same for
     every yellow filter header (one bold "FILTER=opt" per filter, DEFAULT comment). ONE default per switch / filter, no more, no less.
     2c. DEFAULT REPAIR (2026-10-07): a group with NO winner keeps its default — unless it is broken (0 or 2+ YES/bold),
     in which case the writer heals it INSTEAD of failing verification forever: single-bold half-write -> YES completes it; else the
     row matching the venue live default (config.Config / TradierConfig via venue_values); else the first row. Color-blind: broken
     switch groups sit in orange rows too (appended at tab bottoms) and verification/pilot demand YES==bold regardless of fill.
     Repairs are reported as repaired_defaults (never in ledger/promoted_keys: they align the template TO config, nothing syncs
     back). --no-promote disables.
  3. worst_first: white switch groups by mean AVG_DELTA (most negative first, untested groups last), orange filter groups below;
     inside a group rows ascending by AVG_DELTA; whole rows (every cell, yellows, fonts, fills) move together.
Verified before save (row multiset unchanged, no white below orange, one YES == bold per switch name per tab, one bold header per
filter, VECTOR_DELTA empty). A failed check = that template is not saved. --apply writes (backup first); default = report only.
"""
import os as _os, sys as _sys
if _os.path.exists(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "SPREADSHEETS", "TEMPLATES_FROZEN")) and "--dry-run" not in _sys.argv and not _os.environ.get("TEMPLATES_UNFREEZE"):
    _writes = any(a in _sys.argv for a in ("--apply", "--write", "--out-dir")) or "v15_daily_template_update" in __file__ or "restructure" in __file__
    if _writes and ("--apply" in _sys.argv or "v15_daily_template_update" in __file__ or "restructure" in __file__):
        _sys.exit("REFUSED: SPREADSHEETS/TEMPLATES_FROZEN exists (USER 2026-10-01: templates restored to the pre-trainwreck 20:36 version; no writer may touch them until the user approves a proposal). Remove the file only on explicit user approval.")
import argparse
import collections
import copy
import datetime
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from v15_template_restructure import snapshot_row, write_row, is_orange, white_below_orange  # noqa: E402

SPREAD = ROOT / "SPREADSHEETS"
TPL = SPREAD / "TEMPLATE_FINAL_NORM"
AGG = SPREAD / "v15_vector_delta_latest.xlsx"
PROMOTIONS = ROOT / "data" / "cat_side_promotions.json"
ROUND_LEDGER = ROOT / "data" / "avg_delta_round_ledger.json"
POSMAP = None
TEMPLATES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
PILOT_COLS = ["override", "BASELINE", "HUSTLE_DELTA", "VECTOR_DELTA", "LIVE_DELTA", "LIVE_SHARPE", "REAL_COMPLETE", "PER_ROW_FILTERS"]
HDR = 2
EPS = 1e-9


def norm(v) -> str:
    return str(v).strip().lower()


def load_stats(path: Path) -> dict:
    """{cat_side: {(tab, name): (avg, pos, n)}} — name is 'SWITCH=value' for switch rows, 'FILTER=opt' for yellow headers."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = {}
    for cs in TEMPLATES:
        d = {}
        if cs in wb.sheetnames:
            rows = wb[cs].iter_rows(values_only=True)
            hdr = [str(h) for h in next(rows)]
            col = "avg_delta" if "avg_delta" in hdr else "vector_delta"
            ix = {h: hdr.index(h) for h in ("tab", "name", "pos_sym", col, "n")}
            for r in rows:
                if r and r[ix["name"]] is not None and r[ix[col]] is not None:
                    d[(str(r[ix["tab"]]).strip(), str(r[ix["name"]]).strip())] = (float(r[ix[col]]), int(r[ix["pos_sym"]] or 0), int(r[ix["n"]] or 0))
        out[cs] = d
    wb.close()
    return out


def col_of(ws, name: str):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=HDR, column=c).value
        if isinstance(v, str) and v.strip().upper() == name.upper():
            return c
    return None


def key_of(a, b) -> str:
    return f"{str(a).strip()}={b}"


def set_bold(cell, bold):
    f = cell.font
    cell.font = Font(name=(f.name if f else None) or "Arial", size=(f.size if f else None) or 10, bold=bold, italic=f.italic if f else None, color=f.color if f else None)


_FIELD_TYPES = None
def _field_types():
    global _FIELD_TYPES
    if _FIELD_TYPES is None:
        import dataclasses as _dc
        import sys as _sys2
        if str(ROOT) not in _sys2.path:
            _sys2.path.insert(0, str(ROOT))
        _FIELD_TYPES = {}
        try:
            import v12_quick_engine as _VQ
            for f in _dc.fields(_VQ.QuickConfig):
                _FIELD_TYPES.setdefault(f.name, f.type)
        except Exception as _e:
            print(f"[coerce-gate-warn] QuickConfig types unavailable: {_e}")
        try:
            import config as _C
            for f in _dc.fields(_C.Config):
                _FIELD_TYPES.setdefault(f.name, f.type)
        except Exception as _e:
            print(f"[coerce-gate-warn] Config types unavailable: {_e}")
        try:
            import config_tradier as _CT
            for f in _dc.fields(_CT.TradierConfig):
                _FIELD_TYPES.setdefault(f.name, f.type)
        except Exception as _e:
            print(f"[coerce-gate-warn] TradierConfig types unavailable: {_e}")
    return _FIELD_TYPES


def _coercible(name, value):
    # 2026-10-03 (BIBLE §64 3rd enforcement): a promoted value MUST pass the pilot's strict
    # _coerce_override for its field type, else it deterministically voids every defaults-pure set
    # (observed phantom: bold '0.5' for bool WT_EXHAUST_EXIT_REQUIRE_GAIN). Unknown switch = accept.
    t = _field_types().get(str(name))
    if t is None:
        return True
    ts = str(t).lower()
    if "bool" in ts:
        if isinstance(value, bool):
            return True
        if isinstance(value, (int, float)):
            return False
        return str(value).strip().lower() in ("true", "false", "1", "0", "yes", "no", "on", "off")
    if "int" in ts and "float" not in ts:
        if isinstance(value, bool):
            return False
        if isinstance(value, int):
            return True
        if isinstance(value, float):
            return value.is_integer()
        try:
            int(str(value).strip())
            return True
        except Exception:
            return False
    if "float" in ts:
        if isinstance(value, bool):
            return False
        if isinstance(value, (int, float)):
            return True
        try:
            float(str(value).strip())
            return True
        except Exception:
            return False
    return True


def _ablation_refused(name, value):
    # 2026-10-03 (BIBLE §64 4th enforcement): ABLATION_DISABLE_*=True in defaults = entry death
    # (fleet-wide 0-1 trades). The promoter must NEVER emit True here. P0 if stats say so.
    if str(name).startswith("ABLATION_DISABLE"):
        if value is True:
            return True
        return str(value).strip().lower() in ("true", "1", "yes", "on")
    return False


def groups_by_name(ws):
    g = collections.OrderedDict()
    for r in range(HDR + 1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            continue
        g.setdefault(str(a).strip(), []).append(r)
    return g


def filter_cols(ws):
    by_f = collections.OrderedDict()
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=HDR, column=c).value
        if isinstance(h, str) and "=" in h:
            by_f.setdefault(h.split("=", 1)[0].strip(), []).append(c)
    return by_f


def _floor_guard(cat_side, promoted):
    """COMBINED trade-floor check of this cat_side's promotions (tools/v15_promotion_floor_guard.py, isolated subprocess, cached per
    promoted set so the TEMPLATE_FINAL_NORM run reuses the main run's verdict). Errors fail CLOSED (all promotions of the cat_side refused)."""
    import subprocess, tempfile
    key = hashlib.md5(json.dumps({"cs": cat_side, "p": promoted}, sort_keys=True, default=str).encode()).hexdigest()
    cache = Path(tempfile.gettempdir()) / f"v15_floor_guard_{key}.json"
    if cache.exists():
        try:
            return json.loads(cache.read_text())
        except Exception:
            pass
    inp = Path(tempfile.gettempdir()) / f"v15_floor_guard_in_{key}.json"
    inp.write_text(json.dumps({"cat_side": cat_side, "promoted": promoted}, default=str))
    try:
        r = subprocess.run([sys.executable, str(ROOT / "tools" / "v15_promotion_floor_guard.py"), "--check-json", str(inp), "--out", str(cache)], capture_output=True, text=True, timeout=1800, cwd=str(ROOT))
        res = json.loads(cache.read_text())
    except Exception as e:
        res = {"ok": False, "refuse_all": True, "refused_keys": sorted(promoted), "reason": f"floor guard error (fail-closed): {e} {(locals().get('r').stderr[-300:] if locals().get('r') else '')}"}
    return res


def pick_winners(wb, stats: dict, min_pos: int, min_n: int, refused: list):
    """Cross-tab winners so ONE value per switch / filter wins in every tab: {('sw',name): (value_norm, avg)}, {('f',name): ...}."""
    win = {}
    f_opts = collections.defaultdict(list)   # filter -> [set of option values per tab that has it]
    for tab in SWITCH_SHEETS:
        if tab in wb.sheetnames:
            for f, cs in filter_cols(wb[tab]).items():
                f_opts[f].append({norm(wb[tab].cell(row=HDR, column=c).value.split("=", 1)[1]) for c in cs})
    f_common = {f: set.intersection(*v) for f, v in f_opts.items()}
    for tab in SWITCH_SHEETS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        for a, rows in groups_by_name(ws).items():
            for r in rows:
                raw = ws.cell(row=r, column=2).value
                st = stats.get((tab, key_of(a, raw)))
                if not st or st[0] <= EPS:
                    continue
                avg, pos, n = st[0], st[1], st[2] if len(st) > 2 else 0
                if pos < min_pos or n < min_n:
                    refused.append([tab, "sw", key_of(a, raw), round(avg, 4), pos, n, f"breadth pos<n(min {min_pos}/{min_n})"])
                    continue
                if _ablation_refused(a, raw):
                    refused.append([tab, "sw", key_of(a, raw), round(avg, 4), pos, n, "P0 ablation=True"])
                    continue
                if not _coercible(a, raw):
                    refused.append([tab, "sw", key_of(a, raw), round(avg, 4), pos, n, "type-uncoercible"])
                    continue
                k = ("sw", a)
                if k not in win or avg > win[k][1]:
                    win[k] = (norm(raw), avg, raw)
        for f, cs in filter_cols(ws).items():
            for c in cs:
                h = ws.cell(row=HDR, column=c).value.strip()
                st = stats.get((tab, h))
                if not st or st[0] <= EPS:
                    continue
                avg, pos, n = st[0], st[1], st[2] if len(st) > 2 else 0
                opt = h.split("=", 1)[1]
                if norm(opt) not in f_common.get(f, set()):
                    continue  # one default per filter in EVERY tab: the option must exist in all tabs carrying the filter
                if pos < min_pos or n < min_n:
                    refused.append([tab, "f", h, round(avg, 4), pos, n, f"breadth pos<n(min {min_pos}/{min_n})"])
                    continue
                if _ablation_refused(f, opt):
                    refused.append([tab, "f", h, round(avg, 4), pos, n, "P0 ablation=True"])
                    continue
                if not _coercible(f, opt):
                    refused.append([tab, "f", h, round(avg, 4), pos, n, "type-uncoercible"])
                    continue
                k = ("f", f)
                if k not in win or avg > win[k][1]:
                    win[k] = (norm(opt), avg, opt)
    return win


_VENUE_CACHE = {}
_PILOT_MOD = None


def _venue_defaults(cs):
    if cs not in _VENUE_CACHE:
        sys.path.insert(0, str(ROOT / "tools"))
        from build_cat_side_defaults_4 import venue_values
        _VENUE_CACHE[cs] = venue_values(cs.startswith("STOCKS"))[0]
    return _VENUE_CACHE[cs]


def _pilot_mod():
    global _PILOT_MOD
    if _PILOT_MOD is None:
        sys.path.insert(0, str(ROOT))
        import v15_pilot as _P
        _PILOT_MOD = _P
    return _PILOT_MOD


def _repair_default(ws, tab, a, rows, c_isd, cs, rep):
    """Heal a winner-less group with 0 or 2+ defaults (else the gate fails forever). Color-blind on purpose:
    verification and the pilot demand exactly-one YES==bold per A-group regardless of fill, and broken switch
    groups do sit in orange rows (STDEV_*/GOLDEN_RULE_* appended at tab bottoms). Returns True if repaired."""
    yes = [r for r in rows if str(ws.cell(row=r, column=c_isd).value or "").strip().upper() == "YES"]
    bold = [r for r in rows if ws.cell(row=r, column=2).font is not None and ws.cell(row=r, column=2).font.b]
    if len(yes) == 1 and yes == bold:
        return False
    best, how = None, None
    if len(bold) == 1 and not yes:
        best, how = bold[0], "single-bold-completed"
    else:
        cfg = _venue_defaults(cs).get(a)
        if cfg is not None:
            _P = _pilot_mod()
            for r in rows:
                try:
                    pv = _P._parse_opt_value(ws.cell(row=r, column=2).value, cfg)
                except Exception:
                    continue
                if _P._same_default(pv, cfg):
                    best, how = r, "config-default"
                    break
        if best is None:
            best, how = rows[0], "first-row-fallback"
    for r in rows:
        y = r == best
        set_bold(ws.cell(row=r, column=2), y)
        ws.cell(row=r, column=c_isd).value = "YES" if y else "NO"
        ws.cell(row=r, column=c_isd).font = Font(name="Arial", size=10, bold=y)
    rep.setdefault("repaired_defaults", []).append([tab, a, how, ws.cell(row=best, column=2).value])
    return True


def apply_tab(ws, stats, win, ledger_cs, round_changed, rep, ledger_out, cleared_total, cs=None, allow_repair=True):
    tab = ws.title
    c_isd, c_avg, c_pos, c_vec = col_of(ws, "is_default"), col_of(ws, "AVG_DELTA"), col_of(ws, "POS_SYM"), col_of(ws, "VECTOR_DELTA")
    if not (c_isd and c_avg and c_pos and c_vec):
        raise RuntimeError(f"{tab}: missing is_default/AVG_DELTA/POS_SYM/VECTOR_DELTA header")
    # 0. blank the pilot-owned columns
    for name in PILOT_COLS:
        c = col_of(ws, name)
        if not c:
            continue
        for r in range(HDR + 1, ws.max_row + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                ws.cell(row=r, column=c).value = None
                cleared_total[name] += 1
    # 1. AVG_DELTA replace / POS_SYM add (round-idempotent)
    bootstrap = not ledger_cs
    for r in range(HDR + 1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            continue
        k = key_of(a, ws.cell(row=r, column=2).value)
        lk = f"{tab}|{k}"
        st = stats.get((tab, k))
        cur_pos = ws.cell(row=r, column=c_pos).value
        cur_pos = int(cur_pos) if isinstance(cur_pos, (int, float)) else 0
        if round_changed or lk not in ledger_out:
            ledger_out[lk] = 0 if bootstrap else cur_pos   # base before THIS round's positives are added
        base = ledger_out[lk]
        if st:
            ws.cell(row=r, column=c_avg).value = round(st[0], 6)
            ws.cell(row=r, column=c_pos).value = base + st[1]
            rep["rows_with_stats"] += 1
        elif bootstrap:
            ws.cell(row=r, column=c_avg).value = None   # old pooled/contaminated values are not trusted
            ws.cell(row=r, column=c_pos).value = None
    # 1b. AVG2 POS_SYM override from the history json (any valid positive in the window, 30D all runs + 365D)
    if POSMAP is not None:
        for r in range(HDR + 1, ws.max_row + 1):
            a = ws.cell(row=r, column=1).value
            if a in (None, ""):
                continue
            e = POSMAP.get(f"{tab}!{key_of(a, ws.cell(row=r, column=2).value)}")
            if e is not None and e.get("pos_sym") is not None:
                ws.cell(row=r, column=c_pos).value = int(e["pos_sym"])
    # 2. promotion — switches
    for a, rows in groups_by_name(ws).items():
        w = win.get(("sw", a))
        if not w:
            if allow_repair and cs is not None:
                _repair_default(ws, tab, a, rows, c_isd, cs, rep)
            continue
        tgt = [r for r in rows if norm(ws.cell(row=r, column=2).value) == w[0]]
        if not tgt:
            continue
        best = tgt[0]
        cur = [r for r in rows if str(ws.cell(row=r, column=c_isd).value or "").strip().upper() == "YES"]
        if cur == [best]:
            rep["ledger"][a] = {"value": ws.cell(row=best, column=2).value, "avg_delta": w[1], "tab": tab, "at": datetime.datetime.utcnow().isoformat() + "Z"}
            continue
        for r in rows:
            yes = r == best
            set_bold(ws.cell(row=r, column=2), yes)
            ws.cell(row=r, column=c_isd).value = "YES" if yes else "NO"
            ws.cell(row=r, column=c_isd).font = Font(name="Arial", size=10, bold=yes)
        rep["promoted_switch"].append([tab, a, [ws.cell(row=r, column=2).value for r in cur], ws.cell(row=best, column=2).value, round(w[1], 4)])
        rep["ledger"][a] = {"value": ws.cell(row=best, column=2).value, "avg_delta": w[1], "tab": tab, "at": datetime.datetime.utcnow().isoformat() + "Z"}
    # 2b. promotion — yellow filter headers
    for f, cs in filter_cols(ws).items():
        w = win.get(("f", f))
        if not w:
            continue
        tgt = [c for c in cs if norm(ws.cell(row=HDR, column=c).value.split("=", 1)[1]) == w[0]]
        if not tgt:
            continue
        best = tgt[0]
        cur = [c for c in cs if ws.cell(row=HDR, column=c).font is not None and ws.cell(row=HDR, column=c).font.b]
        if cur == [best]:
            rep["ledger"][f] = {"value": w[2], "avg_delta": w[1], "tab": tab, "at": datetime.datetime.utcnow().isoformat() + "Z"}
            continue
        for c in cs:
            h = ws.cell(row=HDR, column=c)
            set_bold(h, c == best)
            if c == best:
                h.comment = Comment("DEFAULT", "v15_daily_template_update")
            elif h.comment is not None and h.comment.text.strip() == "DEFAULT":
                h.comment = None
        rep["promoted_filter"].append([tab, f, [ws.cell(row=HDR, column=c).value for c in cur], ws.cell(row=HDR, column=best).value, round(w[1], 4)])
        rep["ledger"][f] = {"value": w[2], "avg_delta": w[1], "tab": tab, "at": datetime.datetime.utcnow().isoformat() + "Z"}
    # 3. worst_first re-order; whole rows move
    ncol = ws.max_column
    grp = collections.OrderedDict()
    last = None
    for r in range(HDR + 1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            if last is not None:
                grp[last].append(snapshot_row(ws, r, ncol))
            continue
        k = (str(a).strip(), is_orange(ws, r))
        grp.setdefault(k, [])
        grp[k].append(snapshot_row(ws, r, ncol))
        last = k

    def val(cells):
        v = cells[c_avg - 1][0]
        return v if isinstance(v, (int, float)) else None

    for k in grp:
        nonblank = [s for s in grp[k] if str(s[0][0] or "").strip() != ""]
        blanks = [s for s in grp[k] if str(s[0][0] or "").strip() == ""]
        nonblank.sort(key=lambda s: (val(s) is None, val(s) if val(s) is not None else 0.0))
        grp[k] = nonblank + blanks

    def gkey(k):
        vs = [val(s) for s in grp[k] if val(s) is not None]
        return (not vs, sum(vs) / len(vs) if vs else 0.0)

    white = sorted([k for k in grp if not k[1]], key=gkey)
    orange = sorted([k for k in grp if k[1]], key=gkey)
    if white:  # USER 2026-10-01: a tab ALWAYS starts with a DEFAULT row (baseline of tab 1 and the port to the next tab are built from defaults+overrides)
        _rows = grp[white[0]]
        _yes = [x for x in _rows if str(x[c_isd - 1][0] or "").strip().upper() == "YES"]
        if _yes and _rows[0] is not _yes[0]:
            _rows.remove(_yes[0])
            _rows.insert(0, _yes[0])
    r = HDR + 1
    for k in white + orange:
        for cells in grp[k]:
            write_row(ws, r, cells)
            r += 1


def verify(before, after) -> list:
    import template_row_guard as G
    bad = []
    for tab in SWITCH_SHEETS:
        if tab not in after.sheetnames:
            continue
        a, b = before[tab], after[tab]
        ms = lambda ws: collections.Counter((str(ws.cell(row=r, column=1).value), str(ws.cell(row=r, column=2).value)) for r in range(HDR + 1, ws.max_row + 1) if ws.cell(row=r, column=1).value)
        if ms(a) != ms(b):
            bad.append(f"{tab}: row set changed")
        _skip = {col_of(b, n) for n in PILOT_COLS + ["AVG_DELTA", "POS_SYM", "is_default"]}
        try:
            G.assert_rows_intact(a, b, _skip)
        except AssertionError as _e:
            bad.append(str(_e))
        if white_below_orange(b):
            bad.append(f"{tab}: white rows below orange {white_below_orange(b)[:5]}")
        c_isd, c_vec = col_of(b, "is_default"), col_of(b, "VECTOR_DELTA")
        if any(b.cell(row=r, column=c_vec).value not in (None, "") for r in range(HDR + 1, b.max_row + 1)):
            bad.append(f"{tab}: VECTOR_DELTA not empty")
        if tab != "STDEV_SLOPE_SIZING" and not any(is_orange(b, r) for r in range(HDR + 1, b.max_row + 1) if b.cell(row=r, column=1).value not in (None, "")):
            bad.append(f"{tab}: no orange FILTER rows (they must exist in every tab, below the white switch rows)")
        _first = next((r for r in range(HDR + 1, b.max_row + 1) if b.cell(row=r, column=1).value not in (None, "")), None)
        if _first and str(b.cell(row=_first, column=c_isd).value or "").strip().upper() != "YES":
            bad.append(f"{tab}: first data row {_first} is not a default row")
        for sw, rows in groups_by_name(b).items():
            yes = [r for r in rows if str(b.cell(row=r, column=c_isd).value or "").strip().upper() == "YES"]
            bold = [r for r in rows if b.cell(row=r, column=2).font is not None and b.cell(row=r, column=2).font.b]
            if len(yes) != 1 or yes != bold:
                bad.append(f"{tab}!{sw}: YES {yes} bold {bold}")
        for f, cs in filter_cols(b).items():
            if sum(1 for c in cs if b.cell(row=HDR, column=c).font is not None and b.cell(row=HDR, column=c).font.b) != 1:
                bad.append(f"{tab}: filter {f} bold headers != 1")
        # a moved row keeps its yellow cells: compare a deep sample by (A,B)
        # PIPE 2026-10-01: key on the RAW option (repr) so 'True' (text) and True (bool) rows stay two rows (template has such duplicate option rows, e.g. REDUCE_PROFIT_LOCK WT_D_BOUNCE_DD_STOP_ENABLED); whole rows untouched
        ia = {(str(a.cell(row=r, column=1).value), repr(a.cell(row=r, column=2).value)): r for r in range(HDR + 1, a.max_row + 1) if a.cell(row=r, column=1).value}
        ib = {(str(b.cell(row=r, column=1).value), repr(b.cell(row=r, column=2).value)): r for r in range(HDR + 1, b.max_row + 1) if b.cell(row=r, column=1).value}
        skip = {col_of(b, n) for n in PILOT_COLS + ["AVG_DELTA", "POS_SYM", "is_default"]}
        _cnt = collections.Counter((str(a.cell(row=r, column=1).value), repr(a.cell(row=r, column=2).value)) for r in range(HDR + 1, a.max_row + 1) if a.cell(row=r, column=1).value)
        for k in list(ib)[:: max(1, len(ib) // 25)]:
            if k not in ia or _cnt[k] > 1:
                continue  # AUTOPILOT 2026-10-02: duplicate (A,B) rows (e.g. ATR_ADAPTIVE_SIZING_TARGET_PCT=1.5 with two different Vec Hook cells) are keyed ambiguously: row multiset is already verified above
            for c in range(1, min(a.max_column, b.max_column) + 1):
                if c in skip:
                    continue
                ca, cb = a.cell(row=ia[k], column=c), b.cell(row=ib[k], column=c)
                fa, fb = str(getattr(ca.fill.fgColor, "rgb", "") or ""), str(getattr(cb.fill.fgColor, "rgb", "") or "")
                if str(ca.value) != str(cb.value) or fa != fb:
                    if c == 2:
                        continue  # bold flip only changes the font, value must match — value compared above
                    bad.append(f"{tab} moved row {k}: col {c} changed")
                    break
    return bad


TAB_LEVEL_SPEC = ROOT / "data" / "wiring" / "tab_filters" / "tab_level_filters.json"


def _last_data_row(ws):
    return max((r for r in range(HDR + 1, ws.max_row + 1) if ws.cell(row=r, column=1).value not in (None, "")), default=HDR)


def add_tab_level_rows(ws, spec_tab, rep):
    """USER 2026-10-01: append the tab-level filters as orange rows at the bottom (whole new rows only, no existing row touched).
    A filter that already has rows in the tab is reused (missing options appended as rows of the same group). Default option = the bold header
    (== config default) so no baseline shifts; every new group gets exactly one bold YES row."""
    import copy as _copy
    groups = groups_by_name(ws)
    c_isd, c_avg, c_pos = col_of(ws, "is_default"), col_of(ws, "AVG_DELTA"), col_of(ws, "POS_SYM")
    last = _last_data_row(ws)
    ref = next((r for r in range(last, HDR, -1) if is_orange(ws, r)), None)
    if ref is None:
        rep["no_orange_ref"] = rep.get("no_orange_ref", 0) + 1
        return 0
    ncol = min(ws.max_column, (c_pos or 14))
    base = snapshot_row(ws, ref, ncol)
    r = last
    added = 0
    for f, info in spec_tab.items():
        have = {str(ws.cell(row=x, column=2).value).strip() for x in groups.get(f, [])}
        white_group = any(not is_orange(ws, x) for x in groups.get(f, []))
        if white_group:
            rep["reused_white"] = rep.get("reused_white", 0) + 1
            continue
        default = info["default"]
        new_group = not groups.get(f)
        for opt in info["options"]:
            if opt in have:
                continue
            r += 1
            cells = [(None, _copy.copy(fo), _copy.copy(fi), nf, _copy.copy(al)) for (_v, fo, fi, nf, al) in base]
            write_row(ws, r, cells)
            ws.cell(row=r, column=1).value = f
            ws.cell(row=r, column=2).value = opt
            ws.cell(row=r, column=c_isd).value = "YES" if (new_group and opt == default) else "NO"
            for c in range(2, ncol + 1):
                if c != 2 and c != c_isd:
                    ws.cell(row=r, column=c).fill = PatternFill(fill_type=None)
            for c in (1, 2, c_isd):
                set_bold(ws.cell(row=r, column=c), ws.cell(row=r, column=c_isd).value == "YES" if c != 1 else ws.cell(row=r, column=c_isd).value == "YES")
            added += 1
            have.add(opt)
        if not groups.get(f) and default not in info["options"]:
            rep["default_not_in_options"] = rep.get("default_not_in_options", 0) + 1
        rep["filters_added"] = rep.get("filters_added", 0) + (0 if groups.get(f) else 1)
    rep["rows_added"] = rep.get("rows_added", 0) + added
    return added


def verify_tab_level(before, after, expected) -> list:
    import template_row_guard as G
    bad = []
    for tab in SWITCH_SHEETS:
        if tab not in after.sheetnames:
            continue
        a, b = before[tab], after[tab]
        lost = G.fingerprints(a) - G.fingerprints(b)
        if lost:
            bad.append(f"{tab}: {sum(lost.values())} original rows changed/lost")
        n_a = sum(1 for r in range(HDR + 1, a.max_row + 1) if a.cell(row=r, column=1).value not in (None, ""))
        n_b = sum(1 for r in range(HDR + 1, b.max_row + 1) if b.cell(row=r, column=1).value not in (None, ""))
        if n_b - n_a != expected.get(tab, 0):
            bad.append(f"{tab}: rows {n_a}->{n_b}, expected +{expected.get(tab, 0)}")
        for r in range(HDR + 1, a.max_row + 1):  # original rows keep their position (rows only appended)
            if a.cell(row=r, column=1).value != b.cell(row=r, column=1).value or a.cell(row=r, column=2).value != b.cell(row=r, column=2).value:
                bad.append(f"{tab}: original row {r} moved")
                break
        if white_below_orange(b):
            bad.append(f"{tab}: white rows below orange")
        def _viol(ws):
            ci = col_of(ws, "is_default")
            out = set()
            for sw, rows in groups_by_name(ws).items():
                yes = [r for r in rows if str(ws.cell(row=r, column=ci).value or "").strip().upper() == "YES"]
                bold = [r for r in rows if ws.cell(row=r, column=2).font is not None and ws.cell(row=r, column=2).font.b]
                if len(yes) != 1 or [ws.cell(row=r, column=2).value for r in yes] != [ws.cell(row=r, column=2).value for r in bold]:
                    out.add(sw)
            return out
        for sw in sorted(_viol(b) - _viol(a)):  # one YES == bold per group: only violations INTRODUCED here count (pre-existing ones are not this writer's)
            bad.append(f"{tab}!{sw}: one-default rule broken by the new rows")
        if [b.cell(row=HDR, column=c).value for c in range(1, b.max_column + 1)] != [a.cell(row=HDR, column=c).value for c in range(1, a.max_column + 1)]:
            bad.append(f"{tab}: header row changed")
    return bad


def main_tab_level(args):
    import shutil
    import zipfile
    spec = json.loads(Path(args.spec).read_text())["cats"]
    tdir = Path(args.template_dir)
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    for cs in (list(TEMPLATES) if args.cat_side == "ALL" else [args.cat_side]):
        path = tdir / TEMPLATES[cs]
        before = openpyxl.load_workbook(str(path))
        wb = openpyxl.load_workbook(str(path))
        expected, rep = {}, {}
        for tab, fs in spec.get(cs, {}).items():
            if tab in wb.sheetnames:
                expected[tab] = add_tab_level_rows(wb[tab], fs, rep)
        bad = verify_tab_level(before, wb, expected)
        print(f"[{cs}] tab-level rows added {expected} {rep} violations={len(bad)} {bad[:3]}")
        if not args.apply:
            continue
        if bad:
            print(f"[{cs}] NOT saved — verification failed")
            continue
        shutil.copy2(path, ROOT / "backups" / f"before_tab_level_filters_{ts}_{path.name}")
        tmp = path.with_suffix(".tmp.xlsx")
        wb.save(str(tmp))
        if len(zipfile.ZipFile(str(tmp)).namelist()) < 10:
            raise SystemExit(f"{tmp}: truncated zip")
        openpyxl.load_workbook(str(tmp))
        tmp.replace(path)
        print(f"[{cs}] saved {path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--dry-run", action="store_true", help="accepted no-op (report-only is the default; also bypasses the freeze guard)")
    ap.add_argument("--tab-level", action="store_true", help="ONLY append the tab-level filter rows (data/wiring/tab_filters/tab_level_filters.json) as orange rows at the bottom of each tab")
    ap.add_argument("--spec", default=str(TAB_LEVEL_SPEC))
    ap.add_argument("--cat-side", default="ALL")
    ap.add_argument("--agg", default=str(AGG), help="merged aggregate workbook (v15_vector_delta_rebuild output)")
    ap.add_argument("--template-dir", default=str(TPL))
    ap.add_argument("--min-pos-sym", type=int, default=2, help="breadth gate: a value needs >= N positive sym_sides this round to be promoted")
    ap.add_argument("--min-n", type=int, default=3, help="evidence gate (2026-10-03): a value needs >= N contributing sym_sides to be promoted")
    ap.add_argument("--round-id", default=None, help="default = md5 of the aggregate workbook")
    ap.add_argument("--pos-json", default=None, help="AVG2: data/avg_delta_pos_sym.json — POS_SYM of every row = history count (30D all runs + 365D evidence); AVG_DELTA stays from the aggregate")
    ap.add_argument("--no-promote", action="store_true", help="AVG2 2026-10-01: do NOT change any default (bold/is_default); only AVG_DELTA/POS_SYM columns and the worst_first row order")
    ap.add_argument("--fresh-pos", action="store_true", help="AVG2: POS_SYM = this aggregate's positive sym_side count (ledger ignored, bootstrap semantics); rows without stats get blank AVG_DELTA/POS_SYM")
    ap.add_argument("--sync-defaults", action="store_true", help="after --apply also rebuild data/cat_side_defaults_4.json (tools/build_cat_side_defaults_4.py) = pushes the new bold defaults to LIVE + sweep engine; OFF by default, run deliberately")
    args = ap.parse_args()
    if args.tab_level:
        return main_tab_level(args)
    global POSMAP
    POSMAP_ALL = None
    if args.pos_json:
        _pj = json.loads(Path(args.pos_json).read_text())
        POSMAP_ALL = {_cs: _d for _cs, _d in _pj.items() if not _cs.startswith("_")}
    agg = Path(args.agg)
    tdir = Path(args.template_dir)
    round_id = args.round_id or hashlib.md5(agg.read_bytes()).hexdigest()
    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    all_stats = load_stats(agg)
    # SQL-primary → JSON fallback prevents silent loss if files vanish mid-herd
    try:
        import per_sym_store as _pss
        promos = _pss.get_cat_side_promotions()
        rl = _pss.get_avg_delta_round_ledger()
        # ensure file fallback if SQL empty on first boot
        if not promos and PROMOTIONS.exists():
            promos = json.loads(PROMOTIONS.read_text())
        if not rl and ROUND_LEDGER.exists():
            rl = json.loads(ROUND_LEDGER.read_text())
    except Exception:
        promos = json.loads(PROMOTIONS.read_text()) if PROMOTIONS.exists() else {}
        rl = json.loads(ROUND_LEDGER.read_text()) if ROUND_LEDGER.exists() else {}
    saved_any = False
    saved_promos = {}
    report = {"round_id": round_id, "agg": str(agg)}
    for cs in (list(TEMPLATES) if args.cat_side == "ALL" else [args.cat_side]):
        path = tdir / TEMPLATES[cs]
        stats = all_stats[cs]
        before = openpyxl.load_workbook(str(path))
        wb = openpyxl.load_workbook(str(path))
        rep = {"rows_with_stats": 0, "promoted_switch": [], "promoted_filter": [], "ledger": {}}
        cleared = collections.Counter()
        ledger_cs = {} if args.fresh_pos else (rl.get(cs, {}).get("pos_base") or {})
        round_changed = True if args.fresh_pos else (rl.get(cs, {}).get("round_id") != round_id)
        ledger_out = {} if round_changed else dict(ledger_cs)
        if not round_changed:
            ledger_out = dict(ledger_cs)
        refused = []
        win = {} if args.no_promote else pick_winners(wb, stats, args.min_pos_sym, args.min_n, refused)
        report.setdefault("refused", {}).__setitem__(cs, refused)
        if refused:
            print(f"[{cs}] refused promotions: {len(refused)} (first: {refused[:3]})")
        POSMAP = POSMAP_ALL.get(cs) if POSMAP_ALL is not None else None
        for tab in SWITCH_SHEETS:
            if tab in wb.sheetnames:
                apply_tab(wb[tab], stats, win, ledger_cs, round_changed, rep, ledger_out, cleared, cs, not args.no_promote)
        _fg = _floor_guard(cs, {k: v.get("value") for k, v in rep["ledger"].items()}) if (rep["ledger"] and _os.environ.get("V15_PROMO_FLOOR_GUARD", "1") != "0") else None
        if _fg and _fg.get("refused_keys"):  # COMBINED trade-floor guard (director 2026-10-06): re-apply without the refused keys (old bolds kept)
            _drop = set(_fg["refused_keys"])
            print(f"[{cs}] FLOOR-GUARD refused {len(_drop)} promotion(s){' (WHOLE cat_side)' if _fg.get('refuse_all') else ''}: {sorted(_drop)[:12]} | {_fg.get('reason')}")
            win = {k: v for k, v in win.items() if k[1] not in _drop}
            wb = openpyxl.load_workbook(str(path))
            rep = {"rows_with_stats": 0, "promoted_switch": [], "promoted_filter": [], "ledger": {}}
            cleared = collections.Counter()
            ledger_out = {} if round_changed else dict(ledger_cs)
            for tab in SWITCH_SHEETS:
                if tab in wb.sheetnames:
                    apply_tab(wb[tab], stats, win, ledger_cs, round_changed, rep, ledger_out, cleared, cs, not args.no_promote)
        bad = verify(before, wb)
        report[cs] = {"rows_with_stats": rep["rows_with_stats"], "cleared": dict(cleared), "promoted_switch": rep["promoted_switch"], "promoted_filter": rep["promoted_filter"], "repaired_defaults": rep.get("repaired_defaults", []), "violations": bad}
        if _fg is not None:
            report[cs]["floor_guard"] = {k: _fg.get(k) for k in ("ok", "refuse_all", "refused_keys", "reason", "samples", "candidate_fail", "eliminated", "single_key")}
            report[cs]["floor_guard"].update({f"{n}_median": {m: (_fg.get(n) or {}).get(m) for m in ("med_trades", "med_tim")} for n in ("baseline", "candidate", "final")})
        print(f"[{cs}] stats={len(stats)} rows_with_stats={rep['rows_with_stats']} cleared={dict(cleared)} promoted switches={len(rep['promoted_switch'])} filters={len(rep['promoted_filter'])} repaired={len(rep.get('repaired_defaults', []))} violations={len(bad)} {bad[:3]}")
        if not args.apply:
            continue
        if bad:
            print(f"[{cs}] NOT saved — verification failed")
            continue
        import shutil
        shutil.copy2(path, ROOT / "backups" / f"before_daily_template_update_{ts}_{path.name}")
        tmp = path.with_suffix(".tmp.xlsx")
        wb.save(str(tmp))
        openpyxl.load_workbook(str(tmp))
        tmp.replace(path)
        promos.setdefault(cs, {}).update(rep["ledger"])
        rl[cs] = {"round_id": round_id, "pos_base": ledger_out, "at": datetime.datetime.utcnow().isoformat() + "Z"}
        saved_any = True
        saved_promos[cs] = list(rep["ledger"])
        print(f"[{cs}] saved {path}")
    if args.apply and saved_any and tdir.resolve() == SPREAD.resolve():
        PROMOTIONS.write_text(json.dumps(promos, indent=1, default=str))
        ROUND_LEDGER.write_text(json.dumps(rl, indent=1))
        # dual-write: JSON stays as generated view, SQL primary prevents silent fallback to globals on next promotion round
        try:
            import per_sym_store as _pss
            _pss.kv_put(_pss.KV_CAT_SIDE_PROMOTIONS, promos)
            _pss.kv_put(_pss.KV_AVG_DELTA_ROUND_LEDGER, rl)
        except Exception as _e:
            print(f"[kv] promos/ledger kv_put failed (JSON view still written): {_e}")
        try:  # INV 2026-10-01: keep the documentation tab WIRING_INVENTORY in the user finals (refreshed only when missing/stale; never read by pilots)
            import subprocess
            _ir = subprocess.run([sys.executable, str(ROOT / "tools" / "v15_wiring_inventory.py"), "--ensure"], capture_output=True, text=True, check=False, timeout=900)
            if _ir.returncode != 0:
                _tail = " ".join(((_ir.stderr or "") + "\n" + (_ir.stdout or "")).split())[-300:]
                print(f"[inventory] ensure rc={_ir.returncode} (non-fatal, output captured): {_tail}")
        except Exception as _e:
            print(f"[inventory] ensure failed: {' '.join(str(_e).split())}")
        if args.sync_defaults:
            import subprocess
            subprocess.run([sys.executable, str(ROOT / "tools" / "build_cat_side_defaults_4.py")], check=True)
        _pk = sorted({k for _kl in saved_promos.values() for k in _kl})
        report["saved_promos"] = saved_promos
        report["promoted_keys"] = _pk
        _sync_on = _os.environ.get("V15_TEMPLATE_SYNC_SURFACES", "1") != "0"  # S1 daily chain sets 0: S1 never edits config.py/config_tradier.py/v12_quick_engine.py; the Mac apply script syncs report["promoted_keys"]
        if not _sync_on:
            report["parity_sync"] = {"keys": len(_pk), "skipped": "V15_TEMPLATE_SYNC_SURFACES=0 (code surfaces synced on the Mac by tools/v15_daily_chain_mac_apply.sh)"}
            print(f"[parity-sync] SKIPPED (V15_TEMPLATE_SYNC_SURFACES=0) promoted={len(_pk)} keys -> report promoted_keys")
        try:  # USER 2026-10-06 obligated symmetry: promoted bolds auto-sync venue globals + QuickConfig (same run, never manual)
            sys.path.insert(0, str(ROOT))
            import switch_parity as _sp
            if _pk and _sync_on:
                _sr = _sp.sync_default_surfaces(apply=True, confirm_unlocked=True, keys=_pk)
                print(f"[parity-sync] promoted={len(_pk)} planned={len(_sr.get('planned', []))} applied={_sr.get('applied')} manual={len(_sr.get('needs_manual', []))} err={str(_sr.get('error') or '')[:160]}")
                for _pm in _sr.get("needs_manual", [])[:10]:
                    print(f"[parity-sync-MANUAL] {_pm.get('cat_side')} {_pm.get('key')}: {_pm.get('why')}")
                report["parity_sync"] = {"keys": len(_pk), "planned": len(_sr.get("planned", [])), "applied": _sr.get("applied"), "manual": len(_sr.get("needs_manual", [])), "error": str(_sr.get("error") or "")[:200]}
        except Exception as _sp_e:
            print(f"[parity-sync-warn] {_sp_e}")
    out = ROOT / "data" / "reports" / f"v15_daily_template_update_{ts}.json"
    out.write_text(json.dumps(report, indent=1, default=str))
    print(f"[report] {out}")
    if _os.environ.get("V15_SLP_AUTOSET") == "1":  # SLP 2026-10-01: slope-sizing autoset (validated rules only); dry-run unless V15_SLP_AUTOSET_APPLY=1
        try:
            _slp = [sys.executable, str(ROOT / "tools" / "v15_slope_sizing_autoset.py")]
            _dry = [] if _os.environ.get("V15_SLP_AUTOSET_APPLY") == "1" else ["--dry-run"]
            for _v in ("stocks", "crypto"):
                __import__("subprocess").run(_slp + ["decide", "--venue", _v] + _dry, check=False)
            __import__("subprocess").run(_slp + ["apply"] + _dry, check=False)
        except Exception as _e:
            print(f"[slp-autoset] skipped: {_e}")


if __name__ == "__main__":
    main()
