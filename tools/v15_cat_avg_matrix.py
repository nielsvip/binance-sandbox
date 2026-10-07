#!/usr/bin/env python3
"""Per-cat_side avg/pos_sym matrix inventory (USER 2026-10-07).

Builds ONE xlsx per cat_side (CRYPTO_LONG/CRYPTO_SHORT/STOCKS_LONG/STOCKS_SHORT),
each mirroring its TEMPLATE_{cat_side} (12 trading tabs, every white switch row,
every orange filter row, every per-switch filter column O..end) with:
  - row M AVG_DELTA / N POS_SYM = avg + pos_sym of VECTOR_DELTA across sym_sides
  - every yellow/filter cell that EVER held a number in ANY sym_side final xlsx:
    avg_delta (base sheet), pos_sym (_POS sheet), n (_N sheet)
Blanks = never numeric anywhere (never tested / skipped / invalid text only).
Cells filled regardless of yellow paint (decisions need unpainted numbers too).

Source: SPREADSHEEDS/V15_V16_CELL_BY_CELL final xlsx per sym_side (ONE file per
sym_side: newest, with >=50% of that sym_side's best numeric-cell count, else the
next newest -- same rule as tools/v15_avg_delta_rebuild.select_latest).
pos_sym = #sym_sides with delta > 1e-9. avg = mean over numeric values (exact
zeros count in n, not in pos). No medians, no annualization, full precision.

Steps (checkpointed in data/cat_avg/): count -> select -> extract -> build.
"""
import argparse, collections, datetime, json, os, pathlib, re, sys
from multiprocessing import Pool

ROOT = pathlib.Path(__file__).resolve().parents[1]
SRC = ROOT / "SPREADSHEETS" / "V15_V16_CELL_BY_CELL"
OUTDIR = ROOT / "SPREADSHEETS" / "V15_CAT_AVG"
WORK = ROOT / "data" / "cat_avg"
PARTIALS = pathlib.Path("/tmp/catavg_partials")
TABS12 = ["ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES",
    "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE",
    "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER",
    "GLOBAL_RISK_GATES"]
CAT_SIDES = ["CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"]
POS_EPS = 1e-9
FILE_PAT = re.compile(r"^(.+)_(LONG|SHORT)_(?:bh.+_gain.+_)?30d_matrix.*\.xlsx$")

def sym_side_of(fname):
    m = FILE_PAT.match(fname)
    return (m.group(1), m.group(2)) if m else (None, None)

def is_crypto(sym):
    return sym.endswith("USDT") or sym.endswith("USDC")

def cat_side_of(sym, side):
    return ("CRYPTO" if is_crypto(sym) else "STOCKS") + "_" + side

def norm_cand(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if abs(v) < 1e15 and v == int(v) else repr(v)
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()

def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)

def stream_file(path, want_values):
    """Stream the 12 tabs. want_values=False -> counts only. Returns dict."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out_tabs = {} if want_values else None
    n_g = 0
    n_y = 0
    n_str = 0
    dups = 0
    try:
        for tab in TABS12:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            it = ws.iter_rows(values_only=True)
            next(it, None)
            row2 = next(it, None)
            if not row2:
                continue
            r2 = list(row2)
            hmap = {}
            for i, v in enumerate(r2[:14]):
                if isinstance(v, str) and v.strip():
                    hmap[v.strip()] = i
            acol = hmap.get("Switch", 0)
            bcol = hmap.get("default", 1)
            gcol = hmap.get("VECTOR_DELTA", 6)
            ycols = []
            for i in range(14, len(r2)):
                h = r2[i]
                if h is None or (isinstance(h, str) and not h.strip()):
                    continue
                hs = str(h).strip()
                if hs.startswith("WHAT SWITCH"):
                    break
                ycols.append((i, hs))
            headers = [h for _, h in ycols]
            seen = set()
            rows = []
            for row in it:
                aval = row[acol] if acol < len(row) else None
                if aval is None or (isinstance(aval, str) and not aval.strip()):
                    continue
                sw = str(aval).strip()
                cand = norm_cand(row[bcol] if bcol < len(row) else None)
                g = row[gcol] if gcol < len(row) else None
                gnum = is_num(g)
                if gnum:
                    n_g += 1
                elif isinstance(g, str) and g.strip():
                    n_str += 1
                key = (sw, cand)
                if key in seen:
                    dups += 1
                    continue
                seen.add(key)
                if want_values:
                    yvals = []
                    for j, (ci, _) in enumerate(ycols):
                        v = row[ci] if ci < len(row) else None
                        if is_num(v):
                            n_y += 1
                            yvals.append((j, v))
                        elif isinstance(v, str) and v.strip():
                            n_str += 1
                    if gnum or yvals:
                        rows.append([sw, cand, g if gnum else None, yvals])
                else:
                    for ci, _ in ycols:
                        v = row[ci] if ci < len(row) else None
                        if is_num(v):
                            n_y += 1
                        elif isinstance(v, str) and v.strip():
                            n_str += 1
            if want_values:
                out_tabs[tab] = {"headers": headers, "rows": rows}
    finally:
        wb.close()
    st = os.stat(path)
    res = {"file": os.path.basename(path), "mtime": st.st_mtime, "size": st.st_size,
        "n_g": n_g, "n_y": n_y, "n_str": n_str, "dups": dups}
    if want_values:
        res["tabs"] = out_tabs
    return res

def count_one(path):
    try:
        r = stream_file(path, False)
        r["ok"] = True
        return r
    except Exception as e:
        return {"file": os.path.basename(path), "ok": False, "err": str(e)[:200]}

def do_count(args):
    files = sorted([str(p) for p in SRC.glob("*.xlsx")])
    print(f"[count] {len(files)} xlsx, workers={args.workers}", flush=True)
    with Pool(args.workers) as pool:
        res = pool.map(count_one, files, chunksize=8)
    ok = [r for r in res if r.get("ok")]
    bad = [r for r in res if not r.get("ok")]
    print(f"[count] ok={len(ok)} bad={len(bad)}", flush=True)
    for r in bad[:20]:
        print("  BAD", r["file"], r["err"], flush=True)
    WORK.mkdir(parents=True, exist_ok=True)
    (WORK / "counts.json").write_text(json.dumps(res, indent=1))
    tot_g = sum(r["n_g"] for r in ok)
    tot_y = sum(r["n_y"] for r in ok)
    print(f"[count] total numeric G={tot_g} Y={tot_y}", flush=True)
    return res

def do_select(args):
    res = json.loads((WORK / "counts.json").read_text())
    by_ss = collections.defaultdict(list)
    for r in res:
        if not r.get("ok"):
            continue
        sym, side = sym_side_of(r["file"])
        if sym is None:
            continue
        r["cells"] = r["n_g"] + r["n_y"]
        by_ss[(sym, side)].append(r)
    sel = {}
    for ss, cands in by_ss.items():
        cands.sort(key=lambda r: r["mtime"], reverse=True)
        best = max(r["cells"] for r in cands)
        floor = 0.5 * best if best > 0 else 0
        pick = next((r for r in cands if r["cells"] >= floor), cands[0])
        sel[f"{ss[0]}_{ss[1]}"] = {"file": pick["file"], "mtime": pick["mtime"], "size": pick["size"],
            "n_g": pick["n_g"], "n_y": pick["n_y"], "floor": floor, "best": best,
            "ncands": len(cands), "cat": cat_side_of(ss[0], ss[1])}
    cats = collections.Counter(v["cat"] for v in sel.values())
    print(f"[select] sym_sides={len(sel)} cats={dict(cats)}", flush=True)
    zeros = [k for k, v in sel.items() if v["n_g"] + v["n_y"] == 0]
    print(f"[select] selected-with-zero-cells={len(zeros)} {zeros[:10]}", flush=True)
    (WORK / "selection.json").write_text(json.dumps(sel, indent=1))
    return sel

def extract_one(item):
    ss, info = item
    try:
        r = stream_file(str(SRC / info["file"]), True)
        r["ss"] = ss
        r["cat"] = info["cat"]
        p = PARTIALS / f"{ss}.json"
        p.write_text(json.dumps(r, separators=(",", ":")))
        return {"ss": ss, "ok": True, "n_g": r["n_g"], "n_y": r["n_y"]}
    except Exception as e:
        return {"ss": ss, "ok": False, "err": str(e)[:200]}

def do_extract(args):
    sel = json.loads((WORK / "selection.json").read_text())
    PARTIALS.mkdir(parents=True, exist_ok=True)
    items = sorted(sel.items())
    if args.only:
        only = set(args.only.split(","))
        items = [(k, v) for k, v in items if k in only]
    print(f"[extract] {len(items)} files, workers={args.workers}", flush=True)
    with Pool(args.workers) as pool:
        res = pool.map(extract_one, items, chunksize=4)
    ok = [r for r in res if r.get("ok")]
    bad = [r for r in res if not r.get("ok")]
    print(f"[extract] ok={len(ok)} bad={len(bad)}", flush=True)
    for r in bad[:20]:
        print("  BAD", r["ss"], r.get("err"), flush=True)
    return res

def template_rowkeys(tpl_path):
    import openpyxl
    wb = openpyxl.load_workbook(tpl_path, read_only=True, data_only=True)
    keys = {}
    try:
        for tab in TABS12:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            it = ws.iter_rows(values_only=True)
            next(it, None)
            row2 = list(next(it, None) or [])
            hmap = {}
            for i, v in enumerate(row2[:14]):
                if isinstance(v, str) and v.strip():
                    hmap[v.strip()] = i
            acol = hmap.get("Switch", 0)
            bcol = hmap.get("default", 1)
            tk = set()
            for row in it:
                aval = row[acol] if acol < len(row) else None
                if aval is None or (isinstance(aval, str) and not aval.strip()):
                    continue
                tk.add((str(aval).strip(), norm_cand(row[bcol] if bcol < len(row) else None)))
            keys[tab] = tk
    finally:
        wb.close()
    return keys

def pick_template(cat, partial_keys):
    cands = [ROOT / f"SPREADSHEETS/TEMPLATE_{cat}.xlsx",
        ROOT / "SPREADSHEETS/TEMPLATE_FINAL_NORM" / f"TEMPLATE_{cat}.xlsx"]
    best = None
    for c in cands:
        if not c.exists():
            continue
        tk = template_rowkeys(str(c))
        hit = sum(1 for (t, s, d) in partial_keys if (s, d) in tk.get(t, set()))
        tot = len(partial_keys)
        print(f"[build] {cat} {c.parent.name}/{c.name}: rowkey hit {hit}/{tot} ({100.0*hit/max(1,tot):.1f}%)", flush=True)
        if best is None or hit > best[1]:
            best = (str(c), hit)
    return best[0]

def do_build(args):
    import openpyxl
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    sel = json.loads((WORK / "selection.json").read_text())
    by_cat = collections.defaultdict(list)
    for ss, info in sel.items():
        p = PARTIALS / f"{ss}.json"
        if p.exists():
            by_cat[info["cat"]].append(ss)
    OUTDIR.mkdir(parents=True, exist_ok=True)
    GREEN = PatternFill("solid", fgColor="006100")
    RED = PatternFill("solid", fgColor="FFC7CE")
    EXTRA_HDR_FILL = PatternFill("solid", fgColor="D9D1F2")
    for cat in CAT_SIDES:
        sss = by_cat.get(cat, [])
        if args.only:
            only = set(args.only.split(","))
            sss = [s for s in sss if s in only]
        if not sss:
            print(f"[build] {cat}: no partials, skip", flush=True)
            continue
        t0 = datetime.datetime.now().isoformat(timespec="seconds")
        print(f"[build] {cat}: {len(sss)} sym_sides", flush=True)
        row_acc = {}
        cell_acc = {}
        pkeys = set()
        for ss in sss:
            d = json.loads((PARTIALS / f"{ss}.json").read_text())
            for tab, td in d["tabs"].items():
                hdrs = td["headers"]
                for sw, cand, g, yvals in td["rows"]:
                    rk = (tab, sw, cand)
                    pkeys.add(rk)
                    if g is not None:
                        a = row_acc.setdefault(rk, [0.0, 0, 0])
                        a[0] += g
                        a[1] += 1
                        if g > POS_EPS:
                            a[2] += 1
                    for j, y in yvals:
                        ck = (tab, sw, cand, hdrs[j])
                        a = cell_acc.setdefault(ck, [0.0, 0, 0])
                        a[0] += y
                        a[1] += 1
                        if y > POS_EPS:
                            a[2] += 1
        print(f"[build] {cat}: rowkeys={len(row_acc)} cellkeys={len(cell_acc)}", flush=True)
        tpl_path = pick_template(cat, pkeys)
        twb = openpyxl.load_workbook(tpl_path)
        nwb = openpyxl.Workbook()
        nwb.remove(nwb.active)
        cov = nwb.create_sheet("COVER")
        cov["A1"] = f"{cat} — avg/pos_sym matrix inventory"
        cov["A2"] = f"source template: {tpl_path} (picked by rowkey hit rate)"
        cov["A3"] = f"built: {t0} local; sym_sides={len(sss)}; rowkeys={len(row_acc)}; cellkeys={len(cell_acc)}"
        cov["A4"] = "semantics: base sheet=TEMPLATE tab AVG (yellow cells=avg_delta, M=row avg of VECTOR_DELTA, N=row pos_sym); _POS sheet=same grid with pos_sym counts; _N sheet=numeric-value counts. blank=never numeric in any sym_side (never tested/skipped/invalid-text). pos=delta>1e-9. avg=mean over numerics incl exact zeros."
        cov["A5"] = "selection: ONE newest xlsx per sym_side from SPREADSHEETS/V15_V16_CELL_BY_CELL with >=50% of that sym_side's best numeric-cell count (else next newest)."
        try:
            uni = json.loads((ROOT / "data" / "tradeable_universe_20261006.json").read_text())
            trad = []
            for s in uni.get("crypto_sym_sides", []) + uni.get("stocks_sym_sides", []):
                sym, side = s.rsplit("_", 1)
                if cat_side_of(sym, side) == cat:
                    trad.append(s)
            have = set(sss)
            missing = sorted(set(trad) - have)
            extra = sorted(have - set(trad))
            cov["A6"] = f"tradeable coverage: {len(set(trad) & have)}/{len(trad)} tradeable sym_sides included; missing={missing if missing else 'NONE'}; non-tradeable extras included={len(extra)}"
        except Exception as e:
            cov["A6"] = f"tradeable coverage: universe file unreadable ({str(e)[:120]})"
        cov.append([])
        cov.append(["sym_side", "file", "mtime_utc", "size", "num_G", "num_Y"])
        for ss in sorted(sss):
            info = sel[ss]
            cov.append([ss, info["file"], datetime.datetime.fromtimestamp(info["mtime"], datetime.timezone.utc).isoformat(), info["size"], info["n_g"], info["n_y"]])
        for col in range(1, 7):
            cov.column_dimensions[get_column_letter(col)].width = 34 if col == 2 else 16
        unmatched = []
        for tab in TABS12:
            if tab not in twb.sheetnames:
                unmatched.append([tab, "TAB MISSING IN TEMPLATE", "", "", 0, 0])
                continue
            tws = twb[tab]
            maxc = tws.max_column
            maxr = tws.max_row
            hdr2 = [tws.cell(row=2, column=c).value for c in range(1, maxc + 1)]
            yhdr_tpl = []
            for c in range(15, maxc + 1):
                h = hdr2[c - 1]
                if h is None or (isinstance(h, str) and not h.strip()):
                    continue
                hs = str(h).strip()
                if hs.startswith("WHAT SWITCH"):
                    break
                yhdr_tpl.append(hs)
            extra_hdrs = sorted({ck[3] for ck in cell_acc if ck[0] == tab} - set(yhdr_tpl))
            out_hdrs = yhdr_tpl + extra_hdrs
            trows = []
            for r in range(3, maxr + 1):
                a = tws.cell(row=r, column=1).value
                if a is None or (isinstance(a, str) and not a.strip()):
                    continue
                sw = str(a).strip()
                b = tws.cell(row=r, column=2).value
                trows.append({"r": r, "sw": sw, "cand": norm_cand(b), "b": b,
                    "d": tws.cell(row=r, column=4).value, "l": tws.cell(row=r, column=12).value,
                    "afill": tws.cell(row=r, column=1).fill.copy(),
                    "bbold": bool(tws.cell(row=r, column=2).font.bold)})
            tpl_keys = {(x["sw"], x["cand"]) for x in trows}
            extra_rows = sorted({(ck[1], ck[2]) for ck in list(cell_acc) + [(t, s, d, "") for (t, s, d) in row_acc if t == tab] if ck[0] == tab} - tpl_keys)
            n_unseen = sum(1 for x in trows if (tab, x["sw"], x["cand"]) not in row_acc)
            unmatched.append([tab, f"tpl_rows={len(trows)} seen={len(trows)-n_unseen} unseen={n_unseen} extra_rows={len(extra_rows)} tpl_yhdr={len(yhdr_tpl)} extra_yhdr={len(extra_hdrs)}", "", "", 0, 0])
            for x in trows:
                if (tab, x["sw"], x["cand"]) not in row_acc:
                    unmatched.append([tab, "TPL_ROW_NEVER_SEEN", x["sw"], x["cand"], 0, 0])
            for sw, cand in extra_rows:
                unmatched.append([tab, "EXTRA_ROW", sw, cand, 0, 0])
            for h in extra_hdrs:
                unmatched.append([tab, "EXTRA_HDR", h, "", 0, 0])
            widths = {c: tws.column_dimensions[get_column_letter(c)].width for c in range(1, maxc + 1)}
            r2fills = [tws.cell(row=2, column=c).fill.copy() for c in range(1, maxc + 1)]
            r2fonts = [tws.cell(row=2, column=c).font.copy() for c in range(1, maxc + 1)]
            allrows = [(x["sw"], x["cand"], x["b"], x["d"], x["l"], x["afill"], x["bbold"], False) for x in trows]
            if extra_rows:
                allrows.append((None, None, None, None, None, None, None, None))
                for sw, cand in extra_rows:
                    allrows.append((sw, cand, cand, "(EXTRA: in sym files, not in template)", "", None, False, True))
            for variant, vname in (("avg", tab), ("pos", tab + "_POS"), ("n", tab + "_N")):
                ws = nwb.create_sheet(vname)
                ws.freeze_panes = tws.freeze_panes
                ws.sheet_properties.tabColor = tws.sheet_properties.tabColor
                for c in range(1, 15):
                    w = widths.get(c)
                    if w:
                        ws.column_dimensions[get_column_letter(c)].width = w
                for j in range(len(out_hdrs)):
                    c = 15 + j
                    w = widths.get(c) if c <= maxc else 14
                    ws.column_dimensions[get_column_letter(c)].width = w or 14
                for c in range(1, 15):
                    v = hdr2[c - 1] if c - 1 < len(hdr2) else None
                    cell = ws.cell(row=2, column=c, value=v)
                    if c - 1 < len(r2fills):
                        cell.fill = r2fills[c - 1]
                        cell.font = r2fonts[c - 1]
                for j, h in enumerate(out_hdrs):
                    cell = ws.cell(row=2, column=15 + j, value=h)
                    if j < len(yhdr_tpl) and 14 + j < len(r2fills):
                        cell.fill = r2fills[14 + j]
                        cell.font = r2fonts[14 + j]
                    else:
                        cell.fill = EXTRA_HDR_FILL
                r1 = [tws.cell(row=1, column=c).value for c in range(1, 15)]
                for c, v in enumerate(r1, start=1):
                    if v is not None:
                        ws.cell(row=1, column=c, value=v)
                wr = 3
                for sw, cand, b, d, l, afill, bbold, is_extra in allrows:
                    if sw is None:
                        ws.cell(row=wr, column=1, value="EXTRA ROWS (in sym files, not in template)").fill = EXTRA_HDR_FILL
                        wr += 1
                        continue
                    rk = (tab, sw, cand)
                    ra = row_acc.get(rk)
                    ws.cell(row=wr, column=1, value=sw)
                    if afill is not None:
                        ws.cell(row=wr, column=1).fill = afill
                    bc = ws.cell(row=wr, column=2, value=b)
                    if bbold:
                        f = bc.font.copy()
                        f.bold = True
                        bc.font = f
                    ws.cell(row=wr, column=4, value=d)
                    ws.cell(row=wr, column=12, value=l)
                    if ra is not None:
                        avg = ra[0] / ra[1]
                        mcell = ws.cell(row=wr, column=13, value=avg)
                        ncell = ws.cell(row=wr, column=14, value=ra[2])
                        mcell.fill = GREEN if avg > POS_EPS else (RED if avg < -POS_EPS else PatternFill())
                        ncell.fill = GREEN if ra[2] > 0 else RED
                    for j, h in enumerate(out_hdrs):
                        a = cell_acc.get((tab, sw, cand, h))
                        if a is None:
                            continue
                        if variant == "avg":
                            v = a[0] / a[1]
                            cell = ws.cell(row=wr, column=15 + j, value=v)
                            cell.fill = GREEN if v > POS_EPS else (RED if v < -POS_EPS else PatternFill())
                        elif variant == "pos":
                            cell = ws.cell(row=wr, column=15 + j, value=a[2])
                            cell.fill = GREEN if a[2] > 0 else RED
                        else:
                            ws.cell(row=wr, column=15 + j, value=a[1])
                    wr += 1
        un = nwb.create_sheet("UNMATCHED")
        un.append(["tab", "kind", "a", "b", "c", "d"])
        for row in unmatched:
            un.append(row)
        for col in range(1, 5):
            un.column_dimensions[get_column_letter(col)].width = 42 if col >= 3 else 26
        twb.close()
        out = OUTDIR / f"AVG_{cat}.xlsx"
        tmp = OUTDIR / f"AVG_{cat}.tmp.xlsx"
        nwb.save(tmp)
        nwb.close()
        import zipfile
        zf = zipfile.ZipFile(tmp)
        assert len(zf.namelist()) >= 10, f"truncated save? {len(zf.namelist())} entries"
        zf.close()
        os.replace(tmp, out)
        print(f"[build] {cat} -> {out} ({out.stat().st_size/1e6:.1f} MB)", flush=True)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workers", type=int, default=10)
    ap.add_argument("--only", default=None, help="comma sym_sides (extract filter / build filter)")
    ap.add_argument("--steps", default="count,select,extract,build")
    a = ap.parse_args()
    steps = a.steps.split(",")
    if "count" in steps:
        do_count(a)
    if "select" in steps:
        do_select(a)
    if "extract" in steps:
        do_extract(a)
    if "build" in steps:
        do_build(a)

if __name__ == "__main__":
    main()
