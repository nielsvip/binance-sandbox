#!/usr/bin/env python3
"""v15_template_addonly_k — STAGED add-only update on top of the RESTORED pre-trainwreck templates (Agent K, USER 2026-10-01).

Base  = live SPREADSHEETS/TEMPLATE_*.xlsx (user's clean 20:36 set). Output = SPREADSHEETS/TEMPLATE_STAGED/<ts>_K/TEMPLATE_*.xlsx. Never writes live files.
Changes (nothing else): (1) real missing switches (Agent G REVIVE / REVIVE_BATCH1 / PRESENT-but-absent-in-base, rows copied whole from G's staged
set 202610010107, same column layout) appended after the last white row of their tab (white above orange stays); orange rows for filters entirely
absent from a tab's orange block are appended at the end; (2) the 5 LG-10 crypto exits from data/wiring/template_row_spec_lg10.json; (3) every tab starts
with a default row (whole-row move of the first group's default row to the top of the group).
Guarantee: every original row image (all columns: value+fill) is present unchanged in the staged tab (template_row_guard fingerprints, subset check).
NO promotions / bold changes / sorting / yellow repaint / option edits of existing rows.
"""
import argparse
import collections
import copy
import csv
import datetime
import json
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from v15_template_restructure import snapshot_row, write_row, is_orange  # noqa: E402
import template_row_guard as GUARD  # noqa: E402

CATS = ["CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"]
SW = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
HDR = 2
ALLOWED = ("REVIVE", "REVIVE_BATCH1", "PRESENT")
COL_L = 11  # is_default (0-based) — verified by header below


def col_idx(ws, name):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=HDR, column=c).value
        if isinstance(v, str) and v.strip().upper() == name.upper():
            return c - 1
    raise RuntimeError(f"{ws.title}: no header {name}")


def is_yes(cells, c_isd):
    return str(cells[c_isd][0] or "").strip().upper() == "YES"


def bold(font, b):
    return Font(name=(font.name if font else None) or "Arial", size=(font.size if font else None) or 10, bold=b, italic=font.italic if font else None, color=font.color if font else None)


def lg10_rows(template_snap, ws, spec_row, crypto_default):
    """white rows for one LG-10 exit, cloned from an existing white bool row's styles."""
    c_isd = col_idx(ws, "is_default")
    out = []
    for opt in spec_row["options"]:
        cells = list(template_snap)
        for i, (v, f, fl, nf, al) in enumerate(cells):
            if i >= 14:
                cells[i] = (None, f, PatternFill(fill_type=None), nf, al)  # no yellow on a new row: pilot's name-token rule decides
            else:
                cells[i] = (None, f, fl, nf, al)
        isd = bool(opt) == bool(crypto_default)
        v0, f0, fl0, nf0, al0 = cells[0]
        cells[0] = (spec_row["switch"], f0, fl0, nf0, al0)
        v1, f1, fl1, nf1, al1 = cells[1]
        cells[1] = (bool(opt), bold(f1, isd), fl1, nf1, al1)
        v3, f3, fl3, nf3, al3 = cells[3]
        cells[3] = (spec_row.get("family") or "EXIT", f3, fl3, nf3, al3)
        vL, fL, flL, nfL, alL = cells[c_isd]
        cells[c_isd] = ("YES" if isd else "NO", Font(name="Arial", size=10, bold=isd), flL, nfL, alL)
        out.append(cells)
    return out


def build(cs, ts_dir, report):
    R = openpyxl.load_workbook(ROOT / "SPREADSHEETS" / f"TEMPLATE_{cs}.xlsx")
    base = openpyxl.load_workbook(ROOT / "SPREADSHEETS" / f"TEMPLATE_{cs}.xlsx")  # pristine copy for the integrity check
    G = openpyxl.load_workbook(ROOT / "SPREADSHEETS" / "TEMPLATE_STAGED" / "202610010107" / f"TEMPLATE_{cs}.xlsx")
    dec = {r["name"]: r["decision"] for r in csv.DictReader(open(ROOT / "data" / "switch_revival" / "202610010107" / f"decisions_{cs}.csv"))}
    spec = json.load(open(ROOT / "data" / "wiring" / "template_row_spec_lg10.json"))["rows"] if cs.startswith("CRYPTO") else []
    rwhite, rorange = set(), collections.defaultdict(set)
    for t in SW:
        ws = R[t]
        for r in range(HDR + 1, ws.max_row + 1):
            a = ws.cell(row=r, column=1).value
            if a in (None, ""):
                continue
            if is_orange(ws, r):
                rorange[t].add(str(a).strip())
            else:
                rwhite.add(str(a).strip())
    rep = {"added_white": collections.defaultdict(list), "added_orange": collections.defaultdict(list), "lg10": [], "first_row_moved": [], "first_row_no_default": [], "skipped_decisions": collections.Counter()}
    for t in SW:
        ws, wg = R[t], G[t]
        ncol = ws.max_column
        c_isd = col_idx(ws, "is_default")
        snaps, kinds = [], []
        for r in range(HDR + 1, ws.max_row + 1):
            snaps.append(snapshot_row(ws, r, ncol))
            a = ws.cell(row=r, column=1).value
            kinds.append("blank" if a in (None, "") else ("orange" if is_orange(ws, r) else "white"))
        # new white rows from G (grouped by switch name, G order)
        new_w, order = collections.OrderedDict(), []
        new_o = collections.OrderedDict()
        for r in range(HDR + 1, wg.max_row + 1):
            a = wg.cell(row=r, column=1).value
            if a in (None, ""):
                continue
            a = str(a).strip()
            if is_orange(wg, r):
                if a not in rorange[t] and (dec.get(a) in ALLOWED or a in rwhite):
                    new_o.setdefault(a, []).append(snapshot_row(wg, r, ncol))
            else:
                d = dec.get(a, "NODEC")
                if a in rwhite:
                    continue
                if d in ALLOWED:
                    new_w.setdefault(a, []).append(snapshot_row(wg, r, ncol))
                else:
                    rep["skipped_decisions"][d] += 1
        # LG-10 exits (crypto only)
        lg = [s for s in spec if s["tab"] == t and s["switch"] not in rwhite and s["switch"] not in new_w]
        if lg:
            tmpl = next((snaps[i] for i, k in enumerate(kinds) if k == "white" and isinstance(snaps[i][1][0], bool)), None)
            if tmpl is None:
                tmpl = next(snaps[i] for i, k in enumerate(kinds) if k == "white")
            for s in lg:
                new_w[s["switch"]] = lg10_rows(tmpl, ws, s, s.get("default_crypto", True))
                rep["lg10"].append((t, s["switch"]))
        i0 = next((i for i, k in enumerate(kinds) if k == "orange"), len(snaps))
        flat_w = [row for rows in new_w.values() for row in rows]
        flat_o = [row for rows in new_o.values() for row in rows]
        final = snaps[:i0] + flat_w + snaps[i0:] + flat_o
        for n, rows in new_w.items():
            rep["added_white"][t].append((n, len(rows)))
        for n, rows in new_o.items():
            rep["added_orange"][t].append((n, len(rows)))
        # default-first: whole-row move inside the first group
        data_idx = [i for i, s in enumerate(final) if str(s[0][0] or "").strip() != ""]
        if data_idx:
            f0 = data_idx[0]
            if not is_yes(final[f0], c_isd):
                name = final[f0][0][0]
                grp = [i for i in data_idx if final[i][0][0] == name and (i - f0) < 10000]
                # contiguous group starting at f0
                g = []
                for i in range(f0, len(final)):
                    if final[i][0][0] == name:
                        g.append(i)
                    else:
                        break
                yes = [i for i in g if is_yes(final[i], c_isd)]
                if yes:
                    row = final.pop(yes[0])
                    final.insert(f0, row)
                    rep["first_row_moved"].append((t, str(name)))
                else:
                    rep["first_row_no_default"].append((t, str(name)))
        for k, cells in enumerate(final):
            write_row(ws, HDR + 1 + k, cells)
        # integrity: every original row image present
        a = GUARD.fingerprints(base[t])
        b = GUARD.fingerprints(ws)
        lost = a - b
        if lost:
            raise AssertionError(f"{cs}/{t}: {sum(lost.values())} original row images lost/altered")
    out = ts_dir / f"TEMPLATE_{cs}.xlsx"
    tmp = ts_dir / f"TEMPLATE_{cs}.tmp.xlsx"
    R.save(str(tmp))
    openpyxl.load_workbook(str(tmp))
    os.replace(tmp, out)
    report[cs] = {"added_white": {t: v for t, v in rep["added_white"].items()}, "added_orange": {t: v for t, v in rep["added_orange"].items()}, "lg10": rep["lg10"],
                  "first_row_moved": rep["first_row_moved"], "first_row_no_default": rep["first_row_no_default"], "skipped_decisions": dict(rep["skipped_decisions"])}
    nw = sum(len(v) for v in rep["added_white"].values())
    nrw = sum(c for v in rep["added_white"].values() for _, c in v)
    nro = sum(c for v in rep["added_orange"].values() for _, c in v)
    print(f"[{cs}] added white switches={nw} rows={nrw} | orange rows={nro} | lg10={len(rep['lg10'])} | first-row moved={len(rep['first_row_moved'])} no-default={rep['first_row_no_default']}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cat-side", default="ALL")
    a = ap.parse_args()
    ts = datetime.datetime.utcnow().strftime("%Y%m%d%H%M") + "_K"
    out = ROOT / "SPREADSHEETS" / "TEMPLATE_STAGED" / ts
    out.mkdir(parents=True, exist_ok=True)
    report = {"ts": ts}
    for cs in (CATS if a.cat_side == "ALL" else [a.cat_side]):
        build(cs, out, report)
    rd = ROOT / "data" / "template_audit" / ts
    rd.mkdir(parents=True, exist_ok=True)
    (rd / "addonly_report.json").write_text(json.dumps(report, indent=1, default=str))
    print("[staged]", out, "[report]", rd / "addonly_report.json")


if __name__ == "__main__":
    main()
