"""Audit every matrix xlsx under SPREADSHEETS (Mac, parallel, read-only).

Checks per file: zip integrity, size, published-filename bh/gain parse,
tab coverage vs the 13 SWITCH_SHEETS, E2/E3 baseline cells, C/F/G/H/I fill
counts, first-row yellow coverage, BASELINE_METRICS B2. Writes JSON only.
"""
import glob
import json
import os
import re
import sys
import zipfile
from multiprocessing import Pool

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(ROOT, "data", "reports", "s5_verify_20261003")
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
PUB_RE = re.compile(r"_bh(m?)([0-9p]+)_gain(m?)([0-9p]+)_")


def num(s):
    return float(s.replace("p", ".").replace("m", "-"))


def audit_one(path):
    import openpyxl
    rec = {"file": os.path.relpath(path, ROOT)}
    try:
        rec["size_kb"] = round(os.path.getsize(path) / 1024.0, 1)
        rec["mtime"] = int(os.path.getmtime(path))
    except OSError as e:
        rec["audit_error"] = f"vanished_or_unreadable: {e}"[:120]
        return rec
    m = PUB_RE.search(os.path.basename(path))
    if m:
        try:
            rec["published"] = True
            rec["fname_bh"] = num(m.group(1) + m.group(2))
            rec["fname_gain"] = num(m.group(3) + m.group(4))
        except ValueError:
            rec["published"] = True
            rec["fname_parse"] = "fail"
    else:
        rec["published"] = False
    try:
        with zipfile.ZipFile(path) as z:
            rec["zip_entries"] = len(z.namelist())
    except Exception as e:
        rec["zip_error"] = f"{type(e).__name__}: {e}"[:120]
        return rec
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as e:
        rec["open_error"] = f"{type(e).__name__}: {e}"[:120]
        return rec
    try:
        names = wb.sheetnames
        rec["tabs_present"] = [t for t in SWITCH_SHEETS if t in names]
        rec["tabs_missing"] = [t for t in SWITCH_SHEETS if t not in names]
        rec["baseline_metrics_tab"] = next((s for s in names if "BASELINE_METRICS" in s), None)
        tabs = {}
        for t in rec["tabs_present"]:
            ws = wb[t]
            try:
                e2 = ws["E2"].value
                e3 = ws["E3"].value
                a1 = ws["A2"].value
            except Exception:
                e2 = e3 = a1 = "read-err"
            f_n = g_n = h_n = i_n = c_n = a_n = 0
            y_num = y_none = y_str = 0
            y_sampled = False
            try:
                for row in ws.iter_rows(min_row=3, values_only=True):
                    a = row[0] if len(row) > 0 else None
                    if a is None or a == "":
                        continue
                    a_n += 1
                    c = row[2] if len(row) > 2 else None
                    f = row[5] if len(row) > 5 else None
                    g = row[6] if len(row) > 6 else None
                    h = row[7] if len(row) > 7 else None
                    iv = row[8] if len(row) > 8 else None
                    if c is not None and c != "":
                        c_n += 1
                    if f is not None:
                        f_n += 1
                    if g is not None:
                        g_n += 1
                    if h is not None:
                        h_n += 1
                    if iv is not None:
                        i_n += 1
                    if not y_sampled and len(row) > 14:
                        for y in row[14:261]:
                            if y is None:
                                y_none += 1
                            elif isinstance(y, (int, float)):
                                y_num += 1
                            else:
                                y_str += 1
                        y_sampled = True
            except Exception as e:
                tabs[t] = {"rowscan_error": f"{type(e).__name__}"[:60]}
                continue
            tabs[t] = {"e2": (e2 if isinstance(e2, (int, float, str)) else str(type(e2).__name__)), "e3": (e3 if isinstance(e3, (int, float, str)) else str(type(e3).__name__)), "a2": (a1 if isinstance(a1, str) else str(a1)[:20]), "rows": a_n, "c": c_n, "f": f_n, "g": g_n, "h": h_n, "i": i_n, "y_num": y_num, "y_none": y_none, "y_str": y_str}
        rec["tabs"] = tabs
        if rec["baseline_metrics_tab"]:
            try:
                ws = wb[rec["baseline_metrics_tab"]]
                rec["metrics_b2"] = ws["B2"].value if isinstance(ws["B2"].value, (int, float, str)) else str(type(ws["B2"].value).__name__)
                kv = {}
                for row in ws.iter_rows(min_row=1, max_row=40, max_col=2, values_only=True):
                    if row[0] and isinstance(row[0], str) and row[0].startswith("final_"):
                        kv[row[0]] = row[1]
                rec["metrics_final"] = kv
            except Exception as e:
                rec["metrics_error"] = f"{type(e).__name__}"[:60]
    finally:
        wb.close()
    return rec


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    files = sorted(glob.glob(os.path.join(ROOT, "SPREADSHEETS", "V15_V16_CELL_BY_CELL", "*.xlsx")))
    files += sorted(glob.glob(os.path.join(ROOT, "SPREADSHEETS", "*matrix*.xlsx")))
    files = sorted(set(files))
    workers = min(8, (os.cpu_count() or 8))
    with Pool(workers) as pool:
        recs = pool.map(audit_one, files)
    out = os.path.join(OUTDIR, "sheet_audit.json")
    tmp = out + ".tmp"
    with open(tmp, "w") as f:
        json.dump({"n": len(recs), "files": recs}, f)
    os.replace(tmp, out)
    bad = sum(1 for r in recs if r.get("zip_error") or r.get("open_error"))
    pub = sum(1 for r in recs if r.get("published"))
    print(f"audited={len(recs)} published={pub} zip_or_open_errors={bad} -> {out}")


if __name__ == "__main__":
    sys.exit(main())
