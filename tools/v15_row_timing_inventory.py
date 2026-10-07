#!/usr/bin/env python3
"""Per-row / per-filter timing inventory from the pilots' delta logs (director, 2026-10-06, USER request).

Every v15_pilot evaluation is one line in {progress_dir}/v15_delta_log/{SS}_{nav}.jsonl with `secs` (worker compute time),
`cached`, the TEMPLATE row (sheet,row,switch,cand) and `label`: label == "switch=cand" is the row's NAKED switch eval, any
other "KEY=VAL" label is a YELLOW filter cell (that filter alone on top of the switch). `ROW_DONE` closes a row; the gap
between consecutive ROW_DONE timestamps of the same file is the row's WALL time (all its evals incl. pool wait/queueing).

  scan (on a server):  python3 tools/v15_row_timing_inventory.py scan OUT.json DIR [DIR ...]
  merge (on the Mac):  python3 tools/v15_row_timing_inventory.py merge OUT_PREFIX IN.json [IN.json ...]
    -> OUT_PREFIX.xlsx (ROWS / FILTERS / SYM_SIDES / SLOWEST), OUT_PREFIX.md
"""
import datetime as dt
import glob
import json
import os
import statistics
import sys
from collections import defaultdict


def _ts(v):
    try:
        return dt.datetime.fromisoformat(str(v).replace("Z", "+00:00")).timestamp()
    except Exception:
        try:
            return float(v)
        except Exception:
            return None


def cat_side(ss):
    sym, _, side = ss.rpartition("_")
    return ("CRYPTO" if sym.endswith(("USDT", "USDC")) else "STOCKS") + "_" + side


def scan(out, dirs):
    rows = defaultdict(lambda: {"walls": [], "naked": [], "yellow": [], "n_yellow_rows": [], "cached": 0, "evals": 0, "err": 0, "ss": set()})
    filt = defaultdict(lambda: {"secs": [], "cached": 0, "evals": 0, "err": 0, "rows": set(), "ss": set()})
    syms = {}
    for d in dirs:
        for f in sorted(glob.glob(os.path.join(os.path.expanduser(d), "*.jsonl"))):
            base = os.path.basename(f)[:-6]
            ss = base.rsplit("_", 1)[0]
            cs = cat_side(ss)
            last_done, cur_yel = None, defaultdict(int)
            s = {"cat_side": cs, "file": f, "rows_done": 0, "evals": 0, "sum_secs": 0.0, "first": None, "last": None, "cached": 0}
            for line in open(f, errors="ignore"):
                try:
                    r = json.loads(line)
                except Exception:
                    continue
                if r.get("nav") not in (None, "sequential") and r.get("label") != "ROW_DONE":
                    continue
                t = _ts(r.get("ts"))
                if t is not None:
                    s["first"] = t if s["first"] is None else min(s["first"], t)
                    s["last"] = t if s["last"] is None else max(s["last"], t)
                sheet, row, sw, cand, lab = r.get("sheet"), r.get("row"), r.get("switch"), str(r.get("cand", "")), str(r.get("label", ""))
                key = (cs, sheet, row, sw, cand)
                if lab == "ROW_DONE":
                    if r.get("row_secs") is not None:
                        rows[key]["walls"].append(float(r["row_secs"]))
                    elif last_done is not None and t is not None:
                        rows[key]["walls"].append(t - last_done)
                    rows[key]["n_yellow_rows"].append(cur_yel.pop(key, 0))
                    rows[key]["ss"].add(ss)
                    last_done = t if t is not None else last_done
                    s["rows_done"] += 1
                    continue
                secs = r.get("secs")
                if secs is None:
                    continue
                s["evals"] += 1
                s["sum_secs"] += float(secs)
                c = bool(r.get("cached"))
                s["cached"] += c
                R = rows[key]
                R["evals"] += 1
                R["cached"] += c
                R["err"] += bool(r.get("err"))
                if lab == f"{sw}={cand}" or ":" in lab or sheet == "DIAGNOSE_REPAIR":
                    R["naked"].append(float(secs))
                else:
                    R["yellow"].append(float(secs))
                    cur_yel[key] += 1
                    fk = lab.split("=", 1)[0]
                    F = filt[(cs, fk)]
                    F["secs"].append(float(secs))
                    F["cached"] += c
                    F["evals"] += 1
                    F["err"] += bool(r.get("err"))
                    F["rows"].add(f"{sheet}!{row}")
                    F["ss"].add(ss)
            syms[ss + "|" + f] = s
    def st(v):
        if not v:
            return {}
        v2 = sorted(v)
        return {"n": len(v), "mean": round(statistics.mean(v), 4), "median": round(statistics.median(v), 4), "p90": round(v2[int(0.9 * (len(v2) - 1))], 4), "max": round(v2[-1], 4), "sum": round(sum(v), 2)}
    res = {"rows": [], "filters": [], "sym_sides": list(syms.values())}
    for (cs, sheet, row, sw, cand), R in rows.items():
        res["rows"].append({"cat_side": cs, "sheet": sheet, "row": row, "switch": sw, "cand": cand, "n_symsides": len(R["ss"]), "wall": st(R["walls"]), "naked": st(R["naked"]), "yellow": st(R["yellow"]),
                            "yellows_per_row": round(statistics.mean(R["n_yellow_rows"]), 1) if R["n_yellow_rows"] else 0, "evals": R["evals"], "cached_share": round(R["cached"] / max(1, R["evals"]), 3), "errors": R["err"],
                            "_walls": R["walls"], "_naked": R["naked"], "_yellow": R["yellow"]})
    for (cs, fk), F in filt.items():
        res["filters"].append({"cat_side": cs, "filter": fk, "secs": st(F["secs"]), "evals": F["evals"], "cached_share": round(F["cached"] / max(1, F["evals"]), 3), "errors": F["err"], "n_rows": len(F["rows"]), "n_symsides": len(F["ss"]), "_secs": F["secs"]})
    json.dump(res, open(out, "w"))
    print(f"scan -> {out}: rows {len(res['rows'])} filters {len(res['filters'])} sym_side files {len(res['sym_sides'])}")


def merge(prefix, ins):
    import openpyxl
    from openpyxl.styles import Font
    rows, filt, syms = defaultdict(lambda: {"walls": [], "naked": [], "yellow": [], "ypr": [], "evals": 0, "cached": 0, "err": 0, "nss": 0}), defaultdict(lambda: {"secs": [], "evals": 0, "cached": 0, "err": 0, "rows": 0, "nss": 0}), []
    for p in ins:
        d = json.load(open(p))
        host = os.path.basename(p).split("_")[0]
        for r in d["rows"]:
            k = (r["cat_side"], r["sheet"], r["row"], r["switch"], r["cand"])
            R = rows[k]
            R["walls"] += r["_walls"]; R["naked"] += r["_naked"]; R["yellow"] += r["_yellow"]
            R["ypr"].append(r["yellows_per_row"]); R["evals"] += r["evals"]; R["cached"] += round(r["cached_share"] * r["evals"]); R["err"] += r["errors"]; R["nss"] += r["n_symsides"]
        for f in d["filters"]:
            F = filt[(f["cat_side"], f["filter"])]
            F["secs"] += f["_secs"]; F["evals"] += f["evals"]; F["cached"] += round(f["cached_share"] * f["evals"]); F["err"] += f["errors"]; F["rows"] = max(F["rows"], f["n_rows"]); F["nss"] += f["n_symsides"]
        for s in d["sym_sides"]:
            s["host"] = host
            syms.append(s)
    def md(v):
        return round(statistics.median(v), 3) if v else None
    def mean(v):
        return round(statistics.mean(v), 3) if v else None
    def p90(v):
        v = sorted(v)
        return round(v[int(0.9 * (len(v) - 1))], 3) if v else None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ROWS"
    hdr = ["cat_side", "tab", "row", "switch", "cand", "n_symsides", "row_wall_median_s", "row_wall_mean_s", "row_wall_p90_s", "row_wall_max_s", "naked_eval_median_s", "yellows_per_row", "yellow_eval_median_s", "yellow_eval_p90_s", "yellow_secs_sum_per_row_mean", "evals", "cached_share", "errors"]
    ws.append(hdr)
    out_rows = []
    for (cs, sheet, row, sw, cand), R in rows.items():
        ypr = mean(R["ypr"]) or 0
        out_rows.append([cs, sheet, row, sw, cand, R["nss"], md(R["walls"]), mean(R["walls"]), p90(R["walls"]), round(max(R["walls"]), 3) if R["walls"] else None, md(R["naked"]), ypr, md(R["yellow"]), p90(R["yellow"]),
                         round(sum(R["yellow"]) / max(1, R["nss"]), 2), R["evals"], round(R["cached"] / max(1, R["evals"]), 3), R["err"]])
    out_rows.sort(key=lambda x: (x[0], str(x[1]), x[2] if isinstance(x[2], int) else 0, str(x[4])))
    for r in out_rows:
        ws.append(r)
    ws2 = wb.create_sheet("FILTERS")
    ws2.append(["cat_side", "filter_switch (yellow)", "evals", "median_s", "mean_s", "p90_s", "max_s", "sum_s", "cached_share", "errors", "rows_it_appears_in", "sym_side_files"])
    out_f = []
    for (cs, fk), F in filt.items():
        v = F["secs"]
        out_f.append([cs, fk, F["evals"], md(v), mean(v), p90(v), round(max(v), 3) if v else None, round(sum(v), 1), round(F["cached"] / max(1, F["evals"]), 3), F["err"], F["rows"], F["nss"]])
    out_f.sort(key=lambda x: (x[0], -(x[7] or 0)))
    for r in out_f:
        ws2.append(r)
    ws3 = wb.create_sheet("SYM_SIDES")
    ws3.append(["host", "cat_side", "file", "rows_done", "evals", "sum_eval_secs", "cached", "wall_hours (first->last line)", "rows_per_hour"])
    for s in sorted(syms, key=lambda s: -(s.get("rows_done") or 0)):
        wall = ((s["last"] - s["first"]) / 3600) if s.get("first") and s.get("last") else None
        ws3.append([s["host"], s["cat_side"], s["file"], s["rows_done"], s["evals"], round(s["sum_secs"], 1), s["cached"], round(wall, 2) if wall else None, round(s["rows_done"] / wall, 1) if wall else None])
    ws4 = wb.create_sheet("SLOWEST_ROWS")
    ws4.append(hdr)
    for r in sorted(out_rows, key=lambda x: -(x[7] or 0))[:200]:
        ws4.append(r)
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
    wb.save(prefix + ".xlsx")
    # markdown summary
    lines = [f"# Row / filter timing inventory ({dt.datetime.utcnow():%Y-%m-%d %H:%MZ})", "", f"Sources: {', '.join(os.path.basename(p) for p in ins)}. `secs` = worker compute per evaluation; row wall = time between consecutive ROW_DONE markers (all evals of that row incl. pool queueing).", ""]
    for cs in sorted({r[0] for r in out_rows}):
        rr = [r for r in out_rows if r[0] == cs and r[7] is not None]
        if not rr:
            continue
        walls = [r[7] for r in rr]
        nrows = len({(r[1], r[2], r[4]) for r in rr})
        lines += [f"## {cs}", f"- template rows timed: {nrows}; row wall mean of means {statistics.mean(walls):.2f}s, median {statistics.median(walls):.2f}s → one full sheet ≈ {sum(statistics.median([x]) for x in walls)/3600:.2f} h of row wall (sum over rows of mean wall)",
                  f"- naked switch eval median {statistics.median([r[10] for r in rr if r[10] is not None] or [0]):.3f}s; yellow filter eval median {statistics.median([r[12] for r in rr if r[12] is not None] or [0]):.3f}s; yellows per row mean {statistics.mean([r[11] for r in rr]):.1f}",
                  "", "Slowest 15 rows (mean wall):", "", "| tab | row | switch=cand | wall mean s | yellows/row | yellow median s | evals |", "|---|---|---|---|---|---|---|"]
        for r in sorted(rr, key=lambda x: -x[7])[:15]:
            lines.append(f"| {r[1]} | {r[2]} | `{r[3]}={r[4]}` | {r[7]} | {r[11]} | {r[12]} | {r[15]} |")
        ff = sorted([f for f in out_f if f[0] == cs], key=lambda x: -(x[7] or 0))[:15]
        lines += ["", "Most expensive yellow filter_switches (total secs):", "", "| filter | evals | median s | p90 s | total s |", "|---|---|---|---|---|"] + [f"| `{f[1]}` | {f[2]} | {f[3]} | {f[5]} | {f[7]} |" for f in ff] + [""]
    lines += ["## Sym_side throughput", "", "| host | file | rows done | wall h | rows/h |", "|---|---|---|---|---|"]
    for s in sorted(syms, key=lambda s: -(s.get("rows_done") or 0))[:40]:
        wall = ((s["last"] - s["first"]) / 3600) if s.get("first") and s.get("last") else None
        lines.append(f"| {s['host']} | {os.path.basename(s['file'])} | {s['rows_done']} | {round(wall, 2) if wall else ''} | {round(s['rows_done'] / wall, 1) if wall else ''} |")
    open(prefix + ".md", "w").write("\n".join(lines) + "\n")
    print(f"merge -> {prefix}.xlsx / .md  rows {len(out_rows)} filters {len(out_f)} files {len(syms)}")


if __name__ == "__main__":
    if sys.argv[1] == "scan":
        scan(sys.argv[2], sys.argv[3:])
    else:
        merge(sys.argv[2], sys.argv[3:])
