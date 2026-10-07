#!/usr/bin/env python3
"""
replay_best_shard.py — replay BEST gain-encoded matrices with new templates
Divides 202 gain-encoded files into 4 shards, 56 workers per server.
Usage: python tools/replay_best_shard.py --shard 0 --total-shards 4 --workers 56 --window-days 30
On S1 (has 30d NPZ), verifies filename bh/gain matches engine replay with new templates.
Writes JSONL per shard to /tmp/replay_shard_<shard>.jsonl and summary to SPREADSHEETS/BEST/replay_report_shard_<shard>.json
"""
import argparse, pathlib, re, sys, json, time, os
from concurrent.futures import ThreadPoolExecutor, as_completed
ROOT = pathlib.Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
BEST = ROOT / "SPREADSHEETS" / "BEST"

def parse_gain(name):
    m=re.search(r"_gain([m]?)([\d]+)p([\d]+)", name)
    if not m: return None
    sign=-1 if m.group(1)=="m" else 1
    return sign*(int(m.group(2))+int(m.group(3))/100)
def parse_bh(name):
    m=re.search(r"_bh([m]?)([\d]+)p([\d]+)", name)
    if not m: return None
    sign=-1 if m.group(1)=="m" else 1
    return sign*(int(m.group(2))+int(m.group(3))/100)

def get_overrides_from_xlsx(xlsx_path):
    import openpyxl
    wb=openpyxl.load_workbook(str(xlsx_path), data_only=False)
    ov={}
    for ws in wb.worksheets:
        if ws.title in ("LEGEND_FILTERS","INSTRUCTIONS","FILTERS_EXPLAINED","INSTRUCTIONS_V2","TEMPLATE_BASELINE_METRICS","FILTER_DICTIONARY_V2","12SYM_PARITY","FORMULAS","Results_Deltas","Results_30d_Deltas","DAYTRADE","DISABLED_TOXIC_0918"):
            continue
        # header row 2 has Switch, default, override...
        for r in range(3, ws.max_row+1):
            k=ws.cell(r,1).value
            o=ws.cell(r,3).value
            if k is None or (isinstance(k,str) and not k.strip()): continue
            k=str(k).strip()
            if k.lower() in ("filter","option value","switch","sheets applicable"): continue
            if o not in (None,"","None"):
                # normalize booleans/numbers as string - keep original
                ov[k]=o
    # also check Results_Deltas for is_non_default? fallback already captured via sheets
    return ov

def replay_one(xlsx_path, window_days=30):
    try:
        fname=xlsx_path.name
        m=re.match(r"(.+)_bh", fname)
        if not m: return {"file": str(xlsx_path), "error": "no bh parse"}
        symside=m.group(1)
        symbol=symside.rsplit("_",1)[0]
        side=symside.rsplit("_",1)[1]
        filename_gain=parse_gain(fname)
        filename_bh=parse_bh(fname)
        is_incomplete="INCOMPLETE" in fname
        overrides=get_overrides_from_xlsx(xlsx_path)
        # check chart exists
        chart = xlsx_path.parent / f"{symside}_30D_REAL_ZOOMABLE.html"
        chart2 = xlsx_path.parent / f"{symside}_INCOMPLETE_30D_REAL_ZOOMABLE.html"
        chart_exists = chart.exists() or chart2.exists()
        # engine replay - try vector engine with new templates
        # Use get_defaults_for_symside + sanitize to ensure template defaults are new
        try:
            from tools.v15_pilot_sheet_runner import get_defaults_for_symside, sanitize_overrides, ensure_npz_for_symside, preload_prepared
            from tools.opt.v12_pilot import evaluate_sanitized, evaluate_prepared_sanitized
            defaults=get_defaults_for_symside(symside)
            clean_ov, warns=sanitize_overrides(dict(overrides), defaults)
            # ensure NPZ (will be no-op if exists)
            try:
                ensure_npz_for_symside(symside, window_days)
            except Exception as e:
                pass
            prepared=preload_prepared(symside, window_days)
            if prepared is not None:
                res=evaluate_prepared_sanitized(prepared, clean_ov, window_days=window_days)
            else:
                res=evaluate_sanitized(symside, clean_ov, window_days=window_days)
            eng_gain=res.get("gain_pct")
            eng_bh=res.get("bh_pct")
            eng_trades=res.get("trades")
            eng_valid=res.get("valid")
            eng_sharpe=res.get("pool_sharpe") or res.get("sharpe")
        except Exception as e:
            eng_gain=None
            eng_bh=None
            eng_trades=None
            eng_valid=False
            eng_sharpe=None
            res_error=str(e)[:500]
        else:
            res_error=None
        # compare within tolerance
        tol=0.5  # allow 0.5 pp rounding + slippage
        gain_match = (eng_gain is not None and filename_gain is not None and abs(eng_gain - filename_gain) <= tol)
        # also check bh match tolerance
        bh_match = (eng_bh is not None and filename_bh is not None and abs(eng_bh - filename_bh) <= tol) if eng_bh is not None else None
        out={
            "file": fname,
            "symside": symside,
            "shard_file": str(xlsx_path),
            "filename_bh": filename_bh,
            "filename_gain": filename_gain,
            "filename_incomplete": is_incomplete,
            "chart_exists": chart_exists,
            "overrides": len(overrides),
            "engine_gain": eng_gain,
            "engine_bh": eng_bh,
            "engine_trades": eng_trades,
            "engine_valid": eng_valid,
            "engine_sharpe": eng_sharpe,
            "gain_match": gain_match,
            "bh_match": bh_match,
            "gain_diff": (eng_gain - filename_gain) if (eng_gain is not None and filename_gain is not None) else None,
            "error": res_error,
        }
        return out
    except Exception as e:
        return {"file": str(xlsx_path), "error": str(e)[:800]}

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--shard", type=int, required=True)
    ap.add_argument("--total-shards", type=int, default=4)
    ap.add_argument("--workers", type=int, default=56)
    ap.add_argument("--window-days", type=int, default=30)
    ap.add_argument("--vector-only", action="store_true")
    args=ap.parse_args()
    gains=sorted(BEST.glob("*/*_gain*30d_matrix.xlsx"))
    print(f"[shard {args.shard}/{args.total_shards}] total gain-encoded {len(gains)}", flush=True)
    my=[p for i,p in enumerate(gains) if i % args.total_shards == args.shard]
    print(f"[shard {args.shard}] assigned {len(my)} files", flush=True)
    # ensure log dir
    out_jsonl=pathlib.Path(f"/tmp/replay_shard_{args.shard}.jsonl")
    out_report=BEST / f"replay_report_shard_{args.shard}.json"
    out_jsonl.parent.mkdir(parents=True, exist_ok=True)
    results=[]
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs={ex.submit(replay_one, p, args.window_days): p for p in my}
        done=0
        for fut in as_completed(futs):
            try:
                r=fut.result(timeout=180)
            except Exception as e:
                r={"file": str(futs[fut]), "error": str(e)[:500]}
            results.append(r)
            done+=1
            if done%10==0:
                print(f"[shard {args.shard}] {done}/{len(my)} done", flush=True)
            # append jsonl
            with open(out_jsonl,"a") as f:
                f.write(json.dumps(r)+"\n")
    # summary
    ok=sum(1 for r in results if r.get("gain_match"))
    mismatch=sum(1 for r in results if r.get("gain_match")==False)
    incomplete=sum(1 for r in results if r.get("filename_incomplete"))
    chart_ok=sum(1 for r in results if r.get("chart_exists"))
    summary={
        "shard": args.shard,
        "total_shards": args.total_shards,
        "workers": args.workers,
        "window_days": args.window_days,
        "total_assigned": len(my),
        "completed": len(results),
        "gain_match": ok,
        "gain_mismatch": mismatch,
        "incomplete": incomplete,
        "chart_exists": chart_ok,
        "chart_missing": len(my)-chart_ok,
        "timestamp": time.time(),
    }
    out_report.write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary, indent=2), flush=True)
    # also print top mismatches
    for r in sorted(results, key=lambda x: abs(x.get("gain_diff") or 0), reverse=True)[:10]:
        if not r.get("gain_match") and r.get("engine_gain") is not None:
            print(f"MISMATCH {r['file']} file_gain {r['filename_gain']} engine {r['engine_gain']} diff {r['gain_diff']:.2f} trades {r['engine_trades']}", flush=True)

if __name__=="__main__":
    main()
