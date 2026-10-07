#!/usr/bin/env python3
"""v15_repair_driver — intelligent TEMPLATE reruns WITHOUT a sheet refill (USER 2026-10-06: "double gains in the next 4 hours
in every backtest … intelligent TEMPLATE_* reruns and function suggestions"). Vector-only (vectorized gains); parity is
fixed in parallel by other lanes — repaired sets are NOT promoted live by this tool.

Per sym_side (queue, N in parallel, each with its own fork pool and NPZ in RAM):
  baseline = better (by tools/v15_diagnose_repair.quality_key) of
             (a) latest round's best result for the sym_side (progress cumulative_overrides, newest progress dir wins)
             (b) cat_side defaults (empty override set = QuickConfig + apply_cat_side_defaults = TEMPLATE bold)
  then tools/v15_diagnose_repair.run: DIAGNOSE -> SOFTEN -> ADD -> TIGHTEN -> POLISH -> REPAIR_365 -> 365D VERIFY,
  candidates = every evaluable TEMPLATE row + yellow filter of the sym_side's cat_side template.
Writes {out}/{SS}.json (full report + best_overrides + gain before/after + 365D) and appends {out}/summary.jsonl.

usage (on S1/S2/S5, repo root):
  .venv/bin/python tools/v15_repair_driver.py --symsides AXSUSDT_LONG,BNBUSDC_LONG --out ~/v15_repair_20261006 --parallel 3 --workers 5 --budget 600
  .venv/bin/python tools/v15_repair_driver.py --all-finished --venue crypto --out ... (every sym_side with a progress final set on this host)
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import glob
import json
import multiprocessing as mp
import os
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
PREP = None


def _w(ov):
    from tools.opt import v12_pilot as VP
    return VP.evaluate_prepared_sanitized(PREP, ov, 30)


def _wl(ov):
    from tools.opt import v12_pilot as VP
    from tools import v15_trade_autopsy as TA
    r = VP.evaluate_prepared_sanitized(PREP, ov, 30, True)
    keep = {k: r.get(k) for k in ("gain_pct", "trades", "tim_pct", "max_dd_pct", "wr_pct", "valid", "invalid_reason", "bh_pct", "behavior_fingerprint")}
    keep["_rows"] = TA.scaled_rows(r)
    return keep


def _is_crypto(ss: str) -> bool:
    return ss.rsplit("_", 1)[0].endswith(("USDT", "USDC", "USD1", "BUSD", "FDUSD"))


def latest_best(ss: str, dirs: list) -> tuple:
    """(overrides, source, final_gain) from the newest progress JSON that carries a final set for ss."""
    best = None
    for d in dirs:
        for p in glob.glob(os.path.join(os.path.expanduser(d), f"{ss}_v14_progress.json")):
            try:
                j = json.load(open(p))
            except Exception:
                continue
            ov = j.get("cumulative_overrides")
            if not ov:
                continue
            mt = os.path.getmtime(p)
            if best is None or mt > best[0]:
                best = (mt, dict(ov), p, j.get("final_gain_fresh_vec"))
    return (best[1], best[2], best[3]) if best else ({}, None, None)


def candidates_for(ss: str, defaults: dict, P) -> list:
    import openpyxl
    cat = "CRYPTO" if _is_crypto(ss) else "STOCKS"
    side = ss.rsplit("_", 1)[1]
    tpl = ROOT / f"SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_{cat}_{side}.xlsx"
    if not tpl.exists():
        tpl = ROOT / f"SPREADSHEETS/TEMPLATE_{cat}_{side}.xlsx"
    wb = openpyxl.load_workbook(str(tpl), read_only=True)
    kcf = P.known_config_fields()
    cands, seen = [], set()
    for sname in P.SWITCH_SHEETS:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        rows = list(ws.iter_rows(min_row=2))
        hdrs = []
        for c in rows[0][11:]:
            hv = c.value
            if isinstance(hv, str) and hv.strip().upper().startswith("WHAT SWITCH"):
                break
            if isinstance(hv, str) and "=" in hv:
                hdrs.append(hv.strip())
        for r in rows[1:]:
            a, b = r[0], r[1] if len(r) > 1 else None
            sw, cand = a.value, (b.value if b is not None else None)
            if sw in (None, "") or cand in (None, ""):
                continue
            sw = str(sw).strip()
            if sw.lower() in ("switch", "general", "blanket", "filter", "option value") or (sw, str(cand)) in seen:
                continue
            try:
                if str(a.font.color.rgb if a.font and a.font.color else "").upper() in P.GREY_SKIP_RGB:
                    continue
            except Exception:
                pass
            if sw in P.UNWIRED_SWITCHES or sw in P.UNWIRED_FILTERS or str(cand).upper().endswith("_ALT"):
                continue
            ov = P._switch_overrides(sw, P._parse_opt_value(cand, defaults.get(sw)))
            if not ov or not all(k in kcf for k in ov) or P._switch_type_violation(sw, cand, defaults):
                continue
            seen.add((sw, str(cand)))
            try:
                orange = str(a.fill.fgColor.rgb or "").upper().endswith("FFE699")
            except Exception:
                orange = False
            cands.append({"tab": sname, "row": a.row, "switch": sw, "cand": cand, "ov": ov, "orange": orange, "blocked": P.promotion_block_reason(sw, sname)})
        for h in hdrs:
            f, o = [x.strip() for x in h.split("=", 1)]
            if (f, o) in seen or f not in kcf or f in P.UNWIRED_FILTERS or f in P.UNWIRED_SWITCHES:
                continue
            v = P._parse_opt_value(o, defaults.get(f))
            if not P._cand_compatible(f, v, defaults)[0]:
                continue
            seen.add((f, o))
            cands.append({"tab": sname, "row": None, "switch": f, "cand": o, "ov": {f: v}, "orange": True, "blocked": P.promotion_block_reason(f, sname)})
    return cands


def same_val(a, b):  # = v15_pilot._spec_fill_workbook._same_val
    if isinstance(a, bool) or isinstance(b, bool) or str(a).lower() in ("true", "false") or str(b).lower() in ("true", "false"):
        return str(a).strip().lower() == str(b).strip().lower()
    try:
        return abs(float(a) - float(b)) < 1e-12
    except Exception:
        return str(a).strip() == str(b).strip()


def run_one(ss: str, out: Path, workers: int, budget: float, dirs: list) -> dict:
    global PREP
    import v15_pilot as P
    from tools import v15_diagnose_repair as DR
    from tools.opt import v12_pilot as VP
    t0 = time.time()
    defaults = P.get_defaults_for_symside(ss)
    san = lambda ov: P.sanitize_overrides(ov, defaults)[0]
    PREP = VP.prepare_batch(ss, 30)
    pool = cf.ProcessPoolExecutor(max_workers=workers, mp_context=mp.get_context("fork"))
    list(pool.map(abs, range(workers * 2)))
    try:
        prev_ov, prev_src, prev_final = latest_best(ss, dirs)
        bases = [("cat_side_defaults", {})] + ([("latest_best", prev_ov)] if prev_ov else [])
        scored = []
        for lab, ov in bases:
            r = VP.evaluate_prepared_sanitized(PREP, san(ov), 30)
            scored.append((DR.quality_key(DR.metrics(r)), lab, ov, r))
        scored.sort(key=lambda x: x[0], reverse=True)
        _, base_lab, base_ov, base_res = scored[0]
        cands = candidates_for(ss, defaults, P)
        span = P._npz_span_days(ss)
        p365 = {}

        def eval365(ov):
            if "p" not in p365:
                p365["p"] = VP.prepare_batch(ss, 365)
            return VP.evaluate_prepared_sanitized(p365["p"], san(ov), 365), span

        def eval_many(items, phase, cum_before, deadline):
            futs = [pool.submit(_w, ov) for _l, ov, _c in items]
            res = []
            wave = min(deadline + 60, time.time() + 10.0 * (-(-len(items) // workers) + 2))
            for f in futs:
                try:
                    res.append((f.result(timeout=max(0.01, wave - time.time())), ""))
                except Exception as e:
                    f.cancel()
                    res.append((None, str(e)[:80]))
            return res
        def eval_many_ledger(items, phase, cum_before, deadline):
            futs = [pool.submit(_wl, ov) for _l, ov, _c in items]
            res = []
            wave = min(deadline + 60, time.time() + 10.0 * (-(-len(items) // workers) + 2))
            for f in futs:
                try:
                    res.append((f.result(timeout=max(0.01, wave - time.time())), ""))
                except Exception as e:
                    f.cancel()
                    res.append((None, str(e)[:80]))
            return res
        npzp = PREP.get("npz_prepared") or {}
        close = [float(x) for x in npzp.get("close")]
        cat = ("CRYPTO" if _is_crypto(ss) else "STOCKS") + "_" + ss.rsplit("_", 1)[1]
        rep = DR.run({"defaults": defaults, "sanitize": san, "same_val": same_val, "candidates": cands, "base_overrides": base_ov,
                      "base_res": base_res, "bh": base_res.get("bh_pct"), "deadline": time.time() + budget, "eval_many": eval_many,
                      "eval_ledger": lambda ov: VP.evaluate_prepared_sanitized(PREP, san(ov), 30, True), "eval_365": eval365,
                      "qualifies_365": P._qualifies_365d, "cat_side": cat, "eval_many_ledger": eval_many_ledger, "close": close, "npz": npzp, "is_long": ss.endswith("_LONG"), "log": lambda m: print(f"{m} [{ss}]", flush=True), "touch": lambda m: None})
        b, a = rep.get("before") or {}, rep.get("after") or {}
        summ = {"symside": ss, "base": base_lab, "prev_final_gain": prev_final, "prev_src": prev_src,
                "bases": [{"label": lab, "gain": (r or {}).get("gain_pct"), "trades": (r or {}).get("trades"), "valid": (r or {}).get("valid")} for _, lab, _, r in scored],
                "gain_before": b.get("gain"), "gain_after": a.get("gain"), "bh": b.get("bh"), "trades_after": a.get("trades"), "tim_after": a.get("tim"), "dd_after": a.get("dd"),
                "accepted": rep.get("accepted"), "accept_reason": rep.get("accept_reason"), "n_changes": len(rep.get("changes") or []),
                "q365_after": next((f["q365"] for f in rep.get("finalists", []) if f["changes"] == rep.get("changes")), None),
                "origin_365": (rep.get("origin_365") or {}).get("m365"), "faults_before": [x[0] for x in rep.get("diagnosis_before", [])],
                "faults_after": [x[0] for x in rep.get("diagnosis_after", [])], "autopsy": rep.get("autopsy"), "surgical_steps": [st["applied"] for st in rep.get("steps", []) if st.get("phase") == "SURGICAL"], "top_rows": [(r["tab"], r.get("row"), r["switch"] + "=" + r["cand"], r["real_d_gain"]) for r in (rep.get("row_recommendations") or [])[:5]], "n_evals": rep.get("n_evals"), "secs": round(time.time() - t0, 1)}
        (out / f"{ss}.json").write_text(json.dumps({**rep, "summary": summ, "base_overrides": base_ov}, default=str))
        with open(out / "summary.jsonl", "a") as fh:
            fh.write(json.dumps(summ, default=str) + "\n")
        return summ
    finally:
        pool.shutdown(wait=False, cancel_futures=True)


def _child(args):
    ss, out, workers, budget, dirs = args
    try:
        return run_one(ss, Path(out), workers, budget, dirs)
    except Exception as e:
        import traceback
        err = {"symside": ss, "error": f"{type(e).__name__}: {e}"[:300], "trace": traceback.format_exc()[-1500:]}
        with open(Path(out) / "summary.jsonl", "a") as fh:
            fh.write(json.dumps(err) + "\n")
        return err


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--symsides", default="")
    ap.add_argument("--all-finished", action="store_true")
    ap.add_argument("--venue", choices=["crypto", "stocks", "all"], default="all")
    ap.add_argument("--out", required=True)
    ap.add_argument("--parallel", type=int, default=2)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--budget", type=float, default=600.0)
    ap.add_argument("--progress-dirs", default="data/reports/lifecycle_pilot,~/v15_run*/progress")
    a = ap.parse_args()
    dirs = []
    for d in a.progress_dirs.split(","):
        dirs += glob.glob(os.path.expanduser(d.strip())) or [os.path.expanduser(d.strip())]
    out = Path(os.path.expanduser(a.out))
    out.mkdir(parents=True, exist_ok=True)
    done = {json.loads(l).get("symside") for l in open(out / "summary.jsonl")} if (out / "summary.jsonl").exists() else set()
    if a.all_finished:
        sss = set()
        for d in dirs:
            for p in glob.glob(os.path.join(d, "*_v14_progress.json")):
                ss = os.path.basename(p).replace("_v14_progress.json", "")
                if (a.venue == "all" or (a.venue == "crypto") == _is_crypto(ss)):
                    sss.add(ss)
        todo = sorted(sss)
    else:
        todo = [s for s in a.symsides.split(",") if s]
    todo = [s for s in todo if s not in done]
    print(f"[repair-driver] {len(todo)} sym_sides parallel={a.parallel} workers={a.workers} budget={a.budget}s out={out}", flush=True)
    # each sym_side in its own spawned process (own NPZ + own fork pool; never fork from a threaded parent)
    with cf.ProcessPoolExecutor(max_workers=a.parallel, mp_context=mp.get_context("spawn")) as ex:
        for r in ex.map(_child, [(s, str(out), a.workers, a.budget, dirs) for s in todo]):
            print(f"[repair-driver] {r.get('symside')} base={r.get('base')} gain {r.get('gain_before')} -> {r.get('gain_after')} accepted={r.get('accepted')} q365={r.get('q365_after')} {r.get('error', '')}", flush=True)


if __name__ == "__main__":
    main()
