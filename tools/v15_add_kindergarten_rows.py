#!/usr/bin/env python3
"""v15_add_kindergarten_rows — ONE-OFF, USER-AUTHORIZED 2026-09-30 (explicit "unlock ... this is an exception"
to the no-add-rows prohibition). Ensure the vec-wired KINDERGARTEN entry gates are present as orange filter
rows in the ENTRY + AUGMENT + REENTRY tabs of ALL 4 templates, so they are swept everywhere (today only
CRYPTO_LONG carries them).

Adds ONLY the vec-wired KG gate groups (NOOP/unwired KG params are NOT added — they would be dead rows):
  * KINDERGARTEN_EMA_GATE_ENABLED  options {False, True}  default = False (crypto) / True (stocks)  [config differs]
  * EMA_9_21_FILTER_ENABLED        options {True, False}   default = True (both venues)
EMA_9_21_FILTER_MIN_TFS + the EMA_9_21_FILTER_FILTER_TF yellow columns already exist in every target tab.

Method (§56-safe): orange rows always sit below white rows, so new orange rows are APPENDED at the bottom of
the tab (can never create a white-below-orange violation). Each new row is a full copy of a same-tab orange KG
donor row (so the 200+ yellow columns/fills match that tab exactly); only A/B/C/D/L + stats are reset. Exactly
one is_default=YES (bold) per group = the venue default. Atomic save with zip-validate (>=10 entries).

Run on copies first:  python3 tools/v15_add_kindergarten_rows.py --template-dir /tmp/kg_test
Apply for real:       python3 tools/v15_add_kindergarten_rows.py --apply
"""
import argparse
import datetime
import os
import shutil
import sys
import zipfile
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from v15_template_restructure import is_orange  # noqa: E402

SPREAD = ROOT / "SPREADSHEETS"
TEMPLATES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx",
             "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
TARGET_TABS = ["ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES",
               "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REENTRY_WINDOWED", "REENTRY_ADAPTIVE"]
HDR = 2  # header row with column names
DONOR_NAMES = ("EMA_9_21_FILTER_MIN_TFS", "KINDERGARTEN_EMA_GATE_ENABLED", "EMA_9_21_FILTER_ENABLED")
# (group -> the two options, in display order); the default is decided per venue below
GROUPS = {"KINDERGARTEN_EMA_GATE_ENABLED": [False, True], "EMA_9_21_FILTER_ENABLED": [True, False]}


def venue_default(cat_side, group):
    if group == "KINDERGARTEN_EMA_GATE_ENABLED":
        return False if cat_side.startswith("CRYPTO") else True  # config.py False / config_tradier True
    return True  # EMA_9_21_FILTER_ENABLED True both venues


def hdr_map(ws):
    return {str(ws.cell(row=HDR, column=c).value).strip(): c for c in range(1, ws.max_column + 1)
            if ws.cell(row=HDR, column=c).value is not None}


def norm(v):
    return str(v).strip().lower()


def copy_row(ws, src_r, dst_r):
    for c in range(1, ws.max_column + 1):
        s = ws.cell(row=src_r, column=c)
        d = ws.cell(row=dst_r, column=c)
        d.value = s.value
        if s.has_style:
            d.font = copy(s.font)
            d.fill = copy(s.fill)
            d.border = copy(s.border)
            d.alignment = copy(s.alignment)
            d.number_format = s.number_format
            d.protection = copy(s.protection)
    if ws.row_dimensions.get(src_r) is not None:
        ws.row_dimensions[dst_r].height = ws.row_dimensions[src_r].height


def set_cell(cell, value, bold):
    cell.value = value
    f = cell.font
    cell.font = copy(f)
    cell.font = openpyxl.styles.Font(name=(f.name or "Arial"), size=(f.size or 10), bold=bold,
                                     italic=f.italic, color=f.color)


def add_tab(ws, cat_side, rep):
    hm = hdr_map(ws)
    cA, cB, cC, cD, cL = 1, hm.get("default", 2), hm.get("override", 3), hm.get("Family", 4), hm.get("is_default")
    cAVG, cPOS = hm.get("AVG_DELTA"), hm.get("POS_SYM")
    if not cL:
        raise RuntimeError(f"{ws.title}: no is_default column")
    # style donor = a same-tab orange KG row (keeps the tab's exact yellow-column layout)
    donor = None
    for r in range(HDR + 1, ws.max_row + 1):
        a = ws.cell(row=r, column=cA).value
        if a and is_orange(ws, r) and str(a).strip() in DONOR_NAMES:
            donor = r
            break
    if donor is None:
        raise RuntimeError(f"{ws.title}: no orange KG donor row found")
    fam = ws.cell(row=donor, column=cD).value
    # existing options per group
    have = {g: {} for g in GROUPS}
    for r in range(HDR + 1, ws.max_row + 1):
        a = str(ws.cell(row=r, column=cA).value or "").strip()
        if a in GROUPS:
            have[a].setdefault(norm(ws.cell(row=r, column=cB).value), r)
    for g, opts in GROUPS.items():
        dflt = venue_default(cat_side, g)
        # add missing options (append at bottom = orange region)
        for opt in opts:
            if norm(opt) not in have[g]:
                nr = ws.max_row + 1
                copy_row(ws, donor, nr)
                ws.cell(row=nr, column=cA).value = g
                set_cell(ws.cell(row=nr, column=cA), g, True)  # name bold like donor
                set_cell(ws.cell(row=nr, column=cB), opt, opt == dflt)
                ws.cell(row=nr, column=cC).value = None
                ws.cell(row=nr, column=cD).value = fam
                for cc in (hm.get("BASELINE"), hm.get("HUSTLE_DELTA"), hm.get("VECTOR_DELTA"),
                           hm.get("LIVE_DELTA"), hm.get("LIVE_SHARPE"), hm.get("REAL_COMPLETE"),
                           hm.get("PER_ROW_FILTERS"), cAVG, cPOS):
                    if cc:
                        ws.cell(row=nr, column=cc).value = None
                have[g][norm(opt)] = nr
                rep.append((ws.title, g, str(opt), "added", opt == dflt))
        # normalise is_default: exactly one YES = venue default, across all rows of the group
        for optn, r in have[g].items():
            yes = optn == norm(dflt)
            set_cell(ws.cell(row=r, column=cB), ws.cell(row=r, column=cB).value, yes)
            set_cell(ws.cell(row=r, column=cL), "YES" if yes else "NO", yes)


def atomic_save(wb, path):
    tmp = str(path) + ".tmp"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as z:
        if len(z.namelist()) < 10:
            os.remove(tmp)
            raise RuntimeError(f"{path}: refused truncated save ({len(z.namelist())} zip entries)")
    os.replace(tmp, path)


def verify(path):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    bad = []
    for cat_side, fn in []:
        pass
    for tab in TARGET_TABS:
        if tab not in wb.sheetnames:
            continue
        ws = wb[tab]
        hm = hdr_map(ws)
        cL = hm.get("is_default")
        for g in GROUPS:
            rows = [r for r in range(HDR + 1, ws.max_row + 1) if str(ws.cell(row=r, column=1).value or "").strip() == g]
            opts = {norm(ws.cell(row=r, column=2).value) for r in rows}
            yes = [r for r in rows if str(ws.cell(row=r, column=cL).value or "").strip().upper() == "YES"]
            if not {"true", "false"} <= opts:
                bad.append(f"{tab}/{g}: options={sorted(opts)} (need true+false)")
            if len(yes) != 1:
                bad.append(f"{tab}/{g}: {len(yes)} is_default=YES (need exactly 1)")
    wb.close()
    return bad


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write real SPREADSHEETS templates (backup first)")
    ap.add_argument("--template-dir", default=None, help="operate on copies in this dir instead of SPREADSHEETS")
    args = ap.parse_args()
    tdir = Path(args.template_dir) if args.template_dir else SPREAD
    if args.template_dir:
        tdir.mkdir(parents=True, exist_ok=True)
        for fn in TEMPLATES.values():
            if not (tdir / fn).exists():
                shutil.copyfile(SPREAD / fn, tdir / fn)
    ts = datetime.datetime.utcnow().strftime("%Y%m%d%H%M")
    rep = []
    for cat_side, fn in TEMPLATES.items():
        path = tdir / fn
        if args.apply and not args.template_dir:
            shutil.copyfile(path, ROOT / "backups" / f"before_add_kindergarten_{cat_side}_{ts}.xlsx")
        wb = openpyxl.load_workbook(str(path))
        for tab in TARGET_TABS:
            if tab in wb.sheetnames:
                add_tab(wb[tab], cat_side, rep)
        if args.apply or args.template_dir:
            atomic_save(wb, path)
        wb.close()
        bad = verify(path)
        added = sum(1 for r in rep if r[0] and r[3] == "added")
        print(f"[{cat_side}] rows added this file: {sum(1 for x in rep if x[3]=='added')} (cumulative) | verify: "
              + ("OK" if not bad else f"{len(bad)} ISSUES"))
        for b in bad[:10]:
            print("   !", b)
    print(f"\nTotal rows added: {sum(1 for x in rep if x[3]=='added')}")
    from collections import Counter
    byg = Counter((x[1]) for x in rep if x[3] == "added")
    for g, n in byg.items():
        print(f"  {g}: +{n} rows")
    if not (args.apply or args.template_dir):
        print("\n(dry preview only — no file written; use --template-dir DIR to test on copies, --apply to write)")


if __name__ == "__main__":
    main()
