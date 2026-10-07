#!/usr/bin/env python3
"""switch_revival_scan (Agent G) — read EVERY available version of TEMPLATE_{cat}.xlsx (backups/, TEMPLATE_STAGED/, data/template_audit orig/base, live)
and extract per file: for each SWITCH tab the white rows (switch, option, family, is_default, bold) and orange rows, + PARKED_ROWS sheet rows.
-> data/switch_revival/scan/<cat>/<file-slug>.json   (resumable; skips existing)"""
import glob, json, multiprocessing as mp, os, re, sys, time
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "switch_revival" / "scan"
TABS = ['STDEV_SLOPE_SIZING', 'ENTRY_REVERSAL_BOUNCE', 'ENTRY_BREAKOUT_CHANNEL', 'ENTRY_CONFIRMATION_GATES', 'EXIT_STRUCTURAL', 'EXIT_VELOCITY', 'REENTRY_WINDOWED', 'REENTRY_ADAPTIVE', 'AUGMENT_TREND', 'AUGMENT_RISK_SIZING', 'REDUCE_PROFIT_LOCK', 'REDUCE_SIGNAL_RATER', 'GLOBAL_RISK_GATES']
CATS = ["CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"]


def slug(p):
    return re.sub(r"[^A-Za-z0-9_.-]", "_", str(Path(p).relative_to(ROOT)))[-150:]


def scan(job):
    cat, path = job
    out = OUT / cat / (slug(path) + ".json")
    if out.exists():
        return str(out), "skip"
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        return path, f"ERR {e}"
    res = {"file": str(Path(path).relative_to(ROOT)), "mtime": os.path.getmtime(path), "size": os.path.getsize(path), "tabs": {}, "parked": []}
    for t in TABS:
        if t not in wb.sheetnames:
            continue
        ws = wb[t]
        hdr = {str(ws.cell(2, c).value).strip().lower(): c for c in range(1, min(ws.max_column, 20) + 1) if ws.cell(2, c).value}
        c_fam = hdr.get("family", 4)
        c_def = hdr.get("is_default", 12)
        rows = []
        for r in range(3, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a in (None, ""):
                continue
            fill = ws.cell(r, 1).fill
            rgb = str(fill.fgColor.rgb) if fill is not None and fill.fill_type == "solid" else ""
            f = ws.cell(r, 1).font
            grey = bool(f is not None and f.color is not None and str(getattr(f.color, "rgb", "") or "").upper().endswith("BFBFBF"))
            rows.append([str(a).strip(), ws.cell(r, 2).value if not isinstance(ws.cell(r, 2).value, (set, list)) else str(ws.cell(r, 2).value), ws.cell(r, c_fam).value, ws.cell(r, c_def).value,
                         bool(ws.cell(r, 2).font and ws.cell(r, 2).font.b), "FFE699" in rgb, grey])
        res["tabs"][t] = rows
    if "PARKED_ROWS" in wb.sheetnames:
        ws = wb["PARKED_ROWS"]
        for row in ws.iter_rows(min_row=1, values_only=True):
            res["parked"].append([None if v is None else str(v) for v in row[:14]])
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(res, default=str))
    return str(out), "ok"


def files_for(cat):
    pats = [f"backups/*TEMPLATE_{cat}*.xlsx", f"SPREADSHEETS/TEMPLATE_{cat}.xlsx", f"SPREADSHEETS/TEMPLATE_STAGED/*/TEMPLATE_{cat}.xlsx", f"data/template_audit/*/orig_TEMPLATE_{cat}.xlsx", f"data/template_audit/*/base/TEMPLATE_{cat}.xlsx",
            f"SPREADSHEETS/*TEMPLATE_{cat}*.xlsx", f"SPREADSHEETS/**/TEMPLATE_{cat}*.xlsx", f"data/switch_revival/gitv/{cat}/*.xlsx"]
    s = set()
    for p in pats:
        for f in glob.glob(str(ROOT / p), recursive=True):
            if f.endswith(".xlsx") and "~$" not in f and "V15_V16_CELL_BY_CELL" not in f:
                s.add(f)
    return sorted(s)


if __name__ == "__main__":
    jobs = [(c, f) for c in CATS for f in files_for(c)]
    print(len(jobs), "files", flush=True)
    t0 = time.time()
    with mp.Pool(int(sys.argv[1]) if len(sys.argv) > 1 else 8) as pool:
        for i, (p, st) in enumerate(pool.imap_unordered(scan, jobs), 1):
            if st.startswith("ERR") or i % 25 == 0:
                print(i, st, p[-80:], round(time.time() - t0), flush=True)
    print("done", round(time.time() - t0), flush=True)
