#!/usr/bin/env python3
"""
v15_progress_board — per-tab truth of how full every V15 workbook is, from the pilot's progress JSON.

Used two ways:
  * herd done-test: `is_complete(symside)` — a workbook is DONE only when every one of the 12 SWITCH_SHEETS
    has every switch row in progress["done"]. (Old test: xlsx >500 KB ⇒ done, which counted a freshly cloned
    template with only a baseline as finished — that is how 148 sheets stayed empty forever.)
  * monitor: `python3 tools/v15_progress_board.py` writes SPREADSHEETS/V15_PROGRESS.md + V15_PROGRESS.csv
    (cron every minute on S1; tools/sync_s1_to_mac.sh brings both to the Mac).

Row state per tab: filled = row key present, red = present with reason "all vectors invalid" (pilot retries
red rows on every resume, see RED-RETRY in v15_pilot.py).
"""
import csv
import json
import os
import pathlib
import subprocess
import sys
import time

ROOT = pathlib.Path(__file__).resolve().parent.parent
PROGRESS_DIR = ROOT / "data" / "reports" / "lifecycle_pilot"
SHEETS_DIR = ROOT / "SPREADSHEETS" / "V15_V16_CELL_BY_CELL"
CACHE_PATH = ROOT / "data" / "reports" / "v15_template_rowcounts.json"
SWITCH_SHEETS = [
    "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES",
    "EXIT_STRUCTURAL", "EXIT_VELOCITY",
    "REENTRY_WINDOWED", "REENTRY_ADAPTIVE",
    "AUGMENT_TREND", "AUGMENT_RISK_SIZING",
    "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER",
    "GLOBAL_RISK_GATES",
]
SHORT_NAMES = ["ENT_REV", "ENT_BRK", "ENT_CNF", "EXT_STR", "EXT_VEL", "RE_WIN", "RE_ADP", "AUG_TRD", "AUG_RSK", "RED_PL", "RED_SR", "GLOBAL"]
_ROWCOUNT_MEMO: dict = {}


def template_for(symside: str) -> pathlib.Path:
    # same selection order as v15_local_herd.launch()
    is_stock = "USDC" not in symside and "USDT" not in symside
    side = "LONG" if symside.endswith("_LONG") else "SHORT"
    venue = "STOCKS" if is_stock else "CRYPTO"
    # USER 2026-09-28: TEMPLATE.xlsx is legacy (replaced by the 4 cat_side templates) — never fall back to it
    for name in (f"TEMPLATE_V15_{venue}_{side}.xlsx", f"TEMPLATE_{venue}_{side}.xlsx"):
        path = ROOT / "SPREADSHEETS" / name
        if path.exists():
            return path
    return ROOT / "SPREADSHEETS" / f"TEMPLATE_{venue}_{side}.xlsx"


def template_rows(template: pathlib.Path) -> dict:
    # {sheet: [row numbers with a switch in col A]} — cached on disk by path+mtime (loading a template takes seconds)
    key = f"{template}|{int(template.stat().st_mtime)}"
    if key in _ROWCOUNT_MEMO:
        return _ROWCOUNT_MEMO[key]
    cache = {}
    try:
        cache = json.loads(CACHE_PATH.read_text())
    except Exception:
        pass
    if key not in cache:
        import openpyxl
        wb = openpyxl.load_workbook(template, read_only=True, data_only=True)
        rows = {}
        for sheet in SWITCH_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            rows[sheet] = [index for index, row in enumerate(wb[sheet].iter_rows(min_row=3, max_col=1, values_only=True), start=3) if row and row[0] not in (None, "")]
        wb.close()
        cache = {k: v for k, v in cache.items() if not k.startswith(f"{template}|")}
        cache[key] = rows
        tmp = f"{CACHE_PATH}.{os.getpid()}.tmp"
        pathlib.Path(tmp).write_text(json.dumps(cache))
        os.replace(tmp, CACHE_PATH)
    _ROWCOUNT_MEMO[key] = cache[key]
    return cache[key]


def template_row_ids(template: pathlib.Path) -> dict:
    # {sheet: [[row, "SWITCH=CAND"], ...]} — identity matching survives row renumbering when the
    # template is edited (USER 2026-09-28: template row changes must add calculations for those
    # rows ONLY, never restart a sheet). Cached like template_rows under an "ids|" key.
    key = f"ids|{template}|{int(template.stat().st_mtime)}"
    if key in _ROWCOUNT_MEMO:
        return _ROWCOUNT_MEMO[key]
    cache = {}
    try:
        cache = json.loads(CACHE_PATH.read_text())
    except Exception:
        pass
    if key not in cache:
        import openpyxl
        def _fmt(v):
            if v is True:
                return "True"
            if v is False:
                return "False"
            if v is None:
                return ""
            return str(v).strip()
        wb = openpyxl.load_workbook(template, read_only=True, data_only=True)
        rows = {}
        for sheet in SWITCH_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            rows[sheet] = [[index, f"{_fmt(row[0])}={_fmt(row[1] if len(row) > 1 else None)}"] for index, row in enumerate(wb[sheet].iter_rows(min_row=3, max_col=2, values_only=True), start=3) if row and row[0] not in (None, "")]
        wb.close()
        cache = {k: v for k, v in cache.items() if not k.startswith(f"ids|{template}|")}
        cache[key] = rows
        tmp = f"{CACHE_PATH}.{os.getpid()}.ids.tmp"
        pathlib.Path(tmp).write_text(json.dumps(cache))
        os.replace(tmp, CACHE_PATH)
    _ROWCOUNT_MEMO[key] = cache[key]
    return cache[key]


def symside_status(symside: str, window_days: int = 30) -> dict:
    progress_path = PROGRESS_DIR / (f"{symside}_v14_progress.json" if window_days == 30 else f"{symside}_{window_days}d_progress.json")
    status = {"symside": symside, "tabs": {}, "filled": 0, "rows": 0, "red": 0, "baseline": None, "cum": None, "updated": None, "has_progress": progress_path.exists()}
    rows_by_sheet = template_rows(template_for(symside))
    done = {}
    if progress_path.exists():
        try:
            progress = json.loads(progress_path.read_text())
            done = progress.get("done", {}) or {}
            status["baseline"] = progress.get("baseline_gain")
            status["cum"] = progress.get("cumulative_gain")
            status["zero_trades"] = bool(progress.get("zero_trades_diagnostic"))
            status["updated"] = progress_path.stat().st_mtime
        except Exception as error:
            status["error"] = f"progress unreadable: {error}"
    done_rows: dict = {}
    done_ids: set = set()
    for key, value in done.items():
        sheet_row, _, ident = key.partition(":")
        if "!" not in sheet_row:
            continue
        sheet, _, row = sheet_row.partition("!")
        # USER 2026-10-03 (complete-rounds): hollow rows are NOT filled — corpses (is_running),
        # undecided blanks, and revoked policy-skips (NOT_WIRED_VEC / SKIPPED_SAMPLING recalc this
        # round). Terminal settles (numeric delta, ZERO_TRADES, TYPE_MISMATCH, NOT_IN_CONFIG,
        # INVENTED_ALT, NO_CANDIDATE, RUNNING_IDENTITY, red rows) still count.
        _vv = value if isinstance(value, dict) else {}
        _vr = str(_vv.get("reason") or "")
        if bool(_vv.get("is_running")) or (not isinstance(_vv.get("delta"), (int, float)) and (_vr == "" or _vr.startswith("NOT_WIRED_VEC") or "SKIPPED_SAMPLING" in _vr)):
            continue
        if ident:
            done_ids.add((sheet, ident))
        if not row.isdigit():
            continue
        is_red = isinstance(value, dict) and value.get("reason") == "all vectors invalid"
        done_rows.setdefault(sheet, {})[int(row)] = is_red
    try:
        ids_by_sheet = template_row_ids(template_for(symside))
    except Exception:
        ids_by_sheet = {}
    for sheet in SWITCH_SHEETS:
        wanted = rows_by_sheet.get(sheet, [])
        got = done_rows.get(sheet, {})
        # a row counts as filled by ROW match or by (switch=cand) IDENTITY match — identity
        # survives template renumbering, so edits add calculations for new rows only
        id_of_row = {row: ident for row, ident in ids_by_sheet.get(sheet, [])}
        filled = sum(1 for row in wanted if row in got or (sheet, id_of_row.get(row, "\x00")) in done_ids)
        red = sum(1 for row in wanted if got.get(row))
        status["tabs"][sheet] = {"rows": len(wanted), "filled": filled, "red": red}
        status["rows"] += len(wanted)
        status["filled"] += filled
        status["red"] += red
    status["complete"] = status["rows"] > 0 and status["filled"] >= status["rows"]
    return status


def is_complete(symside: str) -> bool:
    # zero-trade diagnostics are terminal (herd must never loop on them) — everything else needs all 12 tabs filled
    try:
        status = symside_status(symside)
    except Exception:
        return False
    if bool(status.get("zero_trades")):
        return True
    if not status["complete"]:
        return False
    # 2026-09-28: complete-but-UNPUBLISHED is not done — the DONE stage died before writing the
    # bh/gain final (OOM era); returning False here re-queues it for the pilot's publish-only pass.
    return any(SHEETS_DIR.glob(f"{symside}_bh*_30d_matrix*.xlsx"))


def running_symsides() -> set:
    try:
        out = subprocess.check_output(["pgrep", "-af", "v15_pilot.py"], text=True)
    except Exception:
        return set()
    running = set()
    for line in out.splitlines():
        parts = line.split()
        if "--sym-side" in parts and parts.index("--sym-side") + 1 < len(parts):
            running.add(parts[parts.index("--sym-side") + 1])
    return running


def all_symsides() -> list:
    names = set()
    for path in PROGRESS_DIR.glob("*_v14_progress.json"):
        names.add(path.name.replace("_v14_progress.json", ""))
    for path in SHEETS_DIR.glob("*_30d*.xlsx"):
        if "_2026" in path.name:
            continue
        names.add(path.name.split("_30d")[0])
    return sorted(name for name in names if name.endswith(("_LONG", "_SHORT")))


def tab_cell(tab: dict) -> str:
    if not tab["rows"]:
        return "-"
    if tab["red"]:
        return f"🔴{tab['filled']}/{tab['rows']}"
    if tab["filled"] >= tab["rows"]:
        return "✅"
    if tab["filled"] == 0:
        return "·"
    return f"{100 * tab['filled'] // tab['rows']}%"


def write_board() -> dict:
    running = running_symsides()
    statuses = []
    for symside in all_symsides():
        try:
            statuses.append(symside_status(symside))
        except Exception as error:
            statuses.append({"symside": symside, "error": str(error), "tabs": {}, "filled": 0, "rows": 0, "red": 0, "complete": False})
    for status in statuses:
        status["running"] = status["symside"] in running
        if status.get("error"):
            status["state"] = "ERROR"
        elif status["complete"]:
            status["state"] = "DONE+RED" if status["red"] else "DONE"
        elif status["running"]:
            status["state"] = "RUNNING"
        elif status.get("zero_trades"):
            status["state"] = "ZERO_TRADES"
        elif status["filled"] == 0:
            status["state"] = "EMPTY"
        else:
            status["state"] = "PARTIAL"
    order = {"RUNNING": 0, "ERROR": 1, "PARTIAL": 2, "EMPTY": 3, "ZERO_TRADES": 4, "DONE+RED": 5, "DONE": 6}
    statuses.sort(key=lambda s: (order.get(s["state"], 9), -(s["filled"] / s["rows"] if s["rows"] else 0), s["symside"]))
    counts: dict = {}
    for status in statuses:
        counts[status["state"]] = counts.get(status["state"], 0) + 1
    total_rows = sum(s["rows"] for s in statuses)
    total_filled = sum(s["filled"] for s in statuses)
    total_red = sum(s["red"] for s in statuses)
    now = time.strftime("%Y-%m-%d %H:%M:%S %Z")
    host = os.uname().nodename
    lines = [
        f"# V15 progress board — {now} on {host}",
        "",
        f"**{len(statuses)} workbooks** · " + " · ".join(f"{k} {v}" for k, v in sorted(counts.items(), key=lambda kv: order.get(kv[0], 9))),
        f"**rows filled {total_filled}/{total_rows} ({100 * total_filled / total_rows if total_rows else 0:.1f}%) · red {total_red}** · running now: {', '.join(sorted(running)) or 'none'}",
        "",
        "Legend: ✅ tab full · NN% partly filled · · untouched · 🔴f/n has red rows (retried on next resume) · tabs in pilot order.",
        "",
        "| sym_side | state | filled | red | baseline | cum | updated | " + " | ".join(SHORT_NAMES) + " |",
        "|---|---|---|---|---|---|---|" + "---|" * len(SHORT_NAMES),
    ]
    for status in statuses:
        pct = f"{100 * status['filled'] / status['rows']:.0f}%" if status["rows"] else "-"
        baseline = f"{status['baseline']:.2f}" if isinstance(status.get("baseline"), (int, float)) else "-"
        cum = f"{status['cum']:.2f}" if isinstance(status.get("cum"), (int, float)) else "-"
        updated = time.strftime("%m-%d %H:%M", time.localtime(status["updated"])) if status.get("updated") else "-"
        tabs = " | ".join(tab_cell(status["tabs"][sheet]) if sheet in status["tabs"] else "-" for sheet in SWITCH_SHEETS)
        lines.append(f"| {status['symside']} | {status['state']} | {status['filled']}/{status['rows']} {pct} | {status['red']} | {baseline} | {cum} | {updated} | {tabs} |")
    out_dir = ROOT / "SPREADSHEETS"
    md_tmp = out_dir / f"V15_PROGRESS.md.{os.getpid()}.tmp"
    md_tmp.write_text("\n".join(lines) + "\n")
    os.replace(md_tmp, out_dir / "V15_PROGRESS.md")
    csv_tmp = out_dir / f"V15_PROGRESS.csv.{os.getpid()}.tmp"
    with open(csv_tmp, "w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["symside", "state", "filled", "rows", "red", "baseline", "cum", "updated_epoch"] + [f"{s}_filled" for s in SWITCH_SHEETS] + [f"{s}_red" for s in SWITCH_SHEETS])
        for status in statuses:
            writer.writerow([status["symside"], status["state"], status["filled"], status["rows"], status["red"], status.get("baseline"), status.get("cum"), status.get("updated")] + [status["tabs"].get(s, {}).get("filled", "") for s in SWITCH_SHEETS] + [status["tabs"].get(s, {}).get("red", "") for s in SWITCH_SHEETS])
    os.replace(csv_tmp, out_dir / "V15_PROGRESS.csv")
    return counts


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] != "--write":
        for symside in sys.argv[1:]:
            print(json.dumps(symside_status(symside), indent=1, default=str))
    else:
        print(write_board())
