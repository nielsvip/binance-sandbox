#!/usr/bin/env python3
"""v18_features — the FEATURE MAP of every switch, candidate value, filter and option of a cat_side's TEMPLATE (E8b, 2026-10-20).

WHY. The V18 brain (tools/v18_brain.py) counts links that were measured. To predict a link that was never measured — a pair
never tried, a setting one grid step outside a row's span, a stop reference with no template row yet (doc 17 §3.7) — the
generalising network (tools/v18_linknet.py) needs to know what a key IS, not only its name. This module turns a key into
attributes and sparse features derived ONLY from (a) the name tokens, (b) the TEMPLATE metadata (tab, candidates, bold
default, FILTERS_EXPLAINED type, FILTER_DICTIONARY_V2 sheets), (c) data/SWITCH_BIBLE.json (kind, sides, wiring status),
(d) data/switch_dependencies.json (masters = dependency parents) and (e) data/v17_tf_semantics.json (role). Nothing is
learned or guessed here; unknown tokens are kept as raw `tok:` features AND listed in the coverage report (never dropped).

ATTRIBUTES (per name; the same function serves switches and filters)
  family     indicator family set: dc, bb, wt, ema, sma, rsi, stoch, atr, adx, mfi, macd, volume, sentiment, stdev, ha,
             candle, structure, gap, velocity, regime, sharpe, peak, sma200, golden_rule, delta_engine, kindergarten, ...
  tf         timeframes named in the key (15m, 1h, 4h, D, W ...), class htf/ltf
  role       entry / exit / reentry / augment / reduce / sizing / stop / target / gate / filter / cooldown / window /
             threshold / reference (tab role + name tokens + bible kind + tf_semantics role)
  direction  long / short / both (name tokens + bible sides)
  level      low / high / basis (lower/bottom -> low, upper/top -> high, mid/middle/basis -> basis)
  tab        template tab(s), tab role
  parents    switch_dependencies masters (sub-knob -> switches that must be ON)
VALUE BUCKET (per value of a name)
  bool ON/OFF; TF value + TF rank; numeric: rank quantile among the row's candidates (q in [0,1], outside the span < 0 or
  > 1 for proposals), position vs the bold default (below / default / above), sign, log-magnitude; enum: value tokens
  (family + level of e.g. 'dc_high', 'bb_upper', 'lower').

FEATURE STRINGS (consumed by tools/v18_linknet.py; prefix S. = setting side, F. = filter side)
  ID features first:  S.sw:<NAME>, S.set:<NAME>=<v>  /  F.fl:<NAME>, F.fo:<NAME>=<opt>  (dropped at train time with a
  probability, so the network also learns the ID-free path that generalises to keys it has never seen)
  then  .fam: .tf: .tfc: .role: .dir: .lvl: .tab: .tabrole: .par: .kind: .stat: .sem: .tok:  and the value bucket
  .vb: .vtf: .vtfr: .vq: .vrel: .vsign: .vmag: .venum: .vlvl: .vfam: plus real-valued .vqx (q - 0.5).

CLI
  python3 tools/v18_features.py build [--cat-side CS ...] [--out-dir data/learn]
      -> data/learn/features_<CS>.json (attributes of every template switch/filter, bucket of every candidate/option,
         coverage report: every template key featurised, unknown tokens with counts, names without family/role)
"""
from __future__ import annotations

import argparse
import bisect
import hashlib
import json
import math
import os
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(os.environ.get("V18_ROOT") or Path(__file__).resolve().parents[1])
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

SCHEMA = "v18_features/1"
CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")
TF_ORDER = ("1m", "3m", "5m", "15m", "30m", "1h", "2h", "4h", "6h", "8h", "12h", "D", "3d", "W", "M")
TF_RANK = {t: i for i, t in enumerate(TF_ORDER)}
TF_ALIAS = {"1d": "D", "1D": "D", "DAILY": "D", "1w": "W", "1W": "W", "WEEKLY": "W", "15M": "15m", "1H": "1h", "4H": "4h", "5M": "5m",
            "3M": "3m", "30M": "30m", "2H": "2h", "6H": "6h", "8H": "8h", "12H": "12h", "15": "15m"}
TF_VALUES = frozenset(TF_ORDER) | {"OFF", "1d", "1w"}
HTF = frozenset({"4h", "6h", "8h", "12h", "D", "3d", "W", "M"})
TAB_ROLE = {"STDEV_SLOPE_SIZING": "sizing", "ENTRY_REVERSAL_BOUNCE": "entry", "ENTRY_BREAKOUT_CHANNEL": "entry", "ENTRY_CONFIRMATION_GATES": "gate",
            "EXIT_STRUCTURAL": "exit", "EXIT_VELOCITY": "exit", "REENTRY_WINDOWED": "reentry", "REENTRY_ADAPTIVE": "reentry", "AUGMENT_TREND": "augment",
            "AUGMENT_RISK_SIZING": "augment", "REDUCE_PROFIT_LOCK": "reduce", "REDUCE_SIGNAL_RATER": "reduce", "GLOBAL_RISK_GATES": "gate"}
FAMILY_TOK = {
    "DC": "dc", "DC4": "dc", "DONCHIAN": "dc", "DCHAN": "dc", "CHANNEL": "channel", "BB": "bb", "BOLLINGER": "bb", "BBW": "bb", "BB4H": "bb",
    "WT": "wt", "WAVETREND": "wt", "WT15": "wt", "WT4H": "wt", "WTDC": "wt", "EMA": "ema", "EMA9": "ema", "EMA21": "ema", "EMA50": "ema",
    "EMA200": "ema", "SMA": "sma", "SMA200": "sma200", "SMA50": "sma", "MA": "sma", "RSI": "rsi", "STOCH": "stoch", "STOCHRSI": "stoch",
    "K15M": "stoch", "KD": "stoch", "ATR": "atr", "ADX": "adx", "DMI": "adx", "MFI": "mfi", "MACD": "macd", "VOL": "volume",
    "VOLUME": "volume", "OBV": "volume", "VWAP": "volume", "CVD": "volume", "SENTIMENT": "sentiment", "FUNDING": "sentiment", "OI": "sentiment",
    "LS": "sentiment", "STDEV": "stdev", "STD": "stdev", "SIGMA": "stdev", "ZSCORE": "stdev", "Z": "stdev", "HA": "ha", "HEIKIN": "ha",
    "HAIKU": "ha", "CANDLE": "candle", "PATTERN": "candle", "WICK": "candle", "DOJI": "candle", "ENGULF": "candle", "HAMMER": "candle",
    "LH": "structure", "HL": "structure", "LL": "structure", "HH": "structure", "SWING": "structure", "PIVOT": "structure", "HLR": "structure",
    "BOTA": "structure", "R1": "structure", "R2": "structure", "S1": "structure", "S2": "structure", "GAP": "gap", "VEL": "velocity",
    "VELOCITY": "velocity", "SLOPE": "velocity", "MOMENTUM": "velocity", "FAST": "velocity", "RISER": "velocity", "REGIME": "regime",
    "TREND": "regime", "SHARPE": "sharpe", "CIRCUIT": "sharpe", "PEAK": "peak", "GIVEBACK": "peak", "EROSION": "peak", "GR": "golden_rule",
    "GOLDEN": "golden_rule", "DELTA": "delta_engine", "KINDERGARTEN": "kindergarten", "RALLY": "rally", "BREAKOUT": "breakout",
    "RETEST": "breakout", "BOUNCE": "bounce", "PERCENTILE": "percentile", "PCTL": "percentile", "SATOSHIT": "satoshit", "SCALP": "scalp",
    "DAYTRADE": "daytrade", "DT": "daytrade", "MTF": "mtf", "HTF": "mtf", "LTF": "mtf", "OPEN": "session", "SESSION": "session",
    "MOC": "session", "RTH": "session", "PREMARKET": "session", "NEWBORN": "age", "AGE": "age", "STREAK": "streak", "VIGILANCE": "vigilance",
    "HEDGE": "hedge", "LADDER": "ladder", "FIB": "fib", "ICHIMOKU": "ichimoku", "KELTNER": "keltner", "SUPERTREND": "supertrend",
    "PSAR": "psar", "CCI": "cci", "ROC": "roc", "TRIX": "trix", "WILLIAMS": "williams", "PRICE": "price", "SPREAD": "spread",
    "BAND": "band", "K": "stoch", "K15": "stoch", "K3M": "stoch", "LR": "linreg", "PULLBACK": "pullback", "CORRECTION": "pullback",
    "BREAK": "breakout", "BREACH": "breakout", "RECOVERY": "recovery", "SQUEEZE": "bb", "BBKC": "bb", "PCTB": "bb", "MOM3": "velocity",
    "ACCEL": "velocity", "SURGE": "velocity", "SPIKE": "velocity", "PARABOLIC": "velocity", "SMFI": "mfi", "COUNTER": "counter_trend",
    "CT": "counter_trend", "DD": "drawdown", "BTC": "btc", "OB": "orderblock", "MU": "mean_rev", "MI": "mean_rev", "RZ": "stdev",
    "DIV": "divergence", "PYRAMID": "ladder", "CONSOL": "range", "RANGING": "range", "ZONE": "range", "FRAC": "fractal",
    "FORMATION": "candle", "PSR": "sharpe", "CROWD": "sentiment", "FG": "sentiment", "MACRO": "regime", "TRENDING": "regime",
    "XTREME": "extreme", "EXTREME": "extreme", "STRETCH": "extreme", "EXHAUST": "extreme", "FADE": "fade", "CHURN": "churn",
}
ROLE_TOK = {
    "STOP": "stop", "SL": "stop", "FROZEN": "stop", "TRAIL": "stop", "BREAKEVEN": "stop", "BE": "stop", "FLOOR": "stop", "HOPELESS": "stop",
    "TARGET": "target", "TP": "target", "PROFIT": "target", "TAKE": "target", "GATE": "gate", "BLOCK": "gate", "VETO": "gate",
    "REQUIRE": "gate", "CONFIRM": "gate", "BYPASS": "gate", "GUARD": "gate", "BRAKE": "gate", "THROTTLE": "gate", "FILTER": "filter",
    "COOLDOWN": "cooldown", "WINDOW": "window", "THRESHOLD": "threshold", "THRESH": "threshold", "MIN": "threshold", "MAX": "threshold",
    "PCT": "threshold", "LIMIT": "threshold", "SCORE": "threshold", "MULT": "sizing", "SIZE": "sizing", "SIZING": "sizing", "QTY": "sizing",
    "WEIGHT": "sizing", "ENTRY": "entry", "ENTRIES": "entry", "BUY": "entry", "EXIT": "exit", "CLOSE": "exit", "SELL": "exit",
    "REENTRY": "reentry", "RECLAIM": "reentry", "AUGMENT": "augment", "AUG": "augment", "ADD": "augment", "REDUCE": "reduce",
    "PARTIAL": "reduce", "LOCK": "reduce", "REF": "reference", "FIELD": "reference", "BASIS": "reference", "BUFFER": "threshold",
    "REJECT": "gate", "MANDATORY": "gate", "OBLIGATORY": "gate", "ENFORCE": "gate", "ALLOW": "gate", "WATCHDOG": "gate", "TRAILING": "stop",
    "TIGHTEN": "stop", "HOLD": "window", "LOOKBACK": "window", "TOL": "threshold", "THR": "threshold", "BOOST": "sizing", "BONUS": "sizing",
    "ABLATION": "ablation_lever", "DISABLE": "ablation_lever", "SCORER": "gate", "FLIP": "exit", "RESET": "window", "RETRY": "window",
}
LEVEL_TOK = {"LOW": "low", "LOWER": "low", "BOTTOM": "low", "BOT": "low", "HIGH": "high", "UPPER": "high", "TOP": "high", "BASIS": "basis",
             "MID": "basis", "MIDDLE": "basis", "CENTER": "basis", "MEDIAN": "basis"}
DIR_TOK = {"LONG": "long", "BULL": "long", "BULLISH": "long", "SHORT": "short", "BEAR": "short", "BEARISH": "short"}
# tokens that carry no meaning of their own (still emitted as tok:, never listed as unknown)
GLUE = frozenset({"ENABLED", "ENABLE", "MODE", "USE", "ON", "OFF", "TF", "TFS", "OF", "THE", "AND", "OR", "FOR", "TO", "IN", "AT", "IF", "NO",
                  "V2", "V3", "V4", "V5", "V6", "VEC", "N", "X", "ALL", "ANY", "BOTH", "ONLY", "AFTER", "BEFORE", "BARS", "BAR", "S", "SEC", "H",
                  "HOURS", "MINUTES", "MIN_", "ABS", "NEW", "OLD", "DETAILED", "COMBO", "TRIGGER", "SIGNAL", "STATE", "LEVEL", "VALUE", "TRADIER",
                  "CRYPTO", "STOCKS", "STOCK", "LIVE", "HARD", "SOFT", "TIGHT", "WIDE", "ADAPTIVE", "DYNAMIC", "FIXED", "ABSOLUTE", "COUNT",
                  "RATIO", "RATE", "GAIN", "LOSS", "KILL", "PCT_", "REQUIRED", "DIRECT", "UNCOND", "CROSS", "TURN", "DIVERG", "DIVERGENCE",
                  "TECHNICAL", "ULTIMATE", "GUARANTEED", "CARRY", "WIN", "WINNER", "ONE", "TWO", "FIRST", "LAST", "NEXT", "PREV", "TIME"})


def norm_tf(tok):
    t = str(tok).strip()
    if t in TF_RANK:
        return t
    if t in TF_ALIAS:
        return TF_ALIAS[t]
    if t.upper() in TF_ALIAS:
        return TF_ALIAS[t.upper()]
    if t.upper() == "D":
        return "D"
    if t.upper() == "W":
        return "W"
    return None


def name_tokens(name):
    """'DC_LOW_FROZEN_STOP_TF' -> ['DC','LOW','FROZEN','STOP','TF']; digits glued to letters are split too ('WT15M' -> WT, 15M)."""
    out = []
    for t in re.split(r"[_\s\-\.]+", str(name or "").upper()):
        if not t:
            continue
        m = re.match(r"^([A-Z]+?)(\d+[MHDW])$", t)
        if m and m.group(1) not in ("K",):
            out.extend([m.group(1), m.group(2)])
        else:
            out.append(t)
    return out


def tok_class(t):
    """-> (kind, value) of one name token, kind in family/role/level/dir/tf/glue/num/unknown."""
    if t in FAMILY_TOK:
        return "family", FAMILY_TOK[t]
    if t in LEVEL_TOK:
        return "level", LEVEL_TOK[t]
    if t in DIR_TOK:
        return "dir", DIR_TOK[t]
    tf = norm_tf(t) if (re.match(r"^\d+[MHDW]$", t) or t in ("D", "W", "DAILY", "WEEKLY", "1D", "1W")) else None
    if tf:
        return "tf", tf
    if t in ROLE_TOK:
        return "role", ROLE_TOK[t]
    if t in GLUE:
        return "glue", t
    if re.match(r"^\d+(\.\d+)?$", t):
        return "num", t
    for pref, fam in (("EMA", "ema"), ("SMA", "sma"), ("WT", "wt"), ("BB", "bb"), ("DC", "dc"), ("RSI", "rsi"), ("ATR", "atr")):
        if t.startswith(pref) and re.match(r"^" + pref + r"\d+$", t):
            return "family", fam
    return "unknown", t


def is_num(v):
    try:
        float(str(v))
        return str(v).strip().lower() not in ("nan", "inf", "-inf", "true", "false")
    except Exception:
        return False


def value_kind(v):
    s = str(v if v is not None else "").strip()
    if s in ("True", "False"):
        return "bool"
    if s and all(norm_tf(x) or x == "OFF" for x in re.split(r"[,+]", s) if x):
        return "tf"
    if is_num(s):
        return "num"
    return "enum"


class Context:
    """Read-only metadata of ONE cat_side: template index, filter dictionary, switch bible, dependencies, tf semantics."""

    def __init__(self, cat_side, tidx=None, bible=None, deps=None, tfsem=None, fdict=None):
        self.cs = cat_side
        self.venue = "STOCKS" if cat_side.startswith("STOCKS") else "CRYPTO"
        self.side = cat_side.split("_", 1)[1]
        self.tidx = tidx or {"tabs": {}, "switches": {}, "defaults": {}}
        self.bible = bible if bible is not None else {}
        self.masters = (deps or {}).get("masters") or {}
        self.tfsem = (tfsem or {}).get("keys") or {}
        self.fdict = fdict or {}
        self.tabs_of = {}
        for tab, blk in self.tidx.get("tabs", {}).items():
            for _r, sw, _c, _d in blk["rows"]:
                self.tabs_of.setdefault(sw, [])
                if tab not in self.tabs_of[sw]:
                    self.tabs_of[sw].append(tab)
        self.cands = {sw: list(v) for sw, v in (self.tidx.get("switches") or {}).items()}
        self.defaults = dict(self.tidx.get("defaults") or {})
        self.fopts = {}
        self.ftabs = {}
        for tab, blk in self.tidx.get("tabs", {}).items():
            for y in blk["yellows"]:
                f, _, o = y.partition("=")
                lst = self.fopts.setdefault(f, [])
                if o not in lst:
                    lst.append(o)
                self.ftabs.setdefault(f, [])
                if tab not in self.ftabs[f]:
                    self.ftabs[f].append(tab)
        self.extra_cands = {}

    def add_seen_values(self, name, values):
        """Values seen in the corpus for a name that has no template row (numeric quantiles need a span)."""
        cur = self.extra_cands.setdefault(name, [])
        for v in values:
            if v not in cur:
                cur.append(v)

    def span(self, name, is_filter=False):
        lst = (self.fopts.get(name) if is_filter else self.cands.get(name)) or self.cands.get(name) or self.fopts.get(name) or self.extra_cands.get(name) or []
        return sorted({float(x) for x in lst if is_num(x)})

    def default(self, name, is_filter=False):
        d = self.defaults.get(name)
        if d is None:
            b = self.bible.get(name) or {}
            dd = ((b.get("defaults") or {}).get("cat_side_defaults_4") or {}).get(self.cs)
            d = None if dd is None else (("True" if dd else "False") if isinstance(dd, bool) else str(dd))
        return d


def name_attrs(name, ctx, tab=None, is_filter=False):
    """Attributes of a switch or filter NAME (no value). Every token is classified; unknown tokens are returned."""
    toks = name_tokens(name)
    fam, role, lvl, dirs, tfs, unknown, raw = set(), set(), set(), set(), set(), [], []
    for t in toks:
        k, v = tok_class(t)
        raw.append(t)
        if k == "family":
            fam.add(v)
        elif k == "role":
            role.add(v)
        elif k == "level":
            lvl.add(v)
            if v == "basis":
                role.add("reference")
        elif k == "dir":
            dirs.add(v)
        elif k == "tf":
            tfs.add(v)
        elif k == "unknown":
            unknown.append(t)
    tabs = []
    if is_filter:
        tabs = list(ctx.ftabs.get(name) or [])
    else:
        tabs = list(ctx.tabs_of.get(name) or ([tab] if tab else []))
    if tab and tab not in tabs:
        tabs.append(tab)
    tabroles = sorted({TAB_ROLE.get(t, "other") for t in tabs})
    b = ctx.bible.get(name) or {}
    kind = b.get("kind")
    if kind and kind not in ("unclassified",):
        role.add({"global": "gate"}.get(kind, kind))
    sides = b.get("sides") or []
    if not dirs and sides:
        dirs = {"long"} if sides == ["LONG"] else ({"short"} if sides == ["SHORT"] else set())
    if not dirs:
        dirs = {"both"}
    st = (b.get("status") or {}).get("stocks" if ctx.venue == "STOCKS" else "crypto")
    sem = (ctx.tfsem.get(name) or {}).get("role")
    fd = ctx.fdict.get(name) or {}
    if is_filter or fd.get("type"):
        role.add("filter") if is_filter else None
    if name.endswith("_TF") and not tfs:
        role.add("timeframe_knob")
    parents = sorted(ctx.masters.get(name) or [])
    return {"tokens": raw, "family": sorted(fam), "role": sorted(role), "level": sorted(lvl), "direction": sorted(dirs), "tf": sorted(tfs, key=lambda x: TF_RANK.get(x, 99)),
            "tf_class": sorted({("htf" if t in HTF else "ltf") for t in tfs}), "tabs": tabs, "tab_roles": tabroles, "parents": parents, "kind": kind,
            "status": st, "sem_role": sem, "ftype": fd.get("type"), "unknown_tokens": unknown}


def value_bucket(name, value, ctx, is_filter=False):
    """Bucket of one value of a name: bool / tf / num (quantile among the row's candidates, vs default) / enum tokens."""
    s = "" if value is None else str(value).strip()
    k = value_kind(s)
    out = {"kind": k, "value": s}
    if k == "bool":
        out["vb"] = "ON" if s == "True" else "OFF"
    elif k == "tf":
        parts = [x for x in re.split(r"[,+]", s) if x]
        tfs = [norm_tf(x) or x for x in parts]
        out["tfs"] = tfs
        out["vb"] = "OFF" if tfs == ["OFF"] else "ON"
        ranks = [TF_RANK[t] for t in tfs if t in TF_RANK]
        out["tf_rank"] = max(ranks) if ranks else -1
        out["tf_n"] = len(tfs)
    elif k == "num":
        x = float(s)
        span = ctx.span(name, is_filter)
        if len(span) >= 2:
            lo, hi = span[0], span[-1]
            if x < lo or x > hi:
                q = (x - lo) / (hi - lo) if hi > lo else 0.5
                out["out_of_span"] = True
                out["span_step"] = _step(span)
                out["q"] = round(q, 4)
            else:
                i = bisect.bisect_left(span, x)
                if i < len(span) and span[i] == x:
                    q = i / (len(span) - 1)
                else:
                    q = (i - 1 + (x - span[i - 1]) / (span[i] - span[i - 1])) / (len(span) - 1)
                out["q"] = round(q, 4)
        else:
            out["q"] = 0.5
        d = ctx.default(name, is_filter)
        if d is not None and is_num(d):
            dv = float(d)
            out["rel"] = "default" if abs(x - dv) <= 1e-12 * max(1.0, abs(dv)) else ("below" if x < dv else "above")
        out["sign"] = "zero" if x == 0 else ("neg" if x < 0 else "pos")
        out["mag"] = "0" if x == 0 else str(int(math.floor(math.log10(abs(x)))))
        out["vb"] = "ON" if x != 0 else "OFF"
    else:
        vt = [t for t in re.split(r"[_\s,\-]+", s.upper()) if t]
        out["enum_tokens"] = vt
        out["levels"] = sorted({LEVEL_TOK[t] for t in vt if t in LEVEL_TOK})
        out["families"] = sorted({FAMILY_TOK[t] for t in vt if t in FAMILY_TOK})
        out["vb"] = "OFF" if s.upper() in ("OFF", "NONE", "") else "ON"
    if k != "num":
        d = ctx.default(name, is_filter)
        if d is not None:
            out["rel"] = "default" if str(d) == s else "other"
    return out


def _step(span):
    gaps = sorted(b - a for a, b in zip(span, span[1:]) if b > a)
    return gaps[len(gaps) // 2] if gaps else 1.0


def features(name, value, ctx, tab=None, is_filter=False, attrs=None, bucket=None):
    """-> [(feature, x)]. The first two are the ID features (sw/set or fl/fo); the rest are ID-free (generalising) features."""
    P = "F." if is_filter else "S."
    a = attrs or name_attrs(name, ctx, tab=tab, is_filter=is_filter)
    b = bucket or value_bucket(name, value, ctx, is_filter=is_filter)
    v = b["value"]
    out = [(P + ("fl:" if is_filter else "sw:") + name, 1.0), (P + ("fo:" if is_filter else "set:") + f"{name}={v}", 1.0)]
    add = lambda f, x=1.0: out.append((P + f, x))
    for f in a["family"]:
        add("fam:" + f)
    if not a["family"]:
        add("fam:none")
    for t in a["tf"]:
        add("tf:" + t)
    for t in a["tf_class"]:
        add("tfc:" + t)
    for r in a["role"]:
        add("role:" + r)
    for r in a["level"]:
        add("lvl:" + r)
    for d in a["direction"]:
        add("dir:" + d)
    for t in a["tabs"][:3]:
        add("tab:" + t)
    for t in a["tab_roles"]:
        add("tabrole:" + t)
    for p in a["parents"][:3]:
        add("par:" + p)
    if a.get("kind"):
        add("kind:" + str(a["kind"]))
    if a.get("status"):
        add("stat:" + str(a["status"]))
    if a.get("sem_role"):
        add("sem:" + str(a["sem_role"]))
    if a.get("ftype"):
        add("ftype:" + str(a["ftype"]).split()[0][:16])
    for t in a["tokens"]:
        if tok_class(t)[0] in ("unknown", "glue") and t not in ("ENABLED", "TF"):
            add("tok:" + t, 0.5)
    k = b["kind"]
    add("vk:" + k)
    add("vb:" + b.get("vb", "?"))
    if b.get("rel"):
        add("vrel:" + b["rel"])
    if k == "tf":
        for t in b.get("tfs", []):
            add("vtf:" + str(t))
        add("vtfr:" + str(b.get("tf_rank")))
        if b.get("tf_n", 1) > 1:
            add("vtfn:" + str(min(4, b["tf_n"])))
    elif k == "num":
        q = float(b.get("q", 0.5))
        qb = min(4, max(0, int(round(min(1.0, max(0.0, q)) * 4))))
        add("vq:" + str(qb))
        add("vqx", max(-1.5, min(1.5, q - 0.5)))
        add("vsign:" + b["sign"])
        add("vmag:" + b["mag"])
    elif k == "enum":
        for t in b.get("enum_tokens", [])[:4]:
            add("venum:" + t)
        for t in b.get("levels", []):
            add("vlvl:" + t)
        for t in b.get("families", []):
            add("vfam:" + t)
    if k == "bool" and a["family"]:
        add("vb:" + b["vb"] + "|" + a["family"][0])
    return out


def relation_features(name, value, inforce, src, ctx, is_filter=False):
    """USER LAW 2026-10-20: a candidate is measured against the DEFAULT IN FORCE (the chain's value of that key at eval time:
    the sym_side's last promotion, else the template bold). -> features of the candidate RELATIVE to it: source of the default
    (P = promoted on this sym_side, T = template bold, U = unknown), its value ID (slot 2) and the direction of the change."""
    P = "F." if is_filter else "S."
    inf = None if inforce in (None, "", "None") else str(inforce)
    out = [((P + "inf:" + f"{name}={inf}") if inf is not None else "<pad>", 1.0 if inf is not None else 0.0), (P + "infsrc:" + str(src or "U"), 1.0)]
    if inf is None:
        return out
    v = "" if value is None else str(value).strip()
    kv, ki = value_kind(v), value_kind(inf)
    if v == inf:
        out.append((P + "rinf:same", 1.0))
        return out
    if kv == ki == "bool":
        out.append((P + "rinf:" + ("flip_on" if v == "True" else "flip_off"), 1.0))
    elif kv == ki == "num":
        a, b = float(v), float(inf)
        out.append((P + "rinf:" + ("above" if a > b else "below"), 1.0))
        span = ctx.span(name, is_filter)
        if len(span) >= 2 and span[-1] > span[0]:
            dq = (a - b) / (span[-1] - span[0])
            out.append((P + "rinfx", max(-1.5, min(1.5, dq))))
        if b != 0 and (a / b) < 0:
            out.append((P + "rinf:sign_flip", 1.0))
    elif kv == ki == "tf":
        ra = max([TF_RANK.get(norm_tf(x) or x, -1) for x in re.split(r"[,+]", v) if x] or [-1])
        rb = max([TF_RANK.get(norm_tf(x) or x, -1) for x in re.split(r"[,+]", inf) if x] or [-1])
        if inf == "OFF":
            out.append((P + "rinf:tf_from_off", 1.0))
        elif v == "OFF":
            out.append((P + "rinf:tf_to_off", 1.0))
        else:
            out.append((P + "rinf:" + ("tf_higher" if ra > rb else ("tf_lower" if ra < rb else "tf_same_rank")), 1.0))
    else:
        la = sorted({LEVEL_TOK[t] for t in re.split(r"[_\s,\-]+", v.upper()) if t in LEVEL_TOK})
        lb = sorted({LEVEL_TOK[t] for t in re.split(r"[_\s,\-]+", inf.upper()) if t in LEVEL_TOK})
        out.append((P + "rinf:other_" + kv, 1.0))
        if la != lb and la and lb:
            out.append((P + f"rinflvl:{lb[0]}->{la[0]}", 1.0))
    return out


def load_filter_dict(template_path):
    """{filter: {type, sheets}} from the TEMPLATE's FILTERS_EXPLAINED (Type) and FILTER_DICTIONARY_V2 (Sheets applicable)."""
    out = {}
    try:
        import openpyxl
    except Exception:
        return out
    try:
        wb = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
    except Exception:
        return out
    try:
        if "FILTERS_EXPLAINED" in wb.sheetnames:
            for i, r in enumerate(wb["FILTERS_EXPLAINED"].iter_rows(values_only=True)):
                if i == 0 or not r or not r[0]:
                    continue
                f = str(r[0]).strip()
                t = str(r[1]).strip() if len(r) > 1 and r[1] is not None else None
                if t and not is_num(t):
                    out.setdefault(f, {})["type"] = t
        if "FILTER_DICTIONARY_V2" in wb.sheetnames:
            for i, r in enumerate(wb["FILTER_DICTIONARY_V2"].iter_rows(values_only=True)):
                if i == 0 or not r or len(r) < 6 or not r[1]:
                    continue
                f = str(r[1]).strip()
                out.setdefault(f, {})["sheets"] = str(r[5] or "").strip()[:80]
                if len(r) > 8 and r[8]:
                    out[f]["status"] = str(r[8]).strip()[:40]
    finally:
        wb.close()
    return out


def load_context(cat_side, root=ROOT, use_cache=True):
    from tools import v18_learn_common as C
    tidx = C.template_index(cat_side)
    cache = Path(root) / "data" / "learn" / "_cache" / f"filter_dict_{cat_side}_{tidx['md5'][:12]}.json"
    fdict = None
    if use_cache and cache.exists():
        try:
            fdict = json.loads(cache.read_text())
        except Exception:
            fdict = None
    if fdict is None:
        fdict = load_filter_dict(tidx["path"])
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_text(json.dumps(fdict))
    rd = lambda p: json.loads(Path(p).read_text()) if Path(p).exists() else {}
    bible = (rd(Path(root) / "data" / "SWITCH_BIBLE.json") or {}).get("switches") or {}
    deps = rd(Path(root) / "data" / "switch_dependencies.json")
    tfsem = rd(Path(root) / "data" / "v17_tf_semantics.json")
    return Context(cat_side, tidx=tidx, bible=bible, deps=deps, tfsem=tfsem, fdict=fdict)


def build(cat_side, ctx=None):
    """Feature map + coverage of one cat_side's TEMPLATE. Every switch row, candidate, yellow filter and option is featurised."""
    ctx = ctx or load_context(cat_side)
    sw_attrs, set_b, f_attrs, opt_b = {}, {}, {}, {}
    unknown = Counter()
    no_fam, no_role = [], []
    n_feats = Counter()
    for tab, blk in ctx.tidx["tabs"].items():
        for _r, sw, cand, _d in blk["rows"]:
            if sw not in sw_attrs:
                sw_attrs[sw] = name_attrs(sw, ctx)
                unknown.update(sw_attrs[sw]["unknown_tokens"])
            set_b.setdefault(sw, {})
            if cand not in set_b[sw]:
                set_b[sw][cand] = value_bucket(sw, cand, ctx)
                n_feats["settings"] += 1
                n_feats["setting_features"] += len(features(sw, cand, ctx, attrs=sw_attrs[sw], bucket=set_b[sw][cand]))
        for y in blk["yellows"]:
            f, _, o = y.partition("=")
            if f not in f_attrs:
                f_attrs[f] = name_attrs(f, ctx, is_filter=True)
                unknown.update(f_attrs[f]["unknown_tokens"])
            opt_b.setdefault(f, {})
            if o not in opt_b[f]:
                opt_b[f][o] = value_bucket(f, o, ctx, is_filter=True)
                n_feats["options"] += 1
                n_feats["option_features"] += len(features(f, o, ctx, is_filter=True, attrs=f_attrs[f], bucket=opt_b[f][o]))
    for nm, a in list(sw_attrs.items()) + list(f_attrs.items()):
        if not a["family"]:
            no_fam.append(nm)
        if not a["role"]:
            no_role.append(nm)
    tmpl_keys = sum(len(v) for v in set_b.values())
    tmpl_opts = sum(len(v) for v in opt_b.values())
    expect_settings = len({(sw, c) for blk in ctx.tidx["tabs"].values() for _r, sw, c, _d in blk["rows"]})
    expect_opts = len({y for blk in ctx.tidx["tabs"].values() for y in blk["yellows"]})
    cov = {"switches": len(sw_attrs), "settings": tmpl_keys, "settings_expected": expect_settings, "filters": len(f_attrs), "options": tmpl_opts,
           "options_expected": expect_opts, "all_featurised": tmpl_keys == expect_settings and tmpl_opts == expect_opts,
           "mean_features_per_setting": round(n_feats["setting_features"] / max(1, n_feats["settings"]), 2),
           "mean_features_per_option": round(n_feats["option_features"] / max(1, n_feats["options"]), 2),
           "unknown_tokens": dict(unknown.most_common()), "n_unknown_tokens": len(unknown),
           "names_without_family": sorted(no_fam), "names_without_role": sorted(no_role),
           "note": "unknown tokens are KEPT as tok:<T> features (weight 0.5); they are listed here so the vocabulary can grow, never dropped"}
    meta = {"schema": SCHEMA, "cat_side": cat_side, "template": ctx.tidx.get("path"), "template_md5": ctx.tidx.get("md5"),
            "sources": ["template rows/yellows (tools/v18_learn_common.template_index)", "FILTERS_EXPLAINED / FILTER_DICTIONARY_V2", "data/SWITCH_BIBLE.json",
                        "data/switch_dependencies.json", "data/v17_tf_semantics.json"]}
    return {"meta": meta, "coverage": cov, "switches": sw_attrs, "settings": set_b, "filters": f_attrs, "options": opt_b}


# ------------------------------------------------------------------------------------------------ chain context (E8d, v2)

INACTIVE_VALUES = frozenset({"OFF", "False", "0", "None", "", "<absent>"})
CTX_SKIP_PREFIX = ("ABLATION_",)
CTX_SKIP_KEYS = frozenset({"MODE", "VENUE", "IS_TRADIER", "SIMPLE_PRICE_GT0_ENABLED"})


def canon_value(v):
    """One spelling per value so a store set (python types), a template cell (str) and a delta-log cand compare equal:
    bools -> True/False, numbers -> shortest float form ('1.0' -> '1'), strings stripped. TF words stay as they are."""
    if v is None:
        return "None"
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        s = str(v).strip()
        if s.lower() in ("true", "false"):
            return s.capitalize()
        try:
            f = float(s)
        except ValueError:
            return s
        if not math.isfinite(f) or re.match(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$", s) is None:
            return s
    if not math.isfinite(f):
        return str(f)
    if f.is_integer() and abs(f) < 1e15:
        return str(int(f))
    return f"{f:.10g}"


VALUE_RE = re.compile(r"^[A-Za-z0-9_.,+:\-]{1,40}$")


def plausible_value(cv):
    """a setting value, not a description: board xlsx column C sometimes carries the row's explanation text (seen on stocks
    board_col_c sets) -> such values are ignored (counted), never turned into context features."""
    return bool(VALUE_RE.match(cv))


class ContextEncoder:
    """CHAIN CONTEXT of an eval (E8d): the full set in force, reduced to what differs from the cat_side's TEMPLATE bold,
    over the template's own vocabulary (switch rows + yellow filter keys). Sparse multi-hot features:
      X.f:<FILTER>=<v>   a yellow-filter key whose value in force differs from its bold (an ACTIVE filter option, or a filter
                         the chain switched off where the bold had it on)
      X.s:<SWITCH>=<v>   any other template switch whose value in force differs from its bold
    A key without a known bold is a feature only when its value is active (not OFF/False/0/None). Keys outside the template,
    ABLATION_* levers and MODE/VENUE are ignored. The diff dict is the representation; feats() is its sorted feature tuple."""

    def __init__(self, fctx):
        self.fctx = fctx
        self.filter_keys = frozenset(fctx.fopts)
        self.ignored_values = Counter()
        self.keys = frozenset(k for k in set(fctx.cands) | set(fctx.fopts) if not k.startswith(CTX_SKIP_PREFIX) and k not in CTX_SKIP_KEYS)
        self.bold = {}
        for k in self.keys:
            d = fctx.default(k, is_filter=k in self.filter_keys and k not in fctx.cands)
            self.bold[k] = None if d is None else canon_value(d)

    def differs(self, key, cv):
        b = self.bold.get(key)
        if b is None:
            return cv not in INACTIVE_VALUES
        return cv != b

    def diff(self, full):
        """full set dict -> {key: canon value} of the template keys whose value differs from the bold."""
        out = {}
        for k, v in (full or {}).items():
            if k not in self.keys:
                continue
            cv = canon_value(v)
            if cv == "<absent>":
                continue
            if not plausible_value(cv):
                self.ignored_values[k] += 1
                continue
            if self.differs(k, cv):
                out[k] = cv
        return out

    def overlay(self, diff, mods):
        """diff + chain modifications {key: value} (promotions on the board after the stored base) -> new diff dict."""
        if not mods:
            return diff
        out = dict(diff)
        for k, v in mods.items():
            if k not in self.keys:
                continue
            cv = canon_value(v)
            if cv != "<absent>" and not plausible_value(cv):
                self.ignored_values[k] += 1
                continue
            if cv == "<absent>" or not self.differs(k, cv):
                out.pop(k, None)
            else:
                out[k] = cv
        return out

    def feats(self, diff):
        return tuple(sorted((f"X.f:{k}={v}" if k in self.filter_keys else f"X.s:{k}={v}") for k, v in diff.items()))

    @staticmethod
    def filter_signature(feats):
        """the ACTIVE-FILTER SET of a context (its X.f features) -> the 'never-tried context' split key."""
        return tuple(f for f in feats if f.startswith("X.f:"))


# ------------------------------------------------------------------------------------------------ binding features (V20 v3, 2026-10-22)
# Source: data/learn/binding_map_<CS>.json (tools/v18_binding_map.py, contract binding_map_v1) + the RUNNING SET (the chain in force:
# a board's start set overlaid with its promotions so far; at decision time the pilot's chain). No outcome is read here (no leakage).

BIND_CONTRACT = "binding_map_v1"
BIND_VERDICTS = ("BINDS", "CONDITIONAL_ARMED", "CONDITIONAL_OFF", "INDIRECT", "NEVER", "UNKNOWN")
BIND_STAGES = ("ENTRY", "REENTRY", "EXIT_STOP", "EXIT_TARGET", "EXIT_VEL", "EXIT_STRUCT", "AUGMENT", "REDUCE", "SIZING", "GLOBAL")
FILTER_STAGE_ALIAS = {"EXIT_SIG": "EXIT_VEL", "REDUCE_SIG": "REDUCE", "WIN_TRAIL_EROSION": "EXIT_STOP"}
ENTRY_BLOCKER_GATES = ("ADX_REGIME_FILTER_ENABLED", "COUNTER_TREND_ADD_BLOCK_ENABLED", "HTF_DIRECTION_GATE_ENABLED", "GR_FILTER_ALL_ENTRIES")
TR4H_GATE_RE = re.compile(r"^TR_[A-Z0-9_]*4H[A-Z0-9_]*GATE")
TAB_DIGEST_DIM = 16
BIND_DIM = len(BIND_VERDICTS) + (len(BIND_STAGES) + 1) + (len(BIND_STAGES) + 2) + 1 + 2 + 1 + 1 + TAB_DIGEST_DIM
BIND_FEATURE_NAMES = ([f"bind_verdict:{v}" for v in BIND_VERDICTS] + [f"stage_switch:{s}" for s in BIND_STAGES + ("UNKNOWN",)]
                      + [f"stage_filter:{s}" for s in BIND_STAGES + ("NONE", "UNKNOWN")] + ["same_stage", "masters_armed", "masters_known", "dead_row",
                                                                                            "upstream_gate_active"] + [f"tab_digest:{i}" for i in range(TAB_DIGEST_DIM)])


def _truthy(v):
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return v != 0
    return str(v).strip().lower() not in ("", "off", "false", "0", "0.0", "none", "null")


def _filter_stage(stages):
    s = str((stages or ["UNKNOWN"])[0] or "UNKNOWN").split("+", 1)[0].strip().upper()
    s = FILTER_STAGE_ALIAS.get(s, s)
    return s if (s in BIND_STAGES or s == "NONE") else "UNKNOWN"


def _tf_bucket(cv):
    if cv in ("OFF", "False", "0", "None", ""):
        return "off"
    if cv in ("True",):
        return "bool"
    parts = [norm_tf(x) for x in re.split(r"[,+]", cv) if x]
    if parts and all(parts):
        if len(parts) > 1:
            return "multi"
        return "htf" if parts[0] in HTF else "ltf"
    return "num" if is_num(cv) else "enum"


class BindingIndex:
    """Read-only lookups over ONE cat_side's binding map: cells by (tab, switch, canon cand, filter, canon opt), switch/filter
    stages, master defaults, the filters of each tab and the entry-blocker gate keys. Static parts are memoised per cell."""

    def __init__(self, bm, path=None):
        self.meta = bm.get("meta") or {}
        self.cs = self.meta.get("cat_side")
        self.path = str(path) if path else None
        self.switches = bm.get("switches") or {}
        self.filters = bm.get("filters") or {}
        self.defaults = {k: (v or {}).get("value") for k, v in (self.meta.get("masters_default") or {}).items()}
        self.rows = {}
        self.tab_filters = {}
        for key, row in (bm.get("rows") or {}).items():
            tab, _, kc = key.partition("!")
            sw, _, cand = kc.partition("=")
            cells = {}
            for hdr, c in (row.get("cells") or {}).items():
                f, _, o = hdr.partition("=")
                cells[(f.strip(), canon_value(o))] = c
                self.tab_filters.setdefault(tab, set()).add(f.strip())
            self.rows.setdefault((tab, sw, canon_value(cand)), (row, cells))
        self.gates = tuple(ENTRY_BLOCKER_GATES) + tuple(sorted(f for f, v in self.filters.items() if v.get("class") == "ENTRY_BOOL_GATE" and TR4H_GATE_RE.match(f) and f not in ENTRY_BLOCKER_GATES))
        self.tab_filters = {t: tuple(sorted(v)) for t, v in self.tab_filters.items()}
        self.memo = {}
        self.pair_verdict = None
        try:
            from tools import v18_binding_map as BM
            self.pair_verdict = BM.pair_verdict
            self._bm = bm
        except Exception:
            self._bm = None

    def static(self, tab, sw, cand, hdr):
        """-> (cell dict or None, stage_switch, stage_filter, dead_row) for one cell, memoised."""
        f, _, o = str(hdr).strip().partition("=")
        k = (tab, sw, canon_value(cand), f.strip(), canon_value(o))
        got = self.memo.get(k)
        if got is not None:
            return got
        hit = self.rows.get((tab, sw, k[2]))
        cell = hit[1].get((k[3], k[4])) if hit else None
        if cell is None and self.pair_verdict is not None and self._bm is not None and (sw in self.switches or k[3] in self.filters):
            try:
                cell = self.pair_verdict(self._bm, sw, cand, f"{k[3]}={o}", [tab])
            except Exception:
                cell = None
        swi = self.switches.get(sw) or {}
        fam = swi.get("family")
        st_sw = fam if fam in BIND_STAGES else "UNKNOWN"
        fi = self.filters.get(k[3])
        st_f = _filter_stage(fi.get("stages")) if fi else "UNKNOWN"
        dead = bool((hit[0].get("dead_row") if hit else False) or swi.get("dead"))
        got = (cell, st_sw, st_f, dead)
        if len(self.memo) < 2000000:
            self.memo[k] = got
        return got

    def value(self, key, sw, cand, running_set):
        """effective value of key: the row's own switch=cand, else the running set, else the map's master default (None = unknown)."""
        if key == sw:
            return cand
        if running_set is not None and key in running_set:
            return running_set[key]
        return self.defaults.get(key)

    def cond_ok(self, term, sw, cand, running_set):
        if ">" in term and "=" not in term:
            k, _, thr = term.partition(">")
            v = self.value(k.strip(), sw, cand, running_set)
            try:
                return v is not None and float(v) > float(thr)
            except (TypeError, ValueError):
                return False
        k, _, want = term.partition("=")
        v = self.value(k.strip(), sw, cand, running_set)
        if v is None:
            return False
        return canon_value(v) == canon_value(want) or (canon_value(want) == "True" and _truthy(v) and canon_value(v) not in ("False",))

    def row_context(self, tab, sw, cand, running_set):
        """running-set features shared by every cell of a row: (upstream_gate_active, tab digest Counter)."""
        if running_set is None:
            return 0, Counter()
        gates = sum(1 for g in self.gates if _truthy(self.value(g, sw, cand, running_set)))
        bag = Counter()
        for f in self.tab_filters.get(tab, ()):
            if f != sw and f not in running_set:
                continue
            v = self.value(f, sw, cand, running_set)
            cv = canon_value(v)
            cls = (self.filters.get(f) or {}).get("class", "UNKNOWN")
            bag[f"{cls}|{_tf_bucket(cv)}|{'on' if _truthy(v) else 'off'}"] += 1
        return gates, bag


def load_binding_map(cat_side, root=ROOT, learn_dir=None):
    """-> (BindingIndex | None, why). Own cat_side only: a map built for another cat_side is refused (CAT_SIDE_MISMATCH),
    a wrong contract is refused (CONTRACT_MISMATCH); absent / unreadable -> None. Fail-open: None = every binding feature UNKNOWN."""
    p = Path(learn_dir) if learn_dir else Path(root) / "data" / "learn"
    p = p / f"binding_map_{cat_side}.json"
    if not p.exists():
        return None, f"absent {p}"
    try:
        bm = json.loads(p.read_text())
    except Exception as e:
        return None, f"unreadable {p}: {type(e).__name__}"
    meta = (bm or {}).get("meta") or {}
    if meta.get("cat_side") != cat_side:
        return None, f"CAT_SIDE_MISMATCH want={cat_side} got={meta.get('cat_side')} path={p}"
    if meta.get("contract") != BIND_CONTRACT:
        return None, f"CONTRACT_MISMATCH want={BIND_CONTRACT} got={meta.get('contract')} path={p}"
    return BindingIndex(bm, p), f"ok {p.name} built_at={meta.get('built_at')}"


def binding_features(bi, tab, sw, cand, hdr, running_set=None, row_ctx=None):
    """The v3 binding features of one cell (switch=cand on tab, filter header hdr) under a running set -> dict:
    bind_verdict, stage_switch, stage_filter, same_stage, masters_armed (None = unknown), dead_row, upstream_gate_active,
    tab_digest (Counter of class|tf-bucket|on-off of the tab's in-force filters). bi None (no map) -> everything UNKNOWN / 0.
    running_set None -> CONDITIONAL cells are UNKNOWN, masters_armed unknown, gates 0, empty digest."""
    out = {"bind_verdict": "UNKNOWN", "stage_switch": "UNKNOWN", "stage_filter": "UNKNOWN", "same_stage": False, "masters_armed": None,
           "dead_row": False, "upstream_gate_active": 0, "tab_digest": Counter()}
    if bi is None:
        return out
    cell, st_sw, st_f, dead = bi.static(tab, sw, cand, hdr)
    out.update(stage_switch=st_sw, stage_filter=st_f, dead_row=dead)
    fi = bi.filters.get(str(hdr).strip().partition("=")[0].strip()) or {}
    f_stages = {FILTER_STAGE_ALIAS.get(x, x) for s in (fi.get("stages") or []) for x in str(s).upper().split("+")}
    out["same_stage"] = st_sw != "UNKNOWN" and st_sw in f_stages
    gates, bag = row_ctx if row_ctx is not None else bi.row_context(tab, sw, cand, running_set)
    out["upstream_gate_active"], out["tab_digest"] = gates, bag
    if cell is None:
        return out
    v, keys, reason = cell.get("verdict"), list(cell.get("condition_keys") or []), cell.get("reason")
    master = fi.get("master")
    if v == "BINDS" and master and running_set is not None and master != sw and master in running_set and not _truthy(running_set[master]):
        v, keys = "CONDITIONAL", [f"{master}=True"]
    if v in ("BINDS", "INDIRECT", "NEVER"):
        out["bind_verdict"] = v
        out["masters_armed"] = 1.0 if running_set is not None else None
        return out
    if reason == "UNKNOWN_FAMILY" or not keys or running_set is None:
        return out
    ok = [bi.cond_ok(t, sw, cand, running_set) for t in keys]
    armed = any(ok) if cell.get("condition_mode") == "any" else all(ok)
    out["masters_armed"] = (1.0 if armed else 0.0) if cell.get("condition_mode") == "any" else sum(ok) / len(ok)
    out["bind_verdict"] = "CONDITIONAL_ARMED" if armed else "CONDITIONAL_OFF"
    return out


def _digest_slot(tok):
    return int.from_bytes(hashlib.blake2b(tok.encode(), digest_size=4).digest(), "big") % TAB_DIGEST_DIM


def bind_vector(bf):
    """binding feature dict -> dense list of BIND_DIM floats (order = BIND_FEATURE_NAMES) for the brain's online heads."""
    x = [0.0] * BIND_DIM
    x[BIND_VERDICTS.index(bf["bind_verdict"])] = 1.0
    o = len(BIND_VERDICTS)
    stages_sw = BIND_STAGES + ("UNKNOWN",)
    x[o + stages_sw.index(bf["stage_switch"])] = 1.0
    o += len(stages_sw)
    stages_f = BIND_STAGES + ("NONE", "UNKNOWN")
    x[o + stages_f.index(bf["stage_filter"])] = 1.0
    o += len(stages_f)
    x[o] = 1.0 if bf["same_stage"] else 0.0
    x[o + 1] = bf["masters_armed"] if bf["masters_armed"] is not None else 0.0
    x[o + 2] = 1.0 if bf["masters_armed"] is not None else 0.0
    x[o + 3] = 1.0 if bf["dead_row"] else 0.0
    x[o + 4] = min(4.0, float(bf["upstream_gate_active"])) / 2.0
    o += 5
    tot = sum(bf["tab_digest"].values())
    if tot:
        s = 1.0 / math.sqrt(tot)
        for tok, n in bf["tab_digest"].items():
            x[o + _digest_slot(tok)] += n * s
    return x


def binding_feature_strings(bf):
    """binding feature dict -> [(feature, x)] in the linknet's sparse string form (prefix B.); tab digest tokens are the
    attribute bag (class|tf-bucket|on-off), never raw KEY=value ids (E8D §9)."""
    out = [("B.bv:" + bf["bind_verdict"], 1.0), ("B.ss:" + bf["stage_switch"], 1.0), ("B.sf:" + bf["stage_filter"], 1.0)]
    if bf["same_stage"]:
        out.append(("B.same", 1.0))
    if bf["masters_armed"] is not None:
        out.append(("B.armed", float(bf["masters_armed"])))
    if bf["dead_row"]:
        out.append(("B.dead", 1.0))
    if bf["upstream_gate_active"]:
        out.append(("B.gates", min(4.0, float(bf["upstream_gate_active"])) / 2.0))
    tot = sum(bf["tab_digest"].values())
    for tok, n in sorted(bf["tab_digest"].items()):
        out.append(("B.td:" + tok, n / math.sqrt(tot)))
    return out


def cmd_build(a):
    out_dir = Path(a.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    from tools import v18_learn_common as C
    rc = 0
    for cs in a.cat_side or CAT_SIDES:
        doc = build(cs)
        doc["meta"]["built_at"] = C.utc()
        p = out_dir / f"features_{cs}.json"
        C.write_json(p, doc, indent=None)
        cov = doc["coverage"]
        sz = p.stat().st_size
        print(f"[features] {cs}: switches {cov['switches']} settings {cov['settings']}/{cov['settings_expected']} filters {cov['filters']} options "
              f"{cov['options']}/{cov['options_expected']} all_featurised={cov['all_featurised']} feats/setting {cov['mean_features_per_setting']} "
              f"unknown tokens {cov['n_unknown_tokens']} no-family {len(cov['names_without_family'])} no-role {len(cov['names_without_role'])} -> {p} ({sz} B)", flush=True)
        if sz > 10 * 1000 * 1000 or not cov["all_featurised"]:
            rc = 2
    return rc


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("--cat-side", action="append", choices=CAT_SIDES)
    b.add_argument("--out-dir", default=str(ROOT / "data" / "learn"))
    a = ap.parse_args(argv)
    return cmd_build(a)


if __name__ == "__main__":
    sys.exit(main())
