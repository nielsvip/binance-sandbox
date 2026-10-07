#!/usr/bin/env python3
"""v15_f_refill — validate and repair column F (HUSTLE_DELTA) of run sheets from the progress-JSON truth (NO-LIES: never fabricate).
Match: sheet row (tab, 'A=B') <-> JSON done key 'TAB!row:SWITCH=cand' by NAME (not row number, templates get re-ordered), non-running entries only.
  F==0 and JSON delta_vs_initial exactly 0.0            -> ok (real exact zero)
  F==0 and JSON number != 0                            -> fix to JSON number   (FAKE_ZERO)
  F==0 and JSON None / no entry / running               -> blank                (PLACEHOLDER_ZERO)
  F None and JSON number                                -> fill from JSON       (NONE_FILLABLE)
  F None and JSON None/no entry/running                 -> counted NONE_UNCALC (needs a pilot recalculation, not touched)
  F number != JSON number (|d|>1e-9)                    -> MISMATCH (counted, not touched)
Only column F is written; row guard (all other columns identical) + atomic zip-validated save + backup. Sheets whose pilot is running are skipped.
  python tools/v15_f_refill.py --sheets DIR --progress DIR [--fix] [--backup-dir DIR] [--limit N] [--counter FILE]"""
import argparse, collections, glob, json, os, shutil, subprocess, sys, time, zipfile
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__))))
import template_row_guard as G
SW = ['STDEV_SLOPE_SIZING','ENTRY_REVERSAL_BOUNCE','ENTRY_BREAKOUT_CHANNEL','ENTRY_CONFIRMATION_GATES','EXIT_STRUCTURAL','EXIT_VELOCITY','REENTRY_WINDOWED','REENTRY_ADAPTIVE','AUGMENT_TREND','AUGMENT_RISK_SIZING','REDUCE_PROFIT_LOCK','REDUCE_SIGNAL_RATER','GLOBAL_RISK_GATES']
def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)
def running_syms():
    try:
        o = subprocess.run(["pgrep", "-af", "v15_pilot.py"], capture_output=True, text=True).stdout
    except Exception: return set()
    r = set()
    for l in o.splitlines():
        t = l.split()
        if "--sym-side" in t: r.add(t[t.index("--sym-side") + 1])
    return r
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheets", required=True); ap.add_argument("--progress", required=True); ap.add_argument("--fix", action="store_true")
    ap.add_argument("--backup-dir"); ap.add_argument("--limit", type=int, default=0); ap.add_argument("--counter"); ap.add_argument("--min-age-min", type=float, default=3.0)
    ap.add_argument("--only", help="comma list of sym_sides")
    a = ap.parse_args()
    files = sorted(glob.glob(os.path.join(a.sheets, "*_30d_matrix.xlsx")))
    if a.only:
        keep = set(a.only.split(",")); files = [f for f in files if os.path.basename(f).split("_30d")[0] in keep]
    if a.limit: files = files[:a.limit]
    run = running_syms(); tot = collections.Counter(); per = {}
    for p in files:
        ss = os.path.basename(p).split("_30d")[0]
        c = collections.Counter()
        pj = os.path.join(a.progress, f"{ss}_v14_progress.json")
        if not os.path.exists(pj): tot["no_json_sheets"] += 1; continue
        if ss in run: tot["skipped_running"] += 1; continue
        if (time.time() - os.path.getmtime(p)) / 60 < a.min_age_min and a.fix: tot["skipped_recent_write"] += 1; continue
        try:
            done = json.load(open(pj)).get("done") or {}
            wb = openpyxl.load_workbook(p)
        except Exception as e:
            tot["unreadable"] += 1; continue
        idx = {}
        for k, v in done.items():
            if not isinstance(v, dict): continue
            tab, rest = k.split("!", 1); sw = rest.split(":", 1)[1] if ":" in rest else rest
            idx[(tab, sw)] = v
        before = openpyxl.load_workbook(p) if a.fix else None
        dirty = False
        for tab in SW:
            if tab not in wb.sheetnames: continue
            ws = wb[tab]
            for r in range(3, ws.max_row + 1):
                A, B = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
                if A in (None, "") or B in (None, ""): continue
                cell = ws.cell(row=r, column=6); f = cell.value
                e = idx.get((tab, f"{str(A).strip()}={str(B).strip()}"))
                jv = e.get("delta_vs_initial") if e and not e.get("is_running") else None
                jn = num(jv)
                c["cells"] += 1
                if num(f) and f == 0:
                    if jn and jv == 0.0: c["zero_real"] += 1
                    elif jn: c["zero_fake"] += 1; cell.value = float(jv) if a.fix else f; dirty |= a.fix
                    else: c["zero_placeholder"] += 1; cell.value = None if a.fix else f; dirty |= a.fix
                elif f is None:
                    if jn: c["none_fillable"] += 1; cell.value = float(jv) if a.fix else None; dirty |= a.fix
                    else: c["none_uncalc"] += 1
                elif num(f):
                    c["num_ok" if (not jn or abs(float(f) - float(jv)) <= 1e-9) else "num_mismatch"] += 1
                else: c["other_type"] += 1
        fixed = c["zero_fake"] + c["zero_placeholder"] + c["none_fillable"]
        if a.fix and dirty and fixed:
            for tab in SW:
                if tab in wb.sheetnames: G.assert_rows_intact(before[tab], wb[tab], {6})
            if a.backup_dir:
                os.makedirs(a.backup_dir, exist_ok=True); shutil.copy2(p, os.path.join(a.backup_dir, os.path.basename(p)))
            tmp = p + ".frefill.tmp"; wb.save(tmp)
            if len(zipfile.ZipFile(tmp).namelist()) < 10: raise SystemExit(f"zip validation failed {p}")
            openpyxl.load_workbook(tmp); os.replace(tmp, p); c["fixed"] = fixed
        per[ss] = dict(c); tot.update(c); tot["sheets"] += 1
    rem = tot["zero_fake"] + tot["zero_placeholder"] + tot["none_fillable"] - tot["fixed"]
    out = {"ts": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "host": os.uname().nodename, "cells_total": tot["cells"], "zero": tot["zero_fake"] + tot["zero_placeholder"], "none": tot["none_fillable"],
           "fixed": tot["fixed"], "remaining": rem, "detail": dict(tot)}
    if a.counter: json.dump(out, open(a.counter, "w"), indent=1)
    print(json.dumps(out))
if __name__ == "__main__": main()
