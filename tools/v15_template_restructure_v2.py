#!/usr/bin/env python3
"""v15_template_restructure_v2 — AGENT B (2026-09-30): STAGE a cleaned set of TEMPLATE_*.xlsx (never touches the live templates).

Reads SPREADSHEETS/TEMPLATE_{cat}.xlsx (read-only) and writes <out-dir>/TEMPLATE_{cat}.xlsx + <report-dir>/rows_{cat}.csv + summary JSON:
  1. ONE HOME per switch: a switch (col A) that was broadcast into several tabs is merged into its home tab (tools/v15_template_home_rules.py,
     evidence-based); ENTRY-tab switches whose name is EXIT/REENTRY/AUGMENT/REDUCE are moved to the right lifecycle tab. Whole rows move
     (yellow cells remapped by header name); duplicates go to sheet PARKED_ROWS (nothing is deleted).
  2. SIDE rule: a *_SHORT switch is not carried in a LONG template (dead there) and vice-versa; a missing own-side version is ADDED with the
     donor's options and the venue config default (never invented values).
  3. OPTION repair: options whose type does not fit the config field (e.g. '0.0' for a bool switch) are parked; exactly one default per
     switch (bold == is_default YES); the default option is added when absent; bool switches get both True/False.
  4. GREY = not calculable only: NOT_IN_CONFIG (no field in config/config_tradier/QuickConfig). NO_CONSUMER (only config/wiring stubs, nothing
     reads it in live or vec) is greyed as CONTAMINATION (proposal for erase). Everything else is un-greyed (must be calculated).
  5. ever-yellow json re-keyed to the new tabs (data/yellow_ever_nonzero.staged.json in the out dir).
Usage: python tools/v15_template_restructure_v2.py --out-dir SPREADSHEETS/TEMPLATE_STAGED/<ts> --report-dir data/template_audit/<ts> [--cat-side X]
"""
import os as _os, sys as _sys
if _os.path.exists(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "SPREADSHEETS", "TEMPLATES_FROZEN")) and "--dry-run" not in _sys.argv and not _os.environ.get("TEMPLATES_UNFREEZE"):
    _writes = any(a in _sys.argv for a in ("--apply", "--write", "--out-dir")) or "v15_daily_template_update" in __file__ or "restructure" in __file__
    if _writes and ("--apply" in _sys.argv or "v15_daily_template_update" in __file__ or "restructure" in __file__):
        _sys.exit("REFUSED: SPREADSHEETS/TEMPLATES_FROZEN exists (USER 2026-10-01: templates restored to the pre-trainwreck 20:36 version; no writer may touch them until the user approves a proposal). Remove the file only on explicit user approval.")
import argparse
import collections
import copy
import csv
import dataclasses
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import v15_template_home_rules as H  # noqa: E402

TEMPLATES = {"CRYPTO_LONG": "TEMPLATE_CRYPTO_LONG.xlsx", "CRYPTO_SHORT": "TEMPLATE_CRYPTO_SHORT.xlsx", "STOCKS_LONG": "TEMPLATE_STOCKS_LONG.xlsx", "STOCKS_SHORT": "TEMPLATE_STOCKS_SHORT.xlsx"}
FIXED = ["Switch", "default", "override", "Family", "BASELINE", "HUSTLE_DELTA", "VECTOR_DELTA", "LIVE_DELTA", "LIVE_SHARPE", "REAL_COMPLETE", "PER_ROW_FILTERS", "is_default", "AVG_DELTA", "POS_SYM"]
GREY = "FFBFBFBF"
ORANGE = "FFE699"
YELLOW = "FFFF00"
WIRE = re.compile(r"^(<module>|_ensure_.*|_batch\d.*|_wire_.*|apply_tradier_defaults|__init__|apply|run|_apply_research_only_live_gates)$")
VEC_FILE = "v12_quick_engine.py"   # vec side = v12_quick_engine.py + vec_decisions/*; every other scanned root module is live-side
MISSING = object()
NOT_SIDE = {"ATR_LONG_WINDOW"}
TF_RE = re.compile(r"^(OFF|D|W|M|\d+[mhdwMHDW])(,(OFF|D|W|M|\d+[mhdwMHDW]))*$")


def nrm(v):
    return "" if v is None else str(v).strip()


def coerce(opt, ref, name=""):
    """(ok, normalized key) of option `opt` for a config field whose current value is `ref`."""
    s = nrm(opt)
    if isinstance(ref, bool):
        if s.lower() in ("true", "false"):
            return True, s.lower()
        return False, s
    if isinstance(ref, (int, float)):
        try:
            return True, repr(float(s))
        except Exception:
            return False, s
    if isinstance(ref, str):
        if TF_RE.match(ref) or name.upper().endswith("_TF"):
            return bool(TF_RE.match(s)) or s.upper() in ("OFF", "NONE"), s
        return True, s
    return True, s


def same(a, b):
    if a == b:
        return True
    try:
        return not isinstance(a, bool) and not isinstance(b, bool) and float(a) == float(b)
    except Exception:
        return nrm(a).lower() == nrm(b).lower()


class Fields:
    def __init__(self, cs):
        import config
        import config_tradier
        import v12_quick_engine as V
        self.stocks = cs.startswith("STOCKS")
        self.live = config_tradier.TradierConfig() if self.stocks else config.Config()
        self.qc = V.QuickConfig()
        if self.stocks:
            self.qc.apply_tradier_defaults()
        self.known = {f.name for f in dataclasses.fields(V.QuickConfig)} | {k for k in dir(config.Config) if not k.startswith("_")} | {k for k in dir(config_tradier.TradierConfig) if not k.startswith("_")}
        try:
            self.cs4 = (json.loads((ROOT / "data" / "cat_side_defaults_4.json").read_text()).get(cs) or {})
        except Exception:
            self.cs4 = {}

    def get(self, name):
        for o in (self.live, self.qc):
            try:
                return getattr(o, name)
            except Exception:
                continue
        return MISSING


def load_usage():
    p = ROOT / "data" / "template_audit"
    cands = sorted(p.glob("*/code_usage_full.json"))
    return json.loads(cands[-1].read_text()) if cands else {}


def consumers(usage, name):
    live = vec = 0
    for f, fn, c in usage.get(name, []):
        if WIRE.match(fn):
            continue
        if f == VEC_FILE or f.startswith("vec_decisions/"):
            vec += c
        else:
            live += c
    return live, vec


class Sheet:
    """one tab read into memory: headers by name; rows as {"rec": {hname: (value, style_array, comment_text)}, yellow/orange/grey/bold}"""

    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.hdr = {}
        for c in range(1, ws.max_column + 1):
            h = ws.cell(2, c).value
            if h in (None, ""):
                continue
            h = str(h)
            if h not in self.hdr:
                self.hdr[h] = c
        self.meta_cols = {h for h in self.hdr if h in FIXED or h in ("Live Location", "Vec Hook", "Type", "Correct Options")}
        self.rows = []
        for r in range(3, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a in (None, ""):
                continue
            rec, yel = {}, set()
            for h, c in self.hdr.items():
                cell = ws.cell(r, c)
                rec[h] = (cell.value, copy.copy(cell._style), cell.comment.text if cell.comment else None)
                if h not in self.meta_cols:
                    fl = cell.fill
                    if fl is not None and fl.fill_type == "solid" and str(fl.fgColor.rgb or "").upper().endswith(YELLOW):
                        yel.add(h)
            fa = ws.cell(r, 1).fill
            fo = ws.cell(r, 1).font
            self.rows.append({"rec": rec, "r": r, "tab": ws.title, "yellow": yel,
                              "orange": bool(fa and fa.fgColor and ORANGE in str(fa.fgColor.rgb or "")),
                              "grey": bool(fo and fo.color is not None and str(getattr(fo.color, "rgb", "") or "").upper() == GREY),
                              "bold": bool(ws.cell(r, 2).font.b), "foreign": False})


def is_orange(row):
    return row["orange"]


def yellow_hdrs(row):
    return set(row["yellow"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src-dir", default=str(ROOT / "SPREADSHEETS"))
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--report-dir", required=True)
    ap.add_argument("--cat-side", default="ALL")
    a = ap.parse_args()
    out, rep = Path(a.out_dir), Path(a.report_dir)
    out.mkdir(parents=True, exist_ok=True)
    rep.mkdir(parents=True, exist_ok=True)
    usage = load_usage()
    ever = json.loads((ROOT / "data" / "yellow_ever_nonzero.json").read_text())
    ever_out = {"built_at": ever.get("built_at"), "rule": ever.get("rule"), "symsides": ever.get("symsides"), "cells": {}}
    # donors for missing own-side switches: read all 4 templates' groups once (options only)
    summary = {}
    sheets_all = {}
    for cs in (list(TEMPLATES) if a.cat_side == "ALL" else [a.cat_side]):
        sheets_all[cs] = openpyxl.load_workbook(str(Path(a.src_dir) / TEMPLATES[cs]))
    donor_opts = collections.defaultdict(dict)  # name -> {cs: [(B, bold, recdict)]}
    for cs, wb in sheets_all.items():
        for tab in H.TABS:
            sh = Sheet(wb[tab])
            for row in sh.rows:
                row["foreign"] = True
                n = str(row["rec"]["Switch"][0]).strip()
                donor_opts[n].setdefault(cs, []).append(row)
    for cs, wb in sheets_all.items():
        summary[cs] = build_one(cs, wb, usage, ever, ever_out, donor_opts, out, rep)
    (out / "yellow_ever_nonzero.staged.json").write_text(json.dumps(ever_out))
    (rep / "summary.json").write_text(json.dumps(summary, indent=1, default=str))
    print(json.dumps({k: v["counts"] for k, v in summary.items()}, indent=1))


def build_one(cs, wb, usage, ever, ever_out, donors, out, rep):
    F = Fields(cs)
    side = cs.split("_")[1]
    other = "SHORT" if side == "LONG" else "LONG"
    sheets = {t: Sheet(wb[t]) for t in H.TABS}
    by_sw = collections.OrderedDict()
    for t in H.TABS:
        g = collections.OrderedDict()
        for row in sheets[t].rows:
            g.setdefault(str(row["rec"]["Switch"][0]).strip(), []).append(row)
        for n, rows in g.items():
            by_sw.setdefault(n, []).append((t, rows))
    csv_rows, parked, unclear = [], [], []
    counts = collections.Counter()
    final = {t: [] for t in H.TABS}
    V = lambda row, h: row["rec"][h][0]

    def park(row, reason, name):
        parked.append({"reason": reason, "from_tab": row["tab"], "from_row": row["r"], "Switch": name, "default": nrm(V(row, "default")), "Family": nrm(V(row, "Family")), "is_default": nrm(V(row, "is_default")), "AVG_DELTA": V(row, "AVG_DELTA"), "POS_SYM": V(row, "POS_SYM"), "yellow": len(row["yellow"])})
        csv_rows.append([cs, row["tab"], row["r"], name, nrm(V(row, "default")), "PARKED", reason, "", ""])
        counts["parked:" + reason.split(":")[0]] += 1

    def classify_imp(name, lv, vc, why):
        if name in ("DELTA_GATE_BB_SQUEEZE", "HTF_GATE_D_MANDATORY", "HTF_TREND_VETO_BYPASS_ENABLED", "HTF_TREND_VETO_BYPASS_REASONS", "BTC_ACCEL_RAMP_REQUIRE_POSITIVE") and not (lv or vc):
            return "UNWIRED"
        if "unclear" in why.lower():
            return "UNCLEAR"
        return "HIGH" if (lv and vc) else ("MED" if (lv or vc) else "LOW")

    for name, copies in by_sw.items():
        s_side = None if name in NOT_SIDE else H.side_of(name)
        tabs_now = [t for t, _ in copies]
        if s_side == other:
            for t, rows in copies:
                for row in rows:
                    park(row, f"WRONG_SIDE: {name} is a {s_side} switch; this is a {side} template", name)
            continue
        ref = F.get(name)
        in_cfg = name in F.known
        manual = H.MANUAL_HOME.get(name)
        if manual:
            home, why = manual
        elif len(set(tabs_now)) == 1:
            t0 = tabs_now[0]
            L = H.strong_lifecycle(name)
            if H.LIFE[t0] == "ENTRY" and L and L != "ENTRY":
                home, why = H.default_tab_for(L, name), f"name says {L} but filed in ENTRY tab {t0}"
            else:
                home, why = t0, "single-tab: stays"
                if L and L != H.LIFE[t0]:
                    unclear.append([cs, name, t0, f"name looks {L} (-> {H.default_tab_for(L, name)}) but filed in {H.LIFE[t0]} tab {t0}; left in place"])
        else:
            home, why = sorted(set(tabs_now), key=H.TABS.index)[0], "multi-tab, no rule (UNCLEAR)"
            unclear.append([cs, name, ",".join(sorted(set(tabs_now))), "broadcast into several tabs with no home rule; first tab chosen"])
        allrows = [row for _, rows in copies for row in rows]
        homerows = [row for row in allrows if row["tab"] == home]
        opts = collections.OrderedDict()
        bad = []
        for row in sorted(allrows, key=lambda r: (r["tab"] != home, H.TABS.index(r["tab"]))):
            ok, key = coerce(V(row, "default"), ref, name) if ref is not MISSING else (True, nrm(V(row, "default")))
            if not ok:
                bad.append(row)
                continue
            o = opts.setdefault(key, {"row": row, "yell": set(), "isd": {}, "src": []})
            o["yell"] |= row["yellow"]
            o["src"].append(row)
            o["isd"][row["tab"]] = (str(V(row, "is_default") or "").upper() == "YES") or row["bold"]
        for row in bad:
            park(row, f"OPTION_TYPE_MISMATCH: '{nrm(V(row, 'default'))}' does not fit {type(ref).__name__} field {name}", name)
        tmpl_row = (homerows or allrows)[0]
        dkey = None
        for t in [home] + [x for x in H.TABS if x != home]:
            for k, o in opts.items():
                if o["isd"].get(t):
                    dkey = k
                    break
            if dkey:
                break
        want = MISSING
        if ref is not MISSING:
            want = F.cs4.get(name, ref)
        added = []
        if dkey is None and want is not MISSING:
            ok, k = coerce(want, ref, name)
            if k in opts:
                dkey = k
            else:
                opts[k] = {"row": None, "yell": set(tmpl_row["yellow"]), "isd": {}, "src": [], "new": nrm(want)}
                dkey, _ = k, added.append(nrm(want))
        if dkey is None:
            dkey = next(iter(opts), None)
        if dkey is None:
            continue
        if isinstance(ref, bool):
            for bk in ("true", "false"):
                if bk not in opts:
                    opts[bk] = {"row": None, "yell": set(tmpl_row["yellow"]), "isd": {}, "src": [], "new": "True" if bk == "true" else "False"}
                    added.append(opts[bk]["new"])
        if isinstance(ref, (int, float)) and not isinstance(ref, bool) and len(opts) >= 3:
            try:
                dv = float(dkey)
                for k_ in opts:
                    fv = float(k_)
                    if dv != 0 and fv != 0 and (abs(fv / dv) > 20 or abs(fv / dv) < 0.05):
                        unclear.append([cs, name, home, f"outlier option {k_} vs default {dkey} ({fv / dv:.3g}x): suspicious placeholder value, kept"])
            except Exception:
                pass
        lv, vc = consumers(usage, name)
        if not in_cfg:
            grey, greyw = True, "NOT_IN_CONFIG"
        elif isinstance(ref, (list, dict, tuple, set)):
            grey, greyw = True, "STRUCTURED_VALUE (config field is a list/dict; template options cannot express it)"
        elif lv == 0 and vc == 0 and not name.endswith("_FILTER_TF") and name not in KEEP_CALC:
            grey, greyw = True, "NO_CONSUMER (config/wiring stubs only)"
        else:
            grey, greyw = False, ""
        imp = classify_imp(name, lv, vc, why)
        if greyw.startswith("STRUCTURED"):
            imp = "UNCLEAR"
        orange = (homerows[0] if homerows else tmpl_row)["orange"]
        final[home].append({"name": name, "opts": opts, "dkey": dkey, "orange": orange, "grey": grey, "greyw": greyw, "home": home, "why": why, "ref": ref, "imp": imp, "live": lv, "vec": vc, "from": sorted(set(tabs_now)), "added": added, "tmpl": tmpl_row, "was_grey": all(r["grey"] for r in allrows)})
        counts["switches"] += 1
        if home not in tabs_now:
            counts["moved_to_home"] += 1
        elif len(set(tabs_now)) > 1:
            counts["dedup_broadcast"] += 1
        if grey:
            counts["grey:" + greyw.split(" ")[0]] += 1
        elif all(r["grey"] for r in allrows):
            counts["ungrey"] += 1
        for t, rows in copies:
            for row in rows:
                if row in bad:
                    continue
                act = "KEEP" if t == home else ("MOVED" if len(set(tabs_now)) == 1 else "DUP_MERGED")
                csv_rows.append([cs, t, row["r"], name, nrm(V(row, "default")), act, why if act != "KEEP" else "", home, imp])
    # ---- missing own-side switches (added from donors; options = donor's valid options, default = config value)
    present = {g["name"] for t in final for g in final[t]}
    for name, per in donors.items():
        if name in present or name in NOT_SIDE or H.side_of(name) != side:
            continue
        ref = F.get(name)
        if ref is MISSING or name not in F.known:
            continue
        ven = cs.split("_")[0]
        ov = "STOCKS" if ven == "CRYPTO" else "CRYPTO"
        drows = next((per[d] for d in (f"{ven}_{other}", f"{ov}_{side}", f"{ov}_{other}") if d in per), None)
        if not drows:
            continue
        home = H.MANUAL_HOME.get(name, (None,))[0] or H.default_tab_for(H.strong_lifecycle(name) or "ENTRY", name)
        opts = collections.OrderedDict()
        for row in drows:
            ok, key = coerce(V(row, "default"), ref, name)
            if ok and key not in opts:
                opts[key] = {"row": row, "yell": set(row["yellow"]), "isd": {}, "src": []}
        want = F.cs4.get(name, ref)
        ok, dk = coerce(want, ref, name)
        added = ["(whole switch)"]
        if dk not in opts:
            tmpl0 = next(iter(opts.values()))["row"]
            opts[dk] = {"row": None, "yell": set(tmpl0["yellow"]), "isd": {}, "src": [], "new": nrm(want)}
            added.append(nrm(want))
        tmpl = next(o["row"] for o in opts.values() if o["row"])
        lv, vc = consumers(usage, name)
        final[home].append({"name": name, "opts": opts, "dkey": dk, "orange": False, "grey": False, "greyw": "", "home": home, "why": f"ADDED: own-side ({side}) version missing in this template; donor options, venue config default", "ref": ref, "imp": "HIGH" if (lv and vc) else "MED", "live": lv, "vec": vc, "from": [], "added": added, "tmpl": tmpl, "was_grey": False})
        counts["added_side_switch"] += 1
        csv_rows.append([cs, "", "", name, nrm(want), "ADDED", "missing own-side version", home, "HIGH"])
    # ---- write staged workbook (wb is our private copy; live templates are never written)
    yellow_fill = PatternFill(fill_type="solid", start_color="FFFFFF00", end_color="FFFFFF00")
    nofill = PatternFill(fill_type=None)
    rekey_ever = collections.defaultdict(set)
    hdr_new = {}
    for t in H.TABS:
        ws = wb[t]
        tgt = sheets[t].hdr
        hdr_new[t] = tgt
        plain = next((r for r in sheets[t].rows if not r["foreign"]), None)
        ws.delete_rows(3, ws.max_row)
        r = 3
        for g in sorted(final[t], key=lambda x: x["orange"]):
            keys = list(g["opts"])
            for k in keys:
                o = g["opts"][k]
                is_def = k == g["dkey"]
                srow = o["row"] or g["tmpl"]
                foreign = srow["foreign"]
                styl = plain if foreign else srow
                for h, c in tgt.items():
                    cell = ws.cell(r, c)
                    if styl is not None and h in styl["rec"]:
                        cell._style = copy.copy(styl["rec"][h][1])
                    if h in sheets[t].meta_cols:
                        v = srow["rec"].get(h)
                        cell.value = v[0] if v else None
                    elif not foreign and o["row"] is not None:
                        cell.value = srow["rec"][h][0] if h in srow["rec"] else None
                    else:
                        cell.value = None
                    if h not in sheets[t].meta_cols and (foreign or styl is not srow):
                        cell.fill = copy.copy(nofill)
                for h in o["yell"]:
                    if h in tgt:
                        ws.cell(r, tgt[h]).fill = copy.copy(yellow_fill)
                if o["row"] is None:
                    ws.cell(r, tgt["default"]).value = o.get("new")
                    stat = (None, None)
                elif not o["src"]:
                    stat = (None, None)
                else:
                    bst = max(o["src"], key=lambda s_: (s_["tab"] == t, s_["rec"]["POS_SYM"][0] if isinstance(s_["rec"]["POS_SYM"][0], (int, float)) else -1))
                    stat = (bst["rec"]["AVG_DELTA"][0], bst["rec"]["POS_SYM"][0])
                ws.cell(r, tgt["AVG_DELTA"]).value, ws.cell(r, tgt["POS_SYM"]).value = stat
                for h in ("override", "BASELINE", "HUSTLE_DELTA", "VECTOR_DELTA", "LIVE_DELTA", "LIVE_SHARPE", "REAL_COMPLETE", "PER_ROW_FILTERS"):
                    if h in tgt:
                        ws.cell(r, tgt[h]).value = None
                bf = ws.cell(r, 2).font
                ws.cell(r, 2).font = Font(name=bf.name or "Arial", size=bf.size or 10, bold=is_def, italic=bf.italic, color=bf.color)
                ws.cell(r, tgt["is_default"]).value = "YES" if is_def else "NO"
                ws.cell(r, tgt["is_default"]).font = Font(name="Arial", size=10, bold=is_def)
                af = ws.cell(r, 1).font
                ws.cell(r, 1).font = Font(name=af.name or "Arial", size=af.size or 10, bold=af.bold, italic=af.italic, color=(GREY if g["grey"] else None))
                ws.cell(r, 1).value = g["name"]
                ws.cell(r, 1).fill = PatternFill(fill_type="solid", start_color="FF" + ORANGE, end_color="FF" + ORANGE) if g["orange"] else copy.copy(nofill)
                if k == keys[0] and (g["added"] or g["home"] not in g["from"] or g["grey"] or g["was_grey"] or g["imp"] in ("UNCLEAR", "UNWIRED")):
                    txt = f"v2 staged: {g['why']}" + (f" | ADDED {g['added']}" if g["added"] else "") + (f" | GREY {g['greyw']}" if g["grey"] else "") + (" | was grey, now calculated" if g["was_grey"] and not g["grey"] else "") + (f" | importance {g['imp']}" if g["imp"] in ("UNCLEAR", "UNWIRED") else "")
                    ws.cell(r, 1).comment = Comment(txt[:500], "template_restructure_v2")
                for sr in o["src"]:
                    rekey_ever[(sr["tab"], g["name"], nrm(sr["rec"]["default"][0]))].add((t, nrm(ws.cell(r, 2).value)))
                r += 1
    if "PARKED_ROWS" in wb.sheetnames:
        del wb["PARKED_ROWS"]
    pk = wb.create_sheet("PARKED_ROWS")
    cols = ["reason", "from_tab", "from_row", "Switch", "default", "Family", "is_default", "AVG_DELTA", "POS_SYM", "yellow"]
    pk.append(cols)
    for p_ in parked:
        pk.append([p_[c] for c in cols])
    wb.save(str(out / TEMPLATES[cs]))
    old = ever["cells"].get(cs, [])
    newset, dropped = set(), 0
    hsets = {t: set(hdr_new[t]) for t in H.TABS}
    for key in old:
        tab, sb, hdr = key.split("\t", 2)
        sw, cand = sb.split("=", 1)
        tg = rekey_ever.get((tab, sw, cand.strip()))
        if not tg:
            dropped += 1
            continue
        for nt, nc in tg:
            if hdr in hsets[nt]:
                newset.add(f"{nt}\t{sw}={nc}\t{hdr}")
            else:
                dropped += 1
    ever_out["cells"][cs] = sorted(newset)
    with open(rep / f"rows_{cs}.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["template", "tab_now", "row_now", "switch", "option", "action", "reason", "home_tab", "importance"])
        w.writerows(csv_rows)
    counts.update({"parked_total": len(parked), "ever_yellow_old": len(old), "ever_yellow_new": len(newset), "ever_yellow_dropped": dropped})
    sw_info = [{"name": g["name"], "home": g["home"], "from": g["from"], "grey": g["greyw"], "imp": g["imp"], "live": g["live"], "vec": g["vec"], "added": g["added"], "why": g["why"], "type": ("MISSING" if g["ref"] is MISSING else type(g["ref"]).__name__), "default": nrm(g["opts"][g["dkey"]].get("new") if g["opts"][g["dkey"]]["row"] is None else g["opts"][g["dkey"]]["row"]["rec"]["default"][0]), "n_opts": len(g["opts"]), "n_copies": len(g["from"])} for t in final for g in final[t]]
    with open(rep / f"switches_{cs}.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["switch", "home_tab", "was_in_tabs", "n_options", "default", "cfg_type", "grey", "importance", "live_consumers", "vec_consumers", "added", "note"])
        for x in sw_info:
            w.writerow([x["name"], x["home"], "|".join(x["from"]), x["n_opts"], x["default"], x["type"], x["grey"], x["imp"], x["live"], x["vec"], "|".join(x["added"]), x["why"]])
    return {"counts": dict(counts), "unclear": unclear, "switches": sw_info}


KEEP_CALC = {"HTF_GATE_D_MANDATORY", "HTF_TREND_VETO_BYPASS_ENABLED", "HTF_TREND_VETO_BYPASS_REASONS", "DELTA_GATE_BB_SQUEEZE", "WT_DC_STOCH_THRESHOLD_LONG", "WT_DC_STOCH_THRESHOLD_SHORT", "WT_DC_DC_POS_THRESHOLD_LONG", "WT_DC_DC_POS_THRESHOLD_SHORT", "BTC_ACCEL_RAMP_REQUIRE_POSITIVE"}


if __name__ == "__main__":
    main()
