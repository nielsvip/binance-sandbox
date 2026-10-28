#!/usr/bin/env python3
"""v15_new_rows_yellow_plan — per-row yellow placement plan for the NEW template rows (operator order 2026-10-28).

READ-ONLY on the workbooks. For every switch in scope (0 naked evals in the union corpus, the 61 switches that got a family
only on 10-28, every switch added to TEMPLATE_FINAL_NORM since the 10-22 binding maps, and every never-tried row) it
records what the switch does in LIVE (SWITCH_BIBLE live/vec read sites, verified against the source line text), the stage
it acts in, and — per template row — which FILTER=opt cells can plausibly give a positive delta (binding-map rules of
22_SWITCH_FAMILY_MAP §c evaluated under the cat_side bold defaults, plus measured sibling-family / same-switch evidence)
versus which can NEVER bind. Diff against the row's current shading -> yellows to ADD / REMOVE, each with its reason.

  python3 tools/v15_new_rows_yellow_plan.py --templates DIR --out data/reports/new_rows_yellow_placement_20261028

Output: plan_<CS>.json (consumed by tools/v15_template_bookkeeper.py --paint-yellows) + PLAN.md.
"""
import argparse
import datetime
import gzip
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(1, str(ROOT))
import v18_binding_map as B  # noqa: E402

CAT_SIDES = B.CAT_SIDES
SWITCH_SHEETS = B.SWITCH_SHEETS
FAMILY_MAP = ROOT / "data" / "learn" / "switch_family_map_20261022.json"
FAMILY_EXT = ROOT / "data" / "learn" / "switch_family_map_ext_20261028.json"
OLD_MAPS = ROOT / "data" / "learn" / "v3u3_20261022"
FULL = ROOT / "data" / "learn" / "v20u4_FULL_20261028"
BIBLE = ROOT / "data" / "SWITCH_BIBLE.json"
YELLOW = ("FFFFFF00", "FFFFF2CC")
LIGHT = "FFFFF2CC"
GREY_FONT = "FFBFBFBF"
ORANGE_TOKENS = ("FFC000", "F4B084", "FFE699", "F8CBAD", "FCE4D6", "FFD966")
FILTER_COL0 = 15  # column O
TF_TOKENS = {"15M": "15m", "1H": "1h", "4H": "4h", "D": "D", "1D": "D", "DAILY": "D", "W": "W", "5M": "5m", "30M": "30m"}
IND_TOKENS = ("WT", "DC", "BB", "EMA", "SMA", "RSI", "ADX", "STOCH", "ATR", "MACD", "VOL", "VWAP", "HA", "KG", "KINDERGARTEN", "MOM", "MOM3", "SQUEEZE", "STDEV", "LH", "LL", "HH", "HL", "HLR", "FUNDING", "OI", "BTC", "RATIO", "GR", "GOLDEN", "SPIKE", "DIV", "VEL", "ACCEL", "SLOPE", "RZ", "CCI", "OBV", "PIVOT")
LIVE_FILES = ("ez_manage.py", "tradier_manage.py", "ez_positions_quick.py", "btc_loop.py")
STAGE_OF_KIND = {"entry": "ENTRY", "reentry": "REENTRY", "reduce": "REDUCE", "augment": "AUGMENT", "sizing": "SIZING", "global": "GLOBAL"}
ARTEFACT_SWITCHES = ("MODE", "VENUE", "IS_TRADIER", "SIMPLE_PRICE_GT0_ENABLED")  # tools/v15_trade_autopsy_run.py AUTOPSY_DENY
AMBIG_FN = ("process_position", "_execute_now_inner", "simulate_one", "<module>", "step", "apply", "worker", "run", "main", "__init__")


def md5_12(p):
    return hashlib.md5(Path(p).read_bytes()).hexdigest()[:12]


def vstr(v):
    return B.vstr(v)


def _rgb(cell):
    try:
        f = cell.fill
        if f is None or f.fill_type != "solid":
            return None
        return str(f.fgColor.rgb or "").upper() or None
    except Exception:
        return None


def read_workbook(path):
    """{tab: {"headers": {hdr: col}, "rows": [{row, switch, cand, bold, grey, orange, shaded:{hdr: rgb}}]}}, header_bold {hdr: bool}"""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    out, hbold = {}, {}
    try:
        for tab in SWITCH_SHEETS:
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            hdrs = {}
            for c in range(FILTER_COL0, ws.max_column + 1):
                h = ws.cell(2, c).value
                if isinstance(h, str) and "=" in h and not h.upper().startswith("WHAT SWITCH"):
                    hdrs[h.strip()] = c
                    hbold[h.strip()] = hbold.get(h.strip(), False) or bool(ws.cell(2, c).font is not None and ws.cell(2, c).font.b)
                elif isinstance(h, str) and h.strip().upper().startswith("WHAT SWITCH"):
                    break
            rows = []
            for r in range(3, ws.max_row + 1):
                sw = ws.cell(r, 1).value
                if sw in (None, ""):
                    continue
                a = ws.cell(r, 1)
                fc = str(a.font.color.rgb).upper() if (a.font is not None and a.font.color is not None and isinstance(a.font.color.rgb, str)) else ""
                arg = _rgb(a) or ""
                shaded = {}
                for h, c in hdrs.items():
                    rgb = _rgb(ws.cell(r, c))
                    if rgb in YELLOW:
                        shaded[h] = rgb
                rows.append({"row": r, "switch": str(sw).strip(), "cand": ws.cell(r, 2).value, "bold": bool(ws.cell(r, 2).font is not None and ws.cell(r, 2).font.b), "grey": fc.endswith("BFBFBF"), "orange": any(t in arg for t in ORANGE_TOKENS), "shaded": shaded})
            out[tab] = {"headers": hdrs, "rows": rows}
    finally:
        wb.close()
    return out, hbold


def line_text(rel, line):
    try:
        p = ROOT / rel
        lines = p.read_text(errors="replace").splitlines()
        return lines[line - 1].strip()[:220] if 0 < line <= len(lines) else ""
    except Exception:
        return ""


def grep_name(rel, name, limit=3):
    """non-stub, non-list, non-declaration source lines naming `name` (fallback when a bible line moved)"""
    out = []
    try:
        for i, ln in enumerate((ROOT / rel).read_text(errors="replace").splitlines(), 1):
            if name not in ln:
                continue
            s = ln.strip()
            if s.startswith("#") or re.match(rf"^['\"]{re.escape(name)}['\"],?$", s) or s.startswith("_ = getattr") or re.match(rf"^{re.escape(name)}\s*:", s):
                continue
            out.append({"site": f"{rel}:{i}", "text": s[:220]})
            if len(out) >= limit:
                break
    except Exception:
        pass
    return out


_SITE = re.compile(r"^(?P<file>[^:]+):(?P<line>\d+)(?:\((?P<fn>[^)]*)\))?")


def verify_sites(sites, name):
    """bible sites 'file:line(fn)' -> [{site, fn, text, verified}] ; a moved line is re-located by grep in the same file"""
    out, seen = [], set()
    for s in sites or []:
        m = _SITE.match(str(s))
        if not m or s in seen:
            continue
        seen.add(s)
        rel, ln, fn = m.group("file"), int(m.group("line")), m.group("fn") or ""
        t = line_text(rel, ln)
        if name in t:
            out.append({"site": f"{rel}:{ln}", "fn": fn, "text": t, "verified": True})
            continue
        g = grep_name(rel, name, 1)
        if g:
            out.append({"site": g[0]["site"], "fn": fn, "text": g[0]["text"], "verified": True, "moved_from": f"{rel}:{ln}"})
        else:
            out.append({"site": f"{rel}:{ln}", "fn": fn, "text": t, "verified": False})
    return out


def config_comment(name):
    for rel in ("config.py", "config_tradier.py"):
        try:
            for i, ln in enumerate((ROOT / rel).read_text(errors="replace").splitlines(), 1):
                if re.match(rf"^\s+{re.escape(name)}\s*[:=]", ln):
                    c = ln.split("#", 1)[1].strip() if "#" in ln else ""
                    return {"site": f"{rel}:{i}", "decl": ln.strip()[:200], "comment": c[:200]}
        except Exception:
            continue
    return None


def fn_stage(fn):
    f = (fn or "").lower()
    if not f or any(f.endswith(a.lower()) or f == a.lower() for a in AMBIG_FN):
        return None
    if "compute_reentry_blocks" in f:
        return "ENTRY"  # v12 compute_reentry_blocks = the entry gate blocks (22_SWITCH_FAMILY_MAP §a)
    if "reentry" in f or "re_entry" in f:
        return "REENTRY"
    if "augment" in f:
        return "AUGMENT"
    if "reduce" in f or "profit_lock" in f or "partial" in f:
        return "REDUCE"
    if "exit" in f or "close" in f or "stop" in f:
        return "EXIT"
    if "entry" in f or "enter" in f or "open" in f:
        return "ENTRY"
    if "size" in f or "qty" in f or "sizing" in f:
        return "SIZING"
    return None


def exit_subtype(name):
    u = name.upper()
    if any(t in u for t in ("STOP", "LOSS_KILL", "MAX_PAIN", "NEWBORN")):
        return "EXIT_STOP"
    if any(t in u for t in ("TARGET", "_TP_", "TAKE_PROFIT")):
        return "EXIT_TARGET"
    if any(t in u for t in ("VEL", "_WT_", "WT_", "CROSS", "_DIV", "ACCEL")):
        return "EXIT_VEL"
    return "EXIT_STRUCT"


def condition_of(name, cands):
    toks = [t for t in name.upper().split("_") if t]
    tfs = sorted({TF_TOKENS[t] for t in toks if t in TF_TOKENS})
    for c in cands:
        for part in str(c).split(","):
            p = part.strip()
            if p in ("15m", "1h", "4h", "D", "W", "5m", "30m"):
                tfs.append(p)
    tfs = sorted(set(tfs))
    ind = sorted({t for t in toks if t in IND_TOKENS})
    d = "LONG" if "LONG" in toks else "SHORT" if "SHORT" in toks else "both"
    venue = "stocks" if ("TRADIER" in toks or "STOCKS" in toks or "TR" == toks[0]) else "crypto" if "CRYPTO" in toks else "both"
    return {"tfs": tfs, "indicators": ind, "direction": d, "venue_hint": venue}


def load_positives(cs):
    p = FULL / "stage" / "tools" / f"v18_data_learn_positives_all_{cs}.json.gz"
    try:
        return json.load(gzip.open(p))
    except Exception:
        return {"pairs": [], "rows": []}


def load_needed(cs):
    p = FULL / "stage" / "tools" / f"v18_data_learn_needed_filters_{cs}.json"
    try:
        return json.loads(p.read_text())
    except Exception:
        return {}


def load_linknet(cs):
    p = FULL / "stage" / "tools" / f"v18_data_learn_linknet_proposals_{cs}.json"
    try:
        return json.loads(p.read_text()).get("template_untried") or []
    except Exception:
        return []


def pilot_attach_fn(tab_level_path):
    """-> f(cs, tab, header_map, switch, cand, defaults) = set of headers the CURRENT v15_pilot attaches to the row (name-token rule,
    mandatory bases, WT-core all-columns, ever-yellow, tab-level / unwired / type exclusions) — the cells a board really evaluates."""
    import v15_pilot as P
    if tab_level_path:
        P._TAB_LEVEL_CACHE["spec"] = json.loads(Path(tab_level_path).read_text()).get("cats", {})
    cache = {}

    def f(cs, tab, hm, switch, cand, defaults):
        k = (cs, tab, switch)
        if k not in cache:
            cache[k] = P._header_disc_precompute(tab, switch, hm, P.tab_level_filters(cs, tab), P.mandatory_bases_for_tab(tab, switch), defaults)
        all_cols = switch in P.WT_CORE_MANDATORY_BASES
        return set(P._header_disc_combine(cache[k], tab, switch, cand, P.ever_yellow_cells(cs), all_cols)["hdrs"])
    return f


def build_cs(cs, tpl_path, fm, ext, bible, groups, agk, never_tried, attach=None):
    tabs, hbold = read_workbook(tpl_path)
    # ---- defaults (bold) exactly as v18_binding_map.build_one
    defaults = {}
    for name, e in bible.items():
        v, s = B.bible_default(e, cs)
        if s is not None:
            defaults[name] = v
    for t in tabs.values():
        for r in t["rows"]:
            if r["bold"] and not r["orange"]:
                defaults[r["switch"]] = r["cand"]
    for h, bld in hbold.items():
        f, o = h.split("=", 1)
        if bld and f not in {r["switch"] for t in tabs.values() for r in t["rows"] if r["bold"]}:
            defaults[f] = o
    filt_default_hdr = {}
    for h, bld in hbold.items():
        if bld:
            filt_default_hdr[h.split("=", 1)[0]] = h
    # ---- scope
    old = json.loads((OLD_MAPS / f"binding_map_{cs}.json").read_text())
    old_sw = {k.split("!", 1)[1].split("=", 1)[0] for k in old["rows"]}
    cur_rows = [(tab, r) for tab, t in tabs.items() for r in t["rows"] if not r["orange"]]
    cur_sw = {r["switch"] for _, r in cur_rows}
    nt = never_tried.get(cs) or {}
    s0 = set(nt.get("switches_0_evals") or [])
    r0 = set(nt.get("rows_0_evals_keys") or [])
    new_sw = cur_sw - old_sw
    ext_sw = set(ext) & cur_sw
    scope_sw = (s0 | new_sw | ext_sw) & cur_sw
    reasons = defaultdict(list)
    for s in scope_sw:
        if s in s0:
            reasons[s].append("ZERO_NAKED_EVALS")
        if s in new_sw:
            reasons[s].append("ADDED_SINCE_20261022")
        if s in ext_sw:
            reasons[s].append("FAMILY_EXT_20261028")
    # ---- families
    sw_family, fam_src = {}, {}
    for fam, v in (fm.get("families") or {}).items():
        for s in v.get("switches") or []:
            sw_family.setdefault(s, fam)
            fam_src.setdefault(s, "switch_family_map_20261022")
    for s, v in ext.items():
        if s not in sw_family:
            sw_family[s], fam_src[s] = v["family"], "switch_family_map_ext_20261028"
    dead = set()
    for s in cur_sw:
        e = bible.get(s)
        reach = int(((e or {}).get("vec_read_count") or {}).get("reachable", 0)) if e else 0
        if s in B.DEAD_IN_VEC or reach == 0:
            dead.add(s)
    fmf = fm.get("filters") or {}
    filters = {}
    for f in sorted(set(fmf) | {h.split("=", 1)[0] for h in hbold}):
        v = fmf.get(f) or {}
        filters[f] = {"class": v.get("class", "UNKNOWN"), "stages": list(v.get("stages") or []), "master": v.get("master")}
    # ---- live investigation per scope switch (+ family resolution from live)
    swinfo = {}
    rows_in_scope = [(tab, r) for tab, r in cur_rows if r["switch"] in scope_sw or f"{r['switch']}={vstr(r['cand'])}" in r0]
    for _, r in rows_in_scope:
        s = r["switch"]
        if s in swinfo:
            continue
        if s not in reasons:
            reasons[s].append("NEVER_TRIED_ROW_ONLY")
        e = bible.get(s) or {}
        lr = e.get("live_reads") or {}
        live = {"crypto": verify_sites(lr.get("crypto"), s), "stocks": verify_sites(lr.get("stocks"), s)}
        vec = verify_sites(e.get("vec_reads"), s)
        if not live["crypto"] and not live["stocks"]:
            for rel in LIVE_FILES:
                g = grep_name(rel, s, 2)
                if g:
                    live.setdefault("grep_fallback", []).extend(g)
        fns = [x.get("fn") for x in live["crypto"] + live["stocks"] + vec]
        st = Counter(x for x in (fn_stage(f) for f in fns) if x)
        live_stage = st.most_common(1)[0][0] if st else None
        kind = e.get("kind")
        fam = sw_family.get(s, "UNKNOWN")
        src = fam_src.get(s, "none")
        kind_stage = STAGE_OF_KIND.get(kind) or ("EXIT" if kind == "exit" else None)
        resolved = fam
        why = f"family {fam} ({src})"
        fam_grp = "EXIT" if fam.startswith("EXIT") else fam
        if fam == "UNKNOWN" and (live_stage or kind_stage):
            g = live_stage or kind_stage
            resolved = exit_subtype(s) if g == "EXIT" else g
            why = f"no family; live read stage {g} ({'fn' if live_stage else 'bible kind'})"
        elif fam == "UNKNOWN":
            u = s.upper()
            g = "REENTRY" if "REENTRY" in u else "AUGMENT" if "AUGMENT" in u else "REDUCE" if "REDUCE" in u else "EXIT" if any(t in u for t in ("STOP", "EXIT")) else "ENTRY" if any(t in u for t in ("ENTRY", "OPEN")) else None
            if g:
                resolved = exit_subtype(s) if g == "EXIT" else g
                why = f"no family, no classifiable live read; name tokens -> {g}"
        elif src.endswith("ext_20261028") or fam == "GLOBAL":
            g = live_stage or (kind_stage if kind_stage != "GLOBAL" else None)
            if g and g != fam_grp and g != "SIZING":
                resolved = exit_subtype(s) if g == "EXIT" else g
                why = f"{fam} ({src}) overridden by live read stage {g} ({'fn ' + ','.join(sorted({f for f in fns if fn_stage(f) == g}))[:120] if live_stage else 'bible kind ' + str(kind)})"
        sw_family[s] = resolved
        cands = [x["cand"] for t, x in cur_rows if x["switch"] == s]
        cfgc = config_comment(s)
        swinfo[s] = {"family": resolved, "family_basis": why, "bible_kind": kind, "status": e.get("status"), "live_reads": live, "vec_reads": vec, "vec_read_reachable": int((e.get("vec_read_count") or {}).get("reachable", 0)) if e else 0, "dead_in_vec": s in dead, "config": cfgc, "condition": condition_of(s, cands), "scope": sorted(set(reasons[s])), "tabs": sorted({t for t, x in cur_rows if x["switch"] == s})}
    ctx = B.Ctx(cs, sw_family, filters, defaults, dead, agk)
    # ---- evidence
    pos = load_positives(cs)
    marg_by_fam = defaultdict(set)  # (family, hdr) -> {switch}: pair lift vs its own naked row with CI90 lower bound > 0
    same_sw_need = defaultdict(set)  # switch -> {hdr}
    for p in pos.get("pairs") or []:
        ci = p.get("lift_ci")
        if isinstance(ci, list) and len(ci) == 2 and ci[0] is not None and float(ci[0]) > 0:
            sw = str(p["key"]).split("=", 1)[0]
            marg_by_fam[(sw_family.get(sw, "UNKNOWN"), p["filter"])].add(sw)
            same_sw_need[sw].add(p["filter"])
    needed = load_needed(cs)
    for k, v in needed.items():
        sw = k.split("=", 1)[0]
        for h in list(v.get("needed") or []):
            iv = (v.get("intervals") or {}).get(h)
            if iv and isinstance(iv, list) and len(iv) >= 5 and not str(iv[4]).startswith(("NEEDED", "RESCUE")):
                continue
            same_sw_need[sw].add(h)
    ln_cells = {}
    for x in load_linknet(cs):
        if x.get("filter"):
            ln_cells[(x["switch"], vstr(x["cand"]), f"{x['filter']}={x['opt']}")] = float(x.get("p_pos_pred") or 0.0)
    # ---- per row diff
    rows_out = []
    per_sw = defaultdict(Counter)
    per_tab = defaultdict(Counter)
    skipped = Counter()
    for tab, r in rows_in_scope:
        s, cand = r["switch"], r["cand"]
        rk = f"{tab}!{s}={vstr(cand)}"
        if r["grey"]:
            skipped["GREY_ROW (operator-greyed, never evaluated: no change)"] += 1
            per_sw[s]["grey_rows"] += 1
            continue
        if s in ARTEFACT_SWITCHES:
            skipped["ARTEFACT_ROW (MODE/VENUE class, never a strategy setting: no change)"] += 1
            per_sw[s]["artefact_rows"] += 1
            continue
        if s.startswith("ABLATION_"):
            skipped["ABLATION_ROW (never swept: no change)"] += 1
            per_sw[s]["ablation_rows"] += 1
            continue
        add, rem, keep_ok, keep_other, noted, cond_only = [], [], 0, 0, [], 0
        fam = sw_family.get(s, "UNKNOWN")
        cond = (swinfo.get(s) or {}).get("condition") or {}
        for h, col in tabs[tab]["headers"].items():
            fname, opt = h.split("=", 1)
            c = B.verdict_for(ctx, s, cand, h)
            if fname not in bible and fname not in fmf:
                c = {"verdict": "NEVER", "condition_keys": [], "reason": "NO_CONFIG_FIELD"}
            shaded = h in r["shaded"]
            is_default_opt = filt_default_hdr.get(fname) == h or vstr(defaults.get(fname)) == opt.strip()
            ev = []
            sib = marg_by_fam.get((fam, h), set()) - {s}
            if sib:
                ev.append(f"SIBLING_LIFT_CI90>0 on {len(sib)} {fam} switch(es) ({','.join(sorted(sib)[:3])})")
            if h in same_sw_need.get(s, set()):
                ev.append("SAME_SWITCH_EVIDENCE (lift CI90>0 or NEEDED/RESCUE on another cand of this switch)")
            lp = ln_cells.get((s, vstr(cand), h))
            v = c["verdict"]
            tf_match = bool(cond.get("tfs")) and opt.strip() in cond["tfs"]
            if v == "NEVER":
                if shaded:
                    rem.append({"hdr": h, "col": col, "verdict": v, "reason": c["reason"], "was": r["shaded"][h]})
                if ev:
                    noted.append({"hdr": h, "note": f"NEVER ({c['reason']}) but evidence: {'; '.join(ev)} — not shaded"})
                continue
            plausible = False
            basis = ""
            if is_default_opt:
                basis = "DEFAULT_OPT (filter at its bold default = the naked row)"
            elif v == "BINDS":
                plausible, basis = True, f"BINDS {c['reason']}"
            elif v == "CONDITIONAL" and ev:
                plausible, basis = True, f"CONDITIONAL {c['reason']} needs {'/'.join(c.get('condition_keys') or [])} + evidence"
            elif v == "CONDITIONAL":
                basis = f"CONDITIONAL {c['reason']} needs {'/'.join(c.get('condition_keys') or [])} (master off under the cat_side bold set: identity of the naked row unless the chain arms it; V20 settles it without an eval)"
                cond_only += 1
            elif v == "INDIRECT" and ev:
                plausible, basis = True, "INDIRECT + evidence"
            if plausible:
                if shaded:
                    keep_ok += 1
                else:
                    add.append({"hdr": h, "col": col, "verdict": v, "reason": basis, "evidence": ev, "linknet_p": lp, "tf_match": tf_match})
            else:
                if shaded:
                    keep_other += 1
                if lp is not None and lp >= 0.5 and v == "INDIRECT" and not is_default_opt:
                    noted.append({"hdr": h, "note": f"linknet p_pos_pred {lp:.2f} but INDIRECT (filter-alone echo) — not shaded"})
        if attach is not None:
            att = attach(cs, tab, tabs[tab]["headers"], s, cand, defaults)
            for a_ in add:
                a_["pilot_attaches"] = a_["hdr"] in att
            per_tab[tab]["add_pilot_attaches"] += sum(1 for a_ in add if a_["pilot_attaches"])
            per_sw[s]["add_pilot_attaches"] += sum(1 for a_ in add if a_["pilot_attaches"])
        rows_out.append({"tab": tab, "row": r["row"], "switch": s, "cand": vstr(cand), "row_key": rk, "family": fam, "default_row": r["bold"], "shaded_before": len(r["shaded"]), "add": add, "remove": rem, "kept_plausible": keep_ok, "kept_not_never_not_plausible": keep_other, "conditional_not_shaded": cond_only, "notes": noted})
        per_sw[s]["rows"] += 1
        per_sw[s]["add"] += len(add)
        per_sw[s]["remove"] += len(rem)
        per_tab[tab]["rows"] += 1
        per_tab[tab]["add"] += len(add)
        per_tab[tab]["remove"] += len(rem)
        per_tab[tab]["conditional_not_shaded"] += cond_only
    tot = Counter()
    for c in per_tab.values():
        tot.update(c)
    meta = {"cat_side": cs, "built_at": datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"), "template": str(tpl_path), "template_md5_12": md5_12(tpl_path), "scope": {"zero_naked_eval_switches": len(s0 & cur_sw), "added_since_20261022": len(new_sw), "family_ext_20261028": len(ext_sw), "never_tried_rows": len(r0), "scope_switches": len(swinfo), "scope_rows_white": len(rows_in_scope)}, "skipped_rows": dict(skipped), "totals": dict(tot), "paint": {"add_fill": LIGHT, "add_meaning": "FFF2CC light yellow = plausible, never calculated (painter convention)", "remove_fill": "row's dominant non-yellow filter-cell fill, else none"}, "rules": "tools/v18_binding_map.verdict_for (22_SWITCH_FAMILY_MAP §c) under template-bold defaults; ADD = BINDS | CONDITIONAL+measured evidence | INDIRECT+measured evidence, excluding the filter's bold-default option and self pairs; REMOVE = NEVER and shaded; INDIRECT without evidence and DEFAULT_OPT are left as they are", "evidence_sources": {"positives_all": "FULL stage pairs whose lift vs their own naked row has CI90 lower bound > 0, on a sibling switch of the same family (or another cand of the same switch)", "needed_filters": "FULL stage needed list of other cands of the same switch, interval state NEEDED*/RESCUE* only", "linknet": "FULL template_untried p_pos_pred (annotation only)"}}
    return {"meta": meta, "switches": swinfo, "rows": rows_out, "per_switch": {k: dict(v) for k, v in sorted(per_sw.items())}, "per_tab": {k: dict(v) for k, v in per_tab.items()}}


def write_md(plans, out):
    L = ["# New rows — yellow placement plan (2026-10-28)", "", "Operator order (verbatim): \"make sure the yellow cells for the new rows are in the right places of all the filters that might give a pos delta for that row by investigating what the switch does in live and what filters might help it\"", "", "Built by `tools/v15_new_rows_yellow_plan.py` (read-only). Per cat_side detail in `plan_<CS>.json`.", ""]
    L += ["## Totals per cat_side", "", "| cat_side | template md5 | scope switches | scope rows | rows diffed | yellows to ADD | yellows to REMOVE | grey/ablation rows skipped |", "|---|---|---|---|---|---|---|---|"]
    for cs, p in plans.items():
        m = p["meta"]
        L.append(f"| {cs} | {m['template_md5_12']} | {m['scope']['scope_switches']} | {m['scope']['scope_rows_white']} | {m['totals'].get('rows', 0)} | {m['totals'].get('add', 0)} | {m['totals'].get('remove', 0)} | {sum(m['skipped_rows'].values())} |")
    L += ["", "Rules: " + next(iter(plans.values()))["meta"]["rules"], ""]
    for cs, p in plans.items():
        L += [f"## {cs}", "", "### Per tab", "", "| tab | rows | add | remove |", "|---|---|---|---|"]
        for t in SWITCH_SHEETS:
            c = p["per_tab"].get(t)
            if c:
                L.append(f"| {t} | {c.get('rows', 0)} | {c.get('add', 0)} | {c.get('remove', 0)} |")
        L += ["", "### Per switch", "", "| switch | family (basis) | live read (first verified site) | vec reads | rows | add | remove | scope |", "|---|---|---|---|---|---|---|---|"]
        for s, c in p["per_switch"].items():
            i = p["switches"].get(s) or {}
            lr = (i.get("live_reads") or {})
            first = next((x for x in (lr.get("crypto") or []) + (lr.get("stocks") or []) + (lr.get("grep_fallback") or [])), None)
            site = f"`{first['site']}`" if first else "none"
            L.append(f"| {s} | {i.get('family')} ({(i.get('family_basis') or '')[:60]}) | {site} | {i.get('vec_read_reachable')} | {c.get('rows', 0)} | {c.get('add', 0)} | {c.get('remove', 0)} | {','.join(i.get('scope') or [])} |")
        L.append("")
    Path(out / "PLAN.md").write_text("\n".join(L) + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--templates", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--out", default=str(ROOT / "data" / "reports" / "new_rows_yellow_placement_20261028"))
    ap.add_argument("--cat-side", default="all", choices=CAT_SIDES + ("all",))
    ap.add_argument("--tab-level", default="", help="tab_level_filters.json the fleet pilot reads (S1 copy); enables the pilot_attaches check")
    ap.add_argument("--no-pilot-attach", action="store_true")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    fm = json.loads(FAMILY_MAP.read_text())
    ext = json.loads(FAMILY_EXT.read_text())["switches"]
    bible = json.loads(BIBLE.read_text())["switches"]
    groups = B.literal_assign(B.ABLATION_GROUPS, "GROUPS")
    agk = B.literal_assign(B.EXEMPTIONS, "AUGMENT_GATE_KEYS")
    nt = json.loads((FULL / "reports" / "never_tried_A_latest.json").read_text())
    attach = None if a.no_pilot_attach else pilot_attach_fn(a.tab_level)
    plans = {}
    for cs in (CAT_SIDES if a.cat_side == "all" else (a.cat_side,)):
        p = build_cs(cs, Path(a.templates) / f"TEMPLATE_{cs}.xlsx", fm, ext, bible, groups, agk, nt, attach)
        (out / f"plan_{cs}.json").write_text(json.dumps(p, indent=1, default=str))
        plans[cs] = p
        print(json.dumps({"cat_side": cs, **p["meta"]["scope"], **p["meta"]["totals"], "skipped": p["meta"]["skipped_rows"]}), flush=True)
    write_md(plans, out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
