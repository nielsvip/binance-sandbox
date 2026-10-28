import sys, json, collections, re, argparse
from openpyxl import load_workbook
SW = ["STDEV_SLOPE_SIZING","ENTRY_REVERSAL_BOUNCE","ENTRY_BREAKOUT_CHANNEL","ENTRY_CONFIRMATION_GATES","EXIT_STRUCTURAL","EXIT_VELOCITY","REENTRY_WINDOWED","REENTRY_ADAPTIVE","AUGMENT_TREND","AUGMENT_RISK_SIZING","REDUCE_PROFIT_LOCK","REDUCE_SIGNAL_RATER","GLOBAL_RISK_GATES"]
def scan(path):
    wb = load_workbook(path)
    problems = []
    switch_defaults = collections.defaultdict(list)
    for sh in SW:
        if sh not in wb.sheetnames:
            problems.append(f"missing sheet {sh}")
            continue
        ws = wb[sh]
        opts = collections.defaultdict(list)
        yes = collections.defaultdict(int)
        defaults_v = collections.defaultdict(list)
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name:
                continue
            name = str(name)
            if not re.fullmatch(r"[A-Z][A-Z0-9_]+|[A-Za-z][A-Za-z0-9_]+", name):
                problems.append(f"{sh} r{r}: odd switch name {name!r}")
            val = ws.cell(r, 2).value
            fam = ws.cell(r, 4).value
            key = (repr(val), str(fam or ""))
            opts[name].append((r, key))
            if str(ws.cell(r, 12).value or "").upper() == "YES":
                yes[name] += 1
                defaults_v[name].append(repr(val))
        for name, lst in opts.items():
            keys = [k for _, k in lst]
            dup = [k for k, c in collections.Counter(keys).items() if c > 1]
            if dup:
                problems.append(f"{sh} {name}: duplicate (value, family) rows {dup}")
            if len(set(defaults_v[name])) > 1:
                problems.append(f"{sh} {name}: conflicting default values {sorted(set(defaults_v[name]))}")
            if yes[name] > 1:
                problems.append(f"{sh} {name}: {yes[name]} default (YES) rows")
            if yes[name] == 0:
                problems.append(f"{sh} {name}: no default (YES) row")
            switch_defaults[name].append(sh)
    for name, tabs in switch_defaults.items():
        if len(set(tabs)) > 1:
            problems.append(f"switch {name} appears in multiple tabs: {sorted(set(tabs))}")
    return {"file": path, "switches": len(switch_defaults), "problems": problems}
if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("files", nargs="+"); a = ap.parse_args()
    for f in a.files:
        res = scan(f)
        print(f, "switches", res["switches"], "problems", len(res["problems"]))
        for p in res["problems"][:30]:
            print("   ", p)
