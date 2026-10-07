#!/usr/bin/env python3
"""v15_newx_scan — early-evidence scan of NEW / newly connected template rows (Agent NEWX, 2026-10-01, USER ORDER).

For every sampled sym_side of a cat_side, evaluate with the deployed vector engine (the pilot's own path: tools.opt.v12_pilot.prepare_batch /
evaluate_prepared_sanitized, 30D) on the sym_side's current defaults:
  tier A: each NEW row (switch=cand, template tabs, non-default option) ALONE vs baseline.
  tier B: each NEW row x each applicable filter option (yellow column header FILTER=opt; applicability from data/opportune_filter_map.json plus the
          cat_side's tab-level filters) vs baseline (pilot convention) and vs the row alone.
Output = progress-JSON schema the AVG2 aggregator already reads (tools/v15_avg2_history.py -> v15_vector_delta_rebuild.add_file):
  {"initial_baseline_gain":..., "done": {"TAB!<row>:SWITCH=cand": {delta, delta_invalid, reason, naked_binding, ref_fp, yellows{hdr:delta}, noop_yellows[], yellow_reasons{}}}}
written atomically to <progress-dir>/<SYM_SIDE>_v14_progress.json (default ~/v15_run_newx_20261001/progress, matched by the aggregator's v15_run*/progress glob)
and mirrored under data/newx/<CAT_SIDE>/<SYM_SIDE>_newx.json. NO-LIES: invalid / 0-trade results are recorded invalid (delta_invalid), never 0;
a delta is exactly 0 only if the engine produced identical gain.  Source tag 'newx', engine md5 and htf_align marker recorded.
Resources: nice(15), <=3 threads, memory guard (waits if MemAvailable < --min-mem-gb, exits politely if < --abort-mem-gb).
usage: v15_newx_scan.py --cat CRYPTO_LONG [--tier A|B] [--n-sym 25] [--threads 3] [--out-dir DIR] [--dry-run]
"""
import argparse, concurrent.futures as cf, dataclasses, datetime, hashlib, json, os, sys, threading, time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "tools"))
SWITCH_SHEETS = ["STDEV_SLOPE_SIZING", "ENTRY_REVERSAL_BOUNCE", "ENTRY_BREAKOUT_CHANNEL", "ENTRY_CONFIRMATION_GATES", "EXIT_STRUCTURAL", "EXIT_VELOCITY",
                 "REENTRY_WINDOWED", "REENTRY_ADAPTIVE", "AUGMENT_TREND", "AUGMENT_RISK_SIZING", "REDUCE_PROFIT_LOCK", "REDUCE_SIGNAL_RATER", "GLOBAL_RISK_GATES"]
CATS = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")


def mem_avail_gb():
    for l in open("/proc/meminfo"):
        if l.startswith("MemAvailable"):
            return int(l.split()[1]) / 1048576.0
    return 99.0


def mem_guard(min_gb, abort_gb):
    while True:
        m = mem_avail_gb()
        if m < abort_gb:
            print(f"[newx] MemAvailable {m:.1f} GB < {abort_gb} GB: exiting politely", flush=True)
            sys.exit(3)
        if m >= min_gb:
            return
        print(f"[newx] MemAvailable {m:.1f} GB < {min_gb} GB: waiting 60 s", flush=True)
        time.sleep(60)


def load_template(cat, tdir):
    import openpyxl
    wb = openpyxl.load_workbook(str(tdir / f"TEMPLATE_{cat}.xlsx"), read_only=True)
    rows, fhdr = [], {}
    for tab in SWITCH_SHEETS:
        if tab not in wb.sheetnames:
            continue
        it = ws_rows = list(wb[tab].iter_rows(min_row=2, values_only=True))
        if not ws_rows:
            continue
        hdr = [str(c).strip() if c is not None else "" for c in ws_rows[0]]
        try:
            pos = hdr.index("POS_SYM")
        except ValueError:
            pos = 13
        isd = hdr.index("is_default") if "is_default" in hdr else 11
        fh = [(i, h) for i, h in enumerate(hdr) if i > pos and "=" in h]
        fhdr[tab] = [h for _, h in fh]
        for k, r in enumerate(ws_rows[1:]):
            rn = k + 3
            a = r[0] if len(r) > 0 else None
            b = r[1] if len(r) > 1 else None
            if a in (None, "") or b in (None, ""):
                continue
            rows.append({"tab": tab, "row": rn, "switch": str(a).strip(), "cand": str(b).strip(),
                         "is_default": str(r[isd]).strip().upper() == "YES" if len(r) > isd and r[isd] is not None else False})
    wb.close()
    return rows, fhdr


def coerce(s, ref):
    s = str(s).strip()
    if isinstance(ref, bool):
        if s.lower() in ("true", "false"):
            return s.lower() == "true"
        raise ValueError(s)
    if isinstance(ref, int):
        return int(float(s))
    if isinstance(ref, float):
        return float(s)
    return s


def pick_sample(cat, n):
    d = json.load(open(sorted((ROOT / "data" / "daily_universe").glob("2026*.json"))[-1]))
    syms = sorted(d["crypto"] if cat.startswith("CRYPTO") else d["stocks"])
    syms = [s for s in syms if (ROOT / "backtest_v8" / "indicators" / f"{s}.npz").exists()]
    if len(syms) <= n:
        pick = syms
    else:
        step = len(syms) / float(n)
        pick = [syms[int(i * step)] for i in range(n)]
    side = "LONG" if cat.endswith("LONG") else "SHORT"
    return [f"{s}_{side}" for s in pick]


def atomic_write(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(f".tmp{os.getpid()}")
    with open(tmp, "w") as f:
        json.dump(obj, f)
        f.flush(); os.fsync(f.fileno())
    os.replace(tmp, path)


def fp(r):
    return f"{round(float(r.get('gain_pct') or 0.0), 6)}|{int(r.get('trades') or 0)}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cat", required=True, choices=CATS)
    ap.add_argument("--tier", default="A", choices=("A", "B"))
    ap.add_argument("--n-sym", type=int, default=25)
    ap.add_argument("--threads", type=int, default=3)
    ap.add_argument("--out-dir", default=os.path.expanduser("~/v15_run_newx_20261001/progress"))
    ap.add_argument("--template-dir", default=str(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"))
    ap.add_argument("--json", default=str(ROOT / "data" / "avg_delta_pos_sym.json"))
    ap.add_argument("--names", default=str(ROOT / "data" / "newx" / "connected_names.json"))
    ap.add_argument("--min-mem-gb", type=float, default=10.0)
    ap.add_argument("--abort-mem-gb", type=float, default=6.0)
    ap.add_argument("--only-new", action="store_true", help="tier B: only rows flagged new (not merely connected)")
    ap.add_argument("--max-sym", type=int, default=0)
    ap.add_argument("--shard", default="0/1", help="i/n: process only sym_sides whose index % n == i (work split across hosts, nothing evaluated twice)")
    ap.add_argument("--syms", default="", help="comma list of sym_sides (override sample)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    try:
        os.nice(15)
    except Exception:
        pass
    os.environ.setdefault("V12_NPZ_CACHE", "4")
    import v12_quick_engine as V
    from tools.opt import v12_pilot as P
    eng_md5 = hashlib.md5(open(ROOT / "v12_quick_engine.py", "rb").read()).hexdigest()[:8]
    cat = a.cat
    rows, fhdr = load_template(cat, Path(a.template_dir))
    pj = json.load(open(a.json)).get(cat, {})
    new_keys = {k for k, v in pj.items() if v.get("new")}
    names = set(json.load(open(a.names))) if os.path.exists(a.names) else set()
    sel = []
    for r in rows:
        if r["is_default"]:
            continue
        key = f"{r['tab']}!{r['switch']}={r['cand']}"
        isnew = key in new_keys
        if isnew or r["switch"] in names:
            r["new"] = isnew
            sel.append(r)
    opp = json.load(open(ROOT / "data" / "opportune_filter_map.json")).get(cat, {})
    try:
        tl = json.load(open(ROOT / "data" / "wiring" / "tab_filters" / "tab_level_filters.json")).get("cats", {}).get(cat, {})
    except Exception:
        tl = {}
    syms = [s for s in a.syms.split(",") if s] or pick_sample(cat, a.n_sym)
    _si, _sn = [int(x) for x in a.shard.split("/")]
    if _sn > 1:
        syms = [x for k, x in enumerate(syms) if k % _sn == _si]
    if a.max_sym:
        syms = syms[: a.max_sym]
    print(f"[newx] {cat} tier={a.tier} engine={eng_md5} rows_selected={len(sel)} (new={sum(1 for r in sel if r['new'])}) sym_sides={len(syms)} threads={a.threads}", flush=True)
    if a.dry_run:
        print("[newx] dry-run sample:", syms[:10], "...", "rows sample:", [(r['tab'], r['switch'], r['cand']) for r in sel[:5]])
        return
    qf = {f.name for f in dataclasses.fields(V.QuickConfig)}
    out_dir = Path(a.out_dir)
    mirror = ROOT / "data" / "newx" / cat
    lock = threading.Lock()
    for ss in syms:
        mem_guard(a.min_mem_gb, a.abort_mem_gb)
        t0 = time.time()
        prep = P.prepare_batch(ss, 30)
        if prep is None:
            print(f"[newx] {ss}: prepare failed (no valid NPZ) -> skipped", flush=True)
            continue
        base = P.evaluate_prepared_sanitized(prep, {}, 30)
        htf = base.get("htf_align") or ("see engine" if eng_md5 else "?")
        if not base.get("valid") or int(base.get("trades") or 0) <= 0:
            print(f"[newx] {ss}: baseline invalid/0 trades ({base.get('invalid_reason')}) -> skipped (no fake 0)", flush=True)
            continue
        bgain = float(base["gain_pct"])
        bfp = fp(base)
        cfgd = prep.get("base_cfg_dict", {})
        done, nskip = {}, 0
        jobs = []
        for r in sel:
            sw = r["switch"]
            ref = cfgd.get(sw, None)
            if sw not in qf and sw not in cfgd:
                done[f"{r['tab']}!{r['row']}:{sw}={r['cand']}"] = {"reason": "NOT_IN_CONFIG", "source": "newx"}
                nskip += 1
                continue
            try:
                val = coerce(r["cand"], ref)
            except Exception:
                done[f"{r['tab']}!{r['row']}:{sw}={r['cand']}"] = {"reason": "NO_CANDIDATE", "source": "newx"}
                nskip += 1
                continue
            jobs.append((r, sw, val))

        def run_row(job):
            r, sw, val = job
            key = f"{r['tab']}!{r['row']}:{sw}={r['cand']}"
            res = P.evaluate_prepared_sanitized(prep, {sw: val}, 30)
            ent = {"source": "newx", "engine_md5": eng_md5, "ref_fp": bfp, "is_new_row": bool(r["new"])}
            ok = bool(res.get("valid")) and int(res.get("trades") or 0) > 0
            if ok:
                ent["delta"] = round(float(res["gain_pct"]) - bgain, 6)
                ent["naked_binding"] = fp(res) != bfp
                ent["trades"] = int(res.get("trades") or 0)
                ent["dep_forced"] = res.get("dep_forced")
            else:
                ent["delta_invalid"] = True
                ent["reason"] = str(res.get("invalid_reason") or "invalid")[:120]
            if a.tier == "B" and ok and (r["new"] or not a.only_new):
                bases = set((opp.get(r["tab"]) or {}).get(sw, [])) | set((tl.get(r["tab"]) or {}).keys())
                yel, noop, yr, yvs = {}, [], {}, {}
                for hdr in fhdr.get(r["tab"], []):
                    base_name, _, opt = hdr.partition("=")
                    if base_name not in bases:
                        continue
                    ref2 = cfgd.get(base_name, None)
                    if base_name not in qf and base_name not in cfgd:
                        continue
                    try:
                        v2 = coerce(opt, ref2)
                    except Exception:
                        continue
                    r2 = P.evaluate_prepared_sanitized(prep, {sw: val, base_name: v2}, 30)
                    if r2.get("valid") and int(r2.get("trades") or 0) > 0:
                        d = float(r2["gain_pct"]) - bgain
                        if fp(r2) == fp(res):
                            noop.append(hdr)
                        else:
                            yel[hdr] = round(d, 6)
                            yvs[hdr] = round(float(r2["gain_pct"]) - float(res["gain_pct"]), 6)
                    else:
                        yr[hdr] = str(r2.get("invalid_reason") or "invalid")[:80]
                ent["yellows"], ent["noop_yellows"], ent["yellow_reasons"], ent["yellow_vs_switch"] = yel, noop, yr, yvs
            return key, ent

        with cf.ThreadPoolExecutor(max_workers=max(1, a.threads)) as ex:
            for key, ent in ex.map(run_row, jobs):
                done[key] = ent
        prog = {"initial_baseline_gain": bgain, "baseline_trades": int(base.get("trades") or 0), "bh": prep.get("bh"), "engine_md5": eng_md5, "htf_align": htf,
                "source": "newx", "tier": a.tier, "cat_side": cat, "window_days": 30, "at": datetime.datetime.utcnow().isoformat() + "Z", "done": done}
        atomic_write(out_dir / f"{ss}_v14_progress.json", prog)
        atomic_write(mirror / f"{ss}_newx.json", prog)
        npos = sum(1 for e in done.values() if isinstance(e.get("delta"), float) and e["delta"] > 1e-9)
        print(f"[newx] {ss}: rows={len(jobs)} skipped={nskip} positives={npos} baseline={bgain:.3f} {time.time()-t0:.0f}s", flush=True)


if __name__ == "__main__":
    main()
