#!/usr/bin/env python3
"""Read-only real-engine comparison of published KEEP settings vs legal E7 stop templates.

The input inventory selects every crypto sym_side with at least N losing fixed-stop
exit fills. Each selected side is evaluated on one exact frozen NPZ slice, using the
published active full_config as KEEP and one template-backed stop variation at a time.
This is a diagnostic experiment, not a promotion or deployment tool.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
STOP_FIELDS = (
    "FIXED_STOP_REF_FIELD",
    "FIXED_STOP_REF_TF",
    "FIXED_STOP_BUFFER_PCT",
    "FIXED_STOP_TIGHTEN_MODE",
    "FIXED_STOP_TIGHTEN_TRIGGER_PCT",
)
CRYPTO_TEMPLATE = {
    "LONG": ROOT / "SPREADSHEETS" / "TEMPLATE_CRYPTO_LONG.xlsx",
    "SHORT": ROOT / "SPREADSHEETS" / "TEMPLATE_CRYPTO_SHORT.xlsx",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_json(value) -> str:
    raw = json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode()
    return hashlib.sha256(raw).hexdigest()


def fixed_stop_loser_counts(inventory: dict, minimum: int) -> dict[str, int]:
    counts = {}
    for symbol, symbol_row in (inventory.get("per_symbol") or {}).items():
        for side, side_row in (symbol_row.get("position_sides") or {}).items():
            count = int(side_row.get("fixed_stop_exit_fill_loss_count") or 0)
            if side in ("LONG", "SHORT") and count >= minimum:
                counts[f"{symbol}_{side}"] = count
    return dict(sorted(counts.items()))


def template_values(path: Path) -> dict[str, list]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["EXIT_STRUCTURAL"]
        out = {field: [] for field in STOP_FIELDS}
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row or str(row[0] or "") not in out:
                continue
            field, value = str(row[0]), row[1]
            if value is not None and value not in out[field]:
                out[field].append(value)
        if any(not out[field] for field in STOP_FIELDS):
            raise ValueError(f"template lacks E7 values: {path}")
        return out
    finally:
        book.close()


def _normalized(spec: dict) -> tuple:
    from vec_decisions import fixed_stop as fs

    return fs.normalize_spec(*(spec.get(field) for field in STOP_FIELDS))


def make_variants(base: dict, choices: dict[str, list]) -> list[dict]:
    """Return normalized, template-backed single-field variants plus mode/trigger pairs."""
    from vec_decisions import fixed_stop as fs

    baseline = {field: base[field] for field in STOP_FIELDS}
    base_spec = _normalized(baseline)
    variants = {}

    def add(label_values):
        candidate = dict(baseline)
        candidate.update(label_values)
        norm = fs.normalize_spec(*(candidate[field] for field in STOP_FIELDS))
        if norm != tuple(candidate[field] for field in STOP_FIELDS):
            return
        if norm == base_spec:
            return
        key = json.dumps(candidate, sort_keys=True, separators=(",", ":"), default=str)
        labels = ";".join(f"{field}={candidate[field]}" for field in STOP_FIELDS if candidate[field] != baseline[field])
        variants.setdefault(key, {"label": labels, "overrides": {field: candidate[field] for field in STOP_FIELDS if candidate[field] != baseline[field]}})

    for field in STOP_FIELDS:
        for value in choices[field]:
            if value != baseline[field]:
                add({field: value})
    # Trigger has no effect while tightening is off. These legal two-field template
    # combinations make the trigger values measurable without changing other strategy logic.
    for mode in choices["FIXED_STOP_TIGHTEN_MODE"]:
        if mode == "off":
            continue
        for trigger in choices["FIXED_STOP_TIGHTEN_TRIGGER_PCT"]:
            add({"FIXED_STOP_TIGHTEN_MODE": mode, "FIXED_STOP_TIGHTEN_TRIGGER_PCT": trigger})
    return sorted(variants.values(), key=lambda row: row["label"])


def active_configs(db_path: Path, symsides: list[str]) -> dict:
    uri = f"file:{db_path.resolve()}?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    con.row_factory = sqlite3.Row
    try:
        result = {}
        for ss in symsides:
            row = con.execute(
                "select sym_side,cat_side,updated_at,winning_tag,overrides_json,full_config_json,template_md5 "
                "from per_sym_active where sym_side=?",
                (ss,),
            ).fetchone()
            if row is None:
                result[ss] = {"error": "NO_ACTIVE_PUBLISHED_SET"}
                continue
            try:
                full = json.loads(row["full_config_json"] or "{}")
            except Exception as exc:
                result[ss] = {"error": f"BAD_FULL_CONFIG:{type(exc).__name__}"}
                continue
            try:
                overrides = json.loads(row["overrides_json"] or "{}")
            except Exception:
                overrides = {}
            missing = [field for field in STOP_FIELDS if field not in full]
            if missing:
                result[ss] = {"error": f"PUBLISHED_FULL_CONFIG_MISSING_STOP_FIELDS:{','.join(missing)}"}
                continue
            result[ss] = {
                "sym_side": row["sym_side"], "cat_side": row["cat_side"], "updated_at": row["updated_at"],
                "winning_tag": row["winning_tag"], "template_md5": row["template_md5"],
                "published_overrides_sha256": sha256_json(overrides),
                "published_full_config_sha256": sha256_json(full), "full_config": full,
            }
        return result
    finally:
        con.close()


def summarize_run(result: dict, half_fee_pct: float, window_days: int) -> dict:
    ledger = result.get("execution_ledger") or []
    action_rows = [e for e in ledger if isinstance(e, dict) and "pnl_dollars" not in e]
    action_rows = [e for e in action_rows if str(e.get("type") or "").upper() in ("OPEN", "AUGMENT", "REDUCE", "CLOSE")]
    trade_rows = [e for e in ledger if isinstance(e, dict) and "pnl_dollars" in e and not e.get("seeded_prewindow")]
    turnover = sum(abs(float(e.get("qty") or 0) * float(e.get("price") or 0)) for e in action_rows)
    fee_est = turnover * half_fee_pct / 100.0
    net = sum(float(e.get("pnl_dollars") or 0) for e in trade_rows)
    opens = [e for e in action_rows if str(e.get("type") or "").upper() == "OPEN"]
    stop_rows = [e for e in trade_rows if "FIXED_STOP" in str(e.get("exit_reason") or e.get("reason") or "").upper()]
    rows = []
    for trade in trade_rows:
        entry_bar = int(trade.get("bar_entry", -1))
        exit_bar = int(trade.get("bar_exit", entry_bar))
        trade_turnover = sum(
            abs(float(e.get("qty") or 0) * float(e.get("price") or 0))
            for e in action_rows if entry_bar <= int(e.get("bar", -1)) <= exit_bar
        )
        trade_fee = trade_turnover * half_fee_pct / 100.0
        row = {key: trade.get(key) for key in ("bar_entry", "bar_exit", "entry_price", "exit_price", "qty", "pnl_pct", "entry_reason", "exit_reason", "bars_held")}
        row.update({"net_pnl_engine": float(trade.get("pnl_dollars") or 0), "simulated_fee_estimate": trade_fee, "gross_pnl_estimate": float(trade.get("pnl_dollars") or 0) + trade_fee, "turnover_estimate": trade_turnover})
        rows.append(row)
    data_days = float(result.get("data_days") or 0)
    trades = len(trade_rows)
    diagnostic = data_days <= 365.0 or trades < 30
    return {
        "valid": bool(result.get("valid")), "invalid_reason": result.get("invalid_reason"),
        "status": "DIAGNOSTIC_ONLY" if diagnostic else "DIAGNOSTIC_ONLY_NO_OOS", "window_days": window_days,
        "data_days": data_days, "trades": trades, "entry_count": len(opens),
        "reentry_count": max(0, len(opens) - 1), "closed_fixed_stop_count": len(stop_rows),
        "losing_fixed_stop_count": sum(float(e.get("pnl_dollars") or 0) < 0 for e in stop_rows),
        "net_pnl_engine": net, "simulated_fee_estimate": fee_est, "gross_pnl_estimate": net + fee_est,
        "turnover_estimate": turnover, "commission_model_round_trip_pct": half_fee_pct * 2.0,
        "max_dd_pct_engine": result.get("max_dd_pct"), "trade_rows": rows,
    }


def run_side(ss: str, loss_count: int, active: dict, prepared, choices: dict, template_path: Path, template_file_sha: str, window_days: int) -> dict:
    from tools.opt import evaluate_v12 as engine

    meta = active[ss]
    base = meta["full_config"]
    cfg = {field: base[field] for field in STOP_FIELDS}
    variants = make_variants(cfg, choices)
    runs = []
    base_result = engine.evaluate_prepared(prepared, dict(base), include_ledger=True)
    # evaluate_v12 hard-codes the crypto fee model to this value in prepare().
    base_config = prepared.get("base_cfg_dict") or {}
    round_trip_fee = float(base.get("CRYPTO_ROUND_TRIP_COMMISSION_PCT", base_config.get("CRYPTO_ROUND_TRIP_COMMISSION_PCT", 0.08)) or 0.0)
    keep = summarize_run(base_result, round_trip_fee / 2.0, window_days)
    runs.append({"variant": "KEEP", "stop_spec": cfg, "result": keep})
    for variant in variants:
        changed = dict(base)
        changed.update(variant["overrides"])
        res = engine.evaluate_prepared(prepared, changed, include_ledger=True)
        candidate = summarize_run(res, round_trip_fee / 2.0, window_days)
        candidate["delta_vs_keep"] = {
            "net_pnl_engine": candidate["net_pnl_engine"] - keep["net_pnl_engine"],
            "gross_pnl_estimate": candidate["gross_pnl_estimate"] - keep["gross_pnl_estimate"],
            "simulated_fee_estimate": candidate["simulated_fee_estimate"] - keep["simulated_fee_estimate"],
            "turnover_estimate": candidate["turnover_estimate"] - keep["turnover_estimate"],
            "entry_count": candidate["entry_count"] - keep["entry_count"],
            "reentry_count": candidate["reentry_count"] - keep["reentry_count"],
            "closed_fixed_stop_count": candidate["closed_fixed_stop_count"] - keep["closed_fixed_stop_count"],
            "losing_fixed_stop_count": candidate["losing_fixed_stop_count"] - keep["losing_fixed_stop_count"],
        }
        runs.append({"variant": variant["label"], "stop_spec": {**cfg, **variant["overrides"]}, "result": candidate})
    npz = prepared.get("npz_prepared") or {}
    ts = npz.get("timestamps")
    prov = prepared.get("npz_provenance") or {}
    sym = str(prepared.get("sym") or ss.rsplit("_", 1)[0])
    npz_path = engine._npz_path_for(sym)
    return {
        "sym_side": ss, "inventory_losing_fixed_stop_fills": loss_count,
        "evaluation_api_sha256": sha256_file(Path(engine.__file__)),
        "quick_engine_sha256": sha256_file(ROOT / "v12_quick_engine.py"),
        "fixed_stop_core_sha256": sha256_file(ROOT / "vec_decisions" / "fixed_stop.py"),
        "active_set": {key: meta[key] for key in ("cat_side", "updated_at", "winning_tag", "template_md5", "published_overrides_sha256", "published_full_config_sha256")},
        "template_path": str(template_path), "template_file_sha256": template_file_sha,
        "npz_file": str(npz_path) if npz_path else None, "npz_file_sha256": sha256_file(npz_path) if npz_path else None,
        "npz_provenance": prov, "npz_rows": int(len(ts)) if ts is not None else 0,
        "npz_first_ts": float(ts[0]) if ts is not None and len(ts) else None,
        "npz_last_ts": float(ts[-1]) if ts is not None and len(ts) else None,
        "real_engine": "tools.opt.evaluate_v12.evaluate_prepared -> v12_quick_engine.simulate_one",
        "evaluations": runs,
        "interpretation_limit": "Current active published full_config is replayed over this frozen NPZ window; this is not exchange-fill lifecycle attribution. Historic set-generation differences and actual fill fees/slippage are not reconstructed here.",
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--db", default="data/hourly_reconfig/per_sym_store.db")
    ap.add_argument("--out", required=True)
    ap.add_argument("--window-days", type=int, default=365)
    ap.add_argument("--min-losing-stop-fills", type=int, default=3)
    ap.add_argument("--symsides", default="", help="optional comma-list, intersected with inventory selection")
    args = ap.parse_args()
    if args.window_days not in (30, 365):
        ap.error("--window-days must be 30 or 365")
    inv_path, db_path, out_path = Path(args.inventory), Path(args.db), Path(args.out)
    inventory = json.loads(inv_path.read_text())
    if int(inventory.get("schema_version") or 0) != 1:
        raise ValueError("unsupported inventory schema_version")
    selected = fixed_stop_loser_counts(inventory, args.min_losing_stop_fills)
    if args.symsides:
        wanted = set(filter(None, args.symsides.split(",")))
        selected = {ss: n for ss, n in selected.items() if ss in wanted}
    if not selected:
        raise ValueError("no sym_sides meet the minimum confirmed losing stop-fill count")
    active = active_configs(db_path, list(selected))
    from tools.opt import evaluate_v12 as engine

    errors, sides = {}, {}
    started = dt.datetime.now(dt.timezone.utc).isoformat()
    db_hash_before = sha256_file(db_path)
    for index, (ss, loss_count) in enumerate(selected.items(), 1):
        print(f"[fixed-stop-lab] {index}/{len(selected)} {ss} losing_stop_fills={loss_count}", flush=True)
        if active[ss].get("error"):
            errors[ss] = active[ss]["error"]
            continue
        side = ss.rsplit("_", 1)[1]
        template = CRYPTO_TEMPLATE[side]
        choices = template_values(template)
        try:
            prepared = engine.prepare(ss, args.window_days)
            if not prepared:
                errors[ss] = "ENGINE_PREPARE_FAILED"
                continue
            sides[ss] = run_side(ss, loss_count, active, prepared, choices, template, sha256_file(template), args.window_days)
        except Exception as exc:
            errors[ss] = f"{type(exc).__name__}:{exc}"
        if index % 5 == 0 or index == len(selected):
            out_path.parent.mkdir(parents=True, exist_ok=True)
            partial = {"schema_version": 1, "status": "RUNNING", "started_at": started, "updated_at": dt.datetime.now(dt.timezone.utc).isoformat(), "inventory": str(inv_path), "inventory_sha256": sha256_file(inv_path), "db": str(db_path), "db_file_sha256_before": db_hash_before, "window_days": args.window_days, "minimum_losing_stop_fills": args.min_losing_stop_fills, "selected_sym_sides": selected, "completed_sym_sides": sorted(sides), "errors": errors, "sides": sides}
            out_path.write_text(json.dumps(partial, indent=2, default=str))
            print(f"[fixed-stop-lab] checkpoint sides={len(sides)} errors={len(errors)}", flush=True)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = {"schema_version": 1, "status": "COMPLETE", "started_at": started, "completed_at": dt.datetime.now(dt.timezone.utc).isoformat(), "inventory": str(inv_path), "inventory_sha256": sha256_file(inv_path), "db": str(db_path), "db_file_sha256_before": db_hash_before, "db_file_sha256_after": sha256_file(db_path), "window_days": args.window_days, "minimum_losing_stop_fills": args.min_losing_stop_fills, "selected_sym_sides": selected, "completed_sym_sides": sorted(sides), "errors": errors, "sides": sides}
    out_path.write_text(json.dumps(doc, indent=2, default=str))
    print(f"[fixed-stop-lab] COMPLETE sides={len(sides)}/{len(selected)} errors={len(errors)} out={out_path}", flush=True)


if __name__ == "__main__":
    main()
