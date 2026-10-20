"""--extend with a live venue default that is NOT among the existing rows: insert it bold, demote (never delete) the old bold."""
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
TOOL = str(ROOT / "tools" / "v15_switch_add.py")
SW, TAB = "START_POSITION_SIZE", "GLOBAL_RISK_GATES"


def run(*args):
    return subprocess.run([sys.executable, TOOL, *args], capture_output=True, text=True, cwd=str(ROOT))


def rows(path):
    wb = openpyxl.load_workbook(str(path))
    ws = wb[TAB]
    isd = next(c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "is_default")
    out = [(ws.cell(r, 2).value, bool(ws.cell(r, 2).font.b), ws.cell(r, isd).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == SW]
    n = ws.max_row
    wb.close()
    return out, n


class ExtendMissingDefaultTest(unittest.TestCase):
    def test_new_bold_default_old_demoted(self):
        with tempfile.TemporaryDirectory() as td:
            dst = Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx"
            shutil.copy(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_LONG.xlsx", dst)
            wb = openpyxl.load_workbook(str(dst))
            ws = wb[TAB]
            isd = next(c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "is_default")
            hit = [r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == SW]
            for r in sorted(hit[1:], reverse=True):
                ws.delete_rows(r, 1)
            ws.cell(hit[0], 2).value = 28
            ws.cell(hit[0], 2).font = Font(bold=True)
            ws.cell(hit[0], isd).value = "YES"
            wb.save(str(dst))
            wb.close()
            before, n0 = rows(dst)
            self.assertEqual(before, [(28, True, "YES")])
            args = ["--switch", SW, "--tab", TAB, "--candidates", "28,45", "--default", "45", "--venues", "crypto", "--sides", "long", "--extend", "--templates-dir", td]
            dry = run(*args)
            self.assertEqual(dry.returncode, 0, dry.stdout + dry.stderr)
            self.assertIn("old bold demoted", dry.stdout)
            self.assertEqual(rows(dst)[0], before)
            wet = run(*args, "--apply")
            self.assertEqual(wet.returncode, 0, wet.stdout + wet.stderr)
            after, n1 = rows(dst)
            self.assertEqual(n1, n0 + 1)
            self.assertEqual(sorted((float(v), b, y) for v, b, y in after), [(28.0, False, "NO"), (45.0, True, "YES")])
            again = run(*args)
            self.assertEqual(again.returncode, 0)
            self.assertIn("nothing to do", again.stdout)

    def test_other_refusals_kept(self):
        r = run("--switch", SW, "--tab", TAB, "--candidates", "28,45", "--default", "99", "--venues", "crypto", "--extend")
        self.assertNotEqual(r.returncode, 0)
        r = run("--switch", SW, "--tab", TAB, "--candidates", "28", "--default", "28", "--venues", "crypto", "--extend")
        self.assertNotEqual(r.returncode, 0)


if __name__ == "__main__":
    unittest.main()
