import openpyxl
from pathlib import Path

def test_v14_f_e_visible_data_only():
    # Use the latest PBF pilot if exists, else TEMPLATE
    candidates = list(Path("SPREADSHEETS").glob("PBF_LONG_30d_matrix*.xlsx"))
    if not candidates:
        candidates = [Path("SPREADSHEETS/TEMPLATE.xlsx")]
    p = sorted(candidates, key=lambda x: x.stat().st_mtime)[-1]
    wb = openpyxl.load_workbook(str(p), data_only=False)
    ws = wb["ENTRY_REVERSAL_BOUNCE"] if "ENTRY_REVERSAL_BOUNCE" in wb.sheetnames else wb[wb.sheetnames[0]]
    # Check that at least one evaluated row has numeric F
    has_numeric = False
    for r in range(3, 10):
        f = ws.cell(r, 6).value
        if isinstance(f, (int, float)):
            has_numeric = True
            break
    wb.close()
    # Also check data_only view for same file - numeric should be cached
    wb2 = openpyxl.load_workbook(str(p), data_only=True)
    ws2 = wb2["ENTRY_REVERSAL_BOUNCE"] if "ENTRY_REVERSAL_BOUNCE" in wb2.sheetnames else wb2[wb2.sheetnames[0]]
    has_cached = False
    for r in range(3, 10):
        f = ws2.cell(r, 6).value
        if isinstance(f, (int, float)):
            has_cached = True
            break
    wb2.close()
    # At least one of the two should be true after fix; before fix both false
    assert has_numeric or has_cached, f"No numeric F found in {p} (checked rows 3-9)"
