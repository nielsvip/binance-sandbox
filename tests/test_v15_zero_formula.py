"""ZERO_FORMULA skip (USER 2026-10-06): evidence-condemned never-positive rows/cells
are skipped + booked for formula fix. Settled (never refilled) but NOT obsolete."""

import json
import os
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import v15_pilot as PILOT
from tools.v15_row_guards import is_policy_skip_reason, is_structural_skip_reason, row_needs_recalc


class CondemnTest(unittest.TestCase):
    def test_condemned_needs_zero_pos_and_min_n(self):
        self.assertTrue(PILOT._zero_condemned(0, 15, 15))
        self.assertTrue(PILOT._zero_condemned(0, 56, 15))
        self.assertFalse(PILOT._zero_condemned(1, 56, 15))
        self.assertFalse(PILOT._zero_condemned(0, 14, 15))

    def test_condemned_fails_open_on_missing_evidence(self):
        self.assertFalse(PILOT._zero_condemned(None, 56, 15))
        self.assertFalse(PILOT._zero_condemned(0, None, 15))
        self.assertFalse(PILOT._zero_condemned(None, None, 15))
        self.assertFalse(PILOT._zero_condemned("x", "y", 15))


class PartitionTest(unittest.TestCase):
    def test_splits_kept_and_skipped(self):
        ev = {"R@H1": {"pos_sym": 0, "n_sym": 20}, "R@H2": {"pos_sym": 2, "n_sym": 20}, "R@H3": {"pos_sym": 0, "n_sym": 3}}
        kept, skipped = PILOT._zero_partition_hdrs(["H1", "H2", "H3", "H4"], {"H1": "R@H1", "H2": "R@H2", "H3": "R@H3", "H4": "R@H4"}, ev, 10)
        self.assertEqual(skipped, ["H1"])
        self.assertEqual(kept, ["H2", "H3", "H4"])

    def test_empty_evidence_keeps_everything(self):
        kept, skipped = PILOT._zero_partition_hdrs(["H1"], {"H1": "R@H1"}, {}, 10)
        self.assertEqual((kept, skipped), (["H1"], []))


class EnabledTest(unittest.TestCase):
    def test_default_on_with_kill_switch(self):
        old = os.environ.get("V15_ZERO_FORMULA_SKIP")
        try:
            os.environ.pop("V15_ZERO_FORMULA_SKIP", None)
            self.assertTrue(PILOT._zero_enabled())
            os.environ["V15_ZERO_FORMULA_SKIP"] = "0"
            self.assertFalse(PILOT._zero_enabled())
            os.environ["V15_ZERO_FORMULA_SKIP"] = "1"
            self.assertTrue(PILOT._zero_enabled())
        finally:
            if old is not None:
                os.environ["V15_ZERO_FORMULA_SKIP"] = old
            else:
                os.environ.pop("V15_ZERO_FORMULA_SKIP", None)


class LoadEvidenceTest(unittest.TestCase):
    def test_missing_file_fails_open(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            os.environ["V15_ZERO_CELL_JSON"] = str(Path(td) / "nope.json")
            try:
                self.assertEqual(PILOT._zero_load_cell_evidence("CRYPTO_LONG"), {})
            finally:
                os.environ.pop("V15_ZERO_CELL_JSON", None)

    def test_present_file_loads_cells(self):
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "CRYPTO_LONG.json"
            p.write_text(json.dumps({"cells": {"T!S=c@H": {"pos_sym": 0, "n_sym": 12, "avg_delta": -1.5}}}))
            os.environ["V15_ZERO_CELL_JSON"] = str(p)
            try:
                d = PILOT._zero_load_cell_evidence("CRYPTO_LONG")
                self.assertEqual(d["T!S=c@H"]["n_sym"], 12)
            finally:
                os.environ.pop("V15_ZERO_CELL_JSON", None)


class BookEntryTest(unittest.TestCase):
    def test_entry_is_bad_formula_not_obsolete(self):
        e = PILOT._zero_book_entry("row", "TAB", "SW", "cand", "", 0, 20, -2.0)
        self.assertEqual(e["status"], "SKIPPED_BAD_FORMULA")
        self.assertIn("NOT OBSOLETE", e["note"])
        self.assertNotIn("obsolete", e["status"].lower())
        self.assertEqual(e["n_sym"], 20)


class GuardSettlementTest(unittest.TestCase):
    def test_zero_prefix_is_structural_not_policy(self):
        self.assertTrue(is_structural_skip_reason("ZERO_FORMULA_ROW: pos_sym=0 in 20 syms"))
        self.assertFalse(is_policy_skip_reason("ZERO_FORMULA_ROW: pos_sym=0 in 20 syms"))

    def test_zero_skip_record_stands_no_recalc(self):
        rec = {"delta": None, "reason": "ZERO_FORMULA_ROW: pos_sym=0 in 20 syms", "complete": True, "zero_skipped_filters": ["H1"], "policy": {"ps": 1, "tl": 0, "uw": "x", "pilot": "y"}}
        self.assertIsNone(row_needs_recalc(rec))

    def test_policy_skip_still_recalcs(self):
        rec = {"delta": None, "reason": "SKIPPED_SAMPLING(pos_sym=0)", "complete": True}
        self.assertTrue(str(row_needs_recalc(rec)).startswith("policy-skip"))


class BookkeeperTest(unittest.TestCase):
    def test_merge_dedupes_rows_and_headers(self):
        from tools.v15_template_bookkeeper import merge_ledgers
        a = {"symside": "A_LONG", "rows": {"T!S=c": {"switch": "S"}}, "cell_rows": {"T!S=c": ["H1"]}, "cells_by_header": {"H1": ["T!S=c"]}}
        b = {"symside": "B_LONG", "rows": {"T!S=c": {"switch": "S"}}, "cell_rows": {}, "cells_by_header": {"H1": ["T!S=c"], "H2": ["T!S2=c2"]}}
        m = merge_ledgers([a, b])
        self.assertEqual(sorted(m["rows"]["T!S=c"]["syms_seen"]), ["A_LONG", "B_LONG"])
        self.assertEqual(m["headers"], {"H1": 1, "H2": 1})

    def test_write_book_sheet_only_adds_book_sheet(self):
        try:
            import openpyxl
        except Exception:
            self.skipTest("openpyxl unavailable")
        from tools.v15_template_bookkeeper import SHEET, write_book_sheet
        wb = openpyxl.Workbook()
        wb.active.title = "ENTRY_X"
        n = write_book_sheet(wb, "CRYPTO_LONG", {"rows": {"T!S=c": {"tab": "T", "switch": "S", "cand": "c", "pos_sym": 0, "n_sym": 20, "avg_delta": -1.0}}, "headers": {"H1": 3}}, 2)
        self.assertIn(SHEET, wb.sheetnames)
        self.assertIn("ENTRY_X", wb.sheetnames)
        self.assertEqual(wb["ENTRY_X"]["A1"].value, None)
        self.assertGreaterEqual(n, 4)


if __name__ == "__main__":
    unittest.main()
