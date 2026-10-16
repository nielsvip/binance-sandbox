"""MAX_ORDER_VALUE structural guards — the half of the old fishnet pair that SURVIVES.

Replaces tests/test_maxaug_never_zero.py. USER 2026-10-15 deleted MAX_AUGMENTS_PER_POSITION completely (that side is
now covered by tests/test_maxaug_deleted.py) and kept MAX_ORDER_VALUE: "MAX_ORDER_VALUE right now is 6x
START_POSITION_SIZE but it is not yet obsolete". So the remaining invariants are: its template rows stay greyed (an
un-grey row re-enters the delta machinery) and the per-cat_side value keeps real headroom over START_POSITION_SIZE.
"""
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]
GREY = ("FFBFBFBF", "00BFBFBF")
CAT_SIDES = ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT")


def test_max_order_value_rows_stay_greyed_in_all_templates():
    import openpyxl
    n = 0
    for base in ("SPREADSHEETS/TEMPLATE_%s.xlsx", "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_%s.xlsx"):
        for cs in CAT_SIDES:
            wb = openpyxl.load_workbook(str(ROOT / (base % cs)), read_only=True)
            for tab in ("AUGMENT_RISK_SIZING", "GLOBAL_RISK_GATES"):
                if tab not in wb.sheetnames:
                    continue
                ws = wb[tab]
                for r in range(3, ws.max_row + 1):
                    if ws.cell(row=r, column=1).value != "MAX_ORDER_VALUE":
                        continue
                    f = ws.cell(row=r, column=1).font
                    rgb = str(getattr(f.color, "rgb", "") or "").upper() if f and f.color else ""
                    assert rgb in GREY, (base % cs, tab, r, rgb)
                    n += 1
            wb.close()
    assert n > 0, "no MAX_ORDER_VALUE rows found — the switch is not obsolete, its rows must exist"


def test_max_order_value_keeps_headroom_over_start_size():
    import config as C
    import config_tradier as CT
    import per_sym_settings as pss
    _MISS = object()
    for cs, cfg in (("CRYPTO_LONG", C.Config()), ("CRYPTO_SHORT", C.Config()),
                    ("STOCKS_LONG", CT.TradierConfig()), ("STOCKS_SHORT", CT.TradierConfig())):
        venue = "tradier" if cs.startswith("STOCKS") else "crypto"
        sym = "AAPL" if venue == "tradier" else "BTCUSDC"
        side = cs.rsplit("_", 1)[1]
        sps = pss.get_for("START_POSITION_SIZE", sym, side, _MISS, venue=venue)
        mov = pss.get_for("MAX_ORDER_VALUE", sym, side, _MISS, venue=venue)
        if sps is _MISS:
            sps = getattr(cfg, "START_POSITION_SIZE", 0)
        if mov is _MISS:
            mov = getattr(cfg, "MAX_ORDER_VALUE", 0)
        if float(sps or 0) > 0:
            assert float(mov or 0) >= 4 * float(sps), (cs, mov, sps)


def test_zero_max_order_value_never_reaches_live():
    """A per-sym MAX_ORDER_VALUE of 0 blocks every order for that sym_side, so the read-time strip drops it."""
    import sys
    sys.path.insert(0, str(ROOT))
    import per_sym_store as pss
    assert pss._mandate_sanitize({"MAX_ORDER_VALUE": 0}) == {}
    assert pss._mandate_sanitize({"MAX_ORDER_VALUE": -1}) == {}
    assert pss._mandate_sanitize({"MAX_ORDER_VALUE": 300.0}) == {"MAX_ORDER_VALUE": 300.0}
