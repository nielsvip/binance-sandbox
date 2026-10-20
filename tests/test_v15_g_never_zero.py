import pathlib

def test_g_never_zero_for_neg():
    for name in ["v15_pilot.py", "v15_pilot_0914.py"]:
        text = pathlib.Path(name).read_text()
        bad = [i+1 for i,l in enumerate(text.splitlines()) if 'column=7' in l and '.value = 0.0' in l and 'G never' not in l and 'F hustle' not in l]
        assert bad == [], f"{name} still has G=0.0 without never comment at lines {bad}"
        assert text.count("G never 0.0") >= (1 if name == "v15_pilot.py" else 2), f"{name} missing G never sentinel"
        if name == "v15_pilot.py":
            # the spec-fill writes a literal 0.0 into G in exactly one place: the definitional RUNNING_IDENTITY row (USER 2026-10-03)
            assert text.count('_num_cell(ws, rr, cols["G"], 0.0, False)') == 1
            i = text.index('_num_cell(ws, rr, cols["G"], 0.0, False)')
            assert "running-identity" in text[i - 300 : i], "the only literal G=0.0 must be the running-identity row"

def test_templates_sheet_order_worst_first():
    import openpyxl, pathlib
    for tpl in ["SPREADSHEETS/TEMPLATE_0914_STOCKS_LONG.xlsx","SPREADSHEETS/TEMPLATE_0914_STOCKS_SHORT.xlsx"]:
        if not pathlib.Path(tpl).exists():
            continue
        wb = openpyxl.load_workbook(tpl, data_only=False)
        ordering = []
        if "ORDERING_0914_REV2" in wb.sheetnames:
            ws = wb["ORDERING_0914_REV2"]
            for r in range(5,30):
                v = ws.cell(r,2).value
                if v and isinstance(v,str) and v in wb.sheetnames:
                    ordering.append(v)
        # first trading sheet after CATEGORY/ORDERING/LEGEND must be STDEV (worst -4.49)
        cats = [s for s in wb.sheetnames if s.startswith("CATEGORY")]
        before = [s for s in ["ORDERING_0914_REV2","LEGEND_FILTERS","INSTRUCTIONS","FILTERS_EXPLAINED","INSTRUCTIONS_V2"] if s in wb.sheetnames]
        first_trading = wb.sheetnames[len(cats+before)]
        assert first_trading == "STDEV_SLOPE_SIZING", f"{tpl} first trading {first_trading} != STDEV"
        wb.close()
