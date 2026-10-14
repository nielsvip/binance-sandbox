"""2026-10-13 THROUGHPUT (USER ">12 sym_sides per server per hour"): with V15_POOL_BATCH_N=8 the prefetch counted pool
ENTRIES (8 per task) against a 2*workers cap -> one task in flight -> one busy worker (S1 HOOD_SHORT: workers 87% idle,
main thread 88% waiting). Pins: task-counting cap + flush, whole-row presubmit, revived/sacred presubmit, header-map cache."""
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
os.environ.setdefault("V15_ALLOW_MAC", "1")

import openpyxl  # noqa: E402
import v15_pilot as P  # noqa: E402

SRC = (ROOT / "v15_pilot.py").read_text()


def test_prefetch_counts_tasks_and_flushes():
    assert "len({id(f) for f, _ in inflight.values() if not f.done()})" in SRC
    assert 'os.environ.get("V15_PREFETCH_TASKCAP", "0") == "1"' in SRC and 'os.environ.get("V15_PREFETCH_ROWS", "200")' in SRC


def test_presubmit_everywhere_with_kill_switch():
    assert 'os.environ.get("V15_ROW_PRESUBMIT", "0") == "1"' in SRC
    assert "_submit(_ov_ps)" in SRC and "_submit(_rv_ovs[_rv_hdr])" in SRC and "_sacred_presubmit(_sq)" in SRC


def test_hdr_col_map_cache_matches_scan_and_invalidates():
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(["Switch", "Option", "override", "x", "BASELINE", "HUSTLE_DELTA", "VECTOR_DELTA"], start=1):
        ws.cell(row=2, column=c).value = h
    for r in range(3, 50):
        ws.cell(row=r, column=1).value = f"S{r}"
    P._HDR_MAP_CACHE.clear()
    m1 = P._hdr_col_map(ws)
    assert m1 == P._hdr_col_map_scan(ws) and m1["VECTOR_DELTA"] == 7
    m1["VECTOR_DELTA"] = 99  # callers get a copy
    assert P._hdr_col_map(ws)["VECTOR_DELTA"] == 7
    ws.cell(row=2, column=8).value = "V17_NOTE"
    P._HDR_MAP_CACHE.clear()  # what every header write site does
    assert P._hdr_col_map(ws)["V17_NOTE"] == 8
    assert SRC.count("_HDR_MAP_CACHE.clear()  # header written") >= 7


def test_loser_skips_post_and_guard_intervals():
    assert '_tier_name in ("LOST_CAUSE", "LOSER")' in SRC
    assert 'environ.get("V15_SELF_MONITOR_S", "120")' in SRC and 'environ.get("V15_STUCK_GUARD_S", "120")' in SRC
