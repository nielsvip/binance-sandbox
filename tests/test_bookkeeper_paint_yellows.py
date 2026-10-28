import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import v15_template_bookkeeper as BK  # noqa: E402

HEAD = ["Switch", "default", "override", "Family", "BASELINE", "HUSTLE_DELTA", "VECTOR_DELTA", "LIVE_DELTA", "LIVE_SHARPE", "REAL_COMPLETE", "PER_ROW_FILTERS", "is_default", "AVG_DELTA", "POS_SYM", "DC_BREAK_FILTER_TF=1h", "DC_BREAK_FILTER_TF=4h", "EXIT_TOP_FADE_FILTER_TF=1h"]


def _wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ENTRY_CONFIRMATION_GATES"
    for c, h in enumerate(HEAD, 1):
        ws.cell(2, c).value = h
    rows = [("NEW_GATE_ENABLED", False, True, "FFBFBFBF"), ("NEW_GATE_ENABLED", True, False, None), ("NEW_GATE_ENABLED", False, False, None)]
    for i, (sw, v, bold, grey) in enumerate(rows, 3):
        ws.cell(i, 1).value = sw
        if grey:
            ws.cell(i, 1).font = Font(color=grey)
        ws.cell(i, 2).value = v
        ws.cell(i, 2).font = Font(bold=bold)
        ws.cell(i, 12).value = "YES" if bold else "NO"
    ws.cell(4, 17).fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    ws.cell(4, 15).fill = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
    return wb


def _plan():
    return {"meta": {"cat_side": "CRYPTO_LONG"}, "rows": [
        {"tab": "ENTRY_CONFIRMATION_GATES", "row": 5, "switch": "NEW_GATE_ENABLED", "cand": "False", "row_key": "ENTRY_CONFIRMATION_GATES!NEW_GATE_ENABLED=False", "add": [{"hdr": "DC_BREAK_FILTER_TF=4h"}], "remove": []},
        {"tab": "ENTRY_CONFIRMATION_GATES", "row": 4, "switch": "NEW_GATE_ENABLED", "cand": "True", "row_key": "ENTRY_CONFIRMATION_GATES!NEW_GATE_ENABLED=True", "add": [{"hdr": "DC_BREAK_FILTER_TF=4h"}, {"hdr": "NOPE=1"}], "remove": [{"hdr": "EXIT_TOP_FADE_FILTER_TF=1h"}]},
        {"tab": "ENTRY_CONFIRMATION_GATES", "row": 9, "switch": "GONE", "cand": "True", "row_key": "ENTRY_CONFIRMATION_GATES!GONE=True", "add": [{"hdr": "DC_BREAK_FILTER_TF=1h"}], "remove": []},
    ]}


def test_paint_adds_removes_and_skips_grey_duplicate():
    wb = _wb()
    rep = BK.paint_yellows(wb, _plan())
    ws = wb["ENTRY_CONFIRMATION_GATES"]
    assert BK._fill_rgb(ws.cell(5, 16)) == "FFFFF2CC"  # planned row 5 painted, grey duplicate row 3 untouched
    assert BK._fill_rgb(ws.cell(3, 16)) == ""
    assert BK._fill_rgb(ws.cell(4, 16)) == "FFFFF2CC"
    assert BK._fill_rgb(ws.cell(4, 17)) == "FFDDEBF7"  # removed yellow takes the row's non-yellow filter fill
    assert rep["added"] == 2 and rep["removed"] == 1 and rep["header_missing"] == 1
    assert rep["row_missing"] == ["ENTRY_CONFIRMATION_GATES!GONE=True"]


def test_paint_is_idempotent_and_keeps_values():
    wb = _wb()
    BK.paint_yellows(wb, _plan())
    vals = [[c.value for c in r] for r in wb.active.iter_rows()]
    rep = BK.paint_yellows(wb, _plan())
    assert rep["added"] == 0 and rep["removed"] == 0 and rep["already"] == 2 and rep["already_clear"] == 1
    assert vals == [[c.value for c in r] for r in wb.active.iter_rows()]


def _two_tab_wb():
    wb = _wb()
    src = wb.active
    aug = wb.create_sheet("AUGMENT_TREND")
    aug.cell(1, 1).value = "08 AUGMENT_TREND title"
    for c, h in enumerate(HEAD[:15] + ["Live Location"], 1):
        aug.cell(2, c).value = h
    aug.merge_cells(start_row=1, start_column=1, end_row=1, end_column=16)
    aug.cell(3, 1).value = "BOUNCE_AUGMENT_ENABLED"
    aug.cell(3, 2).value = False
    aug.cell(3, 16).value = "ez_manage.py:1"
    src.cell(2, 18).value = "HTF_DIRECTION_GATE_ENABLED=True"
    src.cell(2, 19).value = "HTF_DIRECTION_GATE_ENABLED=False"
    src.cell(2, 19).font = Font(bold=True)
    return wb


def test_add_filter_columns_copies_headers_and_keeps_everything_else():
    wb = _two_tab_wb()
    rep = BK.add_filter_columns(wb, ["AUGMENT_TREND"], ["HTF_DIRECTION_GATE_ENABLED"])
    ws = wb["AUGMENT_TREND"]
    assert rep["AUGMENT_TREND"]["added"] == ["HTF_DIRECTION_GATE_ENABLED=True", "HTF_DIRECTION_GATE_ENABLED=False"]
    assert ws.cell(2, 16).value == "HTF_DIRECTION_GATE_ENABLED=True" and ws.cell(2, 17).font.b
    assert ws.cell(2, 18).value == "Live Location" and ws.cell(3, 18).value == "ez_manage.py:1"
    assert [str(m) for m in ws.merged_cells.ranges] == ["A1:R1"]
    assert BK.add_filter_columns(wb, ["AUGMENT_TREND"], ["HTF_DIRECTION_GATE_ENABLED"])["AUGMENT_TREND"]["added"] == []
