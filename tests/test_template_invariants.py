import openpyxl
import pathlib
import re

BASE = pathlib.Path("/Users/niels/Documents/binance")
TEMPLATES = [
    BASE / "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_LONG.xlsx",
    BASE / "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_SHORT.xlsx",
    BASE / "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_LONG.xlsx",
    BASE / "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_SHORT.xlsx",
]
# Single font/size invariant
TARGET_FONT = "Calibri"
TARGET_SIZE = 11
ORANGE_RGBS = {"FFFFC000","FFED7D31","FFFFA500","FFFF6600","FFC000"}

def is_orange(cell):
    if not cell.fill or not cell.fill.start_color or not cell.fill.start_color.rgb:
        return False
    rgb = str(cell.fill.start_color.rgb).upper()
    return rgb in ORANGE_RGBS or "C000" in rgb or "ED7D31" in rgb

def test_one_font_one_size():
    for path in TEMPLATES:
        wb = openpyxl.load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            if ws.title.startswith("TEMPLATE_") or "BASELINE" in ws.title:
                continue
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    assert c.font.name == TARGET_FONT, f"{path.name} {ws.title} {c.coordinate} font {c.font.name} != {TARGET_FONT}"
                    assert c.font.size == TARGET_SIZE, f"{path.name} {ws.title} {c.coordinate} size {c.font.size} != {TARGET_SIZE}"
        wb.close()

def test_no_blank_row_below_orange():
    for path in TEMPLATES:
        wb = openpyxl.load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            if ws.title in ("LEGEND_FILTERS","INSTRUCTIONS","FILTERS_EXPLAINED","INSTRUCTIONS_V2","FILTER_DICTIONARY_V2","12SYM_PARITY","FORMULAS","ABLATION_LIVE_EXTRA"):
                continue
            if "BASELINE" in ws.title:
                continue
            orange_rows = []
            for r in range(1, ws.max_row+1):
                for c in range(1, 6):
                    if is_orange(ws.cell(row=r, column=c)):
                        orange_rows.append(r)
                        break
            if not orange_rows:
                continue
            max_orange = max(orange_rows)
            # No blank row (col A-C all blank and no fill) immediately below orange
            for r in range(max_orange+1, ws.max_row+1):
                vals = [ws.cell(row=r, column=c).value for c in range(1,4)]
                all_blank = all(v is None or str(v).strip()=="" for v in vals)
                has_fill = any(ws.cell(row=r, column=c).fill.start_color.rgb not in (None,"00000000","FFFFFFFF") for c in range(1,4))
                if all_blank and not has_fill:
                    # Check if this is truly below orange and not just end of sheet with data beyond
                    # If there is any non-blank row below this blank, it's a violation
                    has_data_below = any(ws.cell(row=rr, column=1).value not in (None,"") for rr in range(r+1, ws.max_row+1))
                    assert not has_data_below, f"{path.name} {ws.title} blank row {r} below orange {max_orange} but data exists below"
                    break
        wb.close()

def test_exactly_one_bold_default_per_switch():
    # Parse defaults per category
    import ast
    def parse_defaults(p):
        d={}
        for m in re.finditer(r'^\s*([A-Z0-9_]+)\s*[:=][^=]*=\s*([^\n#]+)', p.read_text(), re.MULTILINE):
            try:
                d[m.group(1)] = ast.literal_eval(m.group(2).split('#')[0].strip())
            except:
                pass
        return d
    crypto = parse_defaults(BASE / "config.py")
    stocks = parse_defaults(BASE / "config_tradier.py")
    cat_defaults = {
        "TEMPLATE_CRYPTO_LONG.xlsx": crypto,
        "TEMPLATE_CRYPTO_SHORT.xlsx": crypto,
        "TEMPLATE_STOCKS_LONG.xlsx": stocks,
        "TEMPLATE_STOCKS_SHORT.xlsx": stocks,
    }
    for path in TEMPLATES:
        defaults = cat_defaults[path.name]
        wb = openpyxl.load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            if ws.title in ("LEGEND_FILTERS","INSTRUCTIONS","FILTERS_EXPLAINED","INSTRUCTIONS_V2","FILTER_DICTIONARY_V2","12SYM_PARITY","FORMULAS","ABLATION_LIVE_EXTRA"):
                continue
            if "BASELINE" in ws.title:
                continue
            # Group by switch name col A
            from collections import defaultdict
            groups = defaultdict(list)
            for r in range(3, ws.max_row+1):
                a = ws.cell(row=r, column=1).value
                b = ws.cell(row=r, column=2).value
                if a and b is not None and str(a).strip():
                    groups[str(a).strip()].append(r)
            for sw, rows in groups.items():
                bolds = [r for r in rows if ws.cell(row=r, column=2).font.bold]
                assert len(bolds)==1, f"{path.name} {ws.title} switch {sw} has {len(bolds)} bold, expected 1 rows {rows} bolds {bolds}"
                # Check that bold matches config default if defined
                if sw in defaults:
                    exp = defaults[sw]
                    bold_val = ws.cell(row=bolds[0], column=2).value
                    # Normalize compare
                    def norm(v):
                        if isinstance(v,bool): return v
                        if isinstance(v,str) and v.lower() in ("true","false"): return v.lower()=="true"
                        try: return float(v)
                        except: return v
                    assert norm(bold_val) == norm(exp), f"{path.name} {ws.title} {sw} bold {bold_val!r} != default {exp!r}"
        wb.close()

def test_no_bold_outside_default_column():
    for path in TEMPLATES:
        wb = openpyxl.load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            if ws.title in ("LEGEND_FILTERS","INSTRUCTIONS","FILTERS_EXPLAINED","INSTRUCTIONS_V2","FILTER_DICTIONARY_V2","12SYM_PARITY","FORMULAS","ABLATION_LIVE_EXTRA"):
                continue
            if "BASELINE" in ws.title:
                continue
            for row in ws.iter_rows(min_row=3):
                for c in row:
                    if c.column != 2 and c.font.bold:
                        # Only col B (2) may be bold, and only for default
                        # Header rows 1-2 may have bold, ignore
                        if c.row >2:
                            assert False, f"{path.name} {ws.title} {c.coordinate} unexpected bold outside col B (default)"
        wb.close()
