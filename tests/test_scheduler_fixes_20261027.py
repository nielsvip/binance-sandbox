import importlib.util
import json
import pathlib
import sys
import time

TOOLS = pathlib.Path(__file__).resolve().parents[1] / "tools"
sys.path.insert(0, str(TOOLS))
import fleet_dispatch as F
import v15_stranded_board_requeue as R


def load(name):
    spec = importlib.util.spec_from_file_location(name, TOOLS / (name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


NOW = 2_000_000_000


def host_with(launch_ss, start_refused=None):
    return {"running": {}, "progress": {}, "launches": {ss: {"ts": NOW - 1000} for ss in launch_ss},
            "start_refused": start_refused or {}}


def test_start_refusal_is_not_a_death():
    d = host_with(["A_LONG", "B_LONG", "C_LONG"], {"A_LONG": NOW - 900, "B_LONG": NOW - 900})
    assert [ss for _, ss in F.unfinished_deaths(d, NOW)] == ["C_LONG"]


def test_board_that_ran_and_vanished_still_counts():
    d = host_with(["A_LONG", "B_LONG"])
    assert len(F.unfinished_deaths(d, NOW)) == 2


def test_refusal_older_than_launch_does_not_excuse_a_death():
    d = host_with(["A_LONG"], {"A_LONG": NOW - 5000})
    assert len(F.unfinished_deaths(d, NOW)) == 1


def test_snapshot_age_not_changed():
    assert "max_age_seconds=900" in (TOOLS.parent / "v15_pilot.py").read_text()


def test_holes_only_refusal_rules():
    wt = {"WT_X": "untested (0 rows, all running/skipped/missing)"}
    prog = {"not_compliant": "ESSENTIAL-ROWS: 1 protection filters without a non-zero number (WT_X) - refusing publish", "unprotected": wt}
    assert R.holes_only_refusal(prog, {"WT_X"})
    assert not R.holes_only_refusal(prog, {"OTHER"})
    assert not R.holes_only_refusal(prog, None)
    assert not R.holes_only_refusal(dict(prog, unprotected={"WT_X": "measured 0.0"}), {"WT_X"})
    assert not R.holes_only_refusal(dict(prog, not_compliant="invalid:trades 0 < 10"), {"WT_X"})
    assert not R.holes_only_refusal(dict(prog, unprotected={}), {"WT_X"})


def test_tradeable_gate():
    assert not R._tradeable("ARM_SHORT", {"AU_SHORT"})
    assert R._tradeable("AU_SHORT", {"AU_SHORT"})
    assert R._tradeable("ARM_SHORT", None)


def test_measurable_switches_skips_grey(tmp_path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A3"], ws["A4"] = "GREY_ONE", "LIVE_ONE"
    ws["A3"].font = Font(color="FFBFBFBF")
    path = tmp_path / "t.xlsx"
    wb.save(path)
    assert R._measurable_switches(path) == {"LIVE_ONE"}


def test_watchdog_uses_lane_workers_and_defers_to_dispatcher(monkeypatch, tmp_path):
    W = load("v15_cpu80_watchdog")
    monkeypatch.delenv("FLEET_WORKERS", raising=False)
    assert W.lane_workers() == 7
    monkeypatch.setenv("FLEET_WORKERS", "5")
    assert W.lane_workers() == 5
    assert '"1", str(dispatch_max)' not in (TOOLS / "v15_cpu80_watchdog.py").read_text()
    monkeypatch.setattr(W.pathlib.Path, "home", staticmethod(lambda: tmp_path))
    monkeypatch.setattr(W, "sh", lambda cmd, timeout=10: "ERR:x")
    assert not W.dispatcher_alive()
    (tmp_path / "fleet_dispatch").mkdir()
    (tmp_path / "fleet_dispatch" / "heartbeat.json").write_text(json.dumps({"ts": time.time()}))
    assert W.dispatcher_alive()
    monkeypatch.setattr(W, "sh", lambda cmd, timeout=10: json.dumps({"ts": time.time()}))
    (tmp_path / "fleet_dispatch" / "heartbeat.json").write_text(json.dumps({"ts": 1}))
    assert W.dispatcher_alive()
