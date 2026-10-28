"""TEMPLATE is the ONLY authority (USER 2026-10-27) — pins v15_pilot.template_bold_defaults().

The old rule kept the CONFIG value whenever the bold differed from it, which discarded the bold in the
only case where it decides anything. These tests pin the inverted behaviour and its three exceptions.
"""
import os
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import v15_pilot as P  # noqa: E402

GREY = "FFBFBFBF"


def _book(tmp_path, rows, tab="GLOBAL_RISK_GATES"):
    """rows = [(switch, value, is_bold, is_default, grey)] -> a minimal one-tab workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tab
    ws.cell(row=2, column=1).value = "Switch"
    ws.cell(row=2, column=2).value = "Option Value"
    ws.cell(row=2, column=12).value = "is_default"
    for i, (sw, val, bold, isd, grey) in enumerate(rows, start=3):
        a = ws.cell(row=i, column=1)
        a.value = sw
        if grey:
            a.font = Font(color=GREY)
        b = ws.cell(row=i, column=2)
        b.value = val
        if bold:
            b.font = Font(bold=True)
        ws.cell(row=i, column=12).value = "YES" if isd else "NO"
    p = tmp_path / "TEMPLATE_TEST.xlsx"
    wb.save(str(p))
    wb.close()
    return p


def _run(path, truth):
    ov, bad = P.template_bold_defaults(path, dict(truth), dict(truth), set())
    return ov, bad, list(P.BOLD_REINSTATED), list(P.UNTRUSTED_BOLD)


def test_bold_that_differs_from_config_wins(tmp_path):
    """The whole point: config says False, the sheet's bold says True -> True is the default."""
    p = _book(tmp_path, [("MY_GATE_ENABLED", True, True, True, False), ("MY_GATE_ENABLED", False, False, False, False)])
    ov, bad, reinstated, untrusted = _run(p, {"MY_GATE_ENABLED": False})
    assert not bad
    assert ov.get("MY_GATE_ENABLED") is True, "the template bold must outrank the config value"
    assert [r[1] for r in reinstated] == ["MY_GATE_ENABLED"]
    assert untrusted == []


def test_kill_switch_restores_the_old_config_first_behaviour(tmp_path, monkeypatch):
    p = _book(tmp_path, [("MY_GATE_ENABLED", True, True, True, False), ("MY_GATE_ENABLED", False, False, False, False)])
    monkeypatch.setattr(P, "BOLD_WINS", False)
    ov, bad, reinstated, untrusted = _run(p, {"MY_GATE_ENABLED": False})
    assert "MY_GATE_ENABLED" not in ov
    assert reinstated == []
    assert [u[1] for u in untrusted] == ["MY_GATE_ENABLED"]


def test_container_config_field_keeps_config(tmp_path):
    """A scalar cell cannot express a dict/list/set, so that is a switch that does not exist yet."""
    p = _book(tmp_path, [("MY_RATE_MAP", "0.5", True, True, False), ("MY_RATE_MAP", "0.9", False, False, False)])
    ov, bad, reinstated, untrusted = _run(p, {"MY_RATE_MAP": {"a": 1.0}})
    assert "MY_RATE_MAP" not in ov
    assert reinstated == []
    assert [u[1] for u in untrusted] == ["MY_RATE_MAP"]


def test_bold_mark_on_an_empty_cell_is_not_a_default(tmp_path):
    """Template defect, found in all 4 real workbooks on NEWBORN_LOSS_KILL_VEL_TF: it published None."""
    p = _book(tmp_path, [("MY_EMPTY_TF", None, True, True, False), ("MY_EMPTY_TF", "15m", False, False, False)])
    ov, bad, reinstated, untrusted = _run(p, {"MY_EMPTY_TF": "1h"})
    assert ov.get("MY_EMPTY_TF", "1h") == "1h", "an empty bold cell must never become the default"
    assert reinstated == []
    assert [u[1] for u in untrusted] == ["MY_EMPTY_TF"]


def test_grey_retired_group_is_not_an_authority(tmp_path):
    """Operator kept DAYTRADE_DC_TARGET_TF / TECHNICAL_DC_TARGET_TF grey: config legitimately keeps the value."""
    p = _book(tmp_path, [("MY_RETIRED_TF", "15m", True, True, True), ("MY_RETIRED_TF", "4h", False, False, True)])
    ov, bad, reinstated, untrusted = _run(p, {"MY_RETIRED_TF": "15m,1h,4h"})
    assert "MY_RETIRED_TF" not in ov
    assert reinstated == []
    assert untrusted == []


def test_bold_equal_to_config_is_not_reported_as_reinstated(tmp_path):
    p = _book(tmp_path, [("MY_AGREE_ENABLED", True, True, True, False), ("MY_AGREE_ENABLED", False, False, False, False)])
    ov, bad, reinstated, untrusted = _run(p, {"MY_AGREE_ENABLED": True})
    assert ov.get("MY_AGREE_ENABLED") is True
    assert reinstated == [] and untrusted == []


def test_one_default_rule_still_fails_closed(tmp_path):
    """Two is_default=YES rows, or a YES that is not the bold row, must still be a violation."""
    p = _book(tmp_path, [("MY_TWO_DEFAULTS", 1, True, True, False), ("MY_TWO_DEFAULTS", 2, True, True, False)])
    ov, bad, _, _ = _run(p, {"MY_TWO_DEFAULTS": 1})
    assert bad, "a group without exactly one bold is_default=YES must be reported"
    assert "MY_TWO_DEFAULTS" not in ov


def test_real_workbooks_have_no_one_default_violations():
    """The live gate: a single bad cell here blocks the whole fleet's defaults rebuild."""
    sys.path.insert(0, str(ROOT / "tools"))
    from v15_switch_add import FILES

    import config as C
    import config_tradier as CT
    import v12_quick_engine as V

    d = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"
    if not d.exists():
        pytest.skip("canonical templates not present on this host")
    for cs, fn in FILES.items():
        stocks = cs.startswith("STOCKS")
        qc = V.QuickConfig()
        if stocks:
            qc.apply_tradier_defaults()
        typed = {k: getattr(qc, k) for k in dir(qc) if k.isupper() and not k.startswith("_")}
        live = CT.TradierConfig() if stocks else C.Config()
        for k in dir(live):
            if k.isupper() and not k.startswith("_"):
                try:
                    typed[k] = getattr(live, k)
                except Exception:
                    pass
        _ov, bad = P.template_bold_defaults(d / fn, typed, typed, set())
        assert not bad, f"{cs} one-default violations would REFUSE the cat_side rebuild: {bad[:3]}"
