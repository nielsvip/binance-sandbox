
import openpyxl

def test_blank_above_orange_and_sorted():
    p="/Users/niels/Documents/binance/SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_LONG.xlsx"
    wb=openpyxl.load_workbook(p, data_only=False)
    ws=wb["ENTRY_REVERSAL_BOUNCE"]
    blanks=[]; oranges=[]
    for r in range(3, ws.max_row+1):
        if ws.cell(r,1).value is None: continue
        fill=ws.cell(r,1).fill.fgColor.rgb if ws.cell(r,1).fill and ws.cell(r,1).fill.fgColor and ws.cell(r,1).fill.fgColor.rgb!="00000000" else None
        delta=ws.cell(r,13).value
        if fill is None:
            if isinstance(delta,(int,float)): blanks.append(delta)
        elif fill=="FFFFE699":
            if isinstance(delta,(int,float)): oranges.append(delta)
    assert blanks==sorted(blanks), f"Blank not sorted {blanks[:5]}"
    assert oranges==sorted(oranges), f"Orange not sorted {oranges[:5]}"
    # Check blank above orange
    blank_rows=[r for r in range(3, ws.max_row+1) if ws.cell(r,1).value and (ws.cell(r,1).fill.fgColor.rgb if ws.cell(r,1).fill and ws.cell(r,1).fill.fgColor and ws.cell(r,1).fill.fgColor.rgb!="00000000" else None) is None]
    orange_rows=[r for r in range(3, ws.max_row+1) if ws.cell(r,1).value and (ws.cell(r,1).fill.fgColor.rgb if ws.cell(r,1).fill and ws.cell(r,1).fill.fgColor and ws.cell(r,1).fill.fgColor.rgb!="00000000" else None)=="FFFFE699"]
    assert max(blank_rows) < min(orange_rows), "Blank not above orange"
    # Check row intact with yellow
    for r in range(3, ws.max_row+1):
        for c in range(13, ws.max_column+1):
            if ws.cell(r,c).fill.fgColor.rgb=="FFFFFF00":
                assert ws.cell(r,1).value is not None, "Yellow without switch"
    print("test passed")

if __name__=="__main__":
    test_blank_above_orange_and_sorted()
