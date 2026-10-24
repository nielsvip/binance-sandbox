"""DC_HARD_STOP_TF switch: per sym_side D vs 4h, monitor if 4h kills results."""
import pathlib, glob
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]

def test_config_has_switch():
    for p in [ROOT / "config.py", ROOT / "config_tradier.py", ROOT / "v12_quick_engine.py"]:
        src = p.read_text()
        assert "DC_HARD_STOP_TF" in src, f"{p.name} missing DC_HARD_STOP_TF"
        assert 'DC_HARD_STOP_TF: str = "4h"' in src or "DC_HARD_STOP_TF" in src
    # QuickConfig visibility
    qc = (ROOT / "v12_quick_engine.py").read_text()
    assert qc.count("DC_HARD_STOP_TF") >= 3  # definition + 2 uses
    # allowed values
    assert "4h" in (ROOT / "config.py").read_text().split("DC_HARD_STOP_TF")[1][:50]
    assert "D" in (ROOT / "config.py").read_text().split("DC_HARD_STOP_TF")[1][:100]

def test_live_uses_tf():
    for p in [ROOT / "tradier_manage.py", ROOT / "ez_manage.py", ROOT / "ez_positions_quick.py"]:
        src = p.read_text()
        assert "DC_HARD_STOP_TF" in src, f"{p.name} missing TF switch"
        # must handle both D and 4h keys with fallback
        assert "dc_low_D" in src and "dc_low_4h" in src
        assert "dc_high_D" in src and "dc_high_4h" in src
        assert "ULTIMATE_DC_" in src
        # per sym logic: _hs_tf in D vs 4h
        assert "_hs_tf" in src or "_hs_tf_q" in src or "_hs_tf_ez" in src

def test_vector_uses_tf():
    src = (ROOT / "v12_quick_engine.py").read_text()
    assert "_hs_tf_v12" in src
    assert "dc_high_D" in src and "dc_low_D" in src
    assert "ULTIMATE_DC_" in src

def test_template_has_switch():
    # USER 2026-10-24: canonical TEMPLATE.xlsx is legacy/deleted — FINAL_NORM is what
    # v15_pilot reads (v15_pilot.py:761). Check all four FINAL_NORM workbooks.
    for fn in ("TEMPLATE_CRYPTO_LONG.xlsx", "TEMPLATE_CRYPTO_SHORT.xlsx", "TEMPLATE_STOCKS_LONG.xlsx", "TEMPLATE_STOCKS_SHORT.xlsx"):
        path = ROOT / "SPREADSHEETS/TEMPLATE_FINAL_NORM" / fn
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            ws_name = "EXIT_VELOCITY" if "EXIT_VELOCITY" in wb.sheetnames else "EXIT_STRUCTURAL"
            ws = wb[ws_name]
            rows = [(str(ws.cell(r, 1).value or ""), str(ws.cell(r, 2).value or "")) for r in range(1, ws.max_row + 1)]
            assert any("DC_HARD_STOP_TF" in a for a, _ in rows), f"{path} missing DC_HARD_STOP_TF"
            assert any(a == "DC_HARD_STOP_TF" and b == "4h" for a, b in rows), f"{path} missing 4h"
            assert any(a == "DC_HARD_STOP_TF" and b == "D" for a, b in rows), f"{path} missing D"
        finally:
            wb.close()

def test_monitor_counts_stops():
    # v12 quick engine should tag trades with correct TF in reason, enabling per-sym monitoring
    src = (ROOT / "v12_quick_engine.py").read_text()
    assert 'f"ULTIMATE_DC_{_hs_tf_v12}_HARD_STOP"' in src or "ULTIMATE_DC_" in src
