"""Focused coverage for active template rows and the crypto WT exit-chain twin."""

from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
from openpyxl import load_workbook

from vec_decisions import live_exit_chain


ROOT = Path(__file__).resolve().parents[1]
ACTIVE_TEMPLATES = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"
EXIT_SWITCHES = ("WT_CROSS_EXIT_ENABLED", "DC15M_FORCE_EXIT_ENABLED")
GREY_FONT_RGB = {"FFBFBFBF", "00BFBFBF"}


@pytest.mark.parametrize(
    "template_name",
    (
        "TEMPLATE_CRYPTO_LONG.xlsx",
        "TEMPLATE_CRYPTO_SHORT.xlsx",
        "TEMPLATE_STOCKS_LONG.xlsx",
        "TEMPLATE_STOCKS_SHORT.xlsx",
    ),
)
def test_active_template_exit_switch_rows_are_sweepable(template_name):
    workbook = load_workbook(ACTIVE_TEMPLATES / template_name, read_only=False, data_only=False)
    try:
        for switch in EXIT_SWITCHES:
            rows = []
            for sheet in workbook.worksheets:
                for row_index in range(3, sheet.max_row + 1):
                    if sheet.cell(row_index, 1).value == switch:
                        rows.append((sheet, row_index))
            assert len(rows) == 2, (template_name, switch, [(s.title, r) for s, r in rows])
            values = [sheet.cell(row_index, 2).value for sheet, row_index in rows]
            assert sorted(values) == [False, True], (template_name, switch, values)
            bold_defaults = [
                (sheet, row_index)
                for sheet, row_index in rows
                if bool(sheet.cell(row_index, 2).font.bold)
            ]
            assert len(bold_defaults) == 1, (template_name, switch, bold_defaults)
            assert bold_defaults[0][0].cell(bold_defaults[0][1], 2).value in (False, True)
            for sheet, row_index in rows:
                color = sheet.cell(row_index, 1).font.color
                rgb = str(color.rgb).upper() if color and color.type == "rgb" else ""
                assert rgb not in GREY_FONT_RGB, (template_name, switch, sheet.title, row_index, rgb)
    finally:
        workbook.close()


def _cfg(**overrides):
    values = {
        "WT_CROSS_EXIT_ENABLED": True,
        "WT_CROSS_EXIT_REQUIRE_15M_CONFIRM": True,
        "WT_CROSS_EXIT_MIN_AGE_MINUTES": 6.0,
        "WT_CROSS_EXIT_APPLIES_TO_LOSERS": True,
        "WT_CROSS_EXIT_APPLIES_TO_WINNERS": True,
        "WT_CROSS_EXIT_3M_VETO_MAX_AGE": 30.0,
        "LIVE_5m_trading_ENABLED": True,
        "AUGMENTED_DC_BREAK_ENABLED": False,
        "K1M_EXTREME_REVERSE_ENABLED": False,
        "PARABOLIC_EXIT_ENABLED": False,
        "STRUCTURAL_RANGE_SHIFT_EXIT": False,
        "KEY_LEVEL_CRASH_ENABLED": False,
        "HTF_EXIT_VETO_ENABLED": False,
        "BREAKEVEN_GAIN_EROSION_ENABLED": False,
        "BREAKEVEN_DC_LOW4_ENABLED": False,
        "STDEV_REJECT_EXIT_ENABLED": False,
        "PEAK_GIVEBACK_PROTECTION_ENABLED": False,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def _prepared(cfg, *, wt15_against=True, wt3_with=False):
    n = 4
    arrays = {
        "wt1_1h": np.array([10.0] * n),
        "wt2_1h": np.array([20.0] * n),
        "wt1_15m": np.array([10.0] * n if wt15_against else [20.0] * n),
        "wt2_15m": np.array([20.0] * n),
        "wt1_3m": np.array([30.0] * n if wt3_with else [10.0] * n),
        "wt2_3m": np.array([20.0] * n),
    }

    def safe(npz, key, size, default):
        value = np.asarray(npz.get(key, np.full(size, default)))
        return value if value.size == size else np.full(size, default)

    close = np.full(n, 100.0)
    return live_exit_chain.prepare(arrays, n, True, cfg, close, safe)


def _step(prepared, cfg, *, age_min=10.0):
    return live_exit_chain.step(
        prepared,
        cfg,
        2,
        100.0,
        {
            "gain": 0.1,
            "peak": 0.1,
            "age_min": age_min,
            "entry_price": 99.9,
            "position_qty": 1.0,
        },
    )


def test_wt_cross_exit_live_chain_switch_and_confirm_gates():
    on = _cfg()
    hit = _step(_prepared(on), on)
    assert hit is not None and hit[0].startswith("WT_CROSS_EXIT_1h_bear")

    off = _cfg(WT_CROSS_EXIT_ENABLED=False)
    assert _step(_prepared(off), off) is None

    no_confirmation = _cfg()
    assert _step(_prepared(no_confirmation, wt15_against=False), no_confirmation) is None


def test_wt_cross_exit_live_chain_grace_min_age_and_3m_veto():
    cfg = _cfg()
    prepared = _prepared(cfg)
    assert _step(prepared, cfg, age_min=2.0) is None  # live grace window
    assert _step(prepared, cfg, age_min=4.0) is None  # configured six-minute minimum

    vetoed = _prepared(cfg, wt3_with=True)
    assert _step(vetoed, cfg, age_min=10.0) is None

    allowed = _prepared(cfg, wt3_with=False)
    hit = _step(allowed, cfg, age_min=10.0)
    assert hit is not None and hit[0].startswith("WT_CROSS_EXIT_1h_bear")


def test_existing_wt_scalar_numpy_predicate_suite():
    from vec_decisions import test_check_exit_candidates_crypto__wt_cross_exit as parity

    parity.main()
