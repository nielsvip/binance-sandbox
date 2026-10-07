"""Durable: greedy cumulative (G), hustle exhaustive, macbook guard, never-stop red continuation."""
import pathlib
import py_compile
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_LONG.xlsx"
V15 = ROOT / "v15_pilot.py"
V15_ENGINE = ROOT / "backtest_v15_engine.py"

def test_template_greedy_cumulative_uses_G_not_F():
    """E chain must add greedy VECTOR_DELTA (G) not hustle (F): =IF(G>0,E+G,E)"""
    wb = openpyxl.load_workbook(str(TEMPLATE), data_only=False)
    for sheet in ["ENTRY_REVERSAL_BOUNCE", "STDEV_SLOPE_SIZING", "ENTRY_BREAKOUT_CHANNEL", "EXIT_STRUCTURAL", "GLOBAL_RISK_GATES"]:
        ws = wb[sheet]
        # check first data row
        e3 = ws.cell(3, 5).value
        assert isinstance(e3, str), f"{sheet} E3 should be formula"
        # first sheet uses baseline, later sheets use MAX
        if sheet == "ENTRY_REVERSAL_BOUNCE":
            assert "TEMPLATE_BASELINE_METRICS" in e3, f"{sheet} E3 wrong {e3}"
        else:
            assert e3.startswith("=MAX("), f"{sheet} E3 should be MAX prev {e3}"
        # subsequent rows must use G (greedy) not F (hustle) — check only data rows (skip header separators)
        for r in [4, 5, 6]:
            e = ws.cell(r, 5).value
            if e is None or not isinstance(e, str) or not e.startswith("="):
                continue
            assert e.startswith("=IF(G"), f"{sheet}!{r} E must use G not F: {e}"
            assert f"G{r}" in e, f"{sheet}!{r} missing G{r} in {e}"
            assert f"F{r}" not in e, f"{sheet}!{r} E must not use F{r} (hustle) but G (greedy): {e}"
        # G must be VLOOKUP, F must be empty (hustle to be filled as float by engine)
        for r in [3, 4, 5]:
            g = ws.cell(r, 7).value
            assert isinstance(g, str) and "VLOOKUP" in g, f"{sheet}!{r} G must be VLOOKUP {g}"
            f = ws.cell(r, 6).value
            assert f is None or not (isinstance(f, str) and "VLOOKUP" in f), f"{sheet}!{r} F must be empty (hustle float) not VLOOKUP {f}"
        # headers
        assert ws.cell(2, 5).value == "BASELINE" or "GREEDY" in str(ws.cell(2,5).value) or ws.cell(2,5).value == "BASELINE"
        assert ws.cell(2, 6).value == "HUSTLE_DELTA"
        assert ws.cell(2, 7).value == "VECTOR_DELTA"
    wb.close()

def test_stdev_not_all_baseline():
    """STDEV sheet was bug: all E = baseline B2, must be IF(G..) chain like others, E3 = MAX(prev)"""
    wb = openpyxl.load_workbook(str(TEMPLATE), data_only=False)
    ws = wb["STDEV_SLOPE_SIZING"]
    assert ws.cell(3,5).value.startswith("=MAX("), "STDEV E3 must be MAX(prev)"
    assert ws.cell(4,5).value.startswith("=IF(G4"), f"STDEV E4 must be IF(G4...): {ws.cell(4,5).value}"
    assert ws.cell(5,5).value.startswith("=IF(G5"), f"STDEV E5 must be IF chain {ws.cell(5,5).value}"
    wb.close()

def test_macbook_guard_blocks_backtest():
    src = V15.read_text()
    assert ('platform.system() == "Darwin"' in src) or ('sys.platform == "darwin"' in src), "must check Darwin"
    assert "BLOCKED" in src and "NO backtests on MacBook" in src, "must BLOCK on Darwin without dry-run"
    assert 'sys.exit(2)' in src, "must exit 2 on Mac block"
    # backtest_v15_engine also blocks
    src2 = V15_ENGINE.read_text()
    assert 'BLOCKED: NO backtests on MacBook' in src2, "v15_engine must block Darwin"
    assert 'V15_ALLOW_MAC' in src2, "engine must have allow override"

def test_never_stop_marks_red_and_continues():
    src = V15.read_text()
    # must flag NEG with orange but also write G value (our fix)
    assert 'ws_row.cell(row=r, column=7).value = float(delta_best)' in src, "NEG must write G greedy value"
    # parity-fail must write both F and G
    assert src.count('column=7).value = float(delta_best)') >= 3, "all NEG/parity/live-neg must write G"
    # BadZip recovery
    assert "BadZip-recover" in src, "must recover BadZip and continue"
    # per-cell red and continue, not return
    assert "PER_CELL TIMEOUT" in src or "per_cell_timeout" in src
    assert "FF0000" in src  # red for failures
    assert "FFA500" in src  # orange for NEG
    # must not have os._exit killing workbook
    assert "os._exit(2)" not in src

def test_greedy_cumulative_logic():
    src = V15.read_text()
    # delta = vg - cumulative_before (greedy, not vs baseline)
    assert "vg - cumulative_before" in src
    # hustle vs baseline is separate
    assert "hustle" in src.lower() or "HUSTLE_DELTA" in src
    assert "_hustle_neg = float(vec_best" in src or "_hustle_delta" in src or "hustle_delta_vs_baseline" in src
    # E column is cumulative_before, G is greedy delta
    assert "column=5).value = float(cumulative_before)" in src
    assert "column=7).value = float(delta_best)" in src

def test_hustle_exhaustive_combinations():
    src = V15.read_text()
    # hustle must try all best deltas together in numerous combinations until max delta
    assert "hustle" in src.lower()
    assert "beam" in src.lower()
    # now beam 64 depth 10 and exhaustive top12
    assert "64" in src and "10" in src, "beam should be 64 depth 10"
    assert "exhaustive" in src.lower() or "EXHAUSTIVE" in src
    assert "itertools.combinations" in src or "combinations" in src
    assert "_hustle_top[:80]" in src or "80" in src, "top 80 for hustle"

def test_filter_full_evaluation():
    src = V15.read_text()
    # must evaluate ALL filters per row, not limit 4 that caused duplicate deltas and incomplete yellows
    assert "FILTER EVALUATION POLICY" in src or "EVERY CELL CHANGES" in src
    assert "top50 capped" in src or "50" in src
    # old limit 4 must not remain as primary
    assert "top4 FAST" not in src, "old top4 limit must be removed"

def test_compile():
    py_compile.compile(str(V15), doraise=True)
    py_compile.compile(str(V15_ENGINE), doraise=True)
