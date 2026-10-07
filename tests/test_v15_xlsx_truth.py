"""xlsx source-of-truth (2026-10-03): atomic publish copies, manifests, audit/verify.

Regression: renamed-without-regenerated charts (header 0.12% vs filename 12.64),
torn-copy risk at publish, and silent cross-host stragglers after REDO/quarantine.
"""
import json
import sys

sys.path.insert(0, ".")
sys.path.insert(0, "tools")
import v15_pilot as P
from sheet_audit import audit_counts, verify_chart


def test_build_publish_manifest_shape():
    m = P._build_publish_manifest("AU_SHORT", "AU_SHORT_bh13p60_gain22p19_30d_matrix.xlsx", {"gain_pct": 22.19, "trades": 114, "tim_pct": 37.2, "max_dd_pct": 1.0, "valid": True, "invalid_reason": "", "bh": 13.6}, (2175, 52, 39, 2794), "o" * 32, "f" * 32, "AU.npz", "n" * 32, 80.3, "p" * 32, "TEMPLATE_STOCKS_SHORT.xlsx", "s5")
    assert m["symside"] == "AU_SHORT"
    assert m["counts"] == {"F": 2175, "C": 52, "E": 39, "done_n": 2794}
    assert m["metrics"]["gain_pct"] == 22.19
    assert m["file_md5"] == "f" * 32
    assert m["published_utc"] and m["host"] == "s5"


def test_md5_file_roundtrip(tmp_path):
    f = tmp_path / "a.bin"
    f.write_bytes(b"abc" * 1000)
    assert P._md5_file(f) == "8f33d0ecfe648a745b169baf8a77658b"
    assert P._md5_file(tmp_path / "missing") is None


def test_atomic_copy_happy_and_corrupt(tmp_path):
    import zipfile
    src = tmp_path / "good.xlsx"
    with zipfile.ZipFile(str(src), "w") as z:
        for i in range(12):
            z.writestr(f"entry{i}.xml", "<a/>")
    dst = tmp_path / "out.xlsx"
    P._atomic_copy(src, dst)
    assert dst.exists()
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a zip")
    try:
        P._atomic_copy(bad, tmp_path / "out2.xlsx")
        raise AssertionError("corrupt copy must raise")
    except Exception:
        pass
    assert not (tmp_path / "out2.xlsx").exists()


def test_verify_chart_fail_and_pass(tmp_path):
    lie = tmp_path / "X_SHORT_bhm15p84_gain12p64_30d_matrix.html"
    lie.write_text("<html>gain 0.12% &nbsp; BH -15.84% &nbsp; trades 1</html>")
    ok, reasons, _ = verify_chart(lie)
    assert ok is False and any("12.64" in r for r in reasons)
    good = tmp_path / "X_SHORT_bhm15p84_gain11p01_30d_matrix.html"
    good.write_text("<html>gain 11.01% &nbsp; BH -15.84% &nbsp; trades 141</html>")
    ok2, _, det = verify_chart(good)
    assert ok2 is True and det["header_trades"] == 141


def test_audit_counts_synthetic_xlsx(tmp_path):
    import openpyxl
    p = tmp_path / "mini.xlsx"
    wb = openpyxl.Workbook()
    for tab in ("ENTRY_REVERSAL_BOUNCE", "EXIT_VELOCITY"):
        ws = wb.create_sheet(tab)
        for r in range(3, 8):
            ws.cell(r, 3).value = "K=V"
            ws.cell(r, 5).value = 1.5
            ws.cell(r, 6).value = 0.5
            ws.cell(r, 12).value = 0.1
            ws.cell(r, 13).value = 0.2
    wb.save(str(p))
    t = audit_counts(p)["total"]
    assert t == {"C": 10, "E": 10, "F": 10, "Y": 20}, t
