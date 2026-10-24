"""v15_switch_add — the template one-script rule's single exception (USER 2026-10-06).

Proves refusal paths + a real add on throwaway copies (never live templates).
"""

import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TOOL = str(ROOT / "tools" / "v15_switch_add.py")


def run(*args):
    return subprocess.run([sys.executable, TOOL, *args], capture_output=True, text=True, cwd=str(ROOT))


class RefusalTest(unittest.TestCase):
    def test_unknown_tab(self):
        r = run("--switch", "X", "--tab", "NOPE", "--candidates", "a,b", "--default", "a")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("unknown tab", r.stdout + r.stderr)

    def test_single_candidate(self):
        r = run("--switch", "X", "--tab", "GLOBAL_RISK_GATES", "--candidates", "a", "--default", "a")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("at least 2", r.stdout + r.stderr)

    def test_unwired_switch(self):
        r = run("--switch", "DEFINITELY_NOT_A_SWITCH_XYZ", "--tab", "GLOBAL_RISK_GATES", "--candidates", "True,False", "--default", "True")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("not a field of config.py", r.stdout + r.stderr)


class AddOnCopiesTest(unittest.TestCase):
    def test_wired_switch_adds_rows_then_dup_refuses(self):
        with tempfile.TemporaryDirectory() as td:
            shutil.copy(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_LONG.xlsx", Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx")
            import openpyxl as _ox  # 2026-10-13: the live templates now carry CHANNEL_REENTRY_STOP_TF (row queue) -> strip it from the copy first
            _wb = _ox.load_workbook(str(Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx"))
            _ws = _wb["REENTRY_ADAPTIVE"]
            for _r in range(_ws.max_row, 2, -1):
                if _ws.cell(_r, 1).value == "CHANNEL_REENTRY_STOP_TF":
                    _ws.delete_rows(_r, 1)
            _wb.save(str(Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx"))
            _wb.close()
            base = ["--switch", "CHANNEL_REENTRY_STOP_TF", "--tab", "REENTRY_ADAPTIVE", "--candidates", "1h,OFF,15m", "--default", "1h", "--venues", "crypto", "--sides", "long", "--templates-dir", td]
            dry = run(*base)
            self.assertEqual(dry.returncode, 0, dry.stdout + dry.stderr)
            self.assertIn("insert 3 rows", dry.stdout)
            wet = run(*base, "--apply")
            self.assertEqual(wet.returncode, 0, wet.stdout + wet.stderr)
            self.assertIn("installed 3 rows", wet.stdout)
            import openpyxl
            wb = openpyxl.load_workbook(str(Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx"), read_only=True)
            ws = wb["REENTRY_ADAPTIVE"]
            hit = [(ws.cell(r, 2).value, ws.cell(r, 2).font.bold) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == "CHANNEL_REENTRY_STOP_TF"]
            wb.close()
            self.assertEqual(hit, [("1h", True), ("OFF", False), ("15m", False)])
            dup = run(*base)
            self.assertNotEqual(dup.returncode, 0)
            self.assertIn("already has rows", dup.stdout + dup.stderr)


REMOVE_TOOL = str(ROOT / "tools" / "v15_switch_remove.py")
GREY_TOOL = str(ROOT / "tools" / "v15_switch_grey.py")


def run_tool(tool, *args):
    return subprocess.run([sys.executable, tool, *args], capture_output=True, text=True, cwd=str(ROOT))


class NoRemovalLawTest(unittest.TestCase):
    """USER 2026-10-24: removing template rows is impossible — grey instead."""

    def test_remover_refuses_everything(self):
        for extra in ([], ["--apply"]):
            r = run_tool(REMOVE_TOOL, "--switch", "ANYTHING", *extra)
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("removal is impossible", r.stdout + r.stderr)

    def test_greyer_greys_and_ungreys_without_touching_values(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            shutil.copy(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_LONG.xlsx", Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx")
            base = ["--switch", "DC_DAYTRADE_TARGET_PCT", "--reason", "no-removal law unit test", "--venues", "CRYPTO", "--sides", "LONG", "--templates-dir", td]
            g = run_tool(GREY_TOOL, *base, "--apply")
            self.assertEqual(g.returncode, 0, g.stdout + g.stderr)
            self.assertIn("greyed 4 rows", g.stdout)
            wb = openpyxl.load_workbook(str(Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx"))
            ws = wb["EXIT_STRUCTURAL"]
            rows = [(r, ws.cell(r, 2).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == "DC_DAYTRADE_TARGET_PCT"]
            self.assertEqual(len(rows), 4, "greying never adds or removes rows")
            for r, _ in rows:
                f = ws.cell(r, 1).font
                self.assertEqual(str(getattr(f.color, "rgb", "") or "").upper(), "FFBFBFBF")
            vals = sorted(ws.cell(r, 2).value for r, _ in rows)
            wb.close()
            self.assertEqual(vals, [0.005, 0.01, 0.015, 0.02], "values untouched by greying")
            u = run_tool(GREY_TOOL, *base, "--ungrey", "--apply")
            self.assertEqual(u.returncode, 0, u.stdout + u.stderr)
            wb = openpyxl.load_workbook(str(Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx"))
            ws = wb["EXIT_STRUCTURAL"]
            n_grey = sum(1 for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == "DC_DAYTRADE_TARGET_PCT" and (lambda f: f is not None and f.color is not None and str(getattr(f.color, "rgb", "") or "").upper() == "FFBFBFBF")(ws.cell(r, 1).font))
            wb.close()
            self.assertEqual(n_grey, 0, "un-grey restores default font")

    def test_adder_refuses_grey_group(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            shutil.copy(ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_LONG.xlsx", Path(td) / "TEMPLATE_CRYPTO_LONG.xlsx")
            run_tool(GREY_TOOL, "--switch", "DC_DAYTRADE_TARGET_PCT", "--reason", "no-removal law unit test", "--venues", "CRYPTO", "--sides", "LONG", "--templates-dir", td, "--apply")
            r = run("--switch", "DC_DAYTRADE_TARGET_PCT", "--tab", "EXIT_STRUCTURAL", "--candidates", "0.005,0.01,0.015,0.02", "--default", "0.01",
                    "--venues", "crypto", "--sides", "long", "--templates-dir", td, "--extend")
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("grey", r.stdout + r.stderr)


if __name__ == "__main__":
    unittest.main()
