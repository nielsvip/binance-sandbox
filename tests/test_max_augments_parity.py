"""Parity test: MAX_AUGMENTS_PER_POSITION must be wired in backtest and live.

Live ez_manage caps AUGMENT via augmented_count >= max (config 999999 = no cap,
REENTRY exempt). Backtest_v12 was stubbed to `_ = 1` (pilot caps 0/0.5/1/2
never blocked). This test ensures the gate is present and templates retain
the variants so future pilots have valid deltas (connect, don't delete).
"""
import pathlib

def test_backtest_v12_wired():
    txt = pathlib.Path("backtest_v12_engine.py").read_text()
    assert "BLOCKED_MAX_AUGMENTS" in txt
    assert "MAX_AUGMENTS_PER_POSITION" in txt
    assert "augmented_count" in txt
    assert "'REENTRY'" in txt  # exempt, mirrors ez_manage

def test_ez_manage_wired():
    txt = pathlib.Path("ez_manage.py").read_text()
    assert "MAX_AUGMENTS_PER_POSITION" in txt
    assert "BLOCKED_MAX_AUGMENTS" in txt

def test_templates_retain_variants():
    import openpyxl
    for name in ["SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_LONG.xlsx", "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_LONG.xlsx",
                 "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_SHORT.xlsx", "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_SHORT.xlsx"]:
        wb = openpyxl.load_workbook(name, read_only=True, data_only=True)
        found = any(row and row[0] == "MAX_AUGMENTS_PER_POSITION"
                    for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
        # crypto_long/stocks_long must have it; shorts may not (0 is not a cap variant for shorts) — presence is ok
        if "LONG" in name:
            assert found, f"{name} missing MAX_AUGMENTS_PER_POSITION — variants were incorrectly removed"

def test_integer_grid():
    """0.5 augments is impossible — grid must be integers 0,1,2,5,10,999999"""
    import openpyxl
    for name in ["SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_LONG.xlsx", "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_LONG.xlsx", "SPREADSHEETS/TEMPLATE.xlsx"]:
        wb = openpyxl.load_workbook(name, read_only=True, data_only=True)
        vals = [row[1] for ws in wb.worksheets for row in ws.iter_rows(values_only=True) if row and row[0]=="MAX_AUGMENTS_PER_POSITION"]
        assert "0.5" not in vals, f"{name} still has fractional 0.5"
        assert "5" in vals and "10" in vals, f"{name} missing 5/10 integer caps"
        assert "999999" in vals

def test_tradier_has_cap():
    txt = pathlib.Path("config_tradier.py").read_text()
    assert "MAX_AUGMENTS_PER_POSITION" in txt
