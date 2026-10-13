"""Persistent template row intake (BIBLE §71 amendment / §94, 2026-10-13).

Proves on throwaway copies (never live templates): queue install + idempotency + refusals, enqueue/merge-inbox,
rescue-lost (the 2026-10-09 13:49Z ENTRY_LH_TRIGGER_* loss class), and that the pilot's candidate scan sees new rows.
"""

import json
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import v15_switch_add as SA  # noqa: E402
import v15_template_row_intake as RI  # noqa: E402

SRC = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_LONG.xlsx"
LH = "ENTRY_LH_TRIGGER_ENABLED"


def _rows(path, sw):
    wb = openpyxl.load_workbook(str(path))
    try:
        out = []
        for t in SA.SWITCH_SHEETS:
            ws = wb[t]
            for r in range(3, ws.max_row + 1):
                if ws.cell(r, 1).value == sw:
                    out.append((t, str(ws.cell(r, 2).value), bool(ws.cell(r, 2).font.bold), ws.cell(r, 12).value))
        return out
    finally:
        wb.close()


def _strip(path, sw):
    """test helper: produce a template copy WITHOUT `sw` (whole-row delete) so the install path can be exercised."""
    wb = openpyxl.load_workbook(str(path))
    for t in SA.SWITCH_SHEETS:
        ws = wb[t]
        for r in range(ws.max_row, 2, -1):
            if ws.cell(r, 1).value == sw:
                ws.delete_rows(r, 1)
    wb.save(str(path))
    wb.close()


@pytest.fixture()
def tdir(tmp_path):
    d = tmp_path / "TEMPLATE_FINAL_NORM"
    d.mkdir()
    shutil.copy(SRC, d / SRC.name)
    _strip(d / SRC.name, LH)
    return d


def _queue(tmp_path, entries):
    q = tmp_path / "queue.json"
    q.write_text(json.dumps({"version": 1, "entries": entries}))
    return q


def test_queue_installs_then_idempotent(tmp_path, tdir, monkeypatch):
    monkeypatch.setattr(SA, "register_allowlist", lambda *a, **k: False)
    (tmp_path / "backups").mkdir()
    ent = {"switch": LH, "tab": "ENTRY_REVERSAL_BOUNCE", "venues": ["CRYPTO"], "sides": ["LONG"], "candidates": ["False", "True"], "default": "False"}
    q = _queue(tmp_path, [ent])
    real_root = SA.ROOT
    s = SA.apply_queue(q, [str(tdir)], apply=False, root=real_root)
    assert s["entries"][LH]["planned"] and not s["entries"][LH]["installed"]
    assert _rows(tdir / SRC.name, LH) == []
    monkeypatch.setattr(SA.shutil, "copy2", lambda *a, **k: None)  # no backups of throwaway copies into the repo
    s = SA.apply_queue(q, [str(tdir)], apply=True, root=real_root)
    assert s["entries"][LH]["installed"], s
    rows = _rows(tdir / SRC.name, LH)
    assert [(t, v, b, y) for t, v, b, y in rows] == [("ENTRY_REVERSAL_BOUNCE", "False", True, "YES"), ("ENTRY_REVERSAL_BOUNCE", "True", False, "NO")]
    qd = json.loads(q.read_text())
    assert qd["entries"][0]["status"] == "INSTALLED"
    s2 = SA.apply_queue(q, [str(tdir)], apply=True, root=real_root)
    assert not s2["entries"][LH]["planned"] and s2["entries"][LH]["present"]
    assert len(_rows(tdir / SRC.name, LH)) == 2


def test_queue_refuses_bad_entries(tmp_path, tdir):
    q = _queue(tmp_path, [
        {"switch": "R2_TF_LIST", "tab": "EXIT_STRUCTURAL", "venues": ["CRYPTO"], "candidates": ["15m", "1h"], "default": "15m"},
        {"switch": LH, "tab": "ENTRY_REVERSAL_BOUNCE", "venues": ["CRYPTO"], "candidates": ["False", "True"], "default": "True"},
        {"switch": "DEFINITELY_NOT_A_SWITCH_XYZ", "tab": "GLOBAL_RISK_GATES", "venues": ["CRYPTO"], "candidates": ["True", "False"], "default": "True"},
    ])
    s = SA.apply_queue(q, [str(tdir)], apply=False)
    assert "not a scalar" in s["entries"]["R2_TF_LIST"]["refused"]["CRYPTO"]
    assert "config default" in s["entries"][LH]["refused"]["CRYPTO"]
    assert "not a field" in s["entries"]["DEFINITELY_NOT_A_SWITCH_XYZ"]["refused"]["CRYPTO"]


def test_enqueue_and_merge_inbox(tmp_path):
    inbox = tmp_path / "inbox"
    SA.enqueue({"switch": LH, "tab": "ENTRY_REVERSAL_BOUNCE", "candidates": ["False", "True"], "default": "False"}, inbox)
    SA.enqueue({"switch": LH, "tab": "ENTRY_REVERSAL_BOUNCE", "candidates": ["False", "True"], "default": "False"}, inbox)
    q = tmp_path / "q.json"
    q.write_text(json.dumps({"entries": [{"switch": "OTHER_SW", "tab": "GLOBAL_RISK_GATES"}]}))
    rc = RI.main(["merge-inbox", "--queue", str(q), "--inbox", str(inbox)])
    assert rc == 0
    d = json.loads(q.read_text())
    assert [e["switch"] for e in d["entries"]] == ["OTHER_SW", LH]
    assert not list(inbox.glob("*.json")) and len(list((inbox / "merged").glob("*.json"))) >= 1


def test_rescue_lost_detects_dropped_switch(tmp_path, tdir):
    before = tmp_path / "before"
    after = tmp_path / "after"
    (before / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM").mkdir(parents=True)
    (after / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM").mkdir(parents=True)
    shutil.copy(SRC, before / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / SRC.name)  # live: carries the LH rows
    shutil.copy(tdir / SRC.name, after / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / SRC.name)  # stripped = the pull that drops them
    if not _rows(before / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / SRC.name, LH):
        pytest.skip("live template has no LH rows yet")
    ents = RI.lost_rows(before, after, flat_before=False)
    e = {x["switch"]: x for x in ents}[LH]
    assert e["tab"] == "ENTRY_REVERSAL_BOUNCE" and e["default"]["CRYPTO"] == "False" and set(e["candidates"]["CRYPTO"]) == {"False", "True"}
    inbox = tmp_path / "inbox"
    rc = RI.main(["rescue-lost", "--before-dir", str(before), "--unflat", "--after-root", str(after), "--inbox", str(inbox)])
    assert rc == 0 and len(list(inbox.glob("rescue_lost_*.json"))) == 1


def test_live_queue_entries_validate():
    """every QUEUED entry of the committed queue passes the adder's validation for its venues (no silent REFUSED rows)."""
    q = json.loads((ROOT / "data" / "template_row_queue.json").read_text())
    bible = json.loads((ROOT / "data" / "SWITCH_BIBLE.json").read_text())["switches"]
    bad = {}
    for e in q["entries"]:
        for v in SA._entry_venues(e):
            try:
                SA.validate_switch(e["switch"], e["tab"], SA._entry_cands(e, v), SA._entry_default(e, v), [v], ["LONG", "SHORT"], bible=bible)
            except SA.Refused as r:
                bad[f"{e['switch']}@{v}"] = str(r)
    assert not bad, bad
    assert {x["switch"] for x in q["excluded"]} >= {"MAX_ORDER_VALUE_FIN", "MAX_ORDER_VALUE_MEN", "MIN_GAIN"}


def test_live_templates_carry_lh_trigger_rows():
    """the INV-0001 rows exist in all 8 live templates (floor, never exact counts — §70)."""
    for rel in RI.TEMPLATE_RELS:
        rows = RI.rows_by_switch(ROOT / rel)
        for sw in ("ENTRY_LH_TRIGGER_ENABLED", "ENTRY_LH_TRIGGER_REGRESS_PCT", "ENTRY_LH_TRIGGER_DIVERG_ENABLED"):
            assert sw in rows and rows[sw]["tab"] == "ENTRY_REVERSAL_BOUNCE" and len(rows[sw]["rows"]) >= 2, (rel, sw)
            assert sum(1 for v, y, b in rows[sw]["rows"] if y and b) == 1, (rel, sw)


def test_cli_queue_dry_run_on_copy(tmp_path, tdir):
    q = _queue(tmp_path, [{"switch": LH, "tab": "ENTRY_REVERSAL_BOUNCE", "venues": ["CRYPTO"], "sides": ["LONG"], "candidates": ["False", "True"], "default": "False"}])
    r = subprocess.run([sys.executable, str(ROOT / "tools" / "v15_switch_add.py"), "--queue", str(q), "--templates-dirs", str(tdir)], capture_output=True, text=True, cwd=str(ROOT))
    assert r.returncode == 0, r.stdout + r.stderr
    assert "1 switch installs planned" in r.stdout
