"""KINDERGARTEN_CUMULATIVE_MODE must be wired for CRYPTO too — live AND vectorized.

USER 2026-10-28: "if a swtich exists in a TEMPLATE_ there is no discussion possible" and "make sure the
live scripts AND vectorized have EVERY SINGLE switch tested and adjusted to the optimal setting".
The switch had a row in all 4 workbooks (True/False, bold True) and a field in config.py, config_tradier.py
and QuickConfig, but NOTHING in the crypto path read it: only the stocks twins did. vec_decisions/kg_entry_gate
is the ONE predicate crypto live (ez_manage fresh-OPEN chokepoint) and the crypto vector (v12_quick_engine
final entry_sig) share, so wiring it there gives both venues the switch simultaneously.
"""
import sys
from pathlib import Path

import numpy as np
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from vec_decisions.kg_entry_gate import cumulative_mode, ema921_pass, ema921_pass_vec  # noqa: E402

TFS = "15m,1h,4h"


def _get(cum, min_tfs=1, strict=""):
    d = {
        "EMA_9_21_FILTER_ENABLED": True,
        "EMA_9_21_FILTER_TFS": TFS,
        "EMA_9_21_TIMEFRAME": "1h",
        "KINDERGARTEN_CUMULATIVE_MODE": cum,
        "KINDERGARTEN_CUMULATIVE_MIN_TFS": min_tfs,
        "KINDERGARTEN_STRICT_TFS": strict,
    }
    return lambda k, default=None: d.get(k, default)


class _Cfg:
    EMA_9_21_FILTER_ENABLED = True
    EMA_9_21_FILTER_TFS = TFS
    EMA_9_21_TIMEFRAME = "1h"
    KINDERGARTEN_CUMULATIVE_MIN_TFS = 1
    KINDERGARTEN_STRICT_TFS = ""

    def __init__(self, cum):
        self.KINDERGARTEN_CUMULATIVE_MODE = cum


def _npz():
    # bar 0: all three agree | bar 1: 2 of 3 | bar 2: 1 of 3 | bar 3: 2 of 3
    return {
        "ema_9_above_21_15m": np.array([1.0, 1.0, 0.0, 1.0]),
        "ema_9_above_21_1h": np.array([1.0, 0.0, 0.0, 1.0]),
        "ema_9_above_21_4h": np.array([1.0, 1.0, 1.0, 0.0]),
    }


def _safe(z, k, n, default):
    return z[k]


def test_default_is_the_template_bold_true():
    assert cumulative_mode(lambda k, d=None: d) is True


def test_scalar_live_quorum_vs_and_differ():
    ind = {"ema_9_above_21_15m": 1.0, "ema_9_above_21_1h": 1.0, "ema_9_above_21_4h": 0.0}
    ok_cum, passing, total = ema921_pass(_get(True), ind, True)
    assert (ok_cum, passing, total) == (True, 2, 3), "quorum of 1 must pass on 2 of 3"
    ok_and, passing, total = ema921_pass(_get(False), ind, True)
    assert (ok_and, passing, total) == (False, 2, 3), "legacy AND must refuse unless every TF agrees"


def test_vector_quorum_vs_and_differ():
    a = ema921_pass_vec(_npz(), 4, True, _Cfg(True), _safe)
    b = ema921_pass_vec(_npz(), 4, True, _Cfg(False), _safe)
    assert a.tolist() == [True, True, True, True]
    assert b.tolist() == [True, False, False, False], "legacy AND passes only where all TFs agree"
    assert a.tolist() != b.tolist(), "the switch must change the mask or it is not wired"


def test_short_side_inverts():
    cfg = _Cfg(False)
    m = ema921_pass_vec(_npz(), 4, False, cfg, _safe)
    # SHORT ok iff value != 1 -> bar 2 (0,0,1) has 15m/1h ok but 4h==1 fails the AND
    assert m.tolist() == [False, False, False, False]


def test_no_data_passes_in_both_modes():
    for cum in (True, False):
        assert ema921_pass_vec({}, 4, True, _Cfg(cum), _safe) is None
        assert ema921_pass(_get(cum), {}, True) == (True, 0, 0)


def test_gate_off_is_inert_in_both_modes():
    for cum in (True, False):
        d = {"EMA_9_21_FILTER_ENABLED": False, "KINDERGARTEN_CUMULATIVE_MODE": cum}
        assert ema921_pass(lambda k, de=None: d.get(k, de), {"ema_9_above_21_1h": 0.0}, True) == (True, 0, 0)


def test_crypto_vector_call_site_is_not_stocks_gated():
    """v12_quick_engine must call the shared predicate on the crypto branch, not only for tradier."""
    src = (ROOT / "v12_quick_engine.py").read_text().split("\n")
    idx = [i for i, l in enumerate(src) if "kg_entry_gate" in l and "import" in l]
    assert idx, "the crypto vector must import the shared kindergarten predicate"
    for i in idx:
        window = "\n".join(src[max(0, i - 8):i])
        assert 'MODE", "crypto")) != "tradier"' in window or "_wd_" in window, (
            "kg_entry_gate call at line %d is not on the crypto branch" % (i + 1)
        )


def test_crypto_live_calls_the_same_predicate():
    src = (ROOT / "ez_manage.py").read_text()
    assert "from vec_decisions.kg_entry_gate import ema921_pass" in src, (
        "crypto live must use the same predicate as the vector, so one switch cannot mean two things"
    )


def test_all_four_templates_carry_two_testable_options():
    sys.path.insert(0, str(ROOT / "tools"))
    from v15_switch_add import FILES

    import openpyxl

    d = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM"
    if not d.exists():
        pytest.skip("canonical templates not present on this host")
    for cs, fn in FILES.items():
        wb = openpyxl.load_workbook(str(d / fn))
        try:
            vals, bolds = [], []
            for tab in wb.sheetnames:
                ws = wb[tab]
                for r in range(3, ws.max_row + 1):
                    if str(ws.cell(r, 1).value or "").strip() == "KINDERGARTEN_CUMULATIVE_MODE":
                        vals.append(ws.cell(r, 2).value)
                        if ws.cell(r, 2).font is not None and ws.cell(r, 2).font.b:
                            bolds.append(ws.cell(r, 2).value)
            assert len(vals) >= 2, f"{cs}: needs >=2 options to be testable, found {vals}"
            assert len(bolds) == 1, f"{cs}: exactly one bold default required, found {bolds}"
        finally:
            wb.close()
