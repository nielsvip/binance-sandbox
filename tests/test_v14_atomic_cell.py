import json, tempfile, pathlib, os
import openpyxl

def test_atomic_progress_never_shrinks(tmp_path):
    """Ensure _atomic_progress_save never overwrites with fewer done entries (7h loss guard)."""
    import sys
    sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
    from tools.opt.v14_sequential_filler_v15 import _atomic_progress_save
    p = tmp_path / "prog.json"
    # initial 10 rows
    p.write_text(json.dumps({"done": {f"k{i}": {"delta": 1} for i in range(10)}, "cumulative_gain": 10}))
    # try to shrink to 3 rows — should be refused, keep 10
    _atomic_progress_save(p, {"done": {f"k{i}": {"delta": 1} for i in range(3)}, "cumulative_gain": 3})
    data = json.loads(p.read_text())
    assert len(data["done"]) == 10, f"shrink guard failed {len(data['done'])} != 10"
    # growing to 12 should succeed
    _atomic_progress_save(p, {"done": {f"k{i}": {"delta": 1} for i in range(12)}, "cumulative_gain": 12})
    data = json.loads(p.read_text())
    assert len(data["done"]) == 12

def test_json_xls_repair(tmp_path):
    """If JSON has cell but XLS Results missing, repair rewrites XLS from JSON (never lose cell)."""
    import sys
    sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
    # create XLS with empty Results
    wb_path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ENTRY_REVERSAL_BOUNCE"
    ws["A1"] = "switch"; ws["B1"] = "value"
    ws2 = wb.create_sheet("Results_Deltas")
    ws2.append(["key","default","override","is_non_default","delta_gain_vs_bh","delta_sharpe","delta_trades","variant_gain"])
    wb.save(str(wb_path))
    # JSON has one done cell
    prog = {"done": {"ENTRY_REVERSAL_BOUNCE!4:WT_15M_BOUNCE_OPEN_ENABLED=True": {"delta": 1.5, "vec": {"gain_pct": 2.0, "trades": 10, "tim_pct": 50, "max_dd_pct": 5}}}}
    # simulate repair logic (same as in filler)
    wb_repair = openpyxl.load_workbook(str(wb_path), data_only=False)
    target = "Results_Deltas"
    rws = wb_repair[target]
    existing = set(str(rws.cell(row=rr, column=1).value or "").strip() for rr in range(2, rws.max_row+1))
    assert "WT_15M_BOUNCE_OPEN_ENABLED=True" not in existing
    # run repair
    for cell_key, info in prog["done"].items():
        results_key = cell_key.split(":",1)[-1]
        if results_key not in existing and info.get("vec"):
            found = rws.max_row+1
            rws.cell(row=found, column=1).value = results_key
            rws.cell(row=found, column=8).value = float(info["vec"]["gain_pct"])
            rws.cell(row=found, column=5).value = float(info["delta"])
    wb_repair.save(str(wb_path))
    wb2 = openpyxl.load_workbook(str(wb_path), data_only=False)
    assert wb2["Results_Deltas"].cell(2,1).value == "WT_15M_BOUNCE_OPEN_ENABLED=True"
    assert wb2["Results_Deltas"].cell(2,5).value == 1.5
