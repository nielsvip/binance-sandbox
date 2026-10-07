"""Row/yellow/NPZ guards: RULE#1-4 (USER 2026-10-03). Yellow = real delta only."""

import math
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from tools.v15_row_guards import incomplete_rows, invalid_settled, is_real_number, mark_unevaluated, npz_changed, npz_identity_for_symside, row_needs_recalc, short_npz_id, yellows_complete, zero_audit_line


class TestYellowCompleteness(unittest.TestCase):
    def test_missing_reported(self):
        self.assertEqual(yellows_complete(["a", "b", "c"], {"a", "c"}), ["b"])

    def test_empty_relevant_complete(self):
        self.assertEqual(yellows_complete([], set()), [])

    def test_none_safe(self):
        self.assertEqual(yellows_complete(None, None), [])


class TestIncompleteRows(unittest.TestCase):
    def test_flags_only_explicit_false(self):
        done = {"a": {"complete": True}, "b": {"complete": False}, "c": {"delta": 1.0}, "d": None}
        self.assertEqual(incomplete_rows(done), ["b"])
        self.assertEqual(incomplete_rows(None), [])
        self.assertEqual(incomplete_rows({}), [])


class TestRealNumbers(unittest.TestCase):
    def test_accepts_floats_incl_zero(self):
        self.assertTrue(is_real_number(0.0))
        self.assertTrue(is_real_number(-3.25))

    def test_rejects_placeholders(self):
        self.assertFalse(is_real_number(None))
        self.assertFalse(is_real_number("0.0"))
        self.assertFalse(is_real_number(True))
        self.assertFalse(is_real_number(float("nan")))
        self.assertFalse(is_real_number(float("inf")))


class TestInvalidVerdicts(unittest.TestCase):
    def test_gate_verdicts_settled(self):
        self.assertTrue(invalid_settled("trades 4 < 10 (sanitized floor for 30d)"))
        self.assertTrue(invalid_settled("TIM 96.4% >80% (vomit)"))
        self.assertTrue(invalid_settled("DD 33.2% >30% (vomit)"))

    def test_transients_unsettled(self):
        self.assertFalse(invalid_settled("prepare failed (no npz)"))
        self.assertFalse(invalid_settled("v12_pilot stall timeout 90s — woke to next cell"))
        self.assertFalse(invalid_settled("YELLOW_TIMEOUT 10.0s"))
        self.assertFalse(invalid_settled("eval error boom"))
        self.assertFalse(invalid_settled(None))
        self.assertFalse(invalid_settled(""))


class TestNpzIdentity(unittest.TestCase):
    def test_identity_and_change(self):
        try:
            import numpy as _np
        except Exception:
            self.skipTest("numpy unavailable")
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "FOOUSDT.npz"
            _np.savez(p, timestamps=_np.array([1.0, 2.0, 3.0]), close=_np.array([1.0, 2.0, 3.0]))
            a = npz_identity_for_symside("FOOUSDT_LONG", indicators_dir=td)
            self.assertIsNotNone(a)
            self.assertEqual(a["n"], 3)
            self.assertEqual(a["ts_last"], 3.0)
            self.assertFalse(npz_changed(None, a))
            self.assertFalse(npz_changed(a, dict(a)))
            b = dict(a)
            b["ts_last"] = 4.0
            self.assertTrue(npz_changed(a, b))
            self.assertIn("ts", short_npz_id(a))

    def test_missing_npz_none(self):
        with tempfile.TemporaryDirectory() as td:
            self.assertIsNone(npz_identity_for_symside("NOPEUSDT_LONG", indicators_dir=td))


class TestWriters(unittest.TestCase):
    def test_mark_unevaluated_blank_red(self):
        try:
            import openpyxl
        except Exception:
            self.skipTest("openpyxl unavailable")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=3, column=12).value = 0.0
        self.assertTrue(mark_unevaluated(ws, 3, 12, "timeout-test"))
        c = ws.cell(row=3, column=12)
        self.assertIsNone(c.value)
        self.assertEqual(c.fill.start_color.rgb, "00FF0000")
        self.assertFalse(mark_unevaluated(None, 3, 12))

    def test_zero_audit_line(self):
        s = zero_audit_line("SH", 3, "SW", "F", "F=1", 1.5, 1.5, 44, "ts1_n2")
        self.assertIn("ZERO-DELTA-AUDIT", s)
        self.assertIn("SH!3", s)


class TestUnsettledNoneHealer(unittest.TestCase):
    def test_prefix_identity_row_drops(self):
        rec = {"delta": None, "reason": "", "is_running": True, "complete": None, "yellows": {}}
        self.assertTrue(str(row_needs_recalc(rec)).startswith("unsettled-none"))

    def test_crash_row_drops(self):
        rec = {"delta": None, "reason": "", "naked_delta": None}
        self.assertTrue(str(row_needs_recalc(rec)).startswith("unsettled-none"))

    def test_zero_trades_stands(self):
        rec = {"delta": None, "reason": "ZERO_TRADES", "complete": True}
        self.assertIsNone(row_needs_recalc(rec))

    def test_structural_stands(self):
        rec = {"delta": None, "reason": "NOT_WIRED_VEC: no reachable vectorized read (v15_zero_audit)"}
        self.assertIsNone(row_needs_recalc(rec))

    def test_engine_rejection_stands(self):
        rec = {"delta": None, "reason": "v12 prepared could not convert string to float: 'OFF'"}
        self.assertIsNone(row_needs_recalc(rec))

    def test_identity_zero_stands(self):
        rec = {"delta": 0.0, "reason": "RUNNING_IDENTITY: candidate == running set", "complete": True, "policy": {"ps": 0, "tl": 1, "uw": None, "pilot": "x"}}
        self.assertIsNone(row_needs_recalc(rec))

    def test_evaluated_row_stands(self):
        rec = {"delta": 1.5, "reason": "", "complete": True, "policy": {"ps": 0, "tl": 1, "uw": None, "pilot": "x"}}
        self.assertIsNone(row_needs_recalc(rec))


if __name__ == "__main__":
    unittest.main()
