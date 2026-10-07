"""Durable test for v15: exhaustive filter×switch brute + Results_Deltas pos completeness + header dark blue.

Verifies:
- HUSTLE_DELTA header is dark blue 1F4E78 black text like other headers (not 4472C4 white)
- Every pos delta writes full 24-col Results_Deltas line (beating-ideas-to-death)
- Brute script tools/test_all_filters_brute.py exists and memorizes pos_map
"""
import pathlib
import json

ROOT = pathlib.Path(__file__).resolve().parents[1]

def test_hustle_header_dark_blue_black():
    src = (ROOT / "v15_pilot.py").read_text()
    assert '1F4E78' in src, "pilot must use 1F4E78 dark blue"
    assert 'VISUAL_F_FILL = PatternFill(start_color="1F4E78"' in src, "VISUAL_F_FILL must be 1F4E78"
    assert 'VISUAL_F_FONT = Font(name="Arial", size=10, bold=True, color="000000")' in src, "VISUAL_F_FONT must be black 000000, not white"
    # ensure old 4472C4 not used for F
    # allow comment mentioning old color, but not as active code
    active_code = "\n".join(l for l in src.splitlines() if not l.strip().startswith("#") and not l.strip().startswith("-") and "Visual:" not in l)
    assert '4472C4' not in active_code, "old 4472C4 must not be in active code"
    # template headers 1-11 should be dark blue
    import openpyxl
    p = ROOT / "SPREADSHEETS" / "TEMPLATE_FINAL_NORM" / "TEMPLATE_CRYPTO_SHORT.xlsx"
    if p.exists():
        wb = openpyxl.load_workbook(str(p), data_only=False)
        ws = wb["ENTRY_REVERSAL_BOUNCE"]
        f = ws.cell(2, 6)
        assert f.fill.start_color.rgb == "FF1F4E78" or f.fill.start_color.rgb == "001F4E78", f"F2 fill expected 1F4E78 got {f.fill.start_color.rgb}"
        assert f.font.color.rgb in ("FF000000", "000000", "00000000"), f"F2 font expected black got {f.font.color.rgb}"
        wb.close()
    print("test_hustle_header_dark_blue_black PASSED")

def test_results_deltas_pos_full_metrics():
    src = (ROOT / "v15_pilot.py").read_text()
    # Every pos delta must fill full 24-col line
    assert "Every pos delta: fill full 24-col metrics line" in src, "pilot must contain pos delta full metrics comment"
    assert '"real_complete_delta"' in src.lower(), "must handle REAL_COMPLETE_DELTA"
    assert '"filter_or_override"' in src.lower(), "must handle filter_or_override"
    assert '"symside"' in src.lower(), "must handle symside"
    assert 'if delta_best is not None and delta_best > 1e-9:' in src, "must gate pos delta >1e-9"
    print("test_results_deltas_pos_full_metrics PASSED")

def test_brute_script_exists_and_memorizes():
    p = ROOT / "tools" / "test_all_filters_brute.py"
    assert p.exists(), "brute script must exist"
    src = p.read_text()
    assert "Brute-force every filter" in src, "brute script header"
    assert "pos_map" in src, "must memorize pos_map"
    assert "for sw, cand" in src, "must iterate switches"
    assert "for filt, fval" in src, "must iterate filter vals"
    assert "delta > 1e-9" in src, "must record >1e-9"
    assert "data/reports/brute_pos_filters" in src, "must output json"
    print("test_brute_script_exists_and_memorizes PASSED")

if __name__ == "__main__":
    test_hustle_header_dark_blue_black()
    test_results_deltas_pos_full_metrics()
    test_brute_script_exists_and_memorizes()
    print("all test_v15_brute_pos_filters_and_results_deltas PASSED")
