"""Parity: config vs v12_quick_engine (vector) vs TEMPLATE_CRYPTO_LONG/SHORT bold defaults.

Prevents parity disasters: START/MIN and new HTF bottom-exit/churn flags must be identical
across live config, vector, and sheet (bold).
"""
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl

def _template_values(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["GLOBAL_RISK_GATES"]
    vals = {}
    bolds = {}
    for r in range(1, ws.max_row+1):
        k = ws.cell(r,1).value
        v = ws.cell(r,2).value
        if k:
            vals[k] = str(v) if v is not None else None
            bolds[k] = bool(ws.cell(r,2).font.bold)
    return vals, bolds

def test_parity_crypto_defaults_config_vec_template():
    import config
    import v12_quick_engine as V
    flags = {
        "BOTTOM_EXIT_HTF_WT_VETO_ENABLED": "True",
        "HTF_WT_CHURN_REENTRY_ENABLED": "True",
        "HTF_WT_CHURN_REENTRY_MAX_AGE_MIN": "120.0",
        "START_POSITION_SIZE": "28.0",
        "MIN_POSITION_SIZE": "1.0",
    }
    for k, expected_str in flags.items():
        cv = getattr(config.Config, k)
        vv = getattr(V.QuickConfig, k)
        # config vs vec identical
        assert cv == vv, f"{k} config {cv!r} != vec {vv!r}"
        # string form matches expected
        assert str(cv) == expected_str, f"{k} config {cv!r} != expected {expected_str}"
        assert str(vv) == expected_str, f"{k} vec {vv!r} != expected {expected_str}"

    for p in ["SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_LONG.xlsx", "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_SHORT.xlsx"]:
        vals, bolds = _template_values(ROOT / p)
        for k, expected_str in flags.items():
            assert k in vals, f"{p} missing {k} in GLOBAL_RISK_GATES"
            assert vals[k] == expected_str, f"{p} {k}={vals[k]!r} != expected {expected_str}"
            assert bolds[k] is True, f"{p} {k} not bold (parity requires bold default)"

def test_no_tradier_contamination():
    # Crypto templates must not carry TRADIER 500 value; vec TRC_START is separate
    import v12_quick_engine as V
    assert V.QuickConfig.START_POSITION_SIZE == 28.0
    assert V.QuickConfig.TRC_START_POSITION_SIZE == 500.0
    vals_long, _ = _template_values(ROOT / "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_CRYPTO_LONG.xlsx")
    assert vals_long["START_POSITION_SIZE"] == "28.0"

def test_parity_stocks_defaults_config_vec_template():
    import config_tradier
    import v12_quick_engine as V
    # HTF flags are shared crypto/stocks — must match vec
    htf_flags = {
        "BOTTOM_EXIT_HTF_WT_VETO_ENABLED": "True",
        "HTF_WT_CHURN_REENTRY_ENABLED": "True",
        "HTF_WT_CHURN_REENTRY_MAX_AGE_MIN": "120.0",
    }
    for k, expected_str in htf_flags.items():
        cv = getattr(config_tradier.TradierConfig, k)
        vv = getattr(V.QuickConfig, k)
        assert cv == vv, f"{k} tradier config {cv!r} != vec {vv!r}"
        assert str(cv) == expected_str, f"{k} tradier config {cv!r} != expected {expected_str}"
    # Sizing is mode-split: crypto 28/1 vs tradier 500/100. Vec primary START 28 is crypto, TRC_START 500 is tradier
    assert str(getattr(config_tradier.TradierConfig, "START_POSITION_SIZE")) == "500.0"
    assert str(getattr(V.QuickConfig, "TRC_START_POSITION_SIZE")) == "500.0"
    assert str(getattr(config_tradier.TradierConfig, "MIN_POSITION_SIZE")) == "100"
    # template vs tradier config for sizing + HTF (all 5)
    flags = {
        "BOTTOM_EXIT_HTF_WT_VETO_ENABLED": "True",
        "HTF_WT_CHURN_REENTRY_ENABLED": "True",
        "HTF_WT_CHURN_REENTRY_MAX_AGE_MIN": "120.0",
        "START_POSITION_SIZE": "500.0",
        "MIN_POSITION_SIZE": "100",
    }
    for k, expected_str in flags.items():
        cv = getattr(config_tradier.TradierConfig, k)
        assert str(cv) == expected_str, f"{k} tradier config {cv!r} != expected {expected_str}"
    for p in ["SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_LONG.xlsx", "SPREADSHEETS/TEMPLATE_FINAL_NORM/TEMPLATE_STOCKS_SHORT.xlsx"]:
        vals, bolds = _template_values(ROOT / p)
        for k, expected_str in flags.items():
            assert k in vals, f"{p} missing {k} in GLOBAL_RISK_GATES"
            assert vals[k] == expected_str, f"{p} {k}={vals[k]!r} != expected {expected_str}"
            assert bolds[k] is True, f"{p} {k} not bold (parity requires bold default for stocks)"
