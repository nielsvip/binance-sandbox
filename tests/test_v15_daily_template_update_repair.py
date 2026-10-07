"""Writer self-healing: winner-less white groups with 0/2+ defaults get repaired (2026-10-07).

Chain step 2 failed on 96 pre-existing default violations (STDEV_*/GOLDEN_RULE_*/VOL_SPIKE_*...):
switches with no promotable winner kept their broken YES/bold state and the gate refused every
template forever. The writer now heals them (single-bold completion, else venue config default,
else first row) and reports repaired_defaults outside ledger/promoted_keys.
"""
import pathlib
import sys
import types

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import v15_daily_template_update as W  # noqa: E402

TAB = "GLOBAL_RISK_GATES"


def _wb(groups):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TAB
    hdr = ["Switch", "Option Value", "is_default", "AVG_DELTA", "POS_SYM", "VECTOR_DELTA"]
    for c, h in enumerate(hdr, 1):
        ws.cell(row=2, column=c).value = h
    r = 3
    for _a, opts in groups:
        for val, isd, bold in opts:
            ws.cell(row=r, column=1).value = _a
            cell = ws.cell(row=r, column=2)
            cell.value = val
            if bold:
                cell.font = openpyxl.styles.Font(bold=True)
            if isd:
                ws.cell(row=r, column=3).value = isd
            r += 1
    return wb


def _run(ws, cs="CRYPTO_LONG", allow_repair=True):
    rep = {"rows_with_stats": 0, "promoted_switch": [], "promoted_filter": [], "ledger": {}}
    W.apply_tab(ws, {}, {}, {}, True, rep, {}, {}, cs, allow_repair)
    return rep


def _state(ws):
    out = {}
    for r in range(3, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a in (None, ""):
            continue
        b = ws.cell(row=r, column=2)
        out.setdefault(a, []).append((b.value, ws.cell(row=r, column=3).value, bool(b.font and b.font.b)))
    return {a: sorted(v, key=lambda t: str(t[0])) for a, v in out.items()}


def _stub_pilot(monkeypatch):
    stub = types.SimpleNamespace(_parse_opt_value=lambda v, d: v, _same_default=lambda a, b: str(a).strip().lower() == str(b).strip().lower())
    monkeypatch.setattr(W, "_pilot_mod", lambda: stub)
    monkeypatch.setitem(W._VENUE_CACHE, "CRYPTO_LONG", {"SW_NODEFAULT": False, "SW_MULTIBOLD": True})


def test_repair_no_default_uses_config(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_NODEFAULT", [(False, None, False), (True, None, False)])])
    rep = _run(wb[TAB])
    st = _state(wb[TAB])["SW_NODEFAULT"]
    assert [(v, i, b) for v, i, b in st] == [(False, "YES", True), (True, "NO", False)]
    assert rep["repaired_defaults"] == [[TAB, "SW_NODEFAULT", "config-default", False]]
    assert rep["ledger"] == {} and rep["promoted_switch"] == []


def test_repair_multibold_uses_config(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_MULTIBOLD", [(False, None, True), (True, None, True)])])
    rep = _run(wb[TAB])
    st = _state(wb[TAB])["SW_MULTIBOLD"]
    assert [(v, i, b) for v, i, b in st] == [(False, "NO", False), (True, "YES", True)]
    assert rep["repaired_defaults"][0][:3] == [TAB, "SW_MULTIBOLD", "config-default"]


def test_repair_single_bold_completed_no_config_needed(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_HALF", [(0.05, None, False), (0.1, None, True)])])
    rep = _run(wb[TAB])
    st = _state(wb[TAB])["SW_HALF"]
    assert [(v, i, b) for v, i, b in st] == [(0.05, "NO", False), (0.1, "YES", True)]
    assert rep["repaired_defaults"][0][:3] == [TAB, "SW_HALF", "single-bold-completed"]


def test_repair_unknown_switch_first_row(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_UNKNOWN_XYZ", [("aa", None, False), ("bb", None, False)])])
    rep = _run(wb[TAB])
    st = _state(wb[TAB])["SW_UNKNOWN_XYZ"]
    assert [(v, i, b) for v, i, b in st] == [("aa", "YES", True), ("bb", "NO", False)]
    assert rep["repaired_defaults"][0][:3] == [TAB, "SW_UNKNOWN_XYZ", "first-row-fallback"]


def test_valid_group_untouched(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_OK", [(False, "YES", True), (True, "NO", False)])])
    rep = _run(wb[TAB])
    st = _state(wb[TAB])["SW_OK"]
    assert [(v, i, b) for v, i, b in st] == [(False, "YES", True), (True, "NO", False)]
    assert rep.get("repaired_defaults", []) == []


def test_no_promote_disables_repair(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_NODEFAULT", [(False, None, False), (True, None, False)])])
    rep = _run(wb[TAB], allow_repair=False)
    st = _state(wb[TAB])["SW_NODEFAULT"]
    assert all(i is None and not b for _, i, b in st)
    assert rep.get("repaired_defaults", []) == []


def test_repair_orange_rows_not_skipped(monkeypatch):
    _stub_pilot(monkeypatch)
    wb = _wb([("SW_NODEFAULT", [(False, None, False), (True, None, False)])])
    ws = wb[TAB]
    orange = openpyxl.styles.PatternFill(start_color="FFE699", fill_type="solid")
    ws.cell(row=3, column=1).fill = orange
    ws.cell(row=4, column=1).fill = orange
    rep = _run(ws)
    st = _state(ws)["SW_NODEFAULT"]
    assert [(v, i, b) for v, i, b in st] == [(False, "YES", True), (True, "NO", False)]
    assert rep["repaired_defaults"][0][:3] == [TAB, "SW_NODEFAULT", "config-default"]


def test_repair_real_template_rows_with_real_config():
    import v15_daily_template_update as W2

    W2._VENUE_CACHE.clear()
    wb = openpyxl.load_workbook(str(ROOT / "SPREADSHEETS" / "TEMPLATE_CRYPTO_LONG.xlsx"))
    ws = wb["ENTRY_REVERSAL_BOUNCE"]
    rep = {"rows_with_stats": 0, "promoted_switch": [], "promoted_filter": [], "ledger": {}}
    W2.apply_tab(ws, {}, {}, {}, True, rep, {}, {}, "CRYPTO_LONG", True)
    got = {a: how for _t, a, how, _v in rep.get("repaired_defaults", [])}
    assert "STDEV_BOUNCE_ENABLED" in got
    groups = W2.groups_by_name(ws)
    for a in ("STDEV_BOUNCE_ENABLED", "STDEV_BOUNCE_PCTB_LONG", "BB_SQUEEZE_COOLDOWN"):
        rows = groups[a]
        c_isd = W2.col_of(ws, "is_default")
        yes = [r for r in rows if str(ws.cell(row=r, column=c_isd).value or "").strip().upper() == "YES"]
        bold = [r for r in rows if ws.cell(row=r, column=2).font is not None and ws.cell(row=r, column=2).font.b]
        assert len(yes) == 1 and yes == bold, a
    W2._VENUE_CACHE.clear()
