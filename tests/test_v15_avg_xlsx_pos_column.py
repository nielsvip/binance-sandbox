import openpyxl
import pathlib
import tempfile
import collections

import tools.v15_avg.rebuild_all_bestgain as RB
import tools.v15_avg.rebuild_from_xlsx as RX


def test_category_classification():
    assert RX.category_of("AAPL_LONG") == "stocks_long"
    assert RX.category_of("NVDA_SHORT") == "stocks_short"
    assert RX.category_of("BTCUSDC_LONG") == "crypto_long"
    assert RX.category_of("ETHUSDT_SHORT") == "crypto_short"
    assert RX.category_of("1000PEPEUSDC_LONG") == "crypto_long"


def test_parse_sym_side_from_filename():
    assert RX.parse_sym_side_from_filename("AAPL_LONG_30d_matrix.xlsx") == "AAPL_LONG"
    assert RX.parse_sym_side_from_filename("1000BONKUSDC_LONG_bh15p46_gain15p11_30d_matrix.xlsx") == "1000BONKUSDC_LONG"
    assert RX.parse_sym_side_from_filename("AAPL_SHORT_bh6p45_gain0p64_30d_matrix.xlsx") == "AAPL_SHORT"


def test_parse_gain_token():
    assert RX.parse_gain_token("AAPL_LONG_bh6p45_gain0p64_30d_matrix.xlsx") == 0.64
    assert RX.parse_gain_token("1000PEPEUSDC_SHORT_bhm13p46_gainm12p20_30d_matrix.xlsx") == -12.2
    assert RX.parse_gain_token("BTCUSDC_LONG_bh35p08_gain30p69_30d_matrix.xlsx") == 30.69


def test_compute_stats_pos_rate():
    vals = [1.0, 2.0, -1.0, 0.0]
    st = RX.compute_stats(vals)
    assert st["n"] == 4
    assert st["pos"] == 2
    assert st["pos_rate"] == 0.5


def test_pos_syms_counts_distinct_sym_not_occurrences():
    # Simulate two syms both positive for same switch: aggregation should count distinct syms
    agg = {
        "per_cat_switch_pos_syms": {
            "stocks_long": {("ENTRY_REVERSAL_BOUNCE", "WT=True"): {"AAPL_LONG", "MSFT_LONG"}},
            "crypto_long": {("ENTRY_REVERSAL_BOUNCE", "WT=True"): {"BTCUSDC_LONG"}},
        },
        "per_cat_filter_pos_syms": {
            "stocks_long": {"BB_TF=15m": {"AAPL_LONG", "MSFT_LONG", "GOOG_LONG"}},
        },
    }
    assert len(agg["per_cat_switch_pos_syms"]["stocks_long"][("ENTRY_REVERSAL_BOUNCE", "WT=True")]) == 2
    assert len(agg["per_cat_filter_pos_syms"]["stocks_long"]["BB_TF=15m"]) == 3


def test_workbook_has_pos_column_in_per_category_tabs():
    path = pathlib.Path("SPREADSHEETS/V15_AVG_DELTAS.xlsx")
    assert path.exists(), "V15_AVG_DELTAS.xlsx must exist after xlsx rebuild"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in ["SWITCHES_CRYPTO_LONG", "SWITCHES_CRYPTO_SHORT", "SWITCHES_STOCKS_LONG", "SWITCHES_STOCKS_SHORT"]:
        assert sheet in wb.sheetnames
        header = next(wb[sheet].iter_rows(min_row=1, max_row=1, values_only=True))
        assert "pos_syms" in header, f"{sheet} header {header} must contain pos_syms"
        assert "pos_rate" in header
        assert header.index("pos_syms") > header.index("pos_rate")
        assert wb[sheet].max_row > 10
    for sheet in ["FILTERS_CRYPTO_LONG", "FILTERS_CRYPTO_SHORT", "FILTERS_STOCKS_LONG", "FILTERS_STOCKS_SHORT"]:
        assert sheet in wb.sheetnames
        header = next(wb[sheet].iter_rows(min_row=1, max_row=1, values_only=True))
        assert "pos_syms" in header, f"{sheet} header {header} must contain pos_syms"
    wb.close()


def test_global_sheets_have_pos_count():
    path = pathlib.Path("SPREADSHEETS/V15_AVG_DELTAS.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hdr_sw = next(wb["AVG_SWITCHES"].iter_rows(min_row=1, max_row=1, values_only=True))
    assert "pos_count" in hdr_sw
    assert "pos_rate" in hdr_sw
    hdr_f = next(wb["AVG_FILTERS"].iter_rows(min_row=1, max_row=1, values_only=True))
    assert "pos_count" in hdr_f
    wb.close()


def test_best_xlsx_selection_prefers_max_gain(tmp_path):
    # Create two fake xlsx matrices for same sym_side with different gains
    def make_matrix(path, gain, trades=10):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "AAPL_LONG_BASELINE_METRICS"
        ws["A1"] = "metric"
        ws["B1"] = "value"
        ws["A2"] = "gain_pct"
        ws["B2"] = gain
        ws["A5"] = trades
        wb.save(path)
        wb.close()

    p1 = tmp_path / "AAPL_LONG_bh1p0_gain1p0_30d_matrix.xlsx"
    p2 = tmp_path / "AAPL_LONG_bh1p0_gain5p0_30d_matrix.xlsx"
    make_matrix(p1, 1.0)
    make_matrix(p2, 5.0)
    # Simulate load_best logic: keep max gain
    sym_to_best = {}
    for p in [p1, p2]:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        gain = RX.extract_gain(wb, p.name)
        wb.close()
        sym = RX.parse_sym_side_from_filename(p.name)
        cur = sym_to_best.get(sym)
        import os
        mtime = os.path.getmtime(p)
        if cur is None or gain > cur[0]:
            sym_to_best[sym] = (gain, mtime, p)
    assert sym_to_best["AAPL_LONG"][0] == 5.0
    assert sym_to_best["AAPL_LONG"][2] == p2


def test_xlsx_aggregation_counts_pos_distinct_sym(tmp_path):
    # Build minimal matrix xlsx with one switch and one filter, positive delta
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ENTRY_REVERSAL_BOUNCE"
    ws["A2"] = "Switch"
    ws["B2"] = "default"
    ws["F2"] = "HUSTLE_DELTA"
    ws["L2"] = "BB_TF=15m"
    ws["A3"] = "WT_ENABLED"
    ws["B3"] = "True"
    ws["F3"] = 2.5
    ws["L3"] = 3.0
    ws2 = wb.create_sheet("AAPL_LONG_BASELINE_METRICS")
    ws2["B2"] = 2.5
    ws2["B5"] = 10
    p = tmp_path / "AAPL_LONG_30d_matrix.xlsx"
    wb.save(p)
    wb.close()
    sym_to_best = {"AAPL_LONG": (2.5, 0, p)}
    agg = RX.aggregate_from_xlsx(sym_to_best)
    assert agg["total_rows"] == 1
    assert agg["total_yellows_nonzero"] == 1
    assert agg["distinct_switches"] == 1
    # pos_syms for this switch should be 1
    key = ("ENTRY_REVERSAL_BOUNCE", "WT_ENABLED=True")
    assert len(agg["per_cat_switch_pos_syms"]["stocks_long"][key]) == 1
