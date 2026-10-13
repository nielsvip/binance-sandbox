"""MAXAUG_NEVER_ZERO + 4x-headroom mandate guards (USER 2026-10-10 no-cap, 2026-10-11 4x45=180).

2026-10-13 phantom post-mortem: "0.5" bypassed the pilot ("0","0.0") guard, got
promoted on 10-12, coerced to int 0; STOCKS_SHORT 0 promoted from pre-guard
evidence although vec never reads the knob. 104 per-sym poison keys stripped
Mac+S1. These tests pin the hardened guards + the structural invariants.
"""
import json
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]


def test_maxaug_is_zero_predicate():
    from v15_pilot import _maxaug_is_zero
    assert _maxaug_is_zero("0") and _maxaug_is_zero("0.0") and _maxaug_is_zero("0.5") and _maxaug_is_zero(0) and _maxaug_is_zero(0.5)
    assert not _maxaug_is_zero("999999") and not _maxaug_is_zero(10) and not _maxaug_is_zero("OFF") and not _maxaug_is_zero("1")


def test_writer_refuses_maxaug_zero():
    import importlib.util
    # NOTE: never sys.path.insert tools/ — tools/cat_side_defaults.py is a Sep-28 fossil that shadows the real shim and silently breaks the chain.
    import sys
    spec = importlib.util.spec_from_file_location("v15_avg_delta_apply", str(ROOT / "tools" / "v15_avg_delta_apply.py"))
    W = importlib.util.module_from_spec(spec)
    _sp = list(sys.path)
    try:
        spec.loader.exec_module(W)  # writer inserts tools/ into sys.path at import; restored below
    finally:
        sys.path[:] = _sp
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AUGMENT_RISK_SIZING"
    ws.cell(row=1, column=1).value = "Switch"
    ws.cell(row=2, column=1).value = "Switch"
    ws.cell(row=2, column=2).value = "Option Value"
    ws.cell(row=2, column=3).value = "is_default"
    ws.cell(row=2, column=4).value = "AVG_DELTA"
    ws.cell(row=2, column=5).value = "POS_SYM"
    ws.cell(row=3, column=1).value = "MAX_AUGMENTS_PER_POSITION"
    ws.cell(row=3, column=2).value = 999999
    ws.cell(row=3, column=3).value = "YES"
    ws.cell(row=4, column=1).value = "MAX_AUGMENTS_PER_POSITION"
    ws.cell(row=4, column=2).value = "0.5"
    ws.cell(row=4, column=3).value = "NO"
    stats = {("AUGMENT_RISK_SIZING", "MAX_AUGMENTS_PER_POSITION=999999"): (0.01, 1, 1), ("AUGMENT_RISK_SIZING", "MAX_AUGMENTS_PER_POSITION=0.5"): (1.09, 8, 11)}
    ledger, rep = {}, {"rows_with_stats": 0, "promoted_switch": []}
    W.apply_tab(ws, stats, ledger, rep)
    assert "MAX_AUGMENTS_PER_POSITION" not in ledger, ledger
    assert ws.cell(row=3, column=3).value == "YES"
    assert rep.get("refused_mandate"), rep


def test_golive_strips_maxaug_zero_source():
    src = (ROOT / "tools" / "v15_persym_golive.py").read_text()
    assert "mandate_stripped" in src and "MAX_AUGMENTS_PER_POSITION" in src


def test_no_maxaug_poison_in_per_sym_store():
    import per_sym_store as pss
    con = pss._connect()
    rows = con.execute("SELECT sym_side FROM per_sym_active").fetchall()
    poison = []
    for (ss,) in rows:
        try:
            ov = pss.get_overrides(ss) or {}
        except Exception:
            continue
        if "MAX_AUGMENTS_PER_POSITION" in ov:
            try:
                v = float(str(ov["MAX_AUGMENTS_PER_POSITION"]))
            except Exception:
                continue
            if v < 1:
                poison.append(ss)
    con.close()
    assert not poison, f"per-sym MAXAUG<1 poison live: {poison[:10]}"


def test_store_honors_maxaug_never_zero():
    store = json.loads((ROOT / "data" / "per_sym_settings.json").read_text())
    for cs in ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"):
        v = (store.get(cs) or {}).get("MAX_AUGMENTS_PER_POSITION")
        assert v is not None and float(v) >= 1, (cs, v)


def test_store_max_order_keeps_4x_headroom():
    store = json.loads((ROOT / "data" / "per_sym_settings.json").read_text())
    for cs in ("CRYPTO_LONG", "CRYPTO_SHORT", "STOCKS_LONG", "STOCKS_SHORT"):
        sps = float((store.get(cs) or {}).get("START_POSITION_SIZE", 0) or 0)
        mov = float((store.get(cs) or {}).get("MAX_ORDER_VALUE", 0) or 0)
        if sps > 0:
            assert mov >= 4 * sps, (cs, mov, sps)
